"""Fixture preparation and targeted native integration evidence. Run once per folder."""
import sys,json,shutil,copy,importlib,unittest
from pathlib import Path
from unittest.mock import patch
from datetime import datetime
from decimal import Decimal
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from akz_stage6 import core as c
E=Path((ROOT/'validation/STAGE6_PART1_CURRENT.txt').read_text(encoding='utf-8'))
run_name=E.name+(sys.argv[1] if len(sys.argv)>1 else '')
F=c.FIXTURES/run_name;F.mkdir(parents=True,exist_ok=False)
def save(name,value):
    (E/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
results=[]
save('tested_product_before.json',{str(p):c.sha256(p) for folder in ('akz_stage6',) for p in (ROOT/folder).glob('*.py')} | {str(ROOT/'stage6.py'):c.sha256(ROOT/'stage6.py'),str(ROOT/'tests/test_stage6.py'):c.sha256(ROOT/'tests/test_stage6.py')})
def check(name,fn):
    fn();results.append({'name':name,'status':'PASS'});print('PASS',name,flush=True)
def yes(value):assert value
def stop(fn,needle):
    try:fn()
    except c.Stage6Error as exc:
        assert needle in str(exc),(needle,str(exc));return
    raise AssertionError('Expected STOP '+needle)
v=json.loads((E/'real_vom_read.json').read_text(encoding='utf-8'))
assessment=json.loads((E/'area_assessment.json').read_text(encoding='utf-8'))
marks=[r['mark'] for r in assessment['valid'][:3]]
catalog_path=F/'TEST_Справочник марок.xlsx'
shutil.copyfile(c.PROJECT/'Справочник марок.xlsx',catalog_path)
native_opens=[]
original_open=c.native_open
def logged_open(app,path,**kw):
    native_opens.append({'path':str(Path(path).resolve()),**kw})
    return original_open(app,path,**kw)
c.native_open=logged_open
with c.native_excel() as app:
    b=c.native_open(app,catalog_path,read_only=False)
    try:
        s=b.Worksheets(1)
        assigned=[]
        for rn in range(2,197):
            if s.Cells(rn,1).Value2 in marks:
                s.Cells(rn,3).Value2=0
                assigned.append({'mark':s.Cells(rn,1).Value2,'catalog_row':rn,'assignment':0,'mode':'TEST_ONLY'})
        assert len(assigned)==3
        b.Save()
    finally:b.Close(SaveChanges=False)
    # Fully synthetic isolated VOM for visibility/format and rollback boundary checks.
    b=app.Workbooks.Add();s=b.Worksheets(1)
    while b.Worksheets.Count>1:b.Worksheets(b.Worksheets.Count).Delete()
    s.Name='TEST VOM'
    s.Range('A1').Value2='SYNTHETIC TEST ONLY — НЕ ДАННЫЕ РЕАЛЬНОГО ПАСПОРТА'
    for cell,value in [('A2','Отправочная марка'),('B2','Наименование марки'),('C2','Количество'),('D2','Масса, кг'),('F2','Площадь окраски, м2'),('D3','1 дет.'),('E3','Всех'),('F3','1 дет.'),('G3','Всех')]:s.Range(cell).Value2=value
    for a in ('A2:A3','B2:B3','C2:C3','D2:E2','F2:G2'):s.Range(a).Merge()
    fixture_rows=[(marks[0],'TEST visible',2,3.25,6.5,1.25,2.5),
                  (marks[0],'TEST duplicate',1,3.25,3.25,1.25,1.25),
                  (marks[1],'TEST filtered',1,3.25,3.25,1.25,1.25),
                  (marks[2],'TEST manually hidden',1,3.25,3.25,1.25,1.25)]
    s.Range('A4:G7').Value2=tuple(fixture_rows)
    s.Range('A3:G7').AutoFilter(Field=2,Criteria1='<>TEST filtered')
    s.Rows(7).Hidden=True
    xlsx=F/'SYNTHETIC_visibility.xlsx';xls=F/'SYNTHETIC_visibility.xls'
    b.SaveAs(str(xlsx),FileFormat=51);b.SaveAs(str(xls),FileFormat=56)
    b.Close(SaveChanges=False)
    s=None;b=None
save('test_assignments.json',{'fixture':str(catalog_path),'sha256':c.sha256(catalog_path),'assignments':assigned,
     'real_catalog_unchanged':True,'note':'COPY ONLY, 3 TEST assignments, 192 blanks'})
check('01 explicit template pair and hashes',lambda:yes(c.templates()['layers']==1))
check('02 current DEV passport / named date mapping',lambda:yes(c.current_passport('Гараж Казахстан','9/2026')['date']=='2026-10-05'))
p=c.preview('Гараж Казахстан','9/2026',v['path'],catalog_path)
save('preview.json',{'seal':p.seal,'payload':p.payload})
check('03 real visible shipment selected explicitly',lambda:yes(len(p.payload['vom']['rows'])==104))
check('04 397 hidden real rows excluded',lambda:yes(len(p.payload['vom']['hidden_rows'])==397))
check('05 three exact TEST zero assignments / selected rows',lambda:yes([r['row'] for r in p.payload['selected']]==[40,53,54]))
check('06 VOM-only quantity mass area totals',lambda:yes(p.payload['totals']=={'quantity':Decimal(3),'mass_total':Decimal('105.26'),'area_total':Decimal('3.989')}))
check('07 192 catalog assignments remain blank',lambda:yes(sum(x['assignment'] is None for x in c.catalog(catalog_path).values())==192))
for path in (xlsx,xls):
    seen=c.visible_vom(path)
    save(path.stem+path.suffix+'_visibility.json',seen)
    check('08 '+path.suffix+' AutoFilter and manual hidden rows',lambda:yes([r['row'] for r in seen['rows']]==[4,5] and seen['hidden_rows']==[6,7]))
check('08b synthetic source cannot be issued with real passport number',lambda:stop(lambda:c.preview('Гараж Казахстан','9/2026',xlsx,catalog_path),'SYNTHETIC_SOURCE_REQUIRES_SYNTHETIC_MODE'))
synthetic_preview=c.preview('Гараж Казахстан','9/2026',xlsx,catalog_path,synthetic=True)
check('08c synthetic preview has distinct act number and keeps duplicate rows',lambda:yes(synthetic_preview.payload['act_number']=='SYNTHETIC-TEST 9/2026' and [r['row'] for r in synthetic_preview.payload['selected']]==[4,5]))
# File mutation boundaries are isolated copies, never real sources.
changed=F/'TEST_changed_catalog.xlsx';shutil.copyfile(catalog_path,changed)
q=c.preview('Гараж Казахстан','9/2026',v['path'],changed)
with changed.open('ab') as stream:stream.write(b'TEST_MUTATION')
check('09 changed input after preview STOP',lambda:stop(lambda:c.generate(q,'SHOULD_NOT_EXIST_CHANGED'),'INPUT_CHANGED_AFTER_PREVIEW'))
duplicate=F/'TEST_duplicate_catalog.xlsx';shutil.copyfile(catalog_path,duplicate)
with c.native_excel() as app:
    b=c.native_open(app,duplicate,read_only=False)
    s=b.Worksheets(1);s.Range('A197:C197').Value2=((marks[0],'TEST duplicate conflict',1),);b.Save();b.Close(SaveChanges=False)
    s=None;b=None
check('10 conflicting catalog duplicate STOP',lambda:stop(lambda:c.catalog(duplicate),'CATALOG_DUPLICATE_CONFLICT'))
# Template mutation on an isolated file; hardcoded expected SHA retained.
bad=F/'TEST_changed_template.xlsx';shutil.copyfile(c.ACT,bad)
with bad.open('ab') as stream:stream.write(b'TEST_MUTATION')
with patch.object(c,'PINS',{bad:c.PINS[c.ACT]}):
    check('11 changed source template SHA STOP',lambda:stop(c.templates,'TEMPLATE_SHA_CHANGED'))
before_pending=set(c.OUTPUT.glob('.pending-*')) if c.OUTPUT.exists() else set()
name=E.name+'_ROLLBACK_TEST'
with patch.object(c,'write_application',side_effect=c.Stage6Error('INJECTED_SECOND_FILE_FAILURE')):
    check('12 rollback if second workbook fails',lambda:stop(lambda:c.generate(p,name),'INJECTED_SECOND_FILE_FAILURE'))
check('13 no partial published pair or leaked pending dir',lambda:yes(not (c.OUTPUT/name).exists() and set(c.OUTPUT.glob('.pending-*'))==before_pending))
output=c.generate(p,run_name+'_TEST_9_2026')
save('generated.json',{'folder':str(output),'files':{str(x):c.sha256(x) for x in output.iterdir()}})
check('14 complete TEST pair atomically published',lambda:yes(len(list(output.glob('*.xlsx')))==2 and (output/'PAIR.json').is_file()))
hashes={str(x):c.sha256(x) for x in output.iterdir()}
check('15 existing output STOP',lambda:stop(lambda:c.generate(p,output.name),'OUTPUT_ALREADY_EXISTS'))
check('16 collision leaves output unchanged',lambda:yes(all(c.sha256(x)==h for x,h in hashes.items())))
from openpyxl import load_workbook
a=load_workbook(output/'ТЕСТ АКЗ №0.xlsx',data_only=True)
b=load_workbook(output/'ТЕСТ Приложение к АКЗ №0.xlsx',data_only=True)
try:
    s=a['4'];t=b.active
    check('17 pair number date ACT-ID equal',lambda:yes(s['E4'].value=='9/2026 от ' and '9/2026' in t['A1'].value and s['H4'].value.date().isoformat()=='2026-10-05' and '05.10.2026' in t['A1'].value and s.oddFooter.center.text==t.oddFooter.center.text))
    check('18 coating requirements preserved from exact template',lambda:yes('Pirocor Protect 057' in s['E18'].value and s['E21'].value=='RAL 7004' and s['E25'].value==80))
    check('19 historical measurements climate batches dates results removed',lambda:yes(all(s[x].value==c.FACTUAL_PLACEHOLDERS[x] for x in c.FACTUAL_CELLS) and 'Выдача запрещена' in s['B43'].value))
    check('20 application exact selected marks and no historical rows',lambda:yes([t.cell(r,1).value for r in range(4,7)]==marks and all(t.cell(r,col).value is None for r in range(7,21) for col in range(1,8))))
    check('21 cached native recalculated application totals',lambda:yes(abs(t['C21'].value-3)<1e-8 and abs(t['E21'].value-105.26)<1e-8 and abs(t['G21'].value-3.989)<1e-8))
    check('22 TEST marking in both outputs',lambda:yes('ТЕСТ / ЧЕРНОВИК' in s['B3'].value and 'ТЕСТ / ЧЕРНОВИК' in t['A1'].value))
    save('saved_output_checks.json',{'act':{x:s[x].value for x in ['B3','E4','H4','E5','E6','E7','B44','E18','E21','E25',*c.FACTUAL_CELLS]},
        'application':[[t.cell(r,col).value for col in range(1,8)] for r in range(1,22)]})
finally:a.close();b.close()
suite=unittest.defaultTestLoader.discover(str(ROOT/'tests'),pattern='test_stage6.py')
with (E/'unit_tests.txt').open('w',encoding='utf-8') as stream:
    unit=unittest.TextTestRunner(stream=stream,verbosity=2).run(suite)
check('23 pure boundary tests including missing/conflicting area and climate',lambda:yes(unit.wasSuccessful()))
modules=[]
for package in ('passport_app','profiles_stage4','uzk_stage5'):
    for path in (ROOT/package).glob('*.py'):
        module=package if path.stem=='__init__' else package+'.'+path.stem
        importlib.import_module(module);modules.append(module)
for module in ('main','stage2','stage3','stage4','stage5'):
    importlib.import_module(module);modules.append(module)
save('accepted_import_smoke.json',{'modules':modules,'status':'PASS'})
check('24 accepted stages 1-5 import smoke',lambda:yes(len(modules)>30))
from akz_stage6.gui import Launcher
import tkinter as tk
root=tk.Tk();root.withdraw()
try:
    gui=Launcher(root)
    assert str(gui.generate_button['state'])=='disabled'
    gui.checked=p;gui.generate_button.configure(state='normal');gui.vom.set('changed selection')
    check('25 real Tk source change invalidates preview and generation',lambda:yes(gui.checked is None and str(gui.generate_button['state'])=='disabled'))
finally:root.destroy()
baseline=json.loads((E/'baseline.json').read_text(encoding='utf-8'))
changes=[x for x,h in baseline.items() if c.sha256(x)!=h]
check('26 accepted products and all original input fingerprints unchanged',lambda:yes(not changes))
save('source_integrity_after_tests.json',{'before':baseline,'after':{x:c.sha256(x) for x in baseline},'changed':changes})
save('native_open_log.json',native_opens)
save('targeted_validation.json',{'status':'PASS','checks':results,'unit_tests':unit.testsRun,
                               'output':str(output),'fixture':str(F)})
print('ALL_CHECKS_PASS',len(results),'UNIT',unit.testsRun,'OUTPUT',output,flush=True)
