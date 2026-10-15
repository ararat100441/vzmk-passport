import sys
sys.dont_write_bytecode = True
import json, hashlib, re, csv, traceback, os
from pathlib import Path
from datetime import datetime, timezone, timedelta
from collections import Counter, defaultdict
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
DOCS = ROOT.parent
PROJECT = DOCS / 'Гараж казахстан'
PASSPORT = PROJECT / 'Паспорт'
TEMPLATES = DOCS.parent / 'Шаблоны'
OTHER_TEMPLATES = DOCS.parent / 'Другие шаблоны'
FORBIDDEN = {'взмк.xlsx', 'узк.xlsm'}

def allowed(p):
    p = Path(p).resolve()
    assert p.name.casefold() not in FORBIDDEN, str(p)
    assert p.is_relative_to(ROOT) or p.is_relative_to(PROJECT) or p.is_relative_to(TEMPLATES) or (p.parent==OTHER_TEMPLATES and 'Акт АКЗ' in p.name), str(p)
    return p

def sha(p):
    p = allowed(p)
    return hashlib.sha256(p.read_bytes()).hexdigest()

def save(name, data):
    p = (OUT / name).resolve()
    assert p.is_relative_to(OUT)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

def paths():
    items = set()
    for d in [ROOT / 'passport_app', ROOT / 'profiles_stage4', ROOT / 'uzk_stage5', ROOT / 'DEV']:
        for parent, dirs, files in os.walk(d, followlinks=False):
            dirs[:] = [n for n in dirs if n not in {'__pycache__','node_modules'} and not (Path(parent)/n).is_junction() and not (Path(parent)/n).is_symlink()]
            items.update(Path(parent)/n for n in files if not n.startswith('~$'))
    items.update(p for p in ROOT.iterdir() if p.is_file())
    items.update(p for p in PASSPORT.rglob('*') if p.is_file() and not p.name.startswith('~$'))
    items.update(p for p in PROJECT.iterdir() if p.is_file() and p.suffix.lower() in {'.xlsx', '.xls', '.txt', '.xsr'})
    items.update(p for p in TEMPLATES.iterdir() if p.is_file() and 'АКЗ' in p.name)
    items.update(p for p in OTHER_TEMPLATES.glob('1. Акт АКЗ*.xlsx'))
    return sorted(p for p in items if p.name.casefold() not in FORBIDDEN and not p.is_symlink() and (any(p.resolve().is_relative_to(r.resolve()) for r in [ROOT,PROJECT,TEMPLATES]) or p.parent==OTHER_TEMPLATES))

def manifest():
    return [{'path':str(p.resolve()), 'size':p.stat().st_size, 'sha256':sha(p)} for p in paths()]

def extract_xlsx(p, idx):
    p = allowed(p)
    b = openpyxl.load_workbook(p, data_only=False, keep_links=True)
    cache = openpyxl.load_workbook(p, data_only=True, keep_links=True)
    info = {'id':idx, 'path':str(p), 'sha256':sha(p), 'format':p.suffix[1:].upper(),
            'defined_names':{k:str(v) for k,v in b.defined_names.items()},
            'external_links':len(b._external_links), 'sheets':[]}
    lines = []
    for s in b:
        c = cache[s.title]
        cells=[]
        for row in s:
            for cell in row:
                if cell.value is not None:
                    rec={'cell':cell.coordinate,'value':cell.value,'type':cell.data_type,
                         'cached':c[cell.coordinate].value,'format':cell.number_format}
                    if cell.comment: rec['comment']=cell.comment.text
                    cells.append(rec)
        si={'name':s.title,'state':s.sheet_state,'dimension':s.calculate_dimension(),
            'print_area':str(s.print_area),'print_title_rows':s.print_title_rows,
            'merged':list(map(str,s.merged_cells.ranges)),
            'row_breaks':[x.id for x in s.row_breaks.brk], 'col_breaks':[x.id for x in s.col_breaks.brk],
            'page_setup':str(s.page_setup),'page_margins':str(s.page_margins),
            'tables':{k:v.ref for k,v in s.tables.items()} if False else {k:s.tables[k].ref for k in s.tables},
            'data_validations':str(s.data_validations),'conditional_formats':[str(x) for x in s.conditional_formatting],
            'cells':cells, 'formula_count':sum(x['type']=='f' for x in cells)}
        info['sheets'].append(si)
        lines.append('SHEET '+s.title+' '+s.calculate_dimension()+' PRINT '+str(s.print_area))
        for row in s:
            parts=[]
            for cell in row:
                if cell.value is not None:
                    v=str(cell.value).replace('\n',' / ')
                    if cell.data_type=='f': v+=' [cache='+str(c[cell.coordinate].value)+']'
                    parts.append(cell.coordinate+'='+v)
            if parts: lines.append(' | '.join(parts))
    b.close(); cache.close()
    save(f'xlsx_{idx:03}.json',info)
    (OUT/f'xlsx_{idx:03}.txt').write_text('\n'.join(lines),encoding='utf-8')
    return {k:v for k,v in info.items() if k!='sheets'} | {'sheets':[{k:v for k,v in x.items() if k!='cells'} for x in info['sheets']]}

def initial():
    save('source_manifest_before.json',manifest())
    save('discovery.json',{'project':str(PROJECT),'passport_root':str(PASSPORT),
        'project_root_files':[str(p) for p in PROJECT.iterdir() if p.is_file()],
        'passport_files':[str(p) for p in PASSPORT.rglob('*') if p.is_file()],
        'general_akz':[str(p) for p in TEMPLATES.iterdir() if 'АКЗ' in p.name],
        'index_tools':'UNAVAILABLE; direct bounded source inspection; no reindex',
        'timestamp':datetime.now(timezone(timedelta(hours=4))).isoformat()})
    sys.path.insert(0,str(ROOT))
    from passport_app.registry import read_registry, select_project
    from passport_app.config import Settings
    from passport_app.catalog import project_folder
    reg=read_registry(ROOT/'DEV'/'ВЗМК_DEV.xlsx')
    pr=select_project(reg,'Гараж Казахстан')
    settings=Settings.load(ROOT/'settings.json')
    # Do not access the production registry_path held in original settings.
    folder=project_folder(settings,pr.name)
    assert folder == PROJECT.resolve()
    rb=openpyxl.load_workbook(reg.path,data_only=True)
    sheet=rb[reg.sheet]
    headers={c.column:str(c.value) for c in sheet[1] if c.value is not None}
    save('registry_project.json',{'registry_path':str(reg.path),'sha256':reg.digest,
        'selector':'passport_app.registry.read_registry -> select_project',
        'project':pr.__dict__,'folder':str(folder),'rows':reg.projects[pr.name],
        'full_project_rows':[{'row':e['row'],'values':{headers.get(c,str(c)):sheet.cell(e['row'],c).value for c in range(1,sheet.max_column+1)}} for e in reg.projects[pr.name]]})
    rb.close()
    chosen=sorted(set([p for p in PASSPORT.rglob('*.xlsx') if ('АКЗ' in p.name or p.name.startswith('Акт') or 'Приложение № 1' in p.name or 'ВОМ' in p.name)]
        + [p for p in TEMPLATES.glob('*.xlsx') if 'АКЗ' in p.name]
        + [PROJECT/'Справочник марок.xlsx',PROJECT/'ВОМ — АКЗ.xlsx',PROJECT/'Техническое задание - СТК ВК .xlsx']))
    inv=[]
    for i,p in enumerate(chosen,1):
        print('XLSX',i,p.name,flush=True)
        inv.append(extract_xlsx(p,i))
    save('inventory_xlsx.json',inv)
    print('INITIAL_COMPLETE',len(inv),flush=True)

def verify():
    before=json.loads((OUT/'source_manifest_before.json').read_text(encoding='utf-8'))
    after=manifest(); save('source_manifest_after.json',after)
    bm={x['path']:x for x in before}; am={x['path']:x for x in after}
    changes=[{'path':p,'before':bm.get(p),'after':am.get(p)} for p in sorted(set(bm)|set(am)) if (None if p not in bm else (bm[p]['sha256'],bm[p]['size']))!=(None if p not in am else (am[p]['sha256'],am[p]['size']))]
    save('source_integrity.json',{'before_count':len(before),'after_count':len(after),'changes':changes,
         'all_unchanged':not changes,'production_files_opened_or_hashed':False})
    print('SOURCE_INTEGRITY',len(before),len(after),'CHANGES',len(changes))

def extra():
    before=json.loads((OUT/'source_manifest_before.json').read_text(encoding='utf-8'))
    extra=sorted(OTHER_TEMPLATES.glob('1. Акт АКЗ*.xlsx'))
    known={x['path'] for x in before}
    for p in extra:
        if str(p.resolve()) not in known:
            before.append({'path':str(p.resolve()),'size':p.stat().st_size,'sha256':sha(p),'added_before_first_content_read':True})
    save('source_manifest_before.json',before)
    inv=json.loads((OUT/'inventory_xlsx.json').read_text(encoding='utf-8'))
    for i,p in enumerate(extra,39):
        inv.append(extract_xlsx(p,i)); print('EXTRA',i,p.name,flush=True)
    save('inventory_xlsx.json',inv)

if __name__=='__main__':
    {'initial':initial,'verify':verify,'extra':extra}[sys.argv[1]]()
