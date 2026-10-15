"""Reproducible isolated tests. No real catalog writes, production reads or Word output."""
from pathlib import Path
import sys
ROOT = Path(__file__).resolve().parents[2]
RUN = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
from io import BytesIO
import hashlib
import json
import tempfile
from uuid import uuid4
from unittest.mock import patch
from zipfile import ZipFile
import openpyxl
import win32con
import win32file
from passport_app import catalog as cat
from passport_app import catalog_migration as migration
from passport_app.config import Settings, PassportError
from passport_app.excel import locked_file
from uzk_stage5 import core, inputs


def sha(data):
    return hashlib.sha256(data).hexdigest()


def file_sha(path):
    return sha(Path(path).read_bytes())


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf8')


def product_hashes():
    files = [*sorted((ROOT / 'passport_app').glob('*.py')),
             *sorted((ROOT / 'uzk_stage5').glob('*.py')),
             *sorted((ROOT / 'profiles_stage4').glob('*.py')),
             ROOT / 'stage3.py', ROOT / 'stage5.py',
             ROOT / 'CATALOG_SCHEMA_STAGE6_CONTRACT.txt', Path(__file__)]
    return {str(p): file_sha(p) for p in files}


baseline = json.loads((RUN / 'baseline.json').read_text(encoding='utf8'))
source = (RUN / 'source_catalog.xlsx').read_bytes()
assert sha(source) == baseline['source_sha256']
protected = dict(baseline['protected_before'])
# Old acts, templates and the selected real VOM: read-only hash inventory, no registry discovery.
project = Path(baseline['source_catalog']).parent
for p in sorted((project / 'Паспорт').rglob('*')):
    if p.is_file() and not p.name.startswith('~$') and p.name.casefold() not in ('взмк.xlsx', 'узк.xlsm'):
        protected[str(p)] = file_sha(p)
for p in sorted((ROOT / 'DEV').rglob('*.json')):
    protected[str(p)] = file_sha(p)
save_json(RUN / 'protected_before.json', protected)
before = product_hashes()
save_json(RUN / 'product_hashes_BEFORE_TESTS.json', before)
results, audit_reads, denied = [], set(), []
ATTEMPT = RUN / ('attempt_' + uuid4().hex[:8])
ATTEMPT.mkdir()
tempfile.tempdir = str(ATTEMPT)


def audit(event, args):
    if event != 'open' or not args or not isinstance(args[0], (str, bytes)):
        return
    p = Path(args[0].decode() if isinstance(args[0], bytes) else args[0]).absolute()
    if p.name.casefold() in ('взмк.xlsx', 'узк.xlsm'):
        denied.append(str(p))
        raise AssertionError('PRODUCTION_REGISTRY_ACCESS_FORBIDDEN')
    mode = args[1] if len(args) > 1 else None
    flags = args[2] if len(args) > 2 else 0
    write = isinstance(mode, str) and any(c in mode for c in 'wax+') or bool(flags & (1 | 2 | 64 | 512))
    if write and (not p.is_relative_to(RUN) or p.suffix.casefold() == '.docx'):
        raise AssertionError(f'ISOLATED_WRITE_OUTSIDE_RUN: {p}')
    if not write and p.suffix.casefold() in ('.xls', '.xlsx', '.xlsm'):
        audit_reads.add(str(p))


sys.addaudithook(audit)


def test(name, fn):
    try:
        detail = fn()
    except Exception as exc:
        results.append({'test': name, 'status': 'FAIL', 'error': repr(exc)})
        save_json(RUN / 'isolated_tests.json', results)
        raise
    results.append({'test': name, 'status': 'PASS', 'detail': detail})
    print(f'PASS {len(results):02d} {name}', flush=True)


def require(condition, message='assertion failed'):
    assert condition, message


def expect_stop(fn, message=None):
    try:
        fn()
    except (PassportError, OSError) as exc:
        if message:
            assert message in str(exc), repr(exc)
        return str(exc)
    raise AssertionError('Expected safe STOP')


def fixture(rows, legacy=True):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Марки'
    sheet.append(cat.LEGACY_HEADERS if legacy else cat.HEADERS)
    for row in rows:
        sheet.append(row)
    sheet.auto_filter.ref = f'A1:{"D" if legacy else "C"}{sheet.max_row}'
    for col, width in zip('ABCD' if legacy else 'ABC', (25, 38, 28, 18) if legacy else (25, 38, 18)):
        sheet.column_dimensions[col].width = width
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def write_fixture(name, data):
    p = ATTEMPT / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    return p


old = migration.assess(source)
new = migration.migrated_bytes(source)
new_path = write_fixture('isolated_catalog3.xlsx', new)
current = migration.assess(new)
save_json(RUN / 'old_catalog_assessment.json', old)
save_json(RUN / 'isolated_catalog_assessment.json', current)
test('01 migration 4 columns to exactly 3', lambda: require(tuple(current['headers']) == cat.HEADERS))
test('02 all 195 marks and row positions preserved', lambda: require(
    old['records'] == current['records'] == 195 and
    [(r['row'], r['mark']) for r in old['rows']] == [(r['row'], r['mark']) for r in current['rows']]))
test('03 all profiles preserved', lambda: require(
    [r['profile'] for r in old['rows']] == [r['profile'] for r in current['rows']]))
test('04 all AKZ values types styles formats preserved', lambda: require(old['rows'] == current['rows']))
number_fixture = fixture([('ZERO', 'L50X5', None, 0), ('BLANK', 'L50X5', None, None),
                          ('ONE', 'L50X5', None, 1), ('TWO', 'L50X5', None, 2)])
number_new = migration.migrated_bytes(number_fixture)
number_rows = migration.assess(number_new)['rows']
test('05 numeric zero remains numeric assignment', lambda: require(
    number_rows[0]['akz'] == 0 and number_rows[0]['akz_type'] == 'n'))
test('06 blank stays distinct from numeric zero', lambda: require(number_rows[1]['akz'] is None))


def paint_stop():
    for i, value in enumerate(('RAL 9005', 0, '=1+1', ' ')):
        data = fixture([('M', 'L50X5', value, 0)])
        p = write_fixture(f'paint_stop_{i}/catalog.xlsx', data)
        expect_stop(lambda: migration.migrate_catalog(p, sha(data)), 'Краска')
        assert p.read_bytes() == data and not (p.parent / '.stage3-backups').exists()
    mixed = fixture([('BAD AKZ', 'L50X5', None, -1), ('PAINT1', 'L50X5', 'RAL 1', None),
                     ('PAINT2', 'L50X5', 'RAL 2', None)])
    error = expect_stop(lambda: migration.migrated_bytes(mixed), 'Краска')
    assert 'RAL 1' in error and 'RAL 2' in error
    return 'text, numeric zero, formula and whitespace safely STOP before backup/temp'


test('07 nonempty paint never discarded', paint_stop)


def invalid_akz():
    for value in ('0', -1, 1.5, True, '=0', '#VALUE!'):
        data = fixture([('M', 'L50X5', None, value)])
        expect_stop(lambda: migration.migrated_bytes(data), '№ АКЗ')
    return 'text, negative, fraction, boolean, formula, error rejected without coercion'


test('08 invalid AKZ formats STOP', invalid_akz)


def backup_rollback():
    p = write_fixture('rollback/catalog.xlsx', number_fixture)
    result = migration.migrate_catalog(p, sha(number_fixture))
    assert Path(result['backup']).read_bytes() == number_fixture
    assert p.read_bytes() == number_new
    expect_stop(lambda: migration.rollback_catalog(p, result['backup'], '0' * 64,
                                                   sha(number_fixture)), 'изменён')
    assert p.read_bytes() == number_new
    restored = migration.rollback_catalog(p, result['backup'], sha(number_new), sha(number_fixture))
    assert p.read_bytes() == number_fixture
    return {'migration': result, 'rollback': restored}


test('09 backup and exact-byte guarded rollback', backup_rollback)
test('10 native Excel normal-load no repair and no Save', lambda: migration.native_validate(new_path))


def failure_safety():
    def reject(path):
        raise PassportError('injected native failure')
    for i, mode in enumerate(('native', 'atomic')):
        p = write_fixture(f'failure_{mode}/catalog.xlsx', number_fixture)
        if mode == 'native':
            expect_stop(lambda: migration.migrate_catalog(p, sha(number_fixture), validator=reject))
        else:
            with patch.object(migration, 'atomic_replace', side_effect=OSError('injected atomic failure')):
                expect_stop(lambda: migration.migrate_catalog(p, sha(number_fixture)))
        assert p.read_bytes() == number_fixture
        backups = list((p.parent / '.stage3-backups').glob('*.xlsx'))
        assert len(backups) == 1 and backups[0].read_bytes() == number_fixture
        assert not list(p.parent.glob('.stage3-schema-*'))
    return 'native and atomic failure preserve exact original and verified backup'


test('11 injected failures preserve original', failure_safety)


def lock_and_concurrency():
    p = write_fixture('locks/catalog.xlsx', number_fixture)
    expect_stop(lambda: migration.migrate_catalog(p, '0' * 64), 'изменился')
    handle = win32file.CreateFile(str(p), win32con.GENERIC_READ, 0, None, win32con.OPEN_EXISTING, 0, None)
    try:
        expect_stop(lambda: migration.migrate_catalog(p, sha(number_fixture)), 'занят')
    finally:
        handle.Close()
    assert p.read_bytes() == number_fixture and not (p.parent / '.stage3-backups').exists()
    return 'stale SHA and native Excel-equivalent lock block writes'


test('12 concurrent change and file lock STOP', lock_and_concurrency)

# Real CatalogService, explicit DEV registry and isolated project folder.
docs = ATTEMPT / 'documents'
folder = docs / 'Гараж Казахстан'
folder.mkdir(parents=True, exist_ok=True)
catalog_path = folder / cat.CATALOG
catalog_path.write_bytes(fixture([('FILL', None, 0), ('SAME', 'L50X5', 1),
                                ('CONFLICT', 'L50X5', 2), ('CHECK', None, None)], legacy=False))
service = cat.CatalogService(Settings(docs, inputs.REGISTRY))
project_name = next(p for p in service.registry.projects if cat.key(p) == cat.key('Гараж Казахстан'))
text = ('Марка\tПрофиль\nNEW\tL50X5\nFILL\tL50X5\nSAME\tL50X5\n'
        'CONFLICT\tL63X5\nCHECK\tПРОВЕРИТЬ')
plan = service.preview(project_name, text)
by_mark = {r.mark: r for r in plan.rows}
for number, status, mark in ((13, 'NEW', 'NEW'), (14, 'FILL', 'FILL'), (15, 'SAME', 'SAME'),
                             (16, 'CONFLICT', 'CONFLICT'), (17, 'CHECK_REQUIRED', 'CHECK')):
    test(f'{number:02d} Stage3 {status}', lambda s=status, m=mark: require(by_mark[m].status == s))


def apply_preserves():
    result = service.apply(plan, confirmations=(cat.key('CONFLICT'), cat.key('CHECK')))
    rows = migration.assess(catalog_path.read_bytes())['rows']
    assert {r['mark']: r['akz'] for r in rows} == {'FILL': 0, 'SAME': 1, 'CONFLICT': 2,
                                                  'CHECK': None, 'NEW': None}
    assert cat.read_catalog(catalog_path.read_bytes())[0][cat.key('CONFLICT')][2] == 'L63X5'
    assert result['added'] == 1 and result['filled'] == 2 and result['changed'] == 1
    return result


test('18 actual Apply preserves all manual AKZ and new blank', apply_preserves)


def error_stop():
    before_bytes = catalog_path.read_bytes()
    p = service.preview(project_name, 'Марка\tПрофиль\nBROKEN\t\nNEW2\tL50X5')
    assert p.rows[0].status == 'ERROR'
    expect_stop(lambda: service.apply(p), 'ошибки')
    assert catalog_path.read_bytes() == before_bytes


test('19 Stage3 ERROR prevents all writes', error_stop)
test('20 TXT TSV Markdown parsing unchanged', lambda: require(
    [(e.mark, e.profile, e.error) for e in cat.parse_input('Марка\tПрофиль\nМ1\tL50X5')] ==
    [(e.mark, e.profile, e.error) for e in cat.parse_input('| Марка | Профиль |\n|---|---|\n|М1|L50X5|')] and
    cat.parse_input('| Марка | Профиль |\n|---|---|\n|М1|L50X5|')[0].lines == (3,)))


def legacy_apply():
    data = fixture([('LEGACY', 'L50X5', 'KEEP PAINT', 0)])
    catalog_path.write_bytes(data)
    p = service.preview(project_name, 'Марка\tПрофиль\nLEGACY\tL63X5\nNEWLEGACY\tL50X5')
    service.apply(p, confirmations=(cat.key('LEGACY'),))
    book = openpyxl.load_workbook(catalog_path)
    try:
        sheet = book.active
        assert sheet.max_column == 4 and sheet['C2'].value == 'KEEP PAINT' and sheet['D2'].value == 0
        assert sheet['C3'].value is None and sheet['D3'].value is None
        assert sheet.auto_filter.ref == 'A1:D3'
    finally:
        book.close()


test('21 legacy Stage3 Apply preserves paint and AKZ', legacy_apply)


def create_catalog():
    catalog_path.unlink()
    p = service.preview(project_name, 'Марка\tПрофиль\nCREATED\tL50X5')
    service.apply(p)
    a = migration.assess(catalog_path.read_bytes())
    assert tuple(a['headers']) == cat.HEADERS and a['rows'][0]['akz'] is None
    book = openpyxl.load_workbook(catalog_path)
    try:
        assert book.active.auto_filter.ref == 'A1:C2'
    finally:
        book.close()


test('22 new Stage3 catalog has exactly 3 columns', create_catalog)
legacy_path = RUN / 'source_catalog.xlsx'
test('23 Stage5 reads new 3-column mapping', lambda: require(len(inputs.mark_profiles(new_path)) == 195))
test('24 legacy Stage5 and Stage3 read compatibility', lambda: require(
    inputs.mark_profiles(new_path) == inputs.mark_profiles(legacy_path) and
    cat.read_catalog(new)[0] == cat.read_catalog(source)[0]))


def stage5_preview():
    accepted = json.loads(Path(baseline['accepted_four_mark_preview']).read_text(encoding='utf8'))['preview']
    vom = write_fixture('shipment_copy.xls', Path(baseline['vom_source']).read_bytes())
    template = Path(accepted['template'])
    with patch.object(core, 'project_sources', return_value=(project, new_path, template)):
        result = core.preview(accepted['passport']['project'], accepted['passport']['number'], vom)
    assert len(result.data['marks']) == 4 and result.data['controlled_units'] == 33
    expected = {'Бп2-1': (8, [('1280', '5'), ('1024', '8.4')]),
                'Пс1-9': (9, [('3600', '4')]), 'Пс1-10': (15, [('6000', '4')]),
                'Пс2-16': (1, [('480', '3')])}
    from decimal import Decimal
    actual = {m['mark']: (m['Q'], [(str(Decimal(s['L']).normalize()),
                                  str(Decimal(str(s['S'])).normalize())) for s in m['segments']])
              for m in result.data['marks']}
    normalized = lambda pairs: [(Decimal(l), Decimal(s)) for l, s in pairs]
    for mark, (q, segments) in expected.items():
        assert actual[mark][0] == q and normalized(actual[mark][1]) == normalized(segments)
    assert not result.data['network_attempts']
    save_json(RUN / 'stage5_isolated_preview.json', result.data)
    return {'marks': 4, 'items': 33, 'calculations': actual, 'Word_generation': 'NOT_RUN'}


test('25 actual full Stage5 preview 4 marks 33 items expected geometry', stage5_preview)


def gui_compatibility():
    from passport_app.catalog_gui import CatalogApplication
    catalog_path.write_bytes(fixture([('GUI', None, 0)], legacy=False))
    app = CatalogApplication(Settings(docs, inputs.REGISTRY))
    app.withdraw()
    try:
        assert '№ АКЗ' in app.result_var.get() and 'Краска' not in app.result_var.get()
        app.project_var.set(project_name)
        app.input.insert('1.0', 'Марка\tПрофиль\nGUI\tL50X5\nNEWGUI\tL63X5')
        app.update()
        app.check()
        app.update()
        assert app.plan is not None and len(app.tree.get_children()) == 2
        assert {r.status for r in app.plan.rows} == {'FILL', 'NEW'}
        assert str(app.apply_button['state']) == 'normal'
        with patch('passport_app.catalog_gui.messagebox.askyesno', return_value=True) as confirm:
            app.apply()
            assert 'Краска' not in confirm.call_args.args[1] and '№ АКЗ' in confirm.call_args.args[1]
        rows = migration.assess(catalog_path.read_bytes())['rows']
        assert {r['mark']: r['akz'] for r in rows} == {'GUI': 0, 'NEWGUI': None}
        assert app.plan is None and str(app.apply_button['state']) == 'disabled'
        return 'real Tk widgets/check/apply, DEV registry, isolated 3-column file; no clipboard or OS input'
    finally:
        app.destroy()


test('26 Stage3 actual Tk GUI check Apply new schema', gui_compatibility)


def independent_package_check():
    part, old_xml = cat.sheet_part(source)
    _, new_xml = cat.sheet_part(new)
    with ZipFile(BytesIO(source)) as a, ZipFile(BytesIO(new)) as b:
        changed = [name for name in a.namelist() if a.read(name) != b.read(name)]
        assert set(changed) == {part, 'xl/workbook.xml'}
        expected_workbook = a.read('xl/workbook.xml').replace(b'!$A$1:$D$1', b'!$A$1:$C$1')
        assert b.read('xl/workbook.xml') == expected_workbook
    import re
    cells = lambda xml, c: re.findall(rf'<c\b[^>]*\br="{c}\d+"[^>]*(?:/>|>.*?</c>)', xml, re.S)
    assert cells(old_xml, 'A') == cells(new_xml, 'A')
    assert cells(old_xml, 'B') == cells(new_xml, 'B')
    old_book, new_book = openpyxl.load_workbook(BytesIO(source)), openpyxl.load_workbook(BytesIO(new))
    try:
        a, b = old_book.active, new_book.active
        assert b.column_dimensions['C'].width == a.column_dimensions['D'].width == 18
        assert a.column_dimensions['A'].width == b.column_dimensions['A'].width
        assert a.column_dimensions['B'].width == b.column_dimensions['B'].width
        assert a.freeze_panes == b.freeze_panes == 'A2'
        assert a.auto_filter.ref == 'A1:D196' and b.auto_filter.ref == 'A1:C196'
    finally:
        old_book.close()
        new_book.close()
    return {'changed_zip_parts': changed, 'A_B_XML': 'byte-identical', 'AKZ_column_width': 18}


test('27 independent XML package formatting and range preservation', independent_package_check)


def stale_manual_change():
    catalog_path.write_bytes(fixture([('MANUAL', None, 0)], legacy=False))
    p = service.preview(project_name, 'Марка\tПрофиль\nMANUAL\tL50X5')
    edited = fixture([('MANUAL', None, 2)], legacy=False)
    catalog_path.write_bytes(edited)
    expect_stop(lambda: service.apply(p), 'изменился')
    assert catalog_path.read_bytes() == edited
    return 'manual AKZ edit after preview rejects stale Apply; manual 2 preserved'


test('28 manual AKZ change after preview safe STOP', stale_manual_change)
test('29 DEV registries profiles old acts templates VOM unchanged', lambda: require(
    all(file_sha(p) == expected for p, expected in protected.items())))
test('30 no production registry access', lambda: require(not denied and all(
    Path(p).name.casefold() not in ('взмк.xlsx', 'узк.xlsm') for p in audit_reads)))
test('31 real project catalog not migrated by isolated tests', lambda: require(
    file_sha(baseline['source_catalog']) == baseline['source_sha256']))
after = product_hashes()
test('32 tested product bytes unchanged during tests', lambda: require(before == after))
save_json(RUN / 'product_hashes_AFTER_TESTS.json', after)
save_json(RUN / 'isolated_tests.json', results)
save_json(RUN / 'isolation_audit.json', {'production_registry_attempts': denied,
          'python_workbook_read_paths': sorted(audit_reads),
          'native_paths': 'explicit validation XLSX files and shipment_copy.xls only',
          'Word_generation': 'NOT_RUN', 'real_migration': 'NOT_RUN during isolated suite'})
print(f'ISOLATED PASS {len(results)}/{len(results)}', flush=True)
