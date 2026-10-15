from pathlib import Path
import sys,json,hashlib
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT))
from passport_app.registry import read_registry,select_project
from akz_stage6 import core as c
def dump(n,x):(OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
review=ROOT/'validation/stage6_final_targeted_rereview_r2_20261009'
expected=json.loads((review/'tested_product_hashes_independent.json').read_text(encoding='utf-8'))['CURRENT']
current={p:c.sha256(p) for p in expected}
assert current==expected,'CURRENT_PRODUCT_CHANGED_STOP'
r=read_registry(c.REGISTRY);project=select_project(r,'Гараж Казахстан')
for name,path,key in [('DEV_REGISTRY_CHECKPOINT.json',c.REGISTRY,'sha256'),('DEV_UZK_CHECKPOINT.json',ROOT/'DEV/УЗК_DEV.xlsm','dev_initial_sha256')]:
 data=json.loads((ROOT/'DEV'/name).read_text(encoding='utf-8'));assert c.sha256(path)==data[key],name
audit=ROOT/'validation/stage6_new_templates_audit_20261009_125930'
contract=json.loads((audit/'IMPLEMENTATION_CONTRACT.json').read_text(encoding='utf-8'))
for v in contract['variants'].values():
 assert c.sha256(v['act_path'])==v['act_sha256'],'TEMPLATE_CHANGED_STOP'
 assert c.sha256(v['application']['path'])==v['application']['sha256'],'TEMPLATE_CHANGED_STOP'
guard=json.loads((audit/'sources_after.json').read_text(encoding='utf-8'))
baseline={p:c.sha256(p) for p in guard}
dump('sources_before.json',baseline);dump('product_before.json',current)
for p in expected:
 dest=OUT/'BEFORE'/Path(p).relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(Path(p).read_bytes())
vom_path=c.PROJECT/'Паспорт/9 от 06.10.2026/ВОМ отгрузочный счет 183  (5).xls'
vom=c.visible_vom(vom_path)
dump('selected_visible_vom.json',vom)
real=c.PROJECT/'Справочник марок.xlsx'
from openpyxl import load_workbook
b=load_workbook(real,data_only=False);s=b.active
assert [x.value for x in s[1]]==['Марка','Профиль','№ АКЗ']
assert all(row[2].value is None for row in s.iter_rows(min_row=2)),'REAL_ASSIGNMENTS_NOT_BLANK'
b.close()
good=[]
for row in vom['rows']:
 try:
  c.matching({'rows':[row]},{row['mark']:{'assignment':0,'row':2}})
 except c.Stage6Error:continue
 if row['mark'] not in [r['mark'] for r in good]:good.append(row)
chosen=good[:3];assert len(chosen)==3
dump('preflight.json',{'accepted_review':str(review/'REVIEW_MANIFEST.json'),'review_sha256':c.sha256(review/'REVIEW_MANIFEST.json'),'product_matches_review':True,'project':project.__dict__,'dev_checkpoints':'PASS','templates':'ALL_SIX_PIN_MATCH','real_catalog_assignments':'ALL_BLANK','fixture_candidate_rows':chosen,'guard_count':len(baseline)})
print(json.dumps({'product':'MATCH','DEV':'MATCH','templates':'MATCH','chosen':chosen,'guard':len(baseline)},ensure_ascii=False))
