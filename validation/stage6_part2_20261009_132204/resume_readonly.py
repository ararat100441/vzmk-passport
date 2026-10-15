"""Resume evidence only: never imports the product, opens COM, saves XLSX, or runs tests."""
from pathlib import Path
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
import zipfile
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def read(name):
    return json.loads((OUT / name).read_text(encoding='utf-8-sig'))
def dump(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
def require(value, label):
    if not value:
        raise AssertionError(label)

baseline = read('sources_before.json')
allowed = {str(ROOT/'akz_stage6/gui.py'), str(ROOT/'STAGE6_DEV.md')}
current_sources = {p: sha(p) for p in baseline}
changed = [p for p in baseline if baseline[p] != current_sources[p]]
require(set(changed) == allowed, 'UNEXPECTED_SOURCE_CHANGE_STOP')
tested_before = read('tested_product_before.json')
tested_after = read('tested_product_after.json')
current_product = {p: sha(p) for p in tested_after}
require(tested_before == tested_after == current_product, 'TESTED_PRODUCT_BINDING_STOP')
original_product = read('product_before.json')
for p, pin in original_product.items():
    if p not in allowed:
        require(sha(p) == pin, 'PART1_PRODUCT_CHANGED_STOP: '+p)
new_product = [str(ROOT/'akz_stage6/part2.py'), str(ROOT/'tests/test_stage6_part2.py')]
all_product = {p: sha(p) for p in set(current_product) | set(original_product) | set(new_product)}

tests = read('targeted_validation.json')
require(len(tests['results']) == 14 and all(x['status']=='PASS' for x in tests['results']), 'TARGETED_RESULTS_STOP')
unit = (OUT/'unit_tests.txt').read_text(encoding='utf-8')
require('Ran 33 tests' in unit and unit.rstrip().endswith('OK') and unit.count(' ... ok')==33, 'UNIT_RESULTS_STOP')
require(tests['audit']=={'network_events':[], 'forbidden_production_events':[]}, 'AUDIT_EVENTS_STOP')
native = tests['results'][0]['detail']
require(tests['results'][0]['name']=='all_six_real_outputs_current_product_readonly_full_calculation', 'NATIVE_EVIDENCE_STOP')
assignments = read('DEV_assignments.json')
require(sha(assignments['catalog'])==assignments['sha256'], 'DEV_FIXTURE_CHANGED_STOP')
generation = read('real_generation.json')
require(generation['status']=='COMPLETE_ALL_SELECTED', 'PAIR_GENERATION_STOP')
missing = 'НЕТ ДАННЫХ / ТЕСТ'
expected = {1: ('Бп2-5','37.20','1.412',65,66), 2: ('Бп2-18','31.32','1.184',51,52), 3: ('Бп2-19','36.74','1.393',98,99)}
old_tokens = ('1/Б2','1/2025','4/2026','5/2026','6/2026','РУССОКО','Велесстрой','2-ЭГ-2019','ВС.24.ПО','13-25','314-25','Журавлёв','Шляхтуров')
outputs = []
for variant, (mark,mass,area,last,total) in expected.items():
    directory = Path(generation['groups'][str(variant)]['path'])
    act = directory/f'9_2026-{variant} АКЗ TEST.xlsx'
    app = directory/f'9_2026-{variant} Приложение АКЗ TEST.xlsx'
    receipt = json.loads((directory/'PAIR.json').read_text(encoding='utf-8'))
    hashes = {p.name:sha(p) for p in (act,app)}
    prior_native = next(x for x in native if x['variant']==variant)
    require(str(directory)==prior_native['path'] and hashes==receipt['sha256']==prior_native['sha256'], 'EXACT_SIX_HASH_BINDING_STOP')
    require(receipt['official']=='STOP' and receipt['state']=='COMPLETE_TEST_PAIR' and receipt['date']=='2026-10-05', 'RECEIPT_IDENTITY_STOP')
    require(prior_native['native']=={'act_pages':1,'application_pages':1,'service_printed':False,'full_calculation':'PASS','read_only':True,'normal_load':True}, 'NATIVE_RESULT_STOP')
    require(receipt['climate']=={'day':'2026-10-05','air_c':None,'relative_humidity':None,'surface_c':None,'dew_point_c':None,'confirmation':'MISSING'}, 'CLIMATE_STOP')
    books = [load_workbook(p, data_only=False) for p in (act,app)]
    caches = [load_workbook(p, data_only=True) for p in (act,app)]
    try:
        a,b = books; ac,bc = caches; s=a['4']; t=b['Приложение №1']; sc=ac['4']; tc=bc['Приложение №1']
        number=f'9/2026-{variant}'
        require(s['E4'].value==number+' от ' and sc['G8'].value==number+' от ', 'ACT_NUMBER_STOP')
        require(sc['H4'].value==sc['I8'].value==datetime(2026,10,5), 'ACT_DATE_STOP')
        require(t['A1'].value==f'ТЕСТ / ЧЕРНОВИК\nПриложение № {number} к акту АКЗ № {number} от 05.10.2026 г.', 'APP_NUMBER_DATE_STOP')
        require(t['A2'].value=='Объект: '+s['E7'].value, 'ACT_APP_OBJECT_STOP')
        require(s['B3'].value=='ТЕСТ / ЧЕРНОВИК — АКЗ НЕ ПОДТВЕРЖДЕНО' and s['B43'].value=='ТЕСТ: фактические показатели не заполнены. Выдача запрещена.', 'TEST_MARK_STOP')
        require('ТЕСТ / ЧЕРНОВИК' in str(s.oddHeader) and 'ТЕСТ / ЧЕРНОВИК' in str(t.oddHeader), 'TEST_PRINT_MARK_STOP')
        require(s.oddFooter.center.text==t.oddFooter.center.text==f"ACT-ID {receipt['act_id']} | 2026-10-05", 'PAIR_ACT_ID_STOP')
        require([tc.cell(5,c).value for c in range(1,8)]==[mark,'Балка',1,float(mass),float(mass),float(area),float(area)], 'SOURCE_ROW_STOP')
        require(all(tc.cell(r,c).value is None for r in range(6,last+1) for c in range(1,8)), 'HISTORICAL_BODY_STOP')
        require(all(t.row_dimensions[r].hidden==(6<=r<=last) for r in range(1,total+1)), 'ROW_VISIBILITY_STOP')
        require(tc.cell(total,1).value=='Итого' and not t.row_dimensions[total].hidden, 'TOTAL_ROW_STOP')
        for col, val in [('C','1'),('E',mass),('G',area)]:
            require(t[f'{col}{total}'].value==f'=SUM({col}5:{col}{last})', 'TOTAL_FORMULA_STOP')
            require(Decimal(str(tc[f'{col}{total}'].value))==Decimal(val), 'CACHED_TOTAL_STOP')
        require(t.print_title_rows=='$1:$4' and str(t.print_area)==receipt['templates']['application_print_area'], 'APP_PRINT_STOP')
        require(str(s.print_area)==receipt['templates']['act_print_area'], 'ACT_PRINT_STOP')
        require(sorted(map(str,s.merged_cells.ranges))==receipt['templates']['act_merges'] and sorted(map(str,t.merged_cells.ranges))==receipt['templates']['application_merges'], 'MERGES_STOP')
        common=['G11','B13','E13','I13','E15','G37','E39','E40','E41','E42','D49','E46','E47']
        factual = list(common)
        for layer in receipt['templates']['layers']:
            col=layer['column']
            for cell,value in [(layer['ordinal_cell'],layer['ordinal']),(col+'18',layer['material']),(col+'21',layer['color']),(col+'25',layer['required_raw'])]:
                require(s[cell].value==value, 'LAYER_REQUIREMENT_STOP')
            factual += [col+str(r) for r in (19,20,22,23,24,26,27,28,30,31,32,33,34)]
        for cell in factual:
            require(s[cell].value==sc[cell].value==('?' if cell in ('G11','E15','G37') else missing), 'HISTORICAL_FACT_STOP: '+cell)
        if variant==2:
            require(s['I25'].value=='∑1+2=180' and receipt['templates']['layers'][1]['required_individual_microns'] is None and receipt['templates']['layers'][1]['requirement_status']=='NOT_PROVEN', 'INVENTED_SECOND_LAYER_STOP')
            require(a['Лист1'].sheet_state=='veryHidden' and a['Лист1']['A1'].value=='СЛУЖЕБНЫЙ ЛИСТ / '+missing, 'SERVICE_SHEET_STOP')
            require(sum(c.value is not None for row in a['Лист1'] for c in row)==1, 'HISTORICAL_HELPER_STOP')
        require(all(c.value is None for row in s.iter_rows(min_col=12,max_col=15,min_row=1,max_row=240) for c in row), 'HISTORICAL_HELPER_RANGE_STOP')
        for book,path in zip(books,(act,app)):
            formulas=[];strings=[]
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:strings.append(str(cell.value))
                        if cell.data_type=='f':formulas.append({'sheet':sheet.title,'cell':cell.coordinate,'formula':cell.value})
                strings += [str(x) for x in (sheet.oddHeader,sheet.oddFooter,sheet.evenHeader,sheet.evenFooter,sheet.firstHeader,sheet.firstFooter)]
            strings += [str(x) for x in book.defined_names.values()]
            require(not any(tok in text for text in strings for tok in old_tokens), 'HISTORICAL_REQUISITE_STOP')
            if path==act:require({x['cell'] for x in formulas}=={'G8','I8'} and all(x['sheet']=='4' for x in formulas), 'ACT_FORMULA_STOP')
            else:require({x['cell'] for x in formulas}=={f'{c}{total}' for c in 'CEG'}, 'APP_FORMULA_STOP')
            with zipfile.ZipFile(path) as z:
                xml='\n'.join(z.read(name).decode('utf-8') for name in z.namelist() if name.endswith('.xml'))
                require('RANDBETWEEN' not in xml.upper() and '_XLFN.RAND' not in xml.upper(), 'RANDOM_PSEUDOMEASUREMENT_STOP')
            outputs.append({'variant':variant,'kind':'act' if path==act else 'application','path':str(path),'sha256':hashes[path.name],'formulas':formulas,'initial_selection':book.active.sheet_view.selection[0].activeCell,'sheets':{ws.title:ws.sheet_state for ws in book}})
        require(s.sheet_view.selection[0].activeCell==t.sheet_view.selection[0].activeCell=='A1', 'INITIAL_VIEW_STOP')
    finally:
        for book in books+caches:book.close()
    require(hashes=={p.name:sha(p) for p in (act,app)}, 'READ_ONLY_SHA_CHANGED_STOP')

preflight=read('preflight.json')
require(sha(preflight['accepted_review'])==preflight['review_sha256'], 'ACCEPTED_REVIEW_CHANGED_STOP')
report={
    'status':'PASS','created_utc':datetime.now(timezone.utc).isoformat(),
    'scope':'Existing files only; no product import, no COM, no GUI, no regeneration, no test rerun',
    'product_tested_before_after_current_equal':True,'current_product':all_product,
    'sources_compared':len(baseline),'sources_unchanged':len(baseline)-len(changed),
    'authorized_part2_changes':[{ 'path':p,'before':baseline[p],'after':current_sources[p]} for p in changed],
    'unexpected_source_changes':[], 'new_part2_product':new_product,
    'fixture':{'path':assignments['catalog'],'sha256':assignments['sha256'],'unchanged':True},
    'six_outputs':outputs,'targeted_checks_reused':tests['results'],'unit_checks_reused':33,
    'native_excel_evidence':'Reused exact six files; prior normal-load read-only full calculation; current hashes equal tested hashes',
    'native_excel_rerun':'NOT_RUN; sufficient bound existing evidence',
    'official':'STOP','network_production_audit':tests['audit'],
    'audit_limit':'Python audit-hook scope only; not an OS-wide trace',
    'escape_boundary':'Exact keypress time not recorded; all final evidence existed when resume inspected it',
}
dump('sources_after.json',current_sources)
dump('current_product_hashes.json',all_product)
dump('resume_readonly_validation.json',report)
print(json.dumps({'status':'PASS','sources':len(baseline),'unchanged':len(baseline)-len(changed),'authorized_changes':len(changed),'product_binding':'PASS','six_outputs':outputs},ensure_ascii=False,indent=2))
