from review_common import *
from openpyxl import load_workbook
from akz_stage6 import core as c
import tkinter as tk
from akz_stage6.gui import Launcher
import win32file,win32con,pywintypes,unittest,io
S=E/'sandbox';F=S/'fixtures';O=S/'output'
F.mkdir(parents=True);O.mkdir()
results=[]
def run(name,fn):
    try:
        detail=fn();results.append({'name':name,'status':'PASS','detail':detail});print('PASS',name,flush=True)
    except Exception as exc:
        results.append({'name':name,'status':'FAIL','error':str(exc),'traceback':traceback.format_exc()});print('FAIL',name,str(exc),flush=True)
    save('independent_targeted.json',{'results':results,'audit':AUDIT})
def stop(fn,code):
    try:fn()
    except Exception as exc:
        assert code in str(exc),(code,type(exc).__name__,str(exc))
        return {'type':type(exc).__name__,'message':str(exc),'cause':repr(exc.__cause__)}
    raise AssertionError('EXPECTED STOP '+code)
def yes(condition):assert condition
def cp(source,name):
    p=F/name;shutil.copyfile(source,p);return p
reg=cp(c.REGISTRY,'DEV_registry.xlsx');act=cp(c.ACT,'ACT_template.xlsx');application=cp(c.APPLICATION,'APPLICATION_template.xlsx')
old=load(FIX/'current_preview.json')['payload']
catalog=cp(next(p for p in old['sources'] if 'TEST_Справочник' in p),'TEST_catalog.xlsx')
vom=cp(old['vom']['path'],'selected_VOM.xls')
for n in (1,17,18):cp(FIX/f'sandbox_run2/fixtures/SYNTHETIC_{n}.xlsx',f'SYNTHETIC_{n}.xlsx')
initial={str(p):sha(p) for p in F.iterdir()}
def inspect(target,p):
    b=load_workbook(target/APPNAME);cached=load_workbook(target/APPNAME,data_only=True)
    try:
        s=b.active;t=cached.active;n=len(p.payload['selected'])
        hidden=[r for r in range(1,22) if s.row_dimensions[r].hidden]
        assert hidden==list(range(4+n,21))
        for col,key in zip('CEG',('quantity','mass_total','area_total')):
            assert s[col+'21'].value==f'=SUM({col}4:{col}20)'
            assert abs(Decimal(str(t[col+'21'].value))-p.payload['totals'][key])<Decimal('0.00000001')
        assert [s[f'A{r}'].value for r in range(4,4+n)]==[v['mark'] for v in p.payload['selected']]
        assert all(s.cell(r,col).value is None for r in range(4+n,21) for col in range(1,8))
        assert load(target/'PAIR.json')['state']=='COMPLETE_TEST_PAIR'
        return {'target':str(target),'row_count':n,'hidden':hidden,'total_row_visible':not s.row_dimensions[21].hidden,'totals':[t[col+'21'].value for col in 'CEG'],'sha':{str(f):sha(f) for f in target.iterdir()}}
    finally:b.close();cached.close()
with patch.multiple(c,REGISTRY=reg,ACT=act,APPLICATION=application,FIXTURES=F,OUTPUT=O,PINS={act:sha(act),application:sha(application)}):
    p=c.preview('Гараж Казахстан','9/2026',vom,catalog)
    save('isolated_preview.json',{'seal':p.seal,'payload':p.payload})
    def regression():
        rows=p.payload['selected']
        assert [(r['row'],r['mark'],str(r['quantity']),str(r['mass_total']),str(r['area_total'])) for r in rows]==[(40,'Бп2-5','1.0','37.2','1.412'),(53,'Бп2-18','1.0','31.32','1.184'),(54,'Бп2-19','1.0','36.74','1.393')]
        assert p.payload['totals']=={'quantity':Decimal(3),'mass_total':Decimal('105.26'),'area_total':Decimal('3.989')}
        assert len(p.payload['vom']['rows'])==104 and len(p.payload['vom']['hidden_rows'])==397
        return {'rows':rows,'totals':p.payload['totals'],'visible':104,'hidden':397}
    run('three_mark_native_source_regression',regression)
    target=c.generate(p,'GOOD_N3')
    run('R1_N3_complete_hidden_totals',lambda:inspect(target,p))
    goodsha={str(f):sha(f) for f in target.iterdir()}
    for n in (1,17):
        pn=c.preview('Гараж Казахстан','9/2026',F/f'SYNTHETIC_{n}.xlsx',catalog,synthetic=True)
        tn=c.generate(pn,f'BOUNDARY_N{n}')
        run(f'R1_N{n}_complete_hidden_totals',lambda:inspect(tn,pn))
    before=set(O.iterdir())
    def capacity():
        result=stop(lambda:c.preview('Гараж Казахстан','9/2026',F/'SYNTHETIC_18.xlsx',catalog,synthetic=True),'TEMPLATE_CAPACITY_STOP')
        assert before==set(O.iterdir());return result
    run('R1_N18_capacity_STOP_no_publication',capacity)
    original=c.native_open
    def corrupt(excel,path,*,read_only):
        if Path(path).name==APPNAME and not read_only:Path(path).write_bytes(b'NOT A ZIP: isolated corrupt workbook')
        return original(excel,path,read_only=read_only)
    def invalid():
        with patch.object(c,'native_open',side_effect=corrupt):d=stop(lambda:c.generate(p,'CORRUPT'),'INVALID_XLSX')
        assert before==set(O.iterdir()) and not list(O.glob('.pending-*'));return d
    run('R2_corrupt_XLSX_distinct_INVALID_full_rollback',invalid)
    handles=[];proof=[]
    def lock(excel,path,*,read_only):
        if Path(path).name==APPNAME and not read_only:
            healthy=sha(path);assert healthy==sha(application)
            h=win32file.CreateFile(str(path),win32con.GENERIC_READ,0,None,win32con.OPEN_EXISTING,0,None)
            handles.append(h)
            try:return original(excel,path,read_only=read_only)
            except Exception as exc:
                proof.append({'path':str(path),'sha_before_lock':healthy,'share':0,'classified_exception':type(exc).__name__,'code':getattr(exc,'code',None),'message':str(exc),'cause_winerror':getattr(exc.__cause__,'winerror',None)})
                raise
        return original(excel,path,read_only=read_only)
    root=tk.Tk();root.withdraw();gui=Launcher(root);gui.checked=p
    try:
        with patch.object(c,'native_open',side_effect=lock),patch('akz_stage6.gui.messagebox.showerror') as message:
            gui.generate()
        msg=message.call_args.args[1]
        pending=list(O.glob('.pending-*'))
        record={'proof':proof,'GUI_message':msg,'pending_after_GUI':[str(t) for t in pending],'contents':[str(f) for t in pending for f in t.iterdir()],'published_new':[str(t) for t in set(O.iterdir())-before if not t.name.startswith('.pending-')],'lock_still_held':len(handles)==1,'preview_invalidated':gui.checked is None,'old_good_unchanged':all(sha(f)==h for f,h in goodsha.items())}
        save('persistent_lock_GUI_handler.json',record)
        run('R2_persistent_lock_classification_GUI',lambda:yes(proof[0]['cause_winerror']==32 and proof[0]['code']=='FILE_LOCKED' and 'FILE_LOCKED' in msg and proof[0]['path'] in msg and 'Закройте' in msg and 'INVALID_XLSX' not in msg))
        run('R2_persistent_lock_full_rollback',lambda:yes(not pending and before==set(O.iterdir())))
    finally:
        for h in handles:h.Close()
        root.destroy()
    # Capture leaked pending files after release, before review-only cleanup.
    leaked={str(t):{str(f):sha(f) for f in t.iterdir()} for t in O.glob('.pending-*')}
    save('persistent_lock_after_release.json',leaked)
    retry=c.generate(p,'RETRY_AFTER_RELEASE')
    run('R2_release_successful_retry',lambda:inspect(retry,p))
    for t in list(O.glob('.pending-*')):
        assert t.resolve().parent==O.resolve() and O.resolve().is_relative_to(E.resolve()) and t.name.startswith('.pending-')
        shutil.rmtree(t)
    save('review_owned_pending_cleanup.json',{'removed_after_release':list(leaked),'pending_remaining':list(map(str,O.glob('.pending-*'))),'reason':'Remove only newly created review fixture residue after capturing product rollback failure'})
    for path in (reg,act,application,catalog,vom):
        data=path.read_bytes()
        try:
            path.write_bytes(data+b'ISOLATED CHANGED INPUT')
            run('R2_INPUT_CHANGED_'+path.name,lambda:stop(lambda:c.generate(p,'CHANGED'),'INPUT_CHANGED_AFTER_PREVIEW_STOP'))
        finally:path.write_bytes(data)
    run('R2_OUTPUT_EXISTS_preserves_pair',lambda:stop(lambda:c.generate(p,'GOOD_N3'),'OUTPUT_ALREADY_EXISTS_STOP'))
    denied=pywintypes.error(5,'CreateFile','Access denied')
    with patch('win32file.CreateFile',side_effect=denied):
        run('R2_ACCESS_DENIED_error5_unit',lambda:stop(lambda:c.probe_file_access(application,read_only=False),'ACCESS_DENIED'))
    run('R2_valid_XLSX_generic_OTHER',lambda:stop(lambda:c.excel_file_failure(application,RuntimeError('isolated generic Excel failure'),read_only=True,operation='review'),'OTHER_ERROR'))
    assert all(sha(f)==h for f,h in goodsha.items()) and all(sha(f)==h for f,h in initial.items())
suite=unittest.defaultTestLoader.discover(str(ROOT/'tests'),pattern='test_stage6.py')
log=io.StringIO();unit=unittest.TextTestRunner(stream=log,verbosity=2).run(suite)
(E/'unit_tests_independent.txt').write_text(log.getvalue(),encoding='utf-8')
assert unit.wasSuccessful() and unit.testsRun==16
save('independent_targeted.json',{'results':results,'summary':{'PASS':sum(r['status']=='PASS' for r in results),'FAIL':sum(r['status']=='FAIL' for r in results)},'units':{'run':unit.testsRun,'pass':unit.wasSuccessful()},'audit':AUDIT})
print('SUMMARY',len(results),'FAIL',sum(r['status']=='FAIL' for r in results),'UNIT',unit.testsRun,unit.wasSuccessful(),flush=True)
