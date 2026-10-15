from review_common import *
from openpyxl import load_workbook
m=load(FIX/'EVIDENCE_MANIFEST.json');baseline=load(FIX/'baseline.json')
assert {n:sha(PAIR/n) for n in EXPECTED}==EXPECTED,'EXACT NEW PAIR SHA STOP'
maps={'fix_evidence':m['files'],'guarded':m['guarded_current'],'product_current':m['product']['CURRENT']}
checks={}
for label,mapping in maps.items():
    now={p:sha(p) for p in mapping}
    bad=[p for p,h in mapping.items() if now[p]!=h]
    assert not bad,(label,bad)
    checks[label]={'count':len(now),'mismatches':bad}
assert sha(m['review_baseline'])==m['review_baseline_sha256']
assert baseline['product']==m['product']['BEFORE']
assert m['product']['AFTER']==m['product']['CURRENT']==load(FIX/'product_after.json')
assert load(FIX/'tested_product_hashes.json')==m['product']
changed=[p for p,h in baseline['product'].items() if h!=m['product']['CURRENT'][p]]
assert set(Path(p).name for p in changed)=={'core.py','gui.py'} and len(changed)==2
semantic={}
for filename in ('core.py','gui.py'):
    before=(FIX/('BEFORE_'+filename)).read_text('utf-8');after=(ROOT/'akz_stage6'/filename).read_text('utf-8')
    assert sha(FIX/('BEFORE_'+filename))==baseline['product'][str(ROOT/'akz_stage6'/filename)]
    (E/('actual_'+filename+'.diff')).write_text(''.join(difflib.unified_diff(before.splitlines(True),after.splitlines(True),fromfile='BEFORE/'+filename,tofile='CURRENT/'+filename)),encoding='utf-8')
    oldtree=ast.parse(before);newtree=ast.parse(after)
    oldnodes={n.name:ast.dump(n,include_attributes=False) for n in oldtree.body if isinstance(n,(ast.FunctionDef,ast.ClassDef))}
    newnodes={n.name:ast.dump(n,include_attributes=False) for n in newtree.body if isinstance(n,(ast.FunctionDef,ast.ClassDef))}
    semantic[filename]={'changed':[n for n in oldnodes if oldnodes[n]!=newnodes.get(n)],'added':[n for n in newnodes if n not in oldnodes],'unchanged':[n for n in oldnodes if oldnodes[n]==newnodes.get(n)]}
    if filename=='core.py':
        assert set(semantic[filename]['changed'])=={'native_open','write_act','write_application','validate_saved_pair'}
        assert all(n in semantic[filename]['unchanged'] for n in ('number','matching','preview','templates','current_passport','ClimateDay','generate','verify_preview'))
    else:
        class Normalize(ast.NodeTransformer):
            def visit_Call(self,n):
                n=self.generic_visit(n)
                if isinstance(n.func,ast.Attribute) and isinstance(n.func.value,ast.Name) and n.func.value.id=='core' and n.func.attr=='error_message':n.func=ast.Name(id='str',ctx=ast.Load())
                return n
        assert ast.dump(oldtree,include_attributes=False)==ast.dump(Normalize().visit(newtree),include_attributes=False)
tests=load(FIX/'targeted_tests.json');assert len(tests)==16 and all(x['status']=='PASS' for x in tests)
unitfiles=[p for p in FIX.glob('*unit*') if p.is_file()]
unittexts={str(p):p.read_text('utf-8-sig') for p in unitfiles if p.suffix in ('.txt','.log')}
assert any('Ran 16 tests' in t and '\nOK' in t for t in unittexts.values()),unitfiles
assert m['written_last'] and all(Path(p).stat().st_mtime_ns <= (FIX/'EVIDENCE_MANIFEST.json').stat().st_mtime_ns for p in m['files'])
guard={**m['guarded_current'],**m['files'],**m['product']['CURRENT'],str(FIX/'EVIDENCE_MANIFEST.json'):sha(FIX/'EVIDENCE_MANIFEST.json')}
save('baseline_current.json',{'guarded':guard,'product':m['product'],'new_pair':EXPECTED,'checks':checks,'semantic':semantic,'reported_targeted':{'total':len(tests),'passed':16},'reported_unit_files':list(unittexts),'manifest_last_by_mtime':True})
oldpair=ROOT/'DEV/Stage6_Выход/stage6_akz_part1_20261009_085822_release3_TEST_9_2026'
assert sha(oldpair/ACTNAME)=='1a52736427e1529420723ff85f3bfd92340865ece895173ffaaa0dadb48ba06f'
assert sha(oldpair/APPNAME)=='4cd657f84311512255e6a6e39b492c227f18df3ab8b5ab03eb34400f165d9c6d'
result={}
for name in (ACTNAME,APPNAME):
    old=load_workbook(oldpair/name);new=load_workbook(PAIR/name);cache=load_workbook(PAIR/name,data_only=True)
    try:
        x=old.active;y=new.active;z=cache.active
        assert x.max_row==y.max_row and x.max_column==y.max_column
        assert list(map(str,x.merged_cells.ranges))==list(map(str,y.merged_cells.ranges))
        assert str(x.print_area)==str(y.print_area)
        assert str(x.oddHeader)==str(y.oddHeader) and str(x.oddFooter)==str(y.oddFooter)
        for row in x:
            for cell in row:
                other=y[cell.coordinate]
                assert cell.value==other.value,(name,cell.coordinate,cell.value,other.value)
                for attribute in ('border','font','fill','alignment','protection','number_format'):
                    assert copy.copy(getattr(cell,attribute))==copy.copy(getattr(other,attribute)),(name,cell.coordinate,attribute)
        assert {k:v.width for k,v in x.column_dimensions.items()}=={k:v.width for k,v in y.column_dimensions.items()}
        r={'all_cell_values_formulas_styles_equal_to_old':True,'merges':[str(v) for v in y.merged_cells.ranges],'print_area':str(y.print_area),'header':str(y.oddHeader),'footer':str(y.oddFooter)}
        if name==APPNAME:
            hidden={str(i):bool(y.row_dimensions[i].hidden) for i in range(1,22)}
            assert all(hidden[str(i)]==(7<=i<=20) for i in range(1,22))
            assert all(abs(Decimal(str(z[f'{c}21'].value))-Decimal(expected))<Decimal('0.00000001') for c,expected in zip('CEG',('3','105.26','3.989')))
            assert [y[f'{c}21'].value for c in 'CEG']==[f'=SUM({c}4:{c}20)' for c in 'CEG']
            r.update(hidden=hidden,values=list(z.iter_rows(min_row=4,max_row=6,max_col=7,values_only=True)),totals=[z[f'{c}21'].value for c in 'CEG'],title=y['A1'].value)
        else:
            from akz_stage6 import core as c
            assert all(z[cell].value==expected for cell,expected in c.FACTUAL_PLACEHOLDERS.items())
            assert z['E4'].value=='9/2026 от ' and z['H4'].value.date().isoformat()=='2026-10-05'
            assert z['E17'].value==1 and 'Pirocor Protect 057' in z['E18'].value and z['E21'].value=='RAL 7004' and z['E25'].value==80
            r.update(number=z['E4'].value,date=z['H4'].value,coating=[z[v].value for v in ('E17','E18','E21','E25')],factual_placeholders={cell:z[cell].value for cell in c.FACTUAL_CELLS})
        result[name]=r
    finally:old.close();new.close();cache.close()
save('stored_exact_pair.json',result);save('initial_audit.json',AUDIT)
print(json.dumps({'checks':checks,'semantic':semantic,'stored_pair':'PASS','reported_tests':'16 targeted / 16 unit, authenticated evidence'},ensure_ascii=False,indent=2))
