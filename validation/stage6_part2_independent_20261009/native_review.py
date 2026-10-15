from review import *
import win32com.client
from openpyxl import load_workbook
m=json.loads((IMPL/'REVIEW_MANIFEST.json').read_text(encoding='utf-8'))
app=win32com.client.DispatchEx('Excel.Application')
app.AutomationSecurity=3;app.EnableEvents=False;app.AskToUpdateLinks=False;app.DisplayAlerts=True
records=[];book=None
try:
    for x in m['six_outputs']:
        path=Path(x['path']);assert sha(path)==x['sha256']
        book=app.Workbooks.Open(Filename=str(path),ReadOnly=True,UpdateLinks=0,CorruptLoad=0,AddToMru=False,Notify=False,IgnoreReadOnlyRecommended=True)
        assert book.ReadOnly and not book.LinkSources(1) and book.Connections.Count==0
        app.CalculateFullRebuild()
        data={'path':str(path),'sha_before':sha(path),'variant':x['variant'],'kind':x['kind'],'normal_load':True,'ReadOnly':bool(book.ReadOnly),'AutomationSecurity':app.AutomationSecurity,'EnableEvents':app.EnableEvents,'UpdateLinks':0,'CorruptLoad':0,'excel_version':app.Version,'sheets':[]}
        for s in book.Worksheets:
            u=s.UsedRange;values=u.Value2
            cells={}
            if isinstance(values,tuple):
                from openpyxl.utils import get_column_letter
                for i,row in enumerate(values,u.Row):
                    for j,value in enumerate(row,u.Column):
                        if value is not None:cells[f'{get_column_letter(j)}{i}']=value
            data['sheets'].append({'name':s.Name,'visible':s.Visible,'cells_after_full_rebuild':cells,'pages':int(s.PageSetup.Pages.Count) if s.Visible==-1 else None,'print_area':s.PageSetup.PrintArea,'querytables':s.QueryTables.Count})
        n=x['variant'];s=book.Worksheets(1)
        if x['kind']=='act':
            assert s.Range('E4').Value2==f'9/2026-{n} от ' and s.Range('H4').Value2==46300
            assert int(s.PageSetup.Pages.Count)==1
            assert all(sheet['visible']==2 for sheet in data['sheets'][1:])
            raw=load_workbook(path,data_only=False)
            for sh in raw:
                actual=next(y for y in data['sheets'] if y['name']==sh.title)['cells_after_full_rebuild']
                for row in sh:
                    for cell in row:
                        if cell.value in ('НЕТ ДАННЫХ / ТЕСТ','?'):assert actual[cell.coordinate]==cell.value
            raw.close()
        else:
            total={1:66,2:52,3:99}[n];mass={1:37.2,2:31.32,3:36.74}[n];area={1:1.412,2:1.184,3:1.393}[n]
            assert f'9/2026-{n}' in s.Range('A1').Value2 and '05.10.2026' in s.Range('A1').Value2
            assert s.Range(f'C{total}').Value2==1
            assert abs(s.Range(f'E{total}').Value2-mass)<1e-8 and abs(s.Range(f'G{total}').Value2-area)<1e-8
            assert int(s.PageSetup.Pages.Count)==1
        app.Visible=True;book.Activate()
        records.append(data);write('native_exact_six_in_progress.json',records)
        print('READY',n,x['kind'],path.name,flush=True)
        command=input()
        assert command=='next',command
        book.Close(SaveChanges=False);book=None
        data['sha_after']=sha(path);assert data['sha_after']==x['sha256'];data['unchanged']=True
        write('native_exact_six_in_progress.json',records)
    write('native_exact_six.json',records)
    print('NATIVE_EXACT_SIX_PASS',flush=True)
finally:
    if book is not None:book.Close(SaveChanges=False)
    app.Quit()
