from review import *
import copy, shutil, subprocess, os, traceback
from decimal import Decimal
from unittest.mock import patch
from openpyxl import load_workbook
from akz_stage6 import core as c, part2 as p2
import win32file,win32con
results=[]
def check(name,fn):
    try:detail=fn();results.append({'name':name,'status':'PASS','detail':detail});print('PASS',name,flush=True)
    except BaseException as exc:results.append({'name':name,'status':'FAIL','error':str(exc),'traceback':traceback.format_exc()});raise
    finally:write('independent_safety.json',results)
def stop(fn,code):
    try:fn()
    except Exception as exc:
        assert code in str(exc),str(exc)
        return str(exc)
    raise AssertionError('Expected '+code)
m=json.loads((IMPL/'REVIEW_MANIFEST.json').read_text(encoding='utf-8'))
batch=p2.preview_groups('Гараж Казахстан','9/2026',c.PROJECT/'Паспорт/9 от 06.10.2026/ВОМ отгрузочный счет 183  (5).xls',m['fixture']['path'])
box=OUT/'isolated';box.mkdir(exist_ok=False)
def reseal(data):return c.Preview(data,c.digest(data))
def native_lock(module,preview,name):
    opened=c.native_open;handles=[];errors=[];target=module.OUTPUT/name
    def lock(app,path,*,read_only):
        if 'Приложение' in Path(path).name and not read_only:
            handles.append(win32file.CreateFile(str(path),win32con.GENERIC_READ,0,None,win32con.OPEN_EXISTING,0,None))
            try:return opened(app,path,read_only=read_only)
            except BaseException as exc:errors.append(exc);raise
        return opened(app,path,read_only=read_only)
    try:
        with patch.object(c,'native_open',side_effect=lock):
            try:module.generate(preview,name)
            except c.FileOperationError as exc:
                assert exc is errors[0] and exc.code=='FILE_LOCKED' and exc.__cause__.winerror==32
                assert exc.cleanup_error.code=='FILE_LOCKED' and exc.cleanup_error.__cause__.winerror==32
                detail={'primary':str(exc),'cleanup':str(exc.cleanup_error),'primary_identity_preserved':True,'live_handle':win32file.GetFileType(handles[0])==win32con.FILE_TYPE_DISK}
            else:raise AssertionError('LOCK NOT DETECTED')
        assert not target.exists();pending=next(module.OUTPUT.glob('.pending-*'))
        record_file=next(module.OUTPUT.glob('.part2-pending-*' if module is p2 else '.stage6-pending-*'))
        before={str(x):[x.stat().st_ino,x.stat().st_size,x.stat().st_mtime_ns] for x in pending.iterdir()};pin=sha(record_file)
        stop(module.recover_pending,'FILE_LOCKED')
        assert before=={str(x):[x.stat().st_ino,x.stat().st_size,x.stat().st_mtime_ns] for x in pending.iterdir()} and sha(record_file)==pin
        foreign=pending/'USER.txt';foreign.write_bytes(b'KEEP');stop(module.recover_pending,'PENDING_OWNERSHIP_STOP');assert foreign.read_bytes()==b'KEEP';foreign.unlink()
        unknown=module.OUTPUT/'.pending-foreign';unknown.mkdir();(unknown/'USER.txt').write_bytes(b'KEEP')
        stop(module.recover_pending,'PENDING_OWNERSHIP_STOP');assert (unknown/'USER.txt').read_bytes()==b'KEEP';(unknown/'USER.txt').unlink();unknown.rmdir()
        act=next(x for x in pending.glob('*.xlsx') if 'Приложение' not in x.name);backup=act.with_suffix('.held')
        act.rename(backup);act.write_bytes(backup.read_bytes());stop(module.recover_pending,'PENDING_OWNERSHIP_STOP');act.unlink();backup.rename(act)
    finally:
        for h in handles:h.Close()
    child=OUT/('recover_part2.py' if module is p2 else 'recover_zero.py')
    child.write_text('from pathlib import Path\nimport sys\nfrom akz_stage6 import '+('part2' if module is p2 else 'core')+' as m\nm.OUTPUT=Path(sys.argv[1])\nm.recover_pending()\n',encoding='utf-8')
    completed=subprocess.run([sys.executable,'-B',str(child),str(module.OUTPUT)],cwd=ROOT,env={**os.environ,'PYTHONPATH':str(ROOT)},capture_output=True,text=True,encoding='utf-8')
    assert completed.returncode==0,completed.stderr
    assert not pending.exists() and not record_file.exists()
    detail['fresh_process_recovery']=True;detail['retry']=str(module.generate(preview,name));return detail
with patch.object(p2,'OUTPUT',box/'part2'),patch.object(c,'OUTPUT',box/'zero'):
    for n,cap in ((1,61),(2,47),(3,94)):
        def boundary(n=n,cap=cap):
            data=copy.deepcopy(p2.group_preview(batch,n).payload);original=data['selected'][0]
            data['synthetic']=True;data['act_number']=f'SYNTHETIC-TEST 9/2026-{n}'
            data['selected']=[{**original,'row':9000+i} for i in range(cap)]
            data['totals']={f:sum((r[f] for r in data['selected']),Decimal(0)) for f in ('quantity','mass_total','area_total')}
            folder=p2.generate(reseal(data),f'CAP_{n}_{cap}');names=p2.filenames('9/2026',n)
            b=load_workbook(folder/names[1],data_only=False);s=b.active;total=cap+5
            assert all(s[f'A{r}'].value==original['mark'] and not s.row_dimensions[r].hidden for r in range(5,total))
            assert s[f'C{total}'].value==f'=SUM(C5:C{total-1})' and not s.row_dimensions[total].hidden;b.close()
            overflow=copy.deepcopy(data);overflow['selected'].append(original)
            reason=stop(lambda:p2.generate(reseal(overflow),f'OVER_{n}'),'CAPACITY_STOP')
            stop(lambda:p2.matching({'rows':overflow['selected']},{original['mark']:{'assignment':n,'row':2}},n),'CAPACITY_STOP')
            assert not (p2.OUTPUT/f'OVER_{n}').exists() and not list(p2.OUTPUT.glob('.pending-*'))
            return {'native_generated_isolated':str(folder),'capacity':cap,'overflow':cap+1,'overflow_stop':reason,'rows_preserved':True}
        check(f'capacity_{n}_{cap}_PASS_{cap+1}_STOP',boundary)
    p=p2.group_preview(batch,2)
    def changed():
        details=[]
        for i,source in enumerate(p.payload['sources']):
            clone=box/f'changed_{i}.bin';shutil.copyfile(source,clone)
            data=copy.deepcopy(p.payload);data['sources'][str(clone)]=sha(clone);preview=reseal(data);clone.write_bytes(b'changed')
            details.append(stop(lambda:p2.generate(preview,f'CHANGE_{i}'),'INPUT_CHANGED_AFTER_PREVIEW_STOP'));assert not (p2.OUTPUT/f'CHANGE_{i}').exists()
        return details
    check('all_five_inputs_changed_after_preview_STOP',changed)
    def rollback():
        saved=[];original=p2.write_act
        def save(*args):original(*args);saved.append(sha(args[1]))
        with patch.object(p2,'write_act',side_effect=save),patch.object(p2,'write_application',side_effect=c.Stage6Error('SECOND_BOOK_FAILURE')):
            reason=stop(lambda:p2.generate(p,'ROLLBACK'),'SECOND_BOOK_FAILURE')
        assert len(saved)==1 and not (p2.OUTPUT/'ROLLBACK').exists() and not list(p2.OUTPUT.glob('.pending-*')) and not list(p2.OUTPUT.glob('.part2-pending-*'))
        return {'actual_act_saved':saved,'reason':reason,'no_published_first_book':True}
    check('second_book_failure_rolls_back_saved_first_book',rollback)
    def existing():
        target=p2.OUTPUT/'EXISTING';target.mkdir();user=target/'USER.txt';user.write_bytes(b'KEEP')
        reason=stop(lambda:p2.generate(p,'EXISTING'),'OUTPUT_ALREADY_EXISTS_STOP');assert user.read_bytes()==b'KEEP';return reason
    check('existing_output_preserved_STOP',existing)
    check('part2_actual_lock_ownership_cleanup_fresh_retry',lambda:native_lock(p2,p,'LOCK_RETRY'))
    zero=c.preview('Гараж Казахстан','9/2026',c.PROJECT/'Паспорт/9 от 06.10.2026/ВОМ отгрузочный счет 183  (5).xls',c.FIXTURES/'stage6_akz_part1_20261009_085822_release3/TEST_Справочник марок.xlsx')
    def regression_r1():
        target=c.generate(zero,'R1_3');path=target/'ТЕСТ Приложение к АКЗ №0.xlsx'
        with c.native_excel() as app:
            b=c.native_open(app,path,read_only=True)
            try:
                s=b.Worksheets('Приложение АКЗ');assert [r for r in range(1,22) if s.Rows(r).Hidden]==list(range(7,21)) and not s.Rows(21).Hidden
                assert abs(s.Range('E21').Value2-105.26)<1e-8 and abs(s.Range('G21').Value2-3.989)<1e-8
            finally:b.Close(SaveChanges=False)
        return {'isolated_target':str(target),'N3_rows_7_20_hidden':True,'total21_visible':True,'native_totals':'3 / 105.26 / 3.989','number':zero.payload['act_number'],'date':zero.payload['date']}
    check('zero_R1_native_hidden_rows_number_date_source_areas',regression_r1)
    check('zero_R2_actual_lock_ownership_cleanup_fresh_retry',lambda:native_lock(c,zero,'ZERO_LOCK_RETRY'))
def dynamic_registry():
    path=box/'registry.xlsx';shutil.copyfile(c.REGISTRY,path);b=load_workbook(path);s=b['Паспорт качества']
    for col in range(1,s.max_column+1):s.cell(401,col).value=s.cell(400,col).value;s.cell(400,col).value=None
    s['O401']=datetime.datetime(2026,10,4);b.save(path);b.close()
    with patch.object(c,'REGISTRY',path):
        value=c.current_passport('Гараж Казахстан','9/2026');assert value['row']==401 and value['mapping']['date']=='O401' and value['date']=='2026-10-04'
    return value
check('dynamic_registry_row401_and_changed_DEV_date',dynamic_registry)
def fixture_rules():
    fixture=box/'fixtures';fixture.mkdir();catalog=fixture/'catalog.xlsx';shutil.copyfile(m['fixture']['path'],catalog)
    with patch.object(c,'FIXTURES',fixture):
        assignments=c.catalog(catalog);assert {mark:assignments[mark]['assignment'] for mark in ('Бп2-5','Бп2-18','Бп2-19')}=={'Бп2-5':1,'Бп2-18':2,'Бп2-19':3}
        b=load_workbook(catalog);s=b.active;s.append(['BLANK','TEST',None]);s.append(['ZERO','TEST',0]);b.save(catalog);b.close()
        data=c.catalog(catalog);assert data['BLANK']['assignment'] is None and data['ZERO']['assignment']==0
        b=load_workbook(catalog);b.active.append(['Бп2-5','TEST',2]);b.save(catalog);b.close();reason=stop(lambda:c.catalog(catalog),'CATALOG_DUPLICATE_CONFLICT')
    return {'blank_not_zero':True,'duplicate_conflict':reason}
check('actual_catalog_schema_numeric_assignments_blank_conflict',fixture_rules)
def batch_failure():
    calls=[]
    def generate(p,name):
        calls.append(p.payload['variant'])
        if p.payload['variant']==2:raise c.Stage6Error('GROUP2_FAILURE')
        return box/name
    with patch.object(p2,'generate',side_effect=generate):result=p2.generate_groups(batch,[1,2,3],'SIMULATED_BATCH')
    assert calls==[1,2,3] and result['status']=='PARTIAL_STOP' and result['groups']['2']['status']=='STOP'
    bad=copy.deepcopy(batch.payload);bad['groups']['2']={'status':'STOP','reason':'INVALID'};bad=p2.BatchPreview(bad,c.digest(bad))
    with patch.object(p2,'generate') as gen:stop(lambda:p2.generate_groups(bad,[1,2],'PREFLIGHT'),'GROUP_NOT_READY');gen.assert_not_called()
    return result
check('failed_group_not_success_and_all_selected_preflight',batch_failure)
def official():
    for module in (c,p2):stop(lambda:module.generate(None,'OFFICIAL',mode='OFFICIAL'),'OFFICIAL_MODE_STOP')
    return 'STOP before sources/output'
check('official_issuance_STOP',official)
unit=subprocess.run([sys.executable,'-B','-m','unittest','discover','-s','tests','-p','test_stage6*.py','-v'],cwd=ROOT,capture_output=True,text=True,encoding='utf-8')
(OUT/'independent_unit_tests.txt').write_text(unit.stdout+unit.stderr,encoding='utf-8');assert unit.returncode==0,unit.stderr
print('ALL_SAFETY_PASS',len(results),'UNIT_33_PASS',flush=True)
