import sys, json, hashlib, datetime
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path.cwd(); OUT=ROOT/'validation/stage6_part2_independent_20261009'
sys.path.insert(0,str(ROOT))
IMPL=ROOT/'validation/stage6_part2_20261009_132204'
def guard(event,args):
    if event=='open' and isinstance(args[0],(str,bytes)):
        p=Path(args[0].decode() if isinstance(args[0],bytes) else args[0])
        if p.name.casefold() in ('взмк.xlsx','узк.xlsm'): raise RuntimeError('PRODUCTION_READ_FORBIDDEN')
    if event in ('socket.connect','socket.getaddrinfo'): raise RuntimeError('NETWORK_FORBIDDEN')
sys.addaudithook(guard)
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(name,data): (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def gate():
    m=json.loads((IMPL/'REVIEW_MANIFEST.json').read_text(encoding='utf-8'))
    result={'utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'product':{},'sources':{},'outputs':m['six_outputs'],'evidence_errors':[]}
    for p,h in m['product_current_sha256'].items():
        actual=sha(p);result['product'][p]={'expected':h,'actual':actual,'match':actual==h}
    for key in ('tested_before_sha256','tested_after_sha256'):
        assert all(h==result['product'][p]['actual'] for p,h in m[key].items()),key
    assert all(x['match'] for x in result['product'].values()),'PRODUCT_MISMATCH_STOP'
    for p,h in m['source_after_sha256'].items():
        actual=sha(p) if Path(p).is_file() else 'MISSING'
        result['sources'][p]={'expected':h,'actual':actual,'match':actual==h}
    for p,r in m['evidence_files'].items():
        if sha(p)!=r['sha256']:result['evidence_errors'].append({'path':p,'error':'SHA'})
        if Path(p).stat().st_mtime_ns> (IMPL/'REVIEW_MANIFEST.json').stat().st_mtime_ns:result['evidence_errors'].append({'path':p,'error':'AFTER_MANIFEST'})
    write('initial_gate.json',result)
    print('PRODUCT',len(result['product']),'MATCH; SOURCE DRIFT',json.dumps({p:v for p,v in result['sources'].items() if not v['match']},ensure_ascii=False))
    print('OUTPUTS',json.dumps(m['six_outputs'],ensure_ascii=False));print('FIXTURE',json.dumps(m['fixture'],ensure_ascii=False))
    print('EVIDENCE_ERRORS',result['evidence_errors']);print('AUTHORIZED',m['authorized_part2_changes'])
    print('KEYS',list(m['reused_final_evidence']),list(m['validation']))
def baseline():
    m=json.loads((IMPL/'REVIEW_MANIFEST.json').read_text(encoding='utf-8'))
    paths=set(m['source_after_sha256'])|set(m['product_current_sha256'])
    import os
    paths.update(m['evidence_files']);paths.add(str(IMPL/'REVIEW_MANIFEST.json'))
    for base,dirs,files in os.walk(ROOT):
        dirs[:]=[d for d in dirs if d not in ('node_modules','.codebase-index','__pycache__','.git','validation','DEV','audit','e2e_user_stage1','e2e_user_stage2','e2e_user_stage3') and Path(base)/d!=OUT]
        for name in files:
            if not name.startswith('~$') and name.casefold() not in ('взмк.xlsx','узк.xlsм'):paths.add(str(Path(base)/name))
    data={p:sha(p) for p in sorted(paths)};write('protected_before.json',data)
    assert all(sha(x['path'])==x['sha256'] for x in m['six_outputs'])
    before=m['source_before_sha256'];after=m['source_after_sha256']
    changed={p:{'before':before[p],'after':h} for p,h in after.items() if p in before and h!=before[p]}
    added_product=[p for p in paths if Path(p).suffix=='.py' and p not in before and not Path(p).is_relative_to(ROOT/'validation')]
    write('actual_changes.json',{'changed_existing':changed,'product_py_not_in_before':added_product,'protected_count':len(data)})
    print('BASELINE',len(data),'CHANGED',json.dumps(changed,ensure_ascii=False),'NEW',added_product)
def static():
    from akz_stage6 import core as c,part2 as p2
    from openpyxl import load_workbook
    m=json.loads((IMPL/'REVIEW_MANIFEST.json').read_text(encoding='utf-8'))
    catalog=Path(m['fixture']['path']);vom=c.PROJECT/'Паспорт/9 от 06.10.2026/ВОМ отгрузочный счет 183  (5).xls'
    batch=p2.preview_groups('Гараж Казахстан','9/2026',vom,catalog);write('independent_preview.json',batch.payload)
    assert len(batch.payload['groups']['1']['payload']['vom']['rows'])==104
    assert len(batch.payload['groups']['1']['payload']['vom']['hidden_rows'])==397
    assert batch.payload['date']=='2026-10-05'
    expected={1:('Бп2-5','37.2','1.412'),2:('Бп2-18','31.32','1.184'),3:('Бп2-19','36.74','1.393')}
    for n,(mark,mass,area) in expected.items():
        g=p2.group_preview(batch,n).payload;r=g['selected'];assert len(r)==1
        assert r[0]['mark']==mark and r[0]['quantity']==1
        assert float(r[0]['mass_unit'])==float(mass) and float(r[0]['area_unit'])==float(area)
    assert len(batch.payload['climate_days'])==1 and len({c.digest(p2.group_preview(batch,n).payload['climate']) for n in (1,2,3)})==1
    dumps=[];diffs=[]
    contract=json.loads((ROOT/'validation/stage6_new_templates_audit_20261009_125930/IMPLEMENTATION_CONTRACT.json').read_text(encoding='utf-8'))
    for x in m['six_outputs']:
        assert sha(x['path'])==x['sha256'];n=x['variant'];v=contract['variants'][str(n)];kind=x['kind']
        source=v['act_path'] if kind=='act' else v['application']['path']
        assert sha(source)==(v['act_sha256'] if kind=='act' else v['application']['sha256'])
        a=load_workbook(x['path'],data_only=False);sbook=load_workbook(source,data_only=False)
        record={'path':x['path'],'sha256':sha(x['path']),'sheets':{}}
        for s in a:
            cells={cell.coordinate:cell.value for row in s for cell in row if cell.value is not None}
            record['sheets'][s.title]={'visibility':s.sheet_state,'cells':cells,'print_area':str(s.print_area),'merges':sorted(map(str,s.merged_cells.ranges)),'headers':[str(s.oddHeader),str(s.oddFooter)],'formulas':{cell.coordinate:cell.value for row in s for cell in row if cell.data_type=='f'}}
            assert not any('RAND' in str(value).upper() or '#REF!' in str(value).upper() for value in cells.values())
            if s.title in sbook:
                src=sbook[s.title];same={cell.coordinate:cell.value for row in src for cell in row if cell.value is not None and cells.get(cell.coordinate)==cell.value}
                diffs.append({'variant':n,'kind':kind,'sheet':s.title,'unchanged_nonempty':same})
        s=a['4'] if kind=='act' else a.active
        if kind=='act':
            assert s['E4'].value==f'9/2026-{n} от ' and s['H4'].value.date().isoformat()=='2026-10-05'
            assert set(record['sheets']['4']['formulas'])=={'G8','I8'}
            for layer in v['layers']:
                for field in ('material','color'):assert s[layer[field]['cell']].value==layer[field]['value']
                assert s[layer['required_thickness']['cell']].value==layer['required_thickness']['raw']
                for cell in layer['historical_fields'].values():assert s[cell].value=='НЕТ ДАННЫХ / ТЕСТ',cell
            if n==2:assert a['Лист1'].sheet_state=='veryHidden' and set(record['sheets']['Лист1']['cells'])=={'A1'}
        else:
            last=v['application']['body_last'];total=last+1;mark,mass,area=expected[n]
            assert s['A5'].value==mark and s['C5'].value==1
            assert s['D5'].value==s['E5'].value==float(mass) and s['F5'].value==s['G5'].value==float(area)
            assert all(s.cell(r,col).value is None for r in range(6,last+1) for col in range(1,8))
            assert all(s.row_dimensions[r].hidden for r in range(6,last+1)) and not s.row_dimensions[total].hidden
            assert len(record['sheets'][s.title]['formulas'])==3
            for col in 'CEG':assert s[f'{col}{total}'].value==f'=SUM({col}5:{col}{last})'
            assert f'9/2026-{n}' in s['A1'].value and '05.10.2026' in s['A1'].value
            assert sorted(map(str,s.merged_cells.ranges))==sorted(map(str,sbook.active.merged_cells.ranges))
        dumps.append(record);a.close();sbook.close()
    write('workbook_dump.json',dumps);write('unchanged_source_cells.json',diffs)
    print('STATIC SIX PASS; VOM 104/397; groups',expected)
if __name__=='__main__':
    {'gate':gate,'baseline':baseline,'static':static}[sys.argv[1] if len(sys.argv)>1 else 'gate']()
