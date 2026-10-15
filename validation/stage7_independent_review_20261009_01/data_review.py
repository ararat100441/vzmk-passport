import sys
sys.dont_write_bytecode=True
import json,hashlib,re
from pathlib import Path
from datetime import date,datetime
from decimal import Decimal
from collections import defaultdict
from zipfile import ZipFile
from lxml import etree as E
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT))
from passport_app.registry import read_registry,select_project
from profiles_stage4.storage import load_catalog
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def digest(x):return hashlib.sha256(json.dumps(x,ensure_ascii=False,sort_keys=True).encode('utf8')).hexdigest()
def save(n,d):(OUT/n).write_text(json.dumps(d,ensure_ascii=False,indent=2,default=str),encoding='utf8')
m=read(ROOT/'validation/stage7_resume_20261009_175000/ASSEMBLY_01/STAGE7_SELECTED_TEST_MANIFEST.json')
pre=read(ROOT/'validation/stage7_resume_20261009_175000/ASSEMBLY_01/PREFLIGHT.json')
a=read(OUT/'ACTUAL_HASHES.json');native=read(OUT/'NATIVE_REVIEW.json')
checks=[]
def check(name,actual,expected):
    checks.append({'name':name,'result':'PASS' if actual==expected else 'FAIL','actual':actual,'expected':expected})
registry=read_registry(m['membership']['registry_path']);project=select_project(registry,m['project'])
entries=[r for r in registry.projects[project.name] if r['number']==m['passport']]
check('unique registry number',len(entries),1);row=entries[0]['row']
book=load_workbook(registry.path,read_only=True,data_only=True,keep_links=False)
keys=['number','contract','customer','object','cipher','mass','start_date','end_date','km','kmd_company','steel','vik_date','uzk_date','akz_date','connections']
fields=dict(zip(keys,[book[registry.sheet].cell(row,c).value for c in range(2,17)]));book.close()
for k,v in list(fields.items()):
    if isinstance(v,(date,datetime)):fields[k]=v.strftime('%d.%m.%Y')
fields.update(row=row,sheet=registry.sheet,project=project.name)
check('manifest fields match actual registry',m['fields'],fields)
check('current vs next', [project.last_number,project.next_number],['9/2026','10/2026'])
check('registry row independently selected',row,400)

ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def doc(p):
    with ZipFile(p) as z:return E.fromstring(z.read('word/document.xml'))
def txt(n):return ''.join(n.xpath('.//w:t/text()',namespaces=ns))
q=doc(a['quality']['path']);t=doc(a['template']['path'])
tables=q.findall('w:body/w:tbl',ns);rows=tables[0].findall('w:tr',ns)
def cell(r,c=1,table=0):return txt(tables[table].findall('w:tr',ns)[r-1].findall('w:tc',ns)[c-1])
check('quality number',cell(2,4),'9/2026')
check('quality contract',cell(12,4).strip(),'№ '+fields['contract'])
check('quality customer',cell(13,2),fields['customer'])
check('quality object',cell(15),fields['object'])
check('quality KMD',cell(16,2),'Металлоконструкции по шифру: '+fields['cipher']+'Приложение № 1')
check('quality mass',cell(17,2),'15 822,14 кг.')
check('quality start',cell(18,2).strip(),fields['start_date']+' г.')
check('quality end',cell(19,2).strip(),fields['end_date']+' г.')
check('quality KM',txt(rows[19].findall('w:tc',ns)[0].findall('w:p',ns)[1]),fields['km'])
check('quality KMD developer',cell(22).strip(),fields['kmd_company']+' '+fields['cipher'])
check('quality steel',cell(24,2),fields['steel'])
check('quality date',cell(5,1,1).strip(),fields['end_date']+' г.')
check('quality order constant',cell(12,2),'1')
template_rows=t.findall('w:body/w:tbl',ns)[0].findall('w:tr',ns)
block1=b''.join(E.tostring(x,method='c14n') for x in template_rows[40:49])
block2=b''.join(E.tostring(x,method='c14n') for x in rows[40:49])
check('point12 complete canonical XML',block2,block1)
checks[-1]['actual']=hashlib.sha256(block2).hexdigest();checks[-1]['expected']=hashlib.sha256(block1).hexdigest()
def compact(v):return re.sub(r'\s+','',v.replace('\x07',''))
check('point12 remains page1',all(compact(txt(r)) in compact(native['quality']['page_texts'][0]) for r in rows[40:49]),True)
check('native numbering',[p['label'].rstrip('.') for p in native['quality']['lists'] if p['label'].rstrip('.').isdigit()],[str(i) for i in range(1,15)])
check('signature blank and no seal images',[native['quality']['inline_shapes'],native['quality']['shapes']],[0,0])
check('all other ZIP parts unchanged',read(OUT/'OOXML_REVIEW.json')['changed_parts'],['word/document.xml'])
check('point13 exact selected acts',cell(51),'; '.join(r['display_title'] for r in m['records'][1:5])+'.')
check('point14 exact short selection',txt(rows[53].findall('w:tc',ns)[0].findall('w:p',ns)[0]),'; '.join(r['display_title'] for r in m['records'][1:])+'.')

vom=read(OUT/'NATIVE_EXCEL_VOM.json')
app1=load_workbook(a['appendix1']['path'],read_only=True,data_only=True,keep_links=False);s=app1.active
data=list(s.iter_rows(values_only=True));end=next(i for i,r in enumerate(data) if str(r[0] or '').strip().lower()=='итого:') if any(str(r[0] or '').strip().lower()=='итого:' for r in data) else next(i for i,r in enumerate(data) if str(r[0] or '').strip().lower().startswith('итого'))
body1=[(r[0],r[2],r[4]) for r in data[4:end] if r[0] is not None]
def aggregate(seq):
    d=defaultdict(lambda:[Decimal(0),Decimal(0)])
    for mark,q,mass in seq:
        key=str(mark).strip().casefold();d[key][0]+=Decimal(str(q));d[key][1]+=Decimal(str(mass))
    return dict(d)
ag1=aggregate(body1);agv=aggregate((r['mark'],r['quantity'],r['mass_kg']) for r in vom['rows'])
check('appendix1 body rows',len(body1),123)
check('appendix1 current number date object',all(v in str(s[cell].value) for cell,v in [('A1','9/2026'),('A1','06.10.2026'),('A2',fields['object'])]),True)
check('appendix1 native VOM per mark',set(ag1)==set(agv) and all(ag1[k][0]==agv[k][0] and abs(ag1[k][1]-agv[k][1])<=Decimal('.005') for k in agv),True)
check('appendix1 total to 0.01 kg',[data[end][2],Decimal(str(data[end][4])).quantize(Decimal('.01'))],[288,Decimal('15822.14')]);app1.close()

stockbook=load_workbook(ROOT/'DEV/УЗК_DEV.xlsm',read_only=True,data_only=True,keep_links=False)
stock=defaultdict(set)
for proj,mark,qty in stockbook['УЗК'].iter_rows(min_row=2,max_col=3,values_only=True):
    if str(proj or '').strip().casefold()==m['project'].casefold():stock[str(mark).strip().casefold()].add(int(qty) if qty not in (None,'') else 1)
stockbook.close()
raw_mark_path=Path(pre['preview']['mark_catalog']);markbook=load_workbook(raw_mark_path,read_only=True,data_only=True,keep_links=False)
mapping={str(r[0]).strip().casefold():r[1] for r in markbook.active.iter_rows(values_only=True) if r[0] is not None};markbook.close()
catalog=load_catalog(ROOT/'DEV/Справочники/Профили.xlsx')
for p in pre['preview']['marks']:
    key=p['mark'].strip().casefold();check('UZK '+p['mark']+' current DEV remaining',sorted(stock[key]),[p['remaining']])
    check('UZK '+p['mark']+' Q from current sources',min(int(agv[key][0]),next(iter(stock[key]))),p['Q'])
    check('UZK '+p['mark']+' current mark profile',mapping[key],p['raw_profile'])
    lookup=catalog.lookup(mapping[key]);g=lookup['profile']['geometry'];family=lookup['profile']['family']
    check('UZK '+p['mark']+' current geometry',g,p['lookup']['profile']['geometry'])
    if family=='Швеллер':segments=[(Decimal(str(g['h'])),Decimal(str(g['s']))),(2*Decimal(str(g['b'])),Decimal(str(g['t'])))]
    elif family in ('Квадратный профиль','Прямоугольный профиль','Гнутый швеллер'):segments=[(2*(Decimal(str(g['h']))+Decimal(str(g['b']))),Decimal(str(g['t'])))]
    else:raise AssertionError('Unreviewed family '+family)
    check('UZK '+p['mark']+' independent geometry from current catalog',[(l*p['Q'],s) for l,s in segments],[(Decimal(str(x['L'])),Decimal(str(x['S']))) for x in p['segments']])

vik=native['vik']['page_texts']
for i,num in enumerate(['9/2026','9/2026-1']):
    check('VIK page '+str(i+1)+' number/date',num in vik[i] and '29.09.2026' in vik[i],True)
check('VIK current connections',fields['connections'] in vik[1],True)
check('VIK object/cipher',fields['object'] in vik[1] and fields['cipher'] in vik[0],True)
vr=read(m['records'][1]['producer_receipt']['path']);vv=read(vr['validation']['path'])['checks'];vg=read(vv['generation_record'])
check('VIK original registry SHA',vg['registry_sha256'],sha(registry.path))
check('VIK private exact registry',sha(ROOT/'validation/stage7_resume_20261009_175000/ASSEMBLY_01/VIK_ISOLATION/DEV_registry.xlsx'),sha(registry.path))
check('VIK source template',vg['template_sha256'],sha(ROOT.parent/'Гараж казахстан/Паспорт/1_2026 - ВИК.docx'))

uz=doc(a['uzk']['path']);uzrows=uz.findall('w:body/w:tbl',ns)[0].findall('w:tr',ns)
expected={'Бп2-1':(8,'L=1280х5 ммL=1024х8.4 мм','55','7070'),
          'Пс1-9':(9,'L=3600х4 мм','5','70'),'Пс1-10':(15,'L=6000х4 мм','5','70'),'Пс2-16':(1,'L=480х3 мм','5','70')}
actual={}
for r in uzrows:
    cells=r.findall('w:tc',ns);mark=txt(cells[0]) if cells else ''
    if mark in expected:
        actual[mark]=[txt(x) for x in cells]
        check('UZK '+mark+' geometry/frequency/angle/date/results',[txt(cells[i]) for i in [1,3,4,5,6,7]],
            [expected[mark][1],'01.10.2026',expected[mark][2],expected[mark][3],'ДНО','годен'])
check('UZK exact mark set',sorted(actual),sorted(expected))
for mark,v in expected.items():
    p=next(p for p in pre['preview']['marks'] if p['mark']==mark)
    check('UZK '+mark+' actual Q',p['Q'],v[0])
    for segment,unit in zip(p['segments'],p['single']['rows']):
        check('UZK '+mark+' independent L x Q',Decimal(str(segment['L'])),Decimal(str(unit['L']))*p['Q'])
check('UZK all units',sum(expected[k][0] for k in actual),33)

akz=load_workbook(a['akz0']['path'],read_only=True,data_only=True,keep_links=False)
s=akz.active;akz_cells={n:s[n].value for n in ['B3','E4','H4','E17','E18','E21','E25','E26','E27','E28','B43']};akz.close()
check('AKZ number/date',[akz_cells['E4'].strip(),akz_cells['H4'].strftime('%d.%m.%Y')],['9/2026 от','05.10.2026'])
check('AKZ layer/material/thickness',[akz_cells['E17'],'Pirocor Protect 057' in akz_cells['E18'],akz_cells['E25']],[1,True,80])
check('AKZ TEST and no actual measurements',['ТЕСТ' in akz_cells['B3'],all('НЕТ ДАННЫХ / ТЕСТ' in akz_cells[k] for k in ['E26','E27','E28'])],[True,True])
app=load_workbook(a['akz0-appendix']['path'],read_only=True,data_only=True,keep_links=False);s=app.active
body=[[s.cell(r,c).value for c in range(1,8)] for r in range(4,21) if s.cell(r,1).value];head=s['A1'].value;total=s['C21'].value;app.close()
check('AKZ exact 3 marks',[r[0] for r in body],['Бп2-5','Бп2-18','Бп2-19'])
check('AKZ 3 units',[sum(r[2] for r in body),total],[3,3])
check('AKZ mass subtotal',sum(Decimal(str(r[4])) for r in body),Decimal('105.26'))
check('AKZ current visible membership',all(int(r[2])==agv[r[0].casefold()][0] and abs(Decimal(str(r[4]))-agv[r[0].casefold()][1])<=Decimal('.005') for r in body),True)
check('AKZ matching appendix header','Приложение № 9/2026 к акту АКЗ № 9/2026 от 05.10.2026' in head,True)

chains=[]
result=read(ROOT/'validation/stage7_resume_20261009_175000/ASSEMBLY_01/ASSEMBLY_RESULT.json')
for key in ['uzk_receipt','vik_receipt','quality_receipt']:
    ref=result[key];rc=read(ref['path']);start=read(rc['start']['path']);res=read(rc['result']['path']);val=read(rc['validation']['path'])
    for label,v in [('receipt',ref),('start',rc['start']),('result',rc['result']),('validation',rc['validation'])]:
        check(key+' '+label+' SHA',sha(v['path']),v['sha256'])
    check(key+' output SHA',sha(rc['output']),rc['output_sha256'])
    check(key+' identity chain',all(all(v[k]==rc[k] for k in ['assembly_run_id','operation_id','generator','status']) for v in [start,res,val]),True)
    check(key+' event order',start['time']<=res['time']<=val['time']<=rc['time'],True)
    check(key+' input/code chain',[start['inputs']==rc['inputs'],start['code']==rc['code']],[True,True])
    check(key+' all current bindings',all(sha(p)==h for p,h in (rc['inputs']|rc['code']).items()),True)
    check(key+' official blocked',[rc['results_confirmed'],rc['official_issue'],rc['state']],[False,False,'COMPLETE_VALIDATED_TEST'])
    chains.append({'operation':key,'generator':rc['generator'],'assembly_run_id':rc['assembly_run_id'],'operation_id':rc['operation_id'],
        'inputs':len(rc['inputs']),'code':len(rc['code']),'start':start['time'],'result':res['time'],'validation':val['time'],'receipt':rc['time']})
save('DATA_REVIEW.json',{'total':len(checks),'passed':sum(c['result']=='PASS' for c in checks),'checks':checks,
    'registry_fields':fields,'registry_project':project.__dict__,'uzk_rows':actual,'akz_cells':akz_cells,'akz_app_rows':body,'receipts':chains})
print(json.dumps({'checks':len(checks),'failed':[c for c in checks if c['result']!='PASS'],'receipts':chains},ensure_ascii=False,default=str))
