"""Read-only evidence collection. Writes only alongside this script."""
from pathlib import Path
import sys, json, hashlib, re, zipfile
from dataclasses import asdict
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT))
from passport_app.registry import read_registry, select_project
from openpyxl import load_workbook
from akz_stage6.core import native_excel, native_open, current_passport
def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):h.update(chunk)
    return h.hexdigest()
def write(name,data): (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
reg=read_registry(ROOT/'DEV/ВЗМК_DEV.xlsx')
project=select_project(reg,'Гараж Казахстан')
folder=ROOT.parent/'Гараж казахстан'/'Паспорт'
inventory=[]
for i in range(1,4):
    for kind,name in [('act',f'АКЗ-{i}.xlsx'),('application',f'Приложение к АКЗ-{i}.xlsx')]:
        path=folder/name
        alternatives=[str(p) for p in folder.iterdir() if p.is_file() and not p.name.startswith('~$') and p.suffix.lower() in ('.xls','.xlsx','.xlsm') and re.sub(r'[\s_–—-]','',p.stem.casefold())==re.sub(r'[\s_–—-]','',Path(name).stem.casefold())]
        status='AMBIGUOUS' if len(alternatives)>1 else ('FOUND' if path.is_file() else 'MISSING')
        inventory.append(dict(variant=i,kind=kind,path=str(path),status=status,candidates=alternatives))
write('selector_inventory.json',dict(registry_sha256=reg.digest,sheet=reg.sheet,columns=reg.columns,project=asdict(project),current=current_passport('Гараж Казахстан','9/2026'),folder=str(folder),files=inventory))
guard=set()
previous=json.loads((ROOT/'validation/stage6_fix_r2_final_20261009_101859/REVIEW_MANIFEST.json').read_text(encoding='utf-8'))
guard.update(Path(p) for p in previous['guarded_current'])
guard.update(p for p in ROOT.iterdir() if p.is_file())
for package in ('passport_app','akz_stage6','profiles_stage4','uzk_stage5','tests','DEV'):
    guard.update(p for p in (ROOT/package).rglob('*') if p.is_file() and not any(t=='__pycache__' or t.startswith('node_modules') for t in p.parts) and not p.name.startswith('~$'))
for name in ('ТЕСТ АКЗ №0.xlsx','ТЕСТ Приложение к АКЗ №0.xlsx','PAIR.json'):
    guard.add(ROOT/'validation/stage6_fix_r1_r2_20261009_094707/NEW_TEST_9_2026'/name)
for p in folder.iterdir():
    if p.is_file() and not p.name.startswith('~$'): guard.add(p)
guard.add(folder.parent/'Справочник марок.xlsx')
guard.add(folder/'9 от 06.10.2026'/'ВОМ отгрузочный счет 183  (5).xls')
print(json.dumps({'phase':'guard','files':len(guard)},ensure_ascii=False),flush=True)
baseline={str(p):sha(p) for p in sorted(guard)}
write('sources_before.json',baseline)
paths=[Path(x['path']) for x in inventory if x['status']=='FOUND']+[folder/'Акт № 4-2026.xlsx',folder/'1_2026 - Приложение к АКЗ.xlsx']
static=[]
for path in paths:
    b=load_workbook(path,data_only=False); cache=load_workbook(path,data_only=True)
    item=dict(path=str(path),sha256=sha(path),sheetnames=b.sheetnames,defined_names={k:str(v) for k,v in b.defined_names.items()},external_links=len(b._external_links),sheets=[])
    for s in b.worksheets:
        cells={c.coordinate:dict(value=c.value,cached=cache[s.title][c.coordinate].value,type=c.data_type,number_format=c.number_format) for row in s for c in row if c.value is not None}
        formulas={k:v for k,v in cells.items() if v['type']=='f'}
        volatile={k:v for k,v in formulas.items() if re.search(r'\b(RAND|RANDBETWEEN|RANDARRAY|NOW|TODAY|OFFSET|INDIRECT|CELL|INFO)\s*\(',str(v['value']),re.I)}
        item['sheets'].append(dict(name=s.title,state=s.sheet_state,dimensions=s.calculate_dimension(),print_area=str(s.print_area),print_title_rows=s.print_title_rows,print_title_cols=s.print_title_cols,merged=[str(m) for m in s.merged_cells.ranges],hidden_rows=[r for r,d in s.row_dimensions.items() if d.hidden],hidden_columns=[c for c,d in s.column_dimensions.items() if d.hidden],page_setup=str(s.page_setup),cells=cells,formulas=formulas,volatile=volatile,errors={k:v for k,v in cells.items() if v['type']=='e'}))
    with zipfile.ZipFile(path) as z:
        item['package_features']=[n for n in z.namelist() if any(t in n.lower() for t in ('externallink','connection','vbaproject','querytable'))]
    static.append(item)
    b.close(); cache.close()
write('static_workbooks.json',static)
lines=[]
for b in static:
    lines.append('\nFILE '+b['path']+'\nSHA '+b['sha256'])
    for s in b['sheets']:
        lines.append(f"SHEET {s['name']} {s['state']} {s['dimensions']} PRINT {s['print_area']}")
        for c,v in s['cells'].items():lines.append(f"{c}: {v['value']}"+(f" [cached {v['cached']}]" if v['type']=='f' else ''))
(OUT/'cell_dump.txt').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps({'inventory':inventory,'guard_count':len(baseline),'static':'DONE'},ensure_ascii=False),flush=True)
native=[]
try:
    with native_excel() as app:
        for x in inventory:
            if x['status']!='FOUND':continue
            p=Path(x['path']); before=sha(p)
            entry=dict(path=str(p),variant=x['variant'],kind=x['kind'],sha_before=before,macros_security=int(app.AutomationSecurity),events=bool(app.EnableEvents),update_links=0,corrupt_load=0,excel_version=app.Version,sheets=[])
            b=native_open(app,p,read_only=True)
            try:
                entry['read_only']=bool(b.ReadOnly)
                entry['connections']=int(b.Connections.Count)
                entry['links']=b.LinkSources(1)
                for s in b.Worksheets:
                    try:pages=int(s.PageSetup.Pages.Count)
                    except Exception as e:pages={'status':'NOT_VERIFIED','reason':str(e)}
                    entry['sheets'].append(dict(name=s.Name,visible=s.Visible,pages=pages,print_area=s.PageSetup.PrintArea,print_title_rows=s.PageSetup.PrintTitleRows,values=s.UsedRange.Value2,formulas=s.UsedRange.Formula))
            finally:b.Close(SaveChanges=False)
            entry['sha_after']=sha(p); entry['unchanged']=before==entry['sha_after']; native.append(entry)
            write('native_excel.json',native)
            print(json.dumps({'native_path':str(p),'read_only':entry['read_only'],'sheets':[{k:v for k,v in s.items() if k not in ('values','formulas')} for s in entry['sheets']],'unchanged':entry['unchanged']},ensure_ascii=False),flush=True)
finally:
    after={p:sha(Path(p)) if Path(p).is_file() else 'MISSING' for p in baseline}
    changes=[dict(path=p,before=h,after=after[p]) for p,h in baseline.items() if h!=after[p]]
    write('sources_after.json',after)
    write('integrity.json',dict(guarded_files=len(baseline),changes=changes,saves=0,prints=0,generated_acts=0,source_mutations=0 if not changes else 'DETECTED',native_files=len(native)))
    print(json.dumps({'guarded_files':len(baseline),'changes':changes},ensure_ascii=False),flush=True)
