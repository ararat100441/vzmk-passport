# Adapted independent run controller; original SHA 9123c9100912e75af6360e37da5721ebbace5f1de26ef40d13230a6465be2471
"""Controlled Office sentinels foreign to the EXE; never opens user files."""
from common import *
import gc,time,pythoncom,win32process,win32event,win32api
from stage9_app.ownership import launch_child,native_application,request_empty_excel_close,process_image
from openpyxl import Workbook
from docx import Document
RUN=OUT/('foreign_office_'+uuid4().hex[:8]);RUN.mkdir();children=[];records=[];apps=[];cases=[]
pythoncom.CoInitialize()
try:
    for progid,suffix in (('Excel.Application','.xlsx'),('Word.Application','.docx')):
        path=RUN/('FOREIGN_SENTINEL'+suffix)
        if suffix=='.xlsx':
            book=Workbook();book.active['A1']='ARTIFICIAL FOREIGN OFFICE SENTINEL';book.save(path);book.close();book=None
        else:
            doc=Document();doc.add_paragraph('ARTIFICIAL FOREIGN OFFICE SENTINEL');doc.save(path);doc=None
        child,record=launch_child(progid,path,os.getpid());children.append(child);records.append(record)
        deadline=time.monotonic()+30
        while True:
            try:
                with native_application(record,child._handle) as native:app=native
                native=None;break
            except RuntimeError as exc:
                if str(exc)!='OWNED_NATIVEOM_WINDOW_MISSING' or time.monotonic()>=deadline:raise
                pythoncom.PumpWaitingMessages();time.sleep(.1)
        app.AutomationSecurity=3;app.DisplayAlerts=False
        if suffix=='.xlsx':app.UserControl=True;record['sentinel_user_control']=bool(app.UserControl)
        apps.append(app);app=None
    def preserved():
        for child,r,app in zip(children,records,apps):
            assert child.poll() is None and win32process.GetProcessTimes(child._handle)['CreationTime'].isoformat()==r['creation_time']
            if r['progid']=='Excel.Application':assert app.Workbooks.Count==1 and app.Workbooks(1).Worksheets(1).Cells(1,1).Value2=='ARTIFICIAL FOREIGN OFFICE SENTINEL'
            else:assert app.Documents.Count==1 and 'ARTIFICIAL FOREIGN OFFICE SENTINEL' in app.Documents(1).Content.Text
    preserved()
    baseline={str(r['pid']):{'creation_time':r['creation_time'],'image':r['image']} for r in records}
    target=RUN/'SENTINEL_IDENTITIES.json';save(target,baseline)
    work=read(OUT/'LIFECYCLE_FINAL.json')['workspace']
    for n in range(2):
        data,sup,path=task(RUN,'office-probe',workspace=work)
        assert set(int(pid) for pid in baseline).isdisjoint(r['pid'] for r in sup['instances'])
        preserved();cases.append({'normal_probe':str(path),'preserved':True})
        if n==0:
            child_env=os.environ.copy();child_env['STAGE9_FOREIGN_OFFICE']=str(target)
            p=subprocess.run([sys.executable,'-B',str(OUT/'ownership_stop.py')],env=child_env,timeout=300,capture_output=True)
            assert p.returncode==0,p.stderr.decode('utf8','replace');preserved()
    save(RUN/'RESULT.json',{'status':'PASS','exe_sha256':sha(EXE),'controlled_foreign_processes':baseline,'preserved_unsaved_content':True,'cases':cases,'ownership_stop':str(OUT/'OWNERSHIP_STOP_FINAL.json')})
    save(OUT/'FOREIGN_OFFICE_FINAL.json',{'status':'PASS','exe_sha256':sha(EXE),'evidence':str(RUN/'RESULT.json')})
finally:
    for app,r in zip(apps,records):
        if r['progid']=='Excel.Application':
            app.DisplayAlerts=False;app.Workbooks.Close();app.UserControl=False;r['workbooks_before_quit']=int(app.Workbooks.Count)
        else:app.Documents.Close(False)
        app.Quit();r['quit_returned']=True
    app=None;apps.clear();gc.collect()
    for child,r in zip(children,records):
        deadline=time.monotonic()+15
        while child.poll() is None and time.monotonic()<deadline:pythoncom.PumpWaitingMessages();time.sleep(.1)
        if child.poll() is None:
            request_empty_excel_close(r,child._handle,lambda phase,data:save(RUN/(phase+'_'+str(r['pid'])+'.json'),data))
            deadline=time.monotonic()+15
            while child.poll() is None and time.monotonic()<deadline:pythoncom.PumpWaitingMessages();time.sleep(.1)
        if child.poll() is None:
            assert win32process.GetProcessId(child._handle)==r['pid'] and win32process.GetProcessTimes(child._handle)['CreationTime'].isoformat()==r['creation_time'] and process_image(child._handle)==r['image']
            save(RUN/('SOURCE_SENTINEL_EMERGENCY_'+str(r['pid'])+'.json'),dict(r,handle_validated=True))
            win32api.TerminateProcess(child._handle,1);child.wait(10)
        r['final_exit_code']=child.returncode
    save(RUN/'SOURCE_SENTINEL_CLEANUP.json',{'instances':records,'owned_orphans':sum(p.poll() is None for p in children)})
    pythoncom.CoUninitialize()
