import json, hashlib, re
from pathlib import Path
from zipfile import ZipFile
from collections import Counter,defaultdict
from xml.etree import ElementTree as ET
import openpyxl
OUT=Path(__file__).resolve().parent
APP=OUT.parent.parent
DEV=APP/'DEV'/'УЗК_DEV.xlsm'
def save(name,x): (OUT/name).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
books=[openpyxl.load_workbook(DEV,read_only=False,data_only=x,keep_vba=True,keep_links=True) for x in (False,True)]
book,values=books
result={'path':str(DEV),'sheets':[],'defined_names':{k:str(v) for k,v in book.defined_names.items()},'security':str(book.security)}
for s in book:
    cells=[c for row in s for c in row if c.value is not None]
    rows=sorted({c.row for c in cells})
    raw=[{'row':r,'values':{openpyxl.utils.get_column_letter(c):s.cell(r,c).value for c in range(1,s.max_column+1) if s.cell(r,c).value is not None},
          'cached':{openpyxl.utils.get_column_letter(c):values[s.title].cell(r,c).value for c in range(1,s.max_column+1) if s.cell(r,c).value is not None}} for r in rows]
    sheet={'name':s.title,'state':s.sheet_state,'dimension':s.calculate_dimension(),'max_row':s.max_row,'max_column':s.max_column,
           'nonempty_rows':len(rows),'nonempty_cells':len(cells),'actual_first_row':rows[0] if rows else None,'actual_last_row':rows[-1] if rows else None,
           'raw_rows':raw,'formulas':{c.coordinate:c.value for c in cells if c.data_type=='f'},
           'tables':{k:dict(v) for k,v in s.tables.items()} if False else [str(v) for v in s.tables.values()],
           'merged_cells':[str(x) for x in s.merged_cells.ranges], 'auto_filter':str(s.auto_filter),
           'validations':str(s.data_validations),'protection':str(s.protection),'local_names':{k:str(v) for k,v in s.defined_names.items()},
           'hidden_rows':[r for r,d in s.row_dimensions.items() if d.hidden],
           'hidden_columns':[c for c,d in s.column_dimensions.items() if d.hidden]}
    result['sheets'].append(sheet)
with ZipFile(DEV) as z:
    names=z.namelist()
    result['vba']={'present':'xl/vbaProject.bin' in names}
    if result['vba']['present']:
        vb=z.read('xl/vbaProject.bin')
        result['vba'].update(size=len(vb),sha256=hashlib.sha256(vb).hexdigest())
        strings=re.findall(rb'[\x20-\x7e]{5,}',vb)
        result['vba']['safe_ascii_strings']=[x.decode('ascii') for x in strings if any(y in x.lower() for y in [b'sub ',b'function ',b'workbook',b'worksheet',b'project',b'module',b'vb_'])][:150]
    result['external_link_parts']=[x for x in names if 'externalLink' in x]
    result['connections_parts']=[x for x in names if 'connection' in x.lower()]
    result['external_relationships']={n:z.read(n).decode('utf-8') for n in names if n.endswith('.rels') and b'TargetMode="External"' in z.read(n)}
save('registry_structure.json',result)
for b in books:b.close()
print(json.dumps({'sheets':[{'name':s['name'],'dimension':s['dimension'],'rows':s['nonempty_rows'],'formula_count':len(s['formulas']),'sample':s['raw_rows'][:5]} for s in result['sheets']], 'vba':{k:v for k,v in result['vba'].items() if k!='safe_ascii_strings'},'links':result['external_link_parts']},ensure_ascii=True,default=str))
