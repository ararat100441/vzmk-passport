import json,re,hashlib,sys
from pathlib import Path
from collections import Counter,defaultdict
from xml.etree import ElementTree as ET
from decimal import Decimal
import openpyxl
OUT=Path(__file__).resolve().parent;APP=OUT.parent.parent;ROOT=APP.parent
sys.dont_write_bytecode=True;sys.path.insert(0,str(APP))
from passport_app.catalog import key
from profiles_stage4.storage import load_catalog
from profiles_stage4.model import calculate_uzk,frequency_angle
def read(n):return json.loads((OUT/n).read_text(encoding='utf-8-sig'))
def save(n,d):(OUT/n).write_text(json.dumps(d,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
raw=read('registry_structure.json')['sheets'][0]['raw_rows'][1:]
records=[{'row':r['row'],'project':r['values'].get('A'),'mark':r['values'].get('B'),'raw_quantity':r['values'].get('C'),'remaining':r['values'].get('C',1)} for r in raw]
duplicates=defaultdict(list)
for r in records:duplicates[(key(r['project']),key(r['mark']))].append(r['row'])
samples={'first':records[0],'middle':records[(len(records)-1)//2],'last':records[-1],
         'blank_first_5':[r for r in records if r['raw_quantity'] is None][:5],
         'numeric_first_5':[r for r in records if isinstance(r['raw_quantity'],(int,float))][:5],
         'duplicates':[{ 'key':k,'rows':v} for k,v in duplicates.items() if len(v)>1],
         'special':[r for r in records if r['row'] in [9,22,23,24,25,26]]}
save('registry_samples.json',samples)
vom=read('native_vom.json');s=vom['sheets'][0];hidden=set(s['hidden_rows'])
vrows=[]
for r in s['rows']:
    c={x['column']:x['value'] for x in r['cells']}
    if isinstance(c.get(2),str) and isinstance(c.get(4),(int,float)):
        vrows.append({'row':r['row'],'mark':c[2],'name':c.get(3),'quantity':c[4],'mass_total':c.get(9),'hidden':r['row'] in hidden})
bykey=defaultdict(list)
for r in vrows:bykey[key(r['mark'])].append(r)
matched=[]
for r in records:
    if r['project']!='Гараж Казахстан':continue
    vs=bykey[key(r['mark'])]
    assert len(vs)<=1,'STOP VOM normalized mark collision'
    matched.append(r|{'vom_matches':vs,'Q_all_rows':min(r['remaining'],vs[0]['quantity']) if vs else 0,
                      'Q_visible_only':min(r['remaining'],vs[0]['quantity']) if vs and not vs[0]['hidden'] else 0,
                      'profile_status':'MARK_CATALOG_MISSING','lookup_status':'NOT_RUN','geometry_status':'NOT_ESTABLISHED'})
counts={'registry_rows':len(records),'blank_quantity':sum(r['raw_quantity'] is None for r in records),'numeric_quantity':sum(r['raw_quantity'] is not None for r in records),
        'remaining_by_approved_rule':sum(r['remaining'] for r in records),
        'project_counts':dict(Counter(r['project'] for r in records)),
        'VOM_all_mark_rows':len(vrows),'VOM_visible_mark_rows':sum(not r['hidden'] for r in vrows),
        'VOM_all_quantity':sum(r['quantity'] for r in vrows),'VOM_visible_quantity':sum(r['quantity'] for r in vrows if not r['hidden']),
        'VOM_all_mass':sum(Decimal(str(r['mass_total'])) for r in vrows),
        'VOM_visible_mass':sum(Decimal(str(r['mass_total'])) for r in vrows if not r['hidden']),
        'selected_project_registry_rows':len(matched),'matched_any_VOM_rows':sum(bool(r['vom_matches']) for r in matched),
        'matched_visible_VOM_rows':sum(r['Q_visible_only']>0 for r in matched),
        'Q_all_rows':sum(r['Q_all_rows'] for r in matched),'Q_visible_only':sum(r['Q_visible_only'] for r in matched),
        'missing_working_mark_catalog':len(matched),'outside_selected_project':len(records)-len(matched),
        'verified_supported_real_registry_rows':0,'unresolved_profile_rows':len(records),'real_profile_lookups_executed':0}
save('compatibility.json',{'counts':counts,'rows':matched,'VOM_mark_rows':vrows,'scope_status':'USER_DECISION_REQUIRED_ALL_VS_VISIBLE'})
books=[openpyxl.load_workbook(APP/'DEV'/'ВЗМК_DEV.xlsx',data_only=x) for x in (False,True)]
forms,values=books;f=forms['Паспорт качества'];v=values[f.title]
devrows=[]
for ri in range(2,f.max_row+1):
    if f.cell(ri,1).value=='Гараж Казахстан':
        devrows.append({'row':ri,'passport':f.cell(ri,2).value,'end_date':v.cell(ri,9).value,'VIK_date':v.cell(ri,13).value,
                        'UZK_date':v.cell(ri,14).value,'UZK_formula':f.cell(ri,14).value,'object':v.cell(ri,5).value,'code':v.cell(ri,6).value,'mass':v.cell(ri,7).value})
save('header_sources.json',{'sheet':f.title,'columns':{'number':'B','object':'E','code':'F','end_date':'I','VIK_date':'M','UZK_date':'N'},'garage_rows':devrows})
for b in books:b.close()
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def text_lines(xml):
    e=ET.fromstring(xml);chunks=[]
    for node in e.iter():
        if node.tag==W+'t':chunks.append(node.text or '')
        elif node.tag==W+'br':chunks.append('\n')
    return ''.join(chunks)
docs=read('word_structure.json');word_rows=[];mismatches=[]
for doc in docs:
    for r in doc['tables'][0]['rows'][5:]:
        c=r['cells'];geometry=text_lines(c[1]['xml']);freq=text_lines(c[4]['xml']).splitlines();angle=text_lines(c[5]['xml']).splitlines()
        dimensions=re.findall(r'L\s*=\s*([\d.,]+)[хx×]([\d.,]+)',geometry)
        observed={'path':doc['path'],'row':r['row'],'mark':c[0]['text'],'geometry':geometry,'frequency':freq,'angle':angle,'dimensions':dimensions}
        for i,(length,thickness) in enumerate(dimensions):
            expected=frequency_angle(thickness)
            if i<len(freq) and i<len(angle):
                actual={'frequency_mhz':float(freq[i].replace(',','.')),'angle_degrees':float(angle[i].replace(',','.'))}
                if actual!=expected:mismatches.append(observed|{'segment':i+1,'expected':expected,'actual':actual})
        word_rows.append(observed)
save('word_row_analysis.json',{'rows':word_rows,'historical_frequency_angle_mismatches':mismatches})
catalog=load_catalog(APP/'DEV'/'Справочники'/'Профили.xlsx')
lookup=catalog.lookup('35Б1');geometry=calculate_uzk(lookup)
save('accepted_calculation_probe.json',{'profile':'35Б1','lookup':lookup,'one_item':geometry,
      'Q3':[r|{'L':r['L']*3} for r in geometry['rows']], 'boundaries':{'11.5':frequency_angle(11.5),'12':frequency_angle(12)}})
macros=[]
for p in OUT.glob('VBA_*.txt'):
    for i,line in enumerate(p.read_text(encoding='utf-8-sig').splitlines(),1):
        m=re.match(r'\s*(?:(Public|Private|Friend)\s+)?(Sub|Function)\s+(\w+)',line)
        if m:macros.append({'component':p.name,'line':i,'scope':m[1] or 'default','kind':m[2],'name':m[3]})
save('vba_procedures.json',macros)
print(json.dumps({'counts':counts,'visible_registry_matches':[{'mark':r['mark'],'remaining':r['remaining'],'vom':r['vom_matches'],'Q':r['Q_visible_only']} for r in matched if r['Q_visible_only']>0],
                  'date_rows':[r for r in devrows if r['passport'] in ('1/2026','2/2026','9/2026')],
                  'historical_frequency_angle_mismatches':len(mismatches),'vba_procedures':len(macros)},ensure_ascii=True,default=str))
