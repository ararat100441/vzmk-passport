"""S5-R1 rereview: newly authored workbook fixtures; current product bytes only."""
import json,os,shutil,sys,tempfile
from pathlib import Path
from copy import deepcopy
from contextlib import contextmanager,ExitStack
from datetime import datetime
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
RUN=Path(__file__).resolve().parent;ROOT=RUN.parents[1];sys.path.insert(0,str(ROOT))
TEMP=RUN/'independent_tmp';TEMP.mkdir();tempfile.tempdir=str(TEMP)
os.environ['TMP']=os.environ['TEMP']=str(TEMP)
import uzk_stage5.inputs as I
import uzk_stage5.core as C
import uzk_stage5.word as W
from profiles_stage4.model import Catalog
from profiles_stage4.offline import no_network
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def write(n,v):(RUN/n).write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str),encoding='utf8')
PRODUCT=[Path(p) for p in read(RUN/'preflight.json')['product_current']]
EVENTS=[]
def audit(event,args):
    if event!='open' or not isinstance(args[0],(str,bytes)):return
    p=Path(os.fsdecode(args[0])).resolve()
    if p.name.casefold() in ('взмк.xlsx','узк.xlsm'):raise AssertionError('Forbidden production read')
    mode,flags=args[1:3]
    writing=(isinstance(mode,str) and any(c in mode for c in 'wax+')) or (isinstance(flags,int) and flags&(os.O_WRONLY|os.O_RDWR|os.O_CREAT|os.O_TRUNC))
    if writing and not p.is_relative_to(RUN):raise AssertionError('Write outside new rereview: '+str(p))
    EVENTS.append(dict(path=str(p),writing=bool(writing)))
sys.addaudithook(audit)
BASE=RUN/'new_synthetic_base';BASE.mkdir()
dev=BASE/'DEV';dev.mkdir();profiles=dev/'Справочники';profiles.mkdir()
for n in ('Профили.xlsx','profiles.user.json','profiles.catalog.json','profiles.build.json'):
    shutil.copyfile(I.PROFILES.parent/n,profiles/n)
for p in (ROOT/'profiles_stage4').glob('*.py'):
    dest=BASE/'profiles_stage4'/p.name;dest.parent.mkdir(exist_ok=True);shutil.copyfile(p,dest)
dest=BASE/'passport_app/catalog.py';dest.parent.mkdir();shutil.copyfile(ROOT/'passport_app/catalog.py',dest)
project=BASE/'Документы/Гараж казахстан';passport=project/'Паспорт';passport.mkdir(parents=True)
existing=read(ROOT/'validation/stage5_independent_review_20261008_01/independent_inputs.json')
template=next(p for p in existing['input_fingerprints'] if p.endswith('1_2026 - УЗК.docx'))
shutil.copyfile(template,passport/'1_2026 - УЗК.docx')
(BASE/'settings.json').write_text(json.dumps({'documents_root':'Документы','projects':{'Гараж Казахстан':{'passport_root':'Гараж казахстан/Паспорт'}}},ensure_ascii=False),encoding='utf8')
b=Workbook();s=b.active;s.title='Паспорт качества'
header=['Проект','№ паспорта',None,None,'Наименование объекта','шифр КМД','Масса',None,None,None,None,None,None,'УЗК Дата'];s.append(header)
for number in ('1/2026','9/2026'):
    s.append(['Гараж Казахстан',number,None,None,'Объект независимой фикстуры','FIXTURE-КМД',20,None,None,None,None,None,None,datetime(2026,10,1)])
b.save(dev/'ВЗМК_DEV.xlsx');b.close()
shutil.copyfile(I.UZK,dev/'УЗК_DEV.xlsm')
b=load_workbook(dev/'УЗК_DEV.xlsm',keep_vba=True);s=b['УЗК'];s.delete_rows(2,s.max_row)
s.append(['ДРУГОЙ ПРОЕКТ','Бп2-1',8]);s.append(['Гараж Казахстан','ABSENT_FROM_VOM',3]);b.save(dev/'УЗК_DEV.xlsm');b.close()
shutil.copyfile(I.CHECKPOINT,dev/'DEV_UZK_CHECKPOINT.json')
b=Workbook();s=b.active;s.title='ВОМ'
s.append(['Независимая ведомость отправочных марок'])
s.append(['Лист №','Отправочная марка','Наименование марки','Количество',None,None,None,'Масса (изделие), кг',None])
s.append([None,None,None,None,None,None,None,'1 дет.','Всех'])
s.append(['1','Бп2-1','Балка',2,None,None,None,10,20]);b.save(project/'independent_vom.xlsx');b.close()
b=Workbook();s=b.active;s.title='Марки';s.append(['Марка','Профиль','Краска','№ АКЗ'])
for mark,profile in [('Бп2-1','[16П'),('FOREIGN_A','[16П'),(' foreign_a ','Гн.[]120х3'),('FOREIGN_B','Гн.[140X60X4'),('foreign_b','[16П'),('UNRELATED_CHECK','ПРОВЕРИТЬ'),('UNRELATED_MISSING',None)]:s.append([mark,profile,None,None])
b.save(project/'Справочник марок.xlsx');b.close()
def hashes(root):return {str(p.relative_to(root)):I.sha(p) for p in root.rglob('*') if p.is_file()}
def dirs(root):return sorted(str(p.relative_to(root)) for p in root.rglob('*') if p.is_dir())
@contextmanager
def scenario(name,required=False):
    root=RUN/'independent_scenarios'/name;shutil.copytree(BASE,root)
    paths={'ROOT':root,'REGISTRY':root/'DEV/ВЗМК_DEV.xlsx','UZK':root/'DEV/УЗК_DEV.xlsm','PROFILES':root/'DEV/Справочники/Профили.xlsx','CHECKPOINT':root/'DEV/DEV_UZK_CHECKPOINT.json'}
    if required:
        b=load_workbook(paths['UZK'],keep_vba=True);b['УЗК']['A2']='Гараж Казахстан';b.save(paths['UZK']);b.close()
    pins={paths[n]:I.sha(paths[n]) for n in ('REGISTRY','UZK','PROFILES')}
    def refresh():
        for p in pins:
            if p.exists():pins[p]=I.sha(p)
        cp=read(paths['CHECKPOINT']);cp['dev_path']=str(paths['UZK']);cp['dev_initial_sha256']=pins[paths['UZK']]
        paths['CHECKPOINT'].write_text(json.dumps(cp,ensure_ascii=False),encoding='utf8')
    refresh();passports=I.passports;registry=I.registry_rows
    with ExitStack() as stack:
        for module in (I,C):
            for n,p in paths.items():stack.enter_context(patch.object(module,n,p))
        stack.enter_context(patch.object(I,'PINS',pins))
        stack.enter_context(patch.object(I,'passports',side_effect=lambda:passports(paths['REGISTRY'])))
        stack.enter_context(patch.object(C,'registry_rows',side_effect=lambda:registry(paths['UZK'])))
        stack.enter_context(patch.object(C,'product_paths',return_value=PRODUCT))
        yield root,root/'Документы/Гараж казахстан/independent_vom.xlsx',refresh
def preview(vom):return C.preview('Гараж Казахстан','9/2026',vom)
def edit(root,file,fn,vba=False):
    path=root/file;b=load_workbook(path,keep_vba=vba);fn(b.active if not vba else b['УЗК']);b.save(path);b.close()
def stop(fn,code):
    try:fn()
    except Exception as exc:assert code in str(exc),(code,str(exc));return str(exc)
    raise AssertionError('Expected STOP '+code)
def empty(plan):
    assert plan.data['status']=='NO_UZK_REQUIRED' and plan.data['marks']==[] and plan.data['controlled_units']==0
    plan.verify()
RESULTS=[]
def test(name,fn):
    print('START '+name,flush=True)
    try:RESULTS.append(dict(name=name,status='PASS',detail=fn()))
    except Exception as exc:RESULTS.append(dict(name=name,status='FAIL',error=type(exc).__name__+': '+str(exc)))
    write('independent_cases.json',RESULTS);print('END '+name+' '+RESULTS[-1]['status'],flush=True)
def complete_empty():
    with scenario('new_zero_matches') as (root,vom,refresh):
        before=hashes(root);before_dirs=dirs(root);order=[]
        real_vom=C.visible_vom;real_registry=C.registry_rows;real_project=C.project_sources
        def v(*a,**k):order.append('VOM');return real_vom(*a,**k)
        def r(*a,**k):order.append('REGISTRY');return real_registry(*a,**k)
        def p(*a,**k):order.append('PROJECT');return real_project(*a,**k)
        with patch.object(C,'visible_vom',side_effect=v),patch.object(C,'registry_rows',side_effect=r),patch.object(C,'project_sources',side_effect=p),patch.object(C,'mark_profiles',side_effect=AssertionError('No mappings allowed')) as mapping,patch.object(C,'load_catalog',side_effect=AssertionError('No Stage 4 load allowed')) as stage4,patch.object(W,'write_test_document',side_effect=AssertionError('No renderer allowed')) as writer:
            plan=preview(vom);empty(plan);out=root/'validation/no_output'
            assert C.generate(plan,out) is None and C.generate(plan) is None
            mapping.assert_not_called();stage4.assert_not_called();writer.assert_not_called()
        assert order==['PROJECT','VOM','REGISTRY'],order
        assert hashes(root)==before and dirs(root)==before_dirs
        import tkinter as tk
        from uzk_stage5.gui import App
        window=tk.Tk();window.withdraw()
        try:
            app=App(window);app.queue.put(('preview',plan,None));app.poll()
            assert app.status.get()=='Для текущего паспорта УЗК не требуется'
            assert app.generate_button.instate(['disabled'])
            app.checked.set(True);app.update_buttons();assert app.generate_button.instate(['disabled']) and app.output is None
            message=app.status.get()
        finally:window.destroy()
        write('fresh_no_uzk_preview.json',dict(data=plan.data,seal=plan.seal))
        return dict(status=plan.data['status'],marks=0,controlled_units=0,call_order=order,new_DOCX=0,new_directories=0,renderer_calls=0,registry_mutations=0,profile_mapping_calls=0,Stage4_calls=0,GUI_message=message,GUI_generate_disabled=True,fixture_inputs_unchanged=True)
test('Independent new fixture full preview -> generate -> GUI',complete_empty)
def missing():
    with scenario('missing_registry') as (root,vom,refresh):
        (root/'DEV/УЗК_DEV.xlsm').unlink();return stop(lambda:preview(vom),'CHECKPOINT_CHANGED')
test('A Missing registry STOP',missing)
def invalid(kind):
    with scenario('invalid_registry_'+kind) as (root,vom,refresh):
        path=root/'DEV/УЗК_DEV.xlsm'
        if kind=='unreadable':path.write_bytes(b'NOT_A_WORKBOOK')
        elif kind=='header':edit(root,'DEV/УЗК_DEV.xlsm',lambda s:setattr(s['A1'],'value','BAD'),True)
        else:edit(root,'DEV/УЗК_DEV.xlsm',lambda s:setattr(s['C2'],'value',-1),True)
        refresh();return stop(lambda:preview(vom),{'unreadable':'File is not a zip file','header':'UZK_HEADER_CHANGED','quantity':'QUANTITY_INVALID'}[kind])
for kind in ('unreadable','header','quantity'):test('B Invalid registry '+kind+' STOP',lambda k=kind:invalid(k))
def invalid_vom(kind):
    with scenario('invalid_vom_'+kind) as (root,vom,refresh):
        cell,value,code={'header':('B2','BAD','VOM_SHEET_AMBIGUOUS'),'quantity':('D4',0,'QUANTITY_INVALID'),'mass':('I4',19,'VOM_SHIPMENT_MISMATCH')}[kind]
        edit(root,str(vom.relative_to(root)),lambda s:setattr(s[cell],'value',value))
        with patch.object(C,'registry_rows',side_effect=AssertionError('Invalid VOM must stop first')):
            return stop(lambda:preview(vom),code)
for kind in ('header','quantity','mass'):test('C Invalid VOM '+kind+' STOP',lambda k=kind:invalid_vom(k))
def other_project():
    with scenario('other_project') as (root,vom,refresh):
        stock,rows=C.registry_rows();assert ('другой проект','бп2-1') in stock and ('гараж казахстан','бп2-1') not in stock
        plan=preview(vom);empty(plan);return dict(status=plan.data['status'],same_mark_other_project=True)
test('D Same mark other project NO_UZK_REQUIRED',other_project)
def required_conflict():
    with scenario('required_conflict',True) as (root,vom,refresh):
        edit(root,'Документы/Гараж казахстан/Справочник марок.xlsx',lambda s:s.append([' бп2-1 ','Гн.[]120х3',None,None]))
        return stop(lambda:preview(vom),'MARK_PROFILE_CONFLICT')
test('E Required profile mapping conflict STOP',required_conflict)
def unresolved(raw,code):
    with scenario('unresolved_'+code,True) as (root,vom,refresh):
        edit(root,'Документы/Гараж казахстан/Справочник марок.xlsx',lambda s:setattr(s['B2'],'value',raw))
        return stop(lambda:preview(vom),code)
for raw,code in [('UNKNOWN_PROFILE_S5R1','PROFILE_NOT_FOUND'),('ПРОВЕРИТЬ','PROFILE_UNRESOLVED')]:test('F Required '+code+' STOP',lambda r=raw,c=code:unresolved(r,c))
source=deepcopy(C.load_catalog(I.PROFILES).lookup('[16П')['profile']['provenance'][0])
def stage4_failure(kind):
    with scenario('Stage4_'+kind,True) as (root,vom,refresh):
        other=deepcopy(source);other['record_id']='INDEPENDENT';other['geometry']['b']=65
        other['standard']='ГОСТ 99999-2000' if kind=='AMBIGUOUS' else source['standard']
        cat=Catalog([source] if kind=='UZK_RULE_NOT_CONFIGURED' else [source,other])
        with patch.object(C,'load_catalog',return_value=cat):return stop(lambda:preview(vom),kind)
for kind in ('AMBIGUOUS','CONFLICT','UZK_RULE_NOT_CONFIGURED'):test(('G ' if kind=='UZK_RULE_NOT_CONFIGURED' else 'F ')+'Required Stage4 '+kind+' STOP',lambda k=kind:stage4_failure(k))
def unrelated_allowed():
    with scenario('valid_required_unrelated_bad',True) as (root,vom,refresh):
        before=hashes(root);plan=preview(vom)
        assert plan.data['status']=='UZK_REQUIRED' and len(plan.data['marks'])==1 and plan.data['controlled_units']==2
        mark=plan.data['marks'][0];assert mark['raw_profile']=='[16П' and mark['Q']==2
        assert [(s['L'],s['S']) for s in mark['segments']]==[('320',5),('256',8.4)]
        assert hashes(root)==before
        return dict(required_mark=mark['mark'],Q=mark['Q'],unrelated_conflicts_ignored=True,unrelated_CHECK_ignored=True,source_mutations=0)
test('H Required valid, unrelated conflicts and CHECK normal processing',unrelated_allowed)
def actual_four_marks():
    old=read(ROOT/'validation/stage5_independent_review_20261008_01/fresh_preview.json')['data']
    plan=C.preview('Гараж Казахстан','9/2026',old['vom']['path']);plan.verify()
    expected={'Бп2-1':(8,[('1280',5),('1024',8.4)]),'Пс1-9':(9,[('3600',4)]),'Пс1-10':(15,[('6000',4)]),'Пс2-16':(1,[('480',3)])}
    assert len(plan.data['marks'])==4 and plan.data['controlled_units']==33 and plan.data['status']=='UZK_REQUIRED'
    for m in plan.data['marks']:
        q,segs=expected[m['mark']];assert m['Q']==q and [(s['L'],s['S']) for s in m['segments']]==segs
        assert all(s['frequency_mhz']==5 and s['angle_degrees']==70 for s in m['segments'])
    write('actual_four_mark_preview.json',dict(data=plan.data,seal=plan.seal))
    return dict(marks=4,controlled_units=33,visible_rows=plan.data['vom']['visible_count'],mass=plan.data['vom']['mass_kg'],DOCX_generated=False)
test('Existing real inputs four-mark read-only regression',actual_four_marks)
write('independent_access_guard.json',dict(production_reads=0,writes_outside_rereview=0,open_events=EVENTS))
summary=dict(tests=len(RESULTS),passed=sum(x['status']=='PASS' for x in RESULTS),failures=[x for x in RESULTS if x['status']!='PASS'],source_product_hashes={str(p):I.sha(p) for p in PRODUCT})
write('independent_summary.json',summary);print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
raise SystemExit(0 if not summary['failures'] else 1)
