"""Read-only baseline, explicit sources, no production discovery."""
import sys, json, hashlib
from pathlib import Path
from datetime import datetime
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from passport_app.registry import read_registry, select_project
from openpyxl import load_workbook

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
E = ROOT / 'validation' / ('stage6_akz_part1_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
E.mkdir()
(ROOT / 'validation' / 'STAGE6_PART1_CURRENT.txt').write_text(str(E), encoding='utf-8')
PROJECT = ROOT.parent / 'Гараж казахстан'
sources = [ROOT/'DEV/ВЗМК_DEV.xlsx', ROOT/'DEV/УЗК_DEV.xlsm', ROOT/'DEV/Справочники/Профили.xlsx',
           PROJECT/'Справочник марок.xlsx', PROJECT/'Паспорт/Акт № 4-2026.xlsx',
           PROJECT/'Паспорт/1_2026 - Приложение к АКЗ.xlsx',
           PROJECT/'Паспорт/9 от 06.10.2026/ВОМ отгрузочный счет 183  (5).xls']
product = [p for folder in ('passport_app','profiles_stage4','uzk_stage5','tests')
           for p in (ROOT/folder).glob('*.py')]
product += [p for p in ROOT.iterdir() if p.is_file()]
baseline = {str(p):sha(p) for p in sources+product}
(E/'baseline.json').write_text(json.dumps(baseline, ensure_ascii=False, indent=2), encoding='utf-8')
r=read_registry(sources[0]); project=select_project(r,'Гараж Казахстан')
b=load_workbook(sources[0],data_only=True)
s=b[r.sheet]
headers={str(c.value).strip():c.column for c in s[1]}
row=[x for x in r.projects[project.name] if x['number']=='9/2026'][0]['row']
current={str(c.value).strip():s.cell(row,c.column).value for c in s[1] if c.value is not None}
b.close()
(E/'dev_current.json').write_text(json.dumps({'row':row,'headers':headers,'current':current},ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(E)
print(json.dumps(current,ensure_ascii=False,default=str))
print('BASELINE_FILES',len(baseline))
