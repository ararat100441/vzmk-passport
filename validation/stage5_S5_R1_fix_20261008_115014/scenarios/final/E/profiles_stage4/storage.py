"""DEV-only transactional persistence; authoritative user state is separate."""
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import uuid

from .model import Catalog, ProfileError, default_user_state, normalize_profile, normalize_standard, validate_geometry
from .offline import ROOT, load_sources, no_network, read_json, sha
from .windows_generation import GenerationGuard, digest

DEFAULT_WORKBOOK = ROOT / 'DEV/Справочники/Профили.xlsx'
STATE = 'profiles.user.json'
CATALOG = 'profiles.catalog.json'
COMMIT = 'profiles.build.json'
CORE_COLUMNS = ['Профиль', 'Ключ', 'Тип', 'НД', 'h', 'b', 'b2', 's', 't', 'D', 'r',
                'Источник', 'Дата получения', 'Статус', 'УЗК правило']
HEADERS = {
    'Профили': CORE_COLUMNS + ['Радиус R', 'Source ID', 'Evidence reference', 'Conflict ID', 'Примечание'],
    'Источники': ['Source ID', 'Источник', 'Тип источника', 'Family', 'НД', 'URL / файл', 'Дата получения',
                  'Snapshot/cache', 'SHA-256', 'Количество записей', 'Статус', 'Примечание'],
    'Правила УЗК': ['Правило', 'Тип', 'Часть', 'L, мм', 'S, мм', 'Частота, МГц', 'Угол, °', 'Область', 'Статус'],
    'Предпочтения': ['Ключ', 'Предпочтительный НД', 'Основание', 'Статус', 'Дата'],
    'Конфликты': ['Conflict ID', 'Профиль', 'Ключ', 'НД', 'Источник A', 'Значения A', 'Источник B', 'Значения B',
                  'Поля конфликта', 'Статус', 'Решение', 'Дата решения', 'Классификация'],
}


def guard_target(path):
    path = Path(path).resolve()
    # Validation fixtures have an explicit isolated DEV layout; no general destination override.
    fixture = ROOT / 'validation'
    if path == DEFAULT_WORKBOOK.resolve():
        return path
    if path.is_relative_to(fixture.resolve()) and path.name == 'Профили.xlsx' \
            and path.parent.name == 'Справочники' and path.parent.parent.name == 'DEV':
        return path
    raise ProfileError('Only DEV/Справочники/Профили.xlsx or isolated validation DEV fixture is allowed')


def timestamp():
    return datetime.now(timezone.utc).isoformat()


def dump(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def workbook_payload(catalog):
    from .model import RULE, NOT_CONFIGURED, BENT_RULE, CLOSED_RULE, CHANNEL_RULE, APPROVED
    sheets = {name: [headers] for name, headers in HEADERS.items()}
    for row in catalog.rows:
        p = row['provenance']
        source_ids = sorted({r['source_id'] for r in p})
        notes = row['notes'][:]
        if row['geometry'].get('R') is None:
            notes.append('R отсутствует или unresolved')
        if row['geometry'].get('r') is None:
            notes.append('r отсутствует или unresolved')
        sheets['Профили'].append([row['profile'], row['key'], row['family'], row['standard'],
            *(row['geometry'].get(f) for f in ('h', 'b', 'b2', 's', 't', 'D', 'r')),
            ' + '.join(sorted({r['source'] for r in p})), min(r['date'] for r in p), row['status'], row['uzk_rule'],
            row['geometry'].get('R'), '; '.join(source_ids),
            CATALOG + '#' + ','.join(r['record_id'] for r in p), row['conflict_id'], '; '.join(notes)])
    for s in catalog.sources:
        sheets['Источники'].append([s[k] for k in ('id', 'source', 'type', 'family', 'standard', 'reference', 'date',
                                                    'snapshot', 'sha256', 'count', 'status', 'notes')])
    for m in catalog.user_state['manual_rows']:
        sheets['Источники'].append([m['source_id'], 'USER', 'manual state', m['family'], m['standard'], STATE,
                                     m['date'], STATE, '', 1, 'РУЧНОЙ', '; '.join(m['notes'])])
    sheets['Правила УЗК'] += [
        [RULE, 'Двутавр', 'стенка', 'h - 2*t', 's', 'S < 12: 5; S >= 12: 2.5', 'S < 12: 70; S >= 12: 65', 'Одно изделие', 'ПОДТВЕРЖДЕНО ПОЛЬЗОВАТЕЛЕМ'],
        [RULE, 'Двутавр', 'полки', '2*b', 't', 'S < 12: 5; S >= 12: 2.5', 'S < 12: 70; S >= 12: 65', 'Одно изделие', 'ПОДТВЕРЖДЕНО ПОЛЬЗОВАТЕЛЕМ'],
    ]
    definitions = {
        BENT_RULE: [('Гнутый швеллер: подтверждённая identity', 'периметр', '2*(h+b)', 't')],
        CLOSED_RULE: [('Квадратный / прямоугольный; ГОСТ 30245-2003', 'периметр', '2*(h+b)', 't')],
        CHANNEL_RULE: [('Только 16П; ГОСТ 8240-97', 'стенка', 'h', 's'), ('Только 16П; ГОСТ 8240-97', 'полки', '2*b', 't')],
    }
    for rule, definitions_for_rule in definitions.items():
        approvals = [a for a in catalog.user_state.get('uzk_approvals', []) if a.get('rule') == rule and a.get('status') == APPROVED and a.get('basis') and a.get('date')]
        if approvals:
            status = APPROVED + '; ' + '; '.join(sorted({a['date'] for a in approvals}))
            for family, part, length, thickness in definitions_for_rule:
                sheets['Правила УЗК'].append([rule, family, part, length, thickness, 'S < 12: 5; S >= 12: 2.5', 'S < 12: 70; S >= 12: 65', 'Одно изделие', status])
    sheets['Правила УЗК'].append([NOT_CONFIGURED, 'Без отдельного подтверждения правила', '', None, None, None, None, 'Без количества', NOT_CONFIGURED])
    for p in catalog.user_state['preferences']:
        sheets['Предпочтения'].append([p[k] for k in ('key', 'standard', 'basis', 'status', 'date')])
    for c in catalog.conflicts:
        members = c['members']
        a = [m for m in members if m['source'] == members[0]['source']]
        b = [m for m in members if m not in a]
        if not b:  # Internal duplicate disagreement is represented too.
            a, b = members[:1], members[1:]
        def values(items):
            return json.dumps([{'record_id': m['record_id'], 'geometry': m['geometry'], 'quality': m.get('quality', {}),
                                'reference': CATALOG + '#' + m['record_id']} for m in items], ensure_ascii=False, separators=(',', ':'))
        resolution = c['resolution'] or {}
        sheets['Конфликты'].append([c['id'], c['profile'], c['key'], c['standard'],
            ' + '.join(sorted({m['source_id'] for m in a})), values(a),
            ' + '.join(sorted({m['source_id'] for m in b})), values(b), ', '.join(c['fields']), c['status'],
            resolution.get('record_id', ''), resolution.get('date', ''), c['classification']])
    return {'sheets': sheets}


def runtime():
    base = Path(os.environ['USERPROFILE']) / '.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    node = base / 'bin/node.exe'
    package = base / 'node_modules/@oai/artifact-tool'
    if not node.is_file() or not package.is_dir():
        raise ProfileError('Bundled offline artifact-tool/Node runtime unavailable')
    return node


def render_workbook(payload, destination, previews=None):
    cmd = [str(runtime()), str(Path(__file__).with_name('workbook.mjs')), str(payload), str(destination)]
    if previews:
        cmd += ['--preview', str(previews)]
    result = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', timeout=180)
    if result.returncode:
        raise ProfileError('Workbook builder failed: ' + (result.stderr or result.stdout)[-2500:])
    return json.loads(result.stdout.strip().splitlines()[-1])


def verify_workbook(path, expected=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if wb.sheetnames != list(HEADERS):
            raise ProfileError('Workbook sheets differ')
        for name, headers in HEADERS.items():
            ws = wb[name]
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)  # Artifact export may omit optional XML dimension metadata.
            if list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))) != headers:
                raise ProfileError('Workbook headers differ: ' + name)
            if expected is not None:
                matrix = expected['sheets'][name]
                if ws.max_row != len(matrix):
                    raise ProfileError(f'Reopen count differs: {name}: {ws.max_row} vs {len(matrix)}')
                for actual, want in zip(ws.iter_rows(values_only=True), matrix):
                    for a, v, column in zip(actual, want, headers):
                        if column.startswith('Дата') or column == 'Дата':
                            if v not in (None, '') and a is None:
                                raise ProfileError('Lost date')
                        elif isinstance(v, str) and v[:1] in ('=', '+', '@', '-') and a == "'" + v:
                            pass  # Escaped literal text; never a formula.
                        elif (None if v == '' else v) != a:
                            raise ProfileError(f'Reopen value differs: {name}:{column}:{v!r}!={a!r}')
            for row in ws.iter_rows():
                if any(cell.data_type == 'f' for cell in row):
                    raise ProfileError('Unexpected executable formula in value-only catalogue')
        for row in wb['Профили'].iter_rows(min_row=2, values_only=True):
            for i in [4, 5, 6, 7, 8, 9, 10, 15]:
                if row[i] is not None and (isinstance(row[i], bool) or not isinstance(row[i], (int, float))):
                    raise ProfileError('Geometry was exported as text')
        return {'reopened': True, 'sheets': wb.sheetnames, 'rows': wb['Профили'].max_row - 1,
                'numeric_geometry': True, 'formulas': 0}
    finally:
        wb.close()


@contextmanager
def locked(folder):
    path = folder / '.profiles.lock'
    try:
        fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise ProfileError('Another Stage 4 writer is active; no files changed') from exc
    try:
        os.write(fd, str(os.getpid()).encode())
        os.close(fd)
        yield
    finally:
        path.unlink(missing_ok=True)


def validated_generation(data, workbook_name):
    """Validate one protected, coherent snapshot, including the authoritative manifest."""
    if all(value is None for value in data.values()):
        return default_user_state()
    if any(value is None for value in data.values()):
        raise ProfileError('Existing unbound/incomplete generation requires review; no automatic overwrite')
    commit = json.loads(data[COMMIT].decode('utf-8-sig'))
    if set(commit['files']) != {workbook_name, STATE, CATALOG}:
        raise ProfileError('Invalid local build manifest')
    for name, expected in commit['files'].items():
        if digest(data[name]) != expected:
            raise ProfileError('Local generation changed outside Stage 4 service: ' + name)
    state = json.loads(data[STATE].decode('utf-8-sig'))
    catalog = json.loads(data[CATALOG].decode('utf-8-sig'))
    if state != catalog['user_state']:
        raise ProfileError('User state/catalogue generation mismatch')
    return state


def build(workbook=DEFAULT_WORKBOOK, *, user_state=None, previews=None, fail_at=None):
    target = guard_target(workbook)
    folder = target.parent
    folder.mkdir(parents=True, exist_ok=True)
    with locked(folder), no_network() as network:
        files = [target.name, STATE, CATALOG, COMMIT]
        # Capture all four files while external write/delete access is denied.
        # Release during rendering; any intervening edit must survive and abort commit.
        with GenerationGuard(folder, files) as previous:
            previous_state = validated_generation(previous.data, target.name)
            expected_hashes = previous.hashes
        state = deepcopy(previous_state if user_state is None else user_state)
        if user_state is not None and state['revision'] != previous_state['revision']:
            raise ProfileError('User state changed concurrently')
        state['revision'] = previous_state['revision'] + 1
        records, sources = load_sources()
        catalog = Catalog(records, sources, state)
        payload = workbook_payload(catalog)
        stage = Path(tempfile.mkdtemp(prefix='.profiles-stage-', dir=folder))
        backup = None
        try:
            dump(stage / 'payload.json', payload)
            builder = render_workbook(stage / 'payload.json', stage / target.name, previews)
            reopen = verify_workbook(stage / target.name, payload)
            dump(stage / STATE, state)
            dump(stage / CATALOG, {'version': 1, 'records': catalog.raw_records, 'sources': catalog.sources, 'user_state': state})
            dump(stage / COMMIT, {'version': 1, 'built_at': timestamp(), 'counts': catalog.counts(),
                                 'files': {f: sha(stage/f) for f in files if f != COMMIT}, 'network_attempts': len(network),
                                 'renderer_network_attempts': builder['network_attempts'], 'reopen': reopen})
            if fail_at == 'before_commit':
                raise ProfileError('Injected safe-write failure before commit')
            if network or builder['network_attempts']:
                raise ProfileError('Network attempted')
            with GenerationGuard(folder, files) as current:
                if current.hashes != expected_hashes:
                    raise ProfileError('Previous generation changed after preflight; STOP before backup/replacement')
                validated_generation(current.data, target.name)
                if current.handles:
                    backup = folder / 'backups' / (datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S_%fZ') + '_' + uuid.uuid4().hex[:8])
                    pending_backup = stage / 'backups'
                    pending_backup.mkdir()
                    for name, data in current.data.items():
                        (pending_backup/name).write_bytes(data)
                        if sha(pending_backup/name) != expected_hashes[name]:
                            raise ProfileError('Backup verification failed')
                    validated_generation({name:(pending_backup/name).read_bytes() for name in files}, target.name)
                    # Expose only a complete verified backup, never a partial copy.
                    backup.parent.mkdir(exist_ok=True)
                    os.rename(pending_backup, backup)
                # The handles remain held through backup, commit and any rollback.
                current.publish(stage, fail_at)
            return {'workbook': str(target), 'sha256': sha(target), 'counts': catalog.counts(),
                    'reopen': reopen, 'network_attempts': 0, 'backup': str(backup) if backup else None}
        finally:
            # Only the verified private temporary directory beneath the allowed DEV folder.
            if stage.resolve().parent != folder.resolve() or not stage.name.startswith('.profiles-stage-'):
                raise ProfileError('Unsafe temporary cleanup path')
            shutil.rmtree(stage)


def load_catalog(workbook=DEFAULT_WORKBOOK):
    target = guard_target(workbook)
    folder = target.parent
    commit = read_json(folder / COMMIT)
    expected_files = {target.name, STATE, CATALOG}
    if set(commit['files']) != expected_files:
        raise ProfileError('Invalid local build manifest')
    for f, digest in commit['files'].items():
        if sha(folder/f) != digest:
            raise ProfileError('Local catalogue/user state changed outside Stage 4 service: ' + f)
    data = read_json(folder / CATALOG)
    state = read_json(folder / STATE)
    if state != data['user_state']:
        raise ProfileError('User state/catalogue generation mismatch')
    return Catalog(data['records'], data['sources'], state)


def lookup(profile_text, standard=None, workbook=DEFAULT_WORKBOOK):
    with no_network():
        return load_catalog(workbook).lookup(profile_text, standard)


def prepare_manual_state(catalog, profile, family, standard, geometry, note='', *, verification=None):
    """Prepare one explicit user record without writing; enables a single batch generation."""
    state = deepcopy(catalog.user_state)
    if not str(profile).strip() or not str(standard).strip():
        raise ProfileError('Manual profile and standard are required')
    g = validate_geometry(family, geometry)
    key = normalize_profile(profile)
    if any(normalize_profile(m['key']) == key and normalize_standard(m['standard']) == normalize_standard(standard)
           for m in state['manual_rows']):
        raise ProfileError('Manual identity already exists')
    record_id = 'USER-' + uuid.uuid4().hex
    state['manual_rows'].append({'record_id': record_id, 'source_id': record_id, 'source': 'USER',
        'display': profile, 'key': key, 'family': family, 'standard': normalize_standard(standard),
        'geometry': g, 'quality': {}, 'raw': {'profile': profile, 'family': family, 'standard': standard, 'geometry': geometry, 'note': note},
        'date': timestamp(), 'evidence': {'file': STATE, 'record_id': record_id}, 'notes': [note] if note else [],
        'uzk_confirmed': False})
    if verification is not None:
        if verification.get('kind') != 'KMD' or not verification.get('references'):
            raise ProfileError('Verified manual record requires explicit KMD references')
        state['manual_rows'][-1]['verification'] = deepcopy(verification)
        state['manual_rows'][-1]['notes'].append('ПОДТВЕРЖДЕНО ПО КМД')
    return state


def add_manual(profile, family, standard, geometry, note='', *, workbook=DEFAULT_WORKBOOK):
    state = prepare_manual_state(load_catalog(workbook), profile, family, standard, geometry, note)
    return build(workbook, user_state=state)


def set_preference(profile, standard, basis, *, workbook=DEFAULT_WORKBOOK):
    catalog = load_catalog(workbook)
    key, nd = normalize_profile(profile), normalize_standard(standard)
    if catalog.lookup(key, nd)['status'] == 'NOT_FOUND' or not basis.strip():
        raise ProfileError('Preference requires an existing identity and an explicit basis')
    state = deepcopy(catalog.user_state)
    state['preferences'] = [p for p in state['preferences'] if normalize_profile(p['key']) != key]
    state['preferences'].append({'key': key, 'standard': nd, 'basis': basis,
                                'status': 'ПОДТВЕРЖДЕНО ПОЛЬЗОВАТЕЛЕМ', 'date': timestamp()})
    return build(workbook, user_state=state)


def resolve_conflict(conflict_id, record_id, reason, *, workbook=DEFAULT_WORKBOOK):
    catalog = load_catalog(workbook)
    conflict = next((c for c in catalog.conflicts if c['id'] == conflict_id), None)
    if not conflict or not any(m['record_id'] == record_id for m in conflict['members']) or not reason.strip():
        raise ProfileError('Explicit Conflict ID, candidate Record ID and reason are required')
    state = deepcopy(catalog.user_state)
    state['resolutions'][conflict_id] = {'record_id': record_id, 'fingerprint': conflict['fingerprint'],
                                       'reason': reason, 'date': timestamp()}
    return build(workbook, user_state=state)


def confirm_manual_uzk(record_id, basis, *, workbook=DEFAULT_WORKBOOK):
    catalog = load_catalog(workbook)
    state = deepcopy(catalog.user_state)
    row = next((m for m in state['manual_rows'] if m['record_id'] == record_id), None)
    if row is None or row['family'] != 'Двутавр' or not basis.strip():
        raise ProfileError('Only a manual I-beam with explicit rule confirmation is supported')
    row['uzk_confirmed'] = True
    row['notes'].append('USER_CONFIRMED_UZK_RULE: ' + basis)
    return build(workbook, user_state=state)
