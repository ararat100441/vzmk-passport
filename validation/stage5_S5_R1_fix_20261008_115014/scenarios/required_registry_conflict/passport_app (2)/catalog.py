"""Stage 3 only: read-only registry, preview, narrow XLSX cell patch and safe commit."""
from contextlib import contextmanager
from dataclasses import dataclass, replace
from hashlib import sha256 as digest
from io import BytesIO
import os
from pathlib import Path
import re
import unicodedata
from uuid import uuid4
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from .config import PassportError
from .excel import locked_file, atomic_replace, sha256
from .registry import read_registry

HEADERS = ('Марка', 'Профиль', 'Краска', '№ АКЗ')
CATALOG = 'Справочник марок.xlsx'
ARCHIVE_CONTAINER = 'ЗАКРЫТЫЕ ОБЪЕКТЫ'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
LABELS = {'NEW': 'Новая марка', 'FILL': 'Заполнить пустой профиль', 'SAME': 'Без изменений',
          'CONFLICT': 'Конфликт: требуется подтверждение', 'CHECK_REQUIRED': 'Требует проверки',
          'ERROR': 'Ошибка: исправьте входную таблицу'}


def key(value):
    return re.sub(' +', ' ', unicodedata.normalize('NFC', str(value)).strip()).casefold()


def profile_key(value):
    return re.sub(r'(?<=\d)[xх×](?=\d)', 'x', key(value))


def is_check(value):
    return key(value) == 'проверить'


def project_folder(settings, project):
    root = settings.documents_root.resolve()
    if not project or any(c in project for c in '<>:"/\\|?*') or project in ('.', '..'):
        raise PassportError('Небезопасное название проекта. Сопоставление папки не определено.')
    if project in settings.projects:
        passport, _ = settings.project_paths(project)  # Accepted explicit project mapping.
        folder = passport.parent.resolve()
    else:
        roots = [root]
        archive = root / ARCHIVE_CONTAINER
        if archive.exists():
            if not archive.is_dir() or archive.resolve().parent != root:
                raise PassportError('Контейнер ЗАКРЫТЫЕ ОБЪЕКТЫ выходит за разрешённую структуру документов.')
            roots.append(archive.resolve())
        matches = [p for allowed in roots for p in allowed.iterdir()
                   if p.is_dir() and key(p.name) == key(project)]
        if len(matches) > 1:
            raise PassportError(f'Неоднозначная папка проекта «{project}». Импорт остановлен.')
        if not matches:
            raise PassportError(f'Для проекта «{project}» нет папки в разрешённых roots. Требуется explicit mapping; правило создания не определено.')
        folder = matches[0].resolve()
        if folder.parent not in roots:
            raise PassportError('Папка проекта выходит за разрешённые roots.')
    if folder == root or not folder.is_relative_to(root) or not folder.is_dir():
        raise PassportError('Папка проекта отсутствует или выходит за корень документов.')
    return folder


@dataclass(frozen=True)
class Entry:
    mark: str
    profile: str
    lines: tuple
    error: str = ''
    warning: str = ''


@dataclass(frozen=True)
class PreviewRow:
    mark: str
    incoming: str
    current: str
    status: str
    row: int | None
    lines: tuple
    warning: str = ''
    confirmable: bool = False

    @property
    def token(self):
        return key(self.mark)


@dataclass(frozen=True)
class Preview:
    project: str
    path: Path
    catalog_hash: str | None
    rows: tuple
    text: str


def parse_input(text):
    lines = [(n, v.strip()) for n, v in enumerate(text.splitlines(), 1) if v.strip()]
    if not lines:
        raise PassportError('Вставьте таблицу из NotebookLM: Марка | Профиль.')
    markdown = '|' in lines[0][1] and '\t' not in lines[0][1]
    def cells(line):
        if markdown:
            line = line.removeprefix('|').removesuffix('|')
            return [v.strip() for v in line.split('|')]
        return [v.strip() for v in line.split('\t')]
    if tuple(cells(lines[0][1])) != HEADERS[:2]:
        raise PassportError('Нужны ровно два заголовка: Марка и Профиль (Markdown или текст с табуляцией).')
    start = 1
    if markdown:
        if len(lines) < 2 or len(cells(lines[1][1])) != 2 or not all(
                re.fullmatch(r':?-{3,}:?', v) for v in cells(lines[1][1])):
            raise PassportError('В Markdown-таблице отсутствует разделитель |---|---|.')
        start = 2
    if len(lines) == start:
        raise PassportError('В таблице нет марок.')
    entries = []
    positions = {}
    for number, line in lines[start:]:
        parts = cells(line)
        mark = parts[0] if parts else ''
        profile = parts[1] if len(parts) > 1 else ''
        error = ''
        if len(parts) != 2 or not mark or not profile:
            error = f'Строка {number}: нужны две непустые клетки Марка и Профиль.'
        elif any(len(v) > 32767 or any(ord(c) < 32 for c in v) for v in parts):
            error = f'Строка {number}: недопустимый текст или превышен лимит Excel.'
        entry = Entry(mark, profile, (number,), error)
        normalized = key(mark)
        if normalized and normalized in positions:
            i = positions[normalized]
            old = entries[i]
            conflict = old.error or error or (profile_key(old.profile) != profile_key(profile))
            entries[i] = replace(old, lines=old.lines + (number,),
                                 error='Одна марка имеет разные профили или некорректную строку.' if conflict else '',
                                 warning='Повтор марки во входной таблице: строки ' + ', '.join(map(str, old.lines + (number,))))
        else:
            if normalized:
                positions[normalized] = len(entries)
            entries.append(entry)
    return tuple(entries)


def read_catalog(data):
    try:
        book = openpyxl.load_workbook(BytesIO(data), keep_links=True)
    except Exception as exc:
        raise PassportError('Не удалось прочитать справочник XLSX. Файл не изменён.') from exc
    try:
        if len(book.worksheets) != 1 or book.sheetnames != [book.active.title]:
            raise PassportError('Другая структура справочника: требуется один лист с четырьмя столбцами.')
        sheet = book.active
        if tuple(sheet.cell(1, c).value for c in range(1, 5)) != HEADERS or sheet.max_column != 4:
            raise PassportError('Другая структура справочника. Нужны Марка | Профиль | Краска | № АКЗ. Миграция не определена.')
        if sheet.protection.sheet or (book.security and book.security.lockStructure) or sheet.merged_cells.ranges:
            raise PassportError('Справочник защищён или содержит объединённые клетки. Безопасная запись не определена.')
        if sheet.tables:
            raise PassportError('Справочник содержит Excel table. Безопасное расширение такой структуры не определено.')
        rows = {}
        for r in range(2, sheet.max_row + 1):
            mark, profile, paint, akz = [sheet.cell(r, c).value for c in range(1, 5)]
            if mark is None or not str(mark).strip():
                if any(v is not None for v in (profile, paint, akz)):
                    raise PassportError(f'Строка {r}: значения без марки. Импорт остановлен.')
                continue
            if not isinstance(mark, str) or (profile is not None and not isinstance(profile, str)) or any(
                    sheet.cell(r, c).data_type in ('f', 'e') for c in (1, 2)):
                raise PassportError(f'Строка {r}: марка и профиль должны быть обычным текстом.')
            normalized = key(mark)
            if normalized in rows:
                raise PassportError(f'Неоднозначность существующих марок после нормализации: «{mark}», строки {rows[normalized][0]} и {r}.')
            rows[normalized] = (r, mark, profile or '')
        return rows, sheet.max_row
    finally:
        book.close()


def make_rows(entries, existing):
    result = []
    for e in entries:
        row, _, current = existing.get(key(e.mark), (None, '', ''))
        confirmable = False
        warning = e.warning
        if e.error:
            status = 'ERROR'
            warning = (warning + ' ' + e.error).strip()
        elif is_check(e.profile):
            status = 'CHECK_REQUIRED'
            confirmable = row is None or not current.strip()
            if current.strip() and not is_check(current):
                warning = (warning + ' Нормальный профиль не заменяется на ПРОВЕРИТЬ.').strip()
        elif row is None:
            status = 'NEW'
        elif not current.strip():
            status = 'FILL'
        elif profile_key(current) == profile_key(e.profile):
            status = 'SAME'
        else:
            status, confirmable = 'CONFLICT', True
        result.append(PreviewRow(e.mark, e.profile, current, status, row, e.lines, warning, confirmable))
    return tuple(result)


def new_book_bytes():
    # The application uses the existing Python/openpyxl runtime, without new dependencies.
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Марки'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='38536B')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 24
    for col, width in zip('ABCD', (25, 38, 28, 18)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:D1'
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def sheet_part(data):
    with ZipFile(BytesIO(data)) as z:
        wb = ET.fromstring(z.read('xl/workbook.xml'))
        rid = wb.find('s:sheets/s:sheet', NS).get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        target = next(r.get('Target') for r in rels if r.get('Id') == rid)
        if target.startswith('/'):
            part = target.lstrip('/')
        else:
            part = 'xl/' + target
        if '..' in part.split('/'):
            raise PassportError('Неподдерживаемый путь листа XLSX.')
        return part, z.read(part).decode('utf-8')


def xml_cell(address, value, attributes=''):
    attrs = re.sub(r'\s+(?:r|t)="[^"]*"', '', attributes)
    return f'<c r="{address}"{attrs} t="inlineStr"><is><t xml:space="preserve">{escape(value)}</t></is></c>'


def patch_cells(data, changes, new_rows):
    part, xml = sheet_part(data)
    if '<sheetData>' not in xml or '<row ' not in xml or re.search(r'<\w+:', xml):
        raise PassportError('Неподдерживаемая XML-структура листа. Файл не изменён.')
    for row, value in changes:
        pattern = re.compile(rf'<row\b[^>]*\br="{row}"[^>]*>.*?</row>', re.S)
        match = pattern.search(xml)
        if not match:
            raise PassportError('Не найдена исходная строка XLSX.')
        old = match.group()
        cell = re.search(rf'<c\b([^>]*\br="B{row}"[^>]*)(?:/>|>.*?</c>)', old, re.S)
        if cell:
            updated = old[:cell.start()] + xml_cell(f'B{row}', value, cell[1]) + old[cell.end():]
        else:
            following = re.search(r'<c\b[^>]*\br="[CD]\d+"', old)
            pos = following.start() if following else old.index('</row>')
            updated = old[:pos] + xml_cell(f'B{row}', value) + old[pos:]
        xml = xml[:match.start()] + updated + xml[match.end():]
    if new_rows:
        extra = ''.join(f'<row r="{r}">{xml_cell(f"A{r}", mark)}{xml_cell(f"B{r}", profile)}</row>'
                        for r, mark, profile in new_rows)
        xml = xml.replace('</sheetData>', extra + '</sheetData>', 1)
        last = new_rows[-1][0]
        xml = re.sub(r'(<dimension\b[^>]*\bref=")[^"]*(")', rf'\g<1>A1:D{last}\2', xml, count=1)
        xml = re.sub(r'(<autoFilter\b[^>]*\bref=")A1:D\d+(")', rf'\g<1>A1:D{last}\2', xml, count=1)
    output = BytesIO()
    with ZipFile(BytesIO(data)) as old, ZipFile(output, 'w') as new:
        for info in old.infolist():
            new.writestr(info, xml.encode('utf-8') if info.filename == part else old.read(info.filename))
        new.comment = old.comment
    return output.getvalue()


def verify_patch(before, after, changes, new_rows):
    original, last = read_catalog(before)
    actual, _ = read_catalog(after)
    expected = dict(original)
    for row, value in changes:
        for normalized, (r, mark, profile) in original.items():
            if row == r:
                expected[normalized] = (r, mark, value)
    for r, mark, profile in new_rows:
        expected[key(mark)] = (r, mark, profile)
    if actual != expected:
        raise PassportError('Проверка сохранённых марок/профилей не пройдена. Рабочий файл не заменён.')
    part, old_xml = sheet_part(before)
    _, new_xml = sheet_part(after)
    # Every existing cell except the explicitly changed B cells is byte-identical.
    def cells(xml):
        return {m[1]: m[0] for m in re.finditer(r'<c\b[^>]*\br="([A-Z]+\d+)"[^>]*(?:/>|>.*?</c>)', xml, re.S)}
    old_cells, new_cells = cells(old_xml), cells(new_xml)
    allowed = {f'B{r}' for r, _ in changes}
    for address, value in old_cells.items():
        if address not in allowed and new_cells.get(address) != value:
            raise PassportError(f'Изменилась защищённая клетка {address}. Замена отменена.')
    for r, _, _ in new_rows:
        if f'C{r}' in new_cells or f'D{r}' in new_cells:
            raise PassportError('Ручные столбцы новой марки должны быть пустыми.')
    with ZipFile(BytesIO(before)) as old, ZipFile(BytesIO(after)) as new:
        if old.namelist() != new.namelist():
            raise PassportError('Изменилась структура XLSX-пакета.')
        for name in old.namelist():
            if name != part and old.read(name) != new.read(name):
                raise PassportError('Изменена посторонняя часть XLSX-пакета.')


@contextmanager
def catalog_guard(path):
    # A separate native handle serializes our writers even when the catalog is absent.
    import win32con
    import win32file
    import pywintypes
    lock = path.with_name('.stage3-catalog.lock')
    try:
        handle = win32file.CreateFile(str(lock), win32con.GENERIC_READ | win32con.GENERIC_WRITE,
                                      0, None, win32con.OPEN_ALWAYS,
                                      win32con.FILE_FLAG_DELETE_ON_CLOSE, None)
    except pywintypes.error as exc:
        raise PassportError('Справочник занят импортом или папка недоступна для записи.') from exc
    try:
        yield
    finally:
        handle.Close()


def durable_write(path, data):
    with Path(path).open('xb') as file:
        file.write(data)
        file.flush()
        os.fsync(file.fileno())


class CatalogService:
    def __init__(self, settings):
        self.settings = settings
        self.registry = read_registry(settings.registry_path)

    def assert_registry(self):
        if sha256(self.settings.registry_path) != self.registry.digest:
            raise PassportError('STOP: ВЗМК.xlsx изменился после checkpoint. Закройте окно и начните новую проверку.')

    def preview(self, project, text):
        self.assert_registry()
        if project not in self.registry.projects:
            raise PassportError('Выберите зарегистрированный проект.')
        path = project_folder(self.settings, project) / CATALOG
        if path.is_symlink() or path.resolve().parent != path.parent.resolve():
            raise PassportError('Небезопасный путь справочника.')
        data = path.read_bytes() if path.exists() else None
        existing = read_catalog(data)[0] if data else {}
        result = Preview(project, path, digest(data).hexdigest() if data is not None else None,
                         make_rows(parse_input(text), existing), text)
        self.assert_registry()
        return result

    def apply(self, plan, confirmations=()):
        try:
            self.assert_registry()
            # A stale or altered preview cannot become a write plan.
            if self.preview(plan.project, plan.text) != plan:
                raise PassportError('Справочник или preview изменился. Повторите «Проверить импорт».')
            if any(r.status == 'ERROR' for r in plan.rows):
                raise PassportError('Исправьте ошибки входной таблицы и повторите проверку. Ничего не записано.')
            confirmations = frozenset(confirmations)
            eligible = {r.token for r in plan.rows if r.confirmable}
            if not confirmations <= eligible:
                raise PassportError('Подтверждение относится к недоступной операции.')
            changes, additions = [], []
            counts = dict(added=0, filled=0, changed=0, unchanged=0, check_required=0, skipped=0)
            for r in plan.rows:
                if r.status == 'CHECK_REQUIRED':
                    counts['check_required'] += 1
                active = r.status in ('NEW', 'FILL') or (r.confirmable and r.token in confirmations)
                if active:
                    if r.row is None:
                        additions.append((r.mark, r.incoming))
                        counts['added'] += 1
                    else:
                        changes.append((r.row, r.incoming))
                        counts['filled' if not r.current.strip() else 'changed'] += 1
                elif r.status == 'SAME' or (is_check(r.current) and is_check(r.incoming)):
                    counts['unchanged'] += 1
                else:
                    counts['skipped'] += 1
            if not changes and not additions:
                return counts | {'backup': None, 'path': str(plan.path), 'written': False}
            with locked_file(self.settings.registry_path) as registry_bytes:
                if digest(registry_bytes).hexdigest() != self.registry.digest:
                    raise PassportError('STOP: production изменился после checkpoint.')
                with catalog_guard(plan.path):
                    if plan.path.exists():
                        with locked_file(plan.path, writable=True) as data:
                            return self._commit(plan, data, changes, additions, counts)
                    if plan.catalog_hash is not None:
                        raise PassportError('Справочник удалён после preview. Повторите проверку.')
                    return self._commit(plan, None, changes, additions, counts)
        except PassportError:
            raise
        except OSError as exc:
            raise PassportError('Не удалось безопасно записать справочник. Закройте его в Excel и проверьте права папки. Файл не заменён.') from exc

    def _commit(self, plan, data, changes, additions, counts):
        actual = digest(data).hexdigest() if data is not None else None
        if actual != plan.catalog_hash:
            raise PassportError('Справочник изменился после preview. Повторите проверку.')
        source = data if data is not None else new_book_bytes()
        _, last = read_catalog(source)
        new_rows = [(last + i, mark, profile) for i, (mark, profile) in enumerate(additions, 1)]
        token = uuid4().hex
        backups = plan.path.parent / '.stage3-backups'
        if backups.is_symlink() or backups.resolve().parent != plan.path.parent.resolve():
            raise PassportError('Небезопасный путь резервных копий. Запись остановлена.')
        backups.mkdir(exist_ok=True)
        backup = backups / (f'{plan.path.stem}.{token}.xlsx' if data is not None else f'{plan.path.stem}.{token}.absent.txt')
        durable_write(backup, data if data is not None else b'Catalog absent before creation.\n')
        temporary = plan.path.with_name(f'.stage3-{token}.tmp.xlsx')
        try:
            patched = patch_cells(source, changes, new_rows)
            durable_write(temporary, patched)
            saved = temporary.read_bytes()
            verify_patch(source, saved, changes, new_rows)
            if data is not None:
                atomic_replace(temporary, plan.path)
            else:
                # Windows rename is atomic and refuses an existing target.
                os.rename(temporary, plan.path)
            return counts | {'backup': str(backup), 'path': str(plan.path), 'written': True,
                             'sha256': digest(saved).hexdigest()}
        finally:
            if temporary.exists():
                temporary.unlink()
