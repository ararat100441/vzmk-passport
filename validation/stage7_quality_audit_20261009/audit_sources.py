import sys, os, json, hashlib, dataclasses
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT))
from passport_app.registry import read_registry,select_project
import openpyxl
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def save(n,x): (OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def walk(root,excluded=()):
    for base,dirs,files in os.walk(root,followlinks=False):
        dirs[:]=[d for d in dirs if d not in excluded and not (Path(base)/d).is_symlink() and not (Path(base)/d).is_junction()]
        for n in files:
            p=Path(base)/n
            if not p.is_symlink() and not n.startswith('~$'): yield p
registry=read_registry(ROOT/'DEV/ВЗМК_DEV.xlsx')
project=select_project(registry,'Гараж Казахстан')
save('registry_selection.json',{'registry':dataclasses.asdict(registry),'project':dataclasses.asdict(project)})
book=openpyxl.load_workbook(registry.path,data_only=False)
values=openpyxl.load_workbook(registry.path,data_only=True)
s=book[registry.sheet]; v=values[registry.sheet]
header=s[s.tables[registry.table].ref.split(':')[0]].row
rows=[]
for e in registry.projects[project.name]:
    rows.append({'row':e['row'],'cells':[{'cell':c.coordinate,'header':s.cell(header,c.column).value,'formula_or_value':c.value,'cached':v[c.coordinate].value,'format':c.number_format} for c in s[e['row']] if c.value is not None]})
save('registry_rows.json',rows)
book.close();values.close()
projectroot=ROOT.parent/'Гараж казахстан/Паспорт'
files=list(walk(projectroot))
quality=[p for p in files if p.suffix.lower() in ('.doc','.docx') and 'качеств' in p.stem.lower()]
inventory=[{'path':str(p),'format':p.suffix,'size':p.stat().st_size,'sha256':sha(p)} for p in quality]
save('quality_candidates.json',inventory)
save('project_file_inventory.json',[{'path':str(p),'size':p.stat().st_size} for p in files])
selected=[p for p in quality if p.parent==projectroot and p.name.startswith(project.first_number.replace('/','_')+' - ')]
if len(selected)!=1: raise RuntimeError('AMBIGUOUS_FIRST_QUALITY_DOCUMENT_STOP')
save('selected_quality.json',inventory[[p['path'] for p in inventory].index(str(selected[0]))])
baseline={str(p):{'size':p.stat().st_size,'sha256':sha(p)} for p in walk(ROOT,('validation','.codebase-index','node_modules','__pycache__'))}
for p in files:
    if p.suffix.lower() not in ('.pdf',): baseline[str(p)]={'size':p.stat().st_size,'sha256':sha(p)}
save('source_before.json',baseline)
print(json.dumps({'project':dataclasses.asdict(project),'quality_candidates':inventory,'baseline_count':len(baseline)},ensure_ascii=False,default=str,indent=2))
