import json,hashlib,zipfile,re
from pathlib import Path
from xml.etree import ElementTree as E
import openpyxl
OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[1]
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def save(n,x):(OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str),encoding='utf8')
path=Path(load(OUT/'selected_quality.json')['path']);W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}';NS={'w':W[1:-1]}
with zipfile.ZipFile(path) as z:
    xml=E.fromstring(z.read('word/document.xml'));tables=xml.findall('.//w:tbl',NS)
    result={'media_entries':[n for n in z.namelist() if n.startswith('word/media/')],'external_relationships':[],'breaks':[],'row_properties':[],'numbered_paragraphs':[],'protected_point12':{}}
    for n in z.namelist():
        if n.endswith('.rels'):
            for r in E.fromstring(z.read(n)):
                if r.attrib.get('TargetMode')=='External':result['external_relationships'].append({'part':n,**r.attrib})
    for p in xml.findall('.//w:p',NS):
        t=''.join(p.itertext())
        props=p.find('w:pPr',NS)
        if props is not None and props.find('w:numPr',NS) is not None:result['numbered_paragraphs'].append({'text':''.join(p.iter(W+'t')) if False else ''.join(x.text or '' for x in p.iter(W+'t')),'properties':E.tostring(props,encoding='unicode')})
        for b in p.findall('.//w:br',NS):result['breaks'].append({'kind':'br','type':b.get(W+'type'),'near_text':''.join(x.text or '' for x in p.iter(W+'t'))})
        if props is not None and props.find('w:pageBreakBefore',NS) is not None:result['breaks'].append({'kind':'pageBreakBefore','xml':E.tostring(props.find('w:pageBreakBefore',NS),encoding='unicode'),'text':''.join(x.text or '' for x in p.iter(W+'t'))})
    rows=tables[0].findall('w:tr',NS)
    block=b''.join(E.tostring(r,encoding='utf8') for r in rows[40:49])
    result['protected_point12']={'table':1,'rows':[41,49],'word_range':[2036,2702],'xml_sha256':hashlib.sha256(block).hexdigest(),'rows_xml':[E.tostring(r,encoding='unicode') for r in rows[40:49]]}
    for ti,t in enumerate(tables,1):
        for ri,r in enumerate(t.findall('w:tr',NS),1):
            result['row_properties'].append({'table':ti,'row':ri,'properties':E.tostring(r.find('w:trPr',NS),encoding='unicode') if r.find('w:trPr',NS) is not None else None,'cells':[{'gridSpan':c.find('w:tcPr/w:gridSpan',NS).get(W+'val') if c.find('w:tcPr/w:gridSpan',NS) is not None else '1','vMerge':E.tostring(c.find('w:tcPr/w:vMerge',NS),encoding='unicode') if c.find('w:tcPr/w:vMerge',NS) is not None else None} for c in r.findall('w:tc',NS)]})
save('ooxml_structure.json',result)
pr=path.parent
files=[pr/'1_2026 - Приложение № 1.xlsx',pr/'9 от 06.10.2026/9_2026 - Приложение № 1.xlsx',pr/'Акт № 4-2026.xlsx',pr/'9 от 06.10.2026/Акт № 9-2026.xlsx',pr/'9 от 06.10.2026/Акт № 9-1-2026.xlsx']
x=[]
for p in files:
    b=openpyxl.load_workbook(p,data_only=True,read_only=True)
    x.append({'path':str(p),'sha256':sha(p),'sheets':[{'name':s.title,'rows':s.max_row,'cols':s.max_column,'cells':[{'cell':c.coordinate,'value':c.value} for row in s for c in row if c.value is not None]} for s in b]});b.close()
save('historical_xlsx.json',x)
pairs=[]
for p in (ROOT/'DEV/Stage6_Выход').rglob('PAIR.json'):
    d=load(p);pairs.append({'receipt_path':str(p),'receipt_sha256':sha(p),'receipt':d,'outputs':[{'path':str(p.parent/n),'expected_sha256':h,'actual_sha256':sha(p.parent/n) if (p.parent/n).is_file() else None,'match':(p.parent/n).is_file() and sha(p.parent/n)==h} for n,h in d.get('sha256',{}).items()]})
save('existing_akz_pairs.json',pairs)
selection=load(OUT/'registry_selection.json')
print(json.dumps({'numbered_paragraphs':result['numbered_paragraphs'],'breaks':result['breaks'],'pair_count':len(pairs),'all_pair_hashes_match':all(o['match'] for p in pairs for o in p['outputs']),'registry_sha256':selection['registry']['digest']},ensure_ascii=False,indent=2))
