"""Narrow Stage 6 evidence extraction. Never saves a workbook or calls COM."""
import hashlib
import json
import sys
from dataclasses import asdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

sys.dont_write_bytecode = True
OUT = Path(__file__).resolve().parent
BASE = OUT.parent.parent
sys.path.insert(0, str(BASE))

from passport_app.config import Settings
from passport_app.registry import read_registry, select_project
from passport_app.catalog import project_folder, key
from passport_app.appendix import inspect_template
from passport_app.vik import _passport_name_number


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2,
                                      default=str) + '\n', encoding='utf-8')


def listing(root):
    return [{'name': p.name, 'path': str(p.resolve()),
             'kind': 'directory' if p.is_dir() else 'file',
             'symlink': p.is_symlink()}
            for p in sorted(root.iterdir(), key=lambda p: p.name.casefold())]


def extract(path):
    before = digest(path)
    values = openpyxl.load_workbook(path, data_only=True, read_only=True,
                                    keep_links=False)
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True,
                                      keep_links=False)
    try:
        sheets = []
        for sheet in values.worksheets:
            cells = {}
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 51),
                                       min_col=1, max_col=min(sheet.max_column, 12)):
                for cell in row:
                    if cell.value is not None:
                        cells[cell.coordinate] = cell.value
            fs = formulas[sheet.title]
            formula_cells = {cell.coordinate: cell.value
                             for row in fs.iter_rows(min_row=1, max_row=min(fs.max_row, 51),
                                                     min_col=1, max_col=min(fs.max_column, 12))
                             for cell in row if cell.data_type == 'f'}
            sheets.append({'sheet': sheet.title, 'max_row': sheet.max_row,
                           'max_column': sheet.max_column, 'cells': cells,
                           'formulas': formula_cells})
        after = digest(path)
        if after != before:
            raise RuntimeError(f'File changed during read: {path}')
        return {'path': str(path.resolve()), 'name': path.name,
                'extension': path.suffix, 'sha256': before,
                'sha256_after': after, 'sheets': sheets}
    finally:
        values.close()
        formulas.close()


data = json.loads((BASE / 'settings.json').read_text(encoding='utf-8-sig'))
settings = Settings(documents_root=BASE.parent.resolve(),
                    registry_path=(BASE / 'DEV' / 'ВЗМК_DEV.xlsx').resolve(),
                    projects=data.get('projects', {}))
project_dir = project_folder(settings, 'Гараж Казахстан')
passport_root, configured_appendix = settings.project_paths('Гараж Казахстан')

# Preservation scope: product code, tests, launchers/checkpoints, exact DEV input,
# catalog, previous audit report, and all first-passport root files.
# No historical subfolder, other project, production registry or common form opened.
guard = set(p.resolve() for p in BASE.iterdir() if p.is_file())
for name in ('passport_app', 'profiles_stage4', 'uzk_stage5', 'tests'):
    directory = BASE / name
    if directory.exists():
        # Product modules are immediate files. Never walk dependency junctions.
        guard.update(p.resolve() for p in directory.iterdir() if p.is_file()
                     and p.suffix != '.pyc')
guard.update(p.resolve() for p in passport_root.iterdir() if p.is_file())
guard.update((settings.registry_path, project_dir / 'Справочник марок.xlsx',
              BASE / 'validation' / 'stage6_akz_audit_20261008_135318' / 'STAGE6_AKZ_AUDIT.md'))
before = {str(p): {'sha256': digest(p), 'bytes': p.stat().st_size}
          for p in sorted(guard, key=str)}
write('source_manifest_before.json', before)
root_before = listing(passport_root)

registry = read_registry(settings.registry_path)
project = select_project(registry, 'Гараж Казахстан')
if project.first_number != '1/2026':
    raise RuntimeError(f'Unexpected first passport: {project.first_number}')
layout, header, obj = inspect_template(configured_appendix, project)
write('registry_first_passport.json', {
    'timestamp_samara': datetime.now(timezone(timedelta(hours=4))).isoformat(),
    'registry_path': str(registry.path), 'registry_sha256': registry.digest,
    'registry_sheet': registry.sheet, 'registry_table': registry.table,
    'project': asdict(project), 'project_folder': str(project_dir),
    'passport_root': str(passport_root),
    'first_passport_anchor': {'path': str(configured_appendix),
                             'sha256': digest(configured_appendix),
                             'header_cell': header, 'object_cell': obj,
                             'layout': asdict(layout)},
    'first_number_directory_candidates': [p for p in root_before
        if p['kind'] == 'directory' and _passport_name_number(p['name']) is not None
        and _passport_name_number(p['name'])[0] == 1],
    'production_registry_opened': False})

write('passport_root_listing.json', root_before)
root_xlsx = [p for p in passport_root.iterdir()
             if p.is_file() and p.suffix.lower() == '.xlsx' and not p.name.startswith('~$')]
extracted = [extract(p) for p in sorted(root_xlsx, key=lambda p: p.name)]
write('first_passport_workbooks.json', extracted)

catalog_path = project_dir / 'Справочник марок.xlsx'
book = openpyxl.load_workbook(catalog_path, read_only=True, data_only=False,
                              keep_links=False)
try:
    if book.sheetnames != ['Марки']:
        raise RuntimeError(f'Unexpected catalog sheets: {book.sheetnames}')
    sheet = book['Марки']
    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    if headers != ['Марка', 'Профиль', '№ АКЗ']:
        raise RuntimeError(f'Unexpected catalog headers: {headers}')
    rows = [{'row': row[0].row, 'mark': row[0].value,
             'profile': row[1].value, 'akz': row[2].value,
             'akz_type': row[2].data_type}
            for row in sheet.iter_rows(min_row=2, max_col=3)
            if row[0].value is not None]
    write('catalog_current_assessment.json', {
        'path': str(catalog_path), 'sha256': digest(catalog_path),
        'sheet': sheet.title, 'range': f'A1:C{sheet.max_row}', 'headers': headers,
        'mark_count': len(rows), 'unique_marks': len({key(r['mark']) for r in rows}),
        'akz_blank_count': sum(r['akz'] is None or r['akz'] == '' for r in rows),
        'akz_nonblank': [r for r in rows if r['akz'] is not None and r['akz'] != ''],
        'catalog_saved': False, 'test_copy_created': False})
finally:
    book.close()

after = {str(p): {'sha256': digest(p), 'bytes': p.stat().st_size}
         for p in sorted(guard, key=str)}
changed = [p for p in before if before[p] != after[p]]
write('source_manifest_after.json', after)
write('source_integrity.json', {
    'guarded_files': len(before), 'changed': changed,
    'passport_root_listing_unchanged': root_before == listing(passport_root),
    'workbook_saves': 0, 'native_excel_or_word_calls': 0,
    'production_registry_accesses': 0, 'network_calls': 0,
    'method': 'existing read-only selectors; read-only openpyxl; SHA-256 before/after'})
if changed or root_before != listing(passport_root):
    raise RuntimeError('Source or directory changed during check')
print(json.dumps({'first': project.first_number, 'row': project.first_row,
                  'passport_root': str(passport_root),
                  'root_workbooks': [{'name': x['name'], 'sha256': x['sha256']}
                                     for x in extracted],
                  'guarded_files': len(before), 'changed': changed},
                 ensure_ascii=False, indent=2))
