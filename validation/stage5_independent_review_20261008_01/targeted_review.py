"""Review harness only. Product, accepted inputs and reviewed DOCX are never written."""
import importlib.util, json, os, shutil, sys, unittest, tempfile
from contextlib import ExitStack, contextmanager
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock, patch
from types import SimpleNamespace
from zipfile import ZipFile
from lxml import etree as E
from openpyxl import load_workbook

RUN=Path(__file__).resolve().parent
ROOT=RUN.parents[1]
OLD=ROOT/'validation/stage5_implementation_20261008_105100'
(RUN/'tmp').mkdir(exist_ok=True)
tempfile.tempdir=str(RUN/'tmp')
os.environ['TEMP']=os.environ['TMP']=str(RUN/'tmp')
sys.path.insert(0,str(ROOT))
import uzk_stage5.inputs as I
import uzk_stage5.core as C
import uzk_stage5.word as W
from profiles_stage4.model import calculate_uzk
from profiles_stage4.storage import load_catalog
from profiles_stage4.offline import no_network

def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def write(n,v): (RUN/n).write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str),encoding='utf8')

access=[]; active=[True]
def guard(event,args):
    if event!='open' or not active[0] or not isinstance(args[0],(str,bytes)): return
    p=Path(os.fsdecode(args[0])).resolve()
    if p.name.casefold() in ('взмк.xlsx','узк.xlsm'):
        raise AssertionError('PRODUCTION_READ_FORBIDDEN: '+str(p))
    mode=args[1]; flags=args[2]
    writing=(isinstance(mode,str) and any(c in mode for c in 'wax+')) or (isinstance(flags,int) and flags & (os.O_WRONLY|os.O_RDWR|os.O_CREAT|os.O_TRUNC))
    if writing and not p.is_relative_to(RUN): raise AssertionError('WRITE_OUTSIDE_REVIEW: '+str(p))
    access.append(dict(path=str(p),mode=mode,flags=flags,writing=bool(writing)))
sys.addaudithook(guard)

# Bind the existing suite to this NEW evidence folder. Its actual-document
# regeneration test is deliberately excluded; the exact old DOCX stays the target.
for n in ('final_run.json','protected_before.json'):
    shutil.copyfile(OLD/n,RUN/n)
shutil.copyfile(RUN/'native_exact.json',RUN/'native_final_v3.json')
os.environ['STAGE5_VALIDATION_DIR']=str(RUN)
spec=importlib.util.spec_from_file_location('reviewed_stage5_tests',ROOT/'tests/test_stage5.py')
original=importlib.util.module_from_spec(spec);spec.loader.exec_module(original)
methods=unittest.defaultTestLoader.getTestCaseNames(original.Stage5Tests)
selected=[n for n in methods if n!='test_32_33_access_boundary_and_network']
suite=unittest.TestSuite(original.Stage5Tests(n) for n in selected)
if os.environ.get('REVIEW_RESUME')=='1':
    saved=read(RUN/'rerun_35.json');assert saved['successful']
    result=SimpleNamespace(testsRun=saved['tests_run'],wasSuccessful=lambda:True)
else:
    with (RUN/'rerun_35.log').open('w',encoding='utf8') as log:
        result=unittest.TextTestRunner(stream=log,verbosity=2).run(suite)
    write('rerun_35.json',dict(tests_run=result.testsRun,successful=result.wasSuccessful(),failures=[(str(t),v) for t,v in result.failures],errors=[(str(t),v) for t,v in result.errors],skipped=result.skipped,excluded=['test_32_33_access_boundary_and_network: would regenerate another actual 9/2026 document; replaced by fresh read-only preview + isolated generation probes']))

fresh=C.preview('Гараж Казахстан','9/2026',read(OLD/'final_run.json')['preview']['vom']['path'])
fresh.verify()
write('fresh_preview.json',dict(data=fresh.data,seal=fresh.seal,equal_to_implementation_seal=fresh.seal==read(OLD/'final_run.json')['preview_seal']))

# Independent native source reading: known verified B/D/I header locations,
# Excel SpecialCells(visible) instead of the product's row.Hidden loop.
import pythoncom, win32com.client
pythoncom.CoInitialize();app=book=None
try:
    app=win32com.client.DispatchEx('Excel.Application')
    app.Visible=False;app.DisplayAlerts=False;app.AutomationSecurity=3
    app.EnableEvents=False;app.AskToUpdateLinks=False
    book=app.Workbooks.Open(fresh.data['vom']['path'],UpdateLinks=0,ReadOnly=True,AddToMru=False)
    sheet=book.Worksheets(fresh.data['vom']['sheet'])
    assert str(sheet.Range('B2').Value2).strip()=='Отправочная марка'
    assert str(sheet.Range('D2').Value2).strip()=='Количество'
    assert sheet.Range('I3').Value2=='Всех'
    # End(xlUp) skips trailing filtered-out rows. Bound the stored region, then
    # independently identify data cells and actual visible cells below.
    last=int(sheet.UsedRange.Row+sheet.UsedRange.Rows.Count-1)
    vals=sheet.Range(f'B4:I{last}').Value2
    visible={int(c.Row) for area in sheet.Range(f'B4:B{last}').SpecialCells(12).Areas for c in area.Cells}
    rows=[];hidden=[]
    for rn,r in enumerate(vals,4):
        if not isinstance(r[0],str) or not r[0].strip() or r[0].strip().casefold() in ('итого','всего'):continue
        if rn not in visible:hidden.append(rn);continue
        q=Decimal(str(r[2]));assert q>0 and q==q.to_integral()
        rows.append(dict(row=rn,mark=r[0],quantity=int(q),mass_kg=str(r[7])))
    mass=sum((Decimal(r['mass_kg']) for r in rows),Decimal(0))
    independent=dict(method='Native Excel B/D/I cells + SpecialCells(12)',read_only=bool(book.ReadOnly),last_mark_row=last,visible_count=len(rows),hidden_count=len(hidden),mass=mass,rows=rows,hidden_rows=hidden,autofilter=sheet.AutoFilter.Range.Address,macros=app.AutomationSecurity,enable_events=app.EnableEvents,links=book.LinkSources(1),connections=book.Connections.Count)
    write('independent_native_vom.json',independent)
    assert len(rows)==104 and len(hidden)==397 and mass.quantize(Decimal('.01'))==Decimal('15822.14'), (len(rows),len(hidden),str(mass),last,len(visible))
    assert [(r['row'],r['mark'],r['quantity'],Decimal(r['mass_kg'])) for r in rows]==[(r['row'],r['mark'],r['quantity'],Decimal(r['mass_kg'])) for r in fresh.data['vom']['rows']]
    write('independent_native_vom.json',independent)
finally:
    if book is not None:book.Close(False)
    if app is not None:app.Quit()
    pythoncom.CoUninitialize()

# Independent reconciliation from workbook cells + native visible rows + accepted
# Stage 4 lookup. Expected one-item geometry is a separate acceptance oracle.
inputs=read(RUN/'independent_inputs.json'); catalog=load_catalog(I.PROFILES)
oracle={'Бп2-1':[(160,5),(128,8.4)],'Пс1-9':[(400,4)],'Пс1-10':[(400,4)],'Пс2-16':[(480,3)]}
trace=[]
for mark,segments in oracle.items():
    vom=[r for r in rows if r['mark']==mark];vq=sum(r['quantity'] for r in vom)
    reg=[(rn,r) for rn,r in inputs['registry_rows'] if r[0]=='Гараж Казахстан' and r[1]==mark];assert len(reg)==1
    remaining=1 if reg[0][1][2] is None else int(reg[0][1][2]);q=min(vq,remaining)
    cat=[r for r in inputs['catalog_rows'] if r['values'][0]==mark];assert len(cat)==1
    raw=cat[0]['values'][1];lookup=catalog.lookup(raw);single=calculate_uzk(lookup)
    assert lookup['status']=='FOUND'
    assert [(r['L'],r['S']) for r in single['rows']]==segments
    final=[dict(L=l*q,S=s,frequency_mhz=5,angle_degrees=70) for l,s in segments]
    product=next(m for m in fresh.data['marks'] if m['mark']==mark)
    assert product['Q']==q and [(Decimal(r['L']),r['S']) for r in product['segments']]==[(r['L'],r['S']) for r in final]
    trace.append(dict(mark=mark,vom_rows=vom,vom_quantity=vq,registry_row=reg[0][0],remaining=remaining,Q=q,catalog_row=cat[0]['row'],raw_profile=raw,lookup_status=lookup['status'],normalized_key=lookup['key'],standard=lookup['profile']['standard'],rule=single['rule'],source_geometry=lookup['profile']['geometry'],single=segments,final=final))
write('independent_four_mark_trace.json',trace)

# A complete sandbox uses exact copied inputs, including a copied DOCX template;
# patch only module globals in this test process to redirect every input/output.
BASE=RUN/'isolated_base';BASE.mkdir()
for p in C.product_paths()+sorted((ROOT/'profiles_stage4').glob('*.py'))+[ROOT/'passport_app/catalog.py']:
    out=BASE/p.relative_to(ROOT);out.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(p,out)
for p in [I.REGISTRY,I.UZK,I.PROFILES,I.CHECKPOINT,*[I.PROFILES.parent/n for n in ('profiles.user.json','profiles.catalog.json','profiles.build.json')]]:
    out=BASE/p.relative_to(ROOT);out.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(p,out)
project=BASE/'Документы/Гараж казахстан';passport=project/'Паспорт';passport.mkdir(parents=True)
shutil.copyfile(fresh.data['template'],passport/'1_2026 - УЗК.docx')
shutil.copyfile(fresh.data['mark_catalog'],project/'Справочник марок.xlsx')
shutil.copyfile(fresh.data['vom']['path'],project/'source.xls')
(BASE/'settings.json').write_text(json.dumps({'documents_root':'Документы','projects':{'Гараж Казахстан':{'passport_root':'Гараж казахстан/Паспорт'}}},ensure_ascii=False),encoding='utf8')

@contextmanager
def sandbox(name):
    root=RUN/'scenarios'/name;root.parent.mkdir(exist_ok=True);shutil.copytree(BASE,root)
    paths={'ROOT':root,'REGISTRY':root/'DEV/ВЗМК_DEV.xlsx','UZK':root/'DEV/УЗК_DEV.xlsm','PROFILES':root/'DEV/Справочники/Профили.xlsx','CHECKPOINT':root/'DEV/DEV_UZK_CHECKPOINT.json'}
    cp=read(paths['CHECKPOINT']);cp['dev_path']=str(paths['UZK']);cp['dev_initial_sha256']=I.sha(paths['UZK'])
    paths['CHECKPOINT'].write_text(json.dumps(cp,ensure_ascii=False),encoding='utf8')
    pins={paths[k]:I.sha(paths[k]) for k in ('REGISTRY','UZK','PROFILES')}
    original_passports=I.passports;original_registry=I.registry_rows
    with ExitStack() as stack:
        for mod in (I,C):
            for k,v in paths.items():stack.enter_context(patch.object(mod,k,v))
        stack.enter_context(patch.object(I,'PINS',pins))
        stack.enter_context(patch.object(I,'passports',side_effect=lambda:original_passports(paths['REGISTRY'])))
        stack.enter_context(patch.object(C,'registry_rows',side_effect=lambda:original_registry(paths['UZK'])))
        def refresh_pins():
            for p in pins:pins[p]=I.sha(p)
            c=read(paths['CHECKPOINT']);c['dev_initial_sha256']=pins[paths['UZK']];paths['CHECKPOINT'].write_text(json.dumps(c,ensure_ascii=False),encoding='utf8')
        yield root,root/'Документы/Гараж казахстан/source.xls',refresh_pins

probes=[]
def probe(name,fn,expected='PASS'):
    try:
        detail=fn();probes.append(dict(name=name,status='PASS',expected=expected,detail=detail))
    except Exception as e:
        probes.append(dict(name=name,status='FAIL',expected=expected,error=type(e).__name__+': '+str(e)))
def stops(fn,pattern):
    try:fn()
    except Exception as e:
        assert pattern in str(e),(pattern,str(e));return str(e)
    raise AssertionError('Expected STOP: '+pattern)

def missing_project():
    with sandbox('missing_project') as (root,vom,r):
        settings=read(root/'settings.json');settings['projects']['Гараж Казахстан']['passport_root']='MISSING/Паспорт'
        (root/'settings.json').write_text(json.dumps(settings,ensure_ascii=False),encoding='utf8')
        return stops(lambda:C.preview('Гараж Казахстан','9/2026',vom),'PROJECT_PATH_INVALID')
probe('missing_project_STOP',missing_project)
def ambiguous_passport():
    records=I.passports();selected=next(r for r in records if r['project']=='Гараж Казахстан' and r['number']=='9/2026')
    with patch.object(I,'passports',return_value=[selected,deepcopy(selected)]):
        return stops(lambda:I.select_passport('Гараж Казахстан','9/2026'),'PASSPORT_AMBIGUOUS_OR_MISSING')
probe('ambiguous_passport_STOP',ambiguous_passport)
def missing_template():
    with sandbox('missing_template') as (root,vom,r):
        (root/'Документы/Гараж казахстан/Паспорт/1_2026 - УЗК.docx').unlink()
        return stops(lambda:C.preview('Гараж Казахстан','9/2026',vom),'TEMPLATE_MISSING')
probe('missing_template_STOP',missing_template)
def missing_registry():
    with sandbox('missing_registry') as (root,vom,r):
        (root/'DEV/УЗК_DEV.xlsm').unlink()
        return stops(lambda:C.preview('Гараж Казахстан','9/2026',vom),'CHECKPOINT_CHANGED')
probe('missing_registry_STOP_not_no_UZK',missing_registry)
def unreadable_registry():
    with sandbox('unreadable_registry') as (root,vom,r):
        (root/'DEV/УЗК_DEV.xlsm').write_bytes(b'not a workbook');r()
        return stops(lambda:C.preview('Гараж Казахстан','9/2026',vom),'File is not a zip file')
probe('unreadable_registry_STOP_not_no_UZK',unreadable_registry)

def mutate_after_preview(which):
    with sandbox('drift_'+which) as (root,vom,r):
        p=C.preview('Гараж Казахстан','9/2026',vom)
        paths={'vom':vom,'registry':root/'DEV/УЗК_DEV.xlsm','catalog':root/'Документы/Гараж казахстан/Справочник марок.xlsx','profile_database':root/'DEV/Справочники/Профили.xlsx','passport_registry':root/'DEV/ВЗМК_DEV.xlsx','template':root/'Документы/Гараж казахстан/Паспорт/1_2026 - УЗК.docx'}
        with paths[which].open('ab') as f:f.write(b'isolated mutation')
        out=root/'validation/drift_out'
        detail=stops(lambda:C.generate(p,out),'INPUT_CHANGED');assert not out.exists();return detail
for which in ('vom','registry','catalog','profile_database','passport_registry','template'):
    probe(which+'_changed_after_preview_STOP',lambda w=which:mutate_after_preview(w))

def no_uzk(unrelated_conflict=False):
    with sandbox('no_uzk_'+str(unrelated_conflict)) as (root,vom,refresh):
        path=root/'DEV/УЗК_DEV.xlsm';b=load_workbook(path,keep_vba=True)
        s=b['УЗК']
        for row in s.iter_rows(min_row=2,max_col=1):row[0].value='OTHER_VALID_PROJECT'
        b.save(path);b.close();refresh()
        if unrelated_conflict:
            path=root/'Документы/Гараж казахстан/Справочник марок.xlsx';b=load_workbook(path);s=b.active
            s.append(['UNRELATED_MARK','[16П',None,None]);s.append(['UNRELATED_MARK','Гн.[140X60X4',None,None]);b.save(path);b.close()
        before=I.sha(root/'DEV/УЗК_DEV.xlsm')
        p=C.preview('Гараж Казахстан','9/2026',vom);assert p.data['marks']==[]
        out=root/'validation/no_uzk_out';assert C.generate(p,out) is None and not out.exists()
        assert I.sha(root/'DEV/УЗК_DEV.xlsm')==before
        return 'No marks, no output, no registry mutation'
probe('full_no_UZK_pipeline_clean_catalog',no_uzk)
probe('full_no_UZK_pipeline_unrelated_catalog_conflict',lambda:no_uzk(True),'No false error for unrelated catalog rows')

def writer_failure(partial=False):
    with sandbox('writer_failure_'+str(partial)) as (root,vom,refresh):
        plan=C.preview('Гараж Казахстан','9/2026',vom)
        # Distinct fixture-only identity; never regenerate the real DEV artifact.
        data=deepcopy(plan.data);data['passport']['number']='777/2099';data['passport']['object']='ИЗОЛИРОВАННАЯ ТЕСТОВАЯ ФИКСТУРА'
        plan=C.Preview(data,C.digest(data));out=root/'validation/output';name='777_2099 - УЗК (ТЕСТ).docx'
        real_open=Path.open
        class FailStream:
            def __init__(self,s):self.s=s
            def __enter__(self):return self
            def __exit__(self,*a):self.s.close()
            def write(self,b):self.s.write(b[:50]);self.s.flush();raise OSError('Injected write failure in isolated fixture')
        def open_failure(path,*a,**kw):
            if path.name==name and a and a[0]=='xb':
                if partial:return FailStream(real_open(path,*a,**kw))
                raise PermissionError('Injected locked output / Access denied')
            return real_open(path,*a,**kw)
        with patch.object(Path,'open',open_failure):detail=stops(lambda:C.generate(plan,out),'Injected')
        assert not (out/name).exists();return detail
probe('output_locked_save_failure_no_partial_DOCX',writer_failure)
probe('partial_Word_package_save_failure_cleanup',lambda:writer_failure(True))

def isolated_generation():
    with sandbox('isolated_generation') as (root,vom,refresh):
        p=C.preview('Гараж Казахстан','9/2026',vom);data=deepcopy(p.data)
        data['passport']['number']='777/2099';data['passport']['object']='ИЗОЛИРОВАННАЯ ТЕСТОВАЯ ФИКСТУРА'
        p=C.Preview(data,C.digest(data))
        sources={path:I.sha(path) for path in data['fingerprints']}
        output=C.generate(p,root/'validation/output')
        assert all(I.sha(path)==h for path,h in sources.items())
        h=I.sha(output);detail=stops(lambda:C.generate(p,output.parent),'OUTPUT_EXISTS');assert I.sha(output)==h
        return dict(output=output,sha256=h,source_mutations=0,second_generation=detail)
probe('isolated_generate_success_source_integrity_existing_output',isolated_generation)

def excel_open_failure():
    fixture=RUN/'fixtures/com_failure.xls';fixture.write_bytes(b'fake for mocked COM')
    app=MagicMock();app.Workbooks.Open.side_effect=RuntimeError('injected open failure')
    with patch('win32com.client.DispatchEx',return_value=app),patch('pythoncom.CoInitialize') as init,patch('pythoncom.CoUninitialize') as end:
        detail=stops(lambda:I.visible_vom(fixture),'VOM_READ_FAILED');app.Quit.assert_called_once();init.assert_called_once();end.assert_called_once()
    assert app.AutomationSecurity==3 and app.EnableEvents is False and app.AskToUpdateLinks is False
    assert app.Workbooks.Open.call_args.kwargs['ReadOnly'] is True and app.Workbooks.Open.call_args.kwargs['UpdateLinks']==0
    return detail
probe('Excel_COM_open_failure_cleanup_and_security',excel_open_failure)

def gui_open_output():
    from uzk_stage5.gui import App
    app=MagicMock();app.output=Path(read(OLD/'EVIDENCE_MANIFEST.json')['final_output']['path'])
    with patch('uzk_stage5.gui.os.startfile') as start:
        App.open_word(app);start.assert_called_once_with(str(app.output))
    return 'Open button dispatches the exact existing output; actual Word normal-open separately verified'
probe('GUI_open_output_exact_dispatch',gui_open_output)

def regression():
    import importlib
    with no_network() as net:
        for name in ('main','stage2','stage3','stage4'):importlib.import_module(name)
    assert net==[];I.verify_pins();return 'Imports + checkpoint only, no prior workflows'
probe('Stage_1_4_imports_and_checkpoint',regression)

write('supplemental_probes.json',probes)
active[0]=False
write('access_guard.json',dict(production_reads=0,writes_outside_review=0,python_file_events=access,write_events=[e for e in access if e['writing']],native_COM_scope='Excel explicit selected XLS ReadOnly/links0/security3; Word explicit reviewed DOCX ReadOnly, new owned process; native source writes not used'))
summary=dict(original_suite_rerun=result.testsRun,original_suite_success=result.wasSuccessful(),supplemental_count=len(probes),supplemental_pass=sum(p['status']=='PASS' for p in probes),supplemental_fail=[p for p in probes if p['status']!='PASS'],four_mark_units=sum(r['Q'] for r in trace),fresh_preview_matches_original=fresh.seal==read(OLD/'final_run.json')['preview_seal'])
write('review_test_summary.json',summary)
print(json.dumps(summary,ensure_ascii=False,indent=2))
