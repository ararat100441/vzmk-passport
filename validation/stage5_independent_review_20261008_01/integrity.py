import hashlib, json, os, sys
from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
from lxml import etree as E

RUN = Path(__file__).resolve().parent
ROOT = RUN.parents[1]
OLD = ROOT/'validation/stage5_implementation_20261008_105100'
def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(n,v): (RUN/n).write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str),encoding='utf8')
m=read(OLD/'EVIDENCE_MANIFEST.json')
evidence=[]
for n,meta in m['evidence_files'].items():
    p=(OLD/n).resolve()
    assert p.is_relative_to(OLD)
    evidence.append(dict(path=n,match=p.is_file() and sha(p)==meta['sha256'] and p.stat().st_size==meta['bytes'],mtime_utc=datetime.fromtimestamp(p.stat().st_mtime,timezone.utc).isoformat()))
actual_files={str(p.relative_to(OLD)) for p in OLD.rglob('*') if p.is_file() and p.name!='EVIDENCE_MANIFEST.json'}
before=m['product_tested_before']; after=m['product_tested_after']; prior_current=m['product_current']
current={p:sha(p) for p in before}
changed=read(OLD/'changed_files.json')
all_product=[dict(path=p['path'],sha256=sha(p['path']),match=sha(p['path'])==p['sha256']) for p in changed['new_product_files']]
protected=read(OLD/'protected_before.json')
assert all(Path(p).name.casefold() not in ('взмк.xlsx','узк.xlsm') for p in protected)
write('review_initial_hashes.json',{p:sha(p) for p in {*protected,*current,*[p['path'] for p in all_product],str(OLD/'EVIDENCE_MANIFEST.json'),m['final_output']['path']}})
manifest_mtime=(OLD/'EVIDENCE_MANIFEST.json').stat().st_mtime
integrity=dict(output=m['final_output'],output_match=sha(m['final_output']['path'])==m['final_output']['sha256'],evidence_files=evidence,evidence_count=len(evidence),all_evidence_match=all(r['match'] for r in evidence),unlisted_files=sorted(actual_files-set(m['evidence_files'])),listed_missing=sorted(set(m['evidence_files'])-actual_files),tested_before_after_current_equal=before==after==prior_current==current,product_set=all_product,protected_count=len(protected),protected_current_equal=all(sha(p)==h for p,h in protected.items()),manifest_last=all((OLD/n).stat().st_mtime<=manifest_mtime for n in m['evidence_files']),manifest_sha256=sha(OLD/'EVIDENCE_MANIFEST.json'),manifest_created=m['created_utc'])
write('integrity.json',integrity)

ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/package/2006/relationships'}
def content(n): return ''.join(n.xpath('.//w:t/text()',namespaces=ns))
with ZipFile(m['final_output']['path']) as z,ZipFile(next(p for p in m['inputs'] if p.endswith('1_2026 - УЗК.docx'))) as t:
    doc=E.fromstring(z.read('word/document.xml')); template=E.fromstring(t.read('word/document.xml'))
    body=doc.find('w:body',ns); tb=template.find('w:body',ns)
    tables=body.findall('w:tbl',ns); tt=tb.findall('w:tbl',ns)
    rows=tables[0].findall('w:tr',ns); oldrows=tt[0].findall('w:tr',ns)
    oldmarks=[content(r.findall('w:tc',ns)[0]) for r in oldrows[5:]]
    markrows=[]
    for row in rows[5:]:
        cells=row.findall('w:tc',ns)
        markrows.append({'cells':[content(c) for c in cells],'text_lines':[c.xpath('.//w:t/text()',namespaces=ns) for c in cells],'break_counts':[len(c.findall('.//w:br',ns)) for c in cells]})
    parts_changed=[n for n in t.namelist() if t.read(n)!=z.read(n)]
    hyperlink_rels=[];external_pdf=[];fields=[]
    for n in z.namelist():
        if n.endswith('.rels'):
            rels=E.fromstring(z.read(n))
            hyperlink_rels.extend([dict(part=n,**dict(r.attrib)) for r in rels if r.get('Type','').endswith('/hyperlink')])
            external_pdf.extend([dict(part=n,**dict(r.attrib)) for r in rels if r.get('TargetMode')=='External' and '.pdf' in r.get('Target','').casefold()])
        if n.startswith('word/') and n.endswith('.xml'):
            root=E.fromstring(z.read(n))
            if root.xpath('.//w:instrText|.//w:fldSimple|.//w:fldChar|.//w:delText|.//w:vanish',namespaces=ns): fields.append(n)
    word=dict(tables=len(tables),sections=len(doc.findall('.//w:sectPr',ns)),first_table_rows=len(rows),mark_rows=markrows,old_marks=oldmarks,old_marks_absent=all(v not in content(doc) for v in oldmarks),hyperlinks=len(doc.findall('.//w:hyperlink',ns)),hyperlink_relationships=hyperlink_rels,external_pdf_targets=external_pdf,hidden_or_field_parts=fields,changed_parts=parts_changed,signature_table_byte_equal=E.tostring(tables[1])==E.tostring(tt[1]),section_byte_equal=E.tostring(body.find('w:sectPr',ns))==E.tostring(tb.find('w:sectPr',ns)),constant_rows_equal={str(i):E.tostring(rows[i])==E.tostring(oldrows[i]) for i in (0,2,3,4)},body_paragraphs=[content(p) for p in body.findall('w:p',ns)],template_body_paragraphs=[content(p) for p in tb.findall('w:p',ns)],test_label='ТЕСТ / ЧЕРНОВИК' in content(doc))
    write('docx_structure.json',word)

from openpyxl import load_workbook
dev=ROOT/'DEV/ВЗМК_DEV.xlsx'; uzk=ROOT/'DEV/УЗК_DEV.xlsm'
wb=load_workbook(dev,read_only=True,data_only=True,keep_links=False)
s=wb['Паспорт качества']; selected=[(i,r) for i,r in enumerate(s.iter_rows(values_only=True),1) if r[0]=='Гараж Казахстан' and r[1]=='9/2026'];wb.close()
assert len(selected)==1
rn,row=selected[0]
wb=load_workbook(dev,read_only=True,data_only=False,keep_links=False);formula=wb['Паспорт качества'].cell(rn,14).value;wb.close()
wb=load_workbook(uzk,read_only=True,data_only=False,keep_links=False);reg=[(i,list(r)) for i,r in enumerate(wb['УЗК'].iter_rows(min_col=1,max_col=3,values_only=True),1) if i>1 and any(v is not None for v in r)];wb.close()
with ZipFile(uzk) as z:
    filters=[]
    for n in z.namelist():
        if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'):
            filters.extend(E.fromstring(z.read(n)).xpath('//*[local-name()="autoFilter"]/@ref'))
catalog=next(p for p in m['inputs'] if p.endswith('Справочник марок.xlsx'))
wb=load_workbook(catalog,read_only=True,data_only=False,keep_links=False);s=wb.active
target={'Бп2-1','Пс1-9','Пс1-10','Пс2-16'}
cat=[dict(row=i,values=list(r)) for i,r in enumerate(s.iter_rows(values_only=True),1) if r[0] in target];wb.close()
write('independent_inputs.json',dict(passport=dict(row=rn,project=row[0],number=row[1],object=row[4],cipher=row[5],mass=row[6],date_N=row[13],formula_N=formula),registry_rows=reg,registry_count=len(reg),registry_autofilter=filters,catalog_rows=cat,input_fingerprints={p:dict(expected=h,current=sha(p)) for p,h in m['inputs'].items()}))
print(json.dumps({k:v for k,v in integrity.items() if k not in ('evidence_files','product_set','output')},ensure_ascii=False,indent=2))
print(json.dumps({'passport_row':rn,'object':row[4],'cipher':row[5],'date':row[13],'registry_count':len(reg),'registry_filter':filters,'catalog':cat,'markrows':markrows,'word_constants':word['signature_table_byte_equal']},ensure_ascii=False,default=str))
