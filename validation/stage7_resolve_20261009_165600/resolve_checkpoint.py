"""Bounded read-only vom reconciliation. Actual Stage1 algorithm on memory fixtures; native selected VOM read."""
import ast
from datetime import datetime, timezone
from decimal import Decimal
import difflib
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
from unittest.mock import patch

RUN=Path(__file__).resolve().parent
ROOT=RUN.parents[1]
OLD=ROOT/'validation/stage5_refresh_for_stage7_20261009_163408'
MIG=ROOT/'validation/stage3_catalog_schema_20261008_162454'
S6=ROOT/'validation/stage6_final_targeted_rereview_r2_20261009'
FIX=ROOT/'validation/targeted_M1_M2_20261005_231515'
CURRENT=ROOT/'passport_app/vom.py'
ACCEPTED=ROOT/'e2e_user_stage1/app/passport_app/vom.py'
VOM=ROOT.parent/'Гараж казахстан/Паспорт/9 от 06.10.2026/ВОМ отгрузочный счет 183  (5).xls'
sys.path.insert(0,str(ROOT))

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def write(name,value):
    with (RUN/name).open('x',encoding='utf8') as stream:json.dump(value,stream,ensure_ascii=False,indent=2)

paths=set(read(OLD/'SOURCE_BEFORE.json'))
paths.update(str(p) for p in OLD.rglob('*') if p.is_file())
paths.update(str(p) for p in (ROOT/'validation/stage7_quality_audit_20261009').iterdir() if p.is_file())
paths.update([str(CURRENT),str(ACCEPTED),str(FIX/'after.json'),str(FIX/'TARGETED_FIX_REPORT.txt'),str(S6/'REVIEW_REPORT.txt'),str(VOM)])
assert all(Path(p).name.casefold() not in ('взмк.xlsx','узк.xlsm') for p in paths)
before={p:sha(p) for p in sorted(paths)}
product_before={p:sha(p) for p in read(OLD/'PRODUCT_CURRENT.json')}
write('SOURCE_BEFORE.json',before)
write('PRODUCT_BEFORE.json',product_before)
access={'production_open_attempts':[],'network_attempts':[],'outside_write_attempts':[]}
def guard(event,args):
    if event.startswith('socket.') and event!='socket.__new__':
        access['network_attempts'].append(event)
        raise RuntimeError('NETWORK_FORBIDDEN')
    if event!='open' or not isinstance(args[0],(str,bytes)):return
    p=Path(args[0].decode() if isinstance(args[0],bytes) else args[0]).resolve()
    if p.name.casefold() in ('взмк.xlsx','узк.xlsm'):
        access['production_open_attempts'].append(str(p))
        raise RuntimeError('PRODUCTION_FORBIDDEN')
    mode=args[1]
    if isinstance(mode,str) and any(c in mode for c in 'wax+') and not p.is_relative_to(RUN):
        access['outside_write_attempts'].append(str(p))
        raise RuntimeError('SOURCE_WRITE_FORBIDDEN')
sys.addaudithook(guard)
expected='c63bbd362763ba99f85de1610c01622ac230f3303b7a1e297345d02259ea954f'
actual='236ec7287bc4a9c02cd4e8fff84b60334f5d14b74e00588f0cff90e731857a41'
assert sha(ACCEPTED)==expected and sha(CURRENT)==actual
assert read(MIG/'product_hash_comparison.json')[str(CURRENT)]['CURRENT']==expected
assert read(FIX/'after.json')['files_after']['passport_app/vom.py']['sha256']==expected
assert read(FIX/'before.json')['files_before']['passport_app\\vom.py']['sha256']==actual
diff=''.join(difflib.unified_diff(ACCEPTED.read_text(encoding='utf8').splitlines(True),CURRENT.read_text(encoding='utf8').splitlines(True),
    fromfile='accepted_M1_fixed_vom.py',tofile='current_passport_app_vom.py'))
(RUN/'VOM_EXPECTED_TO_CURRENT.diff').write_text(diff,encoding='utf8')
work01=read(RUN/'WORK01_TARGETED_ACCEPTANCE.json')
final=work01['turns'][0]['final_answers'][0]['text']
assert 'M1 RESOLVED' in final and 'strip()' in final and ' Балка , Балка' in final
work06=read(RUN/'WORK06_SCOPE_AND_PARALLEL_CHANGE.json')
assert any('Да, изменения другой работы ожидаемы' in json.dumps(t,ensure_ascii=False) for t in work06['turns'])
from openpyxl import Workbook
from passport_app import vom as current
spec=importlib.util.spec_from_file_location('passport_app.accepted_vom_checkpoint',ACCEPTED)
accepted=importlib.util.module_from_spec(spec)
sys.modules[spec.name]=accepted
spec.loader.exec_module(accepted)

def workbook(names):
    book=Workbook()
    sheet=book.active
    sheet.title='Memory fixture'
    for cell,value in {'A1':'Отправочная марка','B1':'Наименование марки','C1':'Количество','D1':'Масса','F1':'Площадь',
                       'D2':'1 шт','E2':'Всех','F2':'1 шт','G2':'Всех'}.items():sheet[cell]=value
    for region in ('A1:A2','B1:B2','C1:C2','D1:E1','F1:G1'):sheet.merge_cells(region)
    for r,name in enumerate(names,3):
        for c,value in enumerate([f'M{r-2}',name,1,10,10,1,1],1):sheet.cell(r,c,value)
    return book

def run_case(module,names):
    # No source workbook or saved fixture is opened. Only file I/O is replaced;
    # actual locate_layout, row validation, and read_vom string algorithm execute.
    with patch.object(module,'load_book',side_effect=[workbook(names),workbook(names)]),patch.object(module,'sha256',return_value='MEMORY_FIXTURE_NO_FILE_IO'):
        result=module.read_vom(RUN/'MEMORY_FIXTURE_NOT_SAVED.xlsx')
    return dict(connections=result.connections,rows=[r[1] for r in result.rows],quantity=str(result.quantity),mass=str(result.mass))

cases=[('normal',['Балка','Балка'],'Балка'),
       ('edge_spaces',[' Балка '],' Балка '),
       ('exact_unique_first_order',[' Балка ','Балка',' Балка '],' Балка , Балка'),
       ('empty_and_whitespace',[None,'','   ','\t ','Балка'],'Балка')]
results=[]
for name,names,wanted in cases:
    now=run_case(current,names)
    old=run_case(accepted,names)
    results.append(dict(case=name,input=names,expected=wanted,current=now,accepted_checkpoint=old,
        current_pass=now['connections']==wanted,accepted_pass=old['connections']==wanted))
assert all(v['accepted_pass'] for v in results)
assert [v['case'] for v in results if not v['current_pass']]==['edge_spaces','exact_unique_first_order']
write('CONNECTIONS_TARGETED_REPRODUCTION.json',dict(status='CURRENT_FAIL_2_OF_4; ACCEPTED_PASS_4_OF_4',
    fixture='Fresh in-memory openpyxl books; no XLSX serialization; load_book and file digest mocked only',
    actual_code_executed='Both exact SHA modules read_vom + unmodified actual locate_layout',cases=results))

# Actual current shipment verification is a separate accepted read-only COM path.
from uzk_stage5.inputs import visible_vom,verify_pins,PINS
from profiles_stage4.offline import no_network
verify_pins()
with no_network() as attempts:
    selected=visible_vom(VOM)
assert not attempts
assert selected['visible_count']==104 and len(selected['hidden_data_rows'])==397
assert sum(r['quantity'] for r in selected['rows'])==288
assert Decimal(selected['mass_kg'])==Decimal('15822.14')
assert sha(VOM)=='aae5e5b8c0246c8fe97b3dec9cbcc039bc3e3cedb9d245b474b2f871b90596ae'
write('CURRENT_VOM_NATIVE_READONLY.json',dict(status='PASS',source=selected,quantity=288,sha256=sha(VOM),
    scope='Selected real VOM only; native accepted Stage5 adapter; no generation or source save'))
imports={}
for rel in ('uzk_stage5/core.py','uzk_stage5/inputs.py','passport_app/catalog.py'):
    p=ROOT/rel
    tree=ast.parse(p.read_text(encoding='utf8'))
    imports[rel]=dict(sha256=sha(p),imports=[ast.unparse(n) for n in ast.walk(tree) if isinstance(n,(ast.Import,ast.ImportFrom))])
write('STAGE5_DEPENDENCY_ASSESSMENT.json',dict(status='NO_DIRECT_STAGE5_READ_VOM_DEPENDENCY',files=imports,
    fact='Stage5 core.preview calls uzk_stage5.inputs.visible_vom; inputs imports passport_app.catalog, not passport_app.vom',
    limitation='No full dependency/regression claim. No Stage5 preview or generator was run after checkpoint blocker.',
    effect='The M1 regression does not change the selected VOM read just verified; this does not prove acceptance of the reverted Stage1 code.'))
write('VOM_RECONCILIATION.json',dict(status='BLOCKED_ACCEPTANCE_AND_SEMANTIC_REGRESSION',file=str(CURRENT),expected=expected,actual=actual,
    fixed_checkpoint=str(ACCEPTED),diff=str(RUN/'VOM_EXPECTED_TO_CURRENT.diff'),
    provenance=dict(original_M1_fix_report=str(FIX/'TARGETED_FIX_REPORT.txt'),
        original_report_sha256=sha(FIX/'TARGETED_FIX_REPORT.txt'),fix_before=actual,fix_after=expected,
        accepted_work_thread=work01['thread_id'],accepted_work_title=work01['title'],
        accepted_work_turn=work01['turns'][0]['id'],accepted_work_completed_utc=datetime.fromtimestamp(work01['turns'][0]['completedAt'],timezone.utc).isoformat(),
        migration_manifest_created_utc=read(MIG/'REVIEW_MANIFEST.json')['created_utc'],
        stage6_review_created=read(S6/'REVIEW_MANIFEST.json')['created_at'],
        stage6_recorded_drift=True,stage6_authorization='Parallel changes expected; semantic reversal of M1 not explicitly accepted/tested',
        exact_reversion_time='NOT_ESTABLISHED',exact_reversion_author_and_reason='NOT_ESTABLISHED',
        conclusion='Current bytes match the PRE-FIX hash exactly; not a new schema migration change. No later independent acceptance of this reversal established.'),
    current_contract_failure='stripping changes original spelling and merges distinct exact names',
    current_stage5_direct_impact=False,selected_shipment_contract='PASS_104_397_288_15822.14',
    user_stop_rule='If acceptance or correctness unconfirmed: STOP. Stage1-6 code read-only.',
    action='No rollback, no Stage1 fix, no wrapper/product/output generation'))
write('ACCEPTED_INPUT_PINS.json',[dict(path=str(p),expected=h,actual=sha(p)) for p,h in PINS.items()])
after={p:sha(p) for p in before}
current_product={p:sha(p) for p in product_before}
changes={p:dict(before=h,after=after[p]) for p,h in before.items() if h!=after[p]}
assert not changes and product_before==current_product
write('SOURCE_AFTER.json',after)
write('PRODUCT_AFTER.json',current_product)
write('PRODUCT_CURRENT.json',current_product)
write('SOURCE_INTEGRITY.json',dict(status='PASS',protected_files=len(before),changes=changes,
    stage1_6_product_before_after_current_equal=True,production_access='NONE',registry_decrement=False,
    old_receipts_and_stage7_stop_preserved=True,code_fix='NOT_RUN'))
write('ACCESS_GUARD.json',dict(**access,native_access='Only explicit selected VOM via accepted visible_vom, ReadOnly=True, UpdateLinks=0, no save',
    generated_documents=0,full_regression=False))
text=f'''# Stage7 checkpoint resolution — STOP

1. ROOT CAUSE AND FIX

Предыдущий receipt blocker снимается новым контрактом: Stage7 wrapper/provenance recorder разрешён, Stage5 API менять не требуется. Но первый gate не пройден: фактический vom.py вернулся к поведению до принятого M1 fix. Новый wrapper не реализован после этого STOP.

2. vom.py reconciliation

Файл: `{CURRENT}`.
Expected принятый M1 fix/migration checkpoint: `{expected}`.
Actual: `{actual}`.
Exact diff — VOM_EXPECTED_TO_CURRENT.diff. Единственное изменение: original_name/проверка пустоты заменены trimmed/strip-before-dedup. Actual SHA совпадает с PRE-FIX before.json; expected SHA совпадает с after.json и тремя сохранёнными E2E runtime copies.

Независимый «{work01['title']}» принял сохранение исходного написания (M1 RESOLVED,5/5 targeted PASS). Его actual final answer повторно прочитан через read_thread и сохранён в WORK01_TARGETED_ACCEPTANCE.json. Stage6 Work зафиксировал drift и ответ пользователя «Да, изменения другой работы ожидаемы», но его review был S6-P1-R2/minimal related regression; семантический возврат M1 им не принят и не проверен. Прямое подтверждение обоих Work scope сохранено. Точное время/автор/причина возврата неизвестны; они не выводятся из mtime. Временные checkpoint evidence приведены в VOM_RECONCILIATION.json.

Новый targeted proof на actual read_vom:
- `" Балка "`: expected `" Балка "`, actual `"Балка"`.
- `" Балка ", "Балка", " Балка "`: expected `" Балка , Балка"`, actual `"Балка"`.
Текущая версия FAIL2/4; принятая точная версия PASS4/4 на тех же in-memory fixtures. Это регрессия ранее принятого M1, а не только SHA mismatch. Сам дефект был MINOR Stage1; в этом задании он блокирует обязательное принятие checkpoint по §2. Stage1–6 здесь read-only, поэтому исправление/откат не выполнялись.

3. accepted current checkpoints

DEV ВЗМК/DEV УЗК/Профили — pins PASS. Текущий project catalog SHA `{sha(ROOT.parent/'Гараж казахстан/Справочник марок.xlsx')}` совпадает с принятой миграцией. Текущая ВОМ SHA `{sha(VOM)}`: фактический native read-only PASS104 видимых/397 скрытых/288шт./15822.14кг. Производственные реестры не открывались. Проверка именно отгрузки пройдена; она не доказывает принятие возврата кода M1.

4. Stage 7 wrapper design

Контракт исправлен пользователем: новый Stage7 assembly_run_id и START/RESULT/VALIDATION, SHA inputs/code, membership, actual generator call, output validation, repeated fingerprints и самостоятельный SAMPLE_ONLY operation receipt. Принятый Stage5 generator менять нельзя. Реализация NOT_RUN после gate STOP; никакой выдуманной успешной операции нет.

5. actual UZK TEST run and SHA

NOT_RUN; new DOCX path/SHA NOT_APPLICABLE. Stage5 использует собственный visible_vom; прямой зависимости от passport_app.vom.read_vom нет (AST/import proof). Native selected VOM прошёл, но это не отменяет условие принятия checkpoint. Fresh preview/генерация не выполнялись после доказанного semantic blocker.

6. new Stage 7 operation receipt

NOT_CREATED; path/SHA/assembly_run_id NOT_APPLICABLE. Evidence folder не является generation run_id. Старые receipt не переписаны.

7. new VIK

NOT_RUN. Ни старый ВИК с другой registry binding, ни новый непроверенный документ не включены.

8. selected source manifest

NOT_ASSEMBLED. Нет статуса успешного coherent TEST-комплекта. АКЗ1/2/3 не выбирались; новая проверка АКЗ0/Приложения№1 после STOP не выполнялась.

9. quality DOCX and SHA

NOT_CREATED; path/SHA NOT_APPLICABLE. Первый шаблон и пункты12–14 не менялись. Официальная выдача не подразумевается.

10. native Word checks

NOT_RUN: новых документов нет. Native Excel selected VOM read PASS, source SHA unchanged; Word layout этим не подтверждается.

11. targeted test results

Новые4 focused in-memory cases ×2 exact code versions. Actual module read_vom и actual locate_layout выполнялись; заменены только file-loading и file hash для in-memory input. Текущий2/4 PASS и2/4 FAIL; accepted4/4 PASS. Файлы fixtures не сохранялись. Это не full file I/O E2E. Реальная выбранная ВОМ отдельно проверена native accepted adapter. Old Work acceptance и Stage6 drift/scope сверены; expected/current source hashes проверены. Full Stage1–6 regression NOT_RUN.

12. source integrity

PASS: {len(before)} protected sources/evidence, changed0. PRODUCT BEFORE=AFTER=CURRENT. Старые receipts, Stage7 STOP/audit/refresh, DEV, real catalog/profiles/VOM/templates и Stage1–6 product не изменены. Python production/socket/outside-write attempts0; native Excel использовал только explicit selected VOM, macros disabled, UpdateLinks0, close without save. Общесистемный trace не заявляется.

13. changed files

Только новая папка `{RUN}`: read-only resolution helper, Work excerpts, exact diff, targeted reproduction, native VOM evidence, fingerprints, отчёт и manifest. Stage7 product/launcher не созданы. Список CHANGED_FILES.json. Manifest записан последним и проверен read-only.

14. remaining blockers

Один реальный unresolved checkpoint blocker: текущий vom.py `{actual}` возвращает регрессию M1; независимое принятие именно этого возврата не установлено. Требуется отдельное разрешение/решение по Stage1 M1 и его targeted acceptance; текущее разрешение меняет только Stage7. Нельзя объявить regression принятым только потому, что parallel drift ожидался. Receipt architecture больше не самостоятельный blocker: её разрешённая Stage7 реализация ожидает прохождения первого gate.

15. READY FOR WORK 07 REVIEW

NO — Stage7 implementation/output не созданы. Конкретный semantic blocker доказан и доступен для review. После отчёта STOP. Stage8/PDF/EXE/production/decrement не запускались.
'''
(RUN/'RESOLVE_CHECKPOINT_REPORT.md').write_text(text,encoding='utf8')
write('CHANGED_FILES.json',dict(product_changes=[],old_evidence_changes=[],
    added=[str(p) for p in sorted(RUN.iterdir()) if p.is_file()]+[str(RUN/'CHANGED_FILES.json'),str(RUN/'EVIDENCE_MANIFEST.json')]))
manifest=dict(status='STOP_VOM_M1_REGRESSION_NOT_ACCEPTED',created_utc=datetime.now(timezone.utc).isoformat(),
    ready_for_work07_review='NO',receipt_architecture_blocker='RESOLVED_BY_USER_CONTRACT_NOT_IMPLEMENTED',
    product_mutations=0,outputs_created=0,
    evidence=[dict(path=p.name,sha256=sha(p),size=p.stat().st_size) for p in sorted(RUN.iterdir()) if p.is_file()],manifest_written_last=True)
write('EVIDENCE_MANIFEST.json',manifest)
assert all(sha(RUN/v['path'])==v['sha256'] for v in manifest['evidence'])
print(json.dumps(dict(status=manifest['status'],expected=expected,actual=actual,current_cases_passed=2,current_cases_failed=2,
    accepted_cases_passed=4,selected_vom='PASS 104/397/288/15822.14',protected=len(before),changes=0,
    manifest=str(RUN/'EVIDENCE_MANIFEST.json'),manifest_sha256=sha(RUN/'EVIDENCE_MANIFEST.json')),ensure_ascii=False,indent=2))
