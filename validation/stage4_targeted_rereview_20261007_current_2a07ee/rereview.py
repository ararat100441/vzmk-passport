"""Independent targeted M1/M2 review. All mutations are new isolated fixtures."""
from pathlib import Path
from collections import Counter, defaultdict
from copy import deepcopy
from decimal import Decimal
from unittest.mock import patch
import hashlib
import json
import os
import shutil
import subprocess
import sys
import traceback
import zipfile

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
FIX = ROOT/'validation/stage4_fix_m1_m2_20261007_090000'
sys.path.insert(0,str(ROOT))
sys.stdout.reconfigure(encoding='utf-8')
sys.dont_write_bytecode=True
def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(n,x): (OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def require(x,m='assertion failed'):
    if not x: raise AssertionError(m)

manifest=read(FIX/'REVIEW_MANIFEST.json')
tests=[]
network=[]
def audit(event,args):
    if event in ('socket.connect','socket.getaddrinfo','socket.sendto'):
        network.append(event)
        raise RuntimeError('Network forbidden in targeted rereview')
    if event=='open' and len(args)>2:
        p=args[0]
        if isinstance(p,(str,bytes,os.PathLike)):
            try: path=Path(os.fsdecode(p)).resolve()
            except (ValueError,OSError): return
            if path.is_relative_to(ROOT) and not path.is_relative_to(OUT):
                mode,flags=args[1:3]
                writing=(isinstance(mode,str) and any(c in mode for c in 'wax+')) or flags & (os.O_WRONLY|os.O_RDWR|os.O_CREAT|os.O_TRUNC)
                require(not writing,'Write outside independent review: '+str(path))
    if event in ('os.remove','os.rename','os.mkdir','os.rmdir'):
        for p in args[:2] if event=='os.rename' else args[:1]:
            if isinstance(p,(str,bytes,os.PathLike)):
                path=Path(os.fsdecode(p)).resolve()
                require(not path.is_relative_to(ROOT) or path.is_relative_to(OUT),'Mutation outside review: '+str(path))
sys.addaudithook(audit)

expected_paths={ROOT/p:h for p,h in manifest['product_hashes'].items()}
expected_paths.update({Path(p):h for p,h in manifest['protected_hashes'].items()})
expected_paths.update({Path(p):h for p,h in manifest['DEV_artifacts'].items()})
expected_paths.update({FIX/p:h for p,h in manifest['evidence_files'].items()})
expected_paths[FIX/'REVIEW_MANIFEST.json']=sha(FIX/'REVIEW_MANIFEST.json')
before={str(p):sha(p) for p in expected_paths}
save('preservation_before.json',before)
def case(name,func):
    try:
        detail=func()
        tests.append({'name':name,'status':'PASS','detail':detail})
        print('PASS '+name,flush=True)
    except BaseException:
        tests.append({'name':name,'status':'FAIL','traceback':traceback.format_exc()})
        print('FAIL '+name+'\n'+traceback.format_exc(),flush=True)
    save('results.json',{'tests':tests,'passed':sum(t['status']=='PASS' for t in tests),'failed':sum(t['status']=='FAIL' for t in tests)})

def binding():
    mismatches=[str(p) for p,h in expected_paths.items() if sha(p)!=h]
    require(not mismatches,str(mismatches))
    a,b,c=(read(FIX/f'product_hashes_{x}.json') for x in ['before','after','current'])
    require(a==b==c==manifest['product_hashes'])
    require(read(FIX/'protected_before.json')==read(FIX/'protected_after.json')==manifest['protected_hashes'])
    t=read(FIX/'targeted_tests.json')
    require(t['status']=='PASS' and t['passed']==34 and t['failed']==0)
    old=read(ROOT/'validation/stage4_independent_review_20261007_0119/reviewed_hashes_final.json')
    changed={str(Path(p).relative_to(ROOT)) for p,h in old.items() if sha(p)!=h}
    expected_changed={'profiles_stage4/model.py','profiles_stage4/storage.py','profiles_stage4/workbook.mjs','STAGE4_README.md',
                      'DEV/Справочники/Профили.xlsx','DEV/Справочники/profiles.user.json','DEV/Справочники/profiles.catalog.json','DEV/Справочники/profiles.build.json'}
    require({p.replace('\\','/') for p in changed}==expected_changed,str(changed))
    actual_product={str(p.relative_to(ROOT)):sha(p) for p in (ROOT/'profiles_stage4').iterdir() if p.is_file()}
    require('profiles_stage4\\windows_generation.py' in a)
    return {'files_checked':len(expected_paths),'fix_evidence_files':len(manifest['evidence_files']),
            'protected_files':len(manifest['protected_hashes']),'before_after_current':True,
            'changed_from_original_review':sorted(changed),'current_stage4_module_files':actual_product}
case('current artifact and exact fix evidence binding',binding)
require(tests[-1]['status']=='PASS','STOP: artifact/evidence binding failed')

import profiles_stage4.storage as s
from profiles_stage4.model import Catalog,ProfileError,calculate_uzk,frequency_angle,RULE
from profiles_stage4.windows_generation import ProtectedFile,GenerationGuard
from openpyxl import load_workbook
NAMES=[s.DEFAULT_WORKBOOK.name,s.STATE,s.CATALOG,s.COMMIT]
DATA=read(s.DEFAULT_WORKBOOK.parent/s.CATALOG)
STATE=read(s.DEFAULT_WORKBOOK.parent/s.STATE)
CAT=Catalog(DATA['records'],DATA['sources'],STATE)
PRIMARY={'Двутавр':('h','b','s','t'),'Швеллер':('h','b','s','t'),
         'Уголок равнополочный':('b','t'),'Уголок неравнополочный':('b','b2','t'),
         'Квадратный профиль':('h','b','t'),'Прямоугольный профиль':('h','b','t'),'Круглая труба':('D','s')}
BLOCK={'Тип','h','b','b2','s','t','D'}
def differences(members):
    vals=defaultdict(list)
    for m in members:
        for f,v in {**m['geometry'],**m.get('quality',{})}.items():
            if v is not None: vals[f].append(v)
    fields={f for f,v in vals.items() if any(abs(Decimal(str(v[0]))-Decimal(str(x)))>Decimal('0.000000001') for x in v[1:])}
    if len({m['family'] for m in members})>1: fields.add('Тип')
    return fields
def m1_all():
    groups=defaultdict(list)
    for m in DATA['records']: groups[m['key'],m['standard']].append(m)
    disagreed=[]
    for (key,nd),members in groups.items():
        fields=differences(members)
        if not fields: continue
        block=bool(fields & BLOCK)
        r=CAT.lookup(key,nd)
        row=r['profile']
        require(r['status']==('CONFLICT' if block else 'FOUND'))
        require(set(row['unresolved_fields'])==fields)
        require(row['provenance']==members,'Lost/reordered raw provenance')
        require(all(f not in row['geometry'] and f not in row['quality'] for f in fields))
        require(row['disagreement_class']==('BLOCKING' if block else 'NON_BLOCKING'))
        if block: require(not row['geometry'])
        else: require(all(row['geometry'].get(f) is not None for f in PRIMARY[row['family']]))
        disagreed.append({'key':key,'standard':nd,'fields':sorted(fields),'blocking':block,'lookup':r['status'],
                         'geometry':row['geometry'],'raw_members':members})
    require(len(disagreed)==15 and sum(x['blocking'] for x in disagreed)==8)
    require(len(CAT.conflicts)==15)
    save('all_15_disagreements.json',disagreed)
    return {'total':15,'blocking':8,'non_blocking':7,'all_raw_values_and_provenance_preserved':True}
case('M1 independent raw grouping and all 15 lookups',m1_all)

def probes():
    beam=next(x for x in DATA['records'] if x['key']=='35Б1' and x['standard']=='ГОСТ Р 57837-2017')
    answers=[]
    for field in ['R','r','area_cm2','mass_kg_m','reference_probe','h','b','b2','s','t','D','Тип','core_plus_radii']:
        a,b=deepcopy(beam),deepcopy(beam)
        a['record_id'],b['record_id']='INDEP-A','INDEP-B'
        b['source']='INDEP-B'
        b['source_id']='INDEP-B'
        for m in [a,b]:
            for f in ['R','r']: m['geometry'].pop(f,None)
        if field=='Тип': b['family']='Швеллер'
        elif field=='core_plus_radii':
            a['geometry'].update(h=346,R=10,r=2);b['geometry'].update(h=347,R=11,r=3)
        else:
            container='geometry' if field in ['R','r','h','b','b2','s','t','D'] else 'quality'
            start=a[container].get(field) or 10
            a[container][field]=start;b[container][field]=start+1
        c=Catalog([a,b],user_state=deepcopy(STATE))
        block=field in BLOCK or field=='core_plus_radii'
        explicit,preferred=c.lookup('35Б1',beam['standard']),c.lookup('35Б1')
        require(explicit['status']==preferred['status']==('CONFLICT' if block else 'FOUND'))
        if not block:
            require(field not in explicit['profile']['geometry'] and field not in explicit['profile']['quality'])
            require(calculate_uzk(explicit)['status']=='CALCULATED')
        else:
            try: calculate_uzk(explicit)
            except ProfileError: pass
            else: raise AssertionError('Blocking UZK accepted')
        # Include another ND so preference selection really traverses preference branch.
        third=deepcopy(a);third['standard']='INDEPENDENT OTHER ND';third['record_id']='INDEP-C'
        multi=Catalog([a,b,third],user_state=deepcopy(STATE))
        if block: require(multi.lookup('35Б1')['status']=='CONFLICT')
        answers.append({'field':field,'explicit':explicit['status'],'without_ND':preferred['status'],'preferred_same_ND_safety':True})
    return answers
case('M1 independent R/r auxiliary and every blocking field probes',probes)

def counts():
    groups=defaultdict(list)
    for r in DATA['records']:groups[r['key'],r['standard']].append(r)
    status=Counter()
    for members in groups.values():
        fields=differences(members)
        status['КОНФЛИКТ' if fields & BLOCK else 'СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ' if fields else
               'ПРОВЕРЕНО_ДВУМЯ_ИСТОЧНИКАМИ' if len({m['source'] for m in members})>1 else 'ИЗ ИСТОЧНИКА']+=1
    require(len(DATA['records'])==3037 and len(groups)==2849)
    require(dict(status)=={'ИЗ ИСТОЧНИКА':2662,'ПРОВЕРЕНО_ДВУМЯ_ИСТОЧНИКАМИ':172,'СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ':7,'КОНФЛИКТ':8})
    require(Counter(r['status'] for r in CAT.rows)==status)
    ambiguous=sum(CAT.lookup(k)['status']=='AMBIGUOUS' for k in {k for k,n in groups})
    uzk=sum(r['uzk_rule']==RULE for r in CAT.rows)
    require(ambiguous==51 and uzk==480 and len(STATE['manual_rows'])==0)
    save('counts.json',{'raw':3037,'canonical':2849,'statuses':dict(status),'ambiguous_keys':ambiguous,'uzk_enabled':uzk,'manual':0})
    return read(OUT/'counts.json')
case('independent post-fix counts',counts)

def golden():
    result={}
    for nd,geom,uzk in [('ГОСТ Р 57837-2017',[346,174,6,9],[[328,6],[348,9]]),('ГОСТ 26020-83',[346,155,6.2,8.5],[[329,6.2],[310,8.5]])]:
        r=CAT.lookup('35Б1',nd)
        require(r['status']=='FOUND')
        actual=[r['profile']['geometry'][f] for f in ['h','b','s','t']]
        calc=[[r['L'],r['S']] for r in calculate_uzk(r)['rows']]
        require(actual==geom and calc==uzk)
        result[nd]={'lookup':r['status'],'geometry':actual,'UZK':calc}
    require(CAT.lookup('35Б1')['status']=='FOUND_PREFERRED' and CAT.lookup('20Ш1')['status']=='AMBIGUOUS')
    require(frequency_angle(11.5)=={'frequency_mhz':5,'angle_degrees':70})
    require(frequency_angle(12)=={'frequency_mhz':2.5,'angle_degrees':65})
    result.update(plain_35Б1='FOUND_PREFERRED',plain_20Ш1='AMBIGUOUS',thickness_11_5=[5,70],thickness_12=[2.5,65])
    return result
case('golden lookup and UZK regression',golden)

def workbook():
    w=load_workbook(s.DEFAULT_WORKBOOK,read_only=True,data_only=False)
    try:
        require(w.sheetnames==['Профили','Источники','Правила УЗК','Предпочтения','Конфликты'])
        matrix={n:list(w[n].iter_rows()) for n in w.sheetnames}
        require(not any(c.data_type=='f' for rows in matrix.values() for row in rows for c in row))
        profiles=[dict(zip([c.value for c in matrix['Профили'][0]],[c.value for c in row])) for row in matrix['Профили'][1:]]
        require(len(profiles)==2849)
        lookup={(r['Ключ'],r['НД']):r for r in profiles}
        for row in profiles:
            require(all(row[f] is None or type(row[f]) in [int,float] for f in ['h','b','b2','s','t','D','r','R']))
        require(Counter(r['Статус'] for r in profiles)==Counter(r['status'] for r in CAT.rows))
        conflicts=[dict(zip([c.value for c in matrix['Конфликты'][0]],[c.value for c in row])) for row in matrix['Конфликты'][1:]]
        require(Counter(r['Классификация'] for r in conflicts)=={'BLOCKING':8,'NON_BLOCKING':7})
        for c in conflicts:
            row=lookup[c['Ключ'],c['НД']]
            actual=next(x for x in CAT.conflicts if x['id']==c['Conflict ID'])
            raw=json.loads(c['Значения A'])+json.loads(c['Значения B'])
            require({x['record_id'] for x in raw}=={x['record_id'] for x in actual['members']})
            for x in raw:
                original=next(y for y in actual['members'] if y['record_id']==x['record_id'])
                require(x['geometry']==original['geometry'] and x['quality']==original['quality'])
            if c['Классификация']=='NON_BLOCKING':
                require(all(row[f] is not None for f in PRIMARY[row['Тип']]))
                require('unresolved:' in row['Примечание'] and 'NON_BLOCKING' in row['Примечание'])
                require(c['Статус']=='СПРАВОЧНЫЕ ДАННЫЕ НЕ РАЗРЕШЕНЫ')
            else: require(all(row[f] is None for f in ['h','b','b2','s','t','D','r','R']))
        return {'sha256':sha(s.DEFAULT_WORKBOOK),'sheets':w.sheetnames,'profiles':len(profiles),'formulas':0,'numeric_geometry':True,'BLOCKING':8,'NON_BLOCKING':7,'raw_conflict_values_verified':15}
    finally:w.close()
case('current 2a07ee workbook targeted assessment',workbook)

def fixture(name):
    folder=OUT/'fixtures'/name/'DEV/Справочники'
    folder.mkdir(parents=True,exist_ok=False)
    for n in NAMES:shutil.copyfile(s.DEFAULT_WORKBOOK.parent/n,folder/n)
    return folder/NAMES[0]
def generation(t):return {n:sha(t.parent/n) for n in NAMES if (t.parent/n).exists()}
def clean(t):require(not list(t.parent.glob('.profiles-stage-*')) and not (t.parent/'.profiles.lock').exists(),'Partial stage or stale writer lock')
def backup(t,old):
    dirs=list((t.parent/'backups').iterdir()) if (t.parent/'backups').exists() else []
    require(bool(dirs))
    for p in dirs:
        require({n:sha(p/n) for n in NAMES}==old,'Backup bytes not verified previous generation')
        commit=read(p/s.COMMIT)
        require(all(sha(p/n)==h for n,h in commit['files'].items()))
        require(read(p/s.STATE)==read(p/s.CATALOG)['user_state'])
    return str(dirs[-1])
def stopped(func):
    try:func()
    except (OSError,ValueError) as e:return str(e)
    raise AssertionError('Build did not STOP')
def reuse(payload,destination,previews=None):
    require(read(payload)==s.workbook_payload(CAT),'Unexpected fixture payload')
    shutil.copyfile(s.DEFAULT_WORKBOOK,destination)
    return {'network_attempts':0}
def external(mode,t,name=None):
    args=[sys.executable,'-I','-B',str(OUT/'external_probe.py'),mode,str(t.parent)]
    if name:args.append(name)
    r=subprocess.run(args,capture_output=True,text=True,encoding='utf-8',timeout=20)
    require(r.returncode==0,r.stderr)
    return json.loads(r.stdout)

def race_during_render(name,op='write'):
    t=fixture('render_'+str(NAMES.index(name))+'_'+op);old=generation(t);expected={}
    def render(payload,dest,previews=None):
        reuse(payload,dest)
        p=t.parent/name
        if op=='delete':p.unlink();expected['sha']=None
        elif op=='replace':
            temp=t.parent/(name+'.replacement');temp.write_bytes(p.read_bytes()+b' ');os.replace(temp,p);expected['sha']=sha(p)
        elif op=='rename':p.rename(t.parent/(name+'.external-renamed'));expected['sha']=None
        elif name.endswith('.xlsx'):
            with zipfile.ZipFile(p,'a') as z:z.comment=b'INDEPENDENT_VALID_XLSX_MUTATION'
            w=load_workbook(p,read_only=True);w.close();expected['sha']=sha(p)
        else:
            x=read(p);x['independent_external_marker']=name;s.dump(p,x);expected['sha']=sha(p)
        return {'network_attempts':0}
    with patch.object(s,'render_workbook',render): msg=stopped(lambda:s.build(t))
    require(expected,'Mutation not reached')
    now=generation(t)
    require(now.get(name)==expected['sha'])
    require(all(now[n]==old[n] for n in NAMES if n!=name))
    require(not (t.parent/'backups').exists())
    if op=='rename':require(sha(t.parent/(name+'.external-renamed'))==old[name])
    clean(t)
    return {'file':name,'operation':op,'STOP':msg,'external_state_preserved':True,'new_generation_absent':True,'partial_generation_absent':True}
for name in NAMES:case('M2 external mutation during build '+name,lambda n=name:race_during_render(n))
for op in ['delete','rename','replace']:case('M2 render external '+op,lambda o=op:race_during_render(NAMES[0],o))

def boundary():
    t=fixture('boundary');old=generation(t);attempts=[]
    original=ProtectedFile.rename
    def attack(phase):
        a=external('attack',t)
        require(len(a)==16 and all(x['blocked'] and (x['winerror'] in [5,32,33] or x['errno']==13) for x in a),str(a))
        attempts.append({'phase':phase,'attempts':a})
    def rename(handle,destination):
        if not attempts:attack('immediately before first commit rename')
        original(handle,destination)
        if Path(destination)==t and len(attempts)==1:attack('after first new file installed, handles still held')
    with patch.object(s,'render_workbook',reuse),patch.object(ProtectedFile,'rename',rename): r=s.build(t)
    require(len(attempts)==2)
    backup(t,old);s.load_catalog(t);clean(t)
    return {'attempts':attempts,'coherent_backup':True,'normal_commit':True}
case('M2 separate-process commit-boundary write delete rename replace',boundary)

def vacancy(name):
    t=fixture('vacancy_'+str(NAMES.index(name)));old=generation(t);done=[];original=ProtectedFile.rename
    def rename(handle,destination):
        origin=handle.path
        original(handle,destination)
        if not done and origin==t.parent/name and Path(destination).parent.name=='previous':done.extend(external('vacancy',t,name))
    with patch.object(s,'render_workbook',reuse),patch.object(ProtectedFile,'rename',rename):msg=stopped(lambda:s.build(t))
    require(done and (t.parent/name).read_bytes()==b'INDEPENDENT_EXTERNAL_VACANT_NAME')
    now=generation(t)
    require(all(now[n]==old[n] for n in NAMES if n!=name),'Partial new generation survived rollback')
    backup(t,old);clean(t)
    return {'file':name,'STOP':msg,'external_bytes_preserved':True,'other_three_files_rolled_back':True,'verified_complete_previous_backup':True,'new_generation_absent':True}
for name in NAMES:case('M2 free-name race at replacement '+name,lambda n=name:vacancy(n))

def failure(mode):
    t=fixture('failure_'+mode);old=generation(t)
    with patch.object(s,'render_workbook',reuse):
        if mode=='render':
            with patch.object(s,'render_workbook',side_effect=ProfileError('INDEPENDENT_RENDER_FAILURE')):msg=stopped(lambda:s.build(t))
        elif mode=='backup':
            original=Path.write_bytes
            def write(p,data):
                if p.parent.name=='backups':raise OSError('INDEPENDENT_BACKUP_FAILURE')
                return original(p,data)
            with patch.object(Path,'write_bytes',write):msg=stopped(lambda:s.build(t))
            require(not (t.parent/'backups').exists())
        elif mode=='locked':
            with t.open('rb'):msg=stopped(lambda:s.build(t))
        elif mode.startswith('commit_'):
            index=int(mode.split('_')[1]);original=ProtectedFile.rename;done=[]
            def rename(handle,destination):
                if Path(destination)==t.parent/NAMES[index] and not done:
                    done.append(True);raise ProfileError('INDEPENDENT_COMMIT_FAILURE_'+str(index))
                original(handle,destination)
            with patch.object(ProtectedFile,'rename',rename):msg=stopped(lambda:s.build(t))
            require(done);backup(t,old)
        else:msg=stopped(lambda:s.build(t,fail_at='before_commit'))
    require(generation(t)==old,'Previous generation changed after failure')
    s.load_catalog(t);clean(t)
    return {'STOP':msg,'previous_generation_preserved':True,'partial_generation_absent':True}
for mode in ['render','backup','locked','before_commit','commit_0','commit_1','commit_2','commit_3']:
    case('M2 normal failure and rollback '+mode,lambda m=mode:failure(m))

def normal():
    t=fixture('normal_real_renderer');old=generation(t)
    r=s.build(t);backup(t,old);clean(t)
    require(s.load_catalog(t).counts()==CAT.counts())
    save('normal_build.json',r)
    return {'real_renderer':True,'backup':r['backup'],'counts':r['counts'],'reopen':r['reopen'],'network_attempts':r['network_attempts']}
case('normal rebuild with actual offline renderer',normal)

def app_writes():
    t=OUT/'fixtures/normal_real_renderer/DEV/Справочники'/NAMES[0]
    operations=[]
    for label,func in [
        ('preference',lambda:s.set_preference('20Ш1','ГОСТ 26020-83','INDEPENDENT_TARGETED_REVIEW',workbook=t)),
        ('manual',lambda:s.add_manual('INDEPENDENT MANUAL ROW','Двутавр','INDEPENDENT НД',{'h':200,'b':100,'s':5,'t':8},workbook=t))]:
        prev=generation(t);func();b=sorted((t.parent/'backups').iterdir())[-1]
        require({n:sha(b/n) for n in NAMES}==prev);clean(t);operations.append(label)
    current=s.load_catalog(t)
    c=next(c for c in current.conflicts if c['classification']=='BLOCKING')
    s.resolve_conflict(c['id'],c['members'][0]['record_id'],'INDEPENDENT EXPLICIT RESOLUTION',workbook=t)
    require(s.lookup(c['key'],c['standard'],workbook=t)['status']=='FOUND')
    require(s.lookup('20Ш1',workbook=t)['status']=='FOUND_PREFERRED')
    require(s.lookup('INDEPENDENT MANUAL ROW','INDEPENDENT НД',workbook=t)['status']=='FOUND')
    old_state=s.load_catalog(t).user_state
    s.build(t);final=s.load_catalog(t).user_state
    require(final['revision']==old_state['revision']+1)
    require(all(final[f]==old_state[f] for f in ['preferences','manual_rows','resolutions']))
    require(next(p for p in final['preferences'] if p['key']=='35Б1')==next(p for p in STATE['preferences'] if p['key']=='35Б1'))
    clean(t)
    return {'preference_write_rebuild':True,'manual_row_rebuild':True,'conflict_resolution_rebuild':True,'subsequent_rebuild_preserves_user_state':True,'actual_renderer_all_operations':True}
case('normal preference manual resolution and rebuild actual renderer',app_writes)

def preservation():
    after={str(p):sha(p) for p in expected_paths}
    save('preservation_after.json',after)
    require(before==after)
    require(not network)
    return {'BEFORE_EQUALS_AFTER_EQUALS_CURRENT':True,'files':len(after),'protected':len(manifest['protected_hashes']),
            'fix_evidence_unchanged':True,'previous_review_unchanged':True,'source_audits_unchanged':True,'DEV_checkpoint_unchanged':True,
            'real_DEV_writes':0,'network_attempts':len(network),'production_access':0}
case('final protected state and evidence preservation',preservation)
result=read(OUT/'results.json')
result.update(verdict='FAIL' if result['failed'] else 'PASS',current_artifacts=manifest['DEV_artifacts'],
              product_hashes=manifest['product_hashes'],not_run=['User E2E','Stage 5','production access','real DEV rebuild','new source audits','full Stage 4 review','native Excel visual UAT'])
save('TARGETED_REREVIEW.json',result)
print(json.dumps({'verdict':result['verdict'],'passed':result['passed'],'failed':result['failed']},ensure_ascii=False),flush=True)
