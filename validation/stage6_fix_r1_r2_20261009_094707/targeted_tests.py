"""Narrow R1/R2 tests; real Windows locks/native Excel; isolated mutations only."""
from pathlib import Path
import sys,json,shutil,copy,traceback,hashlib
from decimal import Decimal
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
import win32file,win32con,pywintypes
E=Path(__file__).resolve().parent;ROOT=E.parents[1];sys.path.insert(0,str(ROOT))
from akz_stage6 import core as c
from akz_stage6.gui import Launcher
import tkinter as tk

def save(n,v):(E/n).write_text(json.dumps(v,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
results=[]
def check(n,fn):
 try:
  v=fn();results.append({'name':n,'status':'PASS','detail':v});print('PASS',n,flush=True)
 except Exception as exc:
  results.append({'name':n,'status':'FAIL','error':str(exc),'traceback':traceback.format_exc()});raise
 finally:save('targeted_tests.json',results)
def stop(fn,code):
 try:fn()
 except Exception as exc:
  assert code in str(exc),(code,type(exc).__name__,str(exc))
  return {'type':type(exc).__name__,'code':code,'message':str(exc),'cause':repr(exc.__cause__)}
 raise AssertionError('EXPECTED STOP '+code)

audit={'production':[],'network':[]}
def auditor(event,args):
 if event=='open' and isinstance(args[0],(str,bytes)):
  p=Path(args[0].decode() if isinstance(args[0],bytes) else args[0])
  if p.name.casefold() in ('взмк.xlsx','узк.xlsm'):
   audit['production'].append(str(p));raise AssertionError('PRODUCTION DENIED')
 if event in ('socket.connect','socket.getaddrinfo'):
  audit['network'].append(str(args));raise AssertionError('NETWORK DENIED')
sys.addaudithook(auditor)
product=json.loads((E/'baseline.json').read_text('utf-8'))['product']
after={p:sha(p) for p in product};save('product_after.json',after)
F=E/'sandbox_run2/fixtures';O=E/'sandbox_run2/output';F.mkdir(parents=True);O.mkdir()
old=json.loads((ROOT/'validation/stage6_akz_part1_20261009_085822/preview.json').read_text('utf-8'))['payload']
cat=ROOT/'DEV/Stage6_Fixtures/stage6_akz_part1_20261009_085822_release3/TEST_Справочник марок.xlsx'
def cp(p,n):q=F/n;shutil.copyfile(p,q);return q
reg=cp(c.REGISTRY,'DEV_registry.xlsx');act=cp(c.ACT,'ACT_template.xlsx');application=cp(c.APPLICATION,'APPLICATION_template.xlsx')
catalog=cp(cat,'TEST_catalog.xlsx');vom=cp(old['vom']['path'],'selected_VOM.xls')
initial={str(p):sha(p) for p in (reg,act,application,catalog,vom)}
for count in (1,17,18):
 b=Workbook();s=b.active;s['A1']='SYNTHETIC isolated R1 boundary'
 for col,v in enumerate(['Отправочная марка','Наименование марки','Количество','Масса','', 'Площадь',''],1):s.cell(2,col,v)
 for col,v in ((4,'1 шт'),(5,'Всех'),(6,'1 шт'),(7,'Всех')):s.cell(3,col,v)
 for r in range(4,count+4):
  for col,v in enumerate(['Бп2-5',f'SYNTHETIC {r-3}',1,1.23,1.23,.456,.456],1):s.cell(r,col,v)
 b.save(F/f'SYNTHETIC_{count}.xlsx');b.close()

def inspect_pair(target,p):
 a=target/'ТЕСТ АКЗ №0.xlsx';b=target/'ТЕСТ Приложение к АКЗ №0.xlsx';n=len(p.payload['selected'])
 with c.native_excel() as excel:
  c.validate_saved_pair(excel,a,b,p)
  book=c.native_open(excel,b,read_only=True)
  try:
   s=book.Worksheets(1);hidden={str(r):bool(s.Rows(r).Hidden) for r in range(1,22)}
   assert all(hidden[str(r)]==(4+n<=r<=20) for r in range(1,22))
   assert int(s.PageSetup.Pages.Count)==1
   formulas={col+'21':s.Range(col+'21').Formula for col in ('C','E','G')}
   assert formulas=={col+'21':f'=SUM({col}4:{col}20)' for col in ('C','E','G')}
   values=s.Range('A4:G20').Value2
   totals={col+'21':s.Range(col+'21').Value2 for col in ('C','E','G')}
   native={'hidden':hidden,'formulas':formulas,'values':values,'totals':totals,'pages':1,'print_area':s.PageSetup.PrintArea,'read_only':bool(book.ReadOnly)}
  finally:book.Close(SaveChanges=False)
 # Independent stored XML topology: no deleted rows/merges/borders/widths.
 src=load_workbook(application);out=load_workbook(b)
 try:
  x=src.active;y=out.active
  assert set(map(str,x.merged_cells.ranges))==set(map(str,y.merged_cells.ranges))
  assert str(x.print_area)==str(y.print_area)
  assert all(copy.copy(x.cell(r,col).border)==copy.copy(y.cell(r,col).border) for r in range(1,22) for col in range(1,8))
  assert all(x.column_dimensions[col].width==y.column_dimensions[col].width for col in 'ABCDEFG')
 finally:src.close();out.close()
 native['target']=str(target);native['sha256']={str(f):sha(f) for f in (a,b)}
 return native

with patch.multiple(c,REGISTRY=reg,ACT=act,APPLICATION=application,FIXTURES=F,OUTPUT=O,PINS={act:sha(act),application:sha(application)}):
 p=c.preview('Гараж Казахстан','9/2026',vom,catalog)
 assert len(p.payload['selected'])==3 and len(p.payload['vom']['rows'])==104 and len(p.payload['vom']['hidden_rows'])==397
 expected=[('Бп2-5','37.2','1.412'),('Бп2-18','31.32','1.184'),('Бп2-19','36.74','1.393')]
 for r,(mark,mass,area) in zip(p.payload['selected'],expected):assert r['mark']==mark and r['quantity']==1 and r['mass_total']==Decimal(mass) and r['area_total']==Decimal(area)
 check('three_mark_VOM_regression',lambda:dict(p.payload['totals']))
 original=c.native_open;proof=[]
 def lock_open(excel,path,*,read_only):
  if Path(path).name=='ТЕСТ Приложение к АКЗ №0.xlsx' and not read_only:
   assert sha(path)==sha(application)
   h=win32file.CreateFile(str(path),win32con.GENERIC_READ,0,None,win32con.OPEN_EXISTING,0,None)
   try:
    try:return original(excel,path,read_only=read_only)
    except Exception as exc:
     proof.append({'path':str(path),'share':0,'valid_sha256':sha(application),'type':type(exc).__name__,'code':getattr(exc,'code',None),'message':str(exc),'cause':repr(exc.__cause__),'winerror':getattr(exc.__cause__,'winerror',None)})
     raise
   finally:h.Close()
  return original(excel,path,read_only=read_only)
 root=tk.Tk();root.withdraw()
 try:
  gui=Launcher(root);gui.checked=p
  before=set(O.iterdir())
  with patch.object(c,'native_open',side_effect=lock_open),patch('akz_stage6.gui.messagebox.showerror') as message:
   gui.generate()
  assert message.call_count==1 and 'FILE_LOCKED' in message.call_args.args[1] and 'Закройте' in message.call_args.args[1]
  assert proof[0]['winerror']==32 and proof[0]['path'] in message.call_args.args[1]
  assert before==set(O.iterdir()) and not list(O.glob('.pending-*')) and gui.checked is None
  check('real_exclusive_lock_GUI_handler_STOP_and_full_rollback',lambda:{'proof':proof,'GUI_message':message.call_args.args[1]})
 finally:root.destroy()
 # Same real preview, after releasing the OS lock.
 target=c.generate(p,'RELEASED_LOCK_N3')
 check('release_lock_complete_pair_and_R1_N3',lambda:inspect_pair(target,p))
 for n in (1,17):
  pn=c.preview('Гараж Казахстан','9/2026',F/f'SYNTHETIC_{n}.xlsx',catalog,synthetic=True)
  target=c.generate(pn,f'SYNTHETIC_N{n}')
  check(f'R1_N{n}_complete_visible_rows_totals_formulas_topology',lambda:inspect_pair(target,pn))
 before=set(O.iterdir())
 check('R1_N18_capacity_STOP_no_output',lambda:stop(lambda:c.preview('Гараж Казахстан','9/2026',F/'SYNTHETIC_18.xlsx',catalog,synthetic=True),'TEMPLATE_CAPACITY_STOP'))
 assert before==set(O.iterdir())
 def corrupt_open(excel,path,*,read_only):
  if Path(path).name=='ТЕСТ Приложение к АКЗ №0.xlsx' and not read_only:Path(path).write_bytes(b'CORRUPT XLSX TEST')
  return original(excel,path,read_only=read_only)
 with patch.object(c,'native_open',side_effect=corrupt_open):
  check('corrupt_second_XLSX_INVALID_XLSX_full_rollback',lambda:stop(lambda:c.generate(p,'CORRUPT'),'INVALID_XLSX'))
 assert before==set(O.iterdir()) and not list(O.glob('.pending-*'))
 for path in (reg,act,application,catalog,vom):
  data=path.read_bytes()
  try:
   path.write_bytes(data+b'ISOLATED CHANGED INPUT')
   check('changed_input_STOP_'+path.name,lambda:stop(lambda:c.generate(p,'CHANGED'),'INPUT_CHANGED_AFTER_PREVIEW_STOP'))
  finally:path.write_bytes(data)
 assert before==set(O.iterdir())
 sentinel={str(f):sha(f) for f in (O/'RELEASED_LOCK_N3').iterdir()}
 check('OUTPUT_EXISTS_STOP',lambda:stop(lambda:c.generate(p,'RELEASED_LOCK_N3'),'OUTPUT_ALREADY_EXISTS_STOP'))
 assert all(sha(f)==h for f,h in sentinel.items())
 # Diagnostic categories for denied access and an unclassified Excel failure.
 denied=pywintypes.error(5,'CreateFile','Access is denied')
 with patch('win32file.CreateFile',side_effect=denied):
  check('ACCESS_DENIED_distinct_code_unit',lambda:stop(lambda:c.probe_file_access(application,read_only=False),'ACCESS_DENIED'))
 check('valid_container_unclassified_error_OTHER_ERROR',lambda:stop(lambda:c.excel_file_failure(application,RuntimeError('isolated generic failure'),read_only=True,operation='test'),'OTHER_ERROR'))
 assert all(sha(f)==h for f,h in initial.items())
# Actual new current-three-mark pair from original approved inputs, separate evidence folder.
p=c.preview('Гараж Казахстан','9/2026',old['vom']['path'],cat)
assert p.seal=='e9fc12e4433feafd2982391d9d680838242807d341f425eb653e1afb8e4abbf6'
save('current_preview.json',{'seal':p.seal,'payload':p.payload})
with patch.object(c,'OUTPUT',E):target=c.generate(p,'NEW_TEST_9_2026')
check('new_current_pair_native_reopen',lambda:inspect_pair(target,p))
save('new_pair.json',{'directory':str(target),'files':{str(f):sha(f) for f in target.iterdir()}})
current={p:sha(p) for p in product};assert after==current
save('tested_product_hashes.json',{'BEFORE':product,'AFTER':after,'CURRENT':current})
save('audit.json',audit)
print('ALL TARGETED TESTS PASS',len(results),flush=True)
