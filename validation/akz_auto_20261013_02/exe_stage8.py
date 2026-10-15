from common import *
from stage8_closure.core import confirmation
from stage8_closure.recovery import confirmation as restore_confirmation
from stage8_closure.storage import digest
from openpyxl import load_workbook
from zipfile import ZipFile
import io,olefile
steps=[]
def run(name,expected='PASS',**args):
    d,p,s=task('EXE_STAGE8',name,expected,**args);steps.append({'task':name,'expected':expected,'result':str(p)});return d
def audit(sb,expected,restored=False):
    sb=Path(sb);meta=read(sb/'SANDBOX.json');reg=sb/'registry/УЗК_TEST.xlsm'
    b=load_workbook(reg,read_only=True,data_only=True)
    try:stocks=[b['УЗК'].cell(r,3).value for r in (24,25)]
    finally:b.close()
    assert stocks==expected
    with ZipFile(reg) as z:
        with olefile.OleFileIO(io.BytesIO(z.read('xl/vbaProject.bin'))) as ole:streams={'/'.join(n):hashlib.sha256(ole.openstream(n).read()).hexdigest() for n in ole.listdir()}
    assert streams==meta['native_seed']['vba_streams'] and len(streams)==11
    if restored:assert sha(reg)==meta['initial_registry_sha256']
    states=[];previous=None
    for i,p in enumerate(sorted((sb/'journal').glob('*.json'))):
        event=read(p);assert event['seal']==digest({k:v for k,v in event.items() if k!='seal'}) and event['previous']==previous and event['sequence']==i
        previous=sha(p);states.append(event['state'])
    return {'sandbox':str(sb),'stocks':stocks,'vba_streams':streams,'registry_sha256':sha(reg),'seed_sha256':meta['initial_registry_sha256'],'journal_states':states,'journal_seals_and_chain':'PASS','exact_seed_restored':restored}
fixture=run('fixture');audit(fixture['sandbox'],[12,10])
request=read(fixture['request']);request['passport']='9/2026';blocked=Path(fixture['request']).with_name('OFFICIAL_DENIED.json');save(blocked,request)
denied=run('close-preview','FAIL',sandbox=fixture['sandbox'],request=blocked);assert 'PASSPORT_9_2026' in denied['error'];audit(fixture['sandbox'],[12,10])
plan=run('close-preview',sandbox=fixture['sandbox'],request=fixture['request'])['plan'];p=OUT/'STAGE8_CLOSE_PLAN.json';save(p,plan)
closed=run('close',preview=p,confirm=confirmation(plan));assert closed['operation_status']=='CLOSED_TEST'
audit(fixture['sandbox'],[9,6])
again=run('close',preview=p,confirm=confirmation(plan));assert again['operation_status']=='ALREADY_CLOSED_TEST'
run('recover',sandbox=fixture['sandbox']);committed=audit(fixture['sandbox'],[9,6]);assert committed['journal_states'].count('COMMITTED')==1
torn=run('fixture');plan=run('close-preview',sandbox=torn['sandbox'],request=torn['request'])['plan'];p=OUT/'STAGE8_CRASH_PLAN.json';save(p,plan)
crash=run('close','FAIL',preview=p,confirm=confirmation(plan),test_fault='during_publish');assert crash['worker_returncode']==73
restore=run('restore-preview',sandbox=torn['sandbox'])['plan'];p=OUT/'STAGE8_RESTORE_PLAN.json';save(p,restore)
result=run('restore',preview=p,confirm=restore_confirmation(restore));assert result['operation_status']=='RECOVERED_ABORTED' and result['decrements']==0
recovered=audit(torn['sandbox'],[12,10],True);assert recovered['journal_states'].count('COMMITTED')==0
after=run('close-preview',sandbox=torn['sandbox'],request=torn['request']);assert after['plan']['payload']['status']=='READY_TEST'
save(OUT/'EXE_STAGE8_FINAL.json',{'status':'PASS','exe_sha256':sha(EXE),'steps':steps,'committed':committed,'recovered':recovered,'backup_locks_journal_idempotency_recovery':'PASS_ON_ARTIFICIAL_FIXTURES','real_decrements':0})
print('FROZEN Stage8 regression PASS',len(steps),'tasks',flush=True)
