"""Four original M1 cases + bounded read-only selected VOM / Stage6 smoke."""
import argparse
import ast
import hashlib
import importlib.util
import io
import json
from decimal import Decimal
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

RUN = Path(__file__).resolve().parent
ROOT = RUN.parents[1]
sys.dont_write_bytecode = True
sys.path.insert(0, str(ROOT))
access = {'outside_write_attempts': [], 'network_attempts': []}


def guard(event, args):
    if event.startswith('socket.') and event != 'socket.__new__':
        access['network_attempts'].append(event)
        raise RuntimeError('NETWORK_FORBIDDEN')
    if event == 'open' and isinstance(args[0], (str, bytes)):
        mode = args[1]
        if isinstance(mode, str) and any(c in mode for c in 'wax+'):
            p = Path(args[0].decode() if isinstance(args[0], bytes) else args[0]).resolve()
            if not p.is_relative_to(RUN):
                access['outside_write_attempts'].append(str(p))
                raise RuntimeError('OUTSIDE_WRITE_FORBIDDEN')


sys.addaudithook(guard)


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def write(name, data):
    (RUN / name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    # Compile exact source bytes: do not rely on/write a cached pyc.
    exec(compile(Path(path).read_bytes(), str(path), 'exec'), mod.__dict__)
    return mod


def workbook(names, varied=False, hidden=False):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    for cell, value in {'A1': 'Отправочная марка', 'B1': 'Наименование марки',
                        'C1': 'Количество', 'D1': 'Масса', 'F1': 'Площадь',
                        'D2': '1 шт', 'E2': 'Всех', 'F2': '1 шт', 'G2': 'Всех'}.items():
        sheet[cell] = value
    for region in ('A1:A2', 'B1:B2', 'C1:C2', 'D1:E1', 'F1:G1'):
        sheet.merge_cells(region)
    for offset, name in enumerate(names):
        quantity = offset + 1 if varied else 1
        unit = 10 + offset if varied else 10
        for col, value in enumerate([f'M{offset + 1}', name, quantity, unit, quantity * unit, 1, quantity], 1):
            sheet.cell(offset + 3, col, value)
    if hidden:
        sheet.row_dimensions[4].hidden = True
    return book


def parse(mod, names, **kwargs):
    with patch.object(mod, 'load_book', side_effect=[workbook(names, **kwargs), workbook(names, **kwargs)]), \
         patch.object(mod, 'sha256', return_value='IN_MEMORY_NO_FILE_IO'):
        result = mod.read_vom(RUN / 'MEMORY_FIXTURE_NOT_SAVED.xlsx')
    return {'connections': result.connections, 'rows': [[str(x) if isinstance(x, Decimal) else x for x in row] for row in result.rows],
            'source_rows': list(result.source_rows), 'quantity': str(result.quantity),
            'mass': str(result.mass), 'area': str(result.area)}


def run(phase):
    current_path = ROOT / 'passport_app/vom.py'
    checkpoint = read(RUN / 'PRECHECK.json')
    expected_hash = checkpoint['current_sha256'] if phase == 'before' else checkpoint['accepted_sha256']
    assert sha(current_path) == expected_hash
    current = module('passport_app.m1_restoration_under_test', current_path)
    reference = module('passport_app.m1_accepted_reference', RUN / 'vom.accepted.py')
    cases = [('normal', ['Балка', 'Балка'], 'Балка'),
             ('edge_spaces', [' Балка '], ' Балка '),
             ('exact_unique_first_order', [' Балка ', 'Балка', ' Балка '], ' Балка , Балка'),
             ('empty_and_whitespace', [None, '', '   ', '\t ', 'Балка'], 'Балка')]
    old_cases = read(ROOT / 'validation/stage7_resolve_20261009_165600/CONNECTIONS_TARGETED_REPRODUCTION.json')['cases']
    assert [(c[0], c[1], c[2]) for c in cases] == [(c['case'], c['input'], c['expected']) for c in old_cases]
    tests = []
    for name, names, expected in cases:
        actual, accepted = parse(current, names), parse(reference, names)
        assert accepted['connections'] == expected
        assert all(actual[k] == accepted[k] for k in actual if k != 'connections')
        tests.append({'case': name, 'input': names, 'expected': expected, 'actual': actual,
                      'pass': actual['connections'] == expected, 'non_connections_match_accepted': True})
    varied = parse(current, [' Связь ', 'Колонна', 'Связь', ' Связь '], varied=True, hidden=True)
    accepted_varied = parse(reference, [' Связь ', 'Колонна', 'Связь', ' Связь '], varied=True, hidden=True)
    assert varied['source_rows'] == [3, 4, 5, 6]  # preserve Stage1 treatment, including hidden row 4
    assert varied['quantity'] == '10' and varied['mass'] == '120' and varied['area'] == '10'
    assert all(varied[k] == accepted_varied[k] for k in varied if k != 'connections')
    if phase == 'after':
        assert all(c['pass'] for c in tests)
        assert varied == accepted_varied
        prior = read(RUN / 'TESTS_BEFORE.json')
        for before, after in zip(prior['cases'], tests):
            assert all(before['actual'][k] == after['actual'][k] for k in after['actual'] if k != 'connections')
        assert all(prior['varied_numerics_hidden_row'][k] == varied[k] for k in varied if k != 'connections')
    else:
        assert [c['case'] for c in tests if not c['pass']] == ['edge_spaces', 'exact_unique_first_order']

    from akz_stage6 import core, part2
    baseline = read(ROOT / 'validation/stage7_resolve_20261009_165600/CURRENT_VOM_NATIVE_READONLY.json')
    selected_path = Path(baseline['source']['path'])
    assert sha(selected_path) == baseline['sha256']
    # Dynamic tripwire proves this native read does not call Stage1 read_vom.
    from passport_app import vom as live
    with patch.object(live, 'read_vom', side_effect=AssertionError('UNEXPECTED_STAGE6_READ_VOM_CALL')):
        selected = core.visible_vom(selected_path)
    quantity = sum(Decimal(str(r['quantity'])) for r in selected['rows'])
    raw_mass = sum(Decimal(str(r['mass_total'])) for r in selected['rows'])
    assert len(selected['rows']) == 104 and len(selected['hidden_rows']) == 397
    assert quantity == 288 and raw_mass.quantize(Decimal('0.01')) == Decimal('15822.14')
    assert selected['hidden_rows'] == baseline['source']['hidden_data_rows']
    assert [(r['row'], r['mark'], r['quantity']) for r in selected['rows']] == [
        (r['row'], r['mark'], r['quantity']) for r in baseline['source']['rows']]
    assert sha(selected_path) == baseline['sha256']
    native = {'pass': True, 'visible_rows': 104, 'hidden_rows': 397, 'quantity': str(quantity),
              'mass_kg': str(raw_mass.quantize(Decimal('0.01'))), 'raw_mass': str(raw_mass),
              'sha256_before_after': baseline['sha256'], 'source': selected}
    if phase == 'after':
        assert native == read(RUN / 'SELECTED_VOM_BEFORE.json')
    write(f'SELECTED_VOM_{phase.upper()}.json', native)

    # Only four related existing in-memory tests; never discover/run the full suite.
    s1 = module('stage6_m1_smoke_part1', ROOT / 'tests/test_stage6.py')
    s2 = module('stage6_m1_smoke_part2', ROOT / 'tests/test_stage6_part2.py')
    suite = unittest.TestSuite([
        s1.Stage6Boundaries('test_numeric_zero_and_blank_separately'),
        s1.Stage6Boundaries('test_duplicates_and_vom_order'),
        s1.Stage6Boundaries('test_exact_mark_no_fuzzy_match'),
        s2.Part2Boundaries('test_repeated_rows_order_and_exact_source_totals'),
    ])
    stream = io.StringIO()
    result = unittest.TextTestRunner(stream=stream, verbosity=2).run(suite)
    assert result.wasSuccessful(), stream.getvalue()
    files = [*sorted((ROOT / 'akz_stage6').glob('*.py')), ROOT / 'passport_app/registry.py',
             ROOT / 'passport_app/tables.py', ROOT / 'passport_app/excel.py', ROOT / 'passport_app/config.py']
    imports = {str(p.relative_to(ROOT)): {'sha256': sha(p), 'imports': [ast.unparse(n) for n in
               ast.walk(ast.parse(p.read_text(encoding='utf-8-sig'))) if isinstance(n, (ast.Import, ast.ImportFrom))]} for p in files}
    assert not any('passport_app.vom' in imp or 'read_vom' in imp for f in imports.values() for imp in f['imports'])
    write(f'STAGE6_SMOKE_{phase.upper()}.json', {'pass': True, 'tests_run': result.testsRun,
        'output': stream.getvalue(), 'bounded_import_evidence': imports,
        'native_visible_vom_with_stage1_tripwire': 'PASS', 'generation': 'NOT_RUN',
        'limitation': 'Related native input and in-memory numeric/membership smoke only, not full Stage6 acceptance.'})
    write(f'TESTS_{phase.upper()}.json', {'phase': phase, 'product_sha256': sha(current_path),
        'm1_passed': sum(c['pass'] for c in tests), 'm1_total': 4, 'cases': tests,
        'varied_numerics_hidden_row': varied, 'quantity_mass_area_unchanged': True,
        'stage6_smoke': result.wasSuccessful(), 'selected_vom_pass': True,
        'fixture_method': 'Actual read_vom and locate_layout; only load_book/hash mocked for memory fixtures. Real selected XLS separately read via unchanged Stage6 COM adapter.',
        'access_guard': access})
    print(json.dumps({'phase': phase, 'M1': f"{sum(c['pass'] for c in tests)}/4", 'stage6_smoke': '4/4',
                      'selected_VOM': '104 visible / 397 hidden / 288 / 15822.14'}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('phase', choices=('before', 'after'))
    run(parser.parse_args().phase)
