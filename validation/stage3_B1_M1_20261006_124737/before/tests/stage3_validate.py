"""Targeted Stage 3 validation. All writes are below the Stage 3 validation root."""
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import sys
import time
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from passport_app.catalog import CatalogService, CATALOG, HEADERS, key, new_book_bytes, sheet_part, read_catalog
from passport_app.config import Settings, PassportError
from passport_app.excel import sha256, excel_session, open_excel
import openpyxl

CHECKPOINT = json.loads((ROOT / 'STAGE3_CHECKPOINT.json').read_text(encoding='utf-8'))
OUT = Path(CHECKPOINT['validation_root'])
PRODUCTION = Path(CHECKPOINT['production_path'])
RESULTS = []
REGISTRY = OUT / 'ВЗМК.xlsx'
PROJECT = 'Гараж Казахстан'


def production_guard():
    assert sha256(PRODUCTION) == CHECKPOINT['production_sha256'], 'STOP: production ВЗМК.xlsx изменился'


def md(*rows):
    return '| Марка | Профиль |\n|---|---|\n' + '\n'.join(f'| {m} | {p} |' for m, p in rows)


def case(name, rows=None):
    production_guard()
    folder = OUT / name / 'Документы' / 'Гараж казахстан'
    (folder / 'Паспорт').mkdir(parents=True, exist_ok=False)
    settings = Settings(folder.parent, REGISTRY, {PROJECT: {'passport_root': 'Гараж казахстан/Паспорт'}})
    if rows is not None:
        b = openpyxl.load_workbook(BytesIO(new_book_bytes()))
        for row in rows:
            b.active.append(row)
        for row in b.active.iter_rows(min_row=2):
            row[2].number_format = '@'
            row[3].number_format = '0.00'
        b.active.auto_filter.ref = f'A1:D{b.active.max_row}'
        b.save(folder / CATALOG)
        b.close()
    return CatalogService(settings), folder / CATALOG


def values(path):
    b = openpyxl.load_workbook(path)
    try:
        return list(b.active.values)
    finally:
        b.close()


def error(operation, contains=None):
    try:
        operation()
    except PassportError as exc:
        if contains:
            assert contains.casefold() in str(exc).casefold(), str(exc)
        return str(exc)
    raise AssertionError('Expected a safe PassportError')


def check(name, function):
    production_guard()
    evidence = function()
    production_guard()
    RESULTS.append({'scenario': name, 'status': 'PASS', 'evidence': evidence})
    print(name + ': PASS', flush=True)


def creation():
    service, path = case('creation')
    plan = service.preview(PROJECT, md(('К2', '20Ш1'), ('Б4-1', 'Гнз120x60x6')))
    assert [r.status for r in plan.rows] == ['NEW', 'NEW'] and not path.exists()
    result = service.apply(plan)
    assert values(path) == [HEADERS, ('К2', '20Ш1', None, None), ('Б4-1', 'Гнз120x60x6', None, None)]
    assert Path(result['backup']).read_text() == 'Catalog absent before creation.\n'
    b = openpyxl.load_workbook(path)
    assert b.active.freeze_panes == 'A2' and b.active.auto_filter.ref == 'A1:D3'
    assert b.active.max_column == 4
    b.close()
    return result


def same():
    service, path = case('same', [('К1', '35Б1', 'RAL 7004', 0)])
    before = sha256(path)
    plan = service.preview(PROJECT, md(('К1', '35Б1')))
    assert plan.rows[0].status == 'SAME'
    assert service.apply(plan)['unchanged'] == 1 and sha256(path) == before
    assert values(path)[1] == ('К1', '35Б1', 'RAL 7004', 0)
    assert not (path.parent / '.stage3-backups').exists()
    return {'before_after_sha256': before, 'manual_values': values(path)[1][2:], 'no_write': True}


def fill():
    service, path = case('fill', [('К1', None, 'RAL 7004', 0)])
    before = path.read_bytes()
    plan = service.preview(PROJECT, md(('К1', '35Б1')))
    assert plan.rows[0].status == 'FILL'
    result = service.apply(plan)
    assert result['filled'] == 1
    assert values(path)[1] == ('К1', '35Б1', 'RAL 7004', 0)
    assert Path(result['backup']).read_bytes() == before
    return result


def conflicts(confirm):
    service, path = case('conflict_' + str(confirm), [('К1', '35Б1', 'RAL 7004', 0)])
    before = path.read_bytes()
    plan = service.preview(PROJECT, md(('К1', '35Б2')))
    assert plan.rows[0].status == 'CONFLICT'
    result = service.apply(plan, {key('К1')} if confirm else ())
    assert values(path)[1] == ('К1', '35Б2' if confirm else '35Б1', 'RAL 7004', 0)
    if confirm:
        assert result['changed'] == 1 and Path(result['backup']).read_bytes() == before
        _, old = sheet_part(before)
        _, new = sheet_part(path.read_bytes())
        import re
        for address in ('C2', 'D2', 'A2'):
            pattern = rf'<c\b[^>]*\br="{address}"[^>]*(?:/>|>.*?</c>)'
            assert re.search(pattern, old).group() == re.search(pattern, new).group()
    else:
        assert path.read_bytes() == before and result['skipped'] == 1
    return result | {'manual_values': values(path)[1][2:], 'manual_cell_xml_byte_preserved': True}


def check_sentinel():
    service, path = case('check_required', [('К1', '35Б1', 'RAL 7004', 0), ('К3', None, 'Грунт', 7)])
    plan = service.preview(PROJECT, md(('К1', 'ПРОВЕРИТЬ'), ('К2', 'ПРОВЕРИТЬ'), ('К3', 'ПРОВЕРИТЬ')))
    assert all(r.status == 'CHECK_REQUIRED' for r in plan.rows)
    assert [r.confirmable for r in plan.rows] == [False, True, True]
    before = sha256(path)
    service.apply(plan)
    assert sha256(path) == before
    error(lambda: service.apply(plan, {key('К1')}), 'недоступной')
    result = service.apply(plan, {key('К2'), key('К3')})
    assert values(path) == [HEADERS, ('К1', '35Б1', 'RAL 7004', 0),
                            ('К3', 'ПРОВЕРИТЬ', 'Грунт', 7), ('К2', 'ПРОВЕРИТЬ', None, None)]
    return result


def duplicates():
    service, path = case('duplicates')
    plan = service.preview(PROJECT, md(('К1', '35Б1'), ('К1', '35Б1'), ('К2', '20Ш1')))
    assert len(plan.rows) == 2 and plan.rows[0].warning
    result = service.apply(plan)
    assert [r[0] for r in values(path)[1:]] == ['К1', 'К2']
    before = sha256(path)
    bad = service.preview(PROJECT, md(('К1', '35Б1'), ('К1', '35Б2'), ('К9', '45Б1')))
    assert bad.rows[0].status == 'ERROR'
    message = error(lambda: service.apply(bad), 'ошибки')
    assert sha256(path) == before
    return result | {'duplicate_conflict': message, 'warning': plan.rows[0].warning}


def preview_no_write():
    service, path = case('preview', [('К1', '35Б1', 'RAL 7004', 0)])
    before = {str(p): sha256(p) for p in path.parent.rglob('*') if p.is_file()}
    service.preview(PROJECT, md(('К1', '35Б2'), ('К2', '20Ш1')))
    after = {str(p): sha256(p) for p in path.parent.rglob('*') if p.is_file()}
    assert before == after
    return {'file_inventory_hashes_equal': True}


def malformed():
    service, path = case('malformed')
    messages = [error(lambda text=t: service.preview(PROJECT, text)) for t in
                ('', '| Марка | Профиль |\n| К1 | 35Б1 |', 'Марка\tПрофиль\tКраска\nК1\t35Б1\t0', 'текст')]
    for t in (md(('К1', '')), md(('К1', '35Б1')) + '\n| К1 | 35Б1 | лишнее |'):
        plan = service.preview(PROJECT, t)
        assert plan.rows[0].status == 'ERROR'
        messages.append(error(lambda: service.apply(plan)))
    assert not path.exists()
    return messages


def normalization():
    service, path = case('normalization', [('  К  1  ', 'Гнз120х60×6', 'RAL 7004', 0)])
    plan = service.preview(PROJECT, 'Марка\tПрофиль\nк 1\tГнз120x60x6')
    assert plan.rows[0].status == 'SAME'
    service.apply(plan)
    plan = service.preview(PROJECT, md(('к 1', 'Гнз120x60x8')))
    service.apply(plan, {key('к 1')})
    assert values(path)[1] == ('  К  1  ', 'Гнз120x60x8', 'RAL 7004', 0)
    ambiguous, ambiguous_path = case('ambiguous', [('К1', '35Б1', 'A', 0), (' к1 ', '35Б2', 'B', 1)])
    before = sha256(ambiguous_path)
    message = error(lambda: ambiguous.preview(PROJECT, md(('К1', '35Б1'))), 'неоднозначность')
    assert sha256(ambiguous_path) == before
    return {'stored_original_mark': values(path)[1][0], 'ambiguity': message, 'tab_separated': 'PASS'}


def locks():
    import win32con
    import win32file
    service, path = case('locked', [('К1', '35Б1', 'RAL 7004', 0)])
    plan = service.preview(PROJECT, md(('К1', '35Б2')))
    before = sha256(path)
    handle = win32file.CreateFile(str(path), win32con.GENERIC_READ, win32con.FILE_SHARE_READ,
                                  None, win32con.OPEN_EXISTING, 0, None)
    try:
        message = error(lambda: service.apply(plan, {key('К1')}), 'занят')
    finally:
        handle.Close()
    assert sha256(path) == before
    win32file.SetFileAttributes(str(path), win32con.FILE_ATTRIBUTE_READONLY)
    try:
        readonly = error(lambda: service.apply(plan, {key('К1')}))
    finally:
        win32file.SetFileAttributes(str(path), win32con.FILE_ATTRIBUTE_NORMAL)
    assert sha256(path) == before
    return {'native_lock': message, 'native_readonly_attribute': readonly, 'unchanged_sha256': before}


def excel_lock():
    service, path = case('excel_locked', [('К1', '35Б1', 'RAL 7004', 0)])
    before = sha256(path)
    plan = service.preview(PROJECT, md(('К1', '35Б2')))
    with excel_session() as excel:
        book = open_excel(excel, path)
        try:
            message = error(lambda: service.apply(plan, {key('К1')}))
        finally:
            book.Close(SaveChanges=False)
    assert sha256(path) == before
    return {'actual_Excel_COM_open': True, 'message': message, 'unchanged_sha256': before}


def stale():
    service, path = case('stale', [('К1', '35Б1', 'RAL 7004', 0)])
    plan = service.preview(PROJECT, md(('К1', '35Б2')))
    b = openpyxl.load_workbook(path)
    b.active['C2'] = 'Ручная правка'
    b.save(path)
    b.close()
    before = sha256(path)
    message = error(lambda: service.apply(plan, {key('К1')}), 'изменился')
    assert sha256(path) == before
    return message


def failures():
    service, path = case('safe_failure', [('К1', '35Б1', 'RAL 7004', 0)])
    plan = service.preview(PROJECT, md(('К1', '35Б2')))
    before = sha256(path)
    with patch('passport_app.catalog.verify_patch', side_effect=PassportError('Saved verification failure')):
        verification = error(lambda: service.apply(plan, {key('К1')}))
    with patch('passport_app.catalog.atomic_replace', side_effect=PermissionError('Replace denied')):
        replacement = error(lambda: service.apply(plan, {key('К1')}))
    with patch('passport_app.catalog.durable_write', side_effect=PermissionError('Folder write denied')):
        folder = error(lambda: service.apply(plan, {key('К1')}))
    assert sha256(path) == before and not list(path.parent.glob('.stage3-*.tmp.xlsx'))
    assert all(sha256(p) == before for p in (path.parent / '.stage3-backups').glob('*.xlsx'))
    return {'verification_injection': verification, 'replace_injection': replacement,
            'folder_write_injection': folder, 'original_preserved': True, 'temporary_cleaned': True}


def schema_and_folder():
    service, path = case('schema', [('К1', '35Б1', 'A', 0)])
    b = openpyxl.load_workbook(path)
    b.active['E1'] = 'Лишний столбец'
    b.save(path)
    b.close()
    before = sha256(path)
    schema = error(lambda: service.preview(PROJECT, md(('К1', '35Б2'))), 'структура')
    assert sha256(path) == before
    (path.parent.parent / 'ГАРАЖ  КАЗАХСТАН').mkdir()
    ambiguity = error(lambda: service.preview(PROJECT, md(('К1', '35Б2'))), 'неоднозначная папка')
    from passport_app.catalog import project_folder
    missing = error(lambda: project_folder(service.settings, 'Отсутствующий проект'), 'правило создания')
    return {'schema': schema, 'folder_ambiguity': ambiguity, 'missing_folder': missing}


def registry_changed():
    service, path = case('registry_changed')
    # Change a disposable registry, never production.
    copied = path.parent / 'registry_copy.xlsx'
    shutil.copy2(REGISTRY, copied)
    service = CatalogService(Settings(service.settings.documents_root, copied, service.settings.projects))
    plan = service.preview(PROJECT, md(('К1', '35Б1')))
    with copied.open('ab') as f:
        f.write(b'changed')
    message = error(lambda: service.apply(plan), 'STOP')
    assert not path.exists()
    return message


def literal_text():
    service, path = case('literal_text')
    service.apply(service.preview(PROJECT, md(('=К1', '=1+2'), ('+К2', '@профиль'))))
    b = openpyxl.load_workbook(path)
    try:
        assert b.active['A2'].data_type == b.active['B2'].data_type == 's'
        assert b.active['B2'].value == '=1+2'
    finally:
        b.close()
    return {'input_written_as_text': True, 'no_formula_interpretation': True}


def gui_smoke(case_name='gui'):
    from passport_app.catalog_gui import CatalogApplication
    from passport_app.gui import Application
    import passport_app.vik
    import stage2
    import stage3
    service, path = case(case_name, [('К1', '35Б1', 'RAL 7004', 0), ('К3', '45Б1', 'B', 1)])
    app = CatalogApplication(service.settings)
    app.withdraw()
    try:
        app.update()
        assert len(app.combo['values']) == len(CHECKPOINT['projects'])
        app.project_var.set(PROJECT)
        app.combo.event_generate('<<ComboboxSelected>>')
        app.input.insert('1.0', md(('К1', '35Б2'), ('К3', '45Б2'), ('К2', '20Ш1')))
        app.update()
        before = sha256(path)
        app.check_button.invoke()
        app.update()
        assert sha256(path) == before and app.plan is not None
        app.tree.selection_set('0')
        app.toggle_selected()
        assert app.confirmations == {key('К1')}
        with patch('passport_app.catalog_gui.messagebox.askyesno', return_value=True):
            app.apply_button.invoke()
        app.update()
        assert values(path)[1:] == [('К1', '35Б2', 'RAL 7004', 0), ('К3', '45Б1', 'B', 1), ('К2', '20Ш1', None, None)]
        result = app.result_var.get()
        app.check_button.invoke()
        app.update()
        app.input.insert('end', '\n| К4 | 25Ш1 |')
        app.update()
        assert app.plan is None and str(app.apply_button['state']) == 'disabled'
        app.check_button.invoke()
        app.update()
        assert app.plan is not None
        app.project_var.set('Гранит')
        app.combo.event_generate('<<ComboboxSelected>>')
        app.update()
        assert app.plan is None and str(app.apply_button['state']) == 'disabled'
    finally:
        app.destroy()
    settings = Settings.load(ROOT / 'settings.json')
    stage1 = Application(settings)
    stage1.withdraw()
    try:
        deadline = time.monotonic() + 30
        while stage1.busy and time.monotonic() < deadline:
            stage1.update()
            time.sleep(.02)
        assert not stage1.busy
        assert len(stage1.combo['values']) == len(CHECKPOINT['projects'])
        stage1.project_var.set(PROJECT)
        stage1.combo.event_generate('<<ComboboxSelected>>')
        while stage1.busy and time.monotonic() < deadline:
            stage1.update()
            time.sleep(.02)
        assert not stage1.busy and '1/2026' in stage1.passports_var.get()
        smoke = stage1.passports_var.get()
    finally:
        stage1.destroy()
    return {'Stage3_Tk_controls': result, 'input_and_project_invalidation': True,
            'Stage1_start_and_production_read_only_selection': smoke,
            'Stage2_module_imports': 'PASS', 'passport_generation': 'NOT_RUN', 'VIK_generation': 'NOT_RUN',
            'human_GUI_E2E': 'NOT_RUN'}


def main():
    production_guard()
    shutil.copy2(PRODUCTION, REGISTRY)
    assert sha256(REGISTRY) == CHECKPOINT['production_sha256']
    for name, func in [('1,2,3,12,15,20 creation/columns/NEW/order/backup/reopen', creation),
                       ('5,13,14 SAME/manual columns/zero/no rewrite', same),
                       ('4,13,14,15 FILL/manual columns/backup', fill),
                       ('6 CONFLICT without confirmation', lambda: conflicts(False)),
                       ('7,13,14 CONFLICT confirmed/manual cell XML', lambda: conflicts(True)),
                       ('8,9 CHECK_REQUIRED new/empty/normal protected/confirmation', check_sentinel),
                       ('10,11 duplicates same/conflicting', duplicates),
                       ('17 preview no write', preview_no_write),
                       ('18 malformed input', malformed),
                       ('19 normalization ambiguity/original spelling/TSV', normalization),
                       ('16 native file lock/read-only attribute', locks),
                       ('16 actual Excel-open locked workbook', excel_lock),
                       ('stale catalog preview', stale),
                       ('backup/temp/verify/replace failures', failures),
                       ('schema/folder ambiguity/missing folder STOP', schema_and_folder),
                       ('registry checkpoint changed STOP on copy', registry_changed),
                       ('literal text formula safety', literal_text),
                       ('Stage3 GUI controls + Stage1/2 minimal smoke', gui_smoke)]:
        check(name, func)
    production_guard()
    (OUT / 'targeted_results.json').write_text(json.dumps({'status': 'PASS', 'checks': RESULTS,
        'production_before': CHECKPOINT['production_sha256'], 'production_after': sha256(PRODUCTION),
        'required_scenarios': '20/20 PASS', 'test_groups': len(RESULTS)}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'status': 'PASS', 'test_groups': len(RESULTS), 'evidence': str(OUT / 'targeted_results.json')}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        (OUT / 'targeted_failure.json').write_text(json.dumps({'status': 'FAIL', 'completed': RESULTS,
            'error': repr(exc)}, ensure_ascii=False, indent=2), encoding='utf-8')
        raise
