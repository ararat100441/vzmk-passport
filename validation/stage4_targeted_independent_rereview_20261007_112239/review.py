"""Read-only current-artifact review; all mutation/fault tests use fresh fixtures."""
from pathlib import Path
from collections import Counter,defaultdict
from copy import deepcopy
from decimal import Decimal
from datetime import datetime,timezone
from unittest.mock import patch
import hashlib,importlib,io,json,os,shutil,subprocess,sys,traceback,zipfile
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
FIX=ROOT/'validation/stage4_fix_m1_m2_20261007_090000'
OLD=ROOT/'validation/stage4_independent_review_20261007_0119'
sys.path.insert(0,str(ROOT))
sys.stdout.reconfigure(encoding='utf-8')
sys.dont_write_bytecode=True
def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def digest(data): return hashlib.sha256(data).hexdigest()
def sha(p): return digest(Path(p).read_bytes())
def save(n,v): (OUT/n).write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def check(ok,message='Independent assertion failed'):
    if not ok: raise AssertionError(message)
events={'network':[],'production':[],'outside_writes':[]}
def audit(event,args):
    if event in ('socket.connect','socket.getaddrinfo','socket.sendto','socket.connect_ex'):
        events['network'].append(event); raise AssertionError('Offline guard')
    if event in ('open','os.rename','os.remove','os.mkdir','os.rmdir'):
        paths=[a for a in args if isinstance(a,(str,bytes))]
        for arg in paths:
            text=os.fsdecode(arg).replace('\\','/')
            if '/Работа/ВЗМК.xlsx' in text or '/Работа/Профили.xlsx' in text:
                events['production'].append(event); raise AssertionError('Production forbidden')
        writes=event!='open'
        if event=='open':
            mode,flags=args[1:3]
            writes=bool(mode and any(c in mode for c in 'wax+')) or bool(flags & (os.O_WRONLY|os.O_RDWR|os.O_CREAT|os.O_TRUNC|os.O_APPEND))
        if writes:
            for arg in paths[:2] if event=='os.rename' else paths[:1]:
                p=Path(os.fsdecode(arg)).resolve()
                if not p.is_relative_to(OUT):
                    events['outside_writes'].append(str(p)); raise AssertionError('Write outside isolated rereview')
sys.addaudithook(audit)
import profiles_stage4.storage as s
import profiles_stage4.windows_generation as wg
from profiles_stage4.model import Catalog,ProfileError,calculate_uzk,frequency_angle,RULE,NOT_CONFIGURED,default_user_state
from profiles_stage4.offline import load_sources,no_network
from openpyxl import load_workbook
tests=[]
def case(name,fn):
    try:
        result=fn()
        tests.append({'name':name,'status':'PASS','evidence':result})
        print('PASS '+name,flush=True)
    except Exception as e:
        tests.append({'name':name,'status':'FAIL','error':str(e),'traceback':traceback.format_exc()})
        print('FAIL '+name+': '+str(e),flush=True)
    save('tests_progress.json',tests)
def rejected(fn,contains=None):
    try: fn()
    except (ProfileError,OSError,AssertionError) as e:
        if contains: check(contains in str(e),str(e))
        return str(e)
    raise AssertionError('Expected operation to fail closed')

manifest=read(FIX/'REVIEW_MANIFEST.json')
real=s.DEFAULT_WORKBOOK
names=[real.name,s.STATE,s.CATALOG,s.COMMIT]
expected={real.name:'2a07eeacd56793df749ddb848563336a413f54beaae9b6d002e438c35452b98e',s.STATE:'a6031a6f08502f4cd1e6f99ef0ac27d284f2f2a3e5d1578a006fb134448dc67b',s.CATALOG:'350564a44f4c0577ed540cbcf595f397ca2b33ab5e7b7bf2dc4e629d8021aa6d',s.COMMIT:'fb075fd8a622f72d48d231ae9391baba7bcd3d37f297abff38bb3a0da2959508'}
check({n:sha(real.parent/n) for n in names}==expected,'STOP reviewed DEV bytes differ')
protected={str(Path(p) if Path(p).is_absolute() else ROOT/p):h for p,h in manifest['protected_hashes'].items()}
product={str(ROOT/p):h for p,h in manifest['product_hashes'].items()}
evidence={str(FIX/p):(v if isinstance(v,str) else v['sha256']) for p,v in manifest['evidence_files'].items()}
baseline={**protected,**product,**evidence,**{str(real.parent/n):h for n,h in expected.items()},str(FIX/'REVIEW_MANIFEST.json'):sha(FIX/'REVIEW_MANIFEST.json')}
check(all(Path(p).resolve().is_relative_to(ROOT) for p in baseline))
check({p:sha(p) for p in baseline}==baseline,'STOP baseline mismatch')
save('hashes_before.json',baseline)
runtime=[p for p in product if not Path(p).is_relative_to(FIX)]
oldhash=read(FIX/'reviewed_baseline.json')
changes={str(Path(p).relative_to(ROOT)):('ADDED' if p not in oldhash else 'CHANGED') for p in runtime if oldhash.get(p)!=product[p]}
check(set(changes)=={'profiles_stage4/model.py'.replace('/',os.sep),'profiles_stage4/storage.py'.replace('/',os.sep),'profiles_stage4/workbook.mjs'.replace('/',os.sep),'STAGE4_README.md','profiles_stage4/windows_generation.py'.replace('/',os.sep)})
save('actual_state.json',{'workbook':str(real),'DEV_artifacts':expected,'runtime_files':{str(Path(p).relative_to(ROOT)):product[p] for p in runtime},'changed_set':changes,'protected_files':len(protected),'fix_evidence_files':len(evidence),'initial_gate':'PASS'})
cat=s.load_catalog(real)
records=read(real.parent/s.CATALOG)['records']
state=cat.user_state
groups=defaultdict(list)
for r in records: groups[(r['key'],r['standard'])].append(r)
core={'Тип','h','b','b2','s','t','D'}
disputes=[]
for identity,members in sorted(groups.items()):
    vals=defaultdict(set)
    for m in members:
        for k,v in {**m['geometry'],**m.get('quality',{})}.items():
            if v is not None: vals[k].add(Decimal(str(v)))
    diffs=sorted(k for k,v in vals.items() if len(v)>1)
    if len({m['family'] for m in members})>1: diffs.append('Тип')
    if diffs: disputes.append({'key':identity[0],'standard':identity[1],'fields':diffs,'class':'BLOCKING' if core.intersection(diffs) else 'NON_BLOCKING','raw':members})

def m1_current():
    check(len(disputes)==15)
    check(Counter(d['class'] for d in disputes)=={'BLOCKING':8,'NON_BLOCKING':7})
    for d in disputes:
        result=cat.lookup(d['key'],d['standard']); row=result['profile']
        check(row['provenance']==groups[(d['key'],d['standard'])])
        check(sorted(row['unresolved_fields'])==sorted(d['fields']))
        check(row['disagreement_class']==d['class'])
        if d['class']=='BLOCKING':
            check(result['status']=='CONFLICT' and row['geometry']=={} and row['uzk_rule']==NOT_CONFIGURED)
            rejected(lambda:calculate_uzk(result))
        else:
            check(result['status']=='FOUND' and row['status']=='СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ')
            check(all(k not in row['geometry'] and k not in row['quality'] for k in d['fields']))
            for field in ('h','b','b2','s','t','D'):
                values={m['geometry'][field] for m in d['raw'] if m['geometry'].get(field) is not None}
                if values: check(len(values)==1 and row['geometry'][field]==next(iter(values)))
    old=read(FIX/'baseline_generation/profiles.catalog.json')
    check(old['records']==records and old['sources']==cat.sources,'Source/raw regression')
    save('disagreements_independent.json',disputes)
    return {'all':15,'blocking':8,'non_blocking':7,'raw_sources_unchanged':True}
case('M1 all current disputes: independent classification, lookup, provenance',m1_current)

beam=deepcopy(next(r for r in records if r['key']=='35Б1' and r['standard']=='ГОСТ Р 57837-2017'))
def pair_case(fields,blocking=False,family=False):
    a,b=deepcopy(beam),deepcopy(beam)
    a['record_id']='INDEPENDENT-A'; b['record_id']='INDEPENDENT-B'
    b['source']='INDEPENDENT-B'; b['source_id']='INDEPENDENT-B'
    for field in fields:
        place='geometry' if field in ('h','b','b2','s','t','D','R','r') else 'quality'
        a[place][field]=a[place].get(field) or 10
        b[place][field]=a[place][field]+1
    if family: b['family']='Швеллер'
    c=deepcopy(a); c['standard']='ДРУГОЙ НД'; c['record_id']='INDEPENDENT-C'
    user=default_user_state()
    local=Catalog([a,b,c],user_state=user)
    result=local.lookup(a['key'],a['standard'])
    if blocking or family:
        check(result['status']=='CONFLICT' and result['profile']['geometry']=={})
        check(local.lookup(a['key'])['status']=='CONFLICT','Preference must not override source conflict')
        rejected(lambda:calculate_uzk(result))
    else:
        check(result['status']=='FOUND' and result['profile']['uzk_rule']==RULE)
        check(calculate_uzk(result)['rows'][0]['L']==328)
        check(set(fields).issubset(result['profile']['unresolved_fields']))
        check(all(f not in result['profile']['geometry'] and f not in result['profile']['quality'] for f in fields))
        check(len(result['profile']['provenance'])==2)
    return {'fields':fields,'lookup':result['status'],'unresolved':result['profile']['unresolved_fields']}
for f in ['R','r','area_cm2','mass_kg_m','reference_only']:
    case('M1 auxiliary-only '+f,lambda f=f:pair_case([f]))
for f in ['h','b','b2','s','t','D']:
    case('M1 blocking core '+f,lambda f=f:pair_case([f],True))
case('M1 family disagreement',lambda:pair_case([],True,True))
case('M1 mixed core + R/r',lambda:pair_case(['h','R','r'],True))
case('M1 mixed core + area/mass',lambda:pair_case(['b','area_cm2','mass_kg_m'],True))

def counts_golden():
    counts={'raw':len(records),'canonical':len(groups),'raw_by_source':dict(Counter(r['source'] for r in records)),
        'family':dict(Counter(r['family'] for r in cat.rows)),'standard':dict(Counter(r['standard'] for r in cat.rows)),
        'status':dict(Counter(r['status'] for r in cat.rows)), 'disagreements':len(disputes),
        'blocking':sum(d['class']=='BLOCKING' for d in disputes),'non_blocking':sum(d['class']=='NON_BLOCKING' for d in disputes),
        'uzk':sum(r['uzk_rule']==RULE for r in cat.rows),'manual':len(state['manual_rows'])}
    bykey=defaultdict(list)
    for r in cat.rows: bykey[r['key']].append(r)
    prefs={p['key'] for p in state['preferences'] if p['status']=='ПОДТВЕРЖДЕНО ПОЛЬЗОВАТЕЛЕМ'}
    ambiguous=[k for k,rows in bykey.items() if k not in prefs and all(r['status']!='КОНФЛИКТ' for r in rows) and len({(r['family'],*(r['geometry'].get(f) for f in {'Двутавр':('h','b','s','t'),'Швеллер':('h','b','s','t'),'Уголок равнополочный':('b','t'),'Уголок неравнополочный':('b','b2','t'),'Квадратный профиль':('h','b','t'),'Прямоугольный профиль':('h','b','t'),'Круглая труба':('D','s')}[r['family']])) for r in rows})>1]
    counts['ambiguous']=len(ambiguous)
    check(len(ambiguous)==51 and all(cat.lookup(k)['status']=='AMBIGUOUS' for k in ambiguous))
    check((counts['raw'],counts['canonical'],counts['uzk'],counts['manual'])==(3037,2849,480,0))
    check(counts['status']=={'ИЗ ИСТОЧНИКА':2662,'ПРОВЕРЕНО_ДВУМЯ_ИСТОЧНИКАМИ':172,'СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ':7,'КОНФЛИКТ':8})
    reduction=Counter()
    for members in groups.values():
        if len(members)>1: reduction['cross_source' if len({m['source'] for m in members})>1 else 'same_source']+=len(members)-1
    check(reduction=={'cross_source':187,'same_source':1})
    counts['reductions']=dict(reduction)
    golden={}
    for nd,g,uzk in [('ГОСТ Р 57837-2017',[346,174,6,9],[(328,6),(348,9)]),('ГОСТ 26020-83',[346,155,6.2,8.5],[(329,6.2),(310,8.5)])]:
        r=s.lookup('35Б1',nd,workbook=real)
        check(r['status']=='FOUND' and [r['profile']['geometry'][f] for f in ('h','b','s','t')]==g)
        actual=[(v['L'],v['S']) for v in calculate_uzk(r)['rows']]
        check(actual==uzk); golden[nd]={'geometry':g,'UZK':actual}
    check(s.lookup('35Б1',workbook=real)['status']=='FOUND_PREFERRED')
    check(s.lookup('20Ш1',workbook=real)['status']=='AMBIGUOUS')
    check(s.lookup('35Б1','НЕСУЩЕСТВУЮЩИЙ НД',workbook=real)['status']=='NOT_FOUND')
    check(next(p for p in state['preferences'] if p['key']=='35Б1')['basis']=='USER_CONFIRMED_UZK_328x6_348x9')
    for v in ['11.5','11,5','11.999','12.0','12,0','12.001']:
        check(frequency_angle(v)=={'frequency_mhz':5 if Decimal(v.replace(',','.'))<12 else 2.5,'angle_degrees':70 if Decimal(v.replace(',','.'))<12 else 65})
    for row in cat.rows:
        check((row['uzk_rule']==RULE)==(row['family']=='Двутавр' and row['status']!='КОНФЛИКТ' and all(m.get('uzk_confirmed') for m in row['provenance'])))
    save('counts_independent.json',counts);save('golden_independent.json',golden)
    return {'counts':counts,'golden':golden,'20Ш1':'AMBIGUOUS','decimal_boundary':'PASS'}
case('Post-fix counts and standard-aware golden / UZK scope',counts_golden)

def workbook_check():
    wb=load_workbook(real,read_only=True,data_only=False)
    try:
        check(wb.sheetnames==['Профили','Источники','Правила УЗК','Предпочтения','Конфликты'])
        sheets={ws.title:list(ws.values) for ws in wb.worksheets}
        check(len(sheets['Профили'])-1==2849 and len(sheets['Конфликты'])-1==15)
        check(sum(c.data_type=='f' for ws in wb.worksheets for row in ws for c in row)==0)
        for values,row in zip(sheets['Профили'][1:],cat.rows):
            check(values[1]==row['key'] and values[3]==row['standard'])
            for i,f in zip([4,5,6,7,8,9,10,15],['h','b','b2','s','t','D','r','R']):
                check(values[i]==row['geometry'].get(f))
                if values[i] is not None: check(type(values[i]) in (int,float))
            check(all(m['record_id'] in values[17] for m in row['provenance']))
            if row['disagreement_class']=='NON_BLOCKING': check('unresolved' in values[19] and all(f in values[19] for f in row['unresolved_fields']))
        classes=Counter()
        for values in sheets['Конфликты'][1:]:
            d=next(d for d in disputes if (d['key'],d['standard'])==(values[2],values[3]))
            check(values[12]==d['class']);classes[values[12]]+=1
            raws=json.loads(values[5])+json.loads(values[7])
            check({r['record_id'] for r in raws}=={m['record_id'] for m in d['raw']})
            for r in raws:
                origin=next(m for m in d['raw'] if m['record_id']==r['record_id'])
                check(r['geometry']==origin['geometry'] and r['quality']==origin['quality'])
        return {'sheets':wb.sheetnames,'rows':2849,'conflicts':dict(classes),'formulas':0,'numeric_geometry':True,'raw_conflict_provenance':'PASS'}
    finally: wb.close()
case('Exact real workbook structural assessment',workbook_check)

def fixture(label):
    folder=OUT/'fixtures'/label/'DEV/Справочники';folder.mkdir(parents=True)
    for n in names: shutil.copyfile(real.parent/n,folder/n)
    return folder/real.name
def generation(t):return {n:sha(t.parent/n) if (t.parent/n).exists() else None for n in names}
def backups(t):return sorted((t.parent/'backups').glob('*'))
def verified_backup(t,old):
    bs=backups(t);check(len(bs)==1)
    check({n:sha(bs[0]/n) for n in names}==old)
    s.validated_generation({n:(bs[0]/n).read_bytes() for n in names},t.name)
    return str(bs[0])
template=real.read_bytes()
def copied_render(payload,destination,previews=None):
    Path(destination).write_bytes(template)
    return {'network_attempts':0,'independent_fixture_renderer':'copies exact reviewed workbook; real verify_workbook still executes'}
def normal_operations():
    t=fixture('normal_application')
    old=generation(t); result=s.build(t)
    verified_backup(t,old);check(result['network_attempts']==0 and result['reopen']['reopened'])
    s.set_preference('20Ш1','ГОСТ 26020-83','INDEPENDENT_ISOLATED_PREFERENCE',workbook=t)
    check(s.lookup('20Ш1',workbook=t)['status']=='FOUND_PREFERRED')
    s.add_manual('Независимый ручной профиль','Двутавр','РУЧНОЙ НД',{'h':200,'b':100,'s':5,'t':8},'ISOLATED_ONLY',workbook=t)
    manual=s.lookup('Независимый ручной профиль','РУЧНОЙ НД',workbook=t)
    check(manual['status']=='FOUND' and manual['profile']['status']=='РУЧНОЙ' and manual['profile']['uzk_rule']==NOT_CONFIGURED)
    c=next(c for c in s.load_catalog(t).conflicts if c['classification']=='BLOCKING')
    s.resolve_conflict(c['id'],c['members'][0]['record_id'],'INDEPENDENT_EXPLICIT_ISOLATED_RESOLUTION',workbook=t)
    check(s.lookup(c['key'],c['standard'],workbook=t)['status']=='FOUND')
    old_state=deepcopy(s.load_catalog(t).user_state)
    s.build(t); new_state=s.load_catalog(t).user_state
    for field in ('preferences','manual_rows','resolutions'): check(old_state[field]==new_state[field])
    check(new_state['revision']==old_state['revision']+1)
    check(next(p for p in new_state['preferences'] if p['key']=='35Б1')==next(p for p in state['preferences'] if p['key']=='35Б1'))
    return {'real_renderer_builds':5,'normal_preference_manual_resolution_rebuild':'PASS','user_state_preserved':True,'manual_UZK':NOT_CONFIGURED}
case('M2 normal real-renderer application writes and state preservation',normal_operations)

def render_race(name):
    t=fixture('render_race_'+str(names.index(name)));old=generation(t);external=None
    def render(payload,destination,previews=None):
        nonlocal external
        result=copied_render(payload,destination)
        p=t.parent/name
        if name==t.name:
            source=zipfile.ZipFile(io.BytesIO(p.read_bytes()));buffer=io.BytesIO()
            with source,zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as z:
                for item in source.infolist(): z.writestr(item,source.read(item.filename))
                z.comment=b'VALID_EXTERNAL_WORKBOOK_EDIT'
            p.write_bytes(buffer.getvalue())
            w=load_workbook(p,read_only=True);w.close()
        else:
            obj=read(p)
            if name==s.STATE: obj['preferences'][0]['basis']+=' EXTERNAL'
            elif name==s.CATALOG: obj['records'][0]['notes'].append('VALID_EXTERNAL_EDIT')
            else: obj['built_at']='EXTERNAL_VALID_MANIFEST_EDIT'
            p.write_text(json.dumps(obj,ensure_ascii=False),encoding='utf-8')
        external=sha(p);check(external!=old[name])
        return result
    with patch.object(s,'render_workbook',render): message=rejected(lambda:s.build(t),'changed after preflight')
    check(sha(t.parent/name)==external and all(generation(t)[n]==old[n] for n in names if n!=name) and not backups(t))
    return {'file':name,'STOP':message,'external_bytes_preserved':True,'other_files_unchanged':True,'no_backup_published':True}
for n in names:case('M2 external valid edit during render '+n,lambda n=n:render_race(n))

def probe(t,mode):
    completed=subprocess.run([sys.executable,'-B',str(OUT/'external_probe.py'),str(t.parent),mode],capture_output=True,text=True,encoding='utf-8',check=True)
    return json.loads(completed.stdout)
def boundary():
    t=fixture('commit_boundary');old=generation(t);phases=[]
    original_publish=wg.GenerationGuard.publish;original_rename=wg.ProtectedFile.rename
    def attacks(label):
        previous=generation(t); results=probe(t,'attacks')
        check(len(results)==16 and all(r['blocked'] and r['winerror'] in (5,32) for r in results),str(results))
        check(generation(t)==previous);phases.append({'phase':label,'attempts':results})
    def publish(self,stage,fail_at=None):
        attacks('before_first_retirement_all_old_protected')
        return original_publish(self,stage,fail_at)
    def rename(self,destination):
        original_rename(self,destination)
        if Path(destination)==t.parent/s.COMMIT: attacks('after_manifest_install_all_new_protected')
    with patch.object(s,'render_workbook',copied_render),patch.object(wg.GenerationGuard,'publish',publish),patch.object(wg.ProtectedFile,'rename',rename): s.build(t)
    check(len(phases)==2);verified_backup(t,old);s.load_catalog(t)
    save('commit_boundary_independent.json',phases)
    return {'separate_process':True,'phases':2,'write_delete_rename_replace_blocked':32,'coherent_backup':True,'commit':'PASS'}
case('M2 actual Windows handles protect commit against separate process',boundary)

def vacancy():
    t=fixture('commit_vacancy');old=generation(t);original=wg.ProtectedFile.rename;created=False
    def rename(self,destination):
        nonlocal created
        source=self.path
        original(self,destination)
        if source==t and Path(destination).parent.name=='previous' and not created:
            check(probe(t,'vacancy')['created']);created=True
    with patch.object(s,'render_workbook',copied_render),patch.object(wg.ProtectedFile,'rename',rename):
        message=rejected(lambda:s.build(t),'external destinations preserved')
    check(t.read_bytes()==b'INDEPENDENT_EXTERNAL_CREATION_MUST_SURVIVE')
    check(all(generation(t)[n]==old[n] for n in names if n!=t.name))
    backup=verified_backup(t,old)
    rejected(lambda:s.load_catalog(t),'changed outside')
    return {'STOP':message,'external_creation_survives':True,'other_files_unchanged':True,'verified_previous_generation':backup,'incoherent_active_folder_lookup':'BLOCKED'}
case('M2 free-name race: no overwrite, coherent backup, lookup fail closed',vacancy)

def failure(mode):
    t=fixture('failure_'+mode);old=generation(t)
    if mode=='render':
        with patch.object(s,'render_workbook',side_effect=OSError('INDEPENDENT_RENDER_FAILURE')): message=rejected(lambda:s.build(t))
    elif mode=='backup':
        write=Path.write_bytes;count=0
        def backup_write(p,data):
            nonlocal count
            if p.parent.name=='backups':
                count+=1
                if count==2:raise OSError('INDEPENDENT_PARTIAL_BACKUP_FAILURE')
            return write(p,data)
        with patch.object(s,'render_workbook',copied_render),patch.object(Path,'write_bytes',backup_write): message=rejected(lambda:s.build(t))
        check(not backups(t))
    elif mode=='locked':
        with t.open('rb'): message=rejected(lambda:s.build(t),'Cannot protect generation file')
    else:
        with patch.object(s,'render_workbook',copied_render):message=rejected(lambda:s.build(t,fail_at=mode))
    check(generation(t)==old and not (t.parent/'.profiles.lock').exists())
    s.load_catalog(t)
    return {'STOP':message,'previous_four_files_exactly_preserved':True,'valid_lookup_after_failure':True,'partial_backup_not_exposed':mode=='backup'}
for mode in ['render','backup','before_commit','after_first_replace','locked']:
    case('M2 failed rebuild preserves generation: '+mode,lambda mode=mode:failure(mode))

def evidence_regression():
    before=read(FIX/'product_hashes_before.json'); after=read(FIX/'product_hashes_after.json');current=read(FIX/'product_hashes_current.json')
    check(before==after==current==manifest['product_hashes'])
    check(read(FIX/'protected_before.json')==read(FIX/'protected_after.json'))
    tested=read(FIX/'targeted_tests.json')
    check(tested['passed']==34 and tested['failed']==0 and all(t['status']=='PASS' for t in tested['tests']))
    latest=max(Path(p).stat().st_mtime_ns for p in evidence)
    check((FIX/'REVIEW_MANIFEST.json').stat().st_mtime_ns>=latest,'Fix manifest not written last')
    for name in ['main','stage2','stage3','passport_app.config','passport_app.catalog']:importlib.import_module(name)
    checkpoint=read(ROOT/'DEV/DEV_REGISTRY_CHECKPOINT.json')
    check(sha(ROOT/'DEV/ВЗМК_DEV.xlsx')==checkpoint['sha256']=='8e56f68518c70cade9efb09084dd67cd5d72e07e3fdc813b8515030e9cb98e80')
    w=load_workbook(ROOT/'DEV/ВЗМК_DEV.xlsx',read_only=True);sheets=w.sheetnames;w.close()
    return {'tested_before_after_current_equal':True,'manifest_last':True,'implementation_tests':34,'protected':423,'DEV_checkpoint':checkpoint['sha256'],'DEV_sheets':sheets,'imports':'PASS'}
case('Evidence binding + protected Stage 1–3 imports and DEV checkpoint',evidence_regression)
after={p:sha(p) for p in baseline}
save('hashes_after.json',after)
check(after==baseline,'Reviewed/protected artifact changed during review')
check(not any(events.values()),str(events))
save('review_result.json',{'verdict':'FAIL' if any(t['status']=='FAIL' for t in tests) else 'PASS','tests':tests,'passed':sum(t['status']=='PASS' for t in tests),'failed':sum(t['status']=='FAIL' for t in tests),'real_DEV_builds':0,'events':events,'hashes_before_after_equal':True,'scope':'M1/M2 and direct regressions only','User_E2E':'NOT_RUN'})
print(json.dumps({'verdict':'FAIL' if any(t['status']=='FAIL' for t in tests) else 'PASS','passed':sum(t['status']=='PASS' for t in tests),'failed':sum(t['status']=='FAIL' for t in tests)},ensure_ascii=False),flush=True)
