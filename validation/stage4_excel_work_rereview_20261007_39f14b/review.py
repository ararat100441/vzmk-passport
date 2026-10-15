"""Read-only Excel defect rereview. Writes only new isolated review evidence."""
from pathlib import Path
from collections import Counter
from copy import deepcopy
from unittest.mock import patch
from xml.etree import ElementTree as ET
import difflib
import hashlib
import json
import posixpath
import re
import shutil
import sys
import traceback
import zipfile

OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[1]
FIX=ROOT/'validation/stage4_excel_table_fix_20261007_121500'
PREV=ROOT/'validation/stage4_fix_m1_m2_20261007_090000'
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(OUT))
sys.dont_write_bytecode=True;sys.stdout.reconfigure(encoding='utf-8')
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(n,x):(OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def require(x,m='assertion failed'):
    if not x:raise AssertionError(m)
tests=[]
def case(name,fn):
    try:
        detail=fn();tests.append({'name':name,'status':'PASS','detail':detail});print('PASS '+name,flush=True)
    except BaseException:
        tests.append({'name':name,'status':'FAIL','traceback':traceback.format_exc()});print(traceback.format_exc(),flush=True)
    save('results.json',{'tests':tests,'passed':sum(t['status']=='PASS' for t in tests),'failed':sum(t['status']=='FAIL' for t in tests)})
manifest=read(FIX/'REVIEW_MANIFEST.json')
paths={str(ROOT/p):h for p,h in manifest['product_hashes'].items()}
paths.update(manifest['protected_hashes']);paths.update(manifest['DEV_artifacts'])
paths.update({str(FIX/p):h for p,h in manifest['evidence_files'].items()})
paths[str(FIX/'REVIEW_MANIFEST.json')]=sha(FIX/'REVIEW_MANIFEST.json')
before={p:sha(p) for p in paths}
save('preservation_before.json',before)
def integrity():
    require(before==paths,'Current bytes differ from Excel fix manifest')
    require(read(FIX/'product_hashes_before.json')==read(FIX/'product_hashes_after.json')==read(FIX/'product_hashes_current.json')==manifest['product_hashes'])
    require(read(FIX/'protected_before.json')==read(FIX/'protected_after.json')==manifest['protected_hashes'])
    old=read(PREV/'REVIEW_MANIFEST.json')['product_hashes']
    changed=[p for p,h in old.items() if sha(ROOT/p)!=h]
    require(set(changed)=={'profiles_stage4\\storage.py','profiles_stage4\\workbook.mjs'},str(changed))
    storage=(ROOT/'profiles_stage4/storage.py').read_bytes()
    restored=storage.replace("CORE_COLUMNS + ['Радиус R'".encode(),b"CORE_COLUMNS + ['R'")
    require(hashlib.sha256(restored).hexdigest()==old['profiles_stage4\\storage.py'],'Storage change not exactly single caption')
    workbook=(ROOT/'profiles_stage4/workbook.mjs').read_bytes()
    block='''  // Excel table headings are unique without regard to case: r and R collide.
  if (headers.some(h=>typeof h!=='string' || !h.trim() || /[\\x00-\\x1f]/.test(h)) ||
      new Set(headers.map(h=>h.toLowerCase())).size!==headers.length) {
    throw new Error(`Invalid or case-insensitive duplicate table headers: ${name}`);
  }
'''.encode()
    require(workbook.count(block)==1,'Exact added guard not found')
    old_w=workbook.replace(block,b'').replace("'r','Радиус R','Количество записей'".encode(),"'r','R','Количество записей'".encode())
    require(hashlib.sha256(old_w).hexdigest()==old['profiles_stage4\\workbook.mjs'],'Other renderer changes exist')
    diff=''.join(difflib.unified_diff(restored.decode().splitlines(True),storage.decode().splitlines(True),fromfile='previous/storage.py',tofile='current/storage.py'))
    diff+=''.join(difflib.unified_diff(old_w.decode().splitlines(True),workbook.decode().splitlines(True),fromfile='previous/workbook.mjs',tofile='current/workbook.mjs'))
    (OUT/'verified_product.diff').write_text(diff,encoding='utf-8')
    for d in ['pre_fix_xml','post_fix_xml']:
        for p in (FIX/d).glob('*.xml'):ET.fromstring(p.read_bytes())
    logs=[]
    for p in (FIX/'repair_logs').glob('*.xml'):
        tree=ET.fromstring(p.read_bytes());text=' '.join(tree.itertext())
        require('/xl/tables/table1.xml' in text);logs.append({'path':str(p),'text':text})
    require(len(logs)==2)
    native={p.parent.name:read(p) for p in (FIX/'native').glob('*/result.json')}
    require(native['isolated_fixed']['validation']==native['real_DEV_fixed']['validation']=='PASS')
    require(native['pre_fix']['validation']=='EXPECTED_DEFECT_REPRODUCED')
    require(read(FIX/'targeted_tests.json')['passed']==14)
    last=(FIX/'REVIEW_MANIFEST.json').stat().st_mtime_ns
    require(all((FIX/p).stat().st_mtime_ns<=last for p in manifest['evidence_files']),'Manifest was not last by file timestamps')
    return {'preserved_files':len(paths),'protected_files':len(manifest['protected_hashes']),'fix_evidence_files':len(manifest['evidence_files']),
            'changed_product_set':changed,'exact_diff_verified_against_prior_hashes':True,'protection_path_unchanged':True,
            'BEFORE_AFTER_CURRENT_equal':True,'manifest_last_timestamp_and_finalizer_write':True,'saved_recovery_logs':logs}
case('artifact evidence binding and exact two-file product diff',integrity)
require(tests[-1]['status']=='PASS','STOP evidence mismatch')

import profiles_stage4.storage as s
from profiles_stage4.model import calculate_uzk
from native_probe import run as native
from openpyxl import load_workbook
MAIN='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
DOCREL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS={'m':MAIN}
CURRENT=s.DEFAULT_WORKBOOK
BASE=FIX/'baseline_generation/Профили.xlsx'
CAT=s.load_catalog()
def structure(path,old=False):
    issues=[];tables=[]
    wb=load_workbook(path,read_only=True,data_only=False)
    with zipfile.ZipFile(path) as z:
        parts=set(z.namelist())
        for n in parts:
            if n.endswith('.xml') or n.endswith('.rels'):ET.fromstring(z.read(n))
        rels={}
        for n in parts:
            if not n.endswith('.rels'):continue
            source='' if n=='_rels/.rels' else n.replace('/_rels/','/')[:-5]
            base=posixpath.dirname(source)
            rels[source]={}
            for r in ET.fromstring(z.read(n)):
                require(r.get('Id') not in rels[source],'Duplicate relationship ID')
                target=r.get('Target')
                target=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(base,target))
                if r.get('TargetMode')!='External':require(target in parts,'Dangling relationship '+n+' '+target)
                rels[source][r.get('Id')]={'target':target,'type':r.get('Type')}
        root=ET.fromstring(z.read('xl/workbook.xml'))
        used=set()
        for sh in root.find('m:sheets',NS):
            name=sh.get('name');rid=sh.get('{'+DOCREL+'}id')
            sheet=rels['xl/workbook.xml'][rid]['target']
            xml=ET.fromstring(z.read(sheet))
            tableparts=xml.find('m:tableParts',NS)
            require(tableparts is not None and int(tableparts.get('count'))==len(tableparts)==1)
            for tp in tableparts:
                r=rels[sheet][tp.get('{'+DOCREL+'}id')]
                require(r['type'].endswith('/table'));part=r['target'];used.add(part)
                node=ET.fromstring(z.read(part));columns=node.find('m:tableColumns',NS)
                headers=[c.get('name') for c in columns];ids=[int(c.get('id')) for c in columns]
                require(int(columns.get('count'))==len(columns))
                require(ids==list(range(1,len(columns)+1)))
                require(all(isinstance(h,str) and h.strip() for h in headers))
                if len({h.casefold() for h in headers})!=len(headers):issues.append(part+':case-insensitive duplicate')
                require(node.find('m:autoFilter',NS).get('ref')==node.get('ref'))
                require(list(next(wb[name].iter_rows(min_row=1,max_row=1,values_only=True)))==headers)
                require(wb[name].tables if not wb.read_only else True)
                tables.append({'sheet':name,'part':part,'ref':node.get('ref'),'headers':headers,'IDs':ids,'autoFilter':node.find('m:autoFilter',NS).get('ref')})
        actual={n for n in parts if re.fullmatch(r'xl/tables/table\d+\.xml',n)}
        require(used==actual and len(used)==5,'Orphan table part')
        require(tables[0]['ref']=='A1:T2850' and len(tables[0]['headers'])==20)
    wb.close()
    require(issues==(['xl/tables/table1.xml:case-insensitive duplicate'] if old else []),str(issues))
    return {'tables':tables,'issues':issues,'XML_parses':True,'all_internal_relationship_targets_exist':True,'tableParts_valid':True}
case('current OOXML all five tables headers refs and relationships',lambda:structure(CURRENT))
case('pre-fix OOXML sole detected structural defect r/R',lambda:structure(BASE,True))

def content():
    old=load_workbook(BASE,read_only=False,data_only=False)
    new=load_workbook(CURRENT,read_only=False,data_only=False)
    changes=[];radius=Counter();formats=[]
    try:
        require(new.sheetnames==list(s.HEADERS))
        require(new['Профили'].max_row==2850 and new['Профили'].max_column==20)
        require(not any(c.data_type=='f' for ws in new for row in ws for c in row))
        for name in new.sheetnames:
            a,b=old[name],new[name]
            require((a.max_row,a.max_column)==(b.max_row,b.max_column))
            require(a.freeze_panes==b.freeze_panes and a.freeze_panes)
            require(a.sheet_view.showGridLines==b.sheet_view.showGridLines)
            require(len(a.tables)==len(b.tables)==1)
            for row in b:
                for cell in row:
                    prev=a[cell.coordinate]
                    if cell.value!=prev.value:changes.append({'sheet':name,'cell':cell.coordinate,'before':prev.value,'after':cell.value})
                    require(cell._style==prev._style,'Cell style changed '+name+'!'+cell.coordinate)
            require({k:v.width for k,v in a.column_dimensions.items()}=={k:v.width for k,v in b.column_dimensions.items()})
            formats.append({'sheet':name,'freeze_panes':b.freeze_panes,'cell_styles_unchanged':True,'column_widths_unchanged':True})
        require(changes==[{'sheet':'Профили','cell':'P1','before':'R','after':'Радиус R'}],str(changes[:10]))
        for cells,record in zip(new['Профили'].iter_rows(min_row=2),CAT.rows):
            require((cells[1].value,cells[3].value)==(record['key'],record['standard']))
            for field,col in [('h',4),('b',5),('b2',6),('s',7),('t',8),('D',9),('r',10),('R',15)]:
                v=cells[col].value
                require(v is None or type(v) in [int,float])
                require(v==record['geometry'].get(field),'Geometry field mapping changed')
                if field in ['r','R'] and v is not None:radius[field]+=1
        previous=read(FIX/'baseline_generation/profiles.catalog.json')
        require(previous['records']==CAT.raw_records and previous['sources']==CAT.sources)
        require(all(CAT.user_state[f]==previous['user_state'][f] for f in ['manual_rows','preferences','resolutions']))
        return {'sheets':new.sheetnames,'profile_rows':2849,'columns':20,'formulas':0,'numeric_geometry':True,
                'all_geometry_cells_mapped_to_internal_keys':True,'r_R_nonempty_counts':dict(radius),'only_value_change':changes,'formats':formats}
    finally:old.close();new.close()
case('all workbook values styles panes numeric geometry and r/R mapping',content)

# First native open MUST use an exact byte copy of current artifact.
copy=OUT/'isolated_current.xlsx';shutil.copyfile(CURRENT,copy)
require(sha(copy)==sha(CURRENT)==manifest['summary']['workbook_sha256'])
case('native Excel exact current isolated copy normal read-only Open',lambda:native(copy,'isolated_current'))
require(tests[-1]['status']=='PASS','STOP: current isolated native Excel failed')
case('native Excel exact REAL DEV normal read-only Open',lambda:native(CURRENT,'real_DEV'))

def header_only():
    with zipfile.ZipFile(BASE) as z:parts={n:z.read(n) for n in z.namelist()}
    original=parts.copy()
    require(parts['xl/tables/table1.xml']==(FIX/'pre_fix_xml/table1.xml').read_bytes())
    t=parts['xl/tables/table1.xml'];require(t.count(b'id="16" name="R"')==1)
    parts['xl/tables/table1.xml']=t.replace(b'id="16" name="R"','id="16" name="Радиус R"'.encode())
    xml=parts['xl/worksheets/sheet1.xml']
    match=re.search(rb'<x:c\b[^>]*\br="P1"[^>]*>.*?</x:c>',xml)
    require(match is not None,'P1 lexical cell not found')
    require(match.group().count(b'<x:v>R</x:v>')==1)
    replacement=match.group().replace(b'<x:v>R</x:v>','<x:v>Радиус R</x:v>'.encode())
    parts['xl/worksheets/sheet1.xml']=xml[:match.start()]+replacement+xml[match.end():]
    changed=[n for n in parts if parts[n]!=original[n]]
    require(set(changed)=={'xl/tables/table1.xml','xl/worksheets/sheet1.xml'})
    path=OUT/'isolated_header_only.xlsx'
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in parts.items():z.writestr(n,data)
    negative=OUT/'isolated_pre_fix.xlsx';shutil.copyfile(BASE,negative)
    a=native(negative,'pre_fix_negative',True)
    b=native(path,'header_only_positive')
    proof={'only_changed_parts':changed,'table_change':'id16 name R -> Радиус R','sheet_change':'P1 header only',
           'all_other_part_bytes_identical':True,'pre_fix_native_opened_normally':a['opened_normally'],
           'header_only_native_opened_normally':b['opened_normally'],'no_repairs_in_positive':not b['repair_prompt'] and not b['recovered_records']}
    save('root_cause_AB.json',proof);return proof
case('native root-cause A/B exact lexical header-only change',header_only)

def headers():
    payload=s.workbook_payload(CAT)
    small={'sheets':{name:rows[:2] for name,rows in payload['sheets'].items()}}
    small['sheets']['Правила УЗК']=payload['sheets']['Правила УЗК']
    results=[]
    for label,value in [('exact_duplicate','r'),('case_duplicate','R'),('empty',''),('whitespace','   '),('invalid_control','bad\x01header'),('invalid_type',None),('valid','Радиус R')]:
        data=deepcopy(small);data['sheets']['Профили'][0][15]=value
        p=OUT/('headers_'+label+'.json');s.dump(p,data)
        dest=OUT/('headers_'+label+'.xlsx')
        try:result=s.render_workbook(p,dest)
        except ValueError as exc:
            require(label!='valid' and 'Invalid or case-insensitive duplicate table headers' in str(exc),str(exc))
            require(not dest.exists());results.append({'case':label,'rejected':True})
        else:
            require(label=='valid' and result['network_attempts']==0)
            require(s.verify_workbook(dest,data)['rows']==1)
            results.append({'case':label,'export_succeeded':True,'numeric_geometry_verified':True})
    return results
case('actual renderer exact case empty invalid and valid header controls',headers)

def regression():
    statuses=Counter()
    for c in CAT.conflicts:
        r=CAT.lookup(c['key'],c['standard']);statuses[c['classification']]+=1
        require(r['status']==('CONFLICT' if c['classification']=='BLOCKING' else 'FOUND'))
        if c['classification']=='NON_BLOCKING':require(r['profile']['geometry'] and r['profile']['unresolved_fields'])
    require(statuses=={'BLOCKING':8,'NON_BLOCKING':7})
    a=CAT.lookup('L125×80×7','ГОСТ 8510-86');b=CAT.lookup('L100×8','ГОСТ 8509-93')
    require(a['status']=='CONFLICT' and b['status']=='FOUND' and b['profile']['status']=='СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ')
    results=[]
    for nd,g,u in [('ГОСТ Р 57837-2017',[346,174,6,9],[[328,6],[348,9]]),('ГОСТ 26020-83',[346,155,6.2,8.5],[[329,6.2],[310,8.5]])]:
        r=CAT.lookup('35Б1',nd);require(r['status']=='FOUND')
        require([r['profile']['geometry'][k] for k in ['h','b','s','t']]==g)
        require([[r['L'],r['S']] for r in calculate_uzk(r)['rows']]==u)
        results.append({'standard':nd,'status':r['status'],'geometry':g,'UZK':u})
    require(CAT.lookup('35Б1')['status']=='FOUND_PREFERRED' and CAT.lookup('20Ш1')['status']=='AMBIGUOUS')
    return {'M1':dict(statuses),'L125x80x7':'CONFLICT','L100x8':'FOUND with warning','golden':results,'plain35':'FOUND_PREFERRED','plain20':'AMBIGUOUS'}
case('bounded M1 examples and golden regression',regression)

def m2():
    folder=OUT/'fixtures/m2_smoke/DEV/Справочники';folder.mkdir(parents=True)
    names=[CURRENT.name,s.STATE,s.CATALOG,s.COMMIT]
    for n in names:shutil.copyfile(CURRENT.parent/n,folder/n)
    old={n:sha(folder/n) for n in names};edited={}
    def render(payload,destination,previews=None):
        shutil.copyfile(CURRENT,destination)
        state=read(folder/s.STATE);state['independent_external_marker']=True;s.dump(folder/s.STATE,state)
        edited['hash']=sha(folder/s.STATE)
        return {'network_attempts':0}
    with patch.object(s,'render_workbook',render):
        try:s.build(folder/CURRENT.name)
        except ValueError as exc:message=str(exc)
        else:raise AssertionError('M2 did not STOP')
    require(edited and sha(folder/s.STATE)==edited['hash']!=old[s.STATE])
    require(all(sha(folder/n)==old[n] for n in names if n!=s.STATE))
    require(not (folder/'backups').exists() and not list(folder.glob('.profiles-stage-*')) and not (folder/'.profiles.lock').exists())
    return {'STOP':message,'external_bytes_preserved':True,'other_files_unchanged':True,'new_or_partial_generation_absent':True,'protection_code_hash_unchanged':True}
case('one isolated M2 external-mutation commit-guard smoke',m2)

def preserved():
    after={p:sha(p) for p in before};save('preservation_after.json',after);require(before==after)
    return {'BEFORE_AFTER_CURRENT_equal':True,'files':len(after),'protected_files':len(manifest['protected_hashes']),
            'all_prior_reviews_and_fix_evidence_unchanged':True,'Stage1_3_sources_DEV_checkpoint_unchanged':True,
            'real_DEV_rebuilds':0,'production_access':0,'network_requests':0}
case('final protected state and exact artifact preservation',preserved)
summary={'verdict':'FAIL' if any(t['status']!='PASS' for t in tests) else 'PASS','passed':sum(t['status']=='PASS' for t in tests),
         'failed':sum(t['status']=='FAIL' for t in tests),'tests':tests,'DEV_artifacts':manifest['DEV_artifacts'],
         'product_hashes':manifest['product_hashes'],'not_run':['User E2E','Stage 5','real DEV rebuild','production access','live sources/network','full Stage 4 review']}
summary.update(findings={'BLOCKER':0,'MAJOR':0 if summary['verdict']=='PASS' else 1,'MINOR':0},required_fix='NONE' if summary['verdict']=='PASS' else 'See failed targeted checks',ready_to_resume_user_E2E=summary['verdict']=='PASS')
save('WORK_REREVIEW.json',summary)
evidence={str(p.relative_to(OUT)):sha(p) for p in OUT.rglob('*') if p.is_file() and p.name!='REVIEW_MANIFEST.json'}
save('REVIEW_MANIFEST.json',{'verdict':summary['verdict'],'passed':summary['passed'],'failed':summary['failed'],'product_hashes':summary['product_hashes'],
     'DEV_artifacts':summary['DEV_artifacts'],'preservation_BEFORE_AFTER_CURRENT':before=={p:sha(p) for p in before},'evidence_files':evidence})
print(json.dumps({k:summary[k] for k in ['verdict','passed','failed','findings','ready_to_resume_user_E2E']},ensure_ascii=False),flush=True)
