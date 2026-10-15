import sys
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from profiles_stage4.offline import load_sources
from profiles_stage4.model import Catalog
from profiles_stage4.storage import workbook_payload, dump, render_workbook
from openpyxl import load_workbook
folder=Path(__file__).resolve().parent/'development'
payload=workbook_payload(Catalog(*load_sources()))
dump(folder/'payload.json',payload)
path=folder/'DEV/Справочники/Профили.xlsx'
print(render_workbook(folder/'payload.json',path))
wb=load_workbook(path,read_only=True)
for ws in wb:
    print(ws.title,ws.max_row,ws.max_column,'expected',len(payload['sheets'][ws.title]))
    print(list(ws.iter_rows(min_row=max(1,ws.max_row-2),values_only=True)))
wb.close()
