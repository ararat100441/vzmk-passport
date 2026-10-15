from common import *
from openpyxl import load_workbook
s=config();d,p=task('IMPORT_0910','import',source=s['source'],parent=s['parent']);work=Path(d['workspace'])
target=work/'DEV/Stage6_Fixtures/assignments.xlsx';old=sha(target);b=load_workbook(target);sheet=b.active
headers={str(c.value).strip():c.column for c in sheet[1]};assert '№ АКЗ' in headers and 'Марка' in headers,headers
assignment={'Бп2-23':0,'Бп2-5':1,'Бп2-18':2,'Бп2-19':3};found={}
for row in range(2,sheet.max_row+1):
 mark=str(sheet.cell(row,headers['Марка']).value).strip()
 sheet.cell(row,headers['№ АКЗ']).value=assignment.get(mark)
 if mark in assignment:found[mark]=row
assert set(found)==set(assignment)
b.save(target);b.close()
save(OUT/'WORKSPACE.json',{'workspace':str(work),'import_result':str(p),'fixture':str(target),'fixture_before':old,'fixture_after':sha(target),'assignments':assignment,'rows':found,'scope':'OWN_ARTIFICIAL_TEST_ASSIGNMENTS_ONLY'})
d,p=task('SETTINGS_INIT','akz-settings-init',workspace=work)
save(OUT/'DEFAULT_SETTINGS.json',read(work/'preferences/akz_auto/settings.json'))
save(OUT/'BOOTSTRAP_FINAL.json',{'status':'PASS','workspace':str(work),'EXE_SHA':EXPECTED,'settings':str(work/'preferences/akz_auto/settings.json'),'settings_init_result':str(p)})
