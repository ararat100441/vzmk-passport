"""Independent read-only XLSX oracle; no product validator used."""
from common import *
from openpyxl import load_workbook
from copy import copy
from decimal import Decimal as D
import re
work=Path(read(OUT/'WORKSPACE.json')['workspace']);final=read(OUT/'FROZEN_AUTO_FINAL.json');m=read(final['full_manifest'])
assert m['official_issue'] is False and m['seal']==digest({k:v for k,v in m.items() if k!='seal'})
unique={r['path']:r for r in m['records']};assert len(unique)==12 and len(m['records'])==13
assert len([p for p in unique if Path(p).suffix=='.docx'])==3 and len([p for p in unique if Path(p).suffix=='.xlsx'])==9
assert m['quality_generated_last'] and m['records'][-1]['type']=='QUALITY'
assert (m['visible_rows'],m['quantity'],m['mass_kg'],m['area_m2'],m['akz_variants'])==(104,'288','15822.14','1913.494',[0,1,2,3])
assert all(sha(p)==r['sha256'] and r['official_confirmed'] is False for p,r in unique.items())
assert all(sha(p)==h for p,h in m['inputs'].items())
# Read current passport directly from the copied registry; do not call product fields().
from passport_app.registry import read_registry,select_project
registry=read_registry(work/'DEV/ВЗМК_DEV.xlsx');select_project(registry,'Гараж Казахстан')
entries=[e for e in registry.projects['Гараж Казахстан'] if e['number']=='9/2026'];assert len(entries)==1
reg=load_workbook(work/'DEV/ВЗМК_DEV.xlsx',read_only=True,data_only=True)
vals=[reg[registry.sheet].cell(entries[0]['row'],c).value for c in range(1,17)];reg.close()
current={'number':vals[1],'contract':vals[2],'customer':vals[3],'object':vals[4],'cipher':vals[5]}
# Read selected VOM and explicit assignments independently from copies.
checked=read(read(OUT/'CHECKPOINT_CHECK_ALL.json')['result']);vom=Path(checked['data']['run'])/'VISIBLE_VOM.xlsx'
b=load_workbook(vom,read_only=True,data_only=True);rows={r[0]:r for r in b.active.iter_rows(min_row=3,values_only=True) if r[0] and r[0]!='Итого'};b.close();assert len(rows)==104
assignments=read(OUT/'WORKSPACE.json')['assignments'];expected={str(n):mark for mark,n in assignments.items()}
assert len(expected)==4
limits={'0':[(80,100)],'1':[(150,180)],'2':[(100,120),(80,100)],'3':[(100,120),(80,100),(60,80)]}
columns={'0':['E'],'1':['E'],'2':['E','I'],'3':['E','H','J']};lasts={'0':20,'1':65,'2':51,'3':98}
audits=[];climates={};style_diffs=[]
sets=[('FIRST',read(OUT/'PAIRS_FIRST.json')),('REPEAT',read(OUT/'PAIRS_REPEAT.json'))]
full_akz=Path(m['akz_auto_pairs']['path']) if 'akz_auto_pairs' in m else None
if full_akz is None:
 # Full manifest AKZ paths are copies of the same newly generated pairs.
 full_groups={}
 for n in range(4):
  act=[Path(p) for p,r in unique.items() if r['type']=='AKZ' and ('№'+str(n)+' ' in Path(p).name or 'АКЗ'+str(n) in Path(p).parts)]
  if len(act)==1 and (act[0].parent/'PAIR.json').exists():full_groups[str(n)]={'path':str(act[0].parent)}
 if full_groups:sets.append(('FULL',{'groups':full_groups}))
else:sets.append(('FULL',read(full_akz/'COMPLETE.json')))
for label,pairs in sets:
 for n,g in pairs['groups'].items():
  folder=Path(g['path']);pair=read(folder/'PAIR.json');assert pair['official_issue'] is False and pair['mode']=='TEST_SIMULATION'
  paths={kind:folder/name for name in pair['files'] for kind in ['application' if name.startswith('Приложение') else 'act']}
  assert all(sha(folder/name)==h for name,h in pair['files'].items())
  for kind,p in paths.items():
   source=Path(pair['templates'][kind]);out=load_workbook(p,data_only=False);cache=load_workbook(p,data_only=True);original=load_workbook(source,data_only=False)
   assert out.sheetnames==original.sheetnames
   for osheet in original:
    sheet=out[osheet.title]
    assert sorted(str(r) for r in osheet.merged_cells.ranges)==sorted(str(r) for r in sheet.merged_cells.ranges)
    assert str(osheet.print_area)==str(sheet.print_area)
    assert {k:v.width for k,v in osheet.column_dimensions.items()}=={k:v.width for k,v in sheet.column_dimensions.items()}
    for row in sheet:
     for c in row:assert c.data_type!='e' and 'RANDBETWEEN' not in str(c.value)
    for row in cache[sheet.title]:
     for c in row:assert c.data_type!='e'
    coords=set(osheet._cells)|set(sheet._cells)
    for address in coords:
     a=osheet.cell(*address);z=sheet.cell(*address)
     for component in ('font','fill','border','alignment','protection','number_format'):
      if copy(getattr(a,component))==copy(getattr(z,component)):continue
      allowed=(kind=='act' and component=='alignment' and a.coordinate in ('E7','B43')) or (kind=='act' and component=='number_format' and n=='2' and a.coordinate=='I27') or (kind=='application' and component=='number_format' and a.column in (3,4,5,6,7) and a.row>=(4 if n=='0' else 5))
      helpers={'2':{'L31','M31'},'3':{'N31','O31'}}
      if kind=='act' and component=='alignment' and a.coordinate in helpers.get(n,set()):
       assert z.value is None and z.column>11
       pattern=r'(?<![A-Z0-9_])\$?'+z.column_letter+r'\$?'+str(z.row)+r'(?!\d)'
       assert not any(re.search(pattern,str(c.value)) for rr in sheet for c in rr if c.data_type=='f');allowed=True
      style_diffs.append({'set':label,'variant':n,'kind':kind,'sheet':sheet.title,'cell':a.coordinate,'component':component,'allowed':allowed})
      assert allowed,style_diffs[-1]
   if kind=='act':
    sheet=out['4'];cached=cache['4'];assert sheet['E4'].value==pair['act_number']+' от '
    for cell,k in [('E5','contract'),('E6','customer'),('E7','object')]:assert sheet[cell].value==current[k]
    assert sheet['B44'].value==current['cipher'] or n=='0' and sheet['B44'].value.startswith(current['cipher']+',')
    assert cached['H4'].value.date().isoformat()==pair['date']
    assert 'ИЗМЕРЕНИЙ НЕТ' in sheet['B43'].value and 'ВЫДАЧА ЗАПРЕЩЕНА' in sheet['B43'].value
    for col in columns[n]:
     for r in (17,18,21,25):assert sheet[col+str(r)].value==original['4'][col+str(r)].value
    if n=='2':assert sheet['I25'].value=='∑1+2=180' and sheet['I27'].number_format=='0' and out['Лист1'].sheet_state=='veryHidden'
    for index,col in enumerate(columns[n]):
     dry=cached[col+'28'].value;lo,hi=limits[n][index];assert lo<=dry<=hi and int(dry)==dry
     day=cached[col+'22'].value.date().isoformat();record=pair['proof']['days'][day];v=record['values']
     climates.setdefault(day,[]).append(v)
     for r,k in [(31,'surface_c'),(32,'air_c'),(33,'relative_humidity'),(34,'dew_point_c')]:assert cached[col+str(r)].value==v[k]
     assert 15<=v['air_c']<=25 and 18<=v['surface_c']<=28 and 40<=v['relative_humidity']<=65
     import math
     gamma=math.log(v['relative_humidity']/100)+17.62*v['air_c']/(243.12+v['air_c']);dew=243.12*gamma/(17.62-gamma)
     assert abs(dew-v['dew_point_raw_c'])<1e-12 and round(dew,1)==v['dew_point_c'] and v['surface_c']>=max(dew,v['dew_point_c'])+3-1e-10
     if original['4'][col+'27'].data_type=='f':assert sheet[col+'27'].value==original['4'][col+'27'].value
    for c in ('E46','E47'):assert sheet[c].value=='ТЕСТ: Н/Д'
    for c in ('F50','I50'):assert 'нет подписи' in str(sheet[c].value)
    text='\n'.join(str(c.value) for sh in out for rr in sh for c in rr if c.value is not None)
    assert not any(x in text for x in ('1/Б2','1/2025','РУССОКО','Велесстрой','2-ЭГ-2019','ВС.24.ПО','13-25','314-25','Журавлёв','Шляхтуров'))
   else:
    sheet=out.active;cached=cache.active;first=4 if n=='0' else 5;last=lasts[n];total=last+1
    selected=[tuple(sheet.cell(r,c).value for c in range(1,8)) for r in range(first,last+1) if sheet.cell(r,1).value]
    assert len(selected)==1 and selected[0][0]==expected[n],(n,selected)
    wanted=rows[expected[n]]
    for col in range(2,8):
     av=selected[0][col-1];ev=wanted[col-1]
     assert av==ev or isinstance(av,(int,float)) and isinstance(ev,(int,float)) and abs(D(str(av))-D(str(ev)))<D('1e-8'),(col,av,ev)
    for col,key in [('C','quantity'),('E','mass_total'),('G','area_total')]:
     assert sheet[f'{col}{total}'].value==original.active[f'{col}{total}'].value
     assert abs(D(str(cached[f'{col}{total}'].value))-D(str(pair['totals'][key])))<D('1e-8')
    if n=='0':assert sheet['C21'].value=='=SUM(C3:C20)'
   audits.append({'set':label,'variant':n,'kind':kind,'path':str(p),'sha256':sha(p),'template':str(source),'template_sha256':sha(source),'sheets':out.sheetnames,'status':'PASS'})
   out.close();cache.close();original.close()
assert all(all(v==values[0] for v in values) for values in climates.values())
save(OUT/'INDEPENDENT_XLSX_AUDIT.json',{'status':'PASS','files':audits,'styles':style_diffs,'climates':climates,'current_requisites':current,'explicit_variant_marks':expected,'blank_assignments_excluded':True,'manifest':final['full_manifest'],'physical':12,'logical':13,'quality_last':True,'separate_stage7_receipts':'not produced by Stage9 path; sealed Stage9 manifest and PAIR receipts checked'})
print('Independent XLSX audit PASS',len(audits),'files, variants0-3,12/13 manifest',flush=True)
