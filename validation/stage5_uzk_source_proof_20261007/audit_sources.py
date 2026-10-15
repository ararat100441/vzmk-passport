import sys,json,re,hashlib,zipfile
from pathlib import Path
from datetime import datetime,timezone
import xml.etree.ElementTree as ET
import openpyxl
import pypdfium2 as pdfium
sys.dont_write_bytecode=True
OUT=Path(__file__).resolve().parent
APP=OUT.parent.parent
ROOT=APP.parent/'Гараж казахстан'
sys.path.insert(0,str(APP))
from profiles_stage4.storage import load_catalog
from profiles_stage4.model import normalize_profile,calculate_uzk

def save(n,x): (OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
def text(e):
    chunks=[]
    for n in e.iter():
        if n.tag==W+'t':chunks.append(n.text or '')
        elif n.tag in (W+'br',W+'cr'):chunks.append('\n')
        elif n.tag==W+'p':
            if chunks:chunks.append('\n')
    return ''.join(chunks).strip()

catalog_path=ROOT/'Справочник марок.xlsx'
docs=sorted(p for p in ROOT.rglob('*.docx') if 'УЗК' in p.name and not p.name.startswith('~$'))
protected=[catalog_path,ROOT/'WAA0007144-AC2-КМД.pdf',ROOT/'Детали.pdf',APP/'DEV'/'УЗК_DEV.xlsm',APP/'DEV'/'DEV_UZK_CHECKPOINT.json']+docs
protected += [p for p in (APP/'DEV'/'Справочники').iterdir() if p.is_file()]
protected += list((APP/'profiles_stage4').glob('*.py'))
before={str(p):sha(p) for p in protected}
assert before[str(APP/'DEV'/'Справочники'/'Профили.xlsx')]=='39f14b040d3a97365f9e44f5fdc01c87e1f2ce068d5d4f26458a91602759ca91','STOP accepted workbook drift'
prior=read(APP/'validation'/'stage5_uzk_audit_20261007_resume'/'preservation_after.json')
for p in [APP/'DEV'/'УЗК_DEV.xlsm',APP/'DEV'/'DEV_UZK_CHECKPOINT.json']:
    assert before[str(p)]==prior[str(p)],'STOP DEV checkpoint drift'
save('preservation_before.json',before)
wb=openpyxl.load_workbook(catalog_path,read_only=True,data_only=False)
catalog_rows=[]
for s in wb:
    for i,row in enumerate(s.iter_rows(values_only=True),1):
        catalog_rows.append({'sheet':s.title,'row':i,'values':row})
wb.close()
save('working_mark_catalog.json',{'path':str(catalog_path),'sha256':before[str(catalog_path)],'rows':catalog_rows})
reports=[]
for p in docs:
    with zipfile.ZipFile(p) as z:
        root=ET.fromstring(z.read('word/document.xml'))
        rels={n.attrib['Id']:n.attrib for n in ET.fromstring(z.read('word/_rels/document.xml.rels'))}
        tables=root.findall('.//'+W+'body/'+W+'tbl')
        body=text(root);act=re.search(r'Акт\s*№\s*(.*?)\s+от\s+(\d{2}\.\d{2}\.\d{4})',body)
        data=[]
        for i,tr in enumerate(tables[0].findall(W+'tr'),1):
            cells=[text(c) for c in tr.findall(W+'tc')]
            if len(cells)<8 or not re.search(r'L\s*=',cells[1]):continue
            links=[]
            for h in tr.findall('.//'+W+'hyperlink'):
                rel=rels.get(h.attrib.get(R+'id'),{})
                links.append({'text':text(h),'target':rel.get('Target'),'mode':rel.get('TargetMode')})
            lengths=re.findall(r'L\s*=\s*([\d.,]+)\s*[хxХX×]\s*([\d.,]+)',cells[1])
            data.append({'table_row':i,'mark':cells[0],'geometry_raw':cells[1],'segments':lengths,'frequency_raw':cells[4],'angle_raw':cells[5],'date':cells[3],'hyperlinks':links,'all_cells':cells})
        reports.append({'path':str(p),'sha256':sha(p),'act':act[1] if act else None,'date':act[2] if act else None,'rows':data})
save('project_uzk_documents.json',reports)
all_marks=sorted({r['mark'] for d in reports for r in d['rows']})
pdf=pdfium.PdfDocument(str(ROOT/'WAA0007144-AC2-КМД.pdf'))
kmd=[]
def repair(t):return re.sub(r'[\u0080-\u00ff]+',lambda m:m[0].encode('latin1').decode('cp1251',errors='replace'),t)
for i in range(len(pdf)):
    page=pdf[i];tp=page.get_textpage();t=repair(tp.get_text_range());tp.close();page.close()
    if 'Спецификация деталей' not in t:continue
    for mark in all_marks:
        pat=r'(?m)^'+re.escape(mark)+r'\s+(\d+)\s+(\d+)\s+(\S+)\s+(\d+)\s+'
        match=re.search(pat,t)
        if match:
            kmd.append({'mark':mark,'page':i+1,'detail':match[1],'part_quantity':int(match[2]),'raw_profile':match[3],'length':int(match[4]),'text':t})
pdf.close()
save('project_kmd_mark_evidence.json',kmd)
by_mark={r['mark']:r for r in kmd}
for d in reports:
    for r in d['rows']:
        r['kmd']=by_mark.get(r['mark'])
save('project_uzk_documents.json',reports)
cat=load_catalog(APP/'DEV'/'Справочники'/'Профили.xlsx')
queries=[('[16П',None),('16П','ГОСТ 8240-97'),('Гн.[140X60X4','ГОСТ 8278-83'),('Гн.[]120х3','ГОСТ 30245-2003'),('120×120×3','ГОСТ 30245-2003'),('120×120×3',None)]
results=[]
for q,std in queries:
    r=cat.lookup(q,std)
    results.append({'raw_input':q,'normalized_input':normalize_profile(q),'standard':std,'lookup':r,'calculation':calculate_uzk(r) if r['status'] in ('FOUND','FOUND_PREFERRED') else None})
save('lookup_results.json',results)
after={str(p):sha(p) for p in protected};assert before==after,'STOP protected file drift'
save('preservation_after.json',after)
print(json.dumps({'catalog_sha':before[str(catalog_path)],'catalog_rows':catalog_rows,'docs':[{'act':d['act'],'date':d['date'],'path':d['path'],'rows':[{'mark':r['mark'],'LS':r['geometry_raw'],'frequency':r['frequency_raw'],'angle':r['angle_raw'],'profile':r['kmd']['raw_profile'] if r['kmd'] else 'UNRESOLVED','page':r['kmd']['page'] if r['kmd'] else None,'hyperlinks':r['hyperlinks']} for r in d['rows']]} for d in reports],'protected':len(before)},ensure_ascii=False))
