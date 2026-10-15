"""Independent read-only Stage 4 review; evidence writes stay in this directory."""
import sys, json, hashlib, re, importlib, importlib.util, os
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
from copy import deepcopy
from unittest.mock import patch

sys.stdout.reconfigure(encoding='utf-8')
OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
sys.path.insert(0, str(ROOT))
events = []
def audit(event, args):
    if event.startswith('socket.') and event in ('socket.connect','socket.getaddrinfo','socket.sendto'):
        events.append(['network', event]); raise RuntimeError('Independent review: network denied')
    if event == 'open' and isinstance(args[0], (str,bytes)):
        p = Path(os.fsdecode(args[0])).resolve()
        if p in (ROOT.parent.parent/'ВЗМК.xlsx', ROOT.parent.parent/'Профили.xlsx'):
            events.append(['production',event]); raise RuntimeError('Production denied')
        mode, flags = args[1], args[2]
        writing = (isinstance(mode,str) and any(x in mode for x in 'wax+')) or (isinstance(flags,int) and bool(flags & (os.O_WRONLY|os.O_RDWR|os.O_CREAT|os.O_TRUNC)))
        if writing and not p.is_relative_to(OUT):
            raise RuntimeError('Independent review: write outside fixture denied: '+str(p))
sys.addaudithook(audit)
def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(n,d): (OUT/n).write_text(json.dumps(d,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
E = ROOT/'validation/stage4_implementation_20261006_223330'
manifest = read(E/'STAGE4_REVIEW_MANIFEST.json')
reviewed = {str(ROOT/p):h for p,h in manifest['tested_product_hashes'].items()}
reviewed.update(read(E/'protected_before.json'))
reviewed.update({str(E/p):m['sha256'] for p,m in manifest['evidence_files'].items()})
reviewed[str(E/'STAGE4_REVIEW_MANIFEST.json')] = sha(E/'STAGE4_REVIEW_MANIFEST.json')
for rel in ['stage4_km_kmd_audit_20261006_204347','stage4_ptoflow_audit_20261006_215239']:
    folder=ROOT/'validation'/rel
    reviewed[str(folder/'AUDIT_MANIFEST.json')]=sha(folder/'AUDIT_MANIFEST.json')
    reviewed.update({str(folder/n):m['sha256'] for n,m in read(folder/'AUDIT_MANIFEST.json')['evidence_files'].items()})
folder=ROOT/'DEV/Справочники'
for n,h in [('Профили.xlsx',manifest['workbook_sha256']),('profiles.user.json',manifest['user_state_sha256']),('profiles.catalog.json',manifest['catalogue_sha256']),('profiles.build.json',manifest['commit_sha256'])]:
    reviewed[str(folder/n)] = h
assert all(Path(p).resolve().is_relative_to(ROOT) for p in reviewed)
assert all(sha(p)==h for p,h in reviewed.items())
save('reviewed_hashes_before.json',reviewed)

from profiles_stage4 import model, offline, storage
from openpyxl import load_workbook
import xlrd
data=read(folder/storage.CATALOG)
state=read(folder/storage.STATE)
assert data['user_state']==state
with offline.no_network() as attempts:
    records,sources=offline.load_sources()
    assert records==data['records'] and sources==data['sources']
    cat=storage.load_catalog()
assert not attempts

# Independent identity counting, using raw catalogue keys and independent normalization.
def nd(v):
    v=re.sub('[‐‑–—−]','-',str(v).strip().upper())
    v=re.sub(r'\s*-\s*','-',v)
    return re.sub(r'\s+',' ',v)
def pk(v):
    v=str(v).strip().upper().replace('\u00a0',' ')
    v=re.sub(r'\s+','',v).replace(',','.')
    v=re.sub(r'^I(?=\d)','',v)
    if re.fullmatch(r'L?\d+(?:\.\d+)?(?:[XХ×]\d+(?:\.\d+)?){1,2}',v):
        prefix='L' if v.startswith('L') else ''
        nums=re.split('[XХ×]',v.removeprefix('L'))
        v=prefix+'×'.join(format(Decimal(x).normalize(),'f') for x in nums)
    return v
groups=defaultdict(list)
for r in records: groups[(pk(r['key']),nd(r['standard']))].append(r)
assert len(groups)==len(cat.rows)
assert set(groups)=={(r['key'],r['standard']) for r in cat.rows}
assert len({r['record_id'] for r in records})==len(records)
assert Counter(p['record_id'] for r in cat.rows for p in r['provenance'])==Counter(r['record_id'] for r in records)
reductions=Counter(); merge_sizes=Counter()
for g in groups.values():
    if len(g)>1:
        reductions['cross_source' if len({r['source'] for r in g})>1 else 'same_source']+=len(g)-1
        merge_sizes[len(g)]+=1
counts={'raw':len(records),'raw_by_source':dict(Counter(r['source'] for r in records)),
    'canonical':len(groups),'reductions':dict(reductions),'merge_group_sizes':dict(merge_sizes),
    'family':dict(Counter(r['family'] for r in cat.rows)),
    'standard':dict(Counter(r['standard'] for r in cat.rows)),
    'status':dict(Counter(r['status'] for r in cat.rows)),
    'ambiguous_keys':[k for k in cat.by_key if cat.lookup(k)['status']=='AMBIGUOUS'],
    'uzk_enabled':sum(r['uzk_rule']==model.RULE for r in cat.rows),
    'both_optional_radii_missing':sum(all(r['geometry'].get(k) is None for k in ('R','r')) for r in cat.rows),
    'manual':len(state['manual_rows']), 'provenance_record_count':sum(len(r['provenance']) for r in cat.rows)}
assert sum(counts['family'].values())==sum(counts['status'].values())==counts['canonical']
assert counts['raw']-sum(reductions.values())==counts['canonical']
save('counts.json',counts)

# Verify Source A directly from each XLS, including raw blanks and inherited origins.
source_a=read(ROOT/offline.A_REL/'parsed_source_records.json')
books={name:xlrd.open_workbook(str(ROOT/offline.A_REL/'extracted'/name),formatting_info=True)
       for name in {r['workbook'] for r in source_a}}
origin_methods=Counter(); a_fields=0
for r in source_a:
    sh=books[r['workbook']].sheet_by_name(r['sheet'])
    title=sh.cell_value(0,0)
    standard=re.sub(r'\s+',' ',re.search(r'(?:ГОСТ|СТО\s+АСЧМ)\s+\d+-\d+',title).group())
    assert standard==r['standard']
    for f,origin in r['provenance'].items():
        rr,cc=origin['cell']; ar,ac=origin['origin']
        value=sh.cell_value(rr-1,cc-1); value=None if value=='' else value
        assert value==r['raw'][f]
        value=sh.cell_value(ar-1,ac-1); value=None if value=='' else value
        assert value==r['effective'][f]
        origin_methods[origin['method']]+=1; a_fields+=1
        if origin['method']=='merged_anchor':
            assert any(lo<=rr-1<hi and clo<=cc-1<chi and (ar,ac)==(lo+1,clo+1) for lo,hi,clo,chi in sh.merged_cells)
# Independent count of data cells: numeric thickness in observed table columns/blocks.
block_specs={'уголки неравнополочные.xls':[(18,51,1,4),(18,45,12,15)],
 'уголки равнополочные.xls':[(21,57,1,3),(6,57,10,12)],
 'швеллера.xls':[(6,23,1,5),(29,46,1,5)],
 'СТО АСЧМ 20-93.xls':[(6,21,1,5),(26,55,1,5),(26,55,10,14)]}
a_counts={}; cell_origins=set()
for name,blocks in block_specs.items():
    sh=books[name].sheet_by_index(0); n=0
    for first,last,start,tcol in blocks:
        for rr in range(first,last+1):
            thickness=sh.cell_value(rr-1,tcol-1)
            assert isinstance(thickness,(int,float)) or re.fullmatch(r'\d+(?:[.,]\d+)?\*?',str(thickness))
            n+=1; cell_origins.add((name,rr,start))
    a_counts[name]=n
assert cell_origins=={(r['workbook'],r['row'],r['start_column']) for r in source_a}
for book in books.values(): book.release_resources()

# Independently decode saved HTML Flight arrays without product flight parser.
from lxml import html
b_evidence=read(ROOT/offline.B_REL/'group_source_evidence.json')
pages={p['url']:p for p in read(ROOT/offline.B_REL/'pages.json') if p['kind']=='group'}
by_url=defaultdict(list)
for r in b_evidence: by_url[r['url']].append(r)
b_counts={}; b_fields=0; rounding=0
for url,rows in by_url.items():
    page=pages[url]; path=ROOT/offline.B_REL/page['local_path']
    assert sha(path)==page['sha256']
    payload=''
    for script in html.fromstring(path.read_bytes()).xpath('//script[not(@src)]/text()'):
        for hit in re.finditer(r'self\.__next_f\.push\(',script):
            arr,_=json.JSONDecoder().raw_decode(script[hit.end():])
            if len(arr)>1 and isinstance(arr[1],str): payload+=arr[1]
    raw_arrays=[]
    for hit in re.finditer(r'"rows":',payload):
        try: arr,_=json.JSONDecoder().raw_decode(payload[hit.end():])
        except ValueError: continue
        if isinstance(arr,list) and arr and isinstance(arr[0],dict) and 'size' in arr[0]: raw_arrays.append(arr)
    assert len(raw_arrays)==1 and len(raw_arrays[0])==len(rows)
    assert {r['embedded_row'] for r in rows}==set(range(1,len(rows)+1))
    b_counts[rows[0]['standard']]=len(rows)
    for r in rows:
        raw=raw_arrays[0][r['embedded_row']-1]; assert raw==r['raw']
        product=next(p for p in records if p['source']=='PTOflow' and p['evidence']['url']==url and p['evidence']['embedded_row']==r['embedded_row'])
        for f,v in product['geometry'].items():
            original={'D':'d','b':'B','b2':'b'}.get(f,f) if product['family']=='Уголок неравнополочный' else {'D':'d'}.get(f,f)
            if original in raw and isinstance(raw[original],dict):
                assert Decimal(str(v))==Decimal(str(raw[original]['sort']))
                b_fields+=1
                shown=str(raw[original]['display']).replace(',','.').replace('\u00a0','').strip()
                if shown and Decimal(str(raw[original]['sort']))!=Decimal(shown): rounding+=1
assert not any(any(x in r['standard'].upper() for x in ('8639','8645','ТУ','ВНЕ ГОСТ')) for r in records)
assert all(r['source_id'].startswith('B-') for r in records if r['source']=='PTOflow')
save('source_verification.json',{'A_workbook_counts':a_counts,'A_direct_fields':a_fields,'A_origin_methods':dict(origin_methods),
 'B_group_counts':b_counts,'B_numeric_fields_directly_checked':b_fields,'B_rounded_fields_checked':rounding,
 'exact_reingestion_equals_current_catalog':True,'offline_attempts':attempts})

conflicts=[]
for identity,members in groups.items():
    vals=defaultdict(list)
    for r in members:
        for k,v in {**r['geometry'],**r.get('quality',{})}.items():
            if v is not None: vals[k].append(v)
    diffs=[k for k,v in vals.items() if len(set(Decimal(str(x)) for x in v))>1]
    if not diffs: continue
    actual=next(c for c in cat.conflicts if (c['key'],c['standard'])==identity)
    assert set(diffs)==set(actual['fields'])
    primary=set(model.PRIMARY[members[0]['family']])
    kind='main_geometry' if set(diffs)&primary else 'radius_only' if set(diffs)<=set(('R','r')) else 'area_mass_only'
    result=cat.lookup(*identity)
    conflicts.append({'id':actual['id'],'key':identity[0],'standard':identity[1],'kind':kind,'fields':diffs,
     'values':{f:[{'record':r['record_id'],'source':r['source'],'value':{**r['geometry'],**r.get('quality',{})}.get(f)} for r in members] for f in diffs},
     'lookup':result['status'],'canonical_geometry':result['profile']['geometry']})
assert len(conflicts)==len(cat.conflicts)
save('conflicts_independent.json',conflicts)

golden={}
for key in ('35Б1','20Ш1','25Б1','30Б2'):
    golden[key]={'without_standard':cat.lookup(key)['status'],'explicit':{}}
    for row in cat.by_key[key]:
        result=storage.lookup(key,row['standard'])
        golden[key]['explicit'][row['standard']]={'lookup':result['status'],'geometry':result['profile']['geometry'],
          'uzk':model.calculate_uzk(result),'provenance':[p['evidence'] for p in row['provenance']]}
assert golden['35Б1']['without_standard']=='FOUND_PREFERRED'
for ndv,g,pairs in [('ГОСТ Р 57837-2017',dict(h=346,b=174,s=6,t=9),[(328,6),(348,9)]),
                   ('ГОСТ 26020-83',dict(h=346,b=155,s=6.2,t=8.5),[(329,6.2),(310,8.5)])]:
    row=golden['35Б1']['explicit'][ndv]; assert row['geometry']==g
    assert [(r['L'],r['S']) for r in row['uzk']['rows']]==pairs
assert cat.lookup('35Б1','ГОСТ 26020-1983')['status']=='NOT_FOUND'
assert golden['20Ш1']['without_standard']=='AMBIGUOUS'
for row in cat.rows:
    assert (row['uzk_rule']==model.RULE)==(row['family']=='Двутавр')
    if row['family']!='Двутавр' and row['status']!='КОНФЛИКТ':
        assert model.calculate_uzk(cat.lookup(row['key'],row['standard']))['status']==model.NOT_CONFIGURED
save('golden_independent.json',golden)

# Full exact XLSX read and comparison, independently of product reopen routine.
wb=load_workbook(folder/'Профили.xlsx',read_only=True,data_only=False)
matrices={}; formula_count=0
for ws in wb:
    matrix=list(ws.iter_rows(values_only=True)); matrices[ws.title]=matrix
    formula_count+=sum(c.data_type=='f' for row in ws.iter_rows() for c in row)
assert list(matrices)==['Профили','Источники','Правила УЗК','Предпочтения','Конфликты']
assert formula_count==0
headers=matrices['Профили'][0]; row_by_id={}
for values in matrices['Профили'][1:]:
    r=dict(zip(headers,values)); row_by_id[(r['Ключ'],r['НД'])]=r
    for f in model.GEOMETRY:
        assert r[f] is None or (isinstance(r[f],(int,float)) and not isinstance(r[f],bool))
    assert r['Evidence reference'].startswith('profiles.catalog.json#')
for r in cat.rows:
    x=row_by_id[(r['key'],r['standard'])]
    assert all(x[f]==r['geometry'].get(f) for f in model.GEOMETRY)
    assert x['Статус']==r['status'] and x['УЗК правило']==r['uzk_rule']
assert len(row_by_id)==len(cat.rows)
summary={'sheets':{k:len(v)-1 for k,v in matrices.items()},'formulas':formula_count,'numeric_geometry':True,
 '35Б1_rows':{ndv:r for (key,ndv),r in row_by_id.items() if key=='35Б1'},
 'decimal_tube_samples':[r for r in row_by_id.values() if r['Тип'] in ('Квадратный профиль','Прямоугольный профиль') and isinstance(r['t'],float)][:5]}
wb.close(); save('workbook_independent.json',summary)

# Contract probes with pure isolated records; no actual state mutation.
r=deepcopy(next(r for r in records if r['key']=='35Б1' and r['standard']=='СТО АСЧМ 20-93'))
r2=deepcopy(r); r2['record_id']='REVIEW-R'; r2['source']='REVIEW'; r2['geometry']['R']=r['geometry']['R']+1
radius= model.Catalog([r,r2],user_state=model.default_user_state()).lookup('35Б1',r['standard'])
probes={'radius_only_same_ND':{'status':radius['status'],'geometry':radius['profile']['geometry']},
 'frequency':{v:model.frequency_angle(v) for v in ('11.5','11.999','12.0','12.001','11,5','12,0')},
 'normalization':{v:model.normalize_profile(v) for v in ('80x40x4','80х40х4','80×40×4',' I 35Б1 ','35Б01','35Б1','35B1','L100х8','Модель A B','Модель AB')}}
assert all(model.number(x) is not None for x in ('6,2','8.5'))
bad=[]
for x in ('NaN','Infinity','-Infinity',True,'oops'):
    try:model.number(x)
    except model.ProfileError:bad.append(str(x))
assert len(bad)==5
probes['invalid_numbers_blocked']=bad
save('policy_probes.json',probes)

# Evidence finalizer: validate only, never execute its writer/main.
spec=importlib.util.spec_from_file_location('review_finalizer',E/'finalize.py')
fin=importlib.util.module_from_spec(spec); spec.loader.exec_module(fin)
assert fin.validate()['status']=='PASS'
negative={}; orig=fin.read_json
for case in ('stale_product_after','failed_test','stale_workbook_hash','missing_product_path'):
    def altered(p):
        d=orig(p)
        if case=='stale_product_after' and Path(p).name=='product_after.json':d[next(iter(d))]='0'*64
        if case=='failed_test' and Path(p).name=='tests.json':d['failed']=1
        if case=='stale_workbook_hash' and Path(p).name=='build_result.json':d['sha256']='0'*64
        if case=='missing_product_path' and Path(p).name=='product_before.json':d.pop(next(iter(d)))
        return d
    with patch.object(fin,'read_json',altered):
        try:fin.validate()
        except ValueError as exc:negative[case]=str(exc)
        else:raise AssertionError('Negative finalizer accepted')
tests=read(E/'tests.json')
before=read(E/'product_before.json'); after=read(E/'product_after.json')
assert before==after==manifest['tested_product_hashes']==tests['tested_product_hashes']
assert len(tests['tests'])==56 and all(t['status']=='PASS' for t in tests['tests'])
newer=[str(p.relative_to(E)) for p in E.rglob('*') if p.is_file() and p!=E/'STAGE4_REVIEW_MANIFEST.json' and p.stat().st_mtime_ns>(E/'STAGE4_REVIEW_MANIFEST.json').stat().st_mtime_ns]
assert not newer
save('evidence_independent.json',{'before_after_current_equal':True,'manifest_written_last':True,'negative_finalizer':negative,
 'reported_tests':56,'independently_reexecuted_original_suite':False,'reason':'Original suite builds real DEV; independent tests use fixtures only'})
for name in ['passport_app.config','passport_app.registry','passport_app.service','passport_app.vik','passport_app.catalog','passport_app.catalog_gui','stage2','stage3']:
    importlib.import_module(name)
checkpoint=read(ROOT/'DEV/DEV_REGISTRY_CHECKPOINT.json')
assert sha(ROOT/'DEV/ВЗМК_DEV.xlsx')==checkpoint['sha256']
registry=load_workbook(ROOT/'DEV/ВЗМК_DEV.xlsx',read_only=True,data_only=False)
assert 'Паспорт качества' in registry.sheetnames; registry.close()
assert all(sha(p)==h for p,h in reviewed.items())
save('reviewed_hashes_after_inspect.json',{p:sha(p) for p in reviewed})
print(json.dumps({'counts':counts,'conflict_kinds':dict(Counter(c['kind'] for c in conflicts)),
 'sources':{'A':a_counts,'B':b_counts},'radius_probe':probes['radius_only_same_ND'],
 'workbook':summary,'network_production_events':events,'preservation':True},ensure_ascii=False,default=str))
