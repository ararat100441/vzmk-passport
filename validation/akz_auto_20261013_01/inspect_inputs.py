import sys,json,hashlib,os,subprocess
from pathlib import Path
sys.dont_write_bytecode=True
OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[1];sys.path.insert(0,str(ROOT))
from openpyxl import load_workbook
from passport_app.registry import read_registry,select_project
from akz_stage6 import core as c,part2
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(p,d):
    with Path(p).open('x',encoding='utf8') as f:json.dump(d,f,ensure_ascii=False,indent=2,default=str)
seed=Path(r'C:\Users\arara\AppData\Local\VZMKStage9_092_9c0a3396\source')
settings=json.loads((seed/'settings.json').read_text(encoding='utf8'))
from passport_app.config import Settings
cfg=Settings.load(seed/'settings.json');reg=read_registry(seed/'DEV/ВЗМК_DEV.xlsx');project=select_project(reg,'Гараж Казахстан');passport_root,_=cfg.project_paths(project.name)
rows=[]
for n in range(4):
    names=(c.ACT.name,c.APPLICATION.name) if n==0 else (f'АКЗ-{n}.xlsx',f'Приложение к АКЗ-{n}.xlsx')
    pins=list(c.PINS.values()) if n==0 else part2.PINS[n]
    for kind,name,pin in zip(('act','application'),names,pins):
        p=passport_root/name;assert sha(p)==pin
        b=load_workbook(p,data_only=False);s=b['4'] if kind=='act' else b.active
        cells=['B3','E4','H4','E5','E6','E7','B44','F50','I50','A1','A2']+[col+str(r) for col in ('E','H','I','J') for r in (17,18,21,25,26,30,31,32,33,34)]
        rows.append({'variant':n,'kind':kind,'path':str(p),'sha256':pin,'sheets':b.sheetnames,'cells':{x:str(s[x].value) for x in cells},'formulas':{sheet.title:{cell.coordinate:cell.value for row in sheet for cell in row if cell.data_type=='f'} for sheet in b}});b.close()
baseline=json.loads((ROOT/'validation/stage9_gui_099_20261012_01/BINDING_AFTER.json').read_text(encoding='utf8'))
accepted=baseline['accepted_files'];assert all(sha(p)==h for p,h in accepted.items())
sources={str(p):sha(p) for p in [ROOT/'stage9.py',ROOT/'stage9_099.spec',*(ROOT/'stage9_app').glob('*.py')]}
save(OUT/'SOURCE_BEFORE.json',sources);save(OUT/'ACCEPTED_BEFORE.json',accepted)
save(OUT/'TEMPLATE_BINDINGS_BEFORE.json',{'project':project.__dict__,'pairs':rows})
save(OUT/'SEED_BEFORE.json',{'source':str(seed),'files':{str(p):sha(p) for p in seed.rglob('*') if p.is_file()}})
print(json.dumps({'project':project.__dict__,'pairs':rows},ensure_ascii=False,indent=2,default=str))
