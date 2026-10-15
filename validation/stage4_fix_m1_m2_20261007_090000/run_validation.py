"""M1/M2 isolated targeted validation; never builds the real DEV catalogue."""
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
import importlib
import json
import shutil
import subprocess
import sys
import traceback
import zipfile
import uuid

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding='utf-8')
import profiles_stage4.storage as s
from profiles_stage4.model import Catalog, ProfileError, RULE, NOT_CONFIGURED, calculate_uzk, frequency_angle, PRIMARY
from profiles_stage4.offline import load_sources, read_json, sha, no_network
from profiles_stage4.windows_generation import ProtectedFile

FILES = [s.DEFAULT_WORKBOOK.name, s.STATE, s.CATALOG, s.COMMIT]
RUN_ID = uuid.uuid4().hex[:12]
tests, network_events, production_events = [], [], []

def audit(event, args):
    if event.startswith('socket.') and event in ('socket.connect','socket.getaddrinfo','socket.sendto'):
        network_events.append(event)
        raise ProfileError('Validation: network prohibited')
    if event in ('open','os.rename','os.remove','os.mkdir'):
        for arg in args:
            if isinstance(arg,(str,bytes)):
                name = str(arg).replace('\\','/')
                if '/Работа/ВЗМК.xlsx' in name or '/Работа/Профили.xlsx' in name:
                    production_events.append(name)
                    raise ProfileError('Validation: production prohibited')
sys.addaudithook(audit)

def save(name, data):
    s.dump(OUT/name, data)

def hashes():
    products = sorted(p for p in (ROOT/'profiles_stage4').glob('*') if p.is_file())
    products += [ROOT/'stage4.py', ROOT/'STAGE4_README.md', ROOT/'Запустить Этап 4.cmd']
    products += sorted(OUT.glob('*.py'))
    return {str(p.relative_to(ROOT)):sha(p) for p in products}

def check(condition, message='Assertion failed'):
    if not condition:
        raise AssertionError(message)

def rejected(call):
    try:
        call()
    except (ValueError, OSError) as exc:
        return str(exc)
    raise AssertionError('Expected fail-closed STOP')

def no_uzk(result):
    if result['status'] not in ('FOUND','FOUND_PREFERRED'):
        rejected(lambda:calculate_uzk(result))
    else:
        check(calculate_uzk(result)['status']!='CALCULATED')

def case(name, call):
    try:
        detail = call()
        tests.append({'name':name,'status':'PASS','detail':detail})
        print('PASS '+name, flush=True)
        save('targeted_tests.json', {'status':'RUNNING','tests':tests})
    except BaseException:
        tests.append({'name':name,'status':'FAIL','traceback':traceback.format_exc()})
        save('targeted_tests.json', {'status':'FAIL','tests':tests})
        raise

def fixture(label):
    folder = OUT/'fixtures'/RUN_ID/label/'DEV/Справочники'
    folder.mkdir(parents=True)
    for name in FILES:
        shutil.copy2(OUT/'baseline_generation'/name, folder/name)
    return folder/FILES[0]

def generation(target):
    return {name:sha(target.parent/name) for name in FILES}

def backups(target):
    folder = target.parent/'backups'
    return sorted(folder.iterdir()) if folder.exists() else []

def backup_matches(path, expected):
    check({name:sha(path/name) for name in FILES} == expected)
    s.validated_generation({name:(path/name).read_bytes() for name in FILES}, FILES[0])

def verify_protected():
    expected = read_json(OUT/'protected_before.json')
    actual = {p:sha(p) for p in expected}
    check(expected == actual, 'Protected file changed')
    check(read_json(OUT/'dev_generation_before.json') == {p:sha(p) for p in read_json(OUT/'dev_generation_before.json')},
          'Real DEV was modified by isolated tests')
    return {'protected_files':len(actual),'real_DEV_generation_unchanged':True}

before = hashes()
save('product_hashes_before.json', before)
case('protected baseline + real DEV unchanged', verify_protected)
with no_network() as ingestion_network:
    records, sources = load_sources()
check(not ingestion_network)
state = read_json(OUT/'baseline_generation'/s.STATE)
cat = Catalog(records,sources,state)

def current_fifteen():
    independent = read_json(ROOT/'validation/stage4_independent_review_20261007_0119/conflicts_independent.json')
    check(len(cat.conflicts) == len(independent) == 15)
    original_data = read_json(OUT/'baseline_generation'/s.CATALOG)
    check(cat.raw_records == original_data['records'], 'Raw/provenance changed')
    actual = []
    for old in independent:
        conflict = next(c for c in cat.conflicts if c['id'] == old['id'])
        expected = 'NON_BLOCKING' if old['kind'] == 'area_mass_only' else 'BLOCKING'
        check(conflict['classification'] == expected)
        check(set(conflict['fields']) == set(old['fields']))
        result = cat.lookup(old['key'],old['standard'])
        row = result['profile']
        check(result['status'] == ('FOUND' if expected == 'NON_BLOCKING' else 'CONFLICT'))
        check(row['unresolved_fields'] == conflict['fields'])
        check(len(row['provenance']) >= 2 and not conflict['resolution'])
        for field in conflict['fields']:
            check(field not in row['geometry'] and field not in row['quality'])
        if expected == 'NON_BLOCKING':
            check(all(row['geometry'].get(f) is not None for f in PRIMARY[row['family']]))
            check(row['status'] == 'СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ')
        else:
            check(not row['geometry'])
            no_uzk(result)
        actual.append({'id':old['id'],'key':old['key'],'standard':old['standard'],
                       'fields':conflict['fields'],'classification':expected,'lookup':result['status'],
                       'geometry':row['geometry'],'raw_members':conflict['members']})
    check(sum(r['classification']=='BLOCKING' for r in actual) == 8)
    check(sum(r['classification']=='NON_BLOCKING' for r in actual) == 7)
    save('disagreements_post_fix.json', actual)
    return {'disagreements':15,'blocking_CONFLICT':8,'non_blocking_FOUND':7,'raw_unchanged':True}
case('M1 all 15 independent classifications, lookups and raw provenance', current_fifteen)

beam = deepcopy(next(r for r in records if r['key']=='35Б1' and r['standard']=='ГОСТ Р 57837-2017'))
def beam_fixture(changes, *, missing=False):
    a,b = deepcopy(beam),deepcopy(beam)
    a['source'],b['source'] = 'FIXTURE_A','FIXTURE_B'
    a['record_id'],b['record_id'] = 'FIXTURE_A:1','FIXTURE_B:1'
    for field in ('R','r'):
        a['geometry'].pop(field,None); b['geometry'].pop(field,None)
    for field, values in changes.items():
        target = 'geometry' if field in ('h','b','b2','s','t','D','R','r') else 'quality'
        a[target][field],b[target][field] = values
    return Catalog([a,b])

def auxiliary(field):
    c = beam_fixture({field:(10,11)})
    result = c.lookup('35Б1','ГОСТ Р 57837-2017')
    check(result['status']=='FOUND')
    row = result['profile']
    check(row['geometry']['h']==346 and row['geometry']['t']==9)
    check(field not in row['geometry'] and field not in row['quality'])
    check(row['unresolved_fields']==[field] and row['disagreement_class']=='NON_BLOCKING')
    check(len(row['provenance'])==2 and field in '; '.join(row['notes']))
    check(calculate_uzk(result)['status']=='CALCULATED')
    check(c.conflicts[0]['members']==row['provenance'])
    return {'lookup':'FOUND','unresolved':field,'UZK':'CALCULATED','provenance_members':2}
for field in ['R','r','area_cm2','mass_kg_m','reference_only']:
    case('M1 non-blocking '+field+' same ND + I-beam UZK', lambda f=field:auxiliary(f))

def blocking(changes):
    c = beam_fixture(changes)
    explicit, plain = c.lookup('35Б1','ГОСТ Р 57837-2017'),c.lookup('35Б1')
    check(explicit['status']==plain['status']=='CONFLICT')
    check(not explicit['profile']['geometry'])
    check(c.conflicts[0]['classification']=='BLOCKING')
    no_uzk(explicit)
    return {'explicit':'CONFLICT','preferred_identity':'CONFLICT','UZK':'BLOCKED'}
case('M1 mixed core + R/r', lambda:blocking({'b':(174,175),'R':(10,11),'r':(4,5)}))
case('M1 mixed core + area/mass', lambda:blocking({'b':(174,175),'area_cm2':(10,11),'mass_kg_m':(12,13)}))
for field in ['h','b','b2','s','t','D']:
    case('M1 blocking field '+field, lambda f=field:blocking({f:(10,11)}))

def type_disagreement():
    a,b = deepcopy(beam),deepcopy(beam)
    b['family']='Швеллер'; b['record_id']='TYPE_FIXTURE'
    c = Catalog([a,b])
    check(c.lookup('35Б1','ГОСТ Р 57837-2017')['status']=='CONFLICT')
    return {'family_disagreement':'CONFLICT'}
case('M1 family/type blocking', type_disagreement)

def missing_radius():
    c = beam_fixture({})
    r = c.lookup('35Б1',beam['standard'])
    check(r['status']=='FOUND' and not c.conflicts)
    check(calculate_uzk(r)['status']=='CALCULATED')
    return {'missing_R_r_tolerated':True}
case('M1 missing radii tolerated', missing_radius)

def golden():
    result={}
    for nd,want,uzk in [('ГОСТ Р 57837-2017',[346,174,6,9],[[328,6],[348,9]]),
                        ('ГОСТ 26020-83',[346,155,6.2,8.5],[[329,6.2],[310,8.5]])]:
        r=cat.lookup('35Б1',nd)
        check(r['status']=='FOUND')
        check([r['profile']['geometry'][f] for f in ('h','b','s','t')]==want)
        u=calculate_uzk(r)
        check([[row['L'],row['S']] for row in u['rows']]==uzk)
        result[nd]={'lookup':r['status'],'geometry':want,'UZK':uzk}
    check(cat.lookup('35Б1')['status']=='FOUND_PREFERRED')
    check(cat.lookup('20Ш1')['status']=='AMBIGUOUS')
    check(frequency_angle(11.5)=={'frequency_mhz':5,'angle_degrees':70})
    check(frequency_angle(12)=={'frequency_mhz':2.5,'angle_degrees':65})
    check(cat.lookup('I 35Б1')['status']=='FOUND_PREFERRED')
    check(cat.lookup('NONEXISTENT')['status']=='NOT_FOUND')
    check(cat.lookup('35Б1','ГОСТ 00000-00')['status']=='NOT_FOUND')
    for family in PRIMARY:
        if family!='Двутавр':
            row=next(r for r in cat.rows if r['family']==family and r['status']!='КОНФЛИКТ')
            check(calculate_uzk(cat.lookup(row['key'],row['standard']))['status']!='CALCULATED')
    result.update({'plain_35Б1':'FOUND_PREFERRED','plain_20Ш1':'AMBIGUOUS','frequency_11.5':[5,70],
                   'frequency_12':[2.5,65],'unsupported_UZK':'NOT_CONFIGURED'})
    save('golden_regression.json',result)
    return result
case('Golden 35Б1 standard-specific / UZK / frequency / unsupported types', golden)

# Actual renderer is used for normal application writes. Fault fixtures reuse its
# verified value-only workbook when the payload is identical, avoiding redundant renders.
normal = fixture('normal_application_writes')
initial_hashes = generation(normal)
real_render = s.render_workbook
initial_payload = s.workbook_payload(Catalog(records,sources,{**state,'revision':state['revision']+1}))

def normal_rebuild():
    result = s.build(normal, previews=OUT/'previews')
    check(result['network_attempts']==0 and result['reopen']['reopened'])
    backup_matches(Path(result['backup']),initial_hashes)
    check(s.load_catalog(normal).counts()==cat.counts())
    check(s.load_catalog(normal).user_state['preferences']==state['preferences'])
    return result
case('M2 real offline normal rebuild + coherent backup + M1 workbook reopen', normal_rebuild)
template = OUT/'rendered_template.xlsx'
shutil.copy2(normal,template)
s.verify_workbook(template,initial_payload)

def workbook_visibility():
    from openpyxl import load_workbook
    wb=load_workbook(normal,read_only=True,data_only=True)
    try:
        rows=list(wb['Конфликты'].iter_rows(values_only=True))
        check(rows[0][-1]=='Классификация')
        check(sum(r[-1]=='BLOCKING' for r in rows[1:])==8)
        check(sum(r[-1]=='NON_BLOCKING' for r in rows[1:])==7)
        check(sum(r[9]=='СПРАВОЧНЫЕ ДАННЫЕ НЕ РАЗРЕШЕНЫ' for r in rows[1:])==7)
        profiles=list(wb['Профили'].iter_rows(min_row=2,values_only=True))
        check(sum(r[13]=='СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ' and 'NON_BLOCKING' in r[19] for r in profiles)==7)
        old=load_workbook(OUT/'baseline_generation'/FILES[0],read_only=True,data_only=True)
        try:
            for name in ['Источники','Правила УЗК','Предпочтения']:
                check(list(wb[name].iter_rows(values_only=True))==list(old[name].iter_rows(values_only=True)),
                      'Unrelated sheet values changed: '+name)
        finally:
            old.close()
        return {'conflicts_sheet':15,'BLOCKING':8,'NON_BLOCKING':7,'visible_profile_notes':7}
    finally:
        wb.close()
case('M1 workbook visibly differentiates classification / unresolved fields', workbook_visibility)

def copied_render(payload,destination,previews=None):
    check(read_json(payload)==initial_payload,'Fixture payload must equal actual verified renderer payload')
    shutil.copy2(template,destination)
    return {'network_attempts':0}

def external_edit(target,name):
    p=target.parent/name
    if name.endswith('.xlsx'):
        with zipfile.ZipFile(p,'a') as archive:
            archive.comment=b'VALID_EXTERNAL_EDIT_DURING_RENDER'
        from openpyxl import load_workbook
        wb=load_workbook(p,read_only=True)
        check(wb.sheetnames==list(s.HEADERS))  # Valid previous-format XLSX, not corrupted bytes.
        wb.close()
    else:
        value=read_json(p)
        if name==s.STATE:
            value['preferences'][0]['basis']='VALID_EXTERNAL_FIXTURE_BASIS'
        elif name==s.CATALOG:
            value['records'][0]['notes'].append('VALID_EXTERNAL_FIXTURE_NOTE')
        else:
            value['built_at']='2026-10-07T00:00:00+00:00'
        s.dump(p,value)
        check(read_json(p)==value)
    return sha(p)

def render_race(name):
    t=fixture('render_race_'+str(FILES.index(name)))
    old=generation(t); edited={}
    def render(payload,destination,previews=None):
        result=copied_render(payload,destination)
        edited['sha256']=external_edit(t,name)
        return result
    with patch.object(s,'render_workbook',render):
        message=rejected(lambda:s.build(t))
    final=generation(t)
    check('sha256' in edited, 'Render/edit not reached: '+message)
    check(final[name]==edited['sha256'] and final[name]!=old[name])
    check(all(final[n]==old[n] for n in FILES if n!=name))
    check(not backups(t), 'Changed generation must STOP before backup')
    detail={'file':name,'valid_external_edit':True,'external_sha256':edited['sha256'],
            'STOP':message,'other_files_unchanged':True,'backup_not_created':True,'external_bytes_survive':True}
    save('race_render_'+str(FILES.index(name))+'.json',detail)
    return detail
for name in FILES:
    case('M2 external valid edit during render '+name, lambda n=name:render_race(n))

def writer(mode,target):
    result=subprocess.run([sys.executable,'-I','-B',str(OUT/'external_writer.py'),mode,str(target.parent)],
                          capture_output=True,text=True,encoding='utf-8',timeout=20)
    check(result.returncode==0,result.stderr)
    return json.loads(result.stdout)

def boundary():
    t=fixture('commit_boundary')
    old=generation(t); attempts=[]
    original=ProtectedFile.rename
    def blocked_by_windows(result):
        # Python's CRT open maps Windows sharing denial to EACCES without winerror.
        return result['blocked'] and (result['winerror'] in (5,32,33) or
            (result['operation']=='write' and result['winerror'] is None and result['errno']==13))
    def rename(handle,destination):
        destination=Path(destination)
        if not attempts:
            found=writer('attack',t)
            check(all(blocked_by_windows(r) for r in found),str(found))
            attempts.append({'boundary':'immediately before first protected rename','attempts':found})
        original(handle,destination)
        if len(attempts)==1 and destination==t:
            found=writer('attack',t)
            check(all(blocked_by_windows(r) for r in found),str(found))
            attempts.append({'boundary':'immediately after first new file publication; commit still in progress','attempts':found})
    with patch.object(s,'render_workbook',copied_render),patch.object(ProtectedFile,'rename',rename):
        result=s.build(t)
    check(len(attempts)==2)
    backup_matches(Path(result['backup']),old)
    s.load_catalog(t); s.verify_workbook(t,initial_payload)
    save('commit_boundary.json',attempts)
    return {'independent_process':True,'boundary_phases':2,'write_delete_replace_attempts':24,
            'all_blocked_by_Windows':True,'coherent_backup':True,'normal_commit':'PASS'}
case('M2 actual commit boundary independent-process write/delete/replace protection', boundary)

def vacancy():
    t=fixture('commit_name_vacancy')
    old=generation(t); created=[]
    original=ProtectedFile.rename
    def rename(handle,destination):
        old_path=handle.path
        original(handle,destination)
        if not created and old_path==t and Path(destination).parent.name=='previous':
            created.append(writer('create_vacant',t))
    with patch.object(s,'render_workbook',copied_render),patch.object(ProtectedFile,'rename',rename):
        message=rejected(lambda:s.build(t))
    check(t.read_bytes()==b'EXTERNAL_CREATION_AT_COMMIT_BOUNDARY')
    check(all(sha(t.parent/n)==old[n] for n in FILES[1:]))
    check(len(backups(t))==1)
    backup_matches(backups(t)[0],old)
    detail={'STOP':message,'external_creation_survives':True,'other_files_unchanged':True,
            'coherent_previous_generation_in_backup':True,'ReplaceIfExists':False}
    save('commit_name_vacancy.json',detail)
    return detail
case('M2 even commit name-vacancy external creation cannot be overwritten', vacancy)

def failure(label,mode):
    t=fixture('failure_'+label); old=generation(t)
    if mode=='render':
        with patch.object(s,'render_workbook',side_effect=ProfileError('Injected renderer failure')):
            message=rejected(lambda:s.build(t))
    elif mode=='backup_write':
        original=Path.write_bytes
        def write(path,data):
            if 'backups' in path.parts:
                raise OSError('Injected backup write failure')
            return original(path,data)
        with patch.object(s,'render_workbook',copied_render),patch.object(Path,'write_bytes',write):
            message=rejected(lambda:s.build(t))
    else:
        with patch.object(s,'render_workbook',copied_render):
            message=rejected(lambda:s.build(t,fail_at=mode))
    check(generation(t)==old)
    if mode=='backup_write':
        check(not backups(t),'Incomplete backup must never be exposed')
    s.load_catalog(t)
    return {'STOP':message,'previous_generation_exactly_preserved':True}
for label,mode in [('render','render'),('before_commit','before_commit'),('backup_write','backup_write'),
                   ('mid_commit','after_first_replace')]:
    case('M2 rollback / no replacement on '+label+' failure',lambda a=label,b=mode:failure(a,b))

def locked_workbook():
    t=fixture('locked_workbook'); old=generation(t)
    # Same sharing restriction as Excel: an existing reader which does not allow DELETE.
    with t.open('rb'):
        message=rejected(lambda:s.build(t))
    check(generation(t)==old)
    return {'STOP':message,'previous_generation_exactly_preserved':True}
case('M2 locked workbook STOP without replacement',locked_workbook)

def app_writes():
    result=s.set_preference('20Ш1','ГОСТ 26020-83','ISOLATED_TEST_EXPLICIT_PREFERENCE',workbook=normal)
    check(s.lookup('20Ш1',workbook=normal)['status']=='FOUND_PREFERRED')
    result=s.add_manual('Ручной FIX M1 M2','Двутавр','РУЧНОЙ НД',{'h':200,'b':100,'s':5,'t':8},
                        'ISOLATED_TEST_MANUAL',workbook=normal)
    r=s.lookup('Ручной FIX M1 M2','РУЧНОЙ НД',workbook=normal)
    check(r['status']=='FOUND' and r['profile']['uzk_rule']==NOT_CONFIGURED)
    current=s.load_catalog(normal)
    c=next(c for c in current.conflicts if c['classification']=='BLOCKING')
    result=s.resolve_conflict(c['id'],c['members'][0]['record_id'],'ISOLATED_TEST_EXPLICIT_RESOLUTION',workbook=normal)
    check(s.lookup(c['key'],c['standard'],workbook=normal)['status']=='FOUND')
    previous=s.load_catalog(normal).user_state
    result=s.build(normal)
    final=s.load_catalog(normal).user_state
    check(final['revision']==previous['revision']+1)
    for field in ('manual_rows','preferences','resolutions'):
        check(previous[field]==final[field])
    check(next(p for p in final['preferences'] if p['key']=='35Б1') == next(p for p in state['preferences'] if p['key']=='35Б1'))
    return {'application_preference':'PASS','manual_add':'PASS','explicit_resolution':'PASS',
            'subsequent_rebuild_preservation':'PASS','35Б1_preference_unchanged':True}
case('M2 normal application preference / manual add / explicit resolution / preservation',app_writes)

def stage123_smoke():
    for name in ['main','stage2','stage3','passport_app.config','passport_app.catalog']:
        importlib.import_module(name)
    checkpoint=read_json(ROOT/'DEV/DEV_REGISTRY_CHECKPOINT.json')
    from openpyxl import load_workbook
    wb=load_workbook(ROOT/'DEV/ВЗМК_DEV.xlsx',read_only=True,data_only=False)
    try:
        names=wb.sheetnames
        check(bool(names))
    finally:
        wb.close()
    check(bool(checkpoint))
    return {'imports':'PASS','DEV_checkpoint_readable':True,'DEV_registry_readable':True,'sheets':names}
case('Stage 1–3 imports + checkpoint readable only',stage123_smoke)
case('protected hashes + real DEV still unchanged after all isolated writes',verify_protected)
after=hashes()
save('product_hashes_after.json',after)
check(before==after,'Product changed during targeted validation')
check(not network_events and not production_events)
save('post_fix_counts.json',cat.counts())
save('targeted_tests.json',{'status':'PASS','tests':tests,'passed':len(tests),'failed':0,
    'product_hashes_before_after_equal':True,'network_events':network_events,'production_events':production_events,
    'real_DEV_builds':0,'isolated_mutation_tests_only':True})
print(json.dumps({'status':'PASS','passed':len(tests),'counts':cat.counts()},ensure_ascii=False),flush=True)
