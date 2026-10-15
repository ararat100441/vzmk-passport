"""Independent review evidence only. No product edits, no production access."""
import sys, json, hashlib, copy, zipfile, traceback, os, shutil
from pathlib import Path
from decimal import Decimal
from datetime import datetime
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[2]
E=Path(__file__).resolve().parent
I=ROOT/'validation/stage6_akz_part1_20261009_085822'
sys.dont_write_bytecode=True
if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0,str(ROOT))
def load(p):return json.loads(Path(p).read_text(encoding='utf-8'))
def save(name,value):
    (E/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def sha(p):
    p=Path(p)
    assert p.name.casefold() not in ('взмк.xlsx','узк.xlsm'),f'PRODUCTION ACCESS DENIED {p}'
    return hashlib.sha256(p.read_bytes()).hexdigest()
def bootstrap():
    m=load(I/'EVIDENCE_MANIFEST.json');tested=load(I/'tested_product_hashes.json')
    sources=load(I/'source_integrity_current.json')
    protected=dict(sources['BEFORE'])
    protected.update(tested['CURRENT']);protected.update(tested['launch_and_documentation_current'])
    protected.update(m['files']);protected[str(I/'EVIDENCE_MANIFEST.json')]=sha(I/'EVIDENCE_MANIFEST.json')
    paths={str(p):sha(p) for p in protected}
    expected_outputs={'ТЕСТ АКЗ №0.xlsx':'1a52736427e1529420723ff85f3bfd92340865ece895173ffaaa0dadb48ba06f',
                      'ТЕСТ Приложение к АКЗ №0.xlsx':'4cd657f84311512255e6a6e39b492c227f18df3ab8b5ab03eb34400f165d9c6d'}
    outputs={str(Path(m['canonical_output'])/n):sha(Path(m['canonical_output'])/n) for n in expected_outputs}
    assert list(outputs.values())==list(expected_outputs.values()),'EXACT OUTPUT HASH MISMATCH STOP'
    mismatches=[{'path':p,'expected':h,'actual':paths[p]} for p,h in protected.items() if h!=paths[p]]
    tested_equality={p:all(tested[k][p]==paths[p] for k in ('BEFORE','AFTER','CURRENT')) for p in tested['CURRENT']}
    save('hash_before.json',{'all':paths,'source_count':len(sources['BEFORE']),'mismatches':mismatches,
                           'tested_equality':tested_equality,'outputs':outputs})
    print(json.dumps({'protected_count':len(paths),'source_count':len(sources['BEFORE']),
                      'mismatches':mismatches,'tested_equality':tested_equality,'outputs':outputs},ensure_ascii=False,indent=2))
def stored_analysis():
    from openpyxl import load_workbook
    from passport_app.registry import read_registry,select_project
    m=load(I/'EVIDENCE_MANIFEST.json');pv=load(I/'preview.json')['payload']
    registry=Path(pv['sources'].__iter__().__next__())
    registry=ROOT/'DEV/ВЗМК_DEV.xlsx'
    accepted=select_project(read_registry(registry),'Гараж Казахстан')
    result={'accepted_selection':accepted.__dict__,'workbooks':{}}
    targets=[registry,Path(pv['templates']['act']),Path(pv['templates']['application']),
             Path(load(I/'test_assignments.json')['fixture']),ROOT.parent/'Гараж казахстан/Справочник марок.xlsx']
    targets+=list(Path(m['canonical_output']).glob('*.xlsx'))
    for path in targets:
        a=load_workbook(path,data_only=False);b=load_workbook(path,data_only=True)
        try:
            sheets={}
            for s in a:
                rows=[1,353,400] if path==registry else range(1,s.max_row+1)
                cells={c.coordinate:{'value':c.value,'cached':b[s.title][c.coordinate].value,'type':c.data_type,
                                     'number_format':c.number_format}
                       for rn in rows for c in s[rn] if c.value is not None}
                sheets[s.title]={'dimensions':s.calculate_dimension(),'cells':cells,
                                 'hidden_rows':[r for r,d in s.row_dimensions.items() if d.hidden],
                                 'merged':list(map(str,s.merged_cells.ranges)),'print_area':str(s.print_area),
                                 'center_header':s.oddHeader.center.text,'center_footer':s.oddFooter.center.text}
            result['workbooks'][str(path)]={'sha256':sha(path),'sheets':sheets}
        finally:a.close();b.close()
    save('stored_analysis.json',result)
    print('STORED ANALYSIS WRITTEN',len(targets))
if __name__=='__main__':
    bootstrap() if len(sys.argv)==1 else stored_analysis()
