"""Independent arithmetic and source preservation using native snapshots."""
from review import *
from openpyxl import load_workbook
from collections import Counter
data=load(E/'native_vom_independent.json')
s=next(s for s in data['sheets'] if s['name']=='!01_ВЗМК_A_Ведомость отправочны')
assert s['used']=='$A$1:$K$507' and data['read_only']
assert [str(s['values'][1][i]).strip() for i in (1,2,3,7,9)]==[
    'Отправочная марка','Наименование марки','Количество','Масса (изделие), кг','Площадь окраски, м2']
assert [str(s['values'][2][i]).strip() for i in (7,8,9,10)]==['1 дет.','Всех','1 дет.','Всех']
hidden=set(s['hidden_rows'])
all_rows=[(n,r) for n,r in enumerate(s['values'],1) if n>=4 and r[1] is not None and str(r[1]).strip().lower() not in ('итого','всего','')]
visible=[(n,r) for n,r in all_rows if n not in hidden]
ind=load(E/'independent_preview.json')['payload']
assert [n for n,r in visible]==[r['row'] for r in ind['vom']['rows']]
for (n,raw),product in zip(visible,ind['vom']['rows']):
    assert [raw[i] for i in (1,2,3,7,8,9,10)]==[product[k] for k in ('mark','name','quantity','mass_unit','mass_total','area_unit','area_total')]
selected=[]
for n,raw in visible:
    if raw[1] not in ('Бп2-5','Бп2-18','Бп2-19'):continue
    row={'row':n,**dict(zip(('mark','name','quantity','mass_unit','mass_total','area_unit','area_total'),[raw[i] for i in (1,2,3,7,8,9,10)]))}
    for f in ('quantity','mass_unit','mass_total','area_unit','area_total'):row[f]=Decimal(str(row[f]))
    assert row['mass_unit']*row['quantity']==row['mass_total']
    assert row['area_unit']*row['quantity']==row['area_total']
    selected.append(row)
totals={f:sum(r[f] for r in selected) for f in ('quantity','mass_total','area_total')}
assert totals==dict(quantity=Decimal(3),mass_total=Decimal('105.26'),area_total=Decimal('3.989'))
assert [r['row'] for r in selected]==[40,53,54]
original=ROOT.parent/'Гараж казахстан/Справочник марок.xlsx'
fixture=Path(load(I/'test_assignments.json')['fixture'])
a=load_workbook(original,data_only=False);b=load_workbook(fixture,data_only=False)
try:
    assert a.sheetnames==b.sheetnames and a.active.max_row==b.active.max_row==196
    assert all(a.active.cell(r,col).value==b.active.cell(r,col).value for r in range(1,197) for col in (1,2))
    assert all(a.active.cell(r,3).value is None for r in range(2,197))
    assignments=[{'mark':b.active.cell(r,1).value,'row':r,'value':b.active.cell(r,3).value,'type':b.active.cell(r,3).data_type}
                 for r in range(2,197) if b.active.cell(r,3).value is not None]
    assert len(assignments)==3 and all(r['value']==0 and r['type']=='n' and isinstance(r['value'],(int,float)) for r in assignments)
    assert {r['mark'] for r in assignments}=={r['mark'] for r in selected}
finally:a.close();b.close()
stored=load(E/'stored_analysis.json');out=Path(load(I/'EVIDENCE_MANIFEST.json')['canonical_output'])
act_path=out/'ТЕСТ АКЗ №0.xlsx';app_path=out/'ТЕСТ Приложение к АКЗ №0.xlsx'
act=stored['workbooks'][str(act_path)]['sheets']['4']['cells']
application=stored['workbooks'][str(app_path)]['sheets']['Приложение АКЗ']
after=load(E/'native_outputs_recalculated.json')['books']
act_native=next(b for b in after if Path(b['path']).name==act_path.name)['sheets'][0]
app_native=next(b for b in after if Path(b['path']).name==app_path.name)['sheets'][0]
factual=('G11','B13','E13','I13','E15','E19','E20','E22','E23','E24','E26','E27','E28',
         'E30','E31','E32','E33','E34','G37','E39','E40','E41','E42','D49')
from openpyxl.utils.cell import coordinate_to_tuple
factual_verified={}
for coord in factual:
    rn,col=coordinate_to_tuple(coord);native=act_native['values'][rn-1][col-1]
    assert act[coord]['value']==act[coord]['cached']==native
    assert native in ('?','НЕТ ДАННЫХ / ТЕСТ')
    factual_verified[coord]=native
formula_cells={cell:v for cell,v in act.items() if v['type']=='f'}
assert set(formula_cells)=={'G8','I8'}
assert all('RANDBETWEEN' not in str(v) for row in act_native['formulas'] for v in row)
assert all(application['cells'][f'A{n}']['value']==r['mark'] for n,r in enumerate(selected,4))
blank_rows=[r for r in range(4,21) if all(app_native['values'][r-1][col] is None for col in range(7))]
assert blank_rows==list(range(7,21)) and not any(r in app_native['hidden_rows'] for r in blank_rows)
assert [application['cells'][c]['value'] for c in ('C21','E21','G21')]==['=SUM(C4:C20)','=SUM(E4:E20)','=SUM(G4:G20)']
for coord,expected in [('C21',3),('E21',105.26),('G21',3.989)]:
    rn,col=coordinate_to_tuple(coord)
    assert abs(app_native['values'][rn-1][col-1]-expected)<1e-8
source_sets={}
for name,p in [('audit',ROOT/'validation/stage6_akz_audit_20261008_135318/source_manifest_before.json'),
               ('final_check',ROOT/'validation/stage6_akz_final_check_20261008_170548/source_manifest_before.json')]:
    old=load(p)
    expected={r['path']:r['sha256'] for r in old} if isinstance(old,list) else {path:v['sha256'] for path,v in old.items()}
    actual={p:sha(p) for p in expected}
    source_sets[name]={'count':len(expected),'current':actual,'mismatches':[p for p,h in expected.items() if h!=actual[p]]}
M=ROOT/'validation/stage3_catalog_schema_20261008_162454'
migration=load(M/'product_hashes_CURRENT.json')
external=load(M/'external_artifacts.json')
print('MIGRATION_EXTERNAL_FORMAT',type(external).__name__)
resolution={}
for path in source_sets['audit']['mismatches']:
    if path in migration:
        assert sha(path)==migration[path]
        resolution[path]={'cause':'accepted Stage3 three-column compatibility migration before Stage6',
                          'current_sha256':sha(path),'evidence':str(M/'product_hashes_CURRENT.json')}
    else:
        assert Path(path)==original and sha(path)=='7fa36f9ecc2ebb722aff89fdb94ffedf330608dc73c16264005e31e11f4b1587'
        assert sha(path) in json.dumps(external)
        resolution[path]={'cause':'accepted Stage3 real catalog schema migration before Stage6',
                          'current_sha256':sha(path),'evidence':str(M/'external_artifacts.json')}
source_sets['audit']['explained_pre_implementation_changes']=resolution
assert len(resolution)==len(source_sets['audit']['mismatches'])==4
save('accepted_historical_hashes.json',source_sets)
save('reconciliation.json',{'vom':{'path':data['path'],'sha256':sha(data['path']),'total_data_rows':len(all_rows),
    'visible_count':len(visible),'hidden_count':len(all_rows)-len(visible),'visible_row_numbers':[n for n,r in visible],
    'filter':data['filter'],'product_matches_all_visible_rows':True},'selected':selected,'totals':totals,
    'catalog':{'original':str(original),'fixture':str(fixture),'count':195,'real_blank':195,'fixture_blank':192,
               'assignments':assignments,'marks_and_profiles_identical':True},
    'factual_cells_stored_and_recalculated':factual_verified,'remaining_act_formulas':formula_cells,
    'application_blank_visible_rows':blank_rows,'historical_hash_summary':{n:{'count':v['count'],'mismatches':v['mismatches']} for n,v in source_sets.items()}})
print(json.dumps({'selected':selected,'totals':totals,'catalog_assignments':assignments,'visible':len(visible),
    'hidden':len(all_rows)-len(visible),'blank_printed_rows':blank_rows,
    'accepted_baseline_checks':{n:{'count':v['count'],'mismatches':v['mismatches']} for n,v in source_sets.items()}},ensure_ascii=False,default=str,indent=2))
