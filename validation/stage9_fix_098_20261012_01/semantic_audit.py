from common import *
from zipfile import ZipFile
import io,olefile
from collections import Counter
from openpyxl import load_workbook
e=read(OUT/'EXE_E2E_FINAL.json');detail=read(e['evidence']);bundles=[];all_files={}
for p in e['manifests']:
    m=read(p);groups={str(Path(r['path']).resolve()).casefold():r for r in m['records']}
    assert len(m['records'])==7 and len(groups)==6
    assert Counter(Path(r['path']).suffix.lower() for r in groups.values())=={'.docx':3,'.xlsx':3}
    assert m['quality_generated_last'] and m['records'][-1]['type']=='QUALITY'
    assert (m['visible_rows'],m['quantity'],m['mass_kg'],m['area_m2'],m['uzk_units'],m['akz_variants'])==(104,'288','15822.14','1913.494',33,[0])
    vik=[r for r in m['records'] if r['type']=='VIK'];assert vik[0]['path']==vik[1]['path'] and [r['number'] for r in vik]==[m['passport'],m['passport']+'-1']
    for r in m['records']:assert r['official_confirmed'] is False and sha(r['path'])==r['sha256']
    for k,r in groups.items():assert k not in all_files;all_files[k]=r['sha256']
    bundles.append({'manifest':p,'manifest_sha256':sha(p),'passport':m['passport'],'physical_files':len(groups),'logical_records':len(m['records']),'numbers_dates':[{k:r[k] for k in ('type','number','date')} for r in m['records']],'native_quality':m['native_quality'],'quality_validation':m['quality_validation']})
closures=[]
for state,key in [('closed','fixture_closed'),('restored','fixture_restored')]:
    sb=Path(detail[key]['sandbox']);meta=read(sb/'SANDBOX.json');reg=sb/'registry/УЗК_TEST.xlsm'
    b=load_workbook(reg,read_only=True,data_only=True)
    try:stocks=[b['УЗК'].cell(i,3).value for i in (24,25)]
    finally:b.close()
    with ZipFile(reg) as z:
        with olefile.OleFileIO(io.BytesIO(z.read('xl/vbaProject.bin'))) as ole:
            streams={'/'.join(n):hashlib.sha256(ole.openstream(n).read()).hexdigest() for n in ole.listdir()}
    assert streams==meta['native_seed']['vba_streams'] and len(streams)==11
    assert stocks==([9,6] if state=='closed' else [12,10])
    if state=='restored':assert sha(reg)==meta['initial_registry_sha256']
    journal=[read(f) for f in sorted((sb/'journal').glob('*.json'))]
    closures.append({'state':state,'sandbox':str(sb),'registry_sha256':sha(reg),'seed_sha256':meta['initial_registry_sha256'],'exact_seed_sha_restored':sha(reg)==meta['initial_registry_sha256'],'stocks':stocks,'all_vba_streams_equal':True,'stream_count':len(streams),'journal':journal})
save(OUT/'INDEPENDENT_DOCUMENT_AND_XLSM_AUDIT.json',{'status':'PASS','exe_sha256':sha(EXE),'frozen_CLI_tasks':len(detail['steps']),'bundles':bundles,'physical_files':len(all_files),'logical_records':sum(x['logical_records'] for x in bundles),'closures':closures})
print('Independent output audit: 18 files / 21 records; exact XLSM restore SHA; all 11 OLE streams PASS',flush=True)
