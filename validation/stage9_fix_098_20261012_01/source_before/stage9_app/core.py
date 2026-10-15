"""Stage 9 orchestration. Accepted algorithms operate on explicit isolated inputs."""
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import shutil
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from .workspace import bind, load, save, PROJECT
from stage8_closure import storage as S

def settings(root, *, registry=None, documents=None):
    from passport_app.config import Settings
    return Settings(documents or root/'documents',registry or root/'DEV/ВЗМК_DEV.xlsx',
        {PROJECT:{'passport_root':'project/Паспорт','template':'project/Паспорт/1_2026 - Приложение № 1.xlsx'}})

def projects(root):
    root,_=load(root)
    from passport_app.registry import read_registry
    return list(read_registry(root/'DEV/ВЗМК_DEV.xlsx').projects)

def passports(root,project):
    root,_=load(root)
    from passport_app.registry import read_registry,select_project
    r=read_registry(root/'DEV/ВЗМК_DEV.xlsx');p=select_project(r,project)
    return [e['number'] for e in r.projects[project]],p.next_number

def fields(registry,project,number):
    from passport_app.registry import read_registry,select_project
    r=read_registry(registry);select_project(r,project)
    entries=[e for e in r.projects[project] if e['number']==number]
    if len(entries)!=1:raise S.Stop('Паспорт не найден однозначно.')
    book=load_workbook(registry,read_only=True,data_only=True,keep_links=False)
    try: values=[book[r.sheet].cell(entries[0]['row'],c).value for c in range(2,17)]
    finally:book.close()
    keys=('number','contract','customer','object','cipher','mass','start_date','end_date','km','kmd_company','steel','vik_date','uzk_date','akz_date','connections')
    f=dict(zip(keys,values))|{'project':project,'row':entries[0]['row']}
    for k in ('start_date','end_date','vik_date','uzk_date','akz_date'):
        if not isinstance(f[k],(date,datetime)):raise S.Stop('Не заполнена дата: '+k)
        f[k]=f[k].strftime('%d.%m.%Y')
    for k in ('contract','customer','object','cipher','km','kmd_company','steel','connections'):
        if not isinstance(f[k],str) or not f[k].strip():raise S.Stop('Не заполнено поле: '+k)
    return f

def normalized_vom(raw,path):
    """Materialize only visible native rows; retain raw values and row order."""
    from passport_app.vom import read_vom
    w=Workbook();s=w.active;s.title='ВОМ'
    s.append(['Отправочная марка','Наименование марки','Количество','Масса, кг',None,'Площадь, м2',None])
    s.append([None,None,None,'1 шт','Всех','1 шт','Всех'])
    for col in ('A','B','C'):s.merge_cells(f'{col}1:{col}2')
    s.merge_cells('D1:E1');s.merge_cells('F1:G1')
    for row in raw['rows']:
        s.append([row[k] for k in ('mark','name','quantity','mass_unit','mass_total','area_unit','area_total')])
    totals=[sum(Decimal(str(row[k])) for row in raw['rows']) for k in ('quantity','mass_total','area_total')]
    s.append(['Итого',None,float(totals[0]),None,float(totals[1]),None,float(totals[2])])
    w.save(path);w.close()
    v=read_vom(path)
    expected=[tuple([r['mark'],r['name'],*[Decimal(str(r[k])) for k in ('quantity','mass_unit','mass_total','area_unit','area_total')]]) for r in raw['rows']]
    if len(v.rows)!=len(expected) or any(a[:3]!=b[:3] or any(abs(x-y)>Decimal('0.00000001') for x,y in zip(a[3:],b[3:])) for a,b in zip(v.rows,expected)):
        raise S.Stop('Потеря данных при подготовке видимой ВОМ.')
    return v

def bindings(root):
    base=Path(__file__).resolve().parents[1]
    result={str(base/'stage9.py'):S.sha(base/'stage9.py'),str(base/'stage7_quality/render_emf.ps1'):S.sha(base/'stage7_quality/render_emf.ps1')}
    for package in ('stage9_app','passport_app','profiles_stage4','uzk_stage5','akz_stage6','stage7_quality','stage8_closure'):
        for p in (Path(__file__).resolve().parents[1]/package).glob('*.py'):result[str(p)]=S.sha(p)
    return result

def check(root,project,number,vom_path,variants=(0,),new=False):
    root,meta=bind(root)
    if project!=PROJECT:raise S.Stop('Для проекта требуется принятая привязка шаблонов и образца УЗК. Доступен Гараж Казахстан.')
    # Fail before creating any document when a first-passport form changed or vanished.
    for relative,ref in meta['inputs'].items():
        if relative.startswith('documents/project/Паспорт/') and (root/relative).suffix.lower() in ('.docx','.xlsx'):
            if S.sha(S.no_links(root/relative))!=ref['sha256']:raise S.Stop('TEMPLATE_OR_SAMPLE_SHA_CONFLICT: '+relative)
    if len(set(variants))!=len(variants) or any(type(n)!=int or n not in range(4) for n in variants):raise S.Stop('Некорректный выбор АКЗ.')
    source=S.no_links(vom_path)
    from akz_stage6 import core as a,part2
    from passport_app.registry import read_registry,select_project
    r=read_registry(root/'DEV/ВЗМК_DEV.xlsx');p=select_project(r,project)
    if new and number!=p.next_number:raise S.Stop('Номер нового паспорта изменился. Повторите выбор.')
    run=root/'validation'/('s9_'+uuid4().hex[:12])
    # Accepted Stage 2 uses this private pending name; native Word has a path
    # limit even when Python/NTFS can create the file. Stop before COM or writes.
    pending=run/'d/p'/(number.replace('/','-')+' от 00.00.0000')/('.vik-pending-'+'0'*32+'.docx')
    if len(str(pending).encode('utf-16-le'))//2>259:
        raise S.Stop('OFFICE_PATH_TOO_LONG: выберите DEV-стенд в более короткой папке. Временный путь Word превышает 259 символов.')
    run.mkdir()
    imported=run/('SOURCE_VOM'+source.suffix.lower())
    with S.frozen([source]):
        original=S.sha(source);S.exclusive_bytes(imported,source.read_bytes())
        if S.sha(imported)!=original:raise S.Stop('SHA ВОМ не совпадает.')
    # Stage 6 accepts only project/fixture VOMs. The immutable import is in fixtures.
    isolated=a.FIXTURES/(run.name+source.suffix.lower())
    S.exclusive_bytes(isolated,imported.read_bytes())
    raw=a.visible_vom(isolated)
    if raw['synthetic_source']:raise S.Stop('Искусственная ВОМ требует отдельного тестового сценария.')
    v=normalized_vom(raw,run/'VISIBLE_VOM.xlsx')
    from passport_app.service import validate
    stage1=validate(settings(root),project,v.path) if new else None
    f=None if new else fields(r.path,project,number)
    if f and abs(v.mass-Decimal(str(f['mass'])))>Decimal('.005'):raise S.Stop('Масса видимых строк ВОМ не совпадает с паспортом.')
    from uzk_stage5.inputs import registry_rows,mark_profiles,key
    from uzk_stage5.core import calculate_marks
    from profiles_stage4.storage import load_catalog
    stock,_=registry_rows(root/'DEV/УЗК_DEV.xlsm')
    required={key(row['mark']) for row in raw['rows'] if (key(project),key(row['mark'])) in stock}
    catalog=a.PROJECT/'Справочник марок.xlsx'
    mappings=mark_profiles(catalog,required_marks=required) if required else {}
    profiles=load_catalog(root/'DEV/Справочники/Профили.xlsx') if required else None
    rows=[{'row':r['row'],'mark':r['mark'],'quantity':r['quantity']} for r in raw['rows']]
    marks,excluded=calculate_marks(project,rows,stock,mappings,profiles)
    # Templates/assignments must pass before Stage 1 can append even a TEST row.
    # For a new passport, dates are recalculated by Stage 1 and previewed again.
    batch=part2.preview_groups(project,p.last_number if new else number,isolated,a.FIXTURES/'assignments.xlsx')
    for n in variants:part2.group_preview(batch,n)
    files=[r.path,source,imported,isolated,v.path,catalog,root/'DEV/УЗК_DEV.xlsm',
           *list((root/'DEV/Справочники').glob('*')),a.FIXTURES/'assignments.xlsx',
           *list((a.PROJECT/'Паспорт').glob('*.xlsx')),*list((a.PROJECT/'Паспорт').glob('*.docx'))]
    from uzk_stage5.result_samples import ACCEPTED_SOURCES
    files += [a.PROJECT/p for p,_,_ in ACCEPTED_SOURCES]
    hashes={str(S.no_links(p)):S.sha(p) for p in files if p.is_file()}|bindings(root)
    data={'schema':'STAGE9_PREVIEW_V1','root':str(root),'run':str(run),'project':project,'number':number,
          'new':new,'variants':list(variants),'source_vom':str(source),'source_sha256':original,
          'quantity':str(v.quantity),'mass_kg':str(v.mass),'area_m2':str(v.area),'visible_rows':len(v.rows),
          'hidden_excluded':len(raw['hidden_rows']),'uzk_units':sum(m['Q'] for m in marks),
          'fingerprints':hashes,'official_issue':False}
    plan={'data':data,'seal':S.digest(data),'vom':v,'raw':raw,'fields':f,'marks':marks,'batch':batch,'stage1':stage1}
    save(run/'PREVIEW.json',data|{'seal':plan['seal']})
    return plan

def verify(plan):
    d=plan['data']
    if d.get('official_issue') is not False or plan['seal']!=S.digest(d):raise S.Stop('Предварительный просмотр изменён.')
    root,_=load(d['root']);run=S.no_links(d['run'])
    if run.parent!=root/'validation':raise S.Stop('Выход вне TEST-стенда.')
    for p,h in d['fingerprints'].items():
        if S.sha(S.no_links(p))!=h:raise S.Stop('SHA_CONFLICT: '+p)

def describe(plan):
    d=plan['data']
    return (f"ТЕСТ — официальная выдача запрещена\nПроект: {d['project']}\nПаспорт: {d['number']}\n"
      f"Видимых строк: {d['visible_rows']}; скрытых исключено: {d['hidden_excluded']}\n"
      f"Количество: {d['quantity']} шт.\nМасса: {d['mass_kg']} кг\nПлощадь: {d['area_m2']} м²\n"
      f"УЗК: {d['uzk_units']} изделий; образец, контроль не подтверждён\nАКЗ: {', '.join(map(str,d['variants'])) or 'не выбраны'}\n"
      'Приложение №1 → ВИК → УЗК → АКЗ → Документ о качестве (последним).\n'
      'Измерения, подписи и подтверждения контроля не создаются.')

def generate(plan,confirmed_test=False):
    if confirmed_test is not True:raise S.Stop('Подтвердите формирование TEST-комплекта.')
    verify(plan);d=plan['data'];root,_=bind(d['root']);run=Path(d['run'])
    if (run/'START.json').exists():raise S.Stop('Запуск уже начат. Повторная запись запрещена.')
    save(run/'START.json',{'preview_seal':plan['seal'],'official_issue':False,'started_at':S.now()})
    try:
        if d['new']:
            from passport_app.service import generate as create
            result=create(settings(root),plan['stage1'])
            save(run/'STAGE1_RESULT.json',result)
            follow=check(root,d['project'],result['number'],d['source_vom'],d['variants'])
            return generate(follow,True)
        return _generate_existing(plan,root,run)
    except Exception as exc:
        save(run/'FAILURE.json',{'status':'FAIL','error':str(exc),'official_issue':False})
        raise

def _generate_existing(plan,root,run):
    d=plan['data'];f=plan['fields'];v=plan['vom']
    docs=run/'d';pp=docs/'p';pp.mkdir(parents=True)
    first=root/'documents/project/Паспорт'
    for name in ('1_2026 - Приложение № 1.xlsx','1_2026 - ВИК.docx'):
        shutil.copyfile(first/name,pp/name)
    registry=run/'registry.xlsx';shutil.copyfile(root/'DEV/ВЗМК_DEV.xlsx',registry)
    folder=pp/(f"{d['number'].replace('/','-')} от {f['end_date']}");folder.mkdir()
    appendix=folder/(d['number'].replace('/','_')+' - Приложение № 1.xlsx')
    shutil.copyfile(pp/'1_2026 - Приложение № 1.xlsx',appendix)
    from passport_app.registry import read_registry,select_project
    from passport_app.appendix import generate_appendix
    from stage8_closure.office import excel_session
    from passport_app.config import Settings
    cfg=Settings(docs,registry,{PROJECT:{'passport_root':'p','template':'p/1_2026 - Приложение № 1.xlsx'}})
    project=select_project(read_registry(registry),d['project'])
    with excel_session(diagnostics=run/'office') as excel:
        generate_appendix(excel,appendix,project,v,{'number':d['number'],'date':datetime.strptime(f['end_date'],'%d.%m.%Y').date(),'object':f['object']})
    from passport_app.vik import generate_vik
    vik=generate_vik(cfg,d['project'],d['number'],test_root=run)
    records=[]
    def record(kind,path,number=None,day=None,**extra):
        records.append({'type':kind,'path':str(path),'sha256':S.sha(path),'number':number or d['number'],
            'date':day or f['end_date'],'status':'SAMPLE_ONLY' if kind=='UZK' else 'TEST','official_confirmed':False,
            'usage':['basis_point13','attachment_point14'] if kind in ('VIK','UZK','AKZ') else [],**extra})
    record('QUALITY_APPENDIX_1',appendix)
    record('VIK',vik['output'],day=f['vik_date']);record('VIK',vik['output'],number=d['number']+'-1',day=f['vik_date'])
    if plan['marks']:
        from uzk_stage5.word import write_test_document
        uzk=folder/(d['number'].replace('/','_')+' - УЗК ТЕСТ.docx')
        data={'kind':'TEST/DRAFT','project_root':str(first.parent),'template':str(first/'1_2026 - УЗК.docx'),
            'passport':{'project':d['project'],'number':d['number'],'date':f['uzk_date'],'object':f['object'],'cipher':f['cipher']},'marks':plan['marks']}
        write_test_document(data,uzk);record('UZK',uzk,day=f['uzk_date'])
    from akz_stage6 import part2
    if d['variants']:
        pairs=part2.generate_groups(plan['batch'],d['variants'],run.name)
        if pairs['status']!='COMPLETE_ALL_SELECTED':raise S.Stop('Неполный АКЗ-комплект: '+str(pairs))
        save(run/'AKZ_RESULTS.json',pairs)
        for n,pair in pairs['groups'].items():
            target=folder/('АКЗ'+n);shutil.copytree(pair['path'],target)
            files=list(target.glob('*.xlsx'))
            acts=[p for p in files if 'Приложение' not in p.name];apps=[p for p in files if 'Приложение' in p.name]
            if len(acts)!=1 or len(apps)!=1:raise S.Stop('Неполная пара АКЗ.')
            day=f['akz_date'];num=d['number']+(('-'+n) if n!='0' else '')
            record('AKZ',acts[0],number=num,day=day,variant=int(n));record('AKZ_APPENDIX',apps[0],number=num,day=day,variant=int(n))
    verify(plan)
    archived_inputs=dict(d['fingerprints'])
    archived_inputs.pop(str(root/'DEV/ВЗМК_DEV.xlsx'))
    archived_inputs[str(registry)]=S.sha(registry)
    manifest={'schema':'STAGE9_TEST_MANIFEST_V1','mode':'DEV_TEST_ONLY','official_issue':False,'assembly_run_id':run.name,
        'project':d['project'],'passport':d['number'],'fields':f,'records':records,'inputs':archived_inputs,
        'registry_source_at_run':{'path':str(root/'DEV/ВЗМК_DEV.xlsx'),'sha256':d['fingerprints'][str(root/'DEV/ВЗМК_DEV.xlsx')]},
        'visible_rows':d['visible_rows'],'quantity':d['quantity'],'mass_kg':d['mass_kg'],'area_m2':d['area_m2'],
        'uzk_units':d['uzk_units'],'akz_variants':d['variants'],
        'warnings':['УЗК SAMPLE_ONLY; результаты не подтверждены.','АКЗ TEST; измерения отсутствуют.','Официальное закрытие запрещено.'],
        'quality_generated_last':True}
    from stage7_quality.word import write_quality,validate_quality
    quality=folder/(d['number'].replace('/','_')+' - Документ о качестве ТЕСТ.docx')
    write_quality(manifest,quality);validation=validate_quality(quality,manifest)
    from stage7_quality.native import inspect
    native=inspect(quality,run,kind='QUALITY',required=[d['number'],f['end_date']])
    record('QUALITY',quality)
    manifest['records']=records;manifest['quality_validation']=validation;manifest['native_quality']=native
    manifest['state']='COMPLETE_TEST';manifest['seal']=S.digest(manifest)
    path=run/'MANIFEST.json';save(path,manifest)
    verify_manifest(path)
    save(run/'RESULT.json',{'status':'PASS','manifest':str(path),'folder':str(folder),'official_issue':False})
    return path

def verify_manifest(path):
    m=S.read(path);payload={k:v for k,v in m.items() if k!='seal'}
    if m.get('schema')!='STAGE9_TEST_MANIFEST_V1' or m.get('seal')!=S.digest(payload) or m.get('official_issue') is not False or m.get('state')!='COMPLETE_TEST':
        raise S.Stop('Неполный или изменённый комплект.')
    if m['records'][-1]['type']!='QUALITY':raise S.Stop('Документ о качестве должен быть последним.')
    kinds=[r['type'] for r in m['records']]
    if any(k not in kinds for k in ('QUALITY_APPENDIX_1','VIK','QUALITY')):raise S.Stop('Отсутствует обязательный документ.')
    if kinds.count('QUALITY')!=1 or kinds.count('QUALITY_APPENDIX_1')!=1 or kinds.count('VIK')!=2 or kinds.count('UZK')!=(1 if m['uzk_units'] else 0):
        raise S.Stop('Неполный или повторяющийся состав комплекта.')
    if [r['variant'] for r in m['records'] if r['type']=='AKZ']!=m['akz_variants'] or [r['variant'] for r in m['records'] if r['type']=='AKZ_APPENDIX']!=m['akz_variants']:
        raise S.Stop('Неполная или смешанная пара АКЗ.')
    for r in m['records']:
        if r.get('official_confirmed') is not False or r['status'] not in ('TEST','SAMPLE_ONLY'):raise S.Stop('TEST/OFFICIAL_CONFLICT')
    for r in m['records']:
        if S.sha(S.no_links(r['path']))!=r['sha256']:raise S.Stop('Изменён документ: '+r['path'])
    for p,h in m['inputs'].items():
        if S.sha(S.no_links(p))!=h:raise S.Stop('Изменён источник: '+p)
    return m

def document_groups(manifest):
    """One physical file can carry several accepted logical acts (VIK)."""
    groups={}
    for record in manifest['records']:
        path=S.no_links(record['path']);key=str(path).casefold()
        group=groups.setdefault(key,{'path':path,'records':[]})
        group['records'].append(record)
    return list(groups.values())

def document_counts(manifest):
    groups=document_groups(manifest)
    return {'documents':len(groups),'physical_files':len(groups),'document_records':len(manifest['records']),
            'word_documents':sum(g['path'].suffix.lower()=='.docx' for g in groups),
            'excel_documents':sum(g['path'].suffix.lower()=='.xlsx' for g in groups),
            'word_records':sum(Path(r['path']).suffix.lower()=='.docx' for r in manifest['records']),
            'excel_records':sum(Path(r['path']).suffix.lower()=='.xlsx' for r in manifest['records'])}

def official_close(*args,**kwargs):
    raise S.Stop('Официальное закрытие запрещено: DEV/TEST, УЗК/АКЗ не подтверждены.')
