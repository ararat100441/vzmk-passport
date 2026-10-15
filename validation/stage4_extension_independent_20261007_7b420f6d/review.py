"""Independent bounded Stage 4 extension review; all writes are isolated evidence."""
from pathlib import Path
from collections import Counter,defaultdict
from copy import deepcopy
from decimal import Decimal
from unittest.mock import patch
from xml.etree import ElementTree as ET
import hashlib,json,os,re,shutil,sys,traceback,zipfile
from datetime import datetime,timezone

OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[1]
FIX=ROOT/'validation/stage4_targeted_extension_20261007_172100'
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(OUT));sys.dont_write_bytecode=True
sys.stdout.reconfigure(encoding='utf-8')
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(n,x):(OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def require(x,m='assertion failed'):
    if not x:raise AssertionError(m)
network=[]
def audit(event,args):
    if event in ['socket.connect','socket.getaddrinfo','socket.sendto']:
        network.append(event);raise AssertionError('No network permitted')
    if event=='open' and isinstance(args[0],(str,bytes,os.PathLike)):
        p=Path(os.fsdecode(args[0])).resolve()
        require(p.name!='ВЗМК.xlsx','Production registry access prohibited')
        mode,flags=args[1:3]
        writes=(isinstance(mode,str) and any(c in mode for c in 'wax+')) or flags&(os.O_WRONLY|os.O_RDWR|os.O_CREAT|os.O_TRUNC)
        if p.is_relative_to(ROOT):require(not writes or p.is_relative_to(OUT),'Write outside isolated review')
sys.addaudithook(audit)
manifest=read(FIX/'REVIEW_MANIFEST.json')
expected={str(ROOT/p):h for p,h in manifest['product_hashes'].items()}
expected.update(manifest['protected_hashes']);expected.update(manifest['DEV_artifacts'])
expected.update({str(FIX/p):v['sha256'] for p,v in manifest['evidence_files'].items()})
expected[str(FIX/'REVIEW_MANIFEST.json')]=sha(FIX/'REVIEW_MANIFEST.json')
require(not any(Path(p).name=='ВЗМК.xlsx' for p in expected),'Manifest contains prohibited production path')
before={p:sha(p) for p in expected};save('preservation_before.json',before)
tests=[]
def case(name,fn):
    try:
        detail=fn();tests.append({'name':name,'status':'PASS','detail':detail});print('PASS '+name,flush=True)
    except BaseException:
        tests.append({'name':name,'status':'FAIL','traceback':traceback.format_exc()});print(traceback.format_exc(),flush=True)
    save('results.json',{'tests':tests,'passed':sum(t['status']=='PASS' for t in tests),'failed':sum(t['status']=='FAIL' for t in tests)})
def binding():
    require(expected==before,'Current artifact/product/evidence mismatch')
    a,b,c=(read(FIX/f'product_hashes_{x}.json') for x in ['before','after','current'])
    require(a==b==c==manifest['product_hashes'])
    prior=read(FIX/'baseline_product.json')
    changed=[p for p,h in prior.items() if sha(ROOT/p)!=h]
    require(set(p.replace('\\','/') for p in changed)=={'profiles_stage4/model.py','profiles_stage4/storage.py','profiles_stage4/workbook.mjs','STAGE4_README.md'},str(changed))
    require(sha(ROOT/'profiles_stage4/windows_generation.py')==prior['profiles_stage4\\windows_generation.py'])
    decision=read(FIX/'user_confirmed_decision.json')
    require(sha(decision['source'])==decision['sha256'])
    source=Path(decision['source']).read_text(encoding='utf-8')
    require(all(x in source for x in ['536 / 527','деталь 718','Гн.[140X60X4','Гн.[]120х3','16П + ГОСТ 8240-97']))
    t=read(FIX/'targeted_tests.json');require(t['status']=='PASS' and t['passed']==33 and t['failed']==0)
    last=(FIX/'REVIEW_MANIFEST.json').stat().st_mtime_ns
    require(all((FIX/p).stat().st_mtime_ns<=last for p in manifest['evidence_files']))
    for suffix in ['DEV/ВЗМК_DEV.xlsx','DEV/DEV_REGISTRY_CHECKPOINT.json','Справочник марок.xlsx','УЗК_DEV.xlsm']:
        require(any(str(p).replace('\\','/').endswith(suffix) for p in expected),'Missing protected '+suffix)
    return {'current_hashes_verified':len(manifest['product_hashes'])+len(manifest['DEV_artifacts']),
            'preserved_files':len(expected),'protected_files':len(manifest['protected_hashes']),'evidence_files':len(manifest['evidence_files']),
            'actual_changed_product_set':changed,'Windows_generation_unchanged':True,'test_binding_BEFORE_AFTER_CURRENT':True,'manifest_last':True,
            'KMD_provenance_basis':'Exact hashed user-confirmed source contract; no KMD re-audit'}
case('exact baseline product sidecars evidence and changed set',binding)
require(tests[-1]['status']=='PASS','STOP baseline mismatch')

import profiles_stage4.storage as s
from profiles_stage4.model import Catalog,ProfileError,calculate_uzk,frequency_angle,normalize_profile,RULE,BENT_RULE,CLOSED_RULE,CHANNEL_RULE,NOT_CONFIGURED,SUPPORTED_RULES
from openpyxl import load_workbook
from native_probe import run as native
CAT=s.load_catalog();STATE=CAT.user_state;CURRENT=s.DEFAULT_WORKBOOK
OLD=read(FIX/'baseline_generation'/s.CATALOG)
NAMES=[CURRENT.name,s.STATE,s.CATALOG,s.COMMIT]
def pairs(r):return [[x['L'],x['S']] for x in calculate_uzk(r)['rows']]
def rejected(fn):
    try:fn()
    except (OSError,ValueError) as e:return str(e)
    raise AssertionError('Expected fail-closed rejection')

def manual_and_counts():
    require(CAT.raw_records==OLD['records'] and CAT.sources==OLD['sources'])
    oldkeys={(r['key'],r['standard']) for r in OLD['records']}
    keys={(r['key'],r['standard']) for r in CAT.raw_records+STATE['manual_rows']}
    require(len(oldkeys)==2849 and len(keys)==len(CAT.rows)==2851 and len(CAT.raw_records)==3037)
    require(len(STATE['manual_rows'])==2 and len(keys-oldkeys)==2)
    expectations=[('Гн.[140X60X4','Гнутый швеллер','ГОСТ 8278-83',{'h':140,'b':60,'t':4},BENT_RULE,[413,414],[536,527]),
                  ('Гн.[]120х3','Квадратный профиль','ГОСТ 30245-2003',{'h':120,'b':120,'t':3},CLOSED_RULE,[458],[718])]
    details=[]
    for raw,family,nd,geometry,rule,pages,parts in expectations:
        row=CAT.lookup(raw,nd)['profile'];m=next(m for m in STATE['manual_rows'] if m['display']==raw)
        require(row['status']=='РУЧНОЙ' and row['family']==family and row['geometry']==geometry and row['uzk_rule']==rule)
        require(m['raw']['profile']==raw and m['raw']['geometry']==geometry and m['key']==normalize_profile(raw))
        require(m['verification']['kind']=='KMD' and m['verification']['confirmed_by']=='USER')
        require(m['verification']['references'][0]['pages']==pages and m['verification']['references'][0]['parts']==parts)
        require(all(Path(r['file']).is_file() for r in m['verification']['references']))
        require(not any(f in row['geometry'] for f in ['R','r']) and row['provenance']==[m])
        details.append({'raw':raw,'key':row['key'],'standard':nd,'geometry':geometry,'rule':rule,'KMD':m['verification']})
    assignment=Counter(r['family'] for r in CAT.rows if r['uzk_rule'] in SUPPORTED_RULES)
    require(assignment=={'Двутавр':480,'Квадратный профиль':124,'Прямоугольный профиль':369,'Швеллер':1,'Гнутый швеллер':1})
    require(sum(assignment.values())==975)
    status=Counter(r['status'] for r in CAT.rows)
    require(status=={'ИЗ ИСТОЧНИКА':2662,'ПРОВЕРЕНО_ДВУМЯ_ИСТОЧНИКАМИ':172,'СОГЛАСОВАНО_С_ЗАМЕЧАНИЯМИ':7,'КОНФЛИКТ':8,'РУЧНОЙ':2})
    return {'previous_canonical':2849,'canonical':len(keys),'source_raw':3037,'manual':2,'UZK_enabled':975,'UZK_by_family':dict(assignment),'manual_records':details}
case('independent counts manual geometry raw and KMD provenance',manual_and_counts)

def required():
    details=[]
    for raw,nd,geometry,want in [('[16П','ГОСТ 8240-97',{'h':160,'b':64,'s':5,'t':8.4},[[160,5],[128,8.4]]),
        ('Гн.[140X60X4','ГОСТ 8278-83',{'h':140,'b':60,'t':4},[[400,4]]),('Гн.[]120х3','ГОСТ 30245-2003',{'h':120,'b':120,'t':3},[[480,3]])]:
        r=CAT.lookup(raw);require(r['status']=='FOUND' and r['profile']['standard']==nd)
        require(all(r['profile']['geometry'][k]==v for k,v in geometry.items()) and pairs(r)==want)
        details.append({'raw':raw,'key':r['key'],'standard':nd,'geometry':r['profile']['geometry'],'UZK':pairs(r)})
    for raw,geometry,want in [('[16П',{'h':162,'b':65,'s':6,'t':9},[[162,6],[130,9]]),
        ('Гн.[140X60X4',{'h':144,'b':63,'t':5},[[414,5]]),('Гн.[]120х3',{'h':122,'b':122,'t':4},[[488,4]]),
        ('80×40×4',{'h':90,'b':50,'t':6},[[280,6]])]:
        r=deepcopy(CAT.lookup(raw));r['profile']['geometry'].update(geometry)
        require(pairs(r)==want,'UZK is not driven by actual geometry')
    return {'required_lookups':details,'four_independent_geometry_sensitivity_probes':'PASS'}
case('required current Catalog lookups and non-hardcoded formula probes',required)

def aliases():
    for sep in ['X','х','×','x','Х']:
        for raw,want in [('Гн.[140'+sep+'60'+sep+'4',BENT_RULE),('Гн.[]120'+sep+'3',CLOSED_RULE)]:
            r=CAT.lookup(raw);require(r['status']=='FOUND' and r['profile']['uzk_rule']==want)
    require(normalize_profile('Гн.[140×60×4')!=normalize_profile('Гн.[]140×60×4'))
    closed=CAT.lookup('Гн.[]140×60×4')
    require(closed['status'] not in ['FOUND','FOUND_PREFERRED'] or closed['profile']['family']=='Прямоугольный профиль')
    for raw in ['[[16П','[16П]','[35Б1','[140×60×4','[Гн.[140×60×4','[16П EXTRA']:
        require(CAT.lookup(raw)['status'] not in ['FOUND','FOUND_PREFERRED'],raw)
    require(CAT.lookup('[16П','ГОСТ 9999-99')['status']=='NOT_FOUND')
    channel=deepcopy(next(r for r in CAT.raw_records if r['key']=='16П'))
    other=deepcopy(channel);other.update(record_id='INDEP_OTHER_ND',standard='INDEPENDENT НД')
    require(Catalog(CAT.raw_records+[other],CAT.sources,STATE).lookup('[16П')['status']=='AMBIGUOUS')
    other=deepcopy(channel);other.update(record_id='INDEP_WRONG_FAMILY',family='Двутавр')
    require(Catalog([other],user_state=STATE).lookup('[16П')['status']=='AMBIGUOUS')
    require(Catalog(CAT.raw_records+[other],CAT.sources,STATE).lookup('[16П')['status']=='AMBIGUOUS')
    exact=deepcopy(channel);exact.update(key='[16П',display='[16П',record_id='INDEP_EXACT',geometry={**channel['geometry'],'h':161})
    r=Catalog(CAT.raw_records+[exact],CAT.sources,STATE).lookup('[16П')
    require(r['status']=='FOUND' and r['key']=='[16П' and r['profile']['geometry']['h']==161)
    return {'separator_variants':10,'open_closed_distinct':True,'arbitrary_brackets_preserved':True,'wrong_family_multiple_ND_collision':'AMBIGUOUS','exact_key_wins':True}
case('alias variants raw preservation type and collision safety',aliases)

def scopes():
    for row in CAT.rows:
        rule=row['uzk_rule']
        if rule==CHANNEL_RULE:require((row['key'],row['standard'],row['family'])==('16П','ГОСТ 8240-97','Швеллер'))
        if rule==BENT_RULE:require(row['key']==normalize_profile('Гн.[140X60X4') and row['standard']=='ГОСТ 8278-83')
        if rule==CLOSED_RULE:require(row['family'] in ['Квадратный профиль','Прямоугольный профиль'] and row['standard']=='ГОСТ 30245-2003')
        if rule in SUPPORTED_RULES:
            result=CAT.lookup(row['key'],row['standard']);calc=calculate_uzk(result)
            require(calc['status']=='CALCULATED')
            if rule in [CLOSED_RULE,BENT_RULE]:require(pairs(result)==[[2*(row['geometry']['h']+row['geometry']['b']),row['geometry']['t']]])
    require(CAT.lookup('20П','ГОСТ 8240-97')['profile']['uzk_rule']==NOT_CONFIGURED)
    state=deepcopy(STATE);state.pop('uzk_approvals')
    off=Catalog(CAT.raw_records,CAT.sources,state)
    for raw in ['[16П','Гн.[140X60X4','Гн.[]120х3']:require(off.lookup(raw)['profile']['uzk_rule']==NOT_CONFIGURED)
    manual=deepcopy(STATE['manual_rows'][0]);manual.update(key='Гн.[150×65×4',display='Гн.[150×65×4',record_id='INDEP_UNAPPROVED_BENT')
    state=deepcopy(STATE);state['manual_rows'].append(manual)
    require(Catalog(CAT.raw_records,CAT.sources,state).lookup(manual['key'])['profile']['uzk_rule']==NOT_CONFIGURED)
    for raw,family,nd,geometry in [('126×126×3','Квадратный профиль','ГОСТ 30245-2003',{'h':126,'b':126,'t':3}),
        ('INDEP ROUND','Круглая труба','ГОСТ 30245-2003',{'D':100,'s':4})]:
        state=s.prepare_manual_state(CAT,raw,family,nd,geometry)
        require(Catalog(CAT.raw_records,CAT.sources,state).lookup(raw,nd)['profile']['uzk_rule']==NOT_CONFIGURED)
    base=deepcopy(next(r for r in CAT.raw_records if r['key']=='80×40×4' and r['standard']=='ГОСТ 30245-2003'))
    base['record_id']='INDEP_CORE_CONFLICT';base['geometry']['t']=3
    r=Catalog(CAT.raw_records+[base],CAT.sources,STATE).lookup(base['key'],base['standard'])
    require(r['status']=='CONFLICT' and r['profile']['uzk_rule']==NOT_CONFIGURED);rejected(lambda:calculate_uzk(r))
    return {'all_975_assignments_checked':True,'only_one_16П':True,'unapproved_bent_disabled':True,'unverified_closed_manual_disabled':True,'other_family_manual_disabled':True,'core_conflict_blocked':True}
case('every rule assignment and unapproved manual/conflict safety',scopes)

def regression():
    previous=Catalog(OLD['records'],OLD['sources'],OLD['user_state'])
    prev={(r['key'],r['standard']):r for r in previous.rows}
    for row in CAT.rows:
        if (row['key'],row['standard']) in prev:
            a,b=deepcopy(row),deepcopy(prev[row['key'],row['standard']]);a.pop('uzk_rule');b.pop('uzk_rule');require(a==b)
    require(CAT.conflicts==previous.conflicts)
    require(Counter(c['classification'] for c in CAT.conflicts)=={'BLOCKING':8,'NON_BLOCKING':7})
    for c in CAT.conflicts:
        r=CAT.lookup(c['key'],c['standard']);require(r['status']==('CONFLICT' if c['classification']=='BLOCKING' else 'FOUND'))
        if c['classification']=='BLOCKING':rejected(lambda:calculate_uzk(r))
    for nd,want in [('ГОСТ Р 57837-2017',[[328,6],[348,9]]),('ГОСТ 26020-83',[[329,6.2],[310,8.5]])]:require(pairs(CAT.lookup('35Б1',nd))==want)
    require(CAT.lookup('35Б1')['status']=='FOUND_PREFERRED' and CAT.lookup('20Ш1')['status']=='AMBIGUOUS')
    require(frequency_angle(11.5)=={'frequency_mhz':5,'angle_degrees':70} and frequency_angle(12)=={'frequency_mhz':2.5,'angle_degrees':65})
    return {'all_2849_previous_geometry_status_provenance_unchanged':True,'M1':[8,7],'golden':'PASS','frequency_boundary':'PASS'}
case('minimal accepted geometry M1 golden and frequency regression',regression)

def preview():
    marks=[]
    for mark,raw,q,want in [('Бп2-1','[16П',8,[[1280,5],[1024,8.4]]),('Пс1-9','Гн.[140X60X4',9,[[3600,4]]),('Пс1-10','Гн.[140X60X4',15,[[6000,4]]),('Пс2-16','Гн.[]120х3',1,[[480,3]])]:
        single=calculate_uzk(CAT.lookup(raw));total=[[int(Decimal(str(r['L']))*q),r['S']] for r in single['rows']]
        require(total==want and all(r['frequency_mhz']==5 and r['angle_degrees']==70 for r in single['rows']))
        marks.append({'mark':mark,'quantity':q,'raw_profile':raw,'single':pairs(CAT.lookup(raw)),'total':total,'MHz':5,'angle':70})
    require(sum(m['quantity'] for m in marks)==33 and len(marks)==4)
    save('four_mark_preview.json',marks);return {'marks':marks,'items':33,'different_marks_kept_separate':True,'S_never_multiplied':True}
case('four-mark independent arithmetic data preview',preview)

def workbook():
    wb=load_workbook(CURRENT,read_only=False,data_only=False)
    try:
        require(wb.sheetnames==list(s.HEADERS) and wb['Профили'].max_row==2852 and wb['Профили'].max_column==20)
        require(not any(c.data_type=='f' for ws in wb for row in ws for c in row))
        for ws in wb:
            require(len(ws.tables)==1 and ws.freeze_panes)
            for t in ws.tables.values():
                names=[c.name for c in t.tableColumns]
                require(len(names)==len(set(n.casefold() for n in names)) and all(names))
                require(names==[c.value for c in ws[1]] and t.autoFilter.ref==t.ref)
        expected=s.workbook_payload(CAT)
        require(s.verify_workbook(CURRENT,expected)['formulas']==0)
        rules=list(wb['Правила УЗК'].iter_rows(min_row=2,values_only=True))
        require({r[0] for r in rules}==SUPPORTED_RULES|{NOT_CONFIGURED})
        for cells,row in zip(wb['Профили'].iter_rows(min_row=2),CAT.rows):
            require(cells[14].value==row['uzk_rule'])
            for col,field in [(4,'h'),(5,'b'),(6,'b2'),(7,'s'),(8,'t'),(9,'D'),(10,'r'),(15,'R')]:
                require(cells[col].value==row['geometry'].get(field) and (cells[col].value is None or type(cells[col].value) in [int,float]))
        return {'sheets':5,'tables':5,'profile_rows':2851,'columns':20,'formulas':0,'numeric_geometry':True,'case_insensitive_headers_unique':True,'user_rules_visible':True}
    finally:wb.close()
case('exact current workbook tables numeric geometry and approved rules',workbook)
case('native Excel exact current read-only alerts enabled no repair',lambda:native(CURRENT,'exact_current'))

def fixture(label):
    folder=OUT/'fixtures'/label/'DEV/Справочники';folder.mkdir(parents=True)
    for n in NAMES:shutil.copyfile(CURRENT.parent/n,folder/n)
    return folder/CURRENT.name
def clean(t):require(not list(t.parent.glob('.profiles-stage-*')) and not (t.parent/'.profiles.lock').exists())
def rebuild():
    t=fixture('rebuild');old={n:sha(t.parent/n) for n in NAMES}
    result=s.build(t);current=s.load_catalog(t)
    require(result['network_attempts']==0)
    require({n:sha(Path(result['backup'])/n) for n in NAMES}==old)
    require(current.user_state['revision']==STATE['revision']+1)
    for field in ['manual_rows','uzk_approvals','preferences','resolutions']:require(current.user_state[field]==STATE[field])
    require(current.rows==CAT.rows and current.raw_records==CAT.raw_records and current.counts()==CAT.counts())
    clean(t);save('isolated_rebuild.json',result)
    return {'real_offline_renderer':True,'manual_records_preserved':2,'rule_approvals_preserved':3,'preference_preserved':True,'coherent_backup':True,'all_rows_unchanged':True,'network_attempts':0}
case('isolated offline rebuild preserves user state and assignments',rebuild)

def race():
    t=fixture('race');old={n:sha(t.parent/n) for n in NAMES};edited={}
    def render(payload,dest,previews=None):
        shutil.copyfile(CURRENT,dest);state=read(t.parent/s.STATE);state['independent_external_marker']=True;s.dump(t.parent/s.STATE,state)
        edited['hash']=sha(t.parent/s.STATE);return {'network_attempts':0}
    with patch.object(s,'render_workbook',render):message=rejected(lambda:s.build(t))
    require(edited and sha(t.parent/s.STATE)==edited['hash']!=old[s.STATE])
    require(all(sha(t.parent/n)==old[n] for n in NAMES if n!=s.STATE) and not (t.parent/'backups').exists())
    clean(t);return {'STOP':message,'external_bytes_preserved':True,'new_partial_generation_absent':True}
case('one isolated M2 external-mutation commit-guard smoke',race)

def preserved():
    after={p:sha(p) for p in expected};save('preservation_after.json',after);require(before==after and not network)
    return {'BEFORE_EQUALS_AFTER_EQUALS_CURRENT':True,'files':len(after),'protected':len(manifest['protected_hashes']),
            'real_DEV_writes':0,'production_access':0,'network_attempts':0}
case('final protected state and evidence preservation',preserved)
failures=[t for t in tests if t['status']=='FAIL']
result={'verdict':'FAIL' if failures else 'PASS','passed':len(tests)-len(failures),'failed':len(failures),'tests':tests,
        'product_hashes':manifest['product_hashes'],'DEV_artifacts':manifest['DEV_artifacts'],
        'findings':[],'not_run':['Stage 5 implementation','UZK.docx generation','real DEV rebuild','production registry access','HTTP/web','User E2E','full Stage 4 review','KMD re-audit'],
        'ready_for_targeted_user_E2E':not failures}
save('INDEPENDENT_REVIEW.json',result)
evidence={str(p.relative_to(OUT)):sha(p) for p in OUT.rglob('*') if p.is_file() and p.name!='REVIEW_MANIFEST.json'}
save('REVIEW_MANIFEST.json',{'created_at':datetime.now(timezone.utc).isoformat(),'verdict':result['verdict'],'passed':result['passed'],'failed':result['failed'],
     'current_product_hashes':manifest['product_hashes'],'current_DEV_artifacts':manifest['DEV_artifacts'],'preserved_files':len(before),
     'BEFORE_AFTER_CURRENT_equal':before=={p:sha(p) for p in before},'evidence_files':evidence})
print(json.dumps({k:result[k] for k in ['verdict','passed','failed','ready_for_targeted_user_E2E']},ensure_ascii=False),flush=True)
