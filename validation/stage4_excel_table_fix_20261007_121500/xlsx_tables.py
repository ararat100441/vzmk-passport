"""Read-only deterministic Excel table package validator; no OOXML rewriting."""
from collections import Counter
from pathlib import Path
import posixpath
import re
from xml.etree import ElementTree as ET
import zipfile
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter

MAIN='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
NS={'s':MAIN,'p':PKG}

def target_part(part,target):
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part),target))

def inspect(path, extraction=None):
    wb=load_workbook(path,read_only=True,data_only=False)
    results=[]; issues=[]
    try:
        with zipfile.ZipFile(path) as package:
            assert package.testzip() is None
            names=package.namelist(); assert len(names)==len(set(names))
            dangling=[]; controls=[]
            for name in names:
                if name.endswith(('.xml','.rels')):
                    text=package.read(name).decode('utf-8-sig')
                    if re.search('[\x00-\x08\x0b\x0c\x0e-\x1f]',text): controls.append(name)
                    root=ET.fromstring(text)
                    if name.endswith('.rels'):
                        parent=posixpath.join(posixpath.dirname(posixpath.dirname(name)),posixpath.basename(name)[:-5])
                        for rel in root:
                            if rel.get('TargetMode')!='External':
                                resolved=target_part(parent,rel.get('Target'))
                                if resolved not in names: dangling.append({'relationship':name,'target':resolved})
            workbook=ET.fromstring(package.read('xl/workbook.xml'))
            workbook_rels={r.get('Id'):r for r in ET.fromstring(package.read('xl/_rels/workbook.xml.rels'))}
            for sheet in workbook.find('s:sheets',NS):
                worksheet=target_part('xl/workbook.xml',workbook_rels[sheet.get('{'+REL+'}id')].get('Target'))
                root=ET.fromstring(package.read(worksheet))
                cells=root.findall('.//s:sheetData/s:row/s:c',NS)
                addresses=[c.get('r') for c in cells]
                last_row=max(range_boundaries(a)[1] for a in addresses)
                last_col=max(range_boundaries(a)[0] for a in addresses)
                used=f'A1:{get_column_letter(last_col)}{last_row}'
                parts=root.find('s:tableParts',NS)
                assert parts is not None and int(parts.get('count'))==len(parts)==1
                rel_path=posixpath.join(posixpath.dirname(worksheet),'_rels',posixpath.basename(worksheet)+'.rels')
                worksheet_rels={r.get('Id'):r for r in ET.fromstring(package.read(rel_path))}
                for part in parts:
                    rid=part.get('{'+REL+'}id')
                    rel=worksheet_rels[rid]
                    assert rel.get('Type').endswith('/table')
                    table_part=target_part(worksheet,rel.get('Target'))
                    raw=package.read(table_part)
                    table=ET.fromstring(raw)
                    columns=table.find('s:tableColumns',NS)
                    column_names=[c.get('name') for c in columns]
                    ids=[int(c.get('id')) for c in columns]
                    duplicates=[n for n,c in Counter(n.casefold() for n in column_names).items() if c>1]
                    invalid=[n for n in column_names if not n or not n.strip() or len(n)>255]
                    ws=wb[sheet.get('name')]
                    if ws.max_row is None or ws.max_column is None: ws.calculate_dimension(force=True)
                    headers=list(next(ws.iter_rows(min_row=1,max_row=1,values_only=True)))
                    checks={'ref_matches_used_range':table.get('ref')==used,
                        'autoFilter_matches_ref':table.find('s:autoFilter',NS).get('ref')==table.get('ref'),
                        'column_count_matches':int(columns.get('count'))==len(columns)==last_col==len(headers),
                        'column_ids_positive_unique_sequential':ids==list(range(1,len(columns)+1)),
                        'column_names_case_insensitive_unique':not duplicates,'column_names_valid':not invalid,
                        'headers_match_table_columns':headers==column_names,
                        'header_totals_settings':table.get('headerRowCount','1')=='1' and table.get('totalsRowCount','0')=='0'}
                    result={'sheet':sheet.get('name'),'worksheet':worksheet,'part':table_part,'table_name':table.get('name'),
                        'ref':table.get('ref'),'autoFilter_ref':table.find('s:autoFilter',NS).get('ref'),
                        'used_range':used,'data_rows':last_row-1,'column_count':len(columns),'column_ids':ids,
                        'column_names':column_names,'case_insensitive_duplicates':duplicates,'invalid_names':invalid,
                        'relationship_id':rid,'tableParts_count':len(parts),'checks':checks}
                    results.append(result)
                    issues.extend({'part':table_part,'check':name} for name,ok in checks.items() if not ok)
                    if extraction:
                        extraction=Path(extraction); extraction.mkdir(parents=True,exist_ok=True)
                        (extraction/posixpath.basename(table_part)).write_bytes(raw)
            if dangling or controls: issues.append({'dangling_relationships':dangling,'invalid_control_characters':controls})
            return {'workbook':str(path),'zip_integrity':True,'xml_parse':True,'dangling_relationships':dangling,
                    'invalid_control_characters':controls,'tables':results,'issues':issues,'valid':not issues}
    finally:
        wb.close()
