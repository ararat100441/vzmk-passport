"""Excel table fix: structural + native Excel + bounded M1/M2/golden regressions."""
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
import io,json,shutil,subprocess,sys,traceback,uuid,zipfile
from xml.etree import ElementTree as ET
OUT=Path(__file__).resolve().parent; ROOT=OUT.parents[1]
sys.path.insert(0,str(ROOT)); sys.path.insert(0,str(OUT))
sys.stdout.reconfigure(encoding='utf-8')
import profiles_stage4.storage as s
from profiles_stage4.model import calculate_uzk
from profiles_stage4.offline import read_json,sha,no_network
from profiles_stage4.windows_generation import ProtectedFile
from native_excel import run as native
from xlsx_tables import inspect, MAIN
from openpyxl import load_workbook

FILES=['Профили.xlsx',s.STATE,s.CATALOG,s.COMMIT]
RUN=uuid.uuid4().hex[:8]
tests=[]; network=[]; production=[]
def audit(event,args):
    if event in ('socket.connect','socket.getaddrinfo','socket.sendto'):
        network.append(event); raise AssertionError('No network allowed')
    if event in ('open','os.rename','os.remove','os.mkdir'):
        for arg in args:
            if isinstance(arg,(str,bytes)) and any(p in str(arg).replace('\\','/') for p in ['/Работа/ВЗМК.xlsx','/Работа/Профили.xlsx']):
                production.append(str(arg)); raise AssertionError('No production access allowed')
sys.addaudithook(audit)
def hashes():
    paths=sorted(p for p in (ROOT/'profiles_stage4').glob('*') if p.is_file())
    paths += [ROOT/'stage4.py',ROOT/'STAGE4_README.md',ROOT/'Запустить Этап 4.cmd']+sorted(OUT.glob('*.py'))
    return {str(p.relative_to(ROOT)):sha(p) for p in paths}
def case(name,fn):
    try:
        detail=fn(); tests.append({'name':name,'status':'PASS','detail':detail})
        print('PASS '+name,flush=True)
    except BaseException:
        tests.append({'name':name,'status':'FAIL','traceback':traceback.format_exc()}); s.dump(OUT/'targeted_tests.json',{'status':'FAIL','tests':tests}); raise
    s.dump(OUT/'targeted_tests.json',{'status':'RUNNING','tests':tests})
def protected():
    expected=read_json(OUT/'protected_before.json'); assert all(sha(p)==h for p,h in expected.items())
    original=read_json(OUT/'generation_before.json'); assert all(sha(p)==h for p,h in original.items())
    return {'protected_files':len(expected),'real_DEV_generation_unchanged':True}
def fixture(name):
    folder=OUT/'fixtures'/RUN/name/'DEV/Справочники'; folder.mkdir(parents=True)
    for n in FILES: shutil.copy2(OUT/'baseline_generation'/n,folder/n)
    return folder/FILES[0]
def rejected(fn):
    try: fn()
    except (ValueError,OSError) as exc: return str(exc)
    raise AssertionError('Expected rejection')
def generation(t): return {n:sha(t.parent/n) for n in FILES}

before=hashes(); s.dump(OUT/'product_hashes_before.json',before)
case('baseline generation / protected evidence',protected)
catalog=s.load_catalog(); payload=s.workbook_payload(catalog)

def golden_and_m1():
    results=[]
    for c in catalog.conflicts:
        r=catalog.lookup(c['key'],c['standard'])
        assert r['status']==('CONFLICT' if c['classification']=='BLOCKING' else 'FOUND')
        assert r['profile']['unresolved_fields']==c['fields']
        results.append({'id':c['id'],'classification':c['classification'],'lookup':r['status']})
    assert len(results)==15 and sum(c['classification']=='BLOCKING' for c in results)==8
    assert sum(c['classification']=='NON_BLOCKING' for c in results)==7
    golden={}
    for nd,geometry,uzk in [('ГОСТ Р 57837-2017',[346,174,6,9],[[328,6],[348,9]]),
                           ('ГОСТ 26020-83',[346,155,6.2,8.5],[[329,6.2],[310,8.5]])]:
        r=catalog.lookup('35Б1',nd); assert r['status']=='FOUND'
        actual=[r['profile']['geometry'][k] for k in ['h','b','s','t']]; assert actual==geometry
        values=[[r['L'],r['S']] for r in calculate_uzk(r)['rows']]; assert values==uzk
        golden[nd]={'status':r['status'],'geometry':actual,'UZK':values}
    assert catalog.lookup('35Б1')['status']=='FOUND_PREFERRED'
    assert catalog.lookup('20Ш1')['status']=='AMBIGUOUS'
    assert catalog.lookup('L125×80×7','ГОСТ 8510-86')['status']=='CONFLICT'
    assert catalog.lookup('L100×8','ГОСТ 8509-93')['status']=='FOUND'
    s.dump(OUT/'M1_golden.json',{'disagreements':results,'golden':golden,'plain35':'FOUND_PREFERRED','plain20':'AMBIGUOUS'})
    return {'blocking_CONFLICT':8,'non_blocking_FOUND':7,'35Б1':'PASS','20Ш1':'AMBIGUOUS'}
case('M1 all 15 + 35Б1 golden + 20Ш1 + accepted examples',golden_and_m1)

isolated=fixture('isolated_native')
def generated():
    old=generation(isolated)
    with no_network() as attempts: result=s.build(isolated,previews=OUT/'previews')
    assert not attempts and result['network_attempts']==0
    backup=Path(result['backup']); assert {n:sha(backup/n) for n in FILES}==old
    current=s.load_catalog(isolated)
    assert current.counts()==catalog.counts()
    assert current.user_state['revision']==catalog.user_state['revision']+1
    for field in ['preferences','manual_rows','resolutions']: assert current.user_state[field]==catalog.user_state[field]
    s.dump(OUT/'isolated_build.json',result)
    return result
case('offline isolated build + coherent backup + current user state preserved',generated)
structure=inspect(isolated,OUT/'post_fix_xml'); assert structure['valid'],structure
s.dump(OUT/'post_fix_structure.json',structure)
case('all 5 table definitions / ranges / headers / XML / relationships',lambda:structure)

def identical_values():
    a=load_workbook(OUT/'baseline_generation/Профили.xlsx',read_only=True,data_only=False)
    b=load_workbook(isolated,read_only=True,data_only=False); differences=[]
    try:
        assert a.sheetnames==b.sheetnames
        for name in a.sheetnames:
            old=list(a[name].iter_rows(values_only=True)); new=list(b[name].iter_rows(values_only=True))
            assert len(old)==len(new)
            for row,(x,y) in enumerate(zip(old,new),1):
                assert len(x)==len(y)
                for col,(v,w) in enumerate(zip(x,y),1):
                    if v!=w: differences.append({'sheet':name,'row':row,'column':col,'before':v,'after':w})
        assert differences==[{'sheet':'Профили','row':1,'column':16,'before':'R','after':'Радиус R'}]
        assert s.verify_workbook(isolated,payload)['formulas']==0
    finally: a.close(); b.close()
    return {'all_values_statuses_provenance_preserved':True,'only_caption_changed':differences}
case('all worksheet values preserved; only technical heading P1 changes',identical_values)

def negative_xml():
    folder=OUT/'invalid_xml'; folder.mkdir(exist_ok=True)
    with zipfile.ZipFile(isolated) as z: original={n:z.read(n) for n in z.namelist()}
    cases={
        'duplicate_names':lambda r:r.find('{'+MAIN+'}tableColumns')[15].set('name','R'),
        'blank_name':lambda r:r.find('{'+MAIN+'}tableColumns')[0].set('name',''),
        'wrong_ref':lambda r:r.set('ref','A1:S2850'),
        'wrong_autofilter':lambda r:r.find('{'+MAIN+'}autoFilter').set('ref','A1:T2849'),
        'wrong_column_count':lambda r:r.find('{'+MAIN+'}tableColumns').set('count','19'),
        'duplicate_id':lambda r:r.find('{'+MAIN+'}tableColumns')[1].set('id','1'),
        'header_mismatch':lambda r:r.find('{'+MAIN+'}tableColumns')[0].set('name','OTHER'),
    }
    result=[]
    for name,mutate in cases.items():
        parts=original.copy(); node=ET.fromstring(parts['xl/tables/table1.xml']); mutate(node)
        parts['xl/tables/table1.xml']=ET.tostring(node,encoding='utf-8',xml_declaration=True)
        path=folder/(name+'.xlsx')
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for n,value in parts.items(): z.writestr(n,value)
        finding=inspect(path); assert not finding['valid'],name
        result.append({'case':name,'rejected':True,'issues':finding['issues']})
    s.dump(OUT/'negative_xml_results.json',result)
    return {'invalid_table_definitions_rejected':len(result)}
case('deterministic XML validator negative controls',negative_xml)

def duplicate_renderer():
    value=deepcopy(payload); value['sheets']['Профили'][0][15]='R'
    path=OUT/'duplicate_renderer_payload.json'; s.dump(path,value)
    message=rejected(lambda:s.render_workbook(path,OUT/'must_not_exist.xlsx'))
    assert 'case-insensitive duplicate' in message and not (OUT/'must_not_exist.xlsx').exists()
    return {'duplicate_r_R_rejected_before_export':True,'message':message}
case('renderer rejects r/R duplicate headings before export',duplicate_renderer)

def native_isolated():
    result=native(isolated,'isolated_fixed')
    assert result['validation']=='PASS' and not result['repair_detected'] and not result['recovered_title']
    return {'status':'PASS','display_alerts':True,'repair':False,'recovered_title':False,'logs':0,'rows':2849,'tables':5}
case('mandatory native Excel isolated normal Open; alerts enabled; no repair',native_isolated)

template=OUT/'verified_template.xlsx'; shutil.copy2(isolated,template)
def copied_render(input,destination,previews=None):
    assert read_json(input)==payload
    shutil.copy2(template,destination); return {'network_attempts':0}
def external_race(name):
    t=fixture('external_'+str(FILES.index(name))); old=generation(t); edited={}
    def render(input,destination,previews=None):
        result=copied_render(input,destination); p=t.parent/name
        if name.endswith('.xlsx'):
            with zipfile.ZipFile(p,'a') as z:z.comment=b'EXTERNAL_VALID_EDIT'
        else:
            value=read_json(p)
            if name==s.STATE:value['preferences'][0]['basis']='EXTERNAL_VALID_EDIT'
            elif name==s.CATALOG:value['records'][0]['notes'].append('EXTERNAL_VALID_EDIT')
            else:value['built_at']='2026-10-07T08:00:00Z'
            s.dump(p,value)
        edited['hash']=sha(p); return result
    with patch.object(s,'render_workbook',render): message=rejected(lambda:s.build(t))
    final=generation(t)
    assert final[name]==edited['hash']!=old[name]
    assert all(final[n]==old[n] for n in FILES if n!=name)
    assert not (t.parent/'backups').exists()
    return {'file':name,'STOP':message,'external_bytes_preserved':True,'other_files_unchanged':True,'backup_not_created':True}
for name in FILES: case('M2 external generation protection '+name,lambda n=name:external_race(n))

def boundary():
    t=fixture('commit_boundary'); original=ProtectedFile.rename; attempts=[]
    def rename(handle,destination):
        if not attempts:
            p=subprocess.run([sys.executable,'-I','-B',str(OUT/'external_writer.py'),str(t.parent)],capture_output=True,text=True,encoding='utf-8',timeout=20)
            assert p.returncode==0,p.stderr
            attempts.extend(json.loads(p.stdout)); assert len(attempts)==12 and all(x['blocked'] for x in attempts)
        return original(handle,destination)
    with patch.object(s,'render_workbook',copied_render),patch.object(ProtectedFile,'rename',rename):result=s.build(t)
    assert s.load_catalog(t).user_state==s.load_catalog(isolated).user_state
    s.dump(OUT/'commit_boundary.json',attempts)
    return {'independent_process_attempts':12,'all_blocked':True,'normal_commit':'PASS'}
case('M2 actual protected commit boundary independent writer',boundary)
case('protected evidence / real DEV still unchanged after tests',protected)
after=hashes(); s.dump(OUT/'product_hashes_after.json',after); assert before==after
assert not network and not production
s.dump(OUT/'targeted_tests.json',{'status':'PASS','tests':tests,'passed':len(tests),'failed':0,
    'product_hashes_before_after_equal':True,'network_events':network,'production_events':production,'real_DEV_rebuilds':0})
print(json.dumps({'status':'PASS','passed':len(tests),'real_DEV_rebuilds':0},ensure_ascii=False),flush=True)
