"""Pre-product A/B proof: only P1/tableColumn16 caption differs."""
from copy import deepcopy
from pathlib import Path
import json,sys
OUT=Path(__file__).resolve().parent; ROOT=OUT.parents[1]
sys.path.insert(0,str(ROOT)); sys.path.insert(0,str(OUT))
sys.stdout.reconfigure(encoding='utf-8')
import profiles_stage4.storage as s
from profiles_stage4.offline import no_network
from xlsx_tables import inspect
from openpyxl import load_workbook
if __name__=='__main__':
    folder=OUT/'header_only_proof'; folder.mkdir()
    with no_network() as events:
        payload=s.workbook_payload(s.load_catalog())
        assert payload['sheets']['Профили'][0][15]=='R'
        payload['sheets']['Профили'][0][15]='Радиус R'
        s.dump(folder/'payload.json',payload)
        result=s.render_workbook(folder/'payload.json',folder/'Профили.xlsx')
    assert not events and result['network_attempts']==0
    before=load_workbook(OUT/'baseline_generation/Профили.xlsx',read_only=True,data_only=False)
    after=load_workbook(folder/'Профили.xlsx',read_only=True,data_only=False)
    differences=[]
    try:
        for name in before.sheetnames:
            for index,(old,new) in enumerate(zip(before[name].iter_rows(values_only=True),after[name].iter_rows(values_only=True)),1):
                for col,(a,b) in enumerate(zip(old,new),1):
                    if a!=b: differences.append({'sheet':name,'row':index,'column':col,'before':a,'after':b})
    finally: before.close(); after.close()
    assert differences==[{'sheet':'Профили','row':1,'column':16,'before':'R','after':'Радиус R'}],differences
    structure=inspect(folder/'Профили.xlsx',OUT/'header_proof_xml')
    assert structure['valid'],structure
    s.dump(OUT/'header_only_proof.json',{'network_attempts':0,'cell_differences':differences,'structure':structure})
    print(json.dumps({'header_only_change':differences,'structure_valid':True},ensure_ascii=False))
