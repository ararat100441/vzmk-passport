"""Explicit DEV sources. No production registry discovery or synchronization."""
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from passport_app.catalog import key, profile_key, HEADERS, LEGACY_HEADERS

ROOT = Path(__file__).resolve().parents[1]
DEV = ROOT / 'DEV'
REGISTRY = DEV / 'ВЗМК_DEV.xlsx'
UZK = DEV / 'УЗК_DEV.xlsm'
PROFILES = DEV / 'Справочники/Профили.xlsx'
CHECKPOINT = DEV / 'DEV_UZK_CHECKPOINT.json'
PINS = {
    REGISTRY: '8e56f68518c70cade9efb09084dd67cd5d72e07e3fdc813b8515030e9cb98e80',
    UZK: '9f25a30139ab9e31b6da7248284c2ae6e9d4327f0d68251d5723293992662302',
    PROFILES: '7b420f6dead96408680a64d220c83e8980aeec1c38635e35b2e4723a98e3831b',
}


class Stage5Error(ValueError):
    pass


def header_key(value):
    return re.sub(r'\s+', ' ', key(value))


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def verify_pins(pins=None):
    for path, expected in (PINS if pins is None else pins).items():
        if not Path(path).is_file() or sha(path) != expected:
            raise Stage5Error(f'CHECKPOINT_CHANGED: {path}. Требуется проверка нового DEV checkpoint.')
    if pins is None:
        cp = json.loads(CHECKPOINT.read_text(encoding='utf-8-sig'))
        if (Path(cp['dev_path']).resolve() != UZK.resolve()
                or cp['dev_initial_sha256'] != PINS[UZK]
                or cp['policy']['auto_sync'] is not False
                or cp['policy']['registry_decrement'] != 'NOT_AUTHORIZED'):
            raise Stage5Error('CHECKPOINT_INVALID: DEV_UZK_CHECKPOINT.json')
    # source_path from checkpoint is metadata only: NEVER opened or hashed.


def positive_integer(value, where, blank_one=False):
    if value is None or isinstance(value, str) and not value.strip():
        if blank_one:
            return 1
        raise Stage5Error(f'QUANTITY_MISSING: {where}')
    if isinstance(value, bool):
        raise Stage5Error(f'QUANTITY_INVALID: {where}: {value!r}')
    try:
        n = Decimal(str(value).strip().replace(',', '.'))
    except InvalidOperation as exc:
        raise Stage5Error(f'QUANTITY_INVALID: {where}: {value!r}') from exc
    if not n.is_finite() or n <= 0 or n != n.to_integral():
        raise Stage5Error(f'QUANTITY_INVALID: {where}: {value!r}')
    return int(n)


def text(value, where):
    if not isinstance(value, str) or not value.strip() or value.startswith(('=', '#')):
        raise Stage5Error(f'TEXT_INVALID: {where}: {value!r}')
    return value


@contextmanager
def workbook(path, *, data_only=False):
    book = load_workbook(path, read_only=True, data_only=data_only, keep_links=False)
    try:
        yield book
    finally:
        book.close()


def passports(path=REGISTRY):
    # N is a formula in the accepted DEV workbook. Read its accepted cached date;
    # openpyxl never evaluates formulas or refreshes external links.
    with workbook(path, data_only=True) as book:
        if 'Паспорт качества' not in book.sheetnames:
            raise Stage5Error('PASSPORT_SHEET_MISSING')
        sheet = book['Паспорт качества']
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        expected = {1: '№ паспорта', 4: 'Наименование объекта', 5: 'шифр КМД', 13: 'УЗК Дата'}
        if any(key(header[c]) != key(v) for c, v in expected.items()):
            raise Stage5Error('PASSPORT_MAPPING_CHANGED: проверьте заголовки B/E/F/N')
        result = []
        for rn, row in enumerate(rows, 2):
            if not row[0] or not row[1]:
                continue
            if not isinstance(row[0], str) or not re.fullmatch(r'\d+/\d{4}', str(row[1])):
                continue
            result.append({'row': rn, 'project': row[0], 'number': str(row[1]),
                           'object': row[4], 'cipher': row[5], 'mass': row[6], 'date': row[13]})
        return result


def select_passport(project, number):
    all_rows = passports()
    selected = [r for r in all_rows if key(r['project']) == key(project) and r['number'] == number]
    if len(selected) != 1:
        raise Stage5Error('PASSPORT_AMBIGUOUS_OR_MISSING: выберите единственный проект / паспорт')
    row = selected[0]
    if not isinstance(row['date'], (date, datetime)):
        raise Stage5Error('UZK_DATE_INVALID: требуется дата из DEV колонки N')
    for field in ('project', 'object', 'cipher'):
        text(row[field], f'Паспорт качества, строка {row["row"]}, {field}')
    number_year = number.split('/')[1]
    numbers = [int(r['number'].split('/')[0]) for r in all_rows
               if key(r['project']) == key(project) and r['number'].endswith('/' + number_year)]
    return {**row, 'date': row['date'].strftime('%d.%m.%Y'), 'first_number': f'{min(numbers)}_{number_year}'}


def project_sources(passport):
    settings = json.loads((ROOT / 'settings.json').read_text(encoding='utf-8-sig'))
    docs = (ROOT / settings['documents_root']).resolve()
    configs = [v for k, v in settings.get('projects', {}).items() if key(k) == key(passport['project'])]
    if len(configs) == 1:
        passport_root = (docs / configs[0]['passport_root']).resolve()
        project_root = passport_root.parent
    elif not configs:
        candidates = [p for p in docs.iterdir() if p.is_dir() and key(p.name) == key(passport['project'])]
        if len(candidates) != 1:
            raise Stage5Error('PROJECT_FOLDER_AMBIGUOUS_OR_MISSING: требуется точное сопоставление проекта')
        project_root = candidates[0].resolve()
        passport_root = project_root / 'Паспорт'
    else:
        raise Stage5Error('PROJECT_CONFIG_AMBIGUOUS')
    if not project_root.is_relative_to(docs) or project_root == ROOT or not passport_root.is_dir():
        raise Stage5Error('PROJECT_PATH_INVALID')
    template = passport_root / (passport['first_number'] + ' - УЗК.docx')
    if not template.is_file():
        # No exact, accepted global fallback is configured in this DEV baseline.
        raise Stage5Error(f'TEMPLATE_MISSING: {template}. Требуется first-passport шаблон.')
    catalog = project_root / 'Справочник марок.xlsx'
    if not catalog.is_file():
        raise Stage5Error(f'MARK_CATALOG_MISSING: {catalog}')
    return project_root, catalog, template


def registry_rows(path=UZK):
    result = {}
    evidence = []
    with workbook(path) as book:
        if 'УЗК' not in book.sheetnames:
            raise Stage5Error('UZK_SHEET_MISSING')
        sheet = book['УЗК']
        rows = sheet.iter_rows(min_col=1, max_col=3, values_only=True)
        header = next(rows)
        if tuple(header_key(v) for v in header) != tuple(header_key(v) for v in ('Заказ', 'Марка элемента', 'Кол-во (шт)')):
            raise Stage5Error('UZK_HEADER_CHANGED')
        for rn, (project, mark, quantity) in enumerate(rows, 2):
            if all(v is None or v == '' for v in (project, mark, quantity)):
                continue
            text(project, f'УЗК A{rn}'); text(mark, f'УЗК B{rn}')
            count = positive_integer(quantity, f'УЗК C{rn}', blank_one=True)
            identity = (key(project), key(mark))
            if identity in result and result[identity]['remaining'] != count:
                raise Stage5Error(f'REGISTRY_CONFLICT: {project} / {mark}, строки {result[identity]["rows"]} и {rn}')
            record = {'row': rn, 'project': project, 'mark': mark, 'raw_quantity': quantity, 'remaining': count}
            evidence.append(record)
            if identity in result:
                result[identity]['rows'].append(rn)  # identical duplicate is not extra stock
            else:
                result[identity] = {**record, 'rows': [rn]}
    return result, evidence


def mark_profiles(path, *, required_marks=None):
    required = None if required_marks is None else {key(mark) for mark in required_marks}
    result = {}
    with workbook(path) as book:
        if len(book.sheetnames) != 1:
            raise Stage5Error('MARK_CATALOG_SHEET_AMBIGUOUS')
        sheet = book.active
        rows = list(sheet.iter_rows(values_only=True))
        headers = [i for i, row in enumerate(rows[:30])
                   if any(tuple(key(v) for v in row[:len(schema)]) == tuple(key(v) for v in schema)
                          and all(v is None or v == '' for v in row[len(schema):])
                          for schema in (HEADERS, LEGACY_HEADERS))]
        if len(headers) != 1:
            raise Stage5Error('MARK_CATALOG_HEADER_AMBIGUOUS')
        for rn, row in enumerate(rows[headers[0] + 1:], headers[0] + 2):
            if not any(v is not None and v != '' for v in row):
                continue
            if required is not None and key(row[0]) not in required:
                continue
            mark = text(row[0], f'Справочник марок A{rn}')
            raw = row[1]  # Missing / ПРОВЕРИТЬ blocks only selected UZK marks.
            identity = key(mark)
            if identity in result and profile_key(result[identity]['profile']) != profile_key(raw):
                raise Stage5Error(f'MARK_PROFILE_CONFLICT: {mark}, строка {rn}')
            result.setdefault(identity, {'mark': mark, 'profile': raw, 'row': rn})
    return result


def visible_vom(path):
    """Native XLS/XLSX adapter: persisted filters + manually hidden rows, no save."""
    import pythoncom
    import win32com.client
    path = Path(path).resolve()
    if path.suffix.lower() not in ('.xls', '.xlsx'):
        raise Stage5Error('VOM_FORMAT: выберите XLS или XLSX')
    before = sha(path)
    pythoncom.CoInitialize()
    app = book = None
    try:
        app = win32com.client.DispatchEx('Excel.Application')
        app.Visible = False
        app.DisplayAlerts = False
        app.AutomationSecurity = 3
        app.EnableEvents = False
        app.AskToUpdateLinks = False
        book = app.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True,
                                  IgnoreReadOnlyRecommended=True, AddToMru=False)
        if not book.ReadOnly:
            raise Stage5Error('VOM_NOT_READ_ONLY')
        external_links = book.LinkSources(1)
        connections = int(book.Connections.Count)
        if external_links or connections or any(int(s.QueryTables.Count) for s in book.Worksheets):
            raise Stage5Error('VOM_EXTERNAL_DATA: внешние ссылки / подключения не поддерживаются offline')
        found = []
        for sheet in book.Worksheets:
            if sheet.Visible != -1:
                continue
            used = sheet.UsedRange
            values = used.Value2
            if not isinstance(values, tuple):
                continue
            for offset, row in enumerate(values[:30]):
                normalized = [key(v) if v is not None else '' for v in row]
                if normalized.count('отправочная марка') == 1 and normalized.count('количество') == 1:
                    mc, qc = normalized.index('отправочная марка'), normalized.index('количество')
                    masscols = [i for i, v in enumerate(normalized) if v.startswith('масса') and 'кг' in v]
                    if len(masscols) != 1 or offset + 1 >= len(values):
                        raise Stage5Error('VOM_MASS_HEADER_AMBIGUOUS')
                    # Verified two-line header: Масса (изделие), кг -> Всех.
                    total = masscols[0] + 1
                    if total >= len(row) or key(values[offset + 1][total]) != 'всех':
                        raise Stage5Error('VOM_TOTAL_MASS_HEADER_CHANGED')
                    found.append((sheet, used.Row, values, offset, mc, qc, total))
        if len(found) != 1:
            raise Stage5Error('VOM_SHEET_AMBIGUOUS: требуется единственная видимая ведомость с заголовками')
        sheet, start, values, header, mc, qc, massc = found[0]
        rows, hidden = [], []
        mass = Decimal(0)
        for offset in range(header + 2, len(values)):
            values_row = values[offset]
            mark = values_row[mc]
            if mark is None or mark == '':
                continue  # footer totals have no mark; never used as shipment mass
            rn = start + offset
            if key(mark) in ('итого', 'всего'):
                continue
            if bool(sheet.Rows(rn).Hidden):
                hidden.append(rn)
                continue
            text(mark, f'{sheet.Name}! строка {rn}, марка')
            qty = positive_integer(values_row[qc], f'{sheet.Name}! строка {rn}, количество')
            try:
                row_mass = Decimal(str(values_row[massc]))
            except InvalidOperation as exc:
                raise Stage5Error(f'VOM_MASS_INVALID: строка {rn}') from exc
            if not row_mass.is_finite() or row_mass < 0:
                raise Stage5Error(f'VOM_MASS_INVALID: строка {rn}')
            mass += row_mass
            rows.append({'row': rn, 'mark': mark, 'quantity': qty, 'mass_kg': str(row_mass)})
        return {'path': str(path), 'sheet': sheet.Name, 'read_only': True,
                'macros_disabled': True, 'update_links': False, 'adapter': 'Excel COM XLS/XLSX',
                'external_link_count': len(external_links) if external_links else 0,
                'connection_count': connections,
                'visible_count': len(rows), 'mass_kg': str(mass.quantize(Decimal('0.01'))),
                'raw_mass_sum': str(mass), 'rows': rows,
                'hidden_data_rows': hidden, 'autofilter_range': sheet.AutoFilter.Range.Address
                if sheet.AutoFilterMode else None}
    except Stage5Error:
        raise
    except Exception as exc:
        raise Stage5Error(f'VOM_READ_FAILED: {exc}') from exc
    finally:
        try:
            if book is not None:
                book.Close(SaveChanges=False)
        finally:
            if app is not None:
                app.Quit()
            book = app = None
            pythoncom.CoUninitialize()
            if sha(path) != before:
                raise Stage5Error('VOM_CHANGED_DURING_READ')
