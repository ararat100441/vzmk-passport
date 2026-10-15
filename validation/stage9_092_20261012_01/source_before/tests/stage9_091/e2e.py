from common import *
import traceback
from stage8_closure.core import confirmation
from stage8_closure.recovery import confirmation as restore_confirmation
RUN=OUT/('e2e_'+uuid4().hex[:8]);RUN.mkdir();steps=[]
def stocks(sandbox):
    from openpyxl import load_workbook
    b=load_workbook(Path(sandbox)/'registry/УЗК_TEST.xlsm',read_only=True,data_only=True)
    try:return [b['УЗК'].cell(r,3).value for r in (24,25)]
    finally:b.close()
def vba(sandbox):
    from zipfile import ZipFile
    import hashlib,io,olefile
    # Excel may repack OLE sectors; every actual stream, including compiled
    # VBA, must remain byte-identical (the accepted Stage 8 contract).
    with ZipFile(Path(sandbox)/'registry/УЗК_TEST.xlsm') as z:binary=z.read('xl/vbaProject.bin')
    with olefile.OleFileIO(io.BytesIO(binary)) as ole:
        return {'/'.join(n):hashlib.sha256(ole.openstream(n).read()).hexdigest() for n in ole.listdir()}
def run(name,expected='PASS',**args):
    data,sup,path=task(RUN,name,expected,**args);steps.append({'task':name,'result':str(path),'supervisor':data['office_supervisor'],'expected':expected})
    save(RUN/f'step_{len(steps):02}.json',steps[-1]);return data
try:
    work=run('import',source=ROOT,parent=Path(os.environ['LOCALAPPDATA'])/'VZMKPassport/workspaces')['workspace']
    smoke=run('smoke',workspace=work);assert smoke['frozen'] and smoke['python_in_path'] is None and smoke['version']=='0.9.1'
    vom=ROOT.parent/'Гараж казахстан/Паспорт/9 от 06.10.2026/ВОМ отгрузочный счет 183  (5).xls'
    outputs=[]
    for number,new in (('9/2026',False),('10/2026',True),('9/2026',False)):
        kw={'new':True} if new else {}
        preview=run('check',workspace=work,vom=vom,passport=number,variants='0',**kw)
        d=preview['data'];assert (d['visible_rows'],d['hidden_excluded'],d['quantity'],d['mass_kg'],d['area_m2'],d['uzk_units'])==(104,397,'288','15822.14','1913.494',33)
        generated=run('generate',workspace=work,preview=preview['preview'],confirm_test=True)
        run('verify',workspace=work,manifest=generated['manifest'])
        native=run('native-verify',workspace=work,manifest=generated['manifest'])
        assert native['rows']==104 and native['connections_m1_equal'] and native['word_documents']==4 and native['excel_documents']==3
        outputs.append(generated['manifest'])
    run('verify',workspace=work,manifest=outputs[0])
    fixture=run('fixture',workspace=work)
    before_vba=vba(fixture['sandbox']);assert stocks(fixture['sandbox'])==[12,10]
    request=read(fixture['request']);request['passport']='9/2026';blocked=Path(fixture['request']).with_name('OFFICIAL_9_REQUEST.json');save(blocked,request)
    denied=run('close-preview','FAIL',workspace=work,sandbox=fixture['sandbox'],request=blocked);assert 'PASSPORT_9_2026' in denied['error'] and stocks(fixture['sandbox'])==[12,10]
    preview=run('close-preview',workspace=work,sandbox=fixture['sandbox'],request=fixture['request'])['plan']
    plan=RUN/'CLOSE_PLAN.json';save(plan,preview)
    closed=run('close',workspace=work,preview=plan,confirm=confirmation(preview));assert closed['operation_status']=='CLOSED_TEST'
    assert stocks(fixture['sandbox'])==[9,6] and vba(fixture['sandbox'])==before_vba
    again=run('close',workspace=work,preview=plan,confirm=confirmation(preview));assert again['operation_status']=='ALREADY_CLOSED_TEST'
    assert stocks(fixture['sandbox'])==[9,6]
    run('recover',workspace=work,sandbox=fixture['sandbox'])
    torn=run('fixture',workspace=work)
    torn_vba=vba(torn['sandbox'])
    preview=run('close-preview',workspace=work,sandbox=torn['sandbox'],request=torn['request'])['plan']
    plan=RUN/'CRASH_PLAN.json';save(plan,preview)
    crash=run('close','FAIL',workspace=work,preview=plan,confirm=confirmation(preview),test_fault='during_publish')
    assert crash['worker_returncode']==73
    rp=run('restore-preview',workspace=work,sandbox=torn['sandbox'])['plan'];rpath=RUN/'RESTORE_PLAN.json';save(rpath,rp)
    restored=run('restore',workspace=work,preview=rpath,confirm=restore_confirmation(rp));assert restored['operation_status']=='RECOVERED_ABORTED' and restored['decrements']==0
    assert stocks(torn['sandbox'])==[12,10] and vba(torn['sandbox'])==torn_vba
    after=run('close-preview',workspace=work,sandbox=torn['sandbox'],request=torn['request']);assert after['plan']['payload']['status']=='READY_TEST'
    run('smoke',workspace=work)
    save(RUN/'RESULT.json',{'status':'PASS','exe_sha256':sha(EXE),'workspace':work,'manifests':outputs,'fixture_closed':fixture,'fixture_restored':torn,'steps':steps,'python_in_path':None})
    save(OUT/'EXE_E2E_FINAL.json',{'status':'PASS','exe_sha256':sha(EXE),'evidence':str(RUN/'RESULT.json'),'workspace':work,'manifests':outputs})
except Exception as exc:
    save(RUN/'FAILURE.json',{'status':'FAIL','exe_sha256':sha(EXE),'error':repr(exc),'traceback':traceback.format_exc(),'steps':steps});raise
