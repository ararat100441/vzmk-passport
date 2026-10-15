"""Stage 9 TEST-only coating simulation; accepted Stage 6 remains unchanged.

Project-bound sample forms supply coating systems, never historical measurements.
Settings and immutable daily climates live outside templates and registries.
"""
from contextlib import contextmanager
from datetime import date, timedelta
from decimal import Decimal, ROUND_CEILING
from pathlib import Path
from uuid import uuid4
import json, math, os, random, re, secrets, shutil
from openpyxl import load_workbook
from stage8_closure import storage as S
from .workspace import bind, load, save, PROJECT
from . import akz_settings as K

SCHEMA='STAGE9_AKZ_AUTO_SETTINGS_V1'
MODE='TEST_SIMULATION'
FORMULA='MAGNUS_WATER_17.62_243.12'
FORMULA_SOURCE='https://etrp.wmo.int/pluginfile.php/97653/mod_folder/content/0/Training%20Material/5_1_Humidity%20Part%201%20Introduction.pdf'
MISSING='ТЕСТ: Н/Д'
DEFAULT_THICKNESS={0:[(80,100)],1:[(150,180)],2:[(100,120),(80,100)],3:[(100,120),(80,100),(60,80)]}

def require(ok, message):
    if not ok:raise S.Stop(message)

def numeric(value):
    require(type(value) in (int,float) and math.isfinite(value),'AKZ_RANGE_NUMBER_INVALID')
    return Decimal(str(value))

def interval(pair,low,high,step=Decimal('.1')):
    require(isinstance(pair,list) and len(pair)==2,'AKZ_RANGE_PAIR_REQUIRED')
    a,b=map(numeric,pair)
    require(Decimal(str(low))<=a<=b<=Decimal(str(high)),'AKZ_RANGE_LIMITS_INVALID')
    require(a/step==(a/step).to_integral() and b/step==(b/step).to_integral(),'AKZ_RANGE_PRECISION_INVALID')
    return a,b

def dew_point(air_c,humidity):
    t,h=numeric(air_c),numeric(humidity)
    require(Decimal(-45)<=t<=Decimal(60) and 0<h<=100,'MAGNUS_INPUT_DOMAIN_STOP')
    gamma=math.log(float(h)/100)+17.62*float(t)/(243.12+float(t))
    result=243.12*gamma/(17.62-gamma)
    require(math.isfinite(result) and -45-1e-10<=result<=60+1e-10,'MAGNUS_RESULT_DOMAIN_STOP')
    return max(-45.0,min(60.0,result))

def settings_path(root):
    root,_=load(root)
    return S.no_links(root/'preferences/akz_auto/settings.json')

@contextmanager
def settings_lock(root):
    import win32file,win32con
    folder=settings_path(root).parent;folder.mkdir(parents=True,exist_ok=True)
    target=S.no_links(folder/'.lock')
    try:handle=win32file.CreateFile(str(target),win32con.GENERIC_READ|win32con.GENERIC_WRITE,0,None,win32con.OPEN_ALWAYS,0,None)
    except Exception as exc:raise S.Stop('AKZ_SETTINGS_OR_GENERATION_LOCKED') from exc
    try:yield folder
    finally:handle.Close()

def profiles(root):
    from akz_stage6 import core as c,part2
    from passport_app.registry import read_registry,select_project
    root,meta=bind(root)
    require(meta['project']==PROJECT,'AKZ_PROJECT_PROFILE_NOT_APPROVED')
    first=select_project(read_registry(root/'DEV/ВЗМК_DEV.xlsx'),meta['project'])
    result={}
    for n in range(4):
        t=c.templates() if n==0 else part2.template(n)
        for field in ('act','application'):
            path=S.no_links(t[field]);relative=path.relative_to(root).as_posix()
            require(path.parent==root/'documents/project/Паспорт' and relative in meta['inputs']
                    and S.sha(path)==meta['inputs'][relative]['sha256'],'AKZ_SAMPLE_BINDING_STOP')
        if n==0:t['layers']=[{'ordinal':1,'ordinal_cell':'E17','column':'E','material':t['material'],'color':t['ral'],'required_raw':t['required_dry_microns'],'required_individual_microns':t['required_dry_microns']}]
        t['variant']=n;t['first_passport']=first.first_number;t['project']=meta['project']
        t['sample_policy']='USER_APPROVED_SYSTEM_SAMPLES_2026_10_13_REPLACE_FOREIGN_REQUISITES'
        t['hashes']={str(Path(t[k])):S.sha(t[k]) for k in ('act','application')}
        result[str(n)]=t
    return result

def defaults(root):
    p=profiles(root)
    return {'schema':SCHEMA,'project':PROJECT,'enabled':True,'mode':MODE,'official_issue':False,
            'template_bindings':{n:t['hashes'] for n,t in p.items()},
            'climate':{'air_c':[15,25],'surface_c':[18,28],'relative_humidity':[40,65],'formula':FORMULA},
            'thickness':{str(n):[list(pair) for pair in pairs] for n,pairs in DEFAULT_THICKNESS.items()}}

def validate_settings(cfg,p):
    require(isinstance(cfg,dict) and set(cfg)=={'schema','project','enabled','mode','official_issue','template_bindings','climate','thickness'},'AKZ_SETTINGS_SCHEMA_STOP')
    require(cfg['schema']==SCHEMA and cfg['project']==PROJECT and type(cfg['enabled']) is bool
            and cfg['mode']==MODE and cfg['official_issue'] is False,'AKZ_OFFICIAL_OR_PROJECT_STOP')
    require(cfg['template_bindings']=={n:t['hashes'] for n,t in p.items()},'AKZ_SETTINGS_TEMPLATE_SHA_STOP')
    c=cfg['climate'];require(isinstance(c,dict) and set(c)=={'air_c','surface_c','relative_humidity','formula'} and c['formula']==FORMULA,'AKZ_FORMULA_NOT_APPROVED')
    air=interval(c['air_c'],-45,60);surface=interval(c['surface_c'],-100,150);h=interval(c['relative_humidity'],Decimal('.1'),100)
    try:min_dew=dew_point(float(air[0]),float(h[0]))
    except S.Stop:min_dew=-45.0
    require(surface[1]>=Decimal(str(min_dew))+3,'AKZ_CLIMATE_RANGE_IMPOSSIBLE')
    # At least the warmest/wettest corner must remain in the approved water domain.
    dew_point(float(air[1]),float(h[1]))
    require(isinstance(cfg['thickness'],dict) and set(cfg['thickness'])==set(p),'AKZ_THICKNESS_VARIANTS_STOP')
    for n,t in p.items():
        pairs=cfg['thickness'][n]
        require(isinstance(pairs,list) and len(pairs)==len(t['layers']),'AKZ_LAYER_COUNT_STOP')
        minima=[]
        for pair,layer in zip(pairs,t['layers']):
            lo,hi=interval(pair,1,5000,step=Decimal(1));minima.append(lo)
            norm=layer['required_individual_microns']
            if norm is not None:require(lo>=Decimal(str(norm)),'AKZ_TEST_RANGE_BELOW_TEMPLATE_NORM')
        if n=='2':require(sum(minima)>=180 and t['layers'][1]['required_raw']=='∑1+2=180','AKZ_CUMULATIVE_180_STOP')
    return cfg

def read_settings(root,*,required=True):
    path=settings_path(root)
    if not path.exists():
        require(not required,'AKZ_SETTINGS_MISSING');return None
    return validate_settings(S.read(path),profiles(root))

def enabled(root):
    cfg=read_settings(root,required=False)
    return bool(cfg and cfg['enabled'])

def store_settings(root,cfg):
    validate_settings(cfg,profiles(root))
    cfg=K.normalize(cfg)
    with settings_lock(root) as folder:
        path=S.no_links(folder/'settings.json');old=S.sha(path) if path.exists() else None
        if path.exists() and K.digest(S.read(path))==K.digest(cfg):
            return {'settings':str(path),'settings_sha256':old,'unchanged':True,'mode':MODE,'summary':summary(cfg,profiles(root))}
        history=folder/'history';history.mkdir(exist_ok=True)
        save(history/(uuid4().hex+'.json'),{'at':S.now(),'previous_sha256':old,'settings':cfg,'official_issue':False})
        temporary=folder/('settings_'+uuid4().hex+'.tmp');save(temporary,cfg);os.replace(temporary,path)
    return {'settings':str(path),'settings_sha256':S.sha(path),'mode':MODE,'summary':summary(cfg,profiles(root))}

def summary(cfg,p):
    lines=['ТЕСТ: диапазоны моделирования, не результаты измерений.','Формула точки росы: Магнус (17,62; 243,12).']
    for label,key in [('Воздух °C','air_c'),('Поверхность °C','surface_c'),('Влажность %','relative_humidity')]:lines.append(label+': '+str(cfg['climate'][key]))
    for n,t in p.items():
        for layer,pair in zip(t['layers'],cfg['thickness'][n]):
            lines.append(f"АКЗ №{n}, слой {layer['ordinal']}: TEST {pair[0]}–{pair[1]} мкм; норматив формы: {layer['required_raw']}")
    lines+=['Один неизменяемый климатический TEST-набор на календарную дату.','Официальная выдача запрещена.']
    return '\n'.join(lines)

def sample_range(rng,pair):
    a,b=interval(pair,-10000,10000)
    return rng.randint(int(a*10),int(b*10))/10

def daily_climate(folder,day,cfg):
    date.fromisoformat(day);directory=S.no_links(folder/'days');directory.mkdir(exist_ok=True)
    path=S.no_links(directory/(day+'.json'));fingerprint=K.digest(cfg)
    if path.exists():
        result=S.read(path);payload={k:v for k,v in result.items() if k!='seal'}
        require(result.get('seal')==S.digest(payload) and result.get('day')==day and result.get('mode')==MODE
                and result.get('formula')==FORMULA and result.get('official_issue') is False,'AKZ_DAY_RECORD_DAMAGED')
        if result.get('settings_digest_schema')==K.CANONICAL:
            require(K.digest(result.get('climate_settings'))==result.get('settings_digest'),'AKZ_DAY_RECORD_DAMAGED')
            compatible=result.get('settings_digest')==fingerprint
        elif 'settings_digest_schema' not in result:
            compatible=result.get('settings_digest') in K.legacy_climate_digests(cfg)
        else:raise S.Stop('AKZ_DAY_RECORD_DAMAGED')
        require(compatible,'AKZ_DAY_SETTINGS_CONFLICT: '+day)
        values=result['values'];expected=dew_point(values['air_c'],values['relative_humidity'])
        require(abs(expected-values['dew_point_raw_c'])<1e-10 and values['surface_c']>=max(expected,values['dew_point_c'])+3-1e-10,'AKZ_DAY_VALUES_DAMAGED')
        for key in ('air_c','surface_c','relative_humidity'):require(cfg[key][0]<=values[key]<=cfg[key][1],'AKZ_DAY_OUTSIDE_RANGE')
        return result
    seed=secrets.token_hex(16);rng=random.Random(seed)
    for _ in range(512):
        air=sample_range(rng,cfg['air_c']);humidity=sample_range(rng,cfg['relative_humidity'])
        try:dew=dew_point(air,humidity)
        except S.Stop:continue
        rounded=round(dew,1);low=max(Decimal(str(cfg['surface_c'][0])),(Decimal(str(max(dew,rounded)))+3).quantize(Decimal('.1'),rounding=ROUND_CEILING))
        if low>Decimal(str(cfg['surface_c'][1])):continue
        surface=sample_range(rng,[float(low),cfg['surface_c'][1]])
        break
    else:raise S.Stop('AKZ_CLIMATE_SAMPLING_IMPOSSIBLE')
    result={'schema':'STAGE9_TEST_CLIMATE_DAY_V1','day':day,'mode':MODE,'official_issue':False,
            'formula':FORMULA,'formula_source':FORMULA_SOURCE,'settings_digest':fingerprint,'seed':seed,
            'settings_digest_schema':K.CANONICAL,'climate_settings':K.normalize(cfg),
            'values':{'air_c':air,'surface_c':surface,'relative_humidity':humidity,'dew_point_c':rounded,'dew_point_raw_c':dew}}
    result['seal']=S.digest(result);save(path,result);return result

def validate_batch(root,batch,variants):
    from akz_stage6 import part2
    cfg=read_settings(root);require(cfg['enabled'],'AKZ_AUTO_DISABLED')
    require(batch.payload['project']==cfg['project'] and variants and len(variants)==len(set(variants)),'AKZ_PROJECT_OR_SELECTION_STOP')
    for n in variants:part2.group_preview(batch,n)
    return cfg

def shape(path):
    b=load_workbook(path,data_only=False,keep_links=False)
    try:
        return {s.title:{'merges':sorted(str(m) for m in s.merged_cells.ranges),'print_area':str(s.print_area),
                        'formulas':{c.coordinate:c.value for row in s for c in row if c.data_type=='f'},
                        'coating':{c.coordinate:c.value for row in s for c in row if c.value is not None and (c.row in (17,18,21,25) or c.row==26 and str(c.value).startswith('∑'))},
                        # Excel may discard explicit default style records when
                        # clearing obsolete helper formulas. Compare the actual
                        # appearance, including any added non-default component.
                        'fonts':{c.coordinate:str(c.font) for row in s for c in row if c.font!=b._fonts[0]},
                        'fills':{c.coordinate:str(c.fill) for row in s for c in row if c.fill!=b._fills[0]},
                        'borders':{c.coordinate:str(c.border) for row in s for c in row if c.border!=b._borders[0]}}
                for s in b}
    finally:b.close()

def write_act(app,path,p,cfg,folder):
    from akz_stage6 import core as c,part2
    data=p.payload;n=data.get('variant',0);columns=['E'] if n==0 else list(part2.layout(n).columns)
    b=c.native_open(app,path,read_only=False);days={};thickness=[];replaced=[]
    try:
        s=b.Worksheets('4');before_formulas=shape(path)['4']['formulas']
        s.Range('B3').Value2='ТЕСТ / МОДЕЛИРОВАНИЕ АКЗ'
        s.Range('E4').Value2=data['act_number']+' от '
        s.Range('H4').Value2=(date.fromisoformat(data['date'])-date(1899,12,30)).days
        for cell,key in [('E5','contract'),('E6','customer'),('E7','object')]:s.Range(cell).Value2=data['passport'][key]
        old_cipher=str(s.Range('B44').Value2)
        suffix=old_cipher[old_cipher.index(','):] if n==0 and ',' in old_cipher and 'ГОСТ' in old_cipher else ''
        s.Range('B44').Value2=data['passport']['cipher']+suffix
        s.Range('E7').WrapText=True;s.Rows(7).RowHeight=max(s.Rows(7).RowHeight,56)
        cells=list(c.FACTUAL_CELLS) if n==0 else part2.factual_cells(part2.layout(n))
        dynamic={col+str(r) for col in columns for r in (22,26,27,28,30,31,32,33,34)}
        for cell in cells:
            if cell not in dynamic and cell not in before_formulas:s.Range(cell).Value2='?' if cell in ('G11','E15','G37') else MISSING
        # Historical climate helper formulas are intentionally replaced; dry/wet
        # thickness, date, identity and application total formulas are retained.
        for cell,formula in before_formulas.items():
            if 'RANDBETWEEN' in formula and cell not in dynamic:
                s.Range(cell).ClearContents();replaced.append(cell)
        s.Range('B43').Value2='ТЕСТ: ИЗМЕРЕНИЙ НЕТ. ВЫДАЧА ЗАПРЕЩЕНА.'
        s.Range('B43').WrapText=True
        for cell in ('E46','E47'):s.Range(cell).Value2=MISSING
        for cell in ('F50','I50'):
            role=str(s.Range(cell).Value2).split('\n')[0];s.Range(cell).Value2=role+'\nТЕСТ: нет подписи'
        if n==2:
            helper=b.Worksheets('Лист1');helper.UsedRange.ClearContents();helper.Range('A1').Value2='СЛУЖЕБНЫЙ ЛИСТ / ТЕСТ';helper.Visible=2
            # This pinned sample formats wet thickness as currency. The copied
            # TEST form must show microns, like the adjacent first-layer cell.
            s.Range('I27').NumberFormat='0'
        app.CalculateFull()
        thickness_seed=secrets.token_hex(16);rng=random.Random(thickness_seed)
        for index,col in enumerate(columns):
            serial=s.Range(col+'22').Value2
            require(type(serial) in (float,int) and serial==int(serial),'AKZ_LAYER_DATE_NOT_PROVEN: '+col+'22')
            day=(date(1899,12,30)+timedelta(days=int(serial))).isoformat()
            climate=daily_climate(folder,day,cfg['climate']);days[day]=climate
            lo,hi=cfg['thickness'][str(n)][index]
            value=rng.randint(int(lo),int(hi));thickness.append(value)
            s.Range(col+'28').Value2=value
            if not str(s.Range(col+'26').Value2).startswith('∑'):
                s.Range(col+'26').Formula='=SUM('+','.join(x+'28' for x in columns[:index+1])+')'
            if col+'27' not in before_formulas:s.Range(col+'27').Value2=MISSING
            if col+'30' not in before_formulas:s.Range(col+'30').Value2=serial
            for row,key in ((31,'surface_c'),(32,'air_c'),(33,'relative_humidity'),(34,'dew_point_c')):
                address=col+str(row)
                if address in before_formulas:replaced.append(address)
                s.Range(address).Value2=climate['values'][key]
        s.PageSetup.CenterHeader='ТЕСТ / МОДЕЛИРОВАНИЕ — НЕ ИЗМЕРЕНИЯ'
        s.PageSetup.CenterFooter='ACT-ID '+p.seal[:16]+' | '+data['date']
        s.PageSetup.Zoom=False;s.PageSetup.FitToPagesWide=1;s.PageSetup.FitToPagesTall=1
        s.Activate();app.ActiveWindow.ScrollRow=1;app.ActiveWindow.ScrollColumn=1;s.Range('A1').Select()
        app.CalculateFull();c.native_save(b,path)
        return {'days':days,'thickness_microns':thickness,'cumulative_thickness_microns':[sum(thickness[:i+1]) for i in range(len(thickness))],'thickness_seed':thickness_seed,'replaced_climate_formulas':replaced,'columns':columns}
    finally:b.Close(SaveChanges=False)

def write_application(app,path,p,source_shape):
    from akz_stage6 import core as c,part2
    (c.write_application if p.payload.get('variant',0)==0 else part2.write_application)(app,path,p)
    book=c.native_open(app,path,read_only=False)
    try:
        for sheet,old in source_shape.items():
            for cell,formula in old['formulas'].items():book.Worksheets(sheet).Range(cell).Formula=formula
        app.CalculateFull();c.native_save(book,path)
    finally:book.Close(SaveChanges=False)

def validate_pair(app,act,application,p,proof,source_shape,application_shape):
    from akz_stage6 import core as c,part2
    d=p.payload;n=d.get('variant',0);a=load_workbook(act,data_only=False);values=load_workbook(act,data_only=True);b=load_workbook(application,data_only=False);bv=load_workbook(application,data_only=True)
    try:
        actual=shape(act)
        require(list(actual)==list(source_shape),'AKZ_SHEETS_CHANGED')
        for sheet,old in source_shape.items():
            now=actual[sheet]
            for key in ('merges','print_area','fonts','fills','borders'):require(old[key]==now[key],'AKZ_FORMAT_CHANGED: '+sheet+'/'+key)
            if sheet=='4':
                require(old['coating']==now['coating'],'AKZ_COATING_SYSTEM_CHANGED')
                replaced=set(proof['replaced_climate_formulas'])
                for cell,formula in old['formulas'].items():
                    if cell not in replaced:require(now['formulas'].get(cell)==formula,'AKZ_FORMULA_CHANGED: '+cell)
        application_actual=shape(application)
        require(list(application_actual)==list(application_shape),'AKZ_APPLICATION_SHEETS_CHANGED')
        for sheet,old in application_shape.items():
            for key in ('merges','print_area','fonts','fills','borders','formulas'):
                require(old[key]==application_actual[sheet][key],'AKZ_APPLICATION_FORMAT_CHANGED: '+sheet+'/'+key)
        for book in (a,values,b,bv):
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        require(cell.data_type!='e','AKZ_EXCEL_FORMULA_ERROR: '+str(cell.value))
                        require('RANDBETWEEN' not in str(cell.value),'AKZ_LEGACY_RANDOM_FORMULA_STOP')
        forbidden=('1/Б2','1/2025','РУССОКО','Велесстрой','2-ЭГ-2019','ВС.24.ПО','13-25','314-25','Журавлёв','Шляхтуров')
        for book in (a,b):
            text='\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)
            text+='\n'+str(list(book.defined_names.values()))
            require(not any(token in text for token in forbidden),'AKZ_FOREIGN_REQUISITES_REMAIN')
        for cell,key in [('E5','contract'),('E6','customer'),('E7','object')]:require(a['4'][cell].value==d['passport'][key],'AKZ_CURRENT_REQUISITE_STOP: '+cell)
        require(a['4']['B44'].value==d['passport']['cipher'] or n==0 and a['4']['B44'].value.startswith(d['passport']['cipher']+','),'AKZ_CURRENT_REQUISITE_STOP: B44')
        require(a['4']['E4'].value==d['act_number']+' от ','AKZ_CURRENT_NUMBER_STOP')
        require(values['4']['H4'].value.date().isoformat()==d['date'],'AKZ_CURRENT_DATE_STOP')
        if n==2:require(a['4']['I27'].number_format=='0','AKZ_THICKNESS_CURRENCY_FORMAT_STOP')
        sheet=b.active;cached=bv.active;last=20 if n==0 else part2.layout(n).last;total=21 if n==0 else last+1
        for col,key in [('C','quantity'),('E','mass_total'),('G','area_total')]:
            require(sheet[f'{col}{total}'].value==application_shape[sheet.title]['formulas'][f'{col}{total}'],'AKZ_TOTAL_FORMULA_CHANGED')
            require(abs(Decimal(str(cached[f'{col}{total}'].value))-d['totals'][key])<=Decimal('.00000001'),'AKZ_VISIBLE_VOM_TOTAL_STOP')
        for index,col in enumerate(proof['columns']):
            layerday=values['4'][col+'22'].value.date().isoformat()
            record=proof['days'][layerday]
            for row,key in ((31,'surface_c'),(32,'air_c'),(33,'relative_humidity'),(34,'dew_point_c')):require(values['4'][col+str(row)].value==record['values'][key],'AKZ_DAY_VALUE_MISMATCH')
            require(values['4'][col+'28'].value==proof['thickness_microns'][index],'AKZ_DRY_THICKNESS_MISMATCH')
    finally:
        for book in (a,values,b,bv):book.close()
    # Normal native readonly reopen confirms the saved package and live formulas.
    native=[]
    for path in (act,application):
        book=c.native_open(app,path,read_only=True)
        try:
            app.CalculateFull();native.append({'path':str(path),'read_only':bool(book.ReadOnly),'normal_load':True,'pages':int(book.Worksheets(1).PageSetup.Pages.Count)})
        finally:book.Close(SaveChanges=False)
    return {'status':'PASS','native':native,'formulas':'PASS','structure_and_format':'PASS'}

def generate_groups(root,batch,variants,output_name,*,mode=MODE,snapshot=None):
    from akz_stage6 import core as c,part2
    require(mode==MODE,'AKZ_OFFICIAL_GENERATION_FORBIDDEN')
    cfg=validate_batch(root,batch,variants);path=settings_path(root);expected=S.sha(path)
    require(re.fullmatch(r's9_[0-9a-f]{12}',output_name) is not None,'AKZ_OUTPUT_ID_INVALID')
    output=S.no_links(Path(root)/'DEV/Stage9_AKZ_Auto');output.mkdir(exist_ok=True)
    final=S.no_links(output/output_name);require(not final.exists(),'AKZ_OUTPUT_EXISTS_NO_OVERWRITE')
    with settings_lock(root) as folder:
        require(S.sha(path)==expected,'AKZ_SETTINGS_CHANGED_AFTER_PREVIEW')
        run=Path(root)/'validation'/output_name
        if snapshot is None:snapshot=K.freeze(root,run,expected)
        frozen_cfg=K.verify_snapshot(root,run,snapshot)
        require(snapshot['sha256']==expected and K.digest(frozen_cfg)==K.digest(cfg),'AKZ_SNAPSHOT_CURRENT_SETTINGS_CONFLICT')
        cfg=frozen_cfg
        previews={str(n):part2.group_preview(batch,n) for n in variants}
        inputs={path:expected,Path(snapshot['path']):snapshot['sha256']}
        for p in previews.values():inputs.update({Path(k):v for k,v in p.payload['sources'].items()})
        with S.frozen(inputs):
            require(all(S.sha(k)==v for k,v in inputs.items()),'AKZ_INPUT_SHA_CONFLICT')
            stage=output/('.pending_'+uuid4().hex);stage.mkdir();groups={}
            try:
                with c.native_excel() as excel:
                    for n,p in previews.items():
                        number=int(n);pair=stage/('АКЗ'+n);pair.mkdir()
                        act=pair/('АКЗ №'+n+' TEST.xlsx');application=pair/('Приложение АКЗ №'+n+' TEST.xlsx')
                        template=p.payload['templates'];shutil.copyfile(template['act'],act);shutil.copyfile(template['application'],application)
                        original=shape(act);application_original=shape(application);proof=write_act(excel,act,p,cfg,folder)
                        write_application(excel,application,p,application_original)
                        validation=validate_pair(excel,act,application,p,proof,original,application_original)
                        receipt={'mode':MODE,'official_issue':False,'variant':number,'act_number':p.payload['act_number'],'date':p.payload['date'],
                                 'project':PROJECT,'settings_sha256':snapshot['sha256'],'settings_snapshot':snapshot,'templates':template,'proof':proof,'totals':{k:str(v) for k,v in p.payload['totals'].items()},
                                 'files':{act.name:S.sha(act),application.name:S.sha(application)},'validation':validation}
                        save(pair/'PAIR.json',receipt)
                        groups[n]={'status':'COMPLETE_TEST_PAIR','path':str(final/pair.name),'proof':proof,'validation':validation}
                require(all(S.sha(k)==v for k,v in inputs.items()),'AKZ_INPUT_CHANGED_DURING_GENERATION')
                complete={'status':'COMPLETE_ALL_SELECTED','mode':MODE,'official_issue':False,'groups':groups,'settings_sha256':snapshot['sha256'],'settings_snapshot':snapshot,
                          'receipt_sha256':{str(p.relative_to(stage)):S.sha(p) for p in stage.glob('*/PAIR.json')}}
                complete['seal']=S.digest(complete);save(stage/'COMPLETE.json',complete)
                os.rename(stage,final)
            except Exception as exc:
                save(stage/'FAILURE.json',{'status':'STOP_NOT_PUBLISHED','error':str(exc),'official_issue':False});raise
    return {'status':'COMPLETE_ALL_SELECTED','mode':MODE,'official_issue':False,'groups':groups,'path':str(final)}
