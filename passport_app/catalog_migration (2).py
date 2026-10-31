"""Explicit, guarded 4 -> 3 catalog migration; never invoked by loading a project."""
from hashlib import sha256
from io import BytesIO
import math
from pathlib import Path
import re
from uuid import uuid4
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from .catalog import (HEADERS, LEGACY_HEADERS, NS, catalog_guard, catalog_headers,
                      durable_write, read_catalog, sheet_part)
from .config import PassportError
from .excel import atomic_replace, excel_session, locked_file


def assess(data):
    """Read every row, retaining zero, blanks, cell types, styles and formats."""
    entries, last = read_catalog(data)
    book = openpyxl.load_workbook(BytesIO(data))
    try:
        sheet = book.active
        headers = catalog_headers(sheet)
        akz_column = len(headers)
        # Scan the complete paint column first: even a bad AKZ in an earlier row
        # must not hide later user paint values from the STOP diagnostic.
        paint = ([{'cell': f'C{r}', 'mark': sheet.cell(r, 1).value,
                   'value': sheet.cell(r, 3).value}
                  for r in range(2, last + 1) if sheet.cell(r, 3).value not in (None, '')]
                 if headers == LEGACY_HEADERS else [])
        if paint:
            raise PassportError(f'STOP: непустая Краска: {paint!r}')
        rows = []
        for r in range(2, last + 1):
            cell = sheet.cell(r, akz_column)
            value = cell.value
            if value not in (None, '') and (isinstance(value, bool)
                    or not isinstance(value, (int, float)) or not math.isfinite(value)
                    or value < 0 or int(value) != value or cell.data_type in ('f', 'e')):
                raise PassportError(f'STOP: {cell.coordinate}: недопустимый № АКЗ {value!r}; '
                                    'нужно целое неотрицательное число или пусто.')
            rows.append({'row': r, 'mark': sheet.cell(r, 1).value,
                         'profile': sheet.cell(r, 2).value, 'akz': value,
                         'akz_type': cell.data_type, 'akz_style': cell.style_id,
                         'akz_format': cell.number_format})
        return {'headers': list(headers), 'last_row': last, 'records': len(entries),
                'rows': rows, 'nonempty_paint': paint}
    finally:
        book.close()


def migrated_bytes(data):
    before = assess(data)
    if tuple(before['headers']) != LEGACY_HEADERS:
        raise PassportError('STOP: миграция требует исходный legacy справочник из четырёх столбцов.')
    if before['nonempty_paint']:
        raise PassportError(f"STOP: непустая Краска: {before['nonempty_paint']!r}")
    part, xml = sheet_part(data)
    root = ET.fromstring(xml)
    allowed = {'sheetPr', 'dimension', 'sheetViews', 'sheetFormatPr', 'cols',
               'sheetData', 'autoFilter', 'pageMargins'}
    if any(node.tag.rsplit('}', 1)[-1] not in allowed for node in root):
        raise PassportError('STOP: сложная структура листа; требуется отдельное правило миграции.')
    if re.search(r'<\w+:', xml) or root.find('s:sheetData', NS) is None:
        raise PassportError('STOP: неподдерживаемый XML листа.')
    # Formulas, names, filters and references may depend on removed coordinates.
    with ZipFile(BytesIO(data)) as archive:
        workbook_xml = archive.read('xl/workbook.xml').decode('utf8')
        workbook = ET.fromstring(workbook_xml)
        names = workbook.find('s:definedNames', NS)
        if names is not None:
            for name in names:
                if (name.get('name') != '_xlnm._FilterDatabase' or name.get('localSheetId') != '0'
                        or not re.fullmatch(r"'[^']+'!\$A\$1:\$D\$[1-9]\d*", name.text or '')):
                    raise PassportError('STOP: именованные диапазоны требуют отдельной миграции.')
        migrated_workbook_xml = re.sub(
            r'(<definedName\b[^>]*name="_xlnm\._FilterDatabase"[^>]*>[^<]*!\$A\$1:)\$D(\$\d+</definedName>)',
            r'\1$C\2', workbook_xml)
        if names is not None and any('$D$' in n.text for n in
                ET.fromstring(migrated_workbook_xml).find('s:definedNames', NS)):
            raise PassportError('STOP: диапазон автофильтра не перенесён.')
    if root.findall('.//s:f', NS):
        raise PassportError('STOP: формулы требуют отдельной миграции.')
    af = root.find('s:autoFilter', NS)
    if af is not None and len(af):
        raise PassportError('STOP: активные условия фильтра требуют отдельной миграции.')
    views = root.find('s:sheetViews', NS)
    if views is not None:
        for node in views.iter():
            for name in ('topLeftCell', 'activeCell', 'sqref'):
                if name in node.attrib and not re.fullmatch(r'[AB][1-9]\d*', node.get(name)):
                    raise PassportError('STOP: позиция просмотра требует отдельной миграции.')
    cell_pattern = re.compile(r'<c\b[^>]*\br="([A-Z]+)(\d+)"[^>]*(?:/>|>.*?</c>)', re.S)
    def shift_cell(match):
        if match[1] == 'C':
            return ''
        if match[1] == 'D':
            return match[0].replace(f'r="D{match[2]}"', f'r="C{match[2]}"', 1)
        return match[0]
    xml = cell_pattern.sub(shift_cell, xml)
    xml = re.sub(r'(<(?:dimension|autoFilter)\b[^>]*\bref=")A1:D(\d+)(")',
                 r'\1A1:C\2\3', xml)
    xml = xml.replace('spans="1:4"', 'spans="1:3"')
    def shift_col(match):
        node = ET.fromstring(match[0])
        first, last = int(node.get('min')), int(node.get('max'))
        if first != last or first not in (1, 2, 3, 4):
            raise PassportError('STOP: сложные диапазоны ширины столбцов.')
        if first == 3:
            return ''
        if first == 4:
            return re.sub(r'\b(min|max)="4"', r'\1="3"', match[0])
        return match[0]
    xml = re.sub(r'<col\b[^>]*/>', shift_col, xml)
    output = BytesIO()
    with ZipFile(BytesIO(data)) as old, ZipFile(output, 'w') as new:
        for info in old.infolist():
            contents = (xml.encode('utf8') if info.filename == part else
                        migrated_workbook_xml.encode('utf8') if info.filename == 'xl/workbook.xml'
                        else old.read(info.filename))
            new.writestr(info, contents)
        new.comment = old.comment
    result = output.getvalue()
    after = assess(result)
    if tuple(after['headers']) != HEADERS or before['rows'] != after['rows']:
        raise PassportError('STOP: значения/типы/форматы строк не сохранились.')
    _, old_xml = sheet_part(data)
    def cells(text):
        return {m[1] + m[2]: m[0] for m in cell_pattern.finditer(text)}
    old_cells, new_cells = cells(old_xml), cells(xml)
    for address, content in old_cells.items():
        if address.startswith('C'):
            continue
        target = 'C' + address[1:] if address.startswith('D') else address
        if new_cells.get(target) != content.replace(f'r="{address}"', f'r="{target}"', 1):
            raise PassportError(f'STOP: XML клетки {address} не сохранился.')
    with ZipFile(BytesIO(data)) as old, ZipFile(BytesIO(result)) as new:
        if old.namelist() != new.namelist() or any(old.read(n) != new.read(n)
                for n in old.namelist() if n not in (part, 'xl/workbook.xml')):
            raise PassportError('STOP: изменилась посторонняя часть workbook.')
    return result


def native_validate(path):
    """Normal-load, read-only, macros/links/events disabled; no repair or Save."""
    path = Path(path).resolve()
    data = path.read_bytes()
    expected = assess(data)
    with excel_session() as excel:
        book = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True,
                                   Notify=False, AddToMru=False, CorruptLoad=0)
        try:
            sheet = book.Worksheets(1)
            width = len(expected['headers'])
            if not book.ReadOnly or book.Worksheets.Count != 1 or sheet.UsedRange.Columns.Count != width:
                raise PassportError('STOP: native Excel: другая структура/режим открытия.')
            values = sheet.Range(sheet.Cells(1, 1), sheet.Cells(expected['last_row'], width)).Value2
            if tuple(values[0]) != tuple(expected['headers']):
                raise PassportError('STOP: native Excel: заголовки отличаются.')
            for row in expected['rows']:
                actual = values[row['row'] - 1]
                wanted = (row['mark'], row['profile'], row['akz'])
                if (actual[0], actual[1], actual[-1]) != wanted:
                    raise PassportError(f"STOP: native Excel: строка {row['row']} отличается.")
            result = {'path': str(path), 'status': 'PASS', 'excel_version': str(excel.Version),
                      'corrupt_load': 0, 'read_only': bool(book.ReadOnly), 'save_called': False,
                      'records_checked': expected['records'], 'columns': width,
                      'sha256': sha256(data).hexdigest()}
        finally:
            book.Close(SaveChanges=False)
    if path.read_bytes() != data:
        raise PassportError('STOP: native Excel изменил проверяемый файл.')
    return result


def _safe_path(path):
    path = Path(path).absolute()
    if path.is_symlink() or path.resolve().parent != path.parent.resolve():
        raise PassportError('STOP: небезопасный путь справочника.')
    return path


def migrate_catalog(path, expected_sha, *, validator=native_validate):
    path = _safe_path(path)
    with catalog_guard(path), locked_file(path, writable=True) as data:
        if sha256(data).hexdigest() != expected_sha:
            raise PassportError('STOP: справочник изменился после текущего assessment.')
        # All validation before any backup/temp creation; nonempty paint never discarded.
        after = migrated_bytes(data)
        backups = path.parent / '.stage3-backups'
        if backups.is_symlink() or backups.resolve().parent != path.parent.resolve():
            raise PassportError('STOP: небезопасная папка backup.')
        backups.mkdir(exist_ok=True)
        token = uuid4().hex
        backup = backups / f'{path.stem}.schema4to3.{token}.xlsx'
        durable_write(backup, data)
        if backup.read_bytes() != data:
            raise PassportError('STOP: проверка backup не пройдена.')
        temporary = path.with_name(f'.stage3-schema-{token}.tmp.xlsx')
        try:
            durable_write(temporary, after)
            if temporary.read_bytes() != after:
                raise PassportError('STOP: временный файл отличается.')
            native = validator(temporary)
            if native.get('status') != 'PASS' or temporary.read_bytes() != after:
                raise PassportError('STOP: временный файл не прошёл native validation.')
            result = {'path': str(path), 'source_sha256': expected_sha,
                      'backup': str(backup), 'backup_sha256': sha256(data).hexdigest(),
                      'new_sha256': sha256(after).hexdigest(), 'native_excel': native,
                      'assessment_before': assess(data), 'assessment_after': assess(after)}
            atomic_replace(temporary, path)
            return result
        finally:
            temporary.unlink(missing_ok=True)


def rollback_catalog(path, backup, expected_current_sha, expected_backup_sha):
    """Explicit rollback only; guard current user work and restore exact backup bytes."""
    path, backup = _safe_path(path), Path(backup).resolve()
    data = backup.read_bytes()
    if sha256(data).hexdigest() != expected_backup_sha:
        raise PassportError('STOP: backup изменён.')
    assess(data)
    with catalog_guard(path), locked_file(path, writable=True) as current:
        if sha256(current).hexdigest() != expected_current_sha:
            raise PassportError('STOP: рабочий справочник изменён; rollback отменён.')
        temporary = path.with_name(f'.stage3-rollback-{uuid4().hex}.tmp.xlsx')
        try:
            durable_write(temporary, data)
            native_validate(temporary)
            if temporary.read_bytes() != data:
                raise PassportError('STOP: временный rollback отличается.')
            atomic_replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)
    return {'status': 'PASS', 'restored_sha256': expected_backup_sha}
