"""Verify persisted staging files before replacing any working document."""
import re
from copy import copy
from io import BytesIO
import openpyxl
from openpyxl.formula.translate import Translator
from .config import PassportError
from .excel import load_book
from .vom import FIELDS


def _formula(value):
    if isinstance(value, str) and value.startswith('='):
        return re.sub(r'\[\[#This Row\],\[([^]]+)\]\]', r'[@[\1]]', value)
    return value


def verify_saved(registry_path, backup_path, appendix_path, template_bytes, plan, new_row, requisites):
    books = []
    try:
        def read(path, values=False):
            book = load_book(path, values=values)
            books.append(book)
            return book
        original = read(backup_path)
        registry = read(registry_path)
        original_sheet = original[plan.registry.sheet]
        sheet = registry[plan.registry.sheet]
        changed = {plan.registry.columns[k] for k in ('number', 'mass', 'date', 'connections')}
        for col in range(1, plan.registry.width + 1):
            source, target = original_sheet.cell(plan.project.first_row, col), sheet.cell(new_row, col)
            expected = source.value
            if source.data_type == 'f':
                expected = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
            if col not in changed and _formula(target.value) != _formula(expected):
                raise PassportError(f'Изменилось постоянное поле или формула новой строки: столбец {col}.')
            same_book_source = sheet.cell(plan.project.first_row, col)
            for prop in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
                if copy(getattr(target, prop)) != copy(getattr(same_book_source, prop)):
                    raise PassportError(f'Не сохранилось оформление новой строки: {target.coordinate}, {prop}.')
        if sheet.row_dimensions[new_row].height != sheet.row_dimensions[plan.project.first_row].height:
            raise PassportError('Не сохранилась высота новой строки реестра.')
        if registry.sheetnames != original.sheetnames:
            raise PassportError('Изменилась структура листов реестра.')
        for old_sheet in original:
            new_sheet = registry[old_sheet.title]
            for row in old_sheet:
                for cell in row:
                    persisted = new_sheet[cell.coordinate]
                    if _formula(cell.value) != _formula(persisted.value):
                        raise PassportError(f'Изменилось существующее поле реестра: {old_sheet.title}!{cell.coordinate}.')
                    for prop in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
                        if copy(getattr(cell, prop)) != copy(getattr(persisted, prop)):
                            raise PassportError(f'Изменилось оформление существующей ячейки: {old_sheet.title}!{cell.coordinate}, {prop}.')
            if {k: v.width for k, v in old_sheet.column_dimensions.items()} != {k: v.width for k, v in new_sheet.column_dimensions.items()}:
                raise PassportError(f'Изменились ширины столбцов реестра: {old_sheet.title}.')
            for row, dimensions in old_sheet.row_dimensions.items():
                saved = new_sheet.row_dimensions[row]
                if dimensions.height != saved.height or dimensions.hidden != saved.hidden:
                    raise PassportError(f'Изменились свойства существующей строки реестра: {old_sheet.title}!{row}.')
        values = read(appendix_path, True)
        formulas = read(appendix_path)
        template = openpyxl.load_workbook(BytesIO(template_bytes), data_only=False)
        books.append(template)
        from .tables import locate_layout
        layout = locate_layout(formulas, appendix=True)
        output = values[layout.sheet]
        total = layout.start + len(plan.vom.rows)
        if layout.total_row != total or output.max_row != total:
            raise PassportError('Количество сохранённых строк приложения не совпадает с ВОМ.')
        for offset, expected in enumerate(plan.vom.rows):
            for i, field in enumerate(FIELDS):
                cell = output.cell(layout.start + offset, layout.columns[field])
                value = float(expected[i]) if i >= 2 else expected[i]
                if cell.value != value:
                    raise PassportError(f'Сохранённая строка приложения отличается от ВОМ: {cell.coordinate}.')
        for field, expected in [('quantity', plan.vom.quantity), ('mass_total', plan.vom.mass), ('area_total', plan.vom.area)]:
            cell = output.cell(total, layout.columns[field])
            if not isinstance(cell.value, (float, int)) or abs(cell.value - float(expected)) > 1e-8:
                raise PassportError(f'Не совпадает сохранённое «Итого»: {field}.')
        mincol, maxcol = min(layout.columns.values()), max(layout.columns.values())
        fs = formulas[layout.sheet]
        if not all(fs.cell(total, col).font.bold for col in range(mincol, maxcol + 1)):
            raise PassportError('Сохранённая строка «Итого» не полностью выделена жирным.')
        header_text = [str(cell.value) for row in fs.iter_rows(max_row=layout.header_row - 1) for cell in row if cell.value]
        numbers = [re.search(r'№\s*(\S+)\s+от\s+(\d{2}\.\d{2}\.\d{4})\s*г\.', text) for text in header_text]
        object_ok = any(re.match(r'^\s*Объект\s*:', text, re.I) and
                        re.sub(r'^\s*Объект\s*:\s*', '', text, flags=re.I).strip() == str(requisites['object']).strip() for text in header_text)
        if not any(m and m[1] == requisites['number'] and m[2] == requisites['date'].strftime('%d.%m.%Y') for m in numbers) or not object_ok:
            raise PassportError('Реквизиты сохранённой шапки приложения не соответствуют новой строке реестра.')
        reference_layout = locate_layout(template, appendix=True)
        reference = template[reference_layout.sheet]
        if {k: v.width for k, v in fs.column_dimensions.items()} != {k: v.width for k, v in reference.column_dimensions.items()}:
            raise PassportError('Изменились ширины столбцов приложения.')
        if fs.page_setup != reference.page_setup or fs.page_margins != reference.page_margins or fs.print_options != reference.print_options:
            raise PassportError('Изменились параметры страницы приложения.')
        if set(map(str, fs.merged_cells.ranges)) != set(map(str, reference.merged_cells.ranges)):
            raise PassportError('Изменились объединения ячеек приложения.')
        for row in range(layout.start, total):
            for field in FIELDS:
                col = layout.columns[field]
                source = reference.cell(reference_layout.start, reference_layout.columns[field])
                target = fs.cell(row, col)
                if any(copy(getattr(source, prop)) != copy(getattr(target, prop)) for prop in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection')):
                    raise PassportError(f'Изменилось оформление строки приложения: {target.coordinate}.')
    finally:
        for book in books:
            book.close()
