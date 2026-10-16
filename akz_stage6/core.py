"""Explicit, fail-closed DEV preview and transactional TEST pair generation.

No production registry parameter or discovery. Native Excel preserves the exact
project templates. Numeric sources are exclusively visible rows of the chosen VOM.
"""
from contextlib import contextmanager
from dataclasses import dataclass, asdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path
import shutil
import stat
import uuid
import time
import threading
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from passport_app.registry import read_registry, select_project
from passport_app.tables import norm
from passport_app.excel import sha256

ROOT = Path(__file__).resolve().parents[1]
DEV = ROOT / 'DEV'
REGISTRY = DEV / 'ВЗМК_DEV.xlsx'
FIXTURES = DEV / 'Stage6_Fixtures'
OUTPUT = DEV / 'Stage6_Выход'
PROJECT = ROOT.parent / 'Гараж казахстан'
ACT = PROJECT / 'Паспорт/Акт № 4-2026.xlsx'
APPLICATION = PROJECT / 'Паспорт/1_2026 - Приложение к АКЗ.xlsx'
PINS = {ACT: '7f6420ac50149805b46454852d62ad68fa9625c53b7c9d375460769b6ae0f8ac',
        APPLICATION: '704d92eb0c5ea87c3951f6643488aa9100cd18f99e177d21ddc85f8843938d6f'}
HISTORICAL_DATE = '2026-10-06'  # accepted audit, never used as output date

class Stage6Error(ValueError):
    pass

class FileOperationError(Stage6Error):
    def __init__(self, code, path, operation):
        self.code = code
        self.path = str(Path(path).resolve())
        self.operation = operation
        messages = {
            'FILE_LOCKED': 'Файл АКЗ занят или заблокирован. Закройте его в Excel и повторите операцию.',
            'ACCESS_DENIED': 'Нет доступа к файлу АКЗ. Проверьте права доступа и повторите операцию.',
            'INVALID_XLSX': 'Файл не является исправной книгой XLSX. Проверьте выбранный файл.',
            'OTHER_ERROR': 'Не удалось выполнить операцию с файлом АКЗ.',
        }
        if code == 'FILE_LOCKED' and Path(path).name == 'ТЕСТ Приложение к АКЗ №0.xlsx':
            messages[code] = 'Файл приложения АКЗ занят или заблокирован. Закройте файл в Excel и повторите операцию.'
        super().__init__(f'{code}: {messages[code]}\nОперация: {operation}\nФайл: {self.path}')

def probe_file_access(path, *, read_only):
    """Ask Windows, never infer a lock from Excel's ambiguous format error."""
    import win32file
    import win32con
    import pywintypes
    access = win32con.GENERIC_READ | (0 if read_only else win32con.GENERIC_WRITE)
    try:
        handle = win32file.CreateFile(str(Path(path).resolve()), access,
            win32con.FILE_SHARE_READ | win32con.FILE_SHARE_WRITE | win32con.FILE_SHARE_DELETE,
            None, win32con.OPEN_EXISTING, 0, None)
    except (OSError, pywintypes.error) as exc:
        code = {32: 'FILE_LOCKED', 33: 'FILE_LOCKED', 5: 'ACCESS_DENIED'}.get(exc.winerror, 'OTHER_ERROR')
        raise FileOperationError(code, path, 'проверка доступа') from exc
    else:
        handle.Close()

def excel_file_failure(path, exc, *, read_only, operation):
    # Recheck after failure too: another process may acquire a lock after preflight.
    probe_file_access(path, read_only=read_only)
    code = {32: 'FILE_LOCKED', 33: 'FILE_LOCKED', 5: 'ACCESS_DENIED'}.get(getattr(exc, 'winerror', None))
    if code is None and Path(path).suffix.lower() == '.xlsx':
        try:
            with ZipFile(path) as archive:
                if archive.testzip() is not None:
                    raise BadZipFile('CRC mismatch')
                for member in ('[Content_Types].xml', 'xl/workbook.xml'):
                    ElementTree.fromstring(archive.read(member))
        except (BadZipFile, KeyError, ElementTree.ParseError):
            code = 'INVALID_XLSX'
    raise FileOperationError(code or 'OTHER_ERROR', path, operation) from exc

def native_save(book, path):
    try:
        book.Save()
    except Exception as exc:
        excel_file_failure(path, exc, read_only=False, operation='сохранение')

def error_message(exc):
    return str(exc) if isinstance(exc, Stage6Error) else f'OTHER_ERROR: {exc}'

def require(condition, message):
    if not condition:
        raise Stage6Error(message)

def digest(data):
    return hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True,
                                    default=str).encode('utf-8')).hexdigest()

def number(value, where, *, positive=False):
    require(not isinstance(value, bool) and isinstance(value, (int, float, Decimal)),
            f'NUMBER_MISSING: {where}: {value!r}')
    try:
        n = Decimal(str(value))
    except InvalidOperation as exc:
        raise Stage6Error(f'NUMBER_INVALID: {where}') from exc
    require(n.is_finite() and (n > 0 if positive else n >= 0), f'NUMBER_INVALID: {where}')
    return n

def text(value, where):
    require(isinstance(value, str) and bool(value.strip()) and not value.startswith(('=', '#')),
            f'TEXT_INVALID: {where}')
    return value  # exact marks; no transliteration, trimming, or fuzzy matching

@dataclass(frozen=True)
class ClimateDay:
    day: str
    air_c: Decimal | None = None
    relative_humidity: Decimal | None = None
    surface_c: Decimal | None = None
    dew_point_c: Decimal | None = None
    confirmation: str = 'MISSING'

    def validate(self):
        date.fromisoformat(self.day)
        require(self.confirmation in ('MISSING', 'TEST'), 'CLIMATE_CONFIRMATION_UNSUPPORTED')
        readings = (self.air_c, self.relative_humidity, self.surface_c, self.dew_point_c)
        if all(x is None for x in readings):
            require(self.confirmation == 'MISSING', 'CLIMATE_STATUS_CONFLICT')
            return
        require(all(x is not None for x in readings) and self.confirmation == 'TEST',
                'CLIMATE_INCOMPLETE: требуется один полный TEST-набор на дату')
        require(all(isinstance(x, Decimal) and x.is_finite() for x in readings), 'CLIMATE_INVALID')
        require(0 <= self.relative_humidity <= 100, 'CLIMATE_HUMIDITY_INVALID')
        require(self.surface_c >= self.dew_point_c + 3, 'CLIMATE_DEW_MARGIN_STOP')

def register_climate(days, record):
    """One immutable record per calendar day; conflict STOP, never random."""
    record.validate()
    require(record.day not in days or days[record.day] == record, 'CLIMATE_DAY_CONFLICT')
    days[record.day] = record

@contextmanager
def native_excel():
    # pythoncom initializes the Python main thread on import. Keep that apartment
    # alive across previews; uninitializing it while caller-held sheet proxies
    # still exist disconnects those proxies and can crash the next session.
    require(threading.current_thread() is threading.main_thread(), 'EXCEL_MAIN_THREAD_REQUIRED')
    import pythoncom
    import win32com.client
    app = None
    try:
        # Retry only COM startup, before any workbook operation or mutation.
        for attempt in range(3):
            try:
                app = win32com.client.DispatchEx('Excel.Application')
                app.Visible = False
                app.DisplayAlerts = False
                app.AutomationSecurity = 3
                app.EnableEvents = False
                app.AskToUpdateLinks = False
                _ = app.Workbooks.Count
                break
            except Exception:
                try:
                    if app is not None: app.Quit()
                except Exception:
                    pass
                app = None
                if attempt == 2: raise
                time.sleep(0.3)
        yield app
    finally:
        try:
            if app is not None:
                try:
                    app.Workbooks.Close()
                finally:
                    app.Quit()
        except Exception:
            pass  # never hide the original failure; sources were never saved
        finally:
            app = None

def native_open(app, path, *, read_only):
    probe_file_access(path, read_only=read_only)
    try:
        book = app.Workbooks.Open(Filename=str(Path(path).resolve()), UpdateLinks=0,
                                 ReadOnly=read_only, IgnoreReadOnlyRecommended=True,
                                 Notify=False, AddToMru=False, CorruptLoad=0)
    except Exception as exc:
        excel_file_failure(path, exc, read_only=read_only, operation='открытие')
    try:
        require(bool(book.ReadOnly) == read_only, 'NATIVE_READ_ONLY_MISMATCH')
        require(not book.LinkSources(1) and not int(book.Connections.Count)
                and not any(int(s.QueryTables.Count) for s in book.Worksheets), 'EXTERNAL_DATA_STOP')
        # CorruptLoad=0 is the normal load path, never xlRepairFile/xlExtractData.
        # Excel exposes no portable Workbook.RepairMode property through this COM API.
        return book
    except Exception:
        book.Close(SaveChanges=False)
        raise

def templates():
    for path, expected in PINS.items():
        require(path.is_file() and sha256(path) == expected, f'TEMPLATE_SHA_CHANGED_STOP: {path}')
    a = load_workbook(ACT, data_only=False)
    b = load_workbook(APPLICATION, data_only=False)
    try:
        require(a.sheetnames == ['4'] and b.sheetnames == ['Приложение АКЗ'], 'TEMPLATE_SHEETS_CHANGED')
        s = a['4']; t = b.active
        require(s['E4'].value == '1/2026 от ' and s['G8'].value == '=E4'
                and s['I8'].value == '=H4', 'TEMPLATE_BINDING_STOP')
        require(s['H4'].value.date() == date(2026, 9, 8)
                and t['A1'].value == 'Приложение № 1/2026 к акту АКЗ № 1/2026  от 08.09.2026 г.',
                'TEMPLATE_PAIR_CONFLICT')
        require(s['E17'].value == 1 and 'Pirocor Protect 057' in s['E18'].value
                and s['E21'].value == 'RAL 7004' and s['E25'].value == 80,
                'TEMPLATE_COATING_CHANGED')
        return {'act': str(ACT), 'application': str(APPLICATION),
                'material': s['E18'].value, 'ral': s['E21'].value,
                'required_dry_microns': s['E25'].value, 'layers': 1}
    finally:
        a.close(); b.close()

def passports(project='Гараж Казахстан'):
    require(project == 'Гараж Казахстан', 'PROJECT_NOT_IMPLEMENTED')
    registry = read_registry(REGISTRY)
    select_project(registry, project)  # accepted DEV selector validates project/first/number uniqueness
    return registry.projects[project]

def current_passport(project, number_):
    rows = passports(project)
    selected = [r for r in rows if r['number'] == number_]
    require(len(selected) == 1, 'PASSPORT_AMBIGUOUS_OR_MISSING')
    rn = selected[0]['row']
    b = load_workbook(REGISTRY, data_only=True)
    try:
        s = b['Паспорт качества']
        mapping = {}
        for c in s[1]:
            mapping.setdefault(norm(c.value), []).append(c.column)
        fields = {'date': 'Дата Акз', 'contract': 'Договор №', 'customer': 'Заказчик',
                  'object': 'Наименование объекта', 'cipher': 'шифр КМД'}
        result = {'project': project, 'number': number_, 'row': rn, 'mapping': {}}
        for k, label in fields.items():
            cols = mapping.get(norm(label), [])
            require(len(cols) == 1, f'DEV_HEADER_MAPPING_STOP: {label}')
            c = s.cell(rn, cols[0])
            result['mapping'][k] = c.coordinate
            if k == 'date':
                require(isinstance(c.value, datetime), 'DEV_AKZ_DATE_MISSING')
                result[k] = c.value.date().isoformat()
            else:
                result[k] = text(c.value, c.coordinate)
        return result
    finally:
        b.close()

def catalog(path):
    path = Path(path).resolve()
    require(path.is_relative_to(FIXTURES.resolve()) and path.suffix.lower() == '.xlsx',
            'DEV_FIXTURE_CATALOG_ONLY')
    b = load_workbook(path, data_only=False)
    try:
        require(len(b.sheetnames) == 1, 'CATALOG_SHEET_AMBIGUOUS')
        s = b.active
        require([c.value for c in s[1]][:3] == ['Марка', 'Профиль', '№ АКЗ'], 'CATALOG_SCHEMA_STOP')
        result = {}
        for rn, row in enumerate(s.iter_rows(min_row=2, max_col=3, values_only=True), 2):
            if all(x is None for x in row):
                continue
            mark = text(row[0], f'catalog row {rn}')
            require(mark not in result, f'CATALOG_DUPLICATE_CONFLICT: {mark}')
            assignment = row[2]
            if assignment is None or isinstance(assignment, str) and not assignment.strip():
                assignment = None
            else:
                n = number(assignment, f'catalog C{rn}')
                require(n == n.to_integral() and n in (0, 1, 2, 3), 'AKZ_ASSIGNMENT_INVALID')
                assignment = int(n)
            result[mark] = {'assignment': assignment, 'row': rn}
        return result
    finally:
        b.close()

def visible_vom(path):
    """Native two-line header mapping; retain all raw numeric values until group validation."""
    path = Path(path).resolve()
    require(path.suffix.lower() in ('.xls', '.xlsx'), 'VOM_FORMAT_STOP')
    require(path.is_relative_to(PROJECT.resolve()) or path.is_relative_to(FIXTURES.resolve()),
            'VOM_SOURCE_OUTSIDE_APPROVED_ROOTS')
    before = sha256(path)
    with native_excel() as app:
        book = native_open(app, path, read_only=True)
        try:
            found = []
            for s in book.Worksheets:
                if s.Visible != -1:
                    continue
                used = s.UsedRange
                values = used.Value2
                if not isinstance(values, tuple):
                    continue
                for offset, raw in enumerate(values[:30]):
                    row = [norm(x) for x in raw]
                    if row.count('отправочная марка') != 1:
                        continue
                    aliases = {'mark': 'отправочная марка', 'name': 'наименование марки', 'quantity': 'количество'}
                    cols = {}
                    for f, label in aliases.items():
                        require(row.count(label) == 1, f'VOM_HEADER_AMBIGUOUS: {label}')
                        cols[f] = row.index(label)
                    require(offset+1 < len(values), 'VOM_SUBHEADER_MISSING')
                    sub = [norm(x) for x in values[offset+1]]
                    for f, label in [('mass', 'масса'), ('area', 'площадь')]:
                        candidates = [i for i, v in enumerate(row) if v.startswith(label)]
                        require(len(candidates) == 1, f'VOM_HEADER_AMBIGUOUS: {label}')
                        i = candidates[0]
                        require(sub[i] in ('1 дет', '1 шт', '1 изделие') and sub[i+1] == 'всех',
                                f'VOM_SUBHEADER_CHANGED: {label}')
                        cols[f+'_unit'], cols[f+'_total'] = i, i+1
                    found.append((s, used.Row, used.Column, values, offset, cols))
            require(len(found) == 1, 'VOM_SHEET_AMBIGUOUS')
            s, start, start_col, values, header, cols = found[0]
            rows, hidden = [], []
            for offset in range(header+2, len(values)):
                raw = values[offset]
                mark = raw[cols['mark']]
                if mark is None or norm(mark) in ('', 'итого', 'всего'):
                    continue
                rn = start + offset
                if bool(s.Rows(rn).Hidden):
                    hidden.append(rn)
                    continue
                record = {f: raw[i] for f, i in cols.items()}
                text(record['mark'], f'VOM row {rn}')
                record['row'] = rn
                rows.append(record)
            require(rows, 'VOM_NO_VISIBLE_ROWS')
            result = {'path': str(path), 'sheet': s.Name, 'rows': rows, 'hidden_rows': hidden,
                      'synthetic_source': any('SYNTHETIC' in str(x).upper() for row in values[:30] for x in row)
                                          or 'SYNTHETIC' in path.name.upper(),
                      'autofilter': s.AutoFilter.Range.Address if s.AutoFilterMode else None,
                      'columns': {f: start_col+i for f,i in cols.items()},
                      'adapter': 'native Excel XLS/XLSX read-only, links/events/macros disabled'}
        finally:
            book.Close(SaveChanges=False)
    require(sha256(path) == before, 'VOM_CHANGED_DURING_READ')
    return result

def matching(vom, assignments):
    selected, excluded = [], []
    for row in vom['rows']:
        assignment = assignments.get(row['mark'])
        n = assignment['assignment'] if assignment is not None else None
        if n != 0:
            excluded.append({'row': row['row'], 'mark': row['mark'], 'assignment': n,
                             'reason': 'NO_ASSIGNMENT' if n is None else 'OTHER_VARIANT_NOT_IMPLEMENTED'})
            continue
        r = dict(row)
        r['name'] = text(r['name'], f'VOM row {r["row"]} name')
        for f in ('quantity', 'mass_unit', 'mass_total', 'area_unit', 'area_total'):
            r[f] = number(r[f], f'VOM row {r["row"]} {f}', positive=True)
        require(r['quantity'] == r['quantity'].to_integral(), 'VOM_QUANTITY_INVALID')
        # No rounding policy is invented: only floating-point transport noise is tolerated.
        for f in ('mass', 'area'):
            expected = r['quantity'] * r[f+'_unit']
            require(abs(expected-r[f+'_total']) <= Decimal('0.00000001'),
                    f'{f.upper()}_CONFLICT_STOP: row {r["row"]}: unit * qty != shipment total')
        selected.append({**r, 'catalog_row': assignment['row']})
    require(selected, 'NO_AKZ_ZERO_ROWS: пустое назначение не означает 0')
    totals = {f: sum((r[f] for r in selected), Decimal(0))
              for f in ('quantity', 'mass_total', 'area_total')}
    return selected, excluded, totals

@dataclass(frozen=True)
class Preview:
    payload: dict
    seal: str

def preview(project, passport, vom_path, catalog_path, *, synthetic=False, climate=None):
    require(isinstance(synthetic, bool), 'SYNTHETIC_MODE_INVALID')
    paths = [REGISTRY, ACT, APPLICATION, Path(vom_path).resolve(), Path(catalog_path).resolve()]
    before = {str(p): sha256(p) for p in paths}
    template = templates()
    current = current_passport(project, passport)
    require(not synthetic or Path(vom_path).resolve().is_relative_to(FIXTURES.resolve()),
            'SYNTHETIC_VOM_MUST_BE_ISOLATED')
    climate = climate or ClimateDay(current['date'])
    climate.validate()
    require(climate.day == current['date'], 'CLIMATE_DATE_CONFLICT')
    vom = visible_vom(vom_path)
    require(synthetic or not vom['synthetic_source'], 'SYNTHETIC_SOURCE_REQUIRES_SYNTHETIC_MODE_STOP')
    selected, excluded, totals = matching(vom, catalog(catalog_path))
    require(all(sha256(p) == h for p,h in before.items()), 'INPUT_CHANGED_DURING_PREVIEW')
    payload = {'mode': 'TEST_ONLY', 'synthetic': synthetic, 'passport': current, 'templates': template,
               'act_number': ('SYNTHETIC-TEST ' if synthetic else '') + passport,
               'date': current['date'], 'sources': before, 'vom': vom,
               'selected': selected, 'excluded': excluded, 'totals': totals,
               'climate': asdict(climate), 'official': 'STOP',
               'warnings': ['ТЕСТ / ЧЕРНОВИК. Не подтверждает выполнение АКЗ.',
                            'Назначения № АКЗ в DEV-каталоге только тестовые.',
                            'Не заполнены фактические измерения, партии, даты/время нанесения и контроль.',
                            'Климат не подтверждён; алгоритм точки росы не выбран.',
                            f'DEV дата АКЗ {current["date"]}; исторический акт 9/2026: {HISTORICAL_DATE}.',
                            'Выбор ВОМ явный. DEV имеет 9/2026 последним; на диске есть 10/2026.']}
    return Preview(payload, digest(payload))

def verify_preview(p):
    require(isinstance(p, Preview) and digest(p.payload) == p.seal, 'PREVIEW_CHANGED_STOP')
    require(p.payload['mode'] == 'TEST_ONLY' and p.payload['official'] == 'STOP', 'OFFICIAL_MODE_STOP')
    for path, expected in p.payload['sources'].items():
        require(sha256(path) == expected, f'INPUT_CHANGED_AFTER_PREVIEW_STOP: {path}')

FACTUAL_CELLS = ('G11','B13','E13','I13','E15','E19','E20','E22','E23','E24','E26','E27','E28',
                 'E30','E31','E32','E33','E34','G37','E39','E40','E41','E42','D49')
FACTUAL_PLACEHOLDERS = {cell: ('?' if cell in ('G11','E15','G37') else 'НЕТ ДАННЫХ / ТЕСТ')
                       for cell in FACTUAL_CELLS}

def write_act(app, path, p):
    b = native_open(app, path, read_only=False)
    try:
        s=b.Worksheets('4'); data=p.payload; passport=data['passport']
        s.Range('B3').Value2 = 'ТЕСТ / ЧЕРНОВИК — АКЗ НЕ ПОДТВЕРЖДЕНО'
        s.Range('E4').Value2 = data['act_number'] + ' от '
        # Value2 serial avoids COM timezone conversion of Python datetime.
        # Keep the template's native date format (portable built-in style 14).
        s.Range('H4').Value2 = (date.fromisoformat(data['date'])-date(1899,12,30)).days
        for c,k in [('E5','contract'),('E6','customer'),('E7','object'),('B44','cipher')]:
            s.Range(c).Value2 = passport[k]
        s.Range('E7').WrapText=True
        s.Rows(7).RowHeight=max(s.Rows(7).RowHeight,56)
        s.Range('G8').Formula = '=E4'; s.Range('I8').Formula = '=H4'
        for c in FACTUAL_CELLS:
            s.Range(c).Value2 = FACTUAL_PLACEHOLDERS[c]
        s.Range('B43').Value2 = 'ТЕСТ: фактические показатели не заполнены. Выдача запрещена.'
        s.Range('B43').WrapText=True
        # Keep original 51 pt height: the template uses 20 pt type here.
        climate=data['climate']
        if climate['confirmation'] == 'TEST':
            s.Range('E30').Value2 = (date.fromisoformat(climate['day'])-date(1899,12,30)).days
            for c,k in [('E31','surface_c'),('E32','air_c'),('E33','relative_humidity'),('E34','dew_point_c')]:
                s.Range(c).Value2 = float(climate[k])
        s.PageSetup.PrintArea = '$B$2:$K$50'
        s.PageSetup.CenterHeader = 'ТЕСТ / ЧЕРНОВИК' + (' — SYNTHETIC' if data['synthetic'] else '')
        s.PageSetup.CenterFooter = 'ACT-ID ' + p.seal[:16] + ' | ' + data['date']
        s.PageSetup.FitToPagesWide=1; s.PageSetup.FitToPagesTall=1; s.PageSetup.Zoom=False
        s.Calculate()
        native_save(b, path)
    finally:
        b.Close(SaveChanges=False)

def write_application(app, path, p):
    """Resize the AKZ-0 appendix to the exact number of selected VOM rows.

    Only the OUTPUT workbook is edited; the pinned template is never changed.
    A physical Excel row insert/delete moves the styled 'Итого' row.
    """
    b=native_open(app,path,read_only=False)
    try:
        s=b.Worksheets('Приложение АКЗ'); data=p.payload
        count=len(data['selected'])
        require(count>=1 and 4+count<=1048576,'AKZ_APPLICATION_EXCEL_ROW_LIMIT')
        first,last,original_total=4,20,21
        total=first+count
        if total>original_total:
            s.Rows(f'{original_total}:{total-1}').Insert()
            # New rows inherit the exact accepted body-row formatting.
            s.Range(f'A{last}:G{last}').Copy()
            s.Range(f'A{last+1}:G{total-1}').PasteSpecial(Paste=-4122)
            app.CutCopyMode=False
            s.Rows(f'{last+1}:{total-1}').RowHeight=s.Rows(last).RowHeight
        elif total<original_total:
            s.Rows(f'{total}:{last}').Delete()
        s.Rows(f'1:{total}').Hidden=False
        s.Range('A1').Value2 = ('ТЕСТ / ЧЕРНОВИК' + (' — SYNTHETIC' if data['synthetic'] else '') + '\n'
            + f'Приложение № {data["act_number"]} к акту АКЗ № {data["act_number"]} от '
            + date.fromisoformat(data['date']).strftime('%d.%m.%Y') + ' г.')
        s.Range('A1').WrapText=True; s.Rows(1).RowHeight=38
        s.Range(f'A{first}:G{total-1}').ClearContents()
        fields=('mark','name','quantity','mass_unit','mass_total','area_unit','area_total')
        for rn,r in enumerate(data['selected'],first):
            s.Range(f'A{rn}:G{rn}').Value2 = (tuple(float(r[f]) if isinstance(r[f],Decimal) else r[f] for f in fields),)
        for col,field in [('C','quantity'),('D','mass_unit'),('E','mass_total'),('F','area_unit'),('G','area_total')]:
            values=[r[field] for r in data['selected']]
            if field in data['totals']:values.append(data['totals'][field])
            # Only E/G totals: two display decimals. Retain full numeric precision.
            places=2 if col in ('E','G') else (0 if field=='quantity' else max(max(0,-v.as_tuple().exponent) for v in values))
            separator=str(app.DecimalSeparator)
            s.Range(f'{col}{first}:{col}{total}').NumberFormatLocal='0'+(separator+'0'*places if places else '')
        for col in ('C','E','G'):
            s.Range(f'{col}{total}').Formula=f'=SUM({col}{first}:{col}{total-1})'
        require(s.Range(f'A{total}').Value2=='Итого','AKZ_TOTAL_ROW_NOT_PRESERVED')
        s.PageSetup.CenterHeader='ТЕСТ / ЧЕРНОВИК'
        s.PageSetup.CenterFooter='ACT-ID ' + p.seal[:16] + ' | ' + data['date']
        s.PageSetup.PrintArea=f'$A$1:$G${total}'
        s.PageSetup.PrintTitleRows='$1:$3'
        s.PageSetup.FitToPagesWide=1; s.PageSetup.FitToPagesTall=False; s.PageSetup.Zoom=False
        s.Calculate()
        for col,field in [('C','quantity'),('E','mass_total'),('G','area_total')]:
            require(abs(Decimal(str(s.Range(f'{col}{total}').Value2))-data['totals'][field])<=Decimal('0.00000001'),
                    'NATIVE_TOTALS_MISMATCH')
        native_save(b,path)
    finally:
        b.Close(SaveChanges=False)

PENDING_FILES = frozenset(('ТЕСТ АКЗ №0.xlsx', 'ТЕСТ Приложение к АКЗ №0.xlsx', 'PAIR.json'))

def pending_identity(path, *, directory=False):
    info = Path(path).lstat()
    require(not (getattr(info, 'st_file_attributes', 0) & 0x400)
            and (stat.S_ISDIR(info.st_mode) if directory else stat.S_ISREG(info.st_mode))
            and info.st_ino != 0, f'PENDING_OWNERSHIP_STOP: принадлежность временного файла не подтверждена: {path}')
    return [info.st_dev, info.st_ino]

def pending_ledger(record):
    return OUTPUT.resolve() / f'.stage6-pending-{record["id"]}.json'

def save_pending_record(record, *, create=False):
    ledger = pending_ledger(record)
    if not create:
        pending_identity(ledger)
    with ledger.open('x' if create else 'w', encoding='utf-8') as stream:
        json.dump(record, stream, ensure_ascii=False, indent=2)

def track_pending_file(record, path):
    require(Path(path).name in PENDING_FILES and Path(path).parent == Path(record['path']),
            'PENDING_OWNERSHIP_STOP: неверный путь временного файла')
    record['files'][Path(path).name] = pending_identity(path)
    save_pending_record(record)

def read_pending_record(ledger):
    """A filename prefix alone never grants permission to delete a directory."""
    try:
        pending_identity(ledger)
        record = json.loads(ledger.read_text(encoding='utf-8'))
        require(isinstance(record, dict) and isinstance(record.get('id'), str),
                'PENDING_OWNERSHIP_STOP: неверная запись принадлежности временного файла')
        token = uuid.UUID(record['id']).hex
        root = OUTPUT.resolve()
        require(record['id'] == token and ledger == root / f'.stage6-pending-{token}.json'
                and record['schema'] == 'STAGE6_OWNED_PENDING_V1' and record['root'] == str(root)
                and record['path'] == str(root / f'.pending-stage6-{token}')
                and isinstance(record['files'], dict) and set(record['files']) <= PENDING_FILES,
                'PENDING_OWNERSHIP_STOP: неверная запись принадлежности временного файла')
        return record
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise Stage6Error(f'PENDING_OWNERSHIP_STOP: не удалось подтвердить принадлежность; '
                          f'файлы не удалены: {ledger}') from exc

def verify_pending_ownership(record):
    temp = Path(record['path'])
    if not temp.exists():
        return []
    require(pending_identity(temp, directory=True) == record['directory_identity'],
            f'PENDING_OWNERSHIP_STOP: временная папка подменена: {temp}')
    files = list(temp.iterdir())
    require(all(f.name in record['files'] for f in files),
            f'PENDING_OWNERSHIP_STOP: в папке есть неподтверждённые файлы; ничего не удалено: {temp}')
    for path in files:
        require(pending_identity(path) == record['files'][path.name],
                f'PENDING_OWNERSHIP_STOP: временный файл подменён: {path}')
    return files

def cleanup_pending(record):
    """Bounded cleanup of registered identities only; no recursive directory delete."""
    files = verify_pending_ownership(record)
    # Check every file before deleting any. An externally held lock leaves the
    # entire private staging directory available for a later controlled retry.
    for path in files:
        probe_file_access(path, read_only=False)
    try:
        for path in files:
            require(pending_identity(path) == record['files'][path.name],
                    f'PENDING_OWNERSHIP_STOP: временный файл подменён: {path}')
            path.unlink()
        temp = Path(record['path'])
        if temp.exists():
            temp.rmdir()  # refuses any unexpected file added since verification
        pending_ledger(record).unlink()
    except OSError as exc:
        code = {32: 'FILE_LOCKED', 33: 'FILE_LOCKED', 5: 'ACCESS_DENIED'}.get(exc.winerror, 'OTHER_ERROR')
        raise FileOperationError(code, exc.filename or record['path'], 'очистка временных файлов') from exc

def recover_pending():
    root = OUTPUT.resolve()
    require(OUTPUT.absolute() == root, 'PENDING_OWNERSHIP_STOP: временный каталог перенаправлен')
    pending_identity(root, directory=True)
    records = [read_pending_record(p) for p in root.glob('.stage6-pending-*.json')]
    owned = {Path(r['path']) for r in records}
    require(all(p in owned for p in root.glob('.pending-*')),
            f'PENDING_OWNERSHIP_STOP: найдена неизвестная временная папка; файлы не удалены: {root}')
    # Reject all unknown contents before cleaning even the first known directory.
    for record in records:
        verify_pending_ownership(record)
    for record in records:
        cleanup_pending(record)

def generate(p, output_name, *, mode='TEST'):
    require(mode == 'TEST', 'OFFICIAL_MODE_STOP')
    verify_preview(p)
    require(bool(output_name) and Path(output_name).name == output_name and output_name not in ('.','..')
            and not any(c in output_name for c in '<>:"/\\|?*'), 'OUTPUT_NAME_INVALID')
    OUTPUT.mkdir(exist_ok=True)
    target=OUTPUT/output_name
    require(not target.exists(), 'OUTPUT_ALREADY_EXISTS_STOP')
    recover_pending()
    token = uuid.uuid4().hex
    temp = OUTPUT.resolve() / f'.pending-stage6-{token}'
    temp.mkdir()
    record = {'schema': 'STAGE6_OWNED_PENDING_V1', 'id': token, 'root': str(OUTPUT.resolve()),
              'path': str(temp), 'directory_identity': pending_identity(temp, directory=True), 'files': {}}
    save_pending_record(record, create=True)
    try:
        act=temp/'ТЕСТ АКЗ №0.xlsx'; application=temp/'ТЕСТ Приложение к АКЗ №0.xlsx'
        shutil.copyfile(ACT,act); track_pending_file(record, act)
        shutil.copyfile(APPLICATION,application); track_pending_file(record, application)
        with native_excel() as app:
            write_act(app,act,p)
            track_pending_file(record, act)  # Excel Save can replace the file identity.
            write_application(app,application,p)
            track_pending_file(record, application)
            validate_saved_pair(app,act,application,p)
        verify_preview(p)
        receipt={'state':'COMPLETE_TEST_PAIR','act_id':p.seal[:16], 'date':p.payload['date'],
                 'number':p.payload['act_number'], 'synthetic':p.payload['synthetic'],
                 'sha256':{x.name:sha256(x) for x in (act,application)}}
        (temp/'PAIR.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
        track_pending_file(record, temp/'PAIR.json')
        # Windows directory rename publishes both files together and refuses existing target.
        os.rename(temp,target)
        pending_ledger(record).unlink()
        return target
    except BaseException as primary:
        try:
            cleanup_pending(record)
        except Exception as cleanup:
            # Keep the original code, path, cause and actionable GUI message.
            # Diagnostics remain inspectable; the durable ownership record allows
            # the next controlled generation to retry cleanup after lock release.
            primary.cleanup_error = cleanup
        raise

def validate_saved_pair(app, act, application, p):
    """Verify saved files before publication, including native read-only reopen."""
    data=p.payload
    a=load_workbook(act,data_only=True)
    try:
        value=a['4']['H4'].value
        require(isinstance(value,datetime) and value.date().isoformat()==data['date'],
                'SAVED_DATE_STYLE_OR_VALUE_STOP')
    finally:a.close()
    a=native_open(app,act,read_only=True)
    b=None
    try:
        b=native_open(app,application,read_only=True)
        require(a.Worksheets.Count==1 and b.Worksheets.Count==1, 'SAVED_EXTRA_SHEETS_STOP')
        s=a.Worksheets('4');t=b.Worksheets('Приложение АКЗ')
        serial=(date.fromisoformat(data['date'])-date(1899,12,30)).days
        require(s.Range('E4').Value2==data['act_number']+' от ' and s.Range('H4').Value2==serial
                and s.Range('G8').Value2==s.Range('E4').Value2 and s.Range('I8').Value2==serial,
                'SAVED_PAIR_IDENTITY_STOP')
        expected_date=date.fromisoformat(data['date']).strftime('%d.%m.%Y')
        require(data['act_number'] in t.Range('A1').Value2 and expected_date in t.Range('A1').Value2
                and s.PageSetup.CenterFooter==t.PageSetup.CenterFooter, 'SAVED_APPLICATION_IDENTITY_STOP')
        require(s.Range('E18').Value2==data['templates']['material'] and s.Range('E21').Value2==data['templates']['ral']
                and s.Range('E25').Value2==data['templates']['required_dry_microns'], 'SAVED_COATING_STOP')
        for cell,key in [('E5','contract'),('E6','customer'),('E7','object'),('B44','cipher')]:
            require(s.Range(cell).Value2==data['passport'][key], 'SAVED_DYNAMIC_FIELDS_STOP')
        require('ТЕСТ / ЧЕРНОВИК' in s.Range('B3').Value2 and 'ТЕСТ / ЧЕРНОВИК' in t.Range('A1').Value2,
                'SAVED_TEST_MARKING_STOP')
        fields=('mark','name','quantity','mass_unit','mass_total','area_unit','area_total')
        for rn,row in enumerate(data['selected'],4):
            actual=t.Range(f'A{rn}:G{rn}').Value2[0]
            for f,x in zip(fields,actual):
                require(abs(Decimal(str(x))-row[f])<=Decimal('0.00000001') if isinstance(row[f],Decimal)
                        else x==row[f], 'SAVED_APPLICATION_ROW_STOP')
        total=4+len(data['selected'])
        require(t.Range(f'A{total}').Value2=='Итого' and
                all(not bool(t.Rows(r).Hidden) for r in range(1,total+1)),
                'SAVED_DYNAMIC_TOTAL_POSITION_STOP')
        # Excel exposes the same area in A1 or R1C1 according to ReferenceStyle.
        require(str(t.PageSetup.PrintArea) in (f'$A$1:$G${total}',f'R1C1:R{total}C7') and
                int(t.PageSetup.Pages.Count)>=1,'SAVED_DYNAMIC_PRINT_AREA_STOP')
        for col,f in [('C','quantity'),('E','mass_total'),('G','area_total')]:
            require(t.Range(f'{col}{total}').Formula==f'=SUM({col}4:{col}{total-1})' and
                    abs(Decimal(str(t.Range(f'{col}{total}').Value2))-data['totals'][f])<=Decimal('0.00000001'),
                    'SAVED_APPLICATION_TOTALS_STOP')
    finally:
        if b is not None:b.Close(SaveChanges=False)
        a.Close(SaveChanges=False)
