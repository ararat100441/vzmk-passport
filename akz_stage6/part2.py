"""Pinned variants 1–3. Part 1 core remains unchanged; all outputs are TEST only."""
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import json
import os
import re
import shutil
import uuid

from openpyxl import load_workbook
from . import core as c

OUTPUT = c.OUTPUT / 'Part2'
MISSING = 'НЕТ ДАННЫХ / ТЕСТ'
FIELDS = ('mark','name','quantity','mass_unit','mass_total','area_unit','area_total')
PINS = {
    1: ('45e14341d244190dee69efe904ad22cdd220dba0e045b55cff3f3b4e57e20741',
        'f0388b627d3a629300d1b6d891f020d4147e0c796469e6f28469100679520706'),
    2: ('cf50e7fc872b82060e44ffabdef998b7dd5bb5647593f9a8654e1b5545425fb2',
        '0297fef23801eff4fad1f952e37f31d997f46d63053161cc849cd6e9b27141bf'),
    3: ('c4f1f3b194ee92bdaaad93ad3c7d7770473e28553c9f1a50a5bd69bb37c8e09e',
        '8dcc7a8f5788132aa7ffc9dd653839911e5ccf69cb03a6b563e393ae6e6f1605'),
}

@dataclass(frozen=True)
class Layout:
    variant: int
    columns: tuple
    ordinals: tuple
    last: int

    @property
    def capacity(self): return self.last - 4
    @property
    def total(self): return self.last + 1
    @property
    def act(self): return c.PROJECT / 'Паспорт' / f'АКЗ-{self.variant}.xlsx'
    @property
    def application(self): return c.PROJECT / 'Паспорт' / f'Приложение к АКЗ-{self.variant}.xlsx'

LAYOUTS = {1:Layout(1,('E',),('E17',),65),
           2:Layout(2,('E','I'),('E17','H17'),51),
           3:Layout(3,('E','H','J'),('E17','H17','J17'),98)}

def layout(variant):
    c.require(type(variant) is int and variant in LAYOUTS, 'VARIANT_INVALID')
    return LAYOUTS[variant]

def template(variant):
    spec=layout(variant)
    for path,pin in zip((spec.act,spec.application),PINS[variant]):
        c.require(path.is_file() and c.sha256(path)==pin, f'TEMPLATE_SHA_CHANGED_STOP: {path}')
    a=load_workbook(spec.act,data_only=False);b=load_workbook(spec.application,data_only=False)
    try:
        c.require(a.sheetnames==(['4','Лист1'] if variant==2 else ['4'])
                  and b.sheetnames==['Приложение №1'], 'TEMPLATE_MAPPING_STOP')
        s=a['4'];t=b.active
        layers=[]
        for n,(col,ordinal) in enumerate(zip(spec.columns,spec.ordinals),1):
            c.require(s[ordinal].value==n, 'TEMPLATE_LAYER_MAPPING_STOP')
            raw=s[col+'25'].value
            layers.append({'ordinal':n,'ordinal_cell':ordinal,'column':col,
                'material':s[col+'18'].value,'color':s[col+'21'].value,
                'required_raw':raw,'required_individual_microns':None if variant==2 and n==2 else raw,
                'requirement_status':'NOT_PROVEN' if variant==2 and n==2 else 'TEMPLATE_EXPLICIT'})
        c.require(t['A'+str(spec.total)].value=='Итого', 'TEMPLATE_TOTAL_MAPPING_STOP')
        # Do not discard a helper with an unknown dependency. Pinned current form
        # has none; changes to formulas/names also fail the SHA gate above.
        if variant==2:
            expressions=[str(cell.value) for sheet in a for row in sheet for cell in row if cell.data_type=='f']
            expressions += [str(v) for v in a.defined_names.values()]
            c.require(not any('Лист1' in value for value in expressions), 'SERVICE_DEPENDENCY_STOP')
        return {'act':str(spec.act),'application':str(spec.application),'layers':layers,
                'layer_count':len(layers),'capacity':spec.capacity,'body_first':5,'body_last':spec.last,
                'total_row':spec.total,'act_print_area':str(s.print_area),
                'application_print_area':str(t.print_area),
                'act_merges':sorted(str(x) for x in s.merged_cells.ranges),
                'application_merges':sorted(str(x) for x in t.merged_cells.ranges)}
    finally:a.close();b.close()

def matching(vom, assignments, variant):
    """Same strict numeric/provenance contract as Part 1; validate only this group."""
    spec=layout(variant);selected=[];excluded=[]
    for row in vom['rows']:
        assignment=assignments.get(row['mark']);n=assignment['assignment'] if assignment else None
        if n!=variant:
            excluded.append({'row':row['row'],'mark':row['mark'],'assignment':n,
                             'reason':'NO_ASSIGNMENT' if n is None else 'OTHER_VARIANT'})
            continue
        # Use accepted row validator, not another source/rounding policy. Each
        # invocation validates one original row, so repeated marks stay repeated.
        one,_,_=c.matching({'rows':[row]}, {row['mark']:{'assignment':0,'row':assignment['row']}})
        selected.extend(one)
    c.require(selected, f'NO_VARIANT_ROWS: №{variant}')
    totals={f:sum((r[f] for r in selected),Decimal(0)) for f in ('quantity','mass_total','area_total')}
    return selected,excluded,totals

@dataclass(frozen=True)
class BatchPreview:
    payload: dict
    seal: str

def preview_groups(project, passport, vom_path, catalog_path, *, synthetic=False):
    c.require(type(synthetic) is bool, 'SYNTHETIC_MODE_INVALID')
    paths=[c.REGISTRY,Path(vom_path).resolve(),Path(catalog_path).resolve()]
    sources={str(p):c.sha256(p) for p in paths}
    current=c.current_passport(project,passport)
    assignments=c.catalog(catalog_path);vom=c.visible_vom(vom_path)
    c.require(not synthetic or Path(vom_path).resolve().is_relative_to(c.FIXTURES.resolve()),
              'SYNTHETIC_VOM_MUST_BE_ISOLATED')
    c.require(synthetic or not vom['synthetic_source'], 'SYNTHETIC_SOURCE_REQUIRES_SYNTHETIC_MODE_STOP')
    days={};climate=c.ClimateDay(current['date']);c.register_climate(days,climate)
    groups={}
    for variant in range(4):
        rows=[r for r in vom['rows'] if assignments.get(r['mark'],{}).get('assignment')==variant]
        if not rows:
            groups[str(variant)]={'status':'EMPTY','reason':'Нет видимых назначенных строк'}
            continue
        try:
            if variant==0:
                p=c.preview(project,passport,vom_path,catalog_path,synthetic=synthetic)
            else:
                inputs=sources | {str(path):c.sha256(path) for path in (layout(variant).act,layout(variant).application)}
                t=template(variant);selected,excluded,totals=matching(vom,assignments,variant)
                c.require(all(c.sha256(path)==pin for path,pin in inputs.items()),'INPUT_CHANGED_DURING_PREVIEW')
                data={'mode':'TEST_ONLY','official':'STOP','synthetic':synthetic,'variant':variant,
                    'passport':current,'act_number':('SYNTHETIC-TEST ' if synthetic else '')+passport+f'-{variant}',
                    'date':current['date'],'sources':inputs,'vom':vom,'templates':t,
                    'selected':selected,'excluded':excluded,'totals':totals,'climate':days[current['date']].__dict__,
                    'warnings':['ТЕСТ / ЧЕРНОВИК. Официальная выдача STOP.',
                                'Измерения и климат отсутствуют; формула Td не утверждена.']}
                p=c.Preview(data,c.digest(data))
            groups[str(variant)]={'status':'READY','payload':p.payload,'seal':p.seal}
        except Exception as exc:
            groups[str(variant)]={'status':'STOP','reason':c.error_message(exc)}
    c.require(all(c.sha256(p)==h for p,h in sources.items()),'INPUT_CHANGED_DURING_PREVIEW')
    data={'project':project,'passport':passport,'date':current['date'],'sources':sources,
          'groups':groups,'climate_days':{k:v.__dict__ for k,v in days.items()},
          'unassigned':[{'row':r['row'],'mark':r['mark']} for r in vom['rows']
                        if assignments.get(r['mark'],{}).get('assignment') is None]}
    return BatchPreview(data,c.digest(data))

def group_preview(batch, variant):
    c.require(isinstance(batch,BatchPreview) and c.digest(batch.payload)==batch.seal,'BATCH_PREVIEW_CHANGED_STOP')
    for path,pin in batch.payload['sources'].items():
        c.require(c.sha256(path)==pin,f'INPUT_CHANGED_AFTER_PREVIEW_STOP: {path}')
    g=batch.payload['groups'].get(str(variant),{})
    c.require(g.get('status')=='READY', f'GROUP_NOT_READY: №{variant}: {g.get("reason")}')
    p=c.Preview(g['payload'],g['seal']);c.verify_preview(p)
    return p

def factual_cells(spec):
    common=['G11','B13','E13','I13','E15','G37','E39','E40','E41','E42','D49','E46','E47']
    return common+[col+str(r) for col in spec.columns for r in (19,20,22,23,24,26,27,28,30,31,32,33,34)]

def header(data):
    return ('ТЕСТ / ЧЕРНОВИК'+(' — SYNTHETIC' if data['synthetic'] else '')+'\n'
        +f'Приложение № {data["act_number"]} к акту АКЗ № {data["act_number"]} от '
        +date.fromisoformat(data['date']).strftime('%d.%m.%Y')+' г.')

def write_act(app, path, p):
    data=p.payload;spec=layout(data['variant']);b=c.native_open(app,path,read_only=False)
    try:
        s=b.Worksheets('4')
        s.Range('B3').Value2='ТЕСТ / ЧЕРНОВИК — АКЗ НЕ ПОДТВЕРЖДЕНО'
        s.Range('E4').Value2=data['act_number']+' от '
        s.Range('H4').Value2=(date.fromisoformat(data['date'])-date(1899,12,30)).days
        for cell,key in [('E5','contract'),('E6','customer'),('E7','object'),('B44','cipher')]:
            s.Range(cell).Value2=data['passport'][key]
        s.Range('E7').WrapText=True;s.Rows(7).RowHeight=max(s.Rows(7).RowHeight,56)
        s.Range('G8').Formula='=E4';s.Range('I8').Formula='=H4'
        for cell in factual_cells(spec):
            s.Range(cell).Value2='?' if cell in ('G11','E15','G37') else MISSING
            s.Range(cell).Font.Size=10
        s.Range('B43').Value2='ТЕСТ: фактические показатели не заполнены. Выдача запрещена.'
        s.Range('B43').WrapText=True
        s.Range('F50').Value2='Начальник ОТК\n'+MISSING
        s.Range('I50').Value2='Начальник малярного участка\n'+MISSING
        s.Range('F50').Font.Size=10;s.Range('I50').Font.Size=10
        # Outside-print helpers contain old solvent fraction and randomized
        # values. No live dependency remains: all factual formulas are replaced.
        s.Range('L1:O240').ClearContents()
        if spec.variant==2:
            helper=b.Worksheets('Лист1')
            helper.UsedRange.ClearContents()
            helper.Range('A1').Value2='СЛУЖЕБНЫЙ ЛИСТ / '+MISSING
            helper.Visible=2  # xlSheetVeryHidden: retained, excluded from ordinary printing
        s.PageSetup.CenterHeader='ТЕСТ / ЧЕРНОВИК'
        s.PageSetup.CenterFooter='ACT-ID '+p.seal[:16]+' | '+data['date']
        s.PageSetup.FitToPagesWide=1;s.PageSetup.FitToPagesTall=1;s.PageSetup.Zoom=False
        s.Activate();app.ActiveWindow.ScrollRow=1;app.ActiveWindow.ScrollColumn=1;s.Range('A1').Select()
        app.CalculateFull()
        c.native_save(b,path)
    finally:b.Close(SaveChanges=False)

def write_application(app,path,p):
    """Grow/shrink only the output appendix; keep the pinned source untouched."""
    data=p.payload;spec=layout(data['variant']);b=c.native_open(app,path,read_only=False)
    try:
        s=b.Worksheets('Приложение №1');count=len(data['selected'])
        c.require(count>=1 and 5+count<=1048576,'AKZ_APPLICATION_EXCEL_ROW_LIMIT')
        first,total=5,5+count
        if total>spec.total:
            s.Rows(f'{spec.total}:{total-1}').Insert()
            s.Range(f'A{spec.last}:G{spec.last}').Copy()
            s.Range(f'A{spec.last+1}:G{total-1}').PasteSpecial(Paste=-4122)
            app.CutCopyMode=False
            s.Rows(f'{spec.last+1}:{total-1}').RowHeight=s.Rows(spec.last).RowHeight
        elif total<spec.total:
            s.Rows(f'{total}:{spec.last}').Delete()
        s.Rows(f'1:{total}').Hidden=False
        s.Range('A1').Value2=header(data);s.Range('A1').WrapText=True;s.Rows(1).RowHeight=max(s.Rows(1).RowHeight,38)
        s.Range('A2').Value2='Объект: '+data['passport']['object']
        s.Range(f'A{first}:G{total-1}').ClearContents()
        for rn,row in enumerate(data['selected'],first):
            s.Range(f'A{rn}:G{rn}').Value2=(tuple(float(row[f]) if isinstance(row[f],Decimal) else row[f] for f in FIELDS),)
        for col,field in zip('CDEFG',FIELDS[2:]):
            values=[r[field] for r in data['selected']]
            if field in data['totals']:values.append(data['totals'][field])
            # Only E/G totals: two display decimals. Retain full numeric precision.
            places=2 if col in ('E','G') else (0 if field=='quantity' else max(max(0,-v.as_tuple().exponent) for v in values))
            s.Range(f'{col}{first}:{col}{total}').NumberFormatLocal='0'+(str(app.DecimalSeparator)+'0'*places if places else '')
        for col in 'CEG':s.Range(f'{col}{total}').Formula=f'=SUM({col}{first}:{col}{total-1})'
        c.require(s.Range(f'A{total}').Value2=='Итого','AKZ_TOTAL_ROW_NOT_PRESERVED')
        s.PageSetup.CenterHeader='ТЕСТ / ЧЕРНОВИК'
        s.PageSetup.CenterFooter='ACT-ID '+p.seal[:16]+' | '+data['date']
        s.PageSetup.PrintArea=f'$A$1:$G${total}'
        s.PageSetup.PrintTitleRows='$1:$4'
        s.PageSetup.FitToPagesWide=1;s.PageSetup.FitToPagesTall=False;s.PageSetup.Zoom=False
        s.Activate();app.ActiveWindow.ScrollRow=1;app.ActiveWindow.ScrollColumn=1;s.Range('A1').Select()
        s.Calculate()
        for col,field in [('C','quantity'),('E','mass_total'),('G','area_total')]:
            c.require(abs(Decimal(str(s.Range(f'{col}{total}').Value2))-data['totals'][field])<=Decimal('0.00000001'),
                      'NATIVE_TOTALS_MISMATCH')
        c.native_save(b,path)
    finally:b.Close(SaveChanges=False)

def validate_saved_pair(app,act,application,p):
    c.verify_preview(p);data=p.payload;spec=layout(data['variant']);tdata=data['templates']
    a=load_workbook(act,data_only=False);b=load_workbook(application,data_only=False)
    try:
        c.require(a.sheetnames==(['4','Лист1'] if spec.variant==2 else ['4']) and b.sheetnames==['Приложение №1'], 'SAVED_SHEETS_STOP')
        s=a['4'];t=b.active
        c.require(all(sheet.sheet_view.selection[0].activeCell=='A1' for sheet in (s,t))
                  and (t.sheet_view.pane is None or t.sheet_view.pane.topLeftCell=='A5'),'SAVED_INITIAL_VIEW_STOP')
        from openpyxl.utils.cell import range_boundaries
        def header_merge(mapping,limit):
            return sorted(x for x in mapping if range_boundaries(x)[3]<limit)
        c.require(sorted(str(x) for x in s.merged_cells.ranges)==tdata['act_merges']
                  and header_merge([str(x) for x in t.merged_cells.ranges],5)==header_merge(tdata['application_merges'],5),
                  'SAVED_MERGES_STOP')
        total=5+len(data['selected'])
        c.require(str(s.print_area)==tdata['act_print_area']
                  and str(t.print_area).replace('$','').upper().endswith(f'G{total}')
                  and t.print_title_rows=='$1:$4','SAVED_PRINT_AREA_STOP')
        for cell in factual_cells(spec):
            c.require(s[cell].value==('?' if cell in ('G11','E15','G37') else MISSING),'SAVED_HISTORICAL_FACT_STOP: '+cell)
        for sheet in a:
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f':c.require(sheet.title=='4' and cell.coordinate in ('G8','I8'),'SAVED_UNAPPROVED_FORMULA_STOP')
        c.require(t[f'A{total}'].value=='Итого','SAVED_DYNAMIC_TOTAL_POSITION_STOP')
        for col in 'CEG':c.require(t[f'{col}{total}'].value==f'=SUM({col}5:{col}{total-1})','SAVED_TOTAL_FORMULA_STOP')
        if spec.variant==2:
            c.require(a['Лист1'].sheet_state=='veryHidden' and a['Лист1']['A1'].value=='СЛУЖЕБНЫЙ ЛИСТ / '+MISSING,'SAVED_SERVICE_SHEET_STOP')
        # Whole-book assertion, including hidden cells, names, headers/footers.
        old_tokens=('1/Б2','1/2025','4/2026','5/2026','6/2026','РУССОКО','Велесстрой',
                    '2-ЭГ-2019','ВС.24.ПО','13-25','314-25','Журавлёв','Шляхтуров')
        for book in (a,b):
            strings=[str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None
                     and not (book is a and sheet.title=='4' and cell.coordinate=='E4')
                     and not (book is b and sheet.title=='Приложение №1' and cell.coordinate=='A1')]
            strings += [str(v) for v in book.defined_names.values()]
            for sheet in book:
                for part in (sheet.oddHeader,sheet.oddFooter,sheet.evenHeader,sheet.evenFooter,sheet.firstHeader,sheet.firstFooter):strings.append(str(part))
            c.require(not any(token in value for value in strings for token in old_tokens),'SAVED_HISTORICAL_REFERENCE_STOP')
    finally:a.close();b.close()
    a=c.native_open(app,act,read_only=True);b=None
    try:
        b=c.native_open(app,application,read_only=True);app.CalculateFull()
        s=a.Worksheets('4');t=b.Worksheets('Приложение №1')
        serial=(date.fromisoformat(data['date'])-date(1899,12,30)).days
        c.require(s.Range('E4').Value2==data['act_number']+' от ' and s.Range('H4').Value2==serial
                  and s.Range('G8').Value2==data['act_number']+' от ' and s.Range('I8').Value2==serial
                  and t.Range('A1').Value2==header(data),'SAVED_PAIR_IDENTITY_STOP')
        for cell,key in [('E5','contract'),('E6','customer'),('E7','object'),('B44','cipher')]:
            c.require(s.Range(cell).Value2==data['passport'][key],'SAVED_REQUISITES_STOP')
        c.require(t.Range('A2').Value2=='Объект: '+data['passport']['object'],'SAVED_OBJECT_STOP')
        for layer in tdata['layers']:
            col=layer['column']
            for cell,value in [(layer['ordinal_cell'],layer['ordinal']),(col+'18',layer['material']),
                               (col+'21',layer['color']),(col+'25',layer['required_raw'])]:
                c.require(s.Range(cell).Value2==value,'SAVED_COATING_STOP: '+cell)
        for cell in factual_cells(spec):
            c.require(s.Range(cell).Value2==('?' if cell in ('G11','E15','G37') else MISSING),'SAVED_RECALC_FACT_STOP')
        for rn,row in enumerate(data['selected'],5):
            for field,value in zip(FIELDS,t.Range(f'A{rn}:G{rn}').Value2[0]):
                c.require(abs(Decimal(str(value))-row[field])<=Decimal('0.00000001') if isinstance(row[field],Decimal)
                          else value==row[field],'SAVED_VOM_ROW_STOP')
        total=5+len(data['selected'])
        c.require(t.Range(f'A{total}').Value2=='Итого' and
                  all(not bool(t.Rows(r).Hidden) for r in range(1,total+1)),
                  'SAVED_DYNAMIC_TOTAL_POSITION_STOP')
        # Excel exposes the same area in A1 or R1C1 according to ReferenceStyle.
        c.require(str(t.PageSetup.PrintArea) in (f'$A$1:$G${total}',f'R1C1:R{total}C7') and
                  int(t.PageSetup.Pages.Count)>=1,'SAVED_DYNAMIC_PRINT_AREA_STOP')
        for col,field in [('C','quantity'),('E','mass_total'),('G','area_total')]:
            c.require(t.Range(f'{col}{total}').Formula==f'=SUM({col}5:{col}{total-1})' and
                      abs(Decimal(str(t.Range(f'{col}{total}').Value2))-data['totals'][field])<=Decimal('0.00000001'),
                      'SAVED_TOTAL_STOP')
        c.require(int(s.PageSetup.Pages.Count)==1,'SAVED_ACT_PAGES_STOP')
        if spec.variant==2:c.require(b is not None and a.Worksheets('Лист1').Visible==2,'SAVED_SERVICE_PRINT_STOP')
        return {'act_pages':int(s.PageSetup.Pages.Count),'application_pages':int(t.PageSetup.Pages.Count),
                'service_printed':False,'full_calculation':'PASS','read_only':True,'normal_load':True}
    finally:
        if b is not None:b.Close(SaveChanges=False)
        a.Close(SaveChanges=False)

def filenames(passport,variant):
    layout(variant)
    c.require(isinstance(passport,str) and re.fullmatch(r'[1-9]\d*/\d{4}',passport),'PASSPORT_NAME_STOP')
    base=passport.replace('/','_')+f'-{variant}'
    return (base+' АКЗ TEST.xlsx',base+' Приложение АКЗ TEST.xlsx','PAIR.json')

def ledger(record):return OUTPUT.resolve()/f'.part2-pending-{record["id"]}.json'

def save_record(record,*,create=False):
    path=ledger(record)
    if not create:c.pending_identity(path)
    with path.open('x' if create else 'w',encoding='utf-8') as f:json.dump(record,f,ensure_ascii=False,indent=2)

def track(record,path):
    c.require(path.name in filenames(record['passport'],record['variant']) and path.parent==Path(record['path']),'PENDING_OWNERSHIP_STOP')
    record['files'][path.name]=c.pending_identity(path);save_record(record)

def read_record(path):
    try:
        c.pending_identity(path);r=json.loads(path.read_text(encoding='utf-8'));token=uuid.UUID(r['id']).hex;root=OUTPUT.resolve()
        c.require(r['id']==token and path==root/f'.part2-pending-{token}.json'
                  and r['schema']=='STAGE6_PART2_OWNED_PENDING_V1' and r['root']==str(root)
                  and r['path']==str(root/f'.pending-part2-{token}') and isinstance(r['files'],dict)
                  and set(r['files'])<=set(filenames(r['passport'],r['variant'])),'PENDING_OWNERSHIP_STOP')
        return r
    except (OSError,ValueError,KeyError,TypeError) as exc:raise c.Stage6Error(f'PENDING_OWNERSHIP_STOP: {path}') from exc

def cleanup(record):
    files=c.verify_pending_ownership(record)
    for path in files:c.probe_file_access(path,read_only=False)
    try:
        for path in files:
            c.require(c.pending_identity(path)==record['files'][path.name],'PENDING_OWNERSHIP_STOP')
            path.unlink()
        path=Path(record['path'])
        if path.exists():path.rmdir()
        ledger(record).unlink()
    except OSError as exc:
        raise c.FileOperationError({32:'FILE_LOCKED',33:'FILE_LOCKED',5:'ACCESS_DENIED'}.get(exc.winerror,'OTHER_ERROR'),exc.filename or record['path'],'очистка временных файлов') from exc

def recover_pending():
    root=OUTPUT.resolve();c.require(OUTPUT.absolute()==root,'PENDING_OWNERSHIP_STOP')
    c.pending_identity(root,directory=True)
    records=[read_record(p) for p in root.glob('.part2-pending-*.json')]
    owned={Path(r['path']) for r in records}
    c.require(all(p in owned for p in root.glob('.pending-*')),'PENDING_OWNERSHIP_STOP: неизвестная папка')
    for r in records:c.verify_pending_ownership(r)
    for r in records:cleanup(r)

def generate(p,output_name,*,mode='TEST'):
    c.require(mode=='TEST','OFFICIAL_MODE_STOP');c.verify_preview(p)
    data=p.payload;spec=layout(data['variant']);names=filenames(data['passport']['number'],spec.variant)
    c.require(1<=len(data['selected']) and 5+len(data['selected'])<=1048576,'AKZ_APPLICATION_EXCEL_ROW_LIMIT')
    c.require(bool(output_name) and Path(output_name).name==output_name and output_name not in ('.','..')
              and not any(x in output_name for x in '<>:"/\\|?*'),'OUTPUT_NAME_INVALID')
    OUTPUT.mkdir(parents=True,exist_ok=True);target=OUTPUT/output_name
    c.require(not target.exists(),'OUTPUT_ALREADY_EXISTS_STOP');recover_pending()
    token=uuid.uuid4().hex;temp=OUTPUT.resolve()/f'.pending-part2-{token}';temp.mkdir()
    record={'schema':'STAGE6_PART2_OWNED_PENDING_V1','id':token,'root':str(OUTPUT.resolve()),'path':str(temp),
            'variant':spec.variant,'passport':data['passport']['number'],
            'directory_identity':c.pending_identity(temp,directory=True),'files':{}}
    save_record(record,create=True)
    try:
        act=temp/names[0];application=temp/names[1]
        shutil.copyfile(spec.act,act);track(record,act)
        shutil.copyfile(spec.application,application);track(record,application)
        with c.native_excel() as app:
            write_act(app,act,p);track(record,act)
            write_application(app,application,p);track(record,application)
            native=validate_saved_pair(app,act,application,p)
        c.verify_preview(p)
        receipt={'state':'COMPLETE_TEST_PAIR','official':'STOP','variant':spec.variant,'act_id':p.seal[:16],
            'number':data['act_number'],'date':data['date'],'templates':data['templates'],
            'climate':data['climate'],'totals':data['totals'],'native':native,
            'sha256':{x.name:c.sha256(x) for x in (act,application)}}
        (temp/names[2]).write_text(json.dumps(receipt,ensure_ascii=False,indent=2,default=str),encoding='utf-8');track(record,temp/names[2])
        os.rename(temp,target);ledger(record).unlink();return target
    except BaseException as primary:
        try:cleanup(record)
        except Exception as secondary:primary.cleanup_error=secondary
        raise

def generate_groups(batch,variants,output_name):
    c.require(variants and len(set(variants))==len(variants) and all(type(n) is int and n in range(4) for n in variants),'GROUP_SELECTION_INVALID')
    # Preflight all selected previews before publishing even the first pair.
    checked={n:group_preview(batch,n) for n in variants}
    results={}
    for n in variants:
        try:
            group_preview(batch,n)  # shared inputs may change between groups
            p=checked[n]
            target=(c.generate if n==0 else generate)(p,output_name+f'_AKZ{n}')
            results[str(n)]={'status':'COMPLETE_TEST_PAIR','path':str(target)}
        except Exception as exc:
            results[str(n)]={'status':'STOP','reason':c.error_message(exc)}
    return {'status':'COMPLETE_ALL_SELECTED' if all(r['status']=='COMPLETE_TEST_PAIR' for r in results.values()) else 'PARTIAL_STOP',
            'groups':results}
