"""Focused Stage 5 validation; all mutations remain inside the named run."""
from copy import deepcopy
from decimal import Decimal
import json
import os
from pathlib import Path
import socket
import sys
import unittest
from zipfile import ZipFile

from lxml import etree as E
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from uzk_stage5.core import Preview, calculate_marks, digest, generate, preview, product_paths
from uzk_stage5.inputs import (PINS, Stage5Error, key, mark_profiles, positive_integer,
                              registry_rows, sha, verify_pins, visible_vom)
from uzk_stage5.word import NS, contents, tag, write_test_document
from profiles_stage4.model import Catalog, ProfileError
from profiles_stage4.offline import no_network
from profiles_stage4.storage import load_catalog


class Stage5Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.run_dir = Path(os.environ['STAGE5_VALIDATION_DIR']).resolve()
        if not cls.run_dir.is_relative_to(ROOT/'validation'):
            raise RuntimeError('Fixtures must remain in validation')
        cls.fixtures = cls.run_dir/'fixtures'
        cls.fixtures.mkdir(exist_ok=True)
        cls.actual = json.loads((cls.run_dir/'final_run.json').read_text(encoding='utf-8-sig'))
        cls.data = cls.actual['preview']
        cls.plan = Preview(cls.data, cls.actual['preview_seal'])
        cls.output = Path(cls.actual['output'])
        cls.catalog = load_catalog()
        cls.stock, cls.evidence = registry_rows()
        cls.mapping = mark_profiles(cls.data['mark_catalog'])
        cls.native = json.loads((cls.run_dir/'native_final_v3.json').read_text(encoding='utf-8-sig'))
        with ZipFile(cls.output) as z:
            cls.doc = E.fromstring(z.read('word/document.xml'))
            cls.rels = z.read('word/_rels/document.xml.rels')
        with ZipFile(cls.data['template']) as z:
            cls.template = E.fromstring(z.read('word/document.xml'))
        cls.rows = cls.doc.find('w:body', NS).findall('w:tbl', NS)[0].findall('w:tr', NS)[5:]
        cls.marks = {m['mark']:m for m in cls.data['marks']}
        cls.binding = {str(p):sha(p) for p in [*product_paths(), Path(__file__)]}
        (cls.run_dir/'tested_product_before.json').write_text(json.dumps(cls.binding,ensure_ascii=False,indent=2),encoding='utf8')

    def fixture_book(self, name, rows):
        path = self.fixtures/name
        book = Workbook()
        sheet = book.active
        sheet.title = 'УЗК'
        for row in rows:
            sheet.append(row)
        book.save(path)
        return path

    def calculate(self, rows=None, stock=None, mappings=None, catalog=None):
        return calculate_marks('Гараж Казахстан', rows if rows is not None else self.data['vom']['rows'],
                               stock if stock is not None else self.stock,
                               mappings if mappings is not None else self.mapping,
                               catalog if catalog is not None else self.catalog)

    def test_01_checkpoint_and_drift(self):
        verify_pins()
        path = self.fixtures/'pin.txt'
        path.write_text('accepted',encoding='utf8')
        pins = {path:sha(path)}
        verify_pins(pins)
        path.write_text('changed',encoding='utf8')
        with self.assertRaisesRegex(Stage5Error,'CHECKPOINT_CHANGED'):
            verify_pins(pins)

    def test_02_actual_readonly_and_visible_controls(self):
        r = visible_vom(self.data['vom']['path'])
        self.assertTrue(r['read_only'])
        self.assertTrue(r['macros_disabled'])
        self.assertFalse(r['update_links'])
        self.assertEqual(r['rows'],self.data['vom']['rows'])
        self.assertEqual(r['visible_count'],104)
        self.assertEqual(Decimal(r['mass_kg']),Decimal('15822.14'))
        self.assertEqual(r['external_link_count'],0)
        self.assertEqual(r['connection_count'],0)
        (self.run_dir/'actual_vom_reread.json').write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf8')

    def test_03_xls_hidden_and_filtered_exclusion(self):
        r=self.data['vom']
        self.assertEqual(Path(r['path']).suffix,'.xls')
        self.assertEqual(len(r['hidden_data_rows']),397)
        self.assertFalse({v['row'] for v in r['rows']} & set(r['hidden_data_rows']))
        self.assertIsNotNone(r['autofilter_range'])

    def test_04_xls_and_xlsx_adapter_fixtures(self):
        import pythoncom
        import win32com.client
        path=self.fixtures/'visible.xlsx'
        xls=self.fixtures/'visible.xls'
        b=Workbook();s=b.active
        s.append(['Ведомость отправочных марок'])
        s.append(['Лист №','Отправочная марка','Наименование марки','Количество',None,None,None,'Масса (изделие), кг'])
        s.append([None,None,None,None,None,None,None,'1 дет.','Всех'])
        s.append(['1','A','Балка',2,None,None,None,10,20])
        s.append(['2','A','Балка',3,None,None,None,10,30])
        s.append(['3','B','Балка',7,None,None,None,10,70])
        b.save(path)
        pythoncom.CoInitialize();app=book=None
        try:
            app=win32com.client.DispatchEx('Excel.Application');app.Visible=False
            app.DisplayAlerts=False;app.AutomationSecurity=3;app.EnableEvents=False
            book=app.Workbooks.Open(str(path),UpdateLinks=0,ReadOnly=False,AddToMru=False)
            sheet=book.Worksheets(1)
            sheet.Range('A2:I6').AutoFilter(Field=2,Criteria1='A')
            sheet.Rows(5).Hidden=True
            book.Save()  # isolated synthetic fixture only
            book.SaveAs(str(xls),FileFormat=56)
        finally:
            if book is not None:book.Close(False)
            if app is not None:app.Quit()
            pythoncom.CoUninitialize()
        results=[]
        for p in (path,xls):
            before=sha(p);r=visible_vom(p)
            self.assertEqual(r['visible_count'],1)
            self.assertEqual(r['rows'][0]['quantity'],2)
            self.assertEqual(r['hidden_data_rows'],[5,6])
            self.assertEqual(Decimal(r['mass_kg']),Decimal(20))
            self.assertEqual(sha(p),before)
            results.append(r)
        (self.run_dir/'adapter_fixture_evidence.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf8')

    def test_05_repeated_vom_marks_aggregate_safely(self):
        rows=[{'row':4,'mark':'Пс1-9','quantity':5},{'row':8,'mark':' пс1-9 ','quantity':20}]
        result,_=self.calculate(rows)
        self.assertEqual(len(result),1)
        self.assertEqual(result[0]['mark'],'Пс1-9')
        self.assertEqual(result[0]['vom_quantity'],25)
        self.assertEqual(result[0]['Q'],9)
        self.assertEqual(result[0]['segments'][0]['L'],'3600')

    def test_06_registry_full_range(self):
        self.assertEqual(len(self.evidence),25)
        self.assertEqual(self.evidence[-1]['row'],26)
        self.assertEqual(sum(r['row']>21 for r in self.evidence),5)

    def test_07_blank_remaining_one_and_numeric(self):
        path=self.fixture_book('blank.xlsx',[['Заказ','Марка элемента','Кол-во (шт)'],['X','A',None],['X','B',8]])
        r,_=registry_rows(path)
        self.assertEqual(r[('x','a')]['remaining'],1)
        self.assertEqual(r[('x','b')]['remaining'],8)
        self.assertTrue(any(r['raw_quantity'] is None and r['remaining']==1 for r in self.evidence))

    def test_08_invalid_quantities_stop(self):
        for bad in (0,-1,1.5,True,'=2+2','bad','NaN','Infinity'):
            with self.subTest(bad=bad),self.assertRaises(Stage5Error):
                positive_integer(bad,'fixture')

    def test_09_minimum(self):
        for quantity,q in ((1,1),(8,8),(50,8)):
            m,_=self.calculate([{'row':1,'mark':'Бп2-1','quantity':quantity}])
            self.assertEqual(m[0]['Q'],q)
            self.assertEqual(Decimal(m[0]['segments'][0]['L']),Decimal(160*q))

    def test_10_project_isolation(self):
        stock=deepcopy(self.stock)
        stock[('other','бп2-1')]={'remaining':900,'rows':[999]}
        marks,_=self.calculate(stock=stock)
        self.assertEqual(next(m for m in marks if m['mark']=='Бп2-1')['Q'],8)
        del stock[(key('Гараж Казахстан'),key('Бп2-1'))]
        marks,excluded=self.calculate(stock=stock)
        self.assertNotIn('Бп2-1',[m['mark'] for m in marks])
        self.assertIn('Бп2-1',excluded)

    def test_11_missing_registry_excluded(self):
        marks,excluded=self.calculate([{'row':1,'mark':'NOT_IN_REGISTRY','quantity':100}])
        self.assertEqual(marks,[])
        self.assertEqual(excluded,['NOT_IN_REGISTRY'])

    def test_12_registry_conflict_and_identical_duplicates(self):
        header=['Заказ','Марка элемента','Кол-во (шт)']
        path=self.fixture_book('conflict.xlsx',[header,['X','A',8],[' x ',' a ',9]])
        with self.assertRaisesRegex(Stage5Error,'REGISTRY_CONFLICT'):registry_rows(path)
        path=self.fixture_book('identical.xlsx',[header,['X','A',8],['x','a',8]])
        r,_=registry_rows(path)
        self.assertEqual(r[('x','a')]['remaining'],8)
        self.assertEqual(r[('x','a')]['rows'],[2,3])

    def test_13_four_mappings_and_unrelated_unresolved(self):
        expected={'Бп2-1':'[16П','Пс1-9':'Гн.[140X60X4','Пс1-10':'Гн.[140X60X4','Пс2-16':'Гн.[]120х3'}
        self.assertEqual({m['mark']:m['raw_profile'] for m in self.data['marks']},expected)
        mapping=deepcopy(self.mapping)
        mapping['unrelated']={'mark':'UNRELATED','profile':'ПРОВЕРИТЬ','row':9999}
        self.assertEqual(len(self.calculate(mappings=mapping)[0]),4)

    def test_14_catalog_conflict_and_missing_profile(self):
        path=self.fixtures/'marks_conflict.xlsx';b=Workbook();s=b.active
        s.append(['Марка','Профиль','Краска','№ АКЗ']);s.append(['A','[16П',None,None]);s.append([' a ','other',None,None]);b.save(path)
        with self.assertRaisesRegex(Stage5Error,'MARK_PROFILE_CONFLICT'):mark_profiles(path)
        for raw in (None,'ПРОВЕРИТЬ','=A1'):
            mappings=deepcopy(self.mapping);mappings[key('Бп2-1')]['profile']=raw
            with self.subTest(raw=raw),self.assertRaises(Stage5Error):self.calculate(mappings=mappings)

    def test_15_stage4_found(self):
        self.assertTrue(all(m['lookup']['status'] in ('FOUND','FOUND_PREFERRED') for m in self.data['marks']))
        self.assertEqual(self.marks['Бп2-1']['lookup']['key'],'16П')

    def test_16_actual_stage4_ambiguous_and_conflict(self):
        source=deepcopy(self.marks['Бп2-1']['lookup']['profile']['provenance'][0])
        other=deepcopy(source);other['record_id']='FIXTURE';other['geometry']['b']=65
        other['standard']='ГОСТ 99999-2000'
        ambiguous=Catalog([source,other])
        self.assertEqual(ambiguous.lookup('[16П')['status'],'AMBIGUOUS')
        with self.assertRaisesRegex(Stage5Error,'PROFILE_AMBIGUOUS'):self.calculate(catalog=ambiguous)
        other['standard']=source['standard']
        conflict=Catalog([source,other])
        self.assertEqual(conflict.lookup('[16П')['status'],'CONFLICT')
        with self.assertRaisesRegex(Stage5Error,'PROFILE_CONFLICT'):self.calculate(catalog=conflict)

    def test_17_unsupported_rule_and_not_found_stop(self):
        source=deepcopy(self.marks['Бп2-1']['lookup']['profile']['provenance'][0])
        unsupported=Catalog([source])  # no extension approval
        with self.assertRaisesRegex(Stage5Error,'UZK_RULE_NOT_CONFIGURED'):self.calculate(catalog=unsupported)
        with self.assertRaisesRegex(Stage5Error,'PROFILE_NOT_FOUND'):self.calculate(catalog=Catalog([]))

    def test_18_channel_segments(self):
        m=self.marks['Бп2-1']
        self.assertEqual([(r['L'],r['S']) for r in m['single']['rows']],[(160,5),(128,8.4)])
        self.assertEqual([(r['L'],r['S']) for r in m['segments']],[('1280',5),('1024',8.4)])

    def test_19_bent_segments(self):
        for mark,length,q in [('Пс1-9','3600',9),('Пс1-10','6000',15)]:
            m=self.marks[mark]
            self.assertEqual(m['Q'],q)
            self.assertEqual([(r['L'],r['S']) for r in m['segments']],[(length,4)])
            self.assertEqual(m['single']['rows'][0]['L'],400)

    def test_20_closed_square(self):
        m=self.marks['Пс2-16']
        self.assertEqual(m['Q'],1)
        self.assertEqual([(r['L'],r['S']) for r in m['segments']],[('480',3)])

    def test_21_frequency_and_angle(self):
        self.assertTrue(all(r['frequency_mhz']==5 and r['angle_degrees']==70 for m in self.data['marks'] for r in m['segments']))

    def test_22_no_cross_mark_merge(self):
        self.assertEqual(len(self.marks),4)
        self.assertEqual(self.data['controlled_units'],33)
        self.assertNotEqual(self.marks['Пс1-9']['segments'][0]['L'],self.marks['Пс1-10']['segments'][0]['L'])

    def test_23_exact_four_physical_rows(self):
        self.assertEqual(len(self.rows),4)
        self.assertEqual([contents(r.findall('w:tc',NS)[0]) for r in self.rows],list(self.marks))
        self.assertTrue(all(len(r.findall('w:tc',NS))==8 for r in self.rows))

    def test_24_multiline_channel(self):
        cells=self.rows[0].findall('w:tc',NS)
        self.assertEqual(cells[1].xpath('.//w:t/text()',namespaces=NS),['L=1280х5 мм','L=1024х8.4 мм'])
        for i in (1,4,5):self.assertEqual(len(cells[i].findall('.//w:br',NS)),1)
        self.assertEqual(cells[4].xpath('.//w:t/text()',namespaces=NS),['5','5'])
        self.assertEqual(cells[5].xpath('.//w:t/text()',namespaces=NS),['70','70'])

    def test_25_requisites_and_draft_results(self):
        body=self.doc.find('w:body',NS)
        self.assertIn('Акт № 9/2026 от 01.10.2026 (ТЕСТ / ЧЕРНОВИК)',contents(body))
        self.assertIn(self.data['passport']['object']+' ('+self.data['passport']['cipher']+')',contents(body))
        for row in self.rows:
            cells=row.findall('w:tc',NS)
            self.assertEqual(contents(cells[3]),'01.10.2026')
            self.assertEqual(contents(cells[6]),'')
            self.assertEqual(contents(cells[7]),'')

    def test_26_no_historical_marks_or_pdf_links(self):
        old=self.template.find('w:body',NS).findall('w:tbl',NS)[0].findall('w:tr',NS)[5:]
        new=contents(self.doc)
        for row in old:self.assertNotIn(contents(row.findall('w:tc',NS)[0]),new)
        self.assertFalse(self.doc.findall('.//w:hyperlink',NS))
        rels=E.fromstring(self.rels)
        self.assertFalse(any(r.get('Type','').endswith('/hyperlink') for r in rels))

    def test_27_package_and_constants_preserved(self):
        with ZipFile(self.data['template']) as a,ZipFile(self.output) as b:
            self.assertEqual(a.namelist(),b.namelist())
            changed=[name for name in a.namelist() if a.read(name)!=b.read(name)]
            self.assertEqual(set(changed),{'word/document.xml','word/_rels/document.xml.rels'})
        a=self.template.find('w:body',NS);b=self.doc.find('w:body',NS)
        ta=a.findall('w:tbl',NS)[0].findall('w:tr',NS);tb=b.findall('w:tbl',NS)[0].findall('w:tr',NS)
        for i in (0,2,3,4):self.assertEqual(E.tostring(ta[i]),E.tostring(tb[i]))
        self.assertEqual(E.tostring(ta[1].findall('w:tc',NS)[0]),E.tostring(tb[1].findall('w:tc',NS)[0]))
        for row in self.rows:self.assertEqual(E.tostring(ta[5].findall('w:tc',NS)[2]),E.tostring(row.findall('w:tc',NS)[2]))
        for i in (0,1,2,3,5,6,7):
            self.assertEqual(E.tostring(a.findall('w:p',NS)[i]),E.tostring(b.findall('w:p',NS)[i]))

    def test_28_signatures_and_page_geometry(self):
        a=self.template.find('w:body',NS);b=self.doc.find('w:body',NS)
        self.assertEqual(E.tostring(a.findall('w:tbl',NS)[1]),E.tostring(b.findall('w:tbl',NS)[1]))
        self.assertEqual(E.tostring(a.find('w:sectPr',NS)),E.tostring(b.find('w:sectPr',NS)))
        self.assertEqual(self.native['sections'],1)
        self.assertEqual(self.native['tables'],2)
        self.assertEqual(self.native['orientation'],1)

    def test_29_no_uzk_no_file(self):
        data=deepcopy(self.data);data['marks']=[];data['controlled_units']=0
        folder=self.run_dir/'no_uzk_output'
        self.assertIsNone(generate(Preview(data,digest(data)),folder))
        self.assertFalse(folder.exists())

    def test_30_31_sources_unchanged(self):
        before=json.loads((self.run_dir/'protected_before.json').read_text(encoding='utf8'))
        self.assertTrue(all(sha(path)==expected for path,expected in before.items()))

    def test_32_33_access_boundary_and_network(self):
        events=[]; active=[False]
        def hook(event,args):
            if event!='open' or not active[0] or not isinstance(args[0],(str,bytes)):return
            path=Path(os.fsdecode(args[0]))
            if path.name.casefold() in ('взмк.xlsx','узк.xlsm'):
                raise AssertionError('Production access forbidden: '+str(path))
            events.append({'path':str(path),'mode':args[1]})
        sys.addaudithook(hook)
        active[0]=True
        try:
            with no_network() as attempts:
                plan=preview('Гараж Казахстан','9/2026',self.data['vom']['path'])
                from uuid import uuid4
                output=generate(plan,self.run_dir/('guard_output_'+uuid4().hex[:8]))
            self.assertEqual(attempts,[])
            self.assertTrue(output.is_file())
        finally:active[0]=False
        self.assertFalse(any(Path(e['path']).name.casefold() in ('взмк.xlsx','узк.xlsm') for e in events))
        writes=[e for e in events if isinstance(e['mode'],str) and any(x in e['mode'] for x in 'wax+')]
        self.assertTrue(writes)
        self.assertTrue(all(Path(e['path']).resolve().is_relative_to(self.run_dir) for e in writes))
        (self.run_dir/'access_network_evidence.json').write_text(json.dumps({'network_attempts':attempts,'python_open_events':events,
            'write_events':writes,'production_denial_guard_active':True,'native_COM_scope':'selected VOM ReadOnly; links=0; connections=0'},ensure_ascii=False,indent=2),encoding='utf8')
        with no_network() as denied:
            with self.assertRaises(ProfileError):socket.getaddrinfo('fixture.invalid',80)
        self.assertEqual(len(denied),1)

    def test_34_native_word_normal_open(self):
        self.assertEqual(self.native['status'],'PASS_NORMAL_READONLY_OPEN')
        self.assertEqual(self.native['before_sha'],sha(self.output))
        self.assertTrue(self.native['unchanged'])
        self.assertTrue(self.native['readonly'])
        self.assertEqual(self.native['first_table_rows'],9)
        self.assertEqual(self.native['hyperlinks'],0)
        self.assertEqual(self.native['pages'],len(self.native['images']))
        self.assertEqual(self.native['pages'],2)

    def test_35_stale_preview_and_output_policies(self):
        with self.assertRaisesRegex(Stage5Error,'OUTPUT_EXISTS'):generate(self.plan,self.output.parent)
        with self.assertRaisesRegex(Stage5Error,'OUTPUT_PATH_FORBIDDEN'):generate(self.plan,ROOT)
        path=self.fixtures/'drift.txt';path.write_text('before',encoding='utf8')
        data=deepcopy(self.data);data['fingerprints'][str(path)]=sha(path)
        plan=Preview(data,digest(data));path.write_text('after',encoding='utf8')
        with self.assertRaisesRegex(Stage5Error,'INPUT_CHANGED'):generate(plan,self.run_dir/'drift_output')
        self.assertFalse((self.run_dir/'drift_output').exists())
        data=deepcopy(self.data);data['marks'][0]['Q']=999
        with self.assertRaisesRegex(Stage5Error,'PREVIEW_CHANGED'):generate(Preview(data,self.plan.seal),self.run_dir/'tampered_output')

    def test_36_word_mapping_failure_isolated(self):
        path=self.fixtures/'ambiguous_template.docx'
        with ZipFile(self.data['template']) as source,ZipFile(path,'w') as dest:
            for p in source.infolist():
                raw=source.read(p.filename)
                if p.filename=='word/document.xml':
                    root=E.fromstring(raw);body=root.find('w:body',NS)
                    title=next(p for p in body.findall('w:p',NS) if contents(p).startswith('Акт №'))
                    body.insert(4,deepcopy(title));raw=E.tostring(root)
                dest.writestr(p,raw)
        data=deepcopy(self.data);data['template']=str(path)
        output=self.fixtures/'must_not_exist.docx'
        with self.assertRaisesRegex(Stage5Error,'WORD_ACT_MAPPING_AMBIGUOUS'):write_test_document(data,output)
        self.assertFalse(output.exists())

    def test_37_gui_generation_gate(self):
        import tkinter as tk
        from uzk_stage5.gui import App
        root=tk.Tk();root.withdraw()
        try:
            app=App(root);root.update_idletasks()
            self.assertTrue(app.generate_button.instate(['disabled']))
            app.plan=self.plan;app.update_buttons()
            self.assertTrue(app.generate_button.instate(['disabled']))
            app.checked.set(True);app.update_buttons()
            self.assertFalse(app.generate_button.instate(['disabled']))
            app.vom.set('changed.xls');app.update_buttons()
            self.assertIsNone(app.plan)
            self.assertTrue(app.generate_button.instate(['disabled']))
            app.plan=self.plan;app.checked.set(True);app.busy=True;app.update_buttons()
            self.assertTrue(app.generate_button.instate(['disabled']))
        finally:root.destroy()

    def test_38_existing_file_race_is_preserved(self):
        from unittest.mock import patch
        from uuid import uuid4
        folder=self.run_dir/('race_'+uuid4().hex[:8])
        def concurrent_create(data,path):
            path.write_bytes(b'created by another process')
            raise FileExistsError(str(path))
        with patch('uzk_stage5.word.write_test_document',side_effect=concurrent_create):
            with self.assertRaises(FileExistsError):generate(self.plan,folder)
        path=folder/self.output.name
        self.assertEqual(path.read_bytes(),b'created by another process')

    @classmethod
    def tearDownClass(cls):
        current={str(p):sha(p) for p in [*product_paths(),Path(__file__)]}
        (cls.run_dir/'tested_product_after.json').write_text(json.dumps(current,ensure_ascii=False,indent=2),encoding='utf8')
        if current!=cls.binding:raise AssertionError('Product changed during tests')


if __name__=='__main__':
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(Stage5Tests)
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    run=Path(os.environ['STAGE5_VALIDATION_DIR'])
    summary={'tests_run':result.testsRun,'successful':result.wasSuccessful(),
             'failures':[(str(t),text) for t,text in result.failures],
             'errors':[(str(t),text) for t,text in result.errors],
             'skipped':[(str(t),why) for t,why in result.skipped]}
    (run/'targeted_tests.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8')
    raise SystemExit(0 if result.wasSuccessful() else 1)
