"""Editable project totals and fail-closed historical accounting reconciliation."""
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from openpyxl import Workbook,load_workbook
from production_app import catalog_ledger as ledger
from production_app.storage import Stop,sha


def appendix(path):
    b=Workbook();s=b.active;s.title='Приложение №1'
    s['A1']='Приложение № 1 к паспорту № 1/2026 от 09.09.2026 г.'
    s['A2']='Объект: SYNTHETIC OBJECT'
    for cell,value in [('A3','Марка'),('B3','Наименование'),('C3','Количество'),('D3','Масса'),('F3','Площадь'),('D4','1 шт'),('E4','Всех'),('F4','1 шт'),('G4','Всех')]:s[cell]=value
    for area in ('A3:A4','B3:B4','C3:C4','D3:E3','F3:G3'):s.merge_cells(area)
    s.append(['M','SYNTHETIC',3,2,6,1,3]);s.append(['N','SYNTHETIC',1,4,4,1,1])
    s.append(['Итого',None,'=SUM(C5:C6)',None,'=SUM(E5:E6)'])
    b.save(path);b.close()


class EditableLedger(unittest.TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name)
        self.folder=self.root/'P'/'Паспорт'/'1 от 09.09.2026';self.folder.mkdir(parents=True)
        self.app=self.folder/'1_2026 - Приложение № 1.xlsx';appendix(self.app)
        self.old=self.root/'old.xlsx';self.catalog=self.root/'catalog.xlsx'
        b=Workbook();s=b.active;s.append(['Марка','Профиль','Краска','№ АКЗ','Количество']);s.append(['M','16П','SYNTHETIC',0,999]);s.append(['N','16П','SYNTHETIC',0,999]);b.save(self.old);b.close()
        self.history=[dict(number='1/2026',date='2026-09-09',complete=True,quantities={'M':3,'N':1},source=str(self.app),sha256=sha(self.app))]
        self.entries=[dict(number='1/2026',date=datetime(2026,9,9),object='SYNTHETIC OBJECT',mass=10)]
        ledger.create_migrated(self.old,self.catalog,'P',deepcopy(self.history))
        self.totals(8,4)
    def tearDown(self):self.tmp.cleanup()
    def edit(self,path,changes):
        b=load_workbook(path)
        for cell,value in changes.items():b.active[cell]=value
        b.save(path);b.close()
    def totals(self,m,n=4):self.edit(self.catalog,{'B2':m,'B3':n})
    def check(self,quantities=None,entries=True):return ledger.validate_stock(self.catalog,'P','2/2026',quantities or {'M':1},self.entries if entries else None)
    def candidate(self,name='candidate.xlsx'):
        p=self.root/name;ledger.shipment_candidate(self.catalog,p,'P','2/2026','2026-10-31',{'M':1},self.entries);return p
    def test_increase_reads_current_B_ignoring_hidden_original(self):
        self.totals(20);self.check({'M':17})
    def test_decrease_reads_current_B_and_rejects_excess(self):
        self.totals(4);self.check()
        with self.assertRaisesRegex(Stop,'остаток 1, запрошено 2'):self.check({'M':2})
    def test_equal_to_shipped_is_valid_zero_stock(self):
        self.totals(3);self.check({'M':0})
        with self.assertRaisesRegex(Stop,'остаток 0'):self.check()
    def test_negative_remaining_reports_total_and_shipped(self):
        self.totals(2)
        with self.assertRaisesRegex(Stop,'проектное количество 2, уже отгружено 3'):self.check()
    def test_negative_unrequested_mark_also_stops(self):
        self.totals(8,0)
        with self.assertRaisesRegex(Stop,'Марка N: отрицательный'):self.check()
    def test_stale_E_does_not_authorize_shipment(self):
        self.totals(3);self.edit(self.catalog,{'E2':999})
        with self.assertRaisesRegex(Stop,'остаток 0'):self.check()
    def test_candidate_recalculates_E_and_preserves_history(self):
        self.edit(self.catalog,{'E2':999});before=sha(self.catalog);p=self.candidate()
        b=load_workbook(p);s=b.active
        self.assertEqual(s['B2'].value,8);self.assertEqual(s['F2'].value,3);self.assertEqual(s['G2'].value,1);self.assertEqual(s['E2'].value,'=B2-SUM(F2:G2)')
        self.assertEqual(ledger.metadata(b)['history'][0],self.history[0]);b.close();self.assertEqual(sha(self.catalog),before)
    def test_normal_resave_reconciles_without_writing_source_or_catalog(self):
        b=load_workbook(self.app);b.properties.title='resaved';b.save(self.app);b.close()
        before={p:sha(p) for p in (self.app,self.catalog)};d=self.check()
        self.assertEqual(d['history'][0]['binding_updates'][0]['reason'],'ACCOUNTING_IDENTICAL_BYTES_CHANGED')
        self.assertEqual({p:sha(p) for p in before},before)
        p=self.candidate();b=load_workbook(p);new=ledger.metadata(b)['history'][0];b.close()
        self.assertEqual(new['sha256'],sha(self.app));self.assertEqual(new['binding_updates'][0]['previous_sha256'],self.history[0]['sha256'])
    def test_changed_quantity_stops_with_exact_mark(self):
        self.edit(self.app,{'C5':4})
        with self.assertRaisesRegex(Stop,'M: было 3, стало 4'):self.check()
    def test_changed_mark_stops_with_removed_and_added_marks(self):
        self.edit(self.app,{'A5':'X'})
        with self.assertRaisesRegex(Stop,'M: было 3, стало 0; X: было 0, стало 3'):self.check()
    def test_changed_mass_stops_against_registry(self):
        self.edit(self.app,{'E5':7})
        with self.assertRaisesRegex(Stop,'масса приложения 11 кг, в реестре 10'):self.check()
    def test_wrong_identity_fields_report_expected_and_actual(self):
        for changes,pattern in [({'A1':'Приложение № 1 к паспорту № 9/2026 от 09.09.2026 г.'},'номер.*9/2026.*1/2026'),({'A1':'Приложение № 1 к паспорту № 1/2026 от 10.09.2026 г.'},'дата 10.09.2026.*09.09.2026'),({'A2':'Объект: OTHER'},'объект.*OTHER.*SYNTHETIC OBJECT')]:
            appendix(self.app);self.edit(self.app,changes)
            with self.subTest(changes=changes),self.assertRaisesRegex(Stop,pattern):self.check()
    def test_changed_total_formula_stops(self):
        self.edit(self.app,{'E7':'=SUM(E5:E6)+1'})
        with self.assertRaisesRegex(Stop,'изменена формула итога'):self.check()
    def test_missing_registry_for_changed_bytes_stops(self):
        self.edit(self.app,{'B5':'resave'})
        with self.assertRaisesRegex(Stop,'необходима строка реестра'):self.check(entries=False)
    def test_renamed_folder_recovers_one_exact_file(self):
        target=self.folder.with_name('renamed');self.folder.rename(target);p=target/self.app.name
        d=self.check();h=d['history'][0]
        self.assertEqual(h['source'],str(p));self.assertEqual(h['binding_updates'][0]['reason'],'SOURCE_RELOCATED_SAME_BYTES')
    def test_ambiguous_relocation_stops_even_for_identical_bytes(self):
        payload=self.app.read_bytes();self.folder.rename(self.folder.with_name('renamed'))
        other=self.folder.parent/'other';other.mkdir();(other/self.app.name).write_bytes(payload)
        with self.assertRaisesRegex(Stop,'найдено.*2'):self.check()
    def test_missing_history_stops_without_creating_candidate(self):
        self.app.unlink()
        with self.assertRaisesRegex(Stop,'найдено.*0'):self.candidate()
        self.assertFalse((self.root/'candidate.xlsx').exists())
    def test_changed_unrequested_history_cell_stops(self):
        self.edit(self.catalog,{'F3':2})
        with self.assertRaisesRegex(Stop,'История|Количество истории'):self.check()
    def test_second_resave_with_different_per_mark_mass_stops(self):
        self.edit(self.app,{'B5':'resaved'});p=self.candidate();b=load_workbook(p);data=ledger.metadata(b);data['history'].pop();b.active.delete_cols(7);ledger.write_metadata(b[ledger.META],data);b.save(self.catalog);b.close()
        self.edit(self.app,{'E5':7,'E6':3})
        with self.assertRaisesRegex(Stop,'Изменена историческая масса.*M: было 6, стало 7'):self.check()


if __name__=='__main__':unittest.main()
