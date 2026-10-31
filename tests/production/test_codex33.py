import unittest,tempfile,hashlib
from pathlib import Path
from unittest.mock import patch
from decimal import Decimal
from openpyxl import Workbook,load_workbook
from production_app.storage import Stop,sha,save,read as json_read
from production_app.shipment_input import parse_tsv,normalize,write_internal
from production_app import model
from production_app.project_catalog import create_catalog,resolve_catalog,catalog_name
from production_app.catalog_ledger import metadata,refresh_history
from production_app.owner_control import declaration,require_declaration,STATEMENT

class ShipmentTests(unittest.TestCase):
    def test_5_columns_unknown_area(self):
        v=parse_tsv('М1\tБалка\t2\t10,2\t20,4');self.assertIsNone(v.area);self.assertEqual(v.mass,Decimal('20.4'))
    def test_6_columns_derives_total(self):self.assertEqual(parse_tsv('М1\tБалка\t2\t10\t20\t1,5').area,3)
    def test_7_columns_total_only(self):self.assertEqual(parse_tsv('М1\tБалка\t2\t10\t20\t\t3').rows[0][5],Decimal('1.5'))
    def test_8_without_header(self):self.assertEqual(parse_tsv('12\tМ1\tБалка\t2\t10\t20\t1\t2').quantity,2)
    def test_8_header(self):
        v=parse_tsv('Лист №\tМарка\tНаименование\tКоличество\tМасса\tВсего\tПлощадь\tВсего\n1\tМ1\tБалка\t2\t10\t20\t1\t2');self.assertEqual(v.area,2)
    def test_ambiguous_8_rejected(self):
        with self.assertRaisesRegex(Stop,'8 колонок'):parse_tsv('x\tМ1\tБалка\t2\t10\t20\t1\t2')
    def test_mixed_rows_rejected(self):
        with self.assertRaises(Stop):parse_tsv('Лист №\tМарка\tНаименование\tКоличество\tМасса\tВсего\tПлощадь\tВсего\nМ1\tБалка\t2\t10\t20')
    def test_mass_mismatch(self):
        with self.assertRaisesRegex(Stop,'масса'):parse_tsv('М1\tБалка\t2\t10\t21')
    def test_area_mismatch(self):
        with self.assertRaisesRegex(Stop,'площадь'):parse_tsv('М1\tБалка\t2\t10\t20\t1\t4')
    def test_zero_area_is_known(self):self.assertEqual(parse_tsv('М1\tБалка\t2\t10\t20\t0\t0').area,0)
    def test_104_and_repeated_marks(self):
        v=parse_tsv('\n'.join('М1\tБалка\t1\t2.5\t2.5\t1\t1' for _ in range(104)));self.assertEqual(v.quantity,104);self.assertEqual(v.unique_marks,1)
    def test_internal_roundtrip_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as d:
            for n,text in enumerate(['М1\tБалка\t2\t10\t20\t1\t2','М1\tБалка\t2\t10\t20']):
                p=Path(d)/f'{n}.xlsx';a=parse_tsv(text);write_internal(a,p);b=model.visible_vom(p);self.assertEqual(a.rows,b.rows);self.assertEqual(a.area,b.area)
                with self.assertRaises(Stop):write_internal(a,p)
    def test_no_formula_or_nonfinite(self):
        for text in ['=HYPERLINK()\tБалка\t1\t1\t1','М1\tБалка\t1\tNaN\t1','М1\tБалка\t1.5\t2\t3']:
            with self.subTest(text=text),self.assertRaises(Stop):parse_tsv(text)

class CatalogTests(unittest.TestCase):
    def test_empty_unregistered_project_existing_folder(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);r=create_catalog(p,'Новый объект');b=load_workbook(r['path']);self.assertEqual(metadata(b)['history'],[]);self.assertEqual(b.worksheets[0].max_column,5);b.close();self.assertEqual(list(p.iterdir()),[p/catalog_name('Новый объект')])
    def test_no_project_creation(self):
        with tempfile.TemporaryDirectory() as d:
            missing=Path(d)/'нет'
            with self.assertRaises(FileNotFoundError):create_catalog(missing,'нет')
            self.assertFalse(missing.exists())
    def test_conflict_and_explicit_choice(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);b=Workbook();b.save(p/'Справочник марок.xlsx');b.save(p/catalog_name('Проект'));b.close()
            with self.assertRaisesRegex(Stop,'Конфликт'):resolve_catalog(p,'Проект')
            self.assertEqual(resolve_catalog(p,'Проект',p/catalog_name('Проект')),p/catalog_name('Проект'))
            with self.assertRaises(Stop):create_catalog(p,'Проект')
    def test_first_history_unknown_total_and_no_duplicate(self):
        h=[{'number':'1/2026','date':'2026-10-01','complete':True,'quantities':{'М1':3}}]
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)
            with patch('production_app.model.read_registry') as reg,patch('production_app.catalog_ledger.registered_history',return_value=h):
                reg.return_value.projects={'Объект':[{}]};r=create_catalog(p,'Объект',{'registry':'fixture'})
            b=load_workbook(r['path']);self.assertIsNone(b.active['B2'].value);self.assertIn('IF(ISNUMBER',b.active['E2'].value);self.assertEqual(b.active['F2'].value,3);b.close()
            refresh_history(r['path'],p/'candidate.xlsx','Объект',h);b=load_workbook(p/'candidate.xlsx');self.assertEqual(len(metadata(b)['history']),1);self.assertEqual(b.active.max_column,6);b.close()

class OwnerAndSourceTests(unittest.TestCase):
    def test_engineering_preview_does_not_invent_positive_findings(self):
        from production_app.owner_control import engineering_facts
        from copy import deepcopy
        ranges={'coats':1,'air':{'min':'25','max':'27'},'surface':{'min':'20','max':'22'},'humidity':{'min':'45','max':'55'},'wet':{'min':'120','max':'140'},'layer_intervals':[{'layer':1,'min':'80','max':'100'}]}
        data={'simple_ui':True,'confirmed':False,'authorized':False,'forms_approved':False,'marks':[{'mark':'М1','profile':'П','akz':0,'uzk_quantity':1}],'akz_drafts':{'0':ranges}}
        before=deepcopy(data)
        with patch('production_app.working_facts.form_material',return_value='Материал образца'):
            result=engineering_facts(data,{})
        self.assertEqual(data,before);self.assertNotIn('vik_result',result['marks'][0]);self.assertNotIn('uzk_result',result['marks'][0]);self.assertEqual(result['evidence'],[]);self.assertFalse(result['confirmed']);self.assertEqual(result['climates'][0]['wet_thickness'],'130')
    def test_generic_approval_does_not_attest_current_control(self):
        from production_app import core
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);save(root/'verified.json',{'fingerprints':{}})
            plan={'seal':'S','run':d,'project':'П','number':'11/2026','facts':{'responsible':'','simple_ui':True,'owner_confirmation_required':True}}
            with patch('production_app.core.load_plan',return_value=plan):
                with self.assertRaisesRegex(Stop,'Явно подтвердите'):core.confirm('plan','',True,True)
                self.assertFalse((root/'confirmation.json').exists())
                core.confirm('plan','',True,True,current_control=True)
                self.assertEqual(require_declaration(plan,json_read(root/'confirmation.json'))['schema'],'VZMK_CURRENT_CONTROL_OWNER_1')
    def test_owner_declaration_pinned_to_exact_plan_and_pdf(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'verified.json';save(path,{'preview':'approved'});p={'project':'П','number':'11/2026','seal':'seal','run':d};c={'verified_sha256':sha(path),'approved':True,'reviewed':True};c['current_control']=declaration(p,c['verified_sha256']);self.assertEqual(require_declaration(p,c)['statement'],STATEMENT)
            with self.assertRaises(Stop):require_declaration(p,{**c,'approved':False})
            with self.assertRaises(Stop):require_declaration({**p,'number':'12/2026'},c)
            save(path,{'preview':'modified'})
            with self.assertRaises(Stop):require_declaration(p,c)
    def test_missing_selected_source_has_no_fallback(self):
        from production_app.working_facts import selected_form
        p={'templates':{'vik':{'path':'other'}},'first_passport_sources':{'first_number':'2/2026','sources':{}}}
        with self.assertRaisesRegex(Stop,'2/2026.*vik'):selected_form(p,'vik')
    def test_incompatible_selected_source_has_no_fallback(self):
        from production_app.working_facts import selected_form
        with tempfile.TemporaryDirectory() as d:
            from docx import Document
            path=Path(d)/'old.docx';Document().save(path)
            p={'templates':{'vik':{'path':'other'}},'first_passport_sources':{'first_number':'2/2026','sources':{'vik':{'path':str(path),'sha256':sha(path)}}}}
            with self.assertRaisesRegex(Stop,'2/2026.*несовместим'):selected_form(p,'vik')
    def test_vik_role_number_independent_of_original_suffix(self):
        from docx import Document
        from test_highlighted_reference import vik_first
        from production_app.highlighted_reference import _set_sdt,_plain
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'vik.docx';vik_first(p);doc=Document(p);_set_sdt(doc,'PASSPORT_NO','11/2026',2)
            texts=[_plain(n) for n in doc.element.body.iter(qn('w:sdt')) if n.find('.//'+qn('w:tag')) is not None and n.find('.//'+qn('w:tag')).get(qn('w:val'))=='PASSPORT_NO'];self.assertEqual(texts,['11/2026','11/2026-1'])
            import re
            numbers=re.findall(r'АКТ № (\S+)', '\n'.join(_plain(p._p) for p in doc.paragraphs))
            self.assertEqual(numbers,['11/2026','11/2026-1'])

class OfficeShutdownTests(unittest.TestCase):
    def test_slow_graceful_exit_observed_while_sta_alive(self):
        from stage9_app.ownership import wait_owned_exit
        clock=[0.0];pumps=[]
        def elapsed(seconds):clock[0]+=seconds
        with patch('stage9_app.ownership.time.monotonic',side_effect=lambda:clock[0]),patch('stage9_app.ownership.time.sleep',side_effect=elapsed),patch('win32event.WaitForSingleObject',side_effect=lambda h,t:0 if clock[0]>=42 else 258),patch('pythoncom.PumpWaitingMessages',side_effect=lambda:pumps.append(clock[0])):
            self.assertTrue(wait_owned_exit([object()],60));self.assertGreater(clock[0],41);self.assertTrue(pumps)
    def test_grace_timeout_remains_a_failure(self):
        from stage9_app.ownership import wait_owned_exit
        clock=[0.0]
        def elapsed(seconds):clock[0]+=seconds
        with patch('stage9_app.ownership.time.monotonic',side_effect=lambda:clock[0]),patch('stage9_app.ownership.time.sleep',side_effect=elapsed),patch('win32event.WaitForSingleObject',return_value=258),patch('pythoncom.PumpWaitingMessages'):
            self.assertFalse(wait_owned_exit([object()],1));self.assertLess(clock[0],1.1)

if __name__=='__main__':unittest.main()
