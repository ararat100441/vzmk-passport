"""Production UX, current evidence, and clean-release safety regressions."""
import inspect
import unittest
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock,patch

from production_app import gui,office
from production_app.storage import Stop,sha
from production_app.working_facts import apply_current_controls,save_controls


class EffectiveForms(unittest.TestCase):
    def test_recognized_reference_needs_no_unused_clean_template_fields(self):
        from test_highlighted_reference import vik_first,uzk_first
        from test_quality_from_first import build_first
        from production_app.working_facts import required_doc_fields,selected_form
        from production_app.real_forms import require_form_inputs
        with TemporaryDirectory() as td:
            root=Path(td);refs={}
            for role,builder in [('vik',vik_first),('uzk',uzk_first),('quality',build_first)]:
                path=root/(role+'.docx');builder(path);refs[role]={'path':str(path),'sha256':sha(path)}
            profile={'templates':{r:{'path':'unused-clean-form.docx'} for r in refs},'first_passport_sources':{'sources':refs}}
            self.assertTrue(all(selected_form(profile,r).get('exact_reference') for r in refs))
            self.assertEqual(required_doc_fields(profile),[])
            self.assertTrue(require_form_inputs(profile,{}))
            Path(refs['vik']['path']).write_bytes(b'changed')
            with self.assertRaisesRegex(Stop,'Изменён'):selected_form(profile,'vik')

    def test_unrecognized_filled_reference_does_not_bypass_clean_form(self):
        from docx import Document
        from production_app.working_facts import selected_form
        with TemporaryDirectory() as td:
            path=Path(td)/'other.docx';d=Document();d.add_paragraph('Other layout');d.save(path)
            clean={'path':'CLEAN'};profile={'templates':{'vik':clean},'first_passport_sources':{'sources':{'vik':{'path':str(path),'sha256':sha(path)}}}}
            self.assertIs(selected_form(profile,'vik'),clean)


class WorkingControls(unittest.TestCase):
    def fixture(self,folder):
        proof=Path(folder)/'current.txt';proof.write_text('TRAINING / synthetic isolated fixture',encoding='utf8')
        controls={role:{'path':str(proof),'result':'ГОДЕН','material':'ЛКМ'} for role in ('VIK','UZK','AKZ0','AKZ1')}
        return {'confirmed':False,'authorized':False,'forms_approved':False,
                'current_controls':save_controls(controls),
                'marks':[{'mark':'M1','profile':'P','akz':'0','uzk_quantity':8},
                         {'mark':'M2','profile':'P','akz':'1','uzk_quantity':0}]},proof

    def test_all_current_proofs_are_role_bound_without_approval(self):
        with TemporaryDirectory() as td:
            data,_=self.fixture(td);before=deepcopy(data)
            result=apply_current_controls(data)
            self.assertEqual(data,before)
            self.assertEqual(result['marks'][0]['vik_result'],'ГОДЕН')
            self.assertEqual(result['marks'][1]['uzk_result'],'НЕ ТРЕБУЕТСЯ')
            self.assertEqual(result['marks'][0]['paint'],'ЛКМ')
            self.assertEqual(result['akz_proofs']['1'][0]['variant'],1)
            self.assertEqual(len(result['evidence']),4)
            for flag in ('confirmed','authorized','forms_approved'):self.assertFalse(result[flag])

    def test_missing_verdict_is_not_inferred_from_old_mark(self):
        with TemporaryDirectory() as td:
            data,_=self.fixture(td)
            data['current_controls']['VIK']['result']=''
            data['marks'][0]['vik_result']='ГОДЕН'
            with self.assertRaisesRegex(Stop,'ВИК.*заключение'):
                apply_current_controls(data)

    def test_negative_verdict_cannot_issue(self):
        with TemporaryDirectory() as td:
            data,_=self.fixture(td);data['current_controls']['UZK']['result']='НЕ ГОДЕН'
            with self.assertRaisesRegex(Stop,'УЗК.*заключение'):apply_current_controls(data)

    def test_unchanged_path_cannot_silently_refresh_changed_protocol(self):
        with TemporaryDirectory() as td:
            data,path=self.fixture(td);old=data['current_controls']
            path.write_text('changed',encoding='utf8')
            data['current_controls']=save_controls(old,old)
            with self.assertRaisesRegex(Stop,'протокол изменён'):apply_current_controls(data)

    def test_explicit_protocol_reselection_gets_fresh_pin(self):
        with TemporaryDirectory() as td:
            data,path=self.fixture(td);old=sha(path)
            path.write_text('TRAINING new current fixture',encoding='utf8')
            data['current_controls']=save_controls(data['current_controls'])
            self.assertNotEqual(old,data['current_controls']['VIK']['sha256'])
            apply_current_controls(data)

    def test_partial_entries_save_without_verdict_or_authorization(self):
        self.assertEqual(save_controls({'VIK':{'path':'','result':''}}),{'VIK':{'path':'','result':''}})

    def test_zero_uzk_does_not_confirm_absence_of_akz(self):
        with TemporaryDirectory() as td:
            data,_=self.fixture(td);data['marks'][1]['akz']=''
            with self.assertRaisesRegex(Stop,'Назначение АКЗ марки M2'):apply_current_controls(data)

    def test_missing_profiles_are_concrete_blockers(self):
        with TemporaryDirectory() as td:
            data,_=self.fixture(td);data['marks'][1]['profile']=''
            with self.assertRaisesRegex(Stop,'Профиль марки M2'):apply_current_controls(data)

    def test_many_missing_assignments_give_a_bounded_readable_dialog(self):
        with TemporaryDirectory() as td:
            data,_=self.fixture(td)
            data['marks']=[{'mark':'M'+str(n),'profile':'','akz':'','uzk_quantity':0} for n in range(49)]
            with self.assertRaises(Stop) as caught:apply_current_controls(data)
            message=str(caught.exception)
            self.assertIn('для 49 марок',message)
            self.assertLess(len(message.splitlines()),10)
            self.assertIn('таблице марок',message)

    def test_no_uzk_protocol_required_when_all_current_plan_is_zero(self):
        with TemporaryDirectory() as td:
            data,_=self.fixture(td)
            for mark in data['marks']:mark['uzk_quantity']=0
            del data['current_controls']['UZK']
            result=apply_current_controls(data)
            self.assertTrue(all(m['uzk_result']=='НЕ ТРЕБУЕТСЯ' for m in result['marks']))


class ReleaseMarks(unittest.TestCase):
    def test_qualification_attestation_is_not_a_test(self):
        office.verify_document_marks('Сведения об аттестации сварщиков. Аттестат № 21',draft=False)

    def test_actual_test_evidence_words_still_block_production_documents(self):
        for word in ('TEST','ТЕСТ','TRAINING','УЧЕБНЫЕ ДАННЫЕ','SAMPLE_ONLY'):
            with self.subTest(word=word),self.assertRaises(Stop):office.verify_document_marks(word)

    def test_preview_mark_is_allowed_only_before_release(self):
        office.verify_document_marks(office.PREVIEW_MARK)
        with self.assertRaises(Stop):office.verify_document_marks(office.PREVIEW_MARK,draft=False)

    def test_legacy_unissued_marks_cannot_survive_final_release(self):
        for word in ('ЧЕРНОВИК','НЕ ДЛЯ ВЫПУСКА','УЧЕБНЫЙ КОМПЛЕКТ'):
            with self.subTest(word=word),self.assertRaises(Stop):
                office.verify_document_marks(word,training=True,draft=False)

    def test_internal_fixture_also_uses_real_cleanup_path(self):
        source=inspect.getsource(office.release_copies)
        self.assertNotIn('if training:return',source)
        self.assertIn('CenterHeader',source)
        self.assertIn("word_preview_mark(out['path'],remove=True)",source)

    def test_word_preview_roundtrip_preserves_body_and_opaque_parts(self):
        from docx import Document
        from zipfile import ZipFile
        with TemporaryDirectory() as td:
            path=Path(td)/'reference-copy.docx'
            doc=Document();doc.add_paragraph('Аттестация сварщиков. Постоянные реквизиты.');doc.add_table(rows=2,cols=3);doc.save(path)
            with ZipFile(path,'a') as z:z.writestr('customXml/keep.xml',b'<original/>')
            original=Document(path).element.body.xml
            with ZipFile(path) as z:pins={n:z.read(n) for n in z.namelist() if n!='word/document.xml'}
            office.word_preview_mark(path)
            self.assertIn(office.PREVIEW_MARK,Document(path).paragraphs[0].text)
            office.word_preview_mark(path,remove=True)
            self.assertEqual(Document(path).element.body.xml,original)
            with ZipFile(path) as z:
                self.assertEqual(pins,{n:z.read(n) for n in pins})

    def test_cannot_remove_an_unknown_word_first_paragraph(self):
        from docx import Document
        with TemporaryDirectory() as td:
            path=Path(td)/'document.docx';doc=Document();doc.add_paragraph('Исходный текст');doc.save(path);before=sha(path)
            with self.assertRaises(Stop):office.word_preview_mark(path,remove=True)
            self.assertEqual(sha(path),before)


class SingleWorkflow(unittest.TestCase):
    def test_managed_install_uses_own_assets_without_changing_working_paths(self):
        from production_app.path_config import installed_assets
        with TemporaryDirectory() as td:
            root=Path(td);new=root/'r5';(new/'Templates').mkdir(parents=True)
            cfg={'auto_first_passport':True,'profiles':{},'documents':str(root/'r4/Templates'),'registry':'WORK.xlsx','output':'WORK DOCUMENTS'}
            out=installed_assets(cfg,new)
            self.assertEqual(out['documents'],str(new/'Templates'))
            self.assertEqual(out['registry'],cfg['registry']);self.assertEqual(out['output'],cfg['output'])
            self.assertEqual(cfg['documents'],str(root/'r4/Templates'))
            for extra in ({'training':True},{'profiles':{'PROJECT':{'root':'CUSTOM'}}},{'documents':str(root/'custom')}):
                value={**cfg,**extra};self.assertEqual(installed_assets(value,new),value)

    def test_one_button_and_no_user_mode_selectors(self):
        source=inspect.getsource(gui.App.__init__)
        self.assertIn("text='Новый паспорт'",source)
        self.assertIn("text='Сформировать пакет'",source)
        for label in ('Учебный режим','Режим черновика','TEST/SAMPLE_ONLY'):
            self.assertNotIn(label,source)

    def test_one_review_confirmation_and_no_second_release_prompt(self):
        source=inspect.getsource(gui.App.review_preview)
        self.assertEqual(source.count('messagebox.askyesno'),1)
        self.assertIn('self.issue()',source)
        self.assertNotIn('messagebox.askyesno',inspect.getsource(gui.App.issue))

    def test_current_controls_are_saved_by_back_not_as_approval(self):
        source=inspect.getsource(gui.App.fact_dialog)
        self.assertIn("draft['current_controls']=save_controls",source)
        self.assertIn("_pending_editor_commit",source)
        self.assertNotIn("draft['confirmed']=True",source)

    def test_missing_current_controls_stop_before_worker(self):
        app=SimpleNamespace(active=None,plan=None,preview_pdf=None,
            cfg={'registry':'R','documents':'D','output':'O'},project=Mock(get=lambda:'P'),
            number=Mock(get=lambda:'11/2026'),vom=Mock(get=lambda:'V'),
            facts={'simple_ui':True,'marks':[{'mark':'M','profile':'P','akz':0,'uzk_quantity':8}]},
            error=Mock(),run=Mock())
        app.shipment_facts=lambda:deepcopy(app.facts)
        gui.App.check_official(app)
        app.run.assert_not_called();app.error.assert_called_once()
        self.assertIn('первичный протокол',str(app.error.call_args))


if __name__=='__main__':unittest.main(verbosity=2)
