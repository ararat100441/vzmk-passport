"""r7 reproductions. Disposable fixtures; no production inspection results."""
import inspect,os,tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from copy import deepcopy
from unittest.mock import patch
from openpyxl import Workbook
from production_app.storage import Stop,sha,safe
from production_app import automatic_acts,gui
from production_app.simple_facts import save_screen_draft,prepare_for_check
from production_app.shared_uzk import verify_planned_quantities
from test_automatic_acts import make_book

def ranges():
    return {'coats':1,'air':{'min':'25,0','max':'27'},
            'surface':{'min':'20','max':'22,0'},'humidity':{'min':'40','max':'60'},
            'wet':{'min':'120','max':'140'},'layer_intervals':[{'layer':1,'min':'80','max':'90'}]}

class Climate(unittest.TestCase):
    def data(self):return {'simple_ui':True,'marks':[{'mark':'M','profile':'16П','akz':0,'uzk_quantity':0}],'akz_drafts':{'0':ranges()}}
    def test_exact_strings_survive_back_json_and_worker_conversion(self):
        import json
        data=self.data();saved=save_screen_draft(data,'P','2/2026',True,data['marks'],{},data['akz_drafts'])
        result=prepare_for_check(json.loads(json.dumps(saved)))['climates'][0]
        self.assertEqual(result['measurement_intervals']['surface'],{'min':'20','max':'22,0'})
        self.assertEqual((result['surface'],result['wet_thickness'],result['thickness']),('21.0','130','85'))
    def test_invalid_range_reports_variant_and_visible_field(self):
        for value in ('','141','NaN','text'):
            with self.subTest(value=value):
                d=self.data();d['akz_drafts']['0']['surface']['min']=value
                with self.assertRaisesRegex(Stop,'АКЗ №0: проверьте температуру поверхности — От'):prepare_for_check(d)
    def test_partial_back_is_allowed_and_no_fake_default(self):
        d=self.data();d['akz_drafts']['0']['surface']['min']=''
        saved=save_screen_draft(d,'P','2/2026',True,d['marks'],{},d['akz_drafts'])
        self.assertEqual(saved['akz_drafts']['0']['surface']['min'],'')
        self.assertFalse(saved['confirmed'])

class Shipment(unittest.TestCase):
    def test_zero_plan_keeps_all_16_marks_and_only_two_selected(self):
        from production_app.shipment_sync import synchronize
        plan={f'M{i}':{'uzk_quantity':1 if i<2 else 0} for i in range(16)}
        d=synchronize({'akz_drafts':{'0':ranges()},'marks':[{'mark':'OLD','uzk_quantity':4}]},'P','2/2026',True,plan,{'M0':{'profile':'16П','akz':0}})
        self.assertEqual(len(d['marks']),16);self.assertEqual(sum(r['uzk_quantity'] for r in d['marks']),2)
        self.assertTrue(verify_planned_quantities(d,plan));self.assertNotIn('OLD',[r['mark'] for r in d['marks']])
        self.assertEqual(d['akz_drafts']['0'],ranges());self.assertEqual(d['marks'][1]['profile'],'');self.assertEqual(d['marks'][1]['akz'],'')
        self.assertFalse(d['confirmed'])
    def test_mismatch_lists_missing_extra_and_counts_and_duplicates(self):
        plan={'M1':{'uzk_quantity':0},'M2':{'uzk_quantity':2}}
        for rows,parts in [([{'mark':'OLD','uzk_quantity':0}],('M1','M2','OLD','ожидалось 2','найдено 1')),
                           ([{'mark':'M1','uzk_quantity':0}]*2,('Повтор','M1'))]:
            with self.subTest(rows=rows):
                with self.assertRaises(Stop) as cm:verify_planned_quantities({'marks':rows},plan)
                for part in parts:self.assertIn(part,str(cm.exception))
    def test_empty_engineering_data_does_not_pass_as_complete(self):
        from production_app.owner_control import engineering_facts
        with self.assertRaisesRegex(Stop,'марки|Фактические'):engineering_facts({}, {})
    def test_repeated_paste_replaces_table(self):
        from production_app.shipment_dialog import replace_rows
        class Grid:
            def __init__(self):self.rows={'old':('OLD',)*7}
            def get_children(self):return list(self.rows)
            def delete(self,*keys):
                for key in keys:self.rows.pop(key)
            def insert(self,*a,values):self.rows[str(len(self.rows))]=values
        g=Grid();replace_rows(g,[('M1',None,1,2,2,None,None),('M2','N',1,1,1,1,1)])
        replace_rows(g,[('M3','N',1,1,1,1,1)])
        self.assertEqual(len(g.rows),1);self.assertEqual(next(iter(g.rows.values()))[0],'M3')
    def test_settings_has_no_standard_editor_and_only_add_row_removed(self):
        self.assertNotIn('ГОСТ / НД текущих марок',inspect.getsource(gui.App.settings))
        from production_app import shipment_dialog
        code=inspect.getsource(shipment_dialog.open_shipment)
        self.assertNotIn("text='Добавить строку'",code)
        self.assertIn("text='Удалить строку'",code)

class Standards(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)/'Профили.xlsx';make_book(self.path,duplicate=True)
        from openpyxl import load_workbook
        w=load_workbook(self.path);w.active['D2']='Z FIRST';w.active['D3']='A SECOND';w.active['H3']=7;w.save(self.path);w.close()
    def tearDown(self):self.tmp.cleanup()
    def selector(self,cfg=None):
        from production_app.geometry_selection import GeometrySelection
        return GeometrySelection(self.path,cfg or {},'P')
    def test_physical_order_not_alphabetical(self):
        self.assertEqual(automatic_acts.profile_standards(self.path,'16П'),['Z FIRST','A SECOND'])
    def test_120_marks_defaults_shared_override_exception_and_reopen(self):
        cfg={};s=self.selector(cfg);before=sha(self.path)
        self.assertEqual({s.resolve(f'M{i}','16П') for i in range(120)},{'Z FIRST'})
        s.choose('M0','16П','Z FIRST',individual=True)
        s.choose('M1','[16П','A SECOND',individual=False)
        import json
        s=self.selector(json.loads(json.dumps(s.settings())))
        self.assertEqual(s.resolve('M0','16П'),'Z FIRST')
        self.assertEqual({s.resolve(f'M{i}','16П') for i in range(1,120)},{'A SECOND'})
        self.assertEqual(sha(self.path),before)
    def test_120_marks_with_different_profiles_use_independent_defaults(self):
        from openpyxl import load_workbook
        w=load_workbook(self.path);values=[w.active.cell(2,c).value for c in range(1,17)]
        values[0]=values[1]='20Б1';values[2]='Двутавр';values[3]='OTHER FIRST';values[14]='ДВУТАВР_ПРОКАТНЫЙ_1';w.active.append(values);w.save(self.path);w.close()
        s=self.selector()
        for i in range(120):self.assertEqual(s.resolve(f'M{i}','16П' if i%2 else '20Б1'),'Z FIRST' if i%2 else 'OTHER FIRST')
    def test_legacy_mark_choices_remain_explicit_exceptions(self):
        s=self.selector({'geometry_choices':{'P':{'M0':'A SECOND'}}})
        self.assertEqual(s.resolve('M0','16П'),'A SECOND')
        s.choose('M1','16П','Z FIRST',individual=False)
        self.assertEqual(s.resolve('M0','16П'),'A SECOND')
        s.choose('M0','16П','Z FIRST',individual=False)
        self.assertEqual(s.resolve('M0','16П'),'Z FIRST')
    def test_saved_invalid_choice_not_silently_replaced(self):
        s=self.selector();s.choose('M1','16П','A SECOND',individual=False);cfg=s.settings()
        from openpyxl import load_workbook
        w=load_workbook(self.path);w.active.delete_rows(3);w.save(self.path);w.close()
        with self.assertRaisesRegex(Stop,'[Ии]змен.*источник|источник.*измен'):self.selector(cfg).resolve('M1','16П')
    def test_missing_profile_is_specific_error(self):
        with self.assertRaisesRegex(Stop,'M1.*НЕИЗВЕСТЕН|НЕИЗВЕСТЕН.*M1'):self.selector().resolve('M1','НЕИЗВЕСТЕН')
    def test_manual_standard_recalculates_characteristics(self):
        from openpyxl import load_workbook
        w=load_workbook(self.path)
        for line,standard in ((2,'ГОСТ 26020-83'),(3,'СТО АСЧМ 20-93')):
            for col,value in ((1,'20Б1'),(2,'20Б1'),(3,'Двутавр'),(4,standard),(15,'ДВУТАВР_ПРОКАТНЫЙ_1')):w.active.cell(line,col,value)
        w.save(self.path);w.close()
        v=SimpleNamespace(rows=[('M1','N',1,1,1,1,1)]);r={'mark':'M1','profile':'20Б1','uzk_quantity':1,'standard':'ГОСТ 26020-83'}
        first=automatic_acts.calculate_design_characteristics(v,{'marks':[r]},self.path)
        r['standard']='СТО АСЧМ 20-93';second=automatic_acts.calculate_design_characteristics(v,{'marks':[r]},self.path)
        self.assertNotEqual([x['thickness_mm'] for x in first],[x['thickness_mm'] for x in second])

@unittest.skipUnless(os.name=='nt','Windows native identity')
class NativeFiles(unittest.TestCase):
    def test_junction_in_parent_is_blocked(self):
        import _winapi
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);target=root/'target';target.mkdir();(target/'file').write_text('unchanged')
            junction=root/'junction';_winapi.CreateJunction(str(target),str(junction))
            try:
                with self.assertRaisesRegex(Stop,'junction|reparse'):safe(junction/'file')
                self.assertEqual((target/'file').read_text(),'unchanged')
            finally:junction.rmdir()
    def test_unknown_stat_link_count_uses_handle_without_modifying_source(self):
        from production_app import storage
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'one.xlsx';p.write_bytes(b'unchanged');before=sha(p)
            real=Path.stat
            def incomplete(path,*args,**kw):
                s=real(path,*args,**kw)
                if path==p:return SimpleNamespace(st_mode=s.st_mode,st_nlink=0,st_file_attributes=s.st_file_attributes)
                return s
            with patch.object(Path,'stat',incomplete):self.assertEqual(safe(p),p.resolve())
            self.assertEqual(sha(p),before)
    def test_real_hardlink_and_symlink_stay_blocked(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'one';p.write_text('unchanged');h=Path(d)/'two';os.link(p,h)
            with self.assertRaisesRegex(Stop,'Hardlink'):safe(p)
            h.unlink();h.symlink_to(p)
            with self.assertRaisesRegex(Stop,'Ссылка|reparse'):safe(h)
            self.assertEqual(p.read_text(),'unchanged')

class WordStartup(unittest.TestCase):
    def test_modal_without_document_window_explains_missing_nativeom(self):
        from stage9_app.ownership import startup_problem
        result=startup_problem([{'class':'#32770','title':'Microsoft Word','visible':True,
                                 'children':[{'class':'Static','text':'Startup diagnostic'}]}],'_WwG')
        self.assertIn('диалог',result);self.assertIn('Startup diagnostic',result)
    def test_no_dialog_and_transient_splash_is_not_classified_as_blocker(self):
        from stage9_app.ownership import startup_problem
        self.assertEqual(startup_problem([{'class':'OpusApp','title':'','visible':False,'children':[]}],'_WwG'),'')

if __name__=='__main__':unittest.main(verbosity=2)
