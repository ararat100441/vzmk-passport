"""Exact annotated FIRST-passport act copies (all data are SYNTHETIC).

User originals are never published to GitHub. These fixtures model the
document structure: 2 VIK numbered/date content-controls, an anchored
VIK product name, a 9-logical-column UZK table with merged characteristic,
and an old external file hyperlink on the mark.
"""
import unittest
from tempfile import TemporaryDirectory
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from production_app.highlighted_reference import (
    create_vik_copy,create_uzk_copy,
)
from production_app.storage import Stop,sha


def _sdt(paragraph,tag,text):
    s=OxmlElement('w:sdt')
    pr=OxmlElement('w:sdtPr')
    t=OxmlElement('w:tag')
    t.set(qn('w:val'),tag)
    pr.append(t)
    s.append(pr)
    content=OxmlElement('w:sdtContent')
    run=OxmlElement('w:r')
    txt=OxmlElement('w:t')
    txt.text=text
    run.append(txt)
    content.append(run)
    s.append(content)
    paragraph._p.append(s)


def vik_first(path):
    doc=Document()
    for act in range(2):
        p=doc.add_paragraph('АКТ № ')
        _sdt(p,'PASSPORT_NO','1/2026')
        p.add_run('-1' if act else '')
        p.add_run(' от ')
        _sdt(p,'VIK_DATE','02.09.2026')
        if act:
            p=doc.add_paragraph()
            p.add_run('НШ, ТШ, СШ соединения ')
            changing=p.add_run('Колонна, Пластина монтажная, Шайба, Элемент монтажный')
            changing.font.highlight_color=WD_COLOR_INDEX.YELLOW
            p.add_run(' МП в соответствии с требованиями технологии сварки')
        else:
            doc.add_paragraph('Контроль выполнен согласно РД 34.10.130-96')
    doc.add_paragraph('Постоянный текст: Костицын Е.А., Журавлев А.Н.')
    doc.save(path)


def _mark_hyperlink(doc,cell,mark):
    par=cell.paragraphs[0]
    rid=doc.part.relate_to('file:///C:/OLD/DRAWINGS/'+mark+'.pdf',
                          RT.HYPERLINK,is_external=True)
    hyperlink=OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'),rid)
    run=OxmlElement('w:r')
    rp=OxmlElement('w:rPr')
    hl=OxmlElement('w:highlight')
    hl.set(qn('w:val'),'yellow')
    rp.append(hl)
    run.append(rp)
    text=OxmlElement('w:t')
    text.text=mark
    run.append(text)
    hyperlink.append(run)
    par._p.append(hyperlink)


def uzk_first(path):
    doc=Document()
    p=doc.add_paragraph()
    p.add_run('Акт № ')
    p.add_run('1/2026')
    p.add_run(' от ')
    p.add_run('04.09.2026')
    doc.add_paragraph('Дефектоскоп А1212 MASTER, заводской № 3121397')
    doc.add_paragraph('Комплект КОУ-2, постоянные реквизиты')
    table=doc.add_table(rows=6,cols=9)
    for i in range(5):
        if i==4:
            table.rows[i].cells[1].text='Характеристика сварного соединения'
    r=table.rows[5]
    r.cells[1].merge(r.cells[2])
    _mark_hyperlink(doc,r.cells[0],'К1-5')
    for index,value in {
        1:'L=810х9 мм\nL=1794х12 мм',
        3:'Иванов С.В., клеймо H630',
        4:'04.09.2026',
        5:'5\n2,5',
        6:'70\n65',
        7:'ДНО',
        8:'годен',
    }.items():
        r.cells[index].text=value
        if index in (1,4,5,6):
            for p in r.cells[index].paragraphs:
                for run in p.runs:
                    run.font.highlight_color=WD_COLOR_INDEX.YELLOW
    other=doc.add_table(rows=1,cols=3)
    other.rows[0].cells[0].text='Контроль выполнил'
    other.rows[0].cells[2].text='Журавлев А.Н.'
    doc.save(path)


class FirstPassportCopiedFields(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.vik=self.root/'1_2026 - ВИК.docx'
        self.uzk=self.root/'1_2026 - УЗК.docx'
        vik_first(self.vik);uzk_first(self.uzk)
    def tearDown(self):self.temp.cleanup()

    def test_vik_changes_only_both_headers_and_product_names(self):
        before=sha(self.vik)
        out=self.root/'11_2026 - ВИК.docx'
        proof=create_vik_copy(
            self.vik,out,number='11/2026',vik_date='2026-10-18',
            item_names=('Балка','Балка'))
        doc=Document(out)
        xml=doc.element.xml
        self.assertEqual(xml.count('11/2026'),2)
        self.assertEqual(xml.count('18.10.2026'),2)
        self.assertEqual(xml.count('Балка'),1)
        self.assertNotIn('Пластина монтажная',xml)
        self.assertIn('НШ, ТШ, СШ соединения',xml)
        self.assertIn(' МП в соответствии',xml)
        self.assertIn('Костицын Е.А.',xml)
        self.assertIn('РД 34.10.130-96',xml)
        self.assertNotIn('w:highlight w:val="yellow"',xml)
        self.assertEqual(proof['item_names'],['Балка'])
        self.assertEqual(sha(self.vik),before)

    def test_uzk_replaces_rows_but_preserves_nonhighlighted_fields(self):
        before=sha(self.uzk)
        out=self.root/'11_2026 - УЗК.docx'
        geometry=[
            {'mark':'Бп2-1','length_mm':'1280','thickness_mm':'5',
             'frequency_mhz':'5','angle_deg':'70'},
            {'mark':'Бп2-1','length_mm':'1024','thickness_mm':'8.4',
             'frequency_mhz':'5','angle_deg':'70'},
        ]
        result=create_uzk_copy(self.uzk,out,number='11/2026',
                               uzk_date='2026-10-20',design_rows=geometry)
        d=Document(out)
        self.assertIn('Акт № 11/2026 от 20.10.2026',d.paragraphs[0].text)
        self.assertEqual(len(d.tables),2)
        self.assertEqual(len(d.tables[0].rows),6)
        row=d.tables[0].rows[5]
        self.assertEqual(row.cells[0].text,'Бп2-1')
        self.assertEqual(row.cells[1].text,
                         'L=1280х5 мм\nL=1024х8,4 мм')
        self.assertEqual(row.cells[2].text,row.cells[1].text)
        self.assertEqual(row.cells[4].text,'20.10.2026')
        self.assertEqual(row.cells[5].text,'5\n5')
        self.assertEqual(row.cells[6].text,'70\n70')
        self.assertEqual(row.cells[3].text,'Иванов С.В., клеймо H630')
        self.assertEqual(row.cells[7].text,'ДНО')
        self.assertEqual(row.cells[8].text,'годен')
        self.assertEqual(d.tables[1].rows[0].cells[2].text,'Журавлев А.Н.')
        self.assertFalse(any(r.reltype.endswith('/hyperlink')
                             for r in d.part.rels.values()))
        self.assertNotIn('K1-5.pdf',d.element.xml)
        self.assertNotIn('К1-5',d.element.xml)
        self.assertEqual(result['marks'],('Бп2-1',))
        self.assertEqual(sha(self.uzk),before)

    def test_multiple_uzk_marks_not_multiplied_by_item_count(self):
        geometry=[
            {'mark':'Бп2-1','length_mm':'1280','thickness_mm':'5',
             'frequency_mhz':'5','angle_deg':'70','quantity':8},
            {'mark':'Бп2-1','length_mm':'1024','thickness_mm':'8.4',
             'frequency_mhz':'5','angle_deg':'70','quantity':8},
            {'mark':'Бп2-19','length_mm':'500','thickness_mm':'6',
             'frequency_mhz':'5','angle_deg':'70','quantity':2},
        ]
        out=self.root/'11_2026 - УЗК.docx'
        create_uzk_copy(self.uzk,out,number='11/2026',
                        uzk_date='2026-10-20',design_rows=geometry)
        d=Document(out)
        self.assertEqual(len(d.tables[0].rows),7)
        self.assertEqual(d.tables[0].rows[6].cells[0].text,'Бп2-19')
        self.assertEqual(d.tables[0].rows[6].cells[1].text,'L=500х6 мм')

    def test_no_reuse_old_date_if_new_date_missing(self):
        with self.assertRaisesRegex(Stop,'Дата акта'):
            create_uzk_copy(self.uzk,self.root/'output.docx',
                            number='11/2026',uzk_date=None,
                            design_rows=[{'mark':'M'}])
        self.assertFalse((self.root/'output.docx').exists())

    def test_no_replace_first_source_or_existing_destination(self):
        with self.assertRaisesRegex(Stop,'отдельный файл'):
            create_vik_copy(self.vik,self.vik,number='11/2026',
                            vik_date='2026-10-18',item_names=['Балка'])
        output=self.root/'exists.docx'
        output.write_bytes(b'PRE-EXISTING')
        with self.assertRaisesRegex(Stop,'отдельный файл'):
            create_uzk_copy(self.uzk,output,number='11/2026',
                            uzk_date='2026-10-20',design_rows=[{}])
        self.assertEqual(output.read_bytes(),b'PRE-EXISTING')

    def test_uzk_no_fake_results_from_geometry(self):
        result=create_uzk_copy(self.uzk,self.root/'preview.docx',
             number='11/2026',uzk_date='2026-10-20',
             design_rows=[{'mark':'M1','length_mm':'160',
                           'thickness_mm':'6','frequency_mhz':'5',
                           'angle_deg':'70'}])
        self.assertIn('UNVERIFIED_CURRENT_CONTROL',result['status'])
        self.assertNotIn('APPROVED',result['status'])

    def test_vik_fails_closed_if_highlight_anchor_missing(self):
        d=Document(self.vik)
        for p in d.paragraphs:
            if 'НШ, ТШ, СШ соединения' in p.text:
                p.text='No approved field'
        d.save(self.vik)
        with self.assertRaisesRegex(Stop,'ровно одно место'):
            create_vik_copy(self.vik,self.root/'new.docx',
                            number='11/2026',vik_date='2026-10-18',
                            item_names=['Балка'])

    def test_no_marks_not_allowed_for_uzk_document(self):
        with self.assertRaisesRegex(Stop,'нет подтверждённых'):
            create_uzk_copy(self.uzk,self.root/'new.docx',
                            number='11/2026',uzk_date='2026-10-20',
                            design_rows=[])


if __name__=='__main__':
    unittest.main(verbosity=2)
