"""Native Office acceptance on explicit registry/template copies ONLY."""
from pathlib import Path
import hashlib,json,os,subprocess,sys,time
from uuid import uuid4
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT))
from production_app.storage import save,read,sha
from production_app import core,model

def main():
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--exe');p.add_argument('--fault',action='store_true');p.add_argument('--reuse-forms');p.add_argument('--root',required=True);args=p.parse_args()
    out=Path(args.root).resolve();out.mkdir(parents=True,exist_ok=False)
    parent=Path(r'C:\Users\arara\Desktop\ВЗМК рабочий стол\Работа')
    registry=parent/'ВЗМК (2).xlsx'
    real_template=parent/'Документы/Гараж казахстан/Паспорт/1_2026 - Приложение № 1.xlsx'
    assert registry.is_file() and real_template.is_file()
    source_pins={str(p):sha(p) for p in (registry,real_template,Path(r'C:\Users\arara\Downloads\ВОМ отгрузочный .xlsx'))}
    env=dict(os.environ);env['LOCALAPPDATA']=str(out/'local')
    if args.exe:env['PATH']=str(Path(os.environ['SystemRoot'])/'System32')
    os.environ['LOCALAPPDATA']=env['LOCALAPPDATA']
    prefix=[str(Path(args.exe).resolve())] if args.exe else [sys.executable,'-B',str(ROOT/'production.py')]
    def run(task,request,expected='PASS',fault=False):
        folder=out/'calls'/uuid4().hex;folder.mkdir(parents=True);save(folder/'request.json',request)
        executable=[sys.executable,'-B',str(ROOT/'tests/production/fault_worker.py')] if fault else prefix
        began=time.monotonic();proc=subprocess.Popen(executable+['--task',task,'--request',str(folder/'request.json'),'--result',str(folder/'result.json')],env=env)
        try:proc.wait(timeout=600)
        except subprocess.TimeoutExpired:raise RuntimeError('Controller deadline; supervisor LEFT ALIVE for safe cleanup: '+str(proc.pid))
        result=read(folder/'result.json');assert result['status']==expected,result
        summary=read(result['office_supervisor'])
        assert summary['owned_orphans']==0 and not summary['ownership_errors'],summary
        assert all(r['exit_code']==0 and not r.get('forced_termination_attempted') and not r['still_running'] for r in summary['instances']),summary
        save(folder/'TIMING.json',{'seconds':time.monotonic()-began})
        print(task,result['status'],round(time.monotonic()-began,3),flush=True)
        return result,str(folder/'result.json')
    project=out/'documents'/'project';project.mkdir(parents=True)
    import shutil
    if args.reuse_forms:shutil.copytree(args.reuse_forms,project/'forms')
    else:run('forms',{'folder':str(project/'forms')})
    shutil.copyfile(real_template,project/'appendix.xlsx')
    templates={'appendix':{'path':str(project/'appendix.xlsx'),'sha256':sha(project/'appendix.xlsx')}}
    for path in (project/'forms').glob('*'):templates[path.stem]={'path':str(path),'sha256':sha(path)}
    cfg={'registry':str(registry),'documents':str(out/'documents'),'output':str(out),'profiles':{'Гараж Казахстан':{'root':str(project),'templates':templates}}}
    cfg=core.training_copy(cfg,out)
    copy_registry=Path(cfg['registry']);initial_sha=sha(copy_registry)
    r=model.read_registry(copy_registry);pr=model.select(r,'Гараж Казахстан');number=pr.next_number
    vom=Path(r'C:\Users\arara\Downloads\ВОМ отгрузочный .xlsx');v=model.visible_vom(vom)
    from decimal import Decimal
    assert len(v.rows)==104 and v.quantity==288 and abs(v.mass-Decimal('15822.14'))<Decimal('.000001')
    evidence=out/'Учебный протокол.txt';evidence.write_text('TRAINING FIXTURE ONLY. NOT ACTUAL CONTROL RESULTS.',encoding='utf8')
    facts={'project':'Гараж Казахстан','number':number,'kind':'TRAINING','confirmed':True,'authorized':True,'forms_approved':True,
        'responsible':'Учебный исполнитель','start_date':'2026-10-01','date':'2026-10-18','vik_date':'2026-10-11','uzk_date':'2026-10-13','akz_date':'2026-10-17',
        'evidence':[{'path':str(evidence),'sha256':sha(evidence)}],
        'marks':[{'mark':mark,'profile':'','paint':'Учебный материал','akz':0,'uzk_quantity':'0','uzk_result':'НЕ ТРЕБУЕТСЯ','basis':'Учебный сценарий на копии'} for mark in dict.fromkeys(str(row[0]) for row in v.rows)],
        'climates':[{'variant':0,'air':'20','surface':'22','humidity':'45','dew':'10','thickness':'100','material':'Учебный материал','basis':'Учебный протокол послойных измерений'}]}
    request={'cfg':cfg,'project':'Гараж Казахстан','number':number,'vom':str(vom),'facts':facts,'new':True}
    check,_=run('check',request);plan=check['plan']
    assert sha(copy_registry)==initial_sha
    run('prepare',{'path':plan});assert sha(copy_registry)==initial_sha
    run('verify',{'path':plan});assert sha(copy_registry)==initial_sha
    run('confirm',{'path':plan,'responsible':facts['responsible'],'reviewed':True})
    _,release_result=run('release',{'path':plan});assert sha(copy_registry)==initial_sha
    if args.fault:
        run('commit',{'path':plan,'release_result':release_result},expected='FAIL',fault=True)
        journal=copy_registry.parent/'.vzmk-operations'/read(plan)['id']/'journal.json'
        assert sha(copy_registry)!=initial_sha
        assert core.journal_read(journal)['phase']=='REPLACE_INTENT'
        run('recover',{'journal':str(journal)})
        assert sha(copy_registry)==initial_sha and not Path(read(plan)['target']).exists()
        run('recover',{'journal':str(journal)})
        run('check',request)
        assert all(sha(path)==h for path,h in source_pins.items())
        save(out/'RESULT.json',{'status':'PASS','fault':'AFTER_REPLACEFILE_BEFORE_JOURNAL_UPDATE','recovery':'SHA_EXACT_ROLLBACK','journal':str(journal),'source_hashes_unchanged':True})
        print('FAULT_RECOVERY_PASS',flush=True);return
    committed,_=run('commit',{'path':plan,'release_result':release_result})
    assert committed['state']!='ГОТОВО'
    readback,readback_result=run('readback',{'journal':committed['journal']})
    complete,_=run('complete',{'journal':committed['journal'],'readback_result':readback_result})
    assert complete['state']=='ГОТОВО'
    assert read(committed['journal'])['phase']=='COMPLETE'
    assert sha(committed['backup'])==initial_sha
    # Final print-ready PDF must be in the SAME new folder, have the exact
    # directory name, include every section in the approved print order,
    # and be sealed by the issued manifest. No standalone PDF fragments.
    from pypdf import PdfReader
    target=Path(complete['folder'])
    print_pdf=Path(complete['print_pdf'])
    assert print_pdf.parent==target and print_pdf.name==target.name+'.pdf'
    assert print_pdf.is_file()
    issued=read(target/'ISSUE.json')
    assert issued['files'][print_pdf.name]==sha(print_pdf)
    release=read(Path(plan).parent/'release.json')
    print_spec=release['print_pdf']
    assert print_spec['sha256']==sha(print_pdf)
    assert len(PdfReader(str(print_pdf)).pages)==print_spec['pages']>0
    assert [entry['role'] for entry in print_spec['sections']] == [
        'quality','appendix','vik','uzk','akz0','akz_appendix0'
    ]
    assert {p.name for p in target.iterdir() if p.suffix.lower()=='.pdf'}=={print_pdf.name}
    run('check',request,expected='FAIL') # changed registry / duplicate number
    run('catalog',{'vom':str(vom),'target':str(out/'Черновик справочника.xlsx')})
    from openpyxl import load_workbook
    w=load_workbook(out/'Черновик справочника.xlsx');s=w.active
    assert s.max_row==106 and all(s.cell(row,col).value is None for row in range(3,107) for col in (4,5,6));w.close()
    assert all(sha(path)==h for path,h in source_pins.items())
    save(out/'RESULT.json',{'status':'PASS','exe':args.exe,'cfg':cfg,'manifest':str(Path(complete['folder'])/'ISSUE.json'),'journal':committed['journal'],
        'marks':104,'quantity':288,'mass_kg':'15822.14','number':number,'training':True,'actual_control_results':False,'source_hashes_unchanged':True})
    print('E2E_PASS',flush=True)

if __name__=='__main__':main()
