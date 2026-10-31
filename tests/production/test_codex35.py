"""CODEX35: conflicting geometry and unclassified source graphics."""
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from zipfile import ZipFile
from unittest.mock import patch
import unittest
from openpyxl import load_workbook
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from production_app import automatic_acts as A,core
from production_app.geometry_selection import GeometrySelection
from production_app.storage import Stop,sha
from test_automatic_acts import make_book

class DuplicateGeometry(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory();self.path=Path(self.temp.name)/'Профили.xlsx';make_book(self.path,duplicate=True)
        w=load_workbook(self.path);w.active.cell(3,4,'ГОСТ 8240-97');w.save(self.path);w.close()
        self.vom=SimpleNamespace(rows=[('M','N',1,1,1,1,1)])
        self.facts={'marks':[{'mark':'M','profile':'16П','uzk_quantity':1,'standard':'ГОСТ 8240-97'}]}
    def tearDown(self):self.temp.cleanup()
    def calculate(self):return A.calculate_design_characteristics(self.vom,self.facts,self.path)
    def change(self,col,value,line=3):
        w=load_workbook(self.path);w.active.cell(line,col,value);w.save(self.path);w.close()
    def test_identical_duplicate_uses_first_physical_row(self):
        before=sha(self.path);result=self.calculate();self.assertEqual({x['geometry_reference']['row'] for x in result},{2});self.assertEqual(sha(self.path),before)
    def test_equivalent_numeric_representations_are_identical(self):
        self.change(8,'5.000');self.assertTrue(self.calculate())
    def test_different_thickness_is_rejected_with_both_rows(self):
        self.change(8,7);before=sha(self.path)
        with self.assertRaisesRegex(Stop,r'Профили.xlsx.*2.*3'):self.calculate()
        self.assertEqual(sha(self.path),before)
    def test_different_dimension_not_used_in_formula_still_rejected(self):
        self.change(10,100)
        with self.assertRaisesRegex(Stop,r'Профили.xlsx.*2.*3'):self.calculate()
    def test_different_rule_is_rejected_before_calculation(self):
        self.change(15,'ДВУТАВР_ПРОКАТНЫЙ_1')
        with self.assertRaisesRegex(Stop,r'Профили.xlsx.*2.*3'):self.calculate()
    def test_different_family_is_rejected(self):
        self.change(3,'Двутавр')
        with self.assertRaisesRegex(Stop,r'Профили.xlsx.*2.*3'):self.calculate()
    def test_conflict_in_third_duplicate_is_not_ignored(self):
        w=load_workbook(self.path);w.active.append([c.value for c in w.active[3]]);w.active.cell(4,8,8);w.save(self.path);w.close()
        with self.assertRaisesRegex(Stop,r'Профили.xlsx.*2.*3.*4'):self.calculate()
    def test_different_standards_still_use_first_gost_for_120_marks(self):
        self.change(4,'OTHER STANDARD');self.change(8,7);s=GeometrySelection(self.path,{},'P')
        self.assertEqual({s.resolve(str(i),'16П') for i in range(120)},{'ГОСТ 8240-97'});self.assertTrue(self.calculate())
    def test_common_choice_and_individual_exception_preserved(self):
        self.change(4,'OTHER STANDARD');s=GeometrySelection(self.path,{},'P');s.choose('M','16П','ГОСТ 8240-97',individual=True);s.choose('N','16П','OTHER STANDARD',individual=False);reopen=GeometrySelection(self.path,s.settings(),'P');self.assertEqual(reopen.resolve('M','16П'),'ГОСТ 8240-97');self.assertEqual(reopen.resolve('N','16П'),'OTHER STANDARD')

class SourceGraphics(unittest.TestCase):
    def setUp(self):self.temp=TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def picture(self):
        p=self.root/'test.png';Image.new('RGB',(40,20),'navy').save(p);return p
    def check(self,p):return core.check_release_graphics([{'role':'vik','path':str(p)}])
    def test_no_graphics_source_is_unchanged(self):
        p=self.root/'clean.docx';Document().save(p);before=sha(p);self.check(p);self.assertEqual(sha(p),before)
    def test_logo_diagram_technical_image_signature_stamp_not_guessed(self):
        for label in ('logo','diagram','technical-image','signature','stamp'):
            with self.subTest(label=label):
                d=Document();d.add_paragraph('SYNTHETIC '+label);d.add_picture(str(self.picture()));p=self.root/(label+'.docx');d.save(p);before=sha(p)
                with self.assertRaises(Stop) as caught:self.check(p)
                self.assertIn(p.name,str(caught.exception));self.assertIn('не подтверждено',str(caught.exception));self.assertEqual(sha(p),before)
    def test_vml_graphic_is_not_silently_copied(self):
        d=Document();d.add_paragraph().add_run()._r.append(OxmlElement('w:pict'));p=self.root/'vml.docx';d.save(p)
        with self.assertRaisesRegex(Stop,'vml.docx'):self.check(p)
    def test_graphic_in_header_cannot_hide_an_old_signature(self):
        d=Document();d.sections[0].header.paragraphs[0].add_run().add_picture(str(self.picture()));p=self.root/'header.docx';d.save(p)
        with self.assertRaisesRegex(Stop,'header.docx'):self.check(p)
    def test_graphic_in_footer_is_also_unclassified(self):
        d=Document();d.sections[0].footer.paragraphs[0].add_run().add_picture(str(self.picture()));p=self.root/'footer.docx';d.save(p)
        with self.assertRaisesRegex(Stop,'footer.docx'):self.check(p)
    def test_graphic_in_footnotes_is_also_unclassified(self):
        p=self.root/'footnote.docx';Document().save(p)
        with ZipFile(p,'a') as package:
            package.writestr('word/footnotes.xml','<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:pict/></w:r></w:p></w:footnote></w:footnotes>')
        with self.assertRaisesRegex(Stop,'footnote.docx'):self.check(p)
    def test_release_checks_graphics_before_writing_output(self):
        d=Document();d.add_picture(str(self.picture()));p=self.root/'source.docx';d.save(p)
        plan={'run':str(self.root),'seal':'seal','facts':{'owner_confirmation_required':True}}
        prepared={'outputs':[{'role':'vik','path':str(p)}]}
        confirmation={'plan_seal':'seal','verified_sha256':'verified'}
        with patch.object(core,'load_plan',return_value=plan),patch.object(core,'read',side_effect=[prepared,{},confirmation]),patch.object(core,'sha',return_value='verified'),patch('production_app.owner_control.require_declaration'),patch.object(core.office,'release_copies') as release:
            with self.assertRaisesRegex(Stop,'source.docx'):core.release('plan.json')
            release.assert_not_called()
        self.assertFalse((self.root/'release').exists())
    def test_all_conflicting_documents_named_together(self):
        outputs=[]
        for role in ('vik','uzk'):
            d=Document();d.add_picture(str(self.picture()));p=self.root/(role+'.docx');d.save(p);outputs.append({'role':role,'path':str(p)})
        with self.assertRaises(Stop) as caught:core.check_release_graphics(outputs)
        self.assertIn('vik.docx',str(caught.exception));self.assertIn('uzk.docx',str(caught.exception))

if __name__=='__main__':unittest.main(verbosity=2)
