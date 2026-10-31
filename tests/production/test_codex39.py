"""Zero planned UZK: document/PDF/accounting regressions on synthetic fixtures.

These tests never stand in for acceptance of a real shipment.
"""
from contextlib import nullcontext
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import MagicMock, patch
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from production_app import office
from production_app.quality_from_first import create_quality_copy
from production_app.real_forms import acceptance_acts, require_form_inputs
from production_app.print_pdf import ordered_sources, create_print_pdf, verify_package_print
from production_app.shared_uzk import decrement_candidate, read_shared_uzk
from production_app.storage import Stop, sha
from test_quality_from_first import build_first


class ZeroUzk(TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def facts(self,quantity=0):
        return dict(number='11/2026',date='2026-10-31',start_date='2026-09-01',vik_date='2026-10-24',uzk_date='2026-10-26',akz_date='2026-10-30',marks=[dict(mark='M',uzk_quantity=quantity)],climates=[dict(variant=0)],evidence=[],responsible='')
    def quality(self,required=True):
        source=self.root/'source.docx';build_first(source)
        target=self.root/'quality.docx';before=sha(source)
        create_quality_copy(source,target,number='11/2026',end_date='2026-10-31',vik_date='2026-10-24',uzk_date='2026-10-26',akz_date='2026-10-30',mass_kg='641.95',variants=[0,2],uzk_required=required)
        self.assertEqual(sha(source),before)
        return source,target
    def test_zero_quality_retains_both_vik_and_all_akz_without_uzk(self):
        source,target=self.quality(False)
        doc=Document(target)
        for row in (50,53):
            text=doc.tables[0].cell(row,0).text
            self.assertNotIn('УЗК',text)
            self.assertEqual(text.count('Акт ВИК №'),2)
            self.assertEqual(text.count('Сертификат качества антикоррозийной обработки №'),2)
            self.assertIn('11/2026-1',text);self.assertIn('11/2026-2',text)
            self.assertNotIn('26.10.2026',text)
            self.assertIn('24.10.2026 г., Сертификат',text)
        self.assertEqual(doc.tables[0].cell(49,0).text,Document(source).tables[0].cell(49,0).text)
        with ZipFile(source) as a,ZipFile(target) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for name in a.namelist():
                if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name),name)
    def test_positive_quality_preserves_uzk_reference(self):
        _,target=self.quality(True)
        for row in (50,53):
            text=Document(target).tables[0].cell(row,0).text
            self.assertEqual(text.count('Акт УЗК № 11/2026 от 26.10.2026 г.'),1)
    def test_zero_dynamic_references_and_positive_mixed_plan(self):
        for marks,required in [([dict(uzk_quantity=0)],False),([dict(uzk_quantity='0')],False),([dict(uzk_quantity=0),dict(uzk_quantity=1)],True)]:
            text=acceptance_acts(self.facts()|{'marks':marks},[0,2])
            self.assertEqual('Акт УЗК №' in text,required)
            self.assertEqual(text.count('Акт ВИК №'),2)
            self.assertEqual(text.count('Сертификат качества АКЗ'),2)
    def test_zero_does_not_require_unused_uzk_form_fields(self):
        profile={'templates':{r:{'path':r+'.docx'} for r in ('vik','uzk','quality')}}
        seen=[]
        def inspect(path):
            seen.append(path)
            return ({'UZK_EQUIPMENT'} if path=='uzk.docx' else {'NUMBER'},set())
        with patch('production_app.word_controls.inspect_word_fields',side_effect=inspect):
            self.assertTrue(require_form_inputs(profile,self.facts()))
            self.assertNotIn('uzk.docx',seen)
            with self.assertRaisesRegex(Stop,'УЗК — дефектоскоп'):require_form_inputs(profile,self.facts(1))
    def documents(self,include_uzk=False):
        roles=['quality','appendix','vik','akz0','akz_appendix0']
        if include_uzk:roles.insert(3,'uzk')
        outputs=[]
        for role in roles:
            path=self.root/(role+('.docx' if role in ('quality','vik','uzk') else '.xlsx'))
            path.write_bytes(b'SYNTHETIC')
            outputs.append(dict(role=role,path=str(path)))
        return outputs
    def test_zero_pdf_has_complete_other_sections_and_verified_manifest(self):
        outputs=self.documents();before={r['path']:sha(r['path']) for r in outputs}
        def export(items):
            for item in items:
                with PdfWriter() as writer:
                    writer.add_blank_page(width=200,height=300)
                    writer.write(item['pdf'])
        with patch('production_app.print_pdf._export_word',side_effect=export),patch('production_app.print_pdf._export_excel',side_effect=export):
            result=create_print_pdf(outputs,self.root,'11_2026',uzk_required=False)
        released=dict(outputs=outputs+[result['output']],print_pdf=result['manifest'])
        self.assertEqual(verify_package_print(released,'11_2026',uzk_required=False)['pages'],5)
        self.assertEqual([s['role'] for s in result['manifest']['sections']],['quality','appendix','vik','akz0','akz_appendix0'])
        self.assertEqual(before,{r['path']:sha(r['path']) for r in outputs})
        with self.assertRaisesRegex(Stop,'uzk'):verify_package_print(released,'11_2026')
    def test_positive_pdf_still_requires_uzk_and_zero_refuses_fictitious_act(self):
        with self.assertRaisesRegex(Stop,'uzk'):ordered_sources(self.documents())
        with self.assertRaisesRegex(Stop,'нулевом плане'):ordered_sources(self.documents(True),uzk_required=False)
    def test_zero_pdf_does_not_accept_missing_vik_or_akz_appendix(self):
        outputs=self.documents()
        for role in ('quality','appendix','vik','akz_appendix0'):
            with self.subTest(role=role),self.assertRaises(Stop):ordered_sources([r for r in outputs if r['role']!=role],uzk_required=False)
    def uzk(self):
        source=self.root/'source.xlsm';b=Workbook();s=b.active;s.title='УЗК'
        s.append(['Заказ','Марка элемента','Кол-во (шт)']);s.append(['P','M',2]);s.append(['OTHER','N',5])
        b.save(source);b.close()
        with ZipFile(source,'a') as z:z.writestr('customXml/opaque.bin',b'PRESERVE')
        return source
    def test_zero_uzk_candidate_is_byte_identical_including_opaque_parts(self):
        source=self.uzk();before=sha(source);target=self.root/'candidate.xlsm'
        result=decrement_candidate(source,target,'P',{'M':0,'ABSENT':'0'})
        self.assertEqual(sha(source),before);self.assertEqual(sha(target),before)
        self.assertEqual(result['changed_parts'],[]);self.assertEqual(result['deleted_rows'],[])
    def test_positive_uzk_still_decrements_only_planned_quantity(self):
        source=self.uzk();before=sha(source);target=self.root/'candidate.xlsm'
        result=decrement_candidate(source,target,'P',{'M':1})
        self.assertEqual(sha(source),before)
        self.assertEqual(read_shared_uzk(target)[('p','m')]['remaining'],1)
        self.assertEqual(read_shared_uzk(target)[('other','n')]['remaining'],5)
        self.assertEqual(result['updated_rows'],{2:1})
        with ZipFile(source) as a,ZipFile(target) as b:self.assertEqual(a.read('customXml/opaque.bin'),b.read('customXml/opaque.bin'))
    def test_zero_renderer_skips_historical_uzk_and_passes_zero_to_quality(self):
        profile=dict(templates={r:{'path':str(self.root/(r+'.docx'))} for r in ('vik','uzk','quality')},first_passport_sources=dict(sources={r:{'path':str(self.root/(r+'.docx')),'sha256':'PIN'} for r in ('vik','uzk','quality')}))
        with patch('production_app.office.document_values',return_value={}),patch('production_app.model.reference_project'),patch('production_app.working_facts.selected_form') as selected,patch('production_app.office.excel_session',return_value=nullcontext(MagicMock())),patch('production_app.office.word_session',return_value=nullcontext(MagicMock())),patch('production_app.office.sha',return_value='PIN'),patch('production_app.highlighted_reference.supports_vik_first',return_value=True),patch('production_app.highlighted_reference.create_vik_copy') as vik,patch('production_app.highlighted_reference.create_uzk_copy') as uzk,patch('production_app.quality_from_first.is_exact_first_quality',return_value=True),patch('production_app.quality_from_first.create_quality_copy') as quality,patch('production_app.office.word_preview_mark'):
            outputs=office.render(profile,None,None,1,SimpleNamespace(rows=[('M','SYNTHETIC',1)],mass=Decimal('641.95')),self.facts()|{'climates':[]},self.root,True)
        self.assertEqual([o['role'] for o in outputs],['vik','quality'])
        vik.assert_called_once();uzk.assert_not_called()
        self.assertFalse(quality.call_args.kwargs['uzk_required'])
        self.assertNotIn('uzk',[c.args[1] for c in selected.call_args_list])
