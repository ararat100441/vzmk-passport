"""CODEX30: stock, opaque XLSM preservation, dates, standards and recovery.

All files are disposable fixtures. These tests do not certify current control.
"""
import json,shutil,sys,unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from datetime import date
from copy import deepcopy
from unittest.mock import patch
from zipfile import ZipFile
from openpyxl import Workbook,load_workbook
from production_app.storage import Stop,sha,read,save
from production_app import catalog_ledger as ledger, shared_uzk as uzk,coordinated,core
from production_app.automatic_acts import calculate_design_characteristics,profile_standards
from production_app.simple_facts import save_screen_draft,prepare_for_check
from production_app.akz_from_first import act_number
from test_automatic_acts import make_book

class SettingsPaths(unittest.TestCase):
    def exercise(self, store=False):
        """Emulate Tcl's variable-name binding, which does not retain StringVar."""
        import gc, weakref
        from production_app import gui, path_config, automatic_acts
        from contextlib import ExitStack
        live=weakref.WeakValueDictionary(); widgets=[]; errors=[]
        class Var:
            def __init__(self,value=''):self.value=value;self.name=str(id(self));live[self.name]=self
            def get(self):return self.value
        class Widget:
            def __init__(self,*a,**kw):
                self.options={k:(v.name if k=='textvariable' else v) for k,v in kw.items()};self.rows={};widgets.append(self)
            def grid(self,*a,**kw):return self
            def pack(self,*a,**kw):return self
            def title(self,*a):pass
            def transient(self,*a):pass
            def grab_set(self):pass
            def heading(self,*a,**kw):pass
            def column(self,*a,**kw):pass
            def insert(self,*a,values):self.rows[str(len(self.rows))]=values
            def get_children(self):return list(self.rows)
            def item(self,item,key):return self.rows[item]
            def destroy(self):pass
        cfg={'registry':'R.xlsx','output':'Documents','uzk_registry':'U.xlsm','training':False}
        with TemporaryDirectory() as tmp, ExitStack() as stack:
            app=SimpleNamespace(active=False,root=object(),cfg=dict(cfg),project=SimpleNamespace(get=lambda:'P'),vom=SimpleNamespace(get=lambda:'V.xlsx'),facts={'marks':[]},fit=lambda *a:None,editable=lambda *a,**kw:None,write=lambda *a:None,error=errors.append,choose=lambda *a:None,settings_path=lambda:Path(tmp)/'settings.json',invalidate=lambda:None,refresh=lambda:None)
            for name in ('Label','Entry','Button','LabelFrame','Treeview'):stack.enter_context(patch.object(gui.ttk,name,Widget))
            stack.enter_context(patch.object(gui.tk,'Toplevel',Widget));stack.enter_context(patch.object(gui.tk,'StringVar',Var))
            stack.enter_context(patch.object(path_config,'visible_settings',return_value={k:v for k,v in cfg.items() if k!='training'}))
            stack.enter_context(patch.object(path_config,'apply_settings',side_effect=lambda old,**kw:dict(old)|{k:kw[k] for k in ('registry','output','uzk_registry')}))
            stack.enter_context(patch.object(gui.model,'profile',return_value={'catalog':'Catalog.xlsx'}))
            stack.enter_context(patch.object(gui.model,'visible_vom',return_value=SimpleNamespace(rows=[('M','',1)])))
            stack.enter_context(patch.object(gui.model,'validate_paths'))
            stack.enter_context(patch.object(automatic_acts,'find_approved_geometry',return_value=Path('Profiles.xlsx')))
            stack.enter_context(patch.object(automatic_acts,'profile_standards',return_value=['APPROVED']))
            stack.enter_context(patch('production_app.catalog_assignments.read_assignments',return_value={'M':{'profile':'16П'}}))
            gui.App.settings(app);gc.collect()
            values=[live[w.options['textvariable']].get() if w.options['textvariable'] in live else None for w in widgets if 'textvariable' in w.options]
            self.assertEqual(values,['R.xlsx','Documents','U.xlsm','Catalog.xlsx','Profiles.xlsx'])
            if store:
                next(w.options['command'] for w in widgets if w.options.get('text')=='Сохранить рабочие пути')()
                self.assertEqual(errors,[])
                self.assertEqual(read(Path(tmp)/'settings.json')['registry'],'R.xlsx')
                # CODEX34: Settings saves paths only; ГОСТ lives in Facts.
                self.assertNotIn('geometry_choices',app.cfg)
    def test_paths_remain_visible_after_settings_callback_returns(self):self.exercise()
    def test_standard_selection_does_not_overwrite_settings_path_variables(self):self.exercise(store=True)

class Ledger(unittest.TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name);self.source=self.root/'old.xlsx';self.new=self.root/'new.xlsx'
        b=Workbook();s=b.active;s.append(['Марка','Профиль','Краска','№ АКЗ','Примечание']);s.append(['M','16П','PAINT',0,'KEEP']);s.append(['N','16П','OTHER',1,'KEEP TOO']);b.save(self.source);b.close()
        self.history=[{'number':'1/2026','date':'2026-09-09','complete':True,'quantities':{'M':8,'N':1}},
                      {'number':'2/2026','date':'2026-09-15','complete':True,'quantities':{'M':4}}]
    def tearDown(self):self.tmp.cleanup()
    def migrate(self):return ledger.create_migrated(self.source,self.new,'P',self.history)
    def total(self,qty):
        b=load_workbook(self.new);b.worksheets[0]['B2']=qty;b.save(self.new);b.close()
    def test_migration_retains_every_assignment_and_additional_source_cell(self):
        before=sha(self.source);result=self.migrate();self.assertEqual(result['marks'],2);self.assertEqual(sha(self.source),before)
        self.assertEqual(ledger.read_assignments(self.source),ledger.read_assignments(self.new))
        b=load_workbook(self.new);self.assertEqual(b[ledger.SOURCE]['E2'].value,'KEEP');self.assertEqual(b[ledger.SOURCE].sheet_state,'hidden');self.assertEqual(tuple(b.worksheets[0].cell(1,c).value for c in range(1,6)),ledger.HEADERS);b.close()
    def test_missing_total_blocks_stock(self):
        self.migrate()
        with self.assertRaisesRegex(Stop,'общее проектное количество'):ledger.validate_stock(self.new,'P','3/2026',{'M':1})
    def test_residual_across_two_passports(self):
        self.migrate();self.total(20);ledger.validate_stock(self.new,'P','3/2026',{'M':8})
        with self.assertRaisesRegex(Stop,'остаток 8, запрошено 9'):ledger.validate_stock(self.new,'P','3/2026',{'M':9})
    def test_zero_balance_retains_mark_and_paint(self):
        self.migrate();self.total(20);candidate=self.root/'candidate.xlsx';ledger.shipment_candidate(self.new,candidate,'P','3/2026','2026-10-29',{'M':8})
        b=load_workbook(candidate);s=b.worksheets[0];self.assertEqual(s['A2'].value,'M');self.assertEqual(s['H2'].value,8);self.assertEqual(s['H1'].value,'Паспорт 3/2026 от 29.10.2026');b.close()
        self.assertEqual(ledger.read_assignments(candidate)['M']['paint'],'PAINT')
    def test_missing_history_stays_blank_and_blocks(self):
        self.history[1].update(complete=False,quantities={},reason='MISSING SOURCE');self.migrate();self.total(20)
        b=load_workbook(self.new);self.assertIsNone(b.worksheets[0]['G2'].value);self.assertIsNone(b.worksheets[0]['E2'].value);b.close()
        with self.assertRaisesRegex(Stop,'отсутствует история'):ledger.validate_stock(self.new,'P','3/2026',{'M':1})
    def test_no_duplicate_migration(self):
        self.migrate()
        with self.assertRaisesRegex(Stop,'уже переведён'):ledger.create_migrated(self.new,self.root/'again.xlsx','P',self.history)
    @unittest.skipUnless(sys.platform=='win32','Native Windows exclusive lock required')
    def test_native_in_place_migration_preserves_backup_without_reopening_locked_source(self):
        before=sha(self.source);assignments=ledger.read_assignments(self.source)
        with patch.object(ledger,'registered_history',return_value=self.history):
            result=ledger.migrate_in_place(self.source,{},'P')
            self.assertEqual(sha(result['backup']),before)
            self.assertEqual(ledger.read_assignments(self.source),assignments)
            repeated=ledger.migrate_in_place(self.source,{},'P')
        b=load_workbook(self.source)
        try:self.assertEqual(len(ledger.metadata(b)['history']),2);self.assertEqual(b.worksheets[0].max_column,7)
        finally:b.close()
        self.assertEqual(ledger.read_assignments(self.source),assignments)
        self.assertEqual(repeated['marks'],2)
    def test_duplicate_history_number_rejected(self):
        self.history.append(deepcopy(self.history[0]))
        with self.assertRaisesRegex(Stop,'дублирующийся'):self.migrate()
    def test_reissue_cannot_duplicate_column(self):
        self.migrate();self.total(20);candidate=self.root/'candidate.xlsx';ledger.shipment_candidate(self.new,candidate,'P','3/2026','2026-10-29',{'M':1})
        with self.assertRaisesRegex(Stop,'уже записан'):ledger.shipment_candidate(candidate,self.root/'again.xlsx','P','3/2026','2026-10-29',{'M':1})
    def test_false_historical_quantity_rejected(self):
        self.migrate();self.total(20);b=load_workbook(self.new);b.worksheets[0]['F2']=0;b.save(self.new);b.close()
        with self.assertRaisesRegex(Stop,'подтверждённым приложением'):ledger.validate_stock(self.new,'P','3/2026',{'M':1})
    def test_wrong_registry_history_rejected(self):
        self.migrate();self.total(20)
        with self.assertRaisesRegex(Stop,'не соответствует'):ledger.validate_stock(self.new,'P','3/2026',{'M':1},[{'number':'1/2026','date':date(2026,9,9)}])
    def test_incomplete_history_can_be_refreshed_without_second_migration(self):
        complete=deepcopy(self.history);self.history[1].update(complete=False,quantities={});self.migrate();self.total(20)
        target=self.root/'refreshed.xlsx';ledger.refresh_history(self.new,target,'P',complete);ledger.validate_stock(target,'P','3/2026',{'M':8})
        self.assertEqual(ledger.read_assignments(self.new),ledger.read_assignments(target))
    def test_long_metadata_is_not_truncated_by_excel_cell_limit(self):
        self.history[0]['reason']='long provenance '*5000;self.migrate();b=load_workbook(self.new)
        try:self.assertEqual(ledger.metadata(b)['history'][0]['reason'],self.history[0]['reason']);self.assertGreater(b[ledger.META].max_row,2)
        finally:b.close()

    def test_repeated_marks_aggregate_only_same_mark(self):
        self.assertEqual(ledger.shipment_quantities(SimpleNamespace(rows=[('M','',3),('M','',2),('N','',1)])),{'M':5,'N':1})
    def test_unknown_mark_never_implicitly_added(self):
        self.migrate();self.total(20)
        with self.assertRaisesRegex(Stop,'отсутствует марка'):ledger.validate_stock(self.new,'P','3/2026',{'ABSENT':1})

class UzkCandidate(unittest.TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name);self.source=self.root/'shared.xlsm';self.target=self.root/'candidate.xlsm'
        b=Workbook();s=b.active;s.title='УЗК';s.append(uzk.HEADERS)
        for r in [('P','M',5),('P','N',None),('OTHER','X',8)]:s.append(r)
        b.create_sheet('OTHER')['A1']='KEEP ALL';b.save(self.source);b.close()
        # Opaque fixture bytes establish preservation independently of Excel.
        with ZipFile(self.source,'a') as z:z.writestr('xl/vbaProject.bin',b'OPAQUE VBA FIXTURE')
    def tearDown(self):self.tmp.cleanup()
    def test_partial_and_zero_deletion_with_other_zip_parts_exact(self):
        before=sha(self.source);r=uzk.decrement_candidate(self.source,self.target,'P',{'M':3,'N':1})
        after=uzk.read_shared_uzk(self.target);self.assertEqual(after[('p','m')]['remaining'],2);self.assertNotIn(('p','n'),after);self.assertEqual(after[('other','x')]['remaining'],8);self.assertEqual(sha(self.source),before)
        with ZipFile(self.source) as a,ZipFile(self.target) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for name in a.namelist():
                if name!=r['sheet_xml']:self.assertEqual(a.read(name),b.read(name),name)
    def test_blank_quantity_is_one_and_deleted_at_zero(self):
        uzk.decrement_candidate(self.source,self.target,'P',{'N':1});self.assertNotIn(('p','n'),uzk.read_shared_uzk(self.target))
    def test_selected_duplicate_blocks_even_identical(self):
        # Rebuild without the intentionally opaque VBA test part.
        b=load_workbook(self.source);b['УЗК'].append(['P','M',5]);b.save(self.source);b.close()
        with self.assertRaisesRegex(Stop,'дубли'):uzk.decrement_candidate(self.source,self.target,'P',{'M':1})
    def test_overconsumption_blocks_before_output(self):
        with self.assertRaisesRegex(Stop,'недостаточный'):uzk.decrement_candidate(self.source,self.target,'P',{'M':6})
        self.assertFalse(self.target.exists())
    def test_formula_depending_on_deleted_row_stops(self):
        b=load_workbook(self.source);b['OTHER']['B1']="='УЗК'!C3";b.save(self.source);b.close()
        with self.assertRaisesRegex(Stop,'формулы'):uzk.decrement_candidate(self.source,self.target,'P',{'N':1})
    def test_unassigned_other_mark_does_not_contaminate_project_plan(self):
        b=load_workbook(self.source);b['УЗК'].append([None,'UNRELATED',3]);b.save(self.source);b.close()
        plan=uzk.planned_control(self.source,'P',SimpleNamespace(rows=[('M','',3)]))
        self.assertEqual(plan['M']['uzk_quantity'],3)
        uzk.decrement_candidate(self.source,self.target,'P',{'M':3})
        b=load_workbook(self.target,read_only=True);self.assertEqual(tuple(b['УЗК'].iter_rows(min_row=5,max_row=5,values_only=True))[0],(None,'UNRELATED',3));b.close()
    def test_unassigned_current_mark_blocks_ambiguous_order(self):
        b=load_workbook(self.source);b['УЗК'].append([None,'M',3]);b.save(self.source);b.close()
        with self.assertRaisesRegex(Stop,'заказ/марка'):uzk.planned_control(self.source,'P',SimpleNamespace(rows=[('M','',3)]))

    def test_simple_validation_and_full_column_conditional_rules_survive_delete(self):
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import Rule
        b=load_workbook(self.source);sheet=b['УЗК']
        v=DataValidation(type='list',formula1='"P,OTHER"');sheet.add_data_validation(v);v.add('A2');v.add('A3')
        sheet.conditional_formatting.add('B1:B1048576',Rule(type='duplicateValues'))
        b.save(self.source);b.close()
        uzk.decrement_candidate(self.source,self.target,'P',{'N':1})
        b=load_workbook(self.target);sheet=b['УЗК'];self.assertEqual(str(sheet.data_validations.dataValidation[0].sqref),'A2');self.assertEqual(str(next(iter(sheet.conditional_formatting)).sqref),'B1:B1048576');b.close()

    def test_own_print_area_adjusts_without_rewriting_other_parts(self):
        b=load_workbook(self.source);b['УЗК'].print_area='A1:C4';b.save(self.source);b.close()
        result=uzk.decrement_candidate(self.source,self.target,'P',{'N':1})
        self.assertIn('xl/workbook.xml',result['changed_parts']);b=load_workbook(self.target);self.assertIn('$A$1:$C$3',str(b['УЗК'].print_area));b.close()

    def test_candidate_cannot_overwrite_source(self):
        with self.assertRaises(Stop):uzk.decrement_candidate(self.source,self.source,'P',{'M':1})

class StandardsAndFacts(unittest.TestCase):
    def setUp(self):self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def test_real_standard_selection_changes_exact_geometry(self):
        p=self.root/'profiles.xlsx';make_book(p,duplicate=True)
        b=load_workbook(p);b.active['H3']=6
        for row,standard in [(2,'ГОСТ 26020-83'),(3,'СТО АСЧМ 20-93')]:
            b.active.cell(row,1,'20Б1');b.active.cell(row,2,'20Б1');b.active.cell(row,3,'Двутавр');b.active.cell(row,4,standard);b.active.cell(row,15,'ДВУТАВР_ПРОКАТНЫЙ_1')
        b.save(p);b.close()
        self.assertEqual(profile_standards(p,'20Б1'),['ГОСТ 26020-83','СТО АСЧМ 20-93'])
        v=SimpleNamespace(rows=[('M','Балка',8)]);f={'marks':[{'mark':'M','profile':'20Б1','uzk_quantity':8,'standard':'СТО АСЧМ 20-93'}]}
        rows=calculate_design_characteristics(v,f,p);self.assertIn('6',[r['thickness_mm'] for r in rows]);self.assertTrue(all(r['geometry_reference']['standard']=='СТО АСЧМ 20-93' for r in rows))
        f['marks'][0]['standard']='INVENTED'
        with self.assertRaises(Stop):calculate_design_characteristics(v,f,p)
    def test_ambiguous_standard_never_first_matches(self):
        p=self.root/'profiles.xlsx';make_book(p,duplicate=True)
        with self.assertRaises(Stop):calculate_design_characteristics(SimpleNamespace(rows=[('M','',1)]),{'marks':[{'mark':'M','profile':'16П','uzk_quantity':1}]},p)
    def test_wet_and_dry_independent_midpoints_preserved_back(self):
        climate={'0':{'coats':1,'air':{'min':'23','max':'25'},'surface':{'min':'20','max':'22'},'humidity':{'min':'45','max':'55'},'wet':{'min':'120','max':'140'},'layer_intervals':[{'layer':1,'min':'80','max':'100'}]}}
        draft=save_screen_draft({},'P','11/2026',True,[{'mark':'M','profile':'16П','standard':'ГОСТ 8240-97','akz':0,'uzk_quantity':8}],{},climate)
        self.assertEqual(draft['akz_drafts']['0']['wet']['min'],'120');self.assertFalse(draft['confirmed']);result=prepare_for_check(draft)['climates'][0];self.assertEqual(float(result['wet_thickness']),130);self.assertEqual(float(result['thickness']),90)
    def test_partial_wet_input_survives_without_certification(self):
        d=save_screen_draft({},'P','11/2026',False,[],{}, {'0':{'coats':1,'wet':{'min':'120','max':''}}});self.assertEqual(d['akz_drafts']['0']['wet']['min'],'120');self.assertFalse(d['authorized'])
    def test_marked_uzk_rows_cannot_split_and_repeat_header(self):
        from test_highlighted_reference import uzk_first
        from production_app.highlighted_reference import _copy_uzk_rows
        from docx import Document
        from docx.oxml.ns import qn
        p=self.root/'first.docx';uzk_first(p);doc=Document(p)
        _copy_uzk_rows(doc,[{'mark':'M','length_mm':'1280','thickness_mm':'5','frequency_mhz':'5','angle_deg':'70'}],'23.10.2026')
        table=doc.tables[0]
        self.assertTrue(all(r._tr.trPr.find(qn('w:tblHeader')) is not None for r in table.rows[:5]))
        self.assertIsNotNone(table.rows[5]._tr.trPr.find(qn('w:cantSplit')))

    def test_numbering_is_independent_of_coats(self):
        self.assertEqual([act_number('11/2026',n) for n in range(4)],['11/2026','11/2026-1','11/2026-2','11/2026-3'])

class QualityAndDates(unittest.TestCase):
    def setUp(self):self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def test_two_akz_references_in_both_quality_points(self):
        from test_quality_from_first import build_first
        from production_app.quality_from_first import create_quality_copy
        from docx import Document
        from docx.enum.text import WD_COLOR_INDEX
        source=self.root/'first.docx';target=self.root/'out.docx';build_first(source)
        create_quality_copy(source,target,number='11/2026',end_date='2026-10-28',vik_date='2026-10-21',uzk_date='2026-10-23',akz_date='2026-10-27',mass_kg='511.98',variants=[0,1])
        doc=Document(target)
        for no in (50,53):
            paragraph=doc.tables[0].rows[no].cells[0].paragraphs[0]
            self.assertIn('обработки № 11/2026 от 27.10.2026',paragraph.text)
            self.assertIn('обработки № 11/2026-1 от 27.10.2026',paragraph.text)
            self.assertFalse(any(r.font.highlight_color==WD_COLOR_INDEX.YELLOW for r in paragraph.runs))
    def test_later_source_without_editing_highlights_is_supported(self):
        from test_quality_from_first import build_first
        from production_app.quality_from_first import create_quality_copy,is_exact_first_quality
        from docx import Document
        source=self.root/'second.docx';target=self.root/'out.docx';build_first(source)
        doc=Document(source)
        for table in doc.tables:
            seen=set()
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc in seen:continue
                    seen.add(cell._tc)
                    for p in cell.paragraphs:
                        for run in p.runs:run.font.highlight_color=None;run.text=run.text.replace('1/2026','2/2026')
        doc.save(source)
        self.assertTrue(is_exact_first_quality(source));create_quality_copy(source,target,number='11/2026',end_date='2026-10-28',vik_date='2026-10-21',uzk_date='2026-10-23',akz_date='2026-10-27',mass_kg=1)
        self.assertEqual(Document(target).tables[0].rows[1].cells[9].text,'11/2026')
    def test_generated_multi_quality_can_be_selected_as_later_sample(self):
        from test_quality_from_first import build_first
        from production_app.quality_from_first import create_quality_copy,is_exact_first_quality
        from docx import Document
        source=self.root/'first.docx';second=self.root/'second.docx';third=self.root/'third.docx';build_first(source)
        kwargs=dict(end_date='2026-10-28',vik_date='2026-10-21',uzk_date='2026-10-23',akz_date='2026-10-27',mass_kg=1)
        create_quality_copy(source,second,number='11/2026',variants=[1,2],**kwargs)
        self.assertTrue(is_exact_first_quality(second))
        create_quality_copy(second,third,number='12/2026',variants=[0,3],**kwargs)
        text=Document(third).tables[0].rows[50].cells[0].text
        self.assertIn('обработки № 12/2026 от',text);self.assertIn('обработки № 12/2026-3 от',text);self.assertNotIn('11/2026',text)
    def test_selected_source_is_exact_and_never_changes_project_first(self):
        from passport_app.registry import Project
        from production_app.first_passport import discover_first_passport_for_config
        p=Project('P',1,'1/2026',date(2026,9,9),'O',2,'2/2026','3/2026')
        cfg={'registry':'private.xlsx','first_passport_documents':str(self.root),'source_passports':{'P':'2/2026'}}
        with patch('passport_app.registry.read_registry',return_value=SimpleNamespace(projects={'P':[{'number':'2/2026','row':2,'date':date(2026,9,15)}]})),patch('production_app.first_passport.discover_first_passport',return_value={}) as discover:
            discover_first_passport_for_config(cfg,p)
            self.assertEqual(discover.call_args.args[1].first_number,'2/2026');self.assertEqual(p.first_number,'1/2026')
        cfg['source_passports']['P']='99/2026'
        with patch('passport_app.registry.read_registry',return_value=SimpleNamespace(projects={'P':[]})),patch('production_app.first_passport.discover_first_passport') as discover:
            with self.assertRaises(Stop):discover_first_passport_for_config(cfg,p)
            discover.assert_not_called()

    def test_user_shipment_date_is_validated_and_retained(self):
        from production_app.gui import App
        f=SimpleNamespace(facts={'marks':[]},new=SimpleNamespace(get=lambda:True),shipment_date=SimpleNamespace(get=lambda:'28.10.2026'),cfg={},project=SimpleNamespace(get=lambda:'P'))
        self.assertEqual(App.shipment_facts(f)['shipment_date'],'2026-10-28')
        f.shipment_date.get=lambda:'31.02.2026'
        with self.assertRaisesRegex(Stop,'ДД.ММ.ГГГГ'):App.shipment_facts(f)
    def test_existing_passport_ignores_proposed_new_date(self):
        from production_app.gui import App
        f=SimpleNamespace(facts={'marks':[]},new=SimpleNamespace(get=lambda:False),shipment_date=SimpleNamespace(get=lambda:'INVALID'),cfg={},project=SimpleNamespace(get=lambda:'P'))
        self.assertNotIn('shipment_date',App.shipment_facts(f))

class Recovery(unittest.TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name);self.stage=self.root/'.vzmk-operations'/'operation';self.stage.mkdir(parents=True);self.journal=self.stage/'journal.json';self.items=[]
        for role,suffix in [('registry','.xlsx'),('catalog','.xlsx'),('uzk','.xlsm')]:
            p=self.root/(role+suffix);backup=self.stage/('backup-'+role+suffix);candidate=self.stage/('candidate-'+role+suffix)
            p.write_bytes(('OLD '+role).encode());backup.write_bytes(p.read_bytes());candidate.write_bytes(('NEW '+role).encode())
            self.items.append({'role':role,'path':str(p),'backup':str(backup),'candidate':str(candidate),'original_sha256':sha(p),'candidate_sha256':sha(candidate)})
        self.state={'id':'operation','phase':'REPLACE_INTENT_CATALOG','registry':self.items[0]['path'],'target':str(self.root/'documents'),'participants':self.items,'new':True}
        core.journal_write(self.journal,self.state)
    def tearDown(self):self.tmp.cleanup()
    def change(self,n):
        for item in self.items[:n]:shutil.copyfile(item['candidate'],item['path'])
    def test_crash_after_first_replacement_restores_all(self):self.check_restore(1)
    def test_crash_after_second_replacement_restores_all(self):self.check_restore(2)
    def test_crash_after_third_replacement_restores_all(self):self.check_restore(3)
    def check_restore(self,n):
        self.change(n);r=core.recover(self.journal);self.assertEqual(r['phase'],'ROLLED_BACK');self.assertTrue(all(sha(i['path'])==i['original_sha256'] for i in self.items));self.assertEqual(core.recover(self.journal)['phase'],'ROLLED_BACK')
    def test_external_change_blocks_every_restore(self):
        self.change(2);Path(self.items[2]['path']).write_bytes(b'EXTERNAL CHANGE');before={i['path']:sha(i['path']) for i in self.items}
        with self.assertRaisesRegex(Stop,'Внешнее изменение'):core.recover(self.journal)
        self.assertEqual(before,{i['path']:sha(i['path']) for i in self.items})
    def test_completed_recovery_never_decrements_again(self):
        self.change(3);self.state['phase']='COMPLETE';core.journal_write(self.journal,self.state);core.recover(self.journal);self.assertTrue(all(sha(i['path'])==i['candidate_sha256'] for i in self.items))
    def test_damaged_backup_blocks_all_before_write(self):
        self.change(3);Path(self.items[2]['backup']).write_bytes(b'BAD');before={i['path']:sha(i['path']) for i in self.items}
        with self.assertRaisesRegex(Stop,'повреждена'):core.recover(self.journal)
        self.assertEqual(before,{i['path']:sha(i['path']) for i in self.items})

class NativeStandardEditor(unittest.TestCase):
    def editor(self,posted):
        from production_app.gui import App
        class Editor:
            def __init__(self):
                self.exists=True;self.bindings={};self.idle=[];self.value=''
                self.tk=SimpleNamespace(call=lambda *args:'.popup' if args[0]=='ttk::combobox::PopdownWindow' else posted,getboolean=bool)
            def set(self,value):self.value=value
            def get(self):return self.value
            def place(self,**kw):pass
            def focus_set(self):pass
            def bind(self,event,callback):self.bindings[event]=callback
            def winfo_exists(self):return self.exists
            def destroy(self):self.exists=False
            def after_idle(self,callback):self.idle.append(callback)
        values=['M','20Б1','ГОСТ 26020-83'];bindings={};entry=Editor()
        def item(key,option=None,**kw):
            if kw:values[:]=kw['values']
            return values
        tree=SimpleNamespace(identify_row=lambda y:'row',identify_column=lambda x:'#3',bbox=lambda *a:(0,0,10,10),item=item,bind=lambda name,cb:bindings.update({name:cb}))
        with patch('production_app.gui.ttk.Combobox',return_value=entry):
            App.editable(SimpleNamespace(),tree,immutable={0,1},choices={2:lambda row:['ГОСТ 26020-83','ГОСТ Р 57837-2017']})
            bindings['<Double-1>'](SimpleNamespace(x=1,y=1))
        return entry,values,tree
    def test_native_popup_focus_does_not_destroy_standard_choice(self):
        entry,values,tree=self.editor(True)
        entry.bindings['<FocusOut>']();self.assertTrue(entry.exists)
        for callback in entry.idle:callback()
        self.assertTrue(entry.exists)
        entry.value='ГОСТ Р 57837-2017';entry.bindings['<<ComboboxSelected>>']()
        self.assertEqual(values[2],'ГОСТ Р 57837-2017');self.assertFalse(entry.exists)
        self.assertIsNone(tree._pending_editor_commit)
    def test_focus_leaving_closed_popup_commits_last_choice(self):
        entry,values,tree=self.editor(False);entry.value='ГОСТ Р 57837-2017'
        entry.bindings['<FocusOut>']()
        for callback in entry.idle:callback()
        self.assertEqual(values[2],'ГОСТ Р 57837-2017');self.assertFalse(entry.exists)

if __name__=='__main__':unittest.main(verbosity=2)
