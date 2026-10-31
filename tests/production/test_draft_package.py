"""Incomplete drafts must never authorize actual findings or registry writes."""
import sys
import unittest
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock,patch
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0,str(Path(__file__).parent))
from test_quality_from_first import build_first
from test_highlighted_reference import uzk_first
from production_app import core,gui,draft_package as draft
from production_app.facts import validate
from production_app.storage import Stop,save,sha
from production_app.print_pdf import ordered_sources


class DraftBoundary(unittest.TestCase):
    def test_full_shipment_summary_keeps_accepted_mass_and_area_precision(self):
        v=SimpleNamespace(rows=[None]*104,quantity=Decimal(288),
            mass=Decimal('15822.13999999999951'),area=Decimal('1021.6439999999999998'))
        result=draft.shipment_summary(v,{'M':{'uzk_quantity':33}})
        self.assertEqual(result['mass_kg'],'15822.14')
        self.assertEqual(result['area_m2'],'1021.644')
        self.assertEqual(result['uzk_quantity'],33)

    def test_no_verdict_or_approval_inherited_from_saved_screen(self):
        rows=draft.collect_marks(None,{'M':{'uzk_quantity':8}},
            {'M':{'profile':'P','akz':None}},
            {'confirmed':True,'marks':[{'mark':'M','uzk_result':'ГОДЕН',
                'vik_result':'ГОДЕН','basis':'OLD','akz':None}]})
        self.assertEqual(rows[0]['uzk_quantity'],8)
        self.assertIsNone(rows[0]['akz'])
        self.assertNotIn('uzk_result',rows[0]);self.assertNotIn('basis',rows[0])

    def test_current_shared_count_overrides_old_saved_quantity(self):
        rows=draft.collect_marks(None,{'M':{'uzk_quantity':8}},
             {'M':{'profile':'P','akz':0}},
             {'marks':[{'mark':'M','uzk_quantity':1000,'akz':'0'}]})
        self.assertEqual(rows[0]['uzk_quantity'],8)
        self.assertEqual(rows[0]['akz'],0)

    def test_different_vom_saved_marks_rejected(self):
        with self.assertRaisesRegex(Stop,'другой ВОМ'):
            draft.collect_marks(None,{'M':{'uzk_quantity':8}},{},
                                {'marks':[{'mark':'OTHER'}]})

    def test_unknown_akz_never_promoted_to_zero(self):
        rows=draft.collect_marks(None,{'M':{'uzk_quantity':0}},{'M':{}},{})
        self.assertIsNone(rows[0]['akz'])
        climates,notes=draft.draft_climates({},rows)
        self.assertEqual(climates,{})

    def test_partial_climate_does_not_fabricate_midrange(self):
        climates,notes=draft.draft_climates(
            {'akz_drafts':{'0':{'coats':1,'air':{'min':'23','max':''}}}},
            [{'mark':'M','akz':0}])
        self.assertEqual(climates,{})
        self.assertTrue(notes)
        self.assertIn('сохранены значения первого акта',notes[0])

    def test_actual_release_gate_still_rejects_missing_facts(self):
        facts={'project':'P','number':'1/2026','kind':'ACTUAL',
               'simple_ui':True,'no_responsible_input':True,
               'marks':[{'mark':'M','profile':'P','akz':0,'uzk_quantity':1}]}
        v=SimpleNamespace(rows=[('M','Part',1,1,1,1,1)])
        with self.assertRaisesRegex(Stop,'результата ВИК'):
            validate(facts,v,'P','1/2026',defer_registry_dates=True,defer_approvals=True)

    def test_draft_manifest_rejected_by_all_plan_stages(self):
        with TemporaryDirectory() as td:
            p=Path(td)/'DRAFT.json'
            save(p,{'schema':'VZMK_DRAFT_PACKAGE_1','official_issue':False})
            before=sha(p)
            for operation in (core.prepare,core.verify,core.release,
                              lambda path:core.confirm(path,'',True,True),
                              lambda path:core.commit(path,'none')):
                with self.assertRaisesRegex(Stop,'не является планом'):
                    operation(p)
            self.assertEqual(sha(p),before)
            self.assertEqual(list(Path(td).iterdir()),[p])

    def test_one_button_enters_protected_workflow(self):
        import inspect
        code=inspect.getsource(gui.App.check)
        self.assertIn('engineering_facts',code)
        self.assertIn("self.run('check'",code)
        self.assertIn('self.preliminary_ready',code)
        self.assertNotIn('self.issue()',code)

    def test_declining_draft_view_cannot_issue(self):
        app=SimpleNamespace(project=Mock(get=lambda:'P'),write=Mock(),
            progress=Mock(),issue=Mock(),run=Mock(),error=Mock())
        result={'preview_pdf':'preview.pdf','folder':'folder','files':6,'pages':8,
                'summary':{'marks':4,'quantity':11,'mass_kg':511.98,
                           'area_m2':19.533,'uzk_quantity':8},'notes':[]}
        with patch.object(gui.messagebox,'askyesno',return_value=False),\
             patch.object(gui.os,'startfile') as opened:
            gui.App.drafted(app,result)
        opened.assert_not_called();app.issue.assert_not_called();app.run.assert_not_called()
        self.assertIsNone(app.plan)

    def test_accepting_draft_view_only_opens_pdf(self):
        app=SimpleNamespace(project=Mock(get=lambda:'P'),write=Mock(),
            progress=Mock(),issue=Mock(),run=Mock(),error=Mock())
        result={'preview_pdf':'preview.pdf','folder':'folder','files':6,'pages':8,
                'summary':{'marks':4,'quantity':11,'mass_kg':511.98,
                           'area_m2':19.533,'uzk_quantity':8},'notes':[]}
        with patch.object(gui.messagebox,'askyesno',return_value=True),\
             patch.object(gui.os,'startfile') as opened:
            gui.App.drafted(app,result)
        opened.assert_called_once_with('preview.pdf')
        app.issue.assert_not_called();app.run.assert_not_called()

    def test_partial_pdf_does_not_relax_release_completeness(self):
        with TemporaryDirectory() as td:
            path=Path(td)/'appendix.xlsx';path.write_bytes(b'fixture')
            files=[{'role':'appendix','path':str(path)}]
            self.assertEqual(ordered_sources(files,partial_draft=True)[0]['role'],'appendix')
            with self.assertRaisesRegex(Stop,'отсутствуют обязательные'):
                ordered_sources(files)

    def test_quality_missing_current_dates_preserves_sample_dates_and_static_fields(self):
        with TemporaryDirectory() as td:
            src=Path(td)/'first.docx';dest=Path(td)/'draft.docx';build_first(src)
            with ZipFile(src,'a') as z:z.writestr('customXml/keep.xml',b'<keep/>')
            before=sha(src)
            draft._quality(src,dest,'11/2026',dict.fromkeys(
                ['date','vik_date','uzk_date','akz_date']),SimpleNamespace(mass=511.98))
            doc=Document(dest);text=doc.element.xml
            self.assertIn('511,98',text)
            for day in ('02.09.2026','04.09.2026','08.09.2026','09.09.2026'):
                self.assertIn(day,text)
            self.assertNotIn('__________',text)
            self.assertEqual(doc.tables[0].cell(49,0).text,Document(src).tables[0].cell(49,0).text)
            with ZipFile(dest) as z:self.assertEqual(z.read('customXml/keep.xml'),b'<keep/>')
            self.assertEqual(sha(src),before)

    def test_uzk_draft_changes_dynamic_columns_preserves_reference_verdict(self):
        with TemporaryDirectory() as td:
            src=Path(td)/'first.docx';dest=Path(td)/'draft.docx';uzk_first(src)
            before=sha(src)
            design=[{'mark':'Бп2-1','thickness_mm':'5','length_mm':'1280',
                     'frequency_mhz':'5','angle_deg':'70'}]
            draft._word_copy(src,dest,'uzk','11/2026',{'uzk_date':None},None,design)
            doc=Document(dest);row=doc.tables[0].rows[5]
            self.assertEqual(row.cells[0].text,'Бп2-1')
            self.assertIn('1280',row.cells[1].text)
            original=Document(src).tables[0].rows[5]
            for column in (3,7,8):self.assertEqual(row.cells[column].text,original.cells[column].text)
            self.assertIn('04.09.2026',doc.element.xml)
            self.assertNotIn('__________',doc.element.xml)
            self.assertEqual(sha(src),before)

    def test_quality_available_registry_dates_replace_only_variable_fields(self):
        from datetime import date
        with TemporaryDirectory() as td:
            src=Path(td)/'first.docx';dest=Path(td)/'draft.docx';build_first(src)
            draft._quality(src,dest,'11/2026',{k:date(2026,10,i) for k,i in
                [('date',27),('vik_date',20),('uzk_date',21),('akz_date',22)]},SimpleNamespace(mass=511.98))
            d=Document(dest)
            self.assertIn('27.10.2026',d.tables[1].cell(4,0).text)
            text=d.tables[0].cell(50,0).text
            for day in ('20.10.2026','21.10.2026','22.10.2026'):self.assertIn(day,text)
            self.assertEqual(d.tables[0].cell(30,0).text,Document(src).tables[0].cell(30,0).text)

    def test_partial_complete_air_interval_is_transferred_without_fabricating_other_values(self):
        climates,notes=draft.draft_climates({'akz_drafts':{'0':{'coats':1,
            'air':{'min':'24','max':'28'}}}},[{'mark':'M','akz':0}])
        self.assertEqual(climates,{0:{'air':'26'}})
        self.assertTrue(notes)

    def test_invalid_entered_interval_is_not_silently_replaced_with_sample(self):
        with self.assertRaisesRegex(Stop,'Начало диапазона'):
            draft.draft_climates({'akz_drafts':{'0':{'air':{'min':'28','max':'24'}}}},[{'mark':'M','akz':0}])

    def test_akz_only_supplied_variable_cells_change_and_all_formulas_survive(self):
        class Cell:
            def __init__(self,value):self.Value2=value;self.Formula='=E28/0.6'
        cells={key:Cell(value) for key,value in {'E4':'1/2026 от ','H4':46273,
            'E28':92,'E31':22.2,'E32':23.1,'E33':53,'E34':17,'E27':153.3,
            'G11':'Подготовка','G37':'Метод','B43':'Подтверждают качество'}.items()}
        sheet=SimpleNamespace(Range=lambda name:cells[name])
        before={k:(v.Value2,v.Formula) for k,v in cells.items()}
        draft._akz_draft(sheet,'11/2026',None,{'air':'26'})
        for key,cell in cells.items():
            self.assertEqual(cell.Formula,before[key][1])
            self.assertEqual(cell.Value2,'11/2026 от ' if key=='E4' else 26 if key=='E32' else before[key][0])

    def test_draft_stamp_preserves_body_and_opaque_parts(self):
        from copy import deepcopy
        from lxml import etree
        with TemporaryDirectory() as td:
            p=Path(td)/'d.docx';uzk_first(p)
            with ZipFile(p,'a') as z:z.writestr('customXml/reference.xml',b'<untouched/>')
            doc=Document(p);body=deepcopy(doc.element.body)
            draft._stamp_word_draft(p)
            result=Document(p)
            stamped_body=deepcopy(result.element.body)
            for e in list(stamped_body.iter(qn('w:headerReference'))):e.getparent().remove(e)
            self.assertEqual(etree.tostring(body,method='c14n',exclusive=True),
                             etree.tostring(stamped_body,method='c14n',exclusive=True))
            self.assertEqual(draft.BANNER,result.sections[0].header.paragraphs[-1].text)
            with ZipFile(p) as z:self.assertEqual(z.read('customXml/reference.xml'),b'<untouched/>')


if __name__=='__main__':unittest.main(verbosity=2)
