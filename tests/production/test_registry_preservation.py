"""Registry date, table boundary and print-area contracts on disposable XLSX."""
import sys,tempfile,unittest,shutil
from pathlib import Path
from datetime import datetime,date
from decimal import Decimal
from types import SimpleNamespace
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT))
from production_app import core,model
from production_app.storage import Stop
from openpyxl import Workbook,load_workbook
from openpyxl.worksheet.table import Table,TableStyleInfo

class Preservation(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.original=self.root/'original.xlsx';self.after=self.root/'after.xlsx'
        b=Workbook();s=b.active;s.title='Реестр'
        headers=list(model.HEADERS.values())+['Примечание'];s.append(headers)
        fields={'NUMBER':'1/2026','CONTRACT':'Договор','CUSTOMER':'Заказчик','OBJECT':'Объект','CIPHER':'Шифр','MASS':10,
                'START_DATE':datetime(2026,10,1),'END_DATE':datetime(2026,10,18),'KM':'КМ','KMD_COMPANY':'КМД','STEEL':'С245',
                'VIK_DATE':datetime(2026,10,11),'UZK_DATE':datetime(2026,10,13),'AKZ_DATE':datetime(2026,10,17),'CONNECTIONS':'Балка'}
        s.append([fields[k] for k in model.HEADERS]+['Проект'])
        t=Table(displayName='Registry',ref='A1:P2');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2');s.add_table(t)
        s.print_area='A1:P2';b.save(self.original)
        s.append([fields[k] for k in model.HEADERS]+['Проект']);s.cell(3,1,'2/2026');s.tables['Registry'].ref='A1:P3';s.print_area='A1:P3';b.save(self.after);b.close()
        self.v=SimpleNamespace(mass=Decimal(10),connections='Балка')
    def tearDown(self):self.temp.cleanup()
    def verify(self):return core.verify_registry(self.original,self.after,'Проект','2/2026',self.v,date(2026,10,18))
    def change(self,fn):
        b=load_workbook(self.after);fn(b.active);b.save(self.after);b.close()
    def test_exact_single_row_extension_passes(self):self.assertEqual(self.verify(),3)
    def test_wrong_new_date_refused(self):
        self.change(lambda s:s.cell(3,8,datetime(2026,10,19)))
        with self.assertRaisesRegex(Stop,'Дата новой строки'):self.verify()
    def test_print_area_removed_refused(self):
        self.change(lambda s:setattr(s,'print_area',None))
        with self.assertRaisesRegex(Stop,'область печати'):self.verify()
    def test_extra_print_row_refused(self):
        self.change(lambda s:setattr(s,'print_area','A1:P4'))
        with self.assertRaisesRegex(Stop,'область печати'):self.verify()
    def test_wrong_table_boundary_refused(self):
        self.change(lambda s:setattr(s.tables['Registry'],'ref','A1:P4'))
        with self.assertRaisesRegex(Stop,'расширение'):self.verify()
    def test_existing_value_change_refused(self):
        self.change(lambda s:s.cell(2,2,'Чужой договор'))
        with self.assertRaisesRegex(Stop,'исходная ячейка'):self.verify()

if __name__=='__main__':unittest.main(verbosity=2)
