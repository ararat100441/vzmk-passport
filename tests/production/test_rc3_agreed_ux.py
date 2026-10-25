"""Synthetic regression for one-button RC3 UX, project data and AKZ ranges."""
import ast
import inspect
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from production_app import gui,core,model
from production_app.akz_ui import available_variants,display_variants,source_coats
from production_app.catalog_assignments import (
    read_assignments,current_assignments,used_akz_variants
)
from production_app.measurement_ranges import (
    measurements_from_ranges,verify_midrange_provenance,midrange
)
from production_app.registry_dates import resolve_registry_dates
from production_app.storage import Stop,sha


def stored_book(path,headers,rows):
    w=Workbook()
    s=w.active
    s.append(headers)
    for row in rows:s.append(row)
    w.save(path);w.close()


def ranged_input():
    return {
        'air':{'min':'19','max':'23'},
        'surface':{'min':'21','max':'27'},
        'humidity':{'min':'50','max':'60'},
        'dew':{'min':'8','max':'12'},
        'layer_intervals':[{'layer':1,'min':'95','max':'125'}],
    }


class AgreedUX(unittest.TestCase):
    def test_catalog_assignments_preserve_variant_zero_and_blanks(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'Справочник марок.xlsx'
            stored_book(path,['Марка','Профиль','Краска','№ АКЗ'],[
                ['Н1-26','Двутавр','ЛКМ-1',0],
                ['Н1-27','Швеллер','ЛКМ-1',None],
                ['Н1-33',None,None,None],
                ['Другая марка','Балка','ЛКМ-2',1],
            ])
            got=read_assignments(path,sha(path))
            self.assertEqual(got['Н1-26']['akz'],0)
            self.assertNotIn('akz',got['Н1-27'])
            self.assertEqual(got['Н1-33'],{})
            active=current_assignments(path,sha(path),
                                       ['Н1-26','Н1-27','НЕИЗВЕСТНАЯ'])
            self.assertEqual(active['НЕИЗВЕСТНАЯ'],{})
            self.assertEqual(used_akz_variants([active['Н1-26'],
                                                active['Н1-27']]),[0])
            with self.assertRaisesRegex(Stop,'Изменился справочник'):
                read_assignments(path,'BAD_HASH')

    def test_duplicate_catalogues_stop(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'duplicate.xlsx'
            stored_book(path,['Марка','АКЗ'],[['М1',0],['М1',1]])
            with self.assertRaisesRegex(Stop,'Дублирующаяся'):
                read_assignments(path)

    def test_only_active_variants_are_displayed(self):
        forms={'templates':{
            'akz0':{},'akz_appendix0':{},
            'akz1':{},'akz_appendix1':{}
        }}
        self.assertEqual(available_variants(forms),[0,1])
        self.assertEqual(display_variants(
            [{'akz':None},{'akz':0},{'akz':''}],forms),[0])
        self.assertEqual(display_variants([{'akz':None}],forms),[])
        only0={'templates':{'akz0':{},'akz_appendix0':{}}}
        self.assertEqual(display_variants([{'akz':None}],only0),[0])
        self.assertEqual(used_akz_variants([{'akz':0}]),[0])
        self.assertEqual(used_akz_variants([{'akz':None}]),[])

    def test_akz_coats_read_from_selected_approved_form(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'akz.xlsx'
            stored_book(p,['Название','Значение'],
                        [['Толщина','{{THICKNESS}}']])
            self.assertEqual(source_coats({'path':str(p),'sha256':sha(p)}),1)
            w=Workbook();w.active['A1']='{{LAYER_3_RANGE}}'
            w.save(p);w.close()
            self.assertEqual(source_coats({'path':str(p),'sha256':sha(p)}),3)

    def test_climate_and_one_coat_midrange_is_deterministic(self):
        c=measurements_from_ranges(0,ranged_input(),coats=1)
        self.assertEqual(c['air'],'21')
        self.assertEqual(c['surface'],'24')
        self.assertEqual(c['humidity'],'55')
        self.assertEqual(c['dew'],'10')
        self.assertEqual(c['thickness'],'110')
        self.assertNotIn('layers',c)
        self.assertTrue(verify_midrange_provenance(c))
        self.assertEqual(c['midrange_source'],'CURRENT_CONFIRMED_MEASURED_RANGE')

    def test_cannot_reuse_design_limit_or_modify_calculated_midrange(self):
        c=measurements_from_ranges(0,ranged_input(),coats=1)
        c['air']='22'
        with self.assertRaisesRegex(Stop,'изменена'):
            verify_midrange_provenance(c)
        c['air']='21'
        c['midrange_source']='PREVIOUS_PASSPORT'
        with self.assertRaisesRegex(Stop,'источника'):
            verify_midrange_provenance(c)

    def test_three_coat_midrange_is_cumulative_in_layer_order(self):
        data=ranged_input()
        data['layer_intervals']=[
            {'layer':3,'min':'80','max':'100'},
            {'layer':1,'min':'30','max':'50'},
            {'layer':2,'min':'40','max':'60'},
        ]
        c=measurements_from_ranges(2,data,coats=3)
        self.assertEqual(
            c['dry_averages'],[
                {'layer':1,'average':'40'},
                {'layer':2,'average':'90'},
                {'layer':3,'average':'180'},
            ])
        self.assertEqual(c['thickness'],'180')
        self.assertTrue(verify_midrange_provenance(c))

    def test_partial_range_and_wrong_two_coat_stop(self):
        data=ranged_input()
        data['air']['max']=''
        with self.assertRaises(Stop):
            measurements_from_ranges(0,data,coats=1)
        data=ranged_input()
        data['layer_intervals'].append(
            {'layer':2,'min':'80','max':'120'})
        with self.assertRaisesRegex(Stop,'Двухслойный'):
            measurements_from_ranges(0,data,coats=2)

    def test_single_package_button_still_runs_all_internal_stages(self):
        source=inspect.getsource(gui.App)
        self.assertIn("'Сформировать пакет'",source)
        for label in ('1. Проверить данные','2. Подготовить документы',
                      '3. Проверить комплект','4. Подтвердить выпуск и запись'):
            self.assertNotIn("text='"+label+"'",source)
        self.assertNotIn('self.buttons[3]',source)
        for task in ("self.run('check'","self.run('prepare'",
                     "self.run('verify'","self.run('confirm'",
                     "self.run('release'","self.run('commit'",
                     "self.run('readback'","self.run('complete'"):
            self.assertIn(task,source)

    def test_preview_decline_does_not_initiate_any_release(self):
        from unittest.mock import Mock
        app=object.__new__(gui.App)
        app.plan='/safe/temporary/job/plan.json'
        app.preview_pdf='/safe/temporary/job/preview.pdf'
        app.cfg={'training':False}
        app.write=Mock()
        app.issue=Mock()
        data={'project':'Гараж Казахстан','number':'11/2026',
              'date':'2026-10-25','registry':'COPY.xlsx','target':'TEMP',
              'new':True}
        with patch.object(gui,'read',return_value=data),\
             patch.object(gui.os,'startfile') as launch,\
             patch.object(gui.messagebox,'askyesno',return_value=False) as confirm:
            app.review_preview()
        launch.assert_called_once_with(app.preview_pdf)
        confirm.assert_called_once()
        app.issue.assert_not_called()
        app.write.assert_called_once()

    def test_preview_yes_has_one_confirmation_dialog(self):
        from unittest.mock import Mock
        app=object.__new__(gui.App)
        app.plan='/safe/temporary/job/plan.json'
        app.preview_pdf='/safe/temporary/job/preview.pdf'
        app.cfg={'training':False}
        app.write=Mock()
        app.issue=Mock()
        data={'project':'Гараж Казахстан','number':'11/2026',
              'date':'2026-10-25','registry':'COPY.xlsx','target':'TEMP',
              'new':True}
        with patch.object(gui,'read',return_value=data),\
             patch.object(gui.os,'startfile'),\
             patch.object(gui.messagebox,'askyesno',return_value=True) as confirm:
            app.review_preview()
        app.issue.assert_called_once()
        self.assertEqual(confirm.call_count,1)
        self.assertIn('11/2026',confirm.call_args.args[1])
        self.assertIn('COPY.xlsx',confirm.call_args.args[1])

    def test_preview_pdf_is_pinned_and_not_published_before_issue(self):
        from production_app.storage import save,pins,read
        with tempfile.TemporaryDirectory() as td:
            folder=Path(td)
            staging=folder/'documents'
            staging.mkdir()
            outputs=[]
            for role,ext in (('quality','.docx'),('appendix','.xlsx'),
                             ('vik','.docx'),('uzk','.docx')):
                p=staging/(role+ext)
                p.write_bytes(('Synthetic '+role).encode())
                outputs.append({'role':role,'path':str(p)})
            registry=folder/'snapshot.xlsx'
            registry.write_bytes(b'SYNTHETIC REGISTRY')
            d={'run':str(folder),'seal':'SEAL','number':'11/2026',
               'cfg':{'training':True},
               'target':str(folder/'Output'/'11_2026 от 25.10.2026')}
            prepared={
                'plan_seal':'SEAL','registry':str(registry),
                'snapshot_sha256':sha(registry),
                'output_fingerprints':pins([x['path'] for x in outputs]),
                'outputs':outputs,
            }
            save(folder/'prepared.json',prepared)

            def fake_pdf(sources,dest,name):
                preview=dest/(name+'.pdf')
                preview.write_bytes(b'%PDF-1.4\\nSYNTHETIC TEST ONLY')
                return {'output':{'role':'print_pdf','path':str(preview)},
                        'manifest':{'filename':preview.name}}
            with patch.object(core,'load_plan',return_value=d),\
                 patch.object(core.office,'verify_documents',return_value=[]),\
                 patch('production_app.print_pdf.create_print_pdf',
                       side_effect=fake_pdf):
                report=core.verify(folder/'plan.json',preview=True)
            self.assertIsNotNone(report['preview_pdf'])
            self.assertTrue(Path(report['preview_pdf']).is_file())
            self.assertFalse((folder/'Output').exists())
            self.assertEqual(registry.read_bytes(),b'SYNTHETIC REGISTRY')
            proof=read(folder/'verified.json')
            self.assertIn(report['preview_pdf'],proof['fingerprints'])
            self.assertIn('preview',report['preview_pdf'])

    def test_existing_registry_row_dates_are_read_only(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'registry.xlsx'
            w=Workbook();s=w.active
            for col,key in enumerate(('START_DATE','END_DATE','VIK_DATE',
                                     'UZK_DATE','AKZ_DATE'),start=1):
                s.cell(2,col).value=date(2026,10,19)
            w.save(p);w.close()
            registry=SimpleNamespace(
                path=p,digest=sha(p),sheet='Sheet',
                projects={'PROJECT':[{'row':2,'number':'11/2026'}]},
            )
            with patch('production_app.registry_dates.read_registry',
                       return_value=registry):
                with patch('production_app.registry_dates.model.select',
                           return_value=SimpleNamespace(next_number='12/2026')):
                    with patch('production_app.registry_dates.model.columns',
                               return_value=dict(zip(
                                    ('START_DATE','END_DATE','VIK_DATE','UZK_DATE','AKZ_DATE'),
                                    range(1,6)))):
                        dates=resolve_registry_dates(
                            {'registry':str(p)},'PROJECT','11/2026',None,False)
            self.assertEqual(dates['date'],'2026-10-19')
            self.assertEqual(dates['akz_date'],'2026-10-19')
            self.assertEqual(dates['registry_date_origin'],'EXACT_EXISTING_REGISTRY_ROW')

    def test_new_passport_must_not_reuse_old_nonformula_inspection_dates(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'registry.xlsx'
            w=Workbook();w.active['A2']=date(2026,9,9)
            w.save(p);w.close()
            reg=SimpleNamespace(path=p,digest=sha(p),sheet='Sheet')
            project=SimpleNamespace(first_row=2,next_number='11/2026')
            with patch('production_app.registry_dates.read_registry',
                       return_value=reg):
                with patch('production_app.registry_dates.model.select',
                           return_value=project):
                    with patch('production_app.registry_dates.model.columns',
                               return_value={
                                  'START_DATE':1,'VIK_DATE':2,
                                  'UZK_DATE':3,'AKZ_DATE':4}):
                        with self.assertRaisesRegex(Stop,'старую фиксированную'):
                            resolve_registry_dates({'registry':str(p)},
                                  'PROJECT','11/2026',None,True,
                                  day=date(2026,10,25))


if __name__=='__main__':
    unittest.main(verbosity=2)
