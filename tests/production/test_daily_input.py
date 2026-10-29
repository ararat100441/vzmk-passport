import tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from datetime import datetime,time
from openpyxl import load_workbook
from production_app import control_book as C
from production_app.storage import Stop,sha

class ControlBookTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.vom=SimpleNamespace(rows=[('M1','Long name',2,10,3)])
        self.facts={'responsible':'Operator','date':'2026-10-24','start_date':'2026-10-01',
          'vik_date':'2026-10-17','uzk_date':'2026-10-19','akz_date':'2026-10-23',
          'marks':[{'mark':'M1','profile':'P','paint':'Paint','akz':0,'uzk_quantity':'1','uzk_result':'ГОДЕН','basis':'Protocol','vik_result':'ГОДЕН','vik_basis':'VIK'}],
          'uzk_details':[dict(zip(C.UZK_COLUMNS,['M1','J1','Welder','2026-10-19','5','70','No defects','OK','Protocol'])),dict(zip(C.UZK_COLUMNS,['M1','J2','Welder','2026-10-19','5','70','No defects','OK','Protocol']))],
          'climates':[{'variant':0,'air':'24','surface':'27','humidity':'55','dew':'14','thickness':'','material':'Paint','basis':'Protocol',
          'layers':[{'layer':n,'min':'80','max':'130'} for n in (1,2,3)],
          'dry_averages':[{'layer':n,'average':str(n*100)} for n in (1,2,3)],
          'layer_details':[{'layer':n,'material':'Paint'+str(n),'paint_batch':'P'+str(n),'hardener_batch':'H'+str(n),'color':'Grey','application_date':'2026-10-23','time_start':'08:00','time_end':'10:30','wet_average':'110'} for n in (1,2,3)]}],
          'doc_fields':{'QUALITY_MANUFACTURER':'Factory','VIK_CONCLUSION':'Current verdict'}}
    def tearDown(self):self.tmp.cleanup()
    def make(self,name='input.xlsx'):
        p=self.root/name;C.write_book(p,'Project','19/1/2026',self.vom,facts=self.facts);return p
    def test_roundtrip_keeps_three_layers_and_repeated_joints_without_approval(self):
        p=self.make();h=sha(p);d=C.read_book(p,'Project','19/1/2026',True)
        self.assertEqual([r['joint'] for r in d['uzk_details']],['J1','J2'])
        self.assertEqual(d['climates'],self.facts['climates']);self.assertEqual(d['marks'],self.facts['marks'])
        self.assertFalse(d['confirmed']);self.assertFalse(d['authorized']);self.assertFalse(d['forms_approved']);self.assertEqual(sha(p),h)
    def test_other_project_refused_even_if_same_mark(self):
        with self.assertRaisesRegex(Stop,'другому проекту'):C.read_book(self.make(),'Other','19/1/2026')
    def test_stale_number_refused(self):
        with self.assertRaisesRegex(Stop,'актуальным реестром'):C.read_book(self.make(),'Project','20/1/2026')
    def test_excel_dates_and_time_are_not_retyped(self):
        p=self.make();b=load_workbook(p);b['Комплект']['B6']=datetime(2026,10,24);b['Слои АКЗ']['K2']=time(8,0);b.save(p);b.close()
        d=C.read_book(p,'Project','19/1/2026');self.assertEqual(d['date'],'2026-10-24');self.assertEqual(d['climates'][0]['layer_details'][0]['time_start'],'08:00')
    def test_uncalculated_formula_stops_instead_of_empty_fact(self):
        p=self.make();b=load_workbook(p);b['Контроль']['E2']='=1+1';b.save(p);b.close()
        with self.assertRaisesRegex(Stop,'Пересчитайте'):C.read_book(p,'Project','19/1/2026')
    def test_preferences_never_reuse_verdicts_or_measurements(self):
        saved=C.preferences(self.facts);self.assertEqual(saved['doc_fields'],{'QUALITY_MANUFACTURER':'Factory'})
        self.assertEqual(set(saved['assignments']['M1']),{'profile','paint','akz'})
        self.assertFalse(set(saved)&{'climates','uzk_details','evidence','confirmed','number'})
    def test_blank_book_does_not_assign_akz_or_results(self):
        p=self.root/'blank.xlsx';C.write_book(p,'Project','19/1/2026',self.vom)
        d=C.read_book(p,'Project','19/1/2026');self.assertIsNone(d['marks'][0]['akz']);self.assertEqual(d['marks'][0]['uzk_result'],'')
        self.assertEqual(d['climates'],[]);self.assertEqual(d['uzk_details'],[])
    def test_no_overwrite_of_input_file(self):
        p=self.make();h=sha(p)
        with self.assertRaises(Stop):C.write_book(p,'Project','19/1/2026',self.vom)
        self.assertEqual(sha(p),h)
    def test_normal_uzk_excel_keeps_repeated_marks_and_native_dates(self):
        from openpyxl import Workbook
        p=self.root/'ledger.xlsx';w=Workbook();s=w.active;s.append(list(C.UZK_HEADERS))
        for joint in ('J1','J2'):s.append(['M1',joint,'Welder',datetime(2026,10,19),'2.5;5','45;60','No defects','OK','Protocol'])
        w.save(p);w.close();h=sha(p)
        rows,proof=C.read_uzk_book(p);self.assertEqual([r['joint'] for r in rows],['J1','J2']);self.assertEqual(rows[0]['date'],'2026-10-19');self.assertEqual(proof['sha256'],h);self.assertEqual(sha(p),h)
        self.assertNotIn('uzk_quantity',rows[0])
    def test_uzk_ledger_without_actual_results_is_refused(self):
        from openpyxl import Workbook
        p=self.root/'assignments.xlsx';w=Workbook();s=w.active;s.append(['Марка','Профиль']);s.append(['M1','Profile']);w.save(p);w.close()
        with self.assertRaisesRegex(Stop,'не найдена подробная'):C.read_uzk_book(p)
    def test_repeated_vom_mark_does_not_duplicate_control_record_or_invent_quantity(self):
        p=self.root/'repeated.xlsx';v=SimpleNamespace(rows=[('M1','A',2),('M1','A',3)])
        C.write_book(p,'Project','19/1/2026',v)
        data=C.read_book(p,'Project','19/1/2026');self.assertEqual(len(data['marks']),1);self.assertEqual(data['marks'][0]['uzk_quantity'],'')
    def test_unused_akz_variant_is_not_required_for_current_marks(self):
        p=self.make();b=load_workbook(p);b['Климат'].append([3,'20','25','50','10','100','Other material','Unused protocol']);b.save(p);b.close()
        data=C.read_book(p,'Project','19/1/2026');self.assertEqual([c['variant'] for c in data['climates']],[0])
    def test_russian_dates_and_short_hour_need_no_retyping(self):
        p=self.make();b=load_workbook(p);b['Комплект']['B6']='24.10.2026';b['Слои АКЗ']['K2']='8:00';b.save(p);b.close()
        d=C.read_book(p,'Project','19/1/2026');self.assertEqual(d['date'],'2026-10-24');self.assertEqual(d['climates'][0]['layer_details'][0]['time_start'],'08:00')
    def test_impossible_date_has_actionable_message(self):
        with self.assertRaisesRegex(Stop,'Пример: 24.10.2026'):C.text('31.02.2026','date')
    def test_explicit_empty_requisite_does_not_restore_stale_cached_value(self):
        p=self.make();b=load_workbook(p)
        for row in b['Реквизиты']:
            if row[0].value==C.DOC_FIELD_LABELS['QUALITY_MANUFACTURER']:row[1].value=None
        b.save(p);b.close()
        d=C.read_book(p,'Project','19/1/2026',defaults={'doc_fields':{'QUALITY_MANUFACTURER':'Old factory'}})
        self.assertNotIn('QUALITY_MANUFACTURER',d['doc_fields'])
    def test_project_assignments_survive_shipment_subset_without_control_results(self):
        previous={'assignments':{'OLD':{'profile':'Old profile','paint':'Paint','akz':2}}}
        result=C.preferences(self.facts,previous);self.assertIn('OLD',result['assignments']);self.assertIn('M1',result['assignments'])
        result['assignments']['OLD']['akz']=3;self.assertEqual(previous['assignments']['OLD']['akz'],2)
        self.assertNotIn('climates',result);self.assertNotIn('uzk_quantity',result['assignments']['M1'])

if __name__=='__main__':unittest.main()
