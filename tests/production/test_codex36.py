"""Synthetic regressions for r9 PDF marking and historical UZK layout."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from io import BytesIO

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader, PdfWriter
from pypdf.generic import RectangleObject
from reportlab.pdfgen import canvas

from production_app.highlighted_reference import create_uzk_copy
from production_app.print_pdf import create_print_pdf, _preview_overlay
from production_app.office import PREVIEW_MARK
from production_app.storage import sha
from test_highlighted_reference import uzk_first


class UzkLayout(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.source=self.root/'source.docx'
        uzk_first(self.source)
        d=Document(self.source)
        p=d.tables[0].rows[4].cells[5].paragraphs[0]
        p.add_run('Рабочая частота, ')
        p.add_run('м').font.size=Pt(10)
        p.add_run('Гц').bold=True
        t=d.tables[1]
        t.add_row().cells[0].text='Уровень квалификации II, TEST удостоверение'
        d.add_paragraph('')
        d.save(self.source)
        self.before=sha(self.source)

    def tearDown(self):self.temp.cleanup()

    def make(self,count=4):
        target=self.root/'output.docx'
        design=[dict(mark='TEST-'+str(n),length_mm='400',thickness_mm='8',
                     frequency_mhz='5',angle_deg='70') for n in range(count)]
        create_uzk_copy(self.source,target,number='11/2026',
                        uzk_date='2026-10-26',design_rows=design)
        self.assertEqual(sha(self.source),self.before)
        return Document(target)

    def test_frequency_unit_across_runs_keeps_formatting(self):
        d=self.make();p=d.tables[0].rows[4].cells[5].paragraphs[0]
        self.assertEqual(p.text,'Рабочая частота, МГц')
        self.assertEqual(p.runs[-2].font.size,Pt(10))
        self.assertTrue(p.runs[-1].bold)

    def test_only_last_mark_is_tethered_to_executor(self):
        d=self.make(40)
        self.assertEqual(len(d.tables[0].rows),45)
        self.assertTrue(all(p.paragraph_format.keep_with_next
                        for c in d.tables[0].rows[-1].cells for p in c.paragraphs))
        self.assertTrue(all(not p.paragraph_format.keep_with_next
                        for c in d.tables[0].rows[-2].cells for p in c.paragraphs))
        executor=d.tables[1]
        self.assertEqual(executor.cell(0,0).text,'Контроль выполнил')
        self.assertEqual(executor.cell(0,2).text,
                         Document(self.source).tables[1].cell(0,2).text)
        self.assertIn('Уровень квалификации II',executor.cell(1,0).text)
        self.assertTrue(all(p.paragraph_format.keep_with_next
                        for row in executor.rows for c in row.cells for p in c.paragraphs))

    def test_margins_widths_and_nonempty_type_sizes_unchanged(self):
        d=self.make();s=Document(self.source)
        self.assertEqual(d.sections[0]._sectPr.xml,s.sections[0]._sectPr.xml)
        self.assertEqual([c.width for c in d.tables[0].columns],
                         [c.width for c in s.tables[0].columns])
        cell=d.tables[0].rows[5].cells[3]
        original=s.tables[0].rows[5].cells[3]
        self.assertEqual(cell.text,original.text)
        self.assertEqual([r._r.xml for p in cell.paragraphs for r in p.runs],
                         [r._r.xml for p in original.paragraphs for r in p.runs])

    def test_akz_reference_total_is_not_replaced_by_measured_thickness(self):
        from test_akz_from_first import Sheet,AKZ_FORMULAS
        from production_app.akz_from_first import fill_first_akz
        s=Sheet();s.data.update(E4='1/2026 от ',E25=80,E26='нет')
        for pos,formula in AKZ_FORMULAS.items():s.Range(pos).Formula=formula
        fill_first_akz(s,number='11/2026',akz_date='2026-10-30',
            climate=dict(variant=0,wet_thickness=130,thickness=85,
                         surface=21,air=26,humidity=50,dew=14.8))
        self.assertEqual(s.Range('E25').Value2,80)
        self.assertEqual(s.Range('E26').Value2,'нет')
        self.assertEqual(s.Range('E27').Value2,130)
        self.assertEqual(s.Range('E28').Value2,85)


class PreviewPDF(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)

    def tearDown(self):self.temp.cleanup()

    @staticmethod
    def export(items):
        for item in items:
            c=canvas.Canvas(str(item['pdf']),pagesize=(595,842))
            for n in range(3):
                if n==1:c.setPageSize((842,595))
                c.drawString(45,100,'CONTENT '+item['role']+' '+str(n))
                c.showPage()
            c.save()

    def create(self,preview):
        folder=self.root/str(preview);folder.mkdir()
        docs=[]
        for role in ('quality','appendix','vik','uzk'):
            p=folder/(role+('.xlsx' if role=='appendix' else '.docx'))
            p.write_bytes(b'SYNTHETIC SOURCE');docs.append(dict(role=role,path=str(p)))
        before={x['path']:sha(x['path']) for x in docs}
        with patch('production_app.print_pdf._export_word',side_effect=self.export), \
             patch('production_app.print_pdf._export_excel',side_effect=self.export):
            result=create_print_pdf(docs,folder,'11_2026',preview=preview)
        self.assertEqual(before,{x['path']:sha(x['path']) for x in docs})
        return PdfReader(result['output']['path'])

    def test_preview_every_page_including_continuations_and_landscape(self):
        r=self.create(True)
        self.assertEqual(len(r.pages),12)
        for p in r.pages:
            text=p.extract_text()
            self.assertEqual(text.count(PREVIEW_MARK),1)
            self.assertIn('CONTENT',text)

    def test_final_has_no_mark_on_any_page(self):
        r=self.create(False)
        for p in r.pages:
            self.assertNotIn(PREVIEW_MARK,p.extract_text())
            self.assertIn('CONTENT',p.extract_text())

    def test_rotated_page_and_offset_crop_are_marked(self):
        raw=BytesIO();c=canvas.Canvas(raw,pagesize=(700,900))
        c.drawString(100,100,'KEEP CONTENT');c.save();raw.seek(0)
        for angle in (0,90,180,270):
            raw.seek(0);p=PdfReader(raw).pages[0]
            p.cropbox=RectangleObject((20,30,680,870));p.rotate(angle)
            _preview_overlay(p)
            w=PdfWriter();w.add_page(p);out=BytesIO();w.write(out);out.seek(0)
            text=PdfReader(out).pages[0].extract_text()
            self.assertIn('KEEP CONTENT',text)
            self.assertEqual(text.count(PREVIEW_MARK),1)


if __name__=='__main__':unittest.main()
