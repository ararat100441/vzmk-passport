"""Current-VOM B gates; unrelated B warnings never relax global F+ checks."""
import test_codex38 as fixtures
from production_app import catalog_ledger as ledger
from production_app.storage import Stop,sha
from openpyxl import load_workbook
from unittest import TestCase


class ShipmentScope(TestCase):
    def setUp(self):
        self.fixture=fixtures.EditableLedger();self.fixture.setUp();self.addCleanup(self.fixture.tearDown)
        self.f=self.fixture
    def test_invalid_unrequested_b_warns_preserves_input_and_candidate_b_history(self):
        for value in (0,1.5,-2,None,'bad','=1+1',True):
            with self.subTest(value=value):
                self.f.edit(self.f.catalog,{'B3':value})
                before=sha(self.f.catalog)
                result=self.f.check();self.assertEqual(len(result['stock_warnings']),1)
                p=self.f.candidate('candidate'+str(len(list(self.f.root.glob('candidate*'))))+'.xlsx')
                b=load_workbook(p);s=b.active
                self.assertEqual(s['B3'].value,value);self.assertEqual(s['F3'].value,1);self.assertEqual(s['G3'].value,0)
                if isinstance(value,(float,int)) and not isinstance(value,bool):self.assertEqual(s['E3'].value,'=B3-SUM(F3:G3)')
                else:self.assertIn('НЕИЗВЕСТНО',s['E3'].value)
                self.assertEqual(s['G2'].value,1);b.close()
                self.assertEqual(sha(self.f.catalog),before)
    def test_same_invalid_b_in_requested_mark_is_blocking(self):
        for value in (2,1.5,-2,None,'bad','=1+1',True):
            with self.subTest(value=value):
                self.f.edit(self.f.catalog,{'B2':value})
                with self.assertRaises(Stop):self.f.check()
    def test_history_mismatch_in_other_mark_still_blocks_with_invalid_b(self):
        for value in (0,1.5,None,'bad'):
            with self.subTest(value=value):
                self.f.edit(self.f.catalog,{'B3':value,'F3':2})
                with self.assertRaisesRegex(Stop,'Количество истории'):self.f.check()
    def test_history_invalid_quantity_in_other_mark_still_blocks(self):
        self.f.edit(self.f.catalog,{'B3':0,'F3':'bad'})
        with self.assertRaisesRegex(Stop,'История'):self.f.check()
    def test_unrequested_history_change_in_appendix_still_blocks(self):
        self.f.edit(self.f.catalog,{'B3':0});self.f.edit(self.f.app,{'C6':2})
        with self.assertRaisesRegex(Stop,'N: было 1, стало 2'):self.f.check()
    def test_warning_is_not_persisted_as_stale_metadata(self):
        self.f.edit(self.f.catalog,{'B3':0});p=self.f.candidate()
        b=load_workbook(p);self.assertNotIn('stock_warnings',ledger.metadata(b));b.close()
