"""Synthetic 2-page quality first-passport contract.

Only six highlighted locations may change. Current act numbers and dates
are provided by the registry; the original form remains unchanged.
"""
import unittest
from zipfile import ZipFile
from tempfile import TemporaryDirectory
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from production_app.quality_from_first import (
    create_quality_copy,is_exact_first_quality
)
from production_app.storage import sha,Stop


def yellow(cell,text,*,paragraph=0):
    p=cell.paragraphs[paragraph]
    r=p.add_run(text)
    r.font.highlight_color=WD_COLOR_INDEX.YELLOW


def build_first(path):
    doc=Document()
    t=doc.add_table(rows=54,cols=10)
    yellow(t.cell(1,9),'1/2026',paragraph=0)
    yellow(t.cell(16,6),'19 407,41')
    t.cell(16,6).paragraphs[0].add_run(' кг.')
    yellow(t.cell(18,8),'09.09.2026')
    t.cell(18,8).paragraphs[0].add_run(' г.')
    for row in (50,53):
        p=t.cell(row,0).paragraphs[0]
        for text in (
            'Акт ВИК № ','1/2026',' от ','02.09.2026',
            ' г., Акт ВИК № ','1/2026','-1 от ','02.09.2026',
            ' г., Акт УЗК № ','1/2026',' от ','04.09.2026',
            ' г., Сертификат качества антикоррозийной обработки № ',
            '1/2026',' от ','08.09.2026',' г.',
        ):
            r=p.add_run(text);r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    t.cell(30,0).text='Дата начала изготовления: 05.08.2026'
    t.cell(49,0).text='Сварщики: ОБРАЗЕЦ_ФИО_02'
    signing=doc.add_table(rows=5,cols=1)
    yellow(signing.cell(4,0),'09.09.2026')
    signing.cell(4,0).paragraphs[0].add_run(' г.')
    doc.save(path)


class ExactQuality(unittest.TestCase):
    def setUp(self):
        self.dir=TemporaryDirectory()
        self.path=Path(self.dir.name)
        self.first=self.path/'first_quality.docx'
        build_first(self.first)
    def tearDown(self):
        self.dir.cleanup()

    def test_opaque_package_parts_remain_byte_identical(self):
        with ZipFile(self.first,'a') as z:
            z.writestr('customXml/preserved-source.xml',b'<source>untouched</source>')
        dest=self.path/'quality.docx'
        create_quality_copy(self.first,dest,number='11/2026',end_date='2026-10-26',vik_date='2026-10-19',uzk_date='2026-10-21',akz_date='2026-10-25',mass_kg='511.98')
        with ZipFile(self.first) as before,ZipFile(dest) as after:
            self.assertEqual(before.namelist(),after.namelist())
            for name in before.namelist():
                if name!='word/document.xml':self.assertEqual(before.read(name),after.read(name),name)

    def test_exact_copy_changes_only_six_groups(self):
        original=sha(self.first)
        dest=self.path/'11_2026_quality.docx'
        result=create_quality_copy(
            self.first,dest,number='11/2026',
            end_date='2026-10-26',vik_date='2026-10-19',
            uzk_date='2026-10-20',akz_date='2026-10-24',
            mass_kg='511.98')
        d=Document(dest)
        t=d.tables[0]
        self.assertEqual(t.cell(1,9).text,'11/2026')
        self.assertEqual(t.cell(16,6).text,'511,98 кг.')
        self.assertEqual(t.cell(18,8).text,'26.10.2026 г.')
        for i in (50,53):
            text=t.cell(i,0).text
            self.assertEqual(text.count('11/2026'),4)
            self.assertEqual(text.count('19.10.2026'),2)
            self.assertEqual(text.count('20.10.2026'),1)
            self.assertEqual(text.count('24.10.2026'),1)
            self.assertNotRegex(text,r'(?<![0-9])1/2026')
            for p in t.cell(i,0).paragraphs[:1]:
                self.assertTrue(all(r.font.highlight_color is None
                                    for r in p.runs))
        self.assertIn('05.08.2026',t.cell(30,0).text)
        self.assertIn('ОБРАЗЕЦ_ФИО_02',t.cell(49,0).text)
        self.assertEqual(d.tables[1].cell(4,0).text,'26.10.2026 г.')
        self.assertEqual(sha(self.first),original)
        self.assertIn('UNVERIFIED',result['status'])

    def test_different_project_layout_rejected(self):
        d=Document(self.first)
        d.tables[0].cell(50,0).text='No valid acts'
        d.save(self.first)
        self.assertFalse(is_exact_first_quality(self.first))
        with self.assertRaisesRegex(Stop,'не совпадает'):
            create_quality_copy(
                self.first,self.path/'new.docx',number='11/2026',
                end_date='2026-10-26',vik_date='2026-10-19',
                uzk_date='2026-10-20',akz_date='2026-10-24',
                mass_kg='511.98')

    def test_missing_registry_date_fails_closed(self):
        with self.assertRaisesRegex(Stop,'дата ВИК'):
            create_quality_copy(
                self.first,self.path/'new.docx',number='11/2026',
                end_date='2026-10-26',vik_date=None,
                uzk_date='2026-10-20',akz_date='2026-10-24',
                mass_kg='511.98')
        self.assertFalse((self.path/'new.docx').exists())

    def test_no_write_to_source_or_existing_target(self):
        with self.assertRaisesRegex(Stop,'отдельный новый'):
            create_quality_copy(
                self.first,self.first,number='11/2026',
                end_date='2026-10-26',vik_date='2026-10-19',
                uzk_date='2026-10-20',akz_date='2026-10-24',
                mass_kg='511.98')
        occupied=self.path/'exists.docx'
        occupied.write_bytes(b'KEEP')
        with self.assertRaisesRegex(Stop,'отдельный новый'):
            create_quality_copy(
                self.first,occupied,number='11/2026',
                end_date='2026-10-26',vik_date='2026-10-19',
                uzk_date='2026-10-20',akz_date='2026-10-24',
                mass_kg='511.98')
        self.assertEqual(occupied.read_bytes(),b'KEEP')


if __name__=='__main__':unittest.main(verbosity=2)
