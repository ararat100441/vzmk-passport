"""Targeted r10 contracts; synthetic fixtures do not prove actual control."""
from contextlib import nullcontext
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from zipfile import ZipFile
from lxml import etree
import tempfile
import unittest
from io import BytesIO
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from docx import Document
from docx.shared import Pt
from production_app.office import word_preview_mark
from production_app.print_pdf import _export_word, _preview_overlay
from production_app.office import PREVIEW_MARK
from production_app.storage import Stop, sha
from passport_app.appendix import generate_appendix


def parts(path):
    with ZipFile(path) as z:return {n:z.read(n) for n in z.namelist()}


class WordExportGeometry(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.original=self.root/'original.docx'
        doc=Document();doc.add_table(2,2).cell(0,0).text='FIRST TABLE'
        p=doc.add_paragraph('BODY AFTER TABLE')
        p.paragraph_format.space_before=Pt(13);p.paragraph_format.space_after=Pt(7)
        doc.sections[0].header.paragraphs[0].text='STATIC HEADER'
        doc.save(self.original)
        self.source=self.root/'draft.docx';self.source.write_bytes(self.original.read_bytes())
        word_preview_mark(self.source);self.before=sha(self.source)

    def tearDown(self):self.temp.cleanup()

    def export(self,preview=True,fail=False):
        opened=[];doc=MagicMock()
        def open_copy(word,path,readonly):
            self.assertTrue(readonly)
            opened.append((Path(path),parts(path)))
            return doc
        if fail:doc.ExportAsFixedFormat.side_effect=RuntimeError('EXPORT FAILURE')
        with patch('production_app.print_pdf.word_session',return_value=nullcontext(object())), \
             patch('production_app.print_pdf.open_word',side_effect=open_copy):
            if fail:
                with self.assertRaisesRegex(RuntimeError,'EXPORT FAILURE'):
                    _export_word([dict(source=self.source,pdf=self.root/'result.pdf',role='quality',preview=preview)])
            else:_export_word([dict(source=self.source,pdf=self.root/'result.pdf',role='quality',preview=preview)])
        self.assertEqual(sha(self.source),self.before)
        return opened,doc

    def test_preview_opens_exact_unmarked_xml_and_preserves_all_opaque_parts(self):
        opened,doc=self.export()
        final=self.root/'final.docx';final.write_bytes(self.source.read_bytes())
        word_preview_mark(final,remove=True)
        self.assertEqual(opened[0][1],parts(final))
        original=parts(self.original);actual=opened[0][1]
        self.assertEqual({k:v for k,v in actual.items() if k!='word/document.xml'},
                         {k:v for k,v in original.items() if k!='word/document.xml'})
        body=lambda xml:etree.tostring(etree.fromstring(xml,etree.XMLParser(remove_blank_text=True)).find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body'),method='c14n',exclusive=True)
        self.assertEqual(body(actual['word/document.xml']),body(original['word/document.xml']))
        self.assertNotEqual(opened[0][0],self.source)
        self.assertFalse(opened[0][0].exists())
        doc.Paragraphs.assert_not_called()
        doc.Close.assert_called_once_with(SaveChanges=False)

    def test_final_opens_source_without_rewriting_it(self):
        opened,doc=self.export(False)
        self.assertEqual(opened[0][0],self.source)
        self.assertEqual(opened[0][1],parts(self.source))

    def test_export_error_closes_document_removes_temporary_copy_keeps_source(self):
        opened,doc=self.export(fail=True)
        self.assertFalse(opened[0][0].exists())
        doc.Close.assert_called_once_with(SaveChanges=False)

    def test_missing_or_altered_preview_mark_fails_before_opening_word_document(self):
        self.source.write_bytes(self.original.read_bytes())
        with patch('production_app.print_pdf.word_session',return_value=nullcontext(object())), \
             patch('production_app.print_pdf.open_word') as opened:
            with self.assertRaises(Stop):
                _export_word([dict(source=self.source,pdf=self.root/'result.pdf',role='quality',preview=True)])
            opened.assert_not_called()
        self.assertFalse(list(self.root.glob('.word-print-*')))


class AppendixAreaPrecision(unittest.TestCase):
    def generate(self,separator):
        fields=('mark','name','quantity','mass_unit','mass_total','area_unit','area_total')
        layout=SimpleNamespace(sheet='Sheet',start=5,total_row=7,columns={v:i+1 for i,v in enumerate(fields)})
        rows=[('TEST-A','TEST',Decimal(1),Decimal('1.234567'),Decimal('1.234567'),Decimal('100.123456'),Decimal('100.123456')),
              ('TEST-B','TEST',Decimal(1),Decimal('2.345678'),Decimal('2.345678'),Decimal('170.507034'),Decimal('170.507034'))]
        vom=SimpleNamespace(rows=rows,quantity=Decimal(2),mass=Decimal('3.580245'),area=Decimal('270.630490'))
        cells={}
        def cell(row,col):
            return cells.setdefault((row,col),SimpleNamespace(Value2={3:2,5:3.580245,7:270.630490}.get(col),Formula=None,NumberFormatLocal='0,00'))
        sheet=MagicMock();sheet.Cells.side_effect=cell
        sheet.PageSetup.PrintArea='$A$1:$G$7'
        ranges={}
        def rng(*args):
            key=tuple(id(x) if isinstance(x,SimpleNamespace) else x for x in args)
            return ranges.setdefault(key,MagicMock(Value2=('Приложение № 1/2026 от 01.01.2026 г.' if args==('A1',) else 'Объект: TEST')))
        sheet.Range.side_effect=rng
        book=MagicMock();book.Worksheets.return_value=sheet
        excel=MagicMock(DecimalSeparator=separator)
        with patch('passport_app.appendix.inspect_template',return_value=(layout,'A1','A2')), \
             patch('passport_app.appendix.open_excel',return_value=book):
            total=generate_appendix(excel,'OUTPUT_COPY.xlsx',{},vom,dict(number='11/2026',date=date(2026,10,31),object='TEST'))
        self.assertEqual(total,7)
        self.assertEqual(cell(7,7).Formula,'=SUM(G5:G6)')
        self.assertEqual(cell(7,7).Value2,270.630490)
        self.assertEqual(cell(7,7).NumberFormatLocal,'0'+separator+'000')
        self.assertTrue(all(c.NumberFormatLocal=='0,00' for k,c in cells.items() if k!=(7,7)))
        self.assertEqual(rng(cell(5,7),cell(6,7)).Value2,((100.123456,),(170.507034,)))
        book.Save.assert_called_once();book.Close.assert_called_once_with(SaveChanges=False)

    def test_ru_display_retains_formula_and_full_input_precision(self):self.generate(',')
    def test_other_excel_decimal_separator_retains_numeric_value(self):self.generate('.')


class PreviewMargin(unittest.TestCase):
    def test_mark_clears_unshifted_first_line_without_moving_it(self):
        data=BytesIO();c=canvas.Canvas(data,pagesize=(595,842))
        c.setFont('Helvetica',10);c.drawString(45,822,'FIRST CONTINUATION LINE');c.save();data.seek(0)
        page=PdfReader(data).pages[0];_preview_overlay(page);positions=[]
        writer=PdfWriter();writer.add_page(page);saved=BytesIO();writer.write(saved);saved.seek(0)
        page=PdfReader(saved).pages[0]
        page.extract_text(visitor_text=lambda text,cm,tm,font,size:positions.append((text.strip(),tm[5],size)))
        body=next(p for p in positions if p[0]=='FIRST CONTINUATION LINE')
        mark=next(p for p in positions if p[0]==PREVIEW_MARK)
        self.assertEqual(body[1],822)
        self.assertGreaterEqual(mark[1]-body[1],10)
        self.assertLessEqual(mark[1]+mark[2]*.8,842)


if __name__=='__main__':unittest.main()
