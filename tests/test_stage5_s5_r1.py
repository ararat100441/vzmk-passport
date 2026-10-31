"""S5-R1 only: full preview -> no-output generate, with isolated file inputs."""
import argparse
from contextlib import contextmanager, ExitStack
from copy import deepcopy
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import BadZipFile

from openpyxl import load_workbook, Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import uzk_stage5.core as C
import uzk_stage5.inputs as I
import uzk_stage5.word as W
from profiles_stage4.model import Catalog
from profiles_stage4.offline import no_network

PARSER = argparse.ArgumentParser()
PARSER.add_argument('--run', required=True)
PARSER.add_argument('--baseline', action='store_true')
PARSER.add_argument('--attempt', default='final')
ARGS = PARSER.parse_args()
RUN = Path(ARGS.run).resolve()
REVIEW = ROOT/'validation/stage5_independent_review_20261008_01'
assert RUN.is_relative_to(ROOT/'validation') and RUN != REVIEW
TEMP = RUN/'tmp'
TEMP.mkdir(exist_ok=True)
tempfile.tempdir = str(TEMP)
os.environ['TEMP'] = os.environ['TMP'] = str(TEMP)
ACTUAL_PRODUCT = [*C.product_paths(), ROOT/'tests/test_stage5.py', ROOT/'STAGE5_DEV.md', Path(__file__).resolve()]
ACCESS = []


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def write(name, value):
    (RUN/name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf8')


def fingerprints(folder):
    return {str(p.relative_to(folder)): I.sha(p) for p in folder.rglob('*') if p.is_file()}


def directories(folder):
    return sorted(str(p.relative_to(folder)) for p in folder.rglob('*') if p.is_dir())


def access_guard(event, args):
    if event != 'open' or not isinstance(args[0], (str, bytes)):
        return
    path = Path(os.fsdecode(args[0])).resolve()
    if path.name.casefold() in ('взмк.xlsx', 'узк.xlsm'):
        raise AssertionError('PRODUCTION_ACCESS_FORBIDDEN: ' + str(path))
    mode, flags = args[1:3]
    writing = (isinstance(mode, str) and any(c in mode for c in 'wax+')) or (
        isinstance(flags, int) and flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC))
    if writing and not path.is_relative_to(RUN):
        raise AssertionError('WRITE_OUTSIDE_FIXTURE_RUN: ' + str(path))
    ACCESS.append({'path': str(path), 'writing': bool(writing)})


sys.addaudithook(access_guard)


@contextmanager
def sandbox(name, *, no_uzk=False):
    source = REVIEW/('scenarios/no_uzk_True' if no_uzk else 'isolated_base')
    root = RUN/'scenarios'/ARGS.attempt/name
    shutil.copytree(source, root)
    paths = {'ROOT': root, 'REGISTRY': root/'DEV/ВЗМК_DEV.xlsx', 'UZK': root/'DEV/УЗК_DEV.xlsm',
             'PROFILES': root/'DEV/Справочники/Профили.xlsx', 'CHECKPOINT': root/'DEV/DEV_UZK_CHECKPOINT.json'}
    pins = {paths[k]: I.sha(paths[k]) for k in ('REGISTRY', 'UZK', 'PROFILES')}
    def refresh():
        for p in pins:
            if p.is_file():
                pins[p] = I.sha(p)
        cp = read(paths['CHECKPOINT'])
        cp['dev_path'] = str(paths['UZK'])
        cp['dev_initial_sha256'] = pins[paths['UZK']]
        paths['CHECKPOINT'].write_text(json.dumps(cp, ensure_ascii=False), encoding='utf8')
    refresh()
    passport_reader, registry_reader = I.passports, I.registry_rows
    with ExitStack() as stack:
        for module in (I, C):
            for key, value in paths.items():
                stack.enter_context(patch.object(module, key, value))
        stack.enter_context(patch.object(I, 'PINS', pins))
        stack.enter_context(patch.object(I, 'passports', side_effect=lambda: passport_reader(paths['REGISTRY'])))
        stack.enter_context(patch.object(C, 'registry_rows', side_effect=lambda: registry_reader(paths['UZK'])))
        # Execute CURRENT product bytes, not the old product copies in fixtures.
        stack.enter_context(patch.object(C, 'product_paths', return_value=ACTUAL_PRODUCT))
        yield root, root/'Документы/Гараж казахстан/source.xls', refresh


def catalog_path(root):
    return root/'Документы/Гараж казахстан/Справочник марок.xlsx'


def add_unrelated(root):
    path = catalog_path(root)
    book = load_workbook(path)
    sheet = book.active
    for mark, profile in [('FOREIGN_A', '[16П'), (' foreign_a ', 'Гн.[]120х3'),
                          ('FOREIGN_B', 'Гн.[140X60X4'), ('foreign_b', '[16П'),
                          ('FOREIGN_CHECK', 'ПРОВЕРИТЬ'), ('FOREIGN_MISSING', None)]:
        sheet.append([mark, profile, None, None])
    book.save(path)
    book.close()


def mutate_registry(root, operation):
    path = root/'DEV/УЗК_DEV.xlsm'
    book = load_workbook(path, keep_vba=True)
    operation(book['УЗК'])
    book.save(path)
    book.close()


def change_required_profile(root, profile):
    path = catalog_path(root)
    book = load_workbook(path)
    for row in book.active:
        if I.key(row[0].value) == I.key('Бп2-1'):
            row[1].value = profile
    book.save(path)
    book.close()


def get_preview(vom):
    return C.preview('Гараж Казахстан', '9/2026', vom)


if ARGS.baseline:
    with sandbox('baseline_reproduction', no_uzk=True) as (root, vom, _):
        before = fingerprints(root)
        stock, rows = C.registry_rows()
        assert len(rows) == 25 and not any(project == I.key('Гараж Казахстан') for project, mark in stock)
        try:
            get_preview(vom)
        except I.Stage5Error as exc:
            assert 'MARK_PROFILE_CONFLICT' in str(exc)
            result = {'finding': 'S5-R1', 'reproduced': True, 'actual': str(exc),
                      'valid_registry_rows': len(rows), 'selected_project_matches': 0}
        else:
            raise AssertionError('Reviewed S5-R1 was not reproduced')
        assert fingerprints(root) == before
        result['fixture_inputs_unchanged'] = True
    write('baseline_reproduction.json', result)
    print(json.dumps(result, ensure_ascii=False))
    raise SystemExit(0)


class S5R1Tests(unittest.TestCase):
    def assert_empty(self, plan):
        self.assertEqual(plan.data['status'], 'NO_UZK_REQUIRED')
        self.assertEqual(plan.data['marks'], [])
        self.assertEqual(plan.data['controlled_units'], 0)
        self.assertTrue(plan.data['excluded_without_registry'])
        plan.verify()

    def test_A_zero_matches_two_unrelated_conflicts(self):
        with sandbox('A', no_uzk=True) as (root, vom, _):
            add_unrelated(root)
            before = fingerprints(root)
            with patch.object(C, 'mark_profiles', side_effect=AssertionError('Catalog must not load')), \
                 patch.object(C, 'load_catalog', side_effect=AssertionError('Stage 4 must not load')):
                plan = get_preview(vom)
            self.assert_empty(plan)
            self.assertEqual(fingerprints(root), before)
            write('A_empty_preview.json', {'data': plan.data, 'seal': plan.seal})

    def test_B_full_preview_generate_no_outputs(self):
        with sandbox('B', no_uzk=True) as (root, vom, _):
            add_unrelated(root)
            before, dirs = fingerprints(root), directories(root)
            with patch.object(W, 'write_test_document', side_effect=AssertionError('No Word renderer call')) as writer:
                plan = get_preview(vom)
                self.assert_empty(plan)
                requested = root/'validation/must_not_exist'
                self.assertIsNone(C.generate(plan, requested))
                self.assertIsNone(C.generate(plan))
                writer.assert_not_called()
            self.assertFalse(requested.exists())
            self.assertFalse((root/'DEV/Stage5_Выход').exists())
            self.assertEqual(directories(root), dirs)
            self.assertEqual(fingerprints(root), before)
            write('B_no_output.json', {'status': plan.data['status'], 'new_docx': 0,
                'new_output_directories': 0, 'source_mutations': 0, 'writer_calls': 0,
                'input_fingerprints': before})

    def test_C_same_project_marks_absent_from_registry(self):
        with sandbox('C', no_uzk=True) as (root, vom, refresh):
            def edit(sheet):
                for row in sheet.iter_rows(min_row=2):
                    row[0].value = 'Гараж Казахстан'
                    row[1].value = 'ABSENT_' + str(row[1].value)
            mutate_registry(root, edit)
            refresh()
            self.assert_empty(get_preview(vom))

    def test_D_same_marks_other_project(self):
        with sandbox('D', no_uzk=True) as (root, vom, _):
            stock, rows = C.registry_rows()
            self.assertEqual(len(rows), 25)
            self.assertTrue(any(mark == I.key('Бп2-1') for project, mark in stock))
            self.assertFalse(any(project == I.key('Гараж Казахстан') for project, mark in stock))
            self.assert_empty(get_preview(vom))

    def test_E_required_mapping_conflict_stops(self):
        with sandbox('E') as (root, vom, _):
            path = catalog_path(root); book = load_workbook(path)
            book.active.append([' бп2-1 ', 'Гн.[]120х3', None, None]); book.save(path); book.close()
            with self.assertRaisesRegex(I.Stage5Error, 'MARK_PROFILE_CONFLICT'):
                get_preview(vom)

    def test_F1_missing_registry_stops(self):
        with sandbox('F_missing', no_uzk=True) as (root, vom, _):
            (root/'DEV/УЗК_DEV.xlsm').unlink()
            with self.assertRaisesRegex(I.Stage5Error, 'CHECKPOINT_CHANGED'):
                get_preview(vom)

    def test_F2_unreadable_registry_stops(self):
        with sandbox('F_unreadable', no_uzk=True) as (root, vom, refresh):
            (root/'DEV/УЗК_DEV.xlsm').write_bytes(b'ISOLATED INVALID WORKBOOK')
            refresh()
            with self.assertRaises(BadZipFile):
                get_preview(vom)

    def test_G_unrelated_conflicts_and_check_do_not_block_selected(self):
        with sandbox('G') as (root, vom, _):
            add_unrelated(root)
            before = fingerprints(root)
            plan = get_preview(vom)
            self.assertEqual(plan.data['status'], 'UZK_REQUIRED')
            self.assertEqual([m['mark'] for m in plan.data['marks']], ['Бп2-1', 'Пс1-9', 'Пс1-10', 'Пс2-16'])
            self.assertEqual(plan.data['controlled_units'], 33)
            self.assertEqual(fingerprints(root), before)
            write('G_required_preview.json', {'data': plan.data, 'seal': plan.seal})

    def test_registry_structure_still_stops(self):
        with sandbox('invalid_registry_structure', no_uzk=True) as (root, vom, refresh):
            mutate_registry(root, lambda s: setattr(s['A1'], 'value', 'WRONG HEADER'))
            refresh()
            with self.assertRaisesRegex(I.Stage5Error, 'UZK_HEADER_CHANGED'):
                get_preview(vom)

    def test_registry_invalid_quantity_still_stops(self):
        with sandbox('invalid_registry_quantity', no_uzk=True) as (root, vom, refresh):
            mutate_registry(root, lambda s: setattr(s['C2'], 'value', -1))
            refresh()
            with self.assertRaisesRegex(I.Stage5Error, 'QUANTITY_INVALID'):
                get_preview(vom)

    def test_required_registry_conflict_still_stops(self):
        with sandbox('required_registry_conflict') as (root, vom, refresh):
            mutate_registry(root, lambda s: s.append(['Гараж Казахстан', 'бп2-1', 9]))
            refresh()
            with self.assertRaisesRegex(I.Stage5Error, 'REGISTRY_CONFLICT'):
                get_preview(vom)

    def test_invalid_vom_still_stops(self):
        with sandbox('invalid_vom', no_uzk=True) as (root, vom, _):
            invalid = vom.with_suffix('.xlsx')
            book = Workbook(); book.active.append(['INVALID VOM HEADERS']); book.save(invalid); book.close()
            with self.assertRaisesRegex(I.Stage5Error, 'VOM_SHEET_AMBIGUOUS'):
                get_preview(invalid)

    def test_required_missing_unresolved_not_found_stop(self):
        for index, (profile, error) in enumerate([(None, 'TEXT_INVALID'), ('ПРОВЕРИТЬ', 'PROFILE_UNRESOLVED'),
                                                 ('UNKNOWN_PROFILE_S5R1', 'PROFILE_NOT_FOUND')]):
            with self.subTest(profile=profile), sandbox('required_profile_' + str(index)) as (root, vom, _):
                change_required_profile(root, profile)
                with self.assertRaisesRegex(I.Stage5Error, error):
                    get_preview(vom)

    def test_required_stage4_ambiguity_conflict_unsupported_stop(self):
        source = deepcopy(C.load_catalog().lookup('[16П')['profile']['provenance'][0])
        other = deepcopy(source); other['record_id'] = 'ISOLATED_S5R1'; other['geometry']['b'] = 65
        other['standard'] = 'ГОСТ 99999-2000'
        cases = [('AMBIGUOUS', Catalog([source, other]))]
        other['standard'] = source['standard']
        cases.extend([('CONFLICT', Catalog([source, other])), ('UZK_RULE_NOT_CONFIGURED', Catalog([source]))])
        for name, catalog in cases:
            with self.subTest(name=name), sandbox('stage4_' + name) as (root, vom, _):
                with patch.object(C, 'load_catalog', return_value=catalog):
                    with self.assertRaisesRegex(I.Stage5Error, name):
                        get_preview(vom)

    def test_filtered_catalog_preserves_default_strictness(self):
        with sandbox('mapping_filter') as (root, vom, _):
            add_unrelated(root)
            with self.assertRaisesRegex(I.Stage5Error, 'MARK_PROFILE_CONFLICT'):
                I.mark_profiles(catalog_path(root))
            mappings = I.mark_profiles(catalog_path(root), required_marks=[' бп2-1 ', 'ПС1-9'])
            self.assertEqual(set(mappings), {I.key('Бп2-1'), I.key('Пс1-9')})

    def test_ui_empty_preview_disables_generation(self):
        with sandbox('ui_empty', no_uzk=True) as (root, vom, _):
            plan = get_preview(vom)
        import tkinter as tk
        from uzk_stage5.gui import App
        window = tk.Tk(); window.withdraw()
        try:
            app = App(window)
            app.queue.put(('preview', plan, None)); app.poll()
            self.assertEqual(app.status.get(), 'Для текущего паспорта УЗК не требуется')
            self.assertTrue(app.generate_button.instate(['disabled']))
            app.checked.set(True); app.update_buttons()
            self.assertTrue(app.generate_button.instate(['disabled']))
            self.assertIsNone(app.output)
        finally:
            window.destroy()

    def test_H_source_and_previous_review_preservation(self):
        baseline = read(RUN/'baseline.json')['protected_before']
        allowed = {str(ROOT/'uzk_stage5/core.py'), str(ROOT/'uzk_stage5/inputs.py'), str(ROOT/'uzk_stage5/gui.py')}
        changed = [p for p, expected in baseline.items() if I.sha(p) != expected]
        self.assertEqual(set(changed), allowed)


binding = {str(p): I.sha(p) for p in ACTUAL_PRODUCT}
write('product_before_final_tests.json', binding)
with no_network() as attempts, patch.object(W, 'write_test_document', side_effect=AssertionError('No new DOCX permitted')):
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(S5R1Tests)
    with (RUN/'targeted_tests.log').open('w', encoding='utf8') as log:
        result = unittest.TextTestRunner(stream=log, verbosity=2).run(suite)
after = {str(p): I.sha(p) for p in ACTUAL_PRODUCT}
assert binding == after, 'Product mutated during tests'
write('product_after_final_tests.json', after)
write('targeted_tests.json', {'tests_run': result.testsRun, 'successful': result.wasSuccessful(),
    'failures': [(str(t), text) for t, text in result.failures],
    'errors': [(str(t), text) for t, text in result.errors], 'skipped': result.skipped,
    'network_attempts': attempts, 'fixture_docx_generation': 0})
write('access_guard.json', {'open_events': ACCESS, 'production_reads': 0, 'writes_outside_run': 0})
print(json.dumps({'tests_run': result.testsRun, 'successful': result.wasSuccessful(), 'network_attempts': attempts}))
raise SystemExit(0 if result.wasSuccessful() else 1)
