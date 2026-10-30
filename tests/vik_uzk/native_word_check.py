"""Targeted read-only Microsoft Word QA of a generated Stage 9 TEST run."""
import argparse
from collections import OrderedDict
import gc
import json
from pathlib import Path
import subprocess
import sys
import traceback
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from lxml import etree as E
from openpyxl import load_workbook
from passport_app.vik import word_session,open_word
from stage9_app.ownership import owned_office
from stage8_closure import storage as S
from uzk_stage5.inputs import registry_rows,key
from uzk_stage5.word import contents,NS
from stage9_app.test_word import TEST_VIK

def check(condition,reason):
    if not condition:raise AssertionError(reason)

def xml(path):
    with ZipFile(path) as z:return E.fromstring(z.read('word/document.xml'))

def structural_checks(run,manifest):
    records={r['type']:Path(r['path']) for r in manifest['records']}
    vik,uzk=records['VIK'],records['UZK']
    with ZipFile(run/'VIK_BEFORE_TEST_MARKER.docx') as old, ZipFile(vik) as new:
        check(set(new.namelist())-set(old.namelist())=={'word/header1.xml'},'VIK_HEADER_ADDITION')
        changed={n for n in old.namelist() if old.read(n)!=new.read(n)}
        check(changed=={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'},'VIK_UNRELATED_PARTS_CHANGED')
    before,after=xml(run/'VIK_BEFORE_TEST_MARKER.docx'),xml(vik)
    check(before.xpath('.//w:t/text()',namespaces=NS)==after.xpath('.//w:t/text()',namespaces=NS),'VIK_BODY_TEXT_CHANGED')
    check(len(after.xpath('.//w:p[starts-with(string(.),"АКТ № ")]',namespaces=NS))==2,'VIK_TWO_ACTS')
    book=load_workbook(run/'VISIBLE_VOM.xlsx',data_only=True)
    shipment=OrderedDict()
    try:
        for row in book.active.iter_rows(min_row=3,values_only=True):
            mark,quantity=row[0],row[2]
            if not mark or mark=='Итого':continue
            shipment[mark]=shipment.get(mark,0)+int(quantity)
    finally:book.close()
    check(len(shipment)==104,'EXPECTED_104_MARKS')
    root=Path(manifest['registry_source_at_run']['path']).parents[1]
    stock,_=registry_rows(root/'DEV/УЗК_DEV.xlsm')
    expected=[]
    for mark,quantity in shipment.items():
        record=stock.get((key(manifest['project']),key(mark)))
        if record is not None:
            expected.append({'mark':mark,'Q':min(quantity,int(record['remaining'])),'vom_quantity':quantity})
    check(len(expected)==4 and sum(m['Q'] for m in expected)==33,'EXPECTED_4_MARKS_33_UNITS')
    before,after=xml(run/'UZK_BEFORE_TEST_LAYOUT.docx'),xml(uzk)
    bt,at=before.findall('w:body/w:tbl',NS),after.findall('w:body/w:tbl',NS)
    check(len(bt)==len(at)==2,'UZK_TABLE_STRUCTURE')
    brows,arows=bt[0].findall('w:tr',NS),at[0].findall('w:tr',NS)
    check(len(brows)==len(arows)==9,'UZK_4_DATA_ROWS')
    for i,(br,ar) in enumerate(zip(brows,arows)):
        bc,ac=br.findall('w:tc',NS),ar.findall('w:tc',NS)
        check(len(bc)==len(ac),'UZK_CELL_COUNT')
        for j,(x,y) in enumerate(zip(bc,ac)):
            if j==0 and i>=4:continue
            check(contents(x)==contents(y),'UZK_OTHER_TEXT_CHANGED')
    check(contents(bt[1])==contents(at[1]),'UZK_SIGNATURE_TEXT_CHANGED')
    check([contents(p) for p in before.findall('w:body/w:p',NS)]==[contents(p) for p in after.findall('w:body/w:p',NS)],'UZK_PARAGRAPHS_CHANGED')
    check(contents(arows[4].findall('w:tc',NS)[0])=='МаркаУЗК/ВОМ, шт.','UZK_HEADER')
    for row,mark in zip(arows[5:],expected):
        cell=row.findall('w:tc',NS)[0]
        check(contents(cell)==mark['mark']+str(mark['Q'])+'/'+str(mark['vom_quantity']),'UZK_PER_MARK_QUANTITY')
        check(len(cell.findall('.//w:br',NS))==1,'UZK_QUANTITY_NEWLINE')
    return records,expected

def inspect(word,path,out,kind):
    out.mkdir(exist_ok=False)
    before=S.sha(path)
    doc=open_word(word,path,readonly=True)
    try:
        doc.ActiveWindow.View.Type=3;doc.Repaginate()
        pages=int(doc.ComputeStatistics(2))
        check(pages==(2 if kind=='VIK' else 1),kind+'_PAGE_COUNT')
        doc.ExportAsFixedFormat(str(out/'print.pdf'),17,OpenAfterExport=False)
        result={'pages':pages,'readonly':bool(doc.ReadOnly),'word_version':word.Version,'word_build':word.Build,
                'text':doc.Content.Text,'page_texts':[],'headers':[],'width':doc.PageSetup.PageWidth,'height':doc.PageSetup.PageHeight}
        for i in range(1,pages+1):
            start=doc.GoTo(1,1,i).Start
            end=doc.GoTo(1,1,i+1).Start if i<pages else doc.Content.End
            text=doc.Range(start,end).Text
            check(bool(text.strip('\r\x07\x0c\n ')),kind+'_BLANK_PAGE')
            result['page_texts'].append(text)
            (out/('page-'+str(i)+'.emf')).write_bytes(bytes(doc.ActiveWindow.Panes(1).Pages(i).EnhMetaFileBits))
        if kind=='VIK':
            result['headers']=[doc.Sections(i).Headers(1).Range.Text for i in range(1,doc.Sections.Count+1)]
            check(all(TEST_VIK in text for text in result['headers']),'VIK_TEST_HEADER_MISSING')
            check(all('АКТ № 9/2026' in text for text in result['page_texts']),'VIK_CURRENT_ACT_NUMBERS')
        else:
            signature=doc.Tables(2).Range
            start=int(doc.Range(signature.Start,signature.Start).Information(3))
            end=int(doc.Range(signature.End-1,signature.End-1).Information(3))
            check(start==end==1,'UZK_SIGNATURE_NOT_COMPLETELY_ON_PAGE_1')
            result['signature_pages']=[start,end]
            del signature
    finally:doc.Close(False)
    check(S.sha(path)==before,'WORD_READONLY_CHANGED_DOCUMENT')
    result['sha256']=before;result['status']='PASS'
    (out/'WORD.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    command=[str(Path(__import__('os').environ['SystemRoot'])/'System32/WindowsPowerShell/v1.0/powershell.exe'),
             '-NoProfile','-File',str(Path(__file__).with_name('render_native.ps1')),
             '-Directory',str(out),'-PageWidth',str(result['width']),'-PageHeight',str(result['height'])]
    subprocess.run(command,check=True,timeout=60,capture_output=True)
    # Page.EnhMetaFileBits can crop a later page's header in Word's display
    # representation. Its print export supplies the authoritative page image.
    import fitz
    with fitz.open(out/'print.pdf') as pdf:
        check(len(pdf)==pages,'WORD_PRINT_PAGE_COUNT')
        result['printed_test_marker']=[]
        for i,page in enumerate(pdf,1):
            if kind=='VIK':
                matches=page.search_for(TEST_VIK)
                check(bool(matches) and all(page.rect.contains(r) for r in matches),'VIK_PRINT_TEST_MARKER_CLIPPED')
                result['printed_test_marker'].append(True)
            page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(out/('print-page-'+str(i)+'.png'))
    (out/'PRINT.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    return result

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifest',type=Path,required=True);parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();out=args.output.resolve();out.mkdir(parents=True,exist_ok=False)
    manifest=S.read(args.manifest);run=args.manifest.parent
    check(manifest['passport']=='9/2026' and manifest['visible_rows']==104 and manifest['uzk_units']==33,'TARGET_MANIFEST')
    records,expected=structural_checks(run,manifest)
    result={'status':'FAIL','manifest':str(args.manifest),'uzk_marks':expected}
    with owned_office(out/'OFFICE.json'):
        with word_session() as word:
            try:
                for kind in ('VIK','UZK'):result[kind]=inspect(word,records[kind],out/kind,kind)
                result['status']='PASS'
            except Exception:result['error']=traceback.format_exc()
        del word;gc.collect()
    (out/'RESULT.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in result.items() if k not in ('VIK','UZK')},ensure_ascii=False,indent=2))
    return 0 if result['status']=='PASS' else 1

if __name__=='__main__':sys.exit(main())
