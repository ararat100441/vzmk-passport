"""Resume the immutable R7 run after the test-only OLE-container comparison fix."""
from common import *
import traceback,io,hashlib,olefile
from zipfile import ZipFile
from stage8_closure.core import confirmation
from stage8_closure.recovery import confirmation as restore_confirmation
RUN=OUT/'e2e_2b9df014'
failure=read(RUN/'FAILURE.json');steps=failure['steps'][:]
assert sha(EXE)==failure['exe_sha256']=='677a111e1762e72949125e74977ca3468f114895a2de92f705c49264aec61c88'
assert len(steps)==19
def result(index):return read(steps[index-1]['result'])
def stocks(sandbox):
    from openpyxl import load_workbook
    b=load_workbook(Path(sandbox)/'registry/УЗК_TEST.xlsm',read_only=True,data_only=True)
    try:return [b['УЗК'].cell(r,3).value for r in (24,25)]
    finally:b.close()
def streams(path):
    with ZipFile(path) as z:binary=z.read('xl/vbaProject.bin')
    with olefile.OleFileIO(io.BytesIO(binary)) as ole:
        return {'/'.join(n):hashlib.sha256(ole.openstream(n).read()).hexdigest() for n in ole.listdir()}
def registry(sandbox):return Path(sandbox)/'registry/УЗК_TEST.xlsm'
def run(name,expected='PASS',**args):
    data,sup,path=task(RUN,name,expected,**args)
    steps.append({'task':name,'result':str(path),'supervisor':data['office_supervisor'],'expected':expected})
    save(RUN/f'step_{len(steps):02}.json',steps[-1]);return data
try:
    work=result(1)['workspace'];outputs=[result(n)['manifest'] for n in (4,8,12)]
    fixture=result(16);plan=RUN/'CLOSE_PLAN.json';preview=read(plan)
    closed=result(19);assert closed['operation_status']=='CLOSED_TEST'
    assert stocks(fixture['sandbox'])==[9,6]
    proof=read(RUN/'VBA_STREAM_READBACK.json');assert proof['status']=='PASS' and proof['all_streams_including_compiled_vba_byte_equal']
    assert streams(registry(fixture['sandbox']))==proof['streams']
    again=run('close',workspace=work,preview=plan,confirm=confirmation(preview))
    assert again['operation_status']=='ALREADY_CLOSED_TEST' and stocks(fixture['sandbox'])==[9,6]
    run('recover',workspace=work,sandbox=fixture['sandbox'])
    torn=run('fixture',workspace=work);before_sha=sha(registry(torn['sandbox']));torn_vba=streams(registry(torn['sandbox']))
    preview=run('close-preview',workspace=work,sandbox=torn['sandbox'],request=torn['request'])['plan']
    plan=RUN/'CRASH_PLAN.json';save(plan,preview)
    crash=run('close','FAIL',workspace=work,preview=plan,confirm=confirmation(preview),test_fault='during_publish')
    assert crash['worker_returncode']==73
    rp=run('restore-preview',workspace=work,sandbox=torn['sandbox'])['plan'];rpath=RUN/'RESTORE_PLAN.json';save(rpath,rp)
    restored=run('restore',workspace=work,preview=rpath,confirm=restore_confirmation(rp))
    assert restored['operation_status']=='RECOVERED_ABORTED' and restored['decrements']==0
    assert stocks(torn['sandbox'])==[12,10] and streams(registry(torn['sandbox']))==torn_vba
    assert sha(registry(torn['sandbox']))==before_sha
    after=run('close-preview',workspace=work,sandbox=torn['sandbox'],request=torn['request'])
    assert after['plan']['payload']['status']=='READY_TEST'
    run('smoke',workspace=work)
    save(RUN/'RESULT.json',{'status':'PASS','exe_sha256':sha(EXE),'workspace':work,'manifests':outputs,'fixture_closed':fixture,'fixture_restored':torn,'steps':steps,'python_in_path':None,'resumed_from_step':19,'prior_test_failure_preserved':str(RUN/'FAILURE.json'),'vba_readback':str(RUN/'VBA_STREAM_READBACK.json'),'restored_sha256':before_sha})
    save(OUT/'EXE_E2E_FINAL.json',{'status':'PASS','exe_sha256':sha(EXE),'evidence':str(RUN/'RESULT.json'),'workspace':work,'manifests':outputs})
except Exception as exc:
    save(RUN/'RESUME_FAILURE.json',{'status':'FAIL','exe_sha256':sha(EXE),'error':repr(exc),'traceback':traceback.format_exc(),'steps':steps});raise
