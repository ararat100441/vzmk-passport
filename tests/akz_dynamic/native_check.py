"""Opt-in AKZ-only native regression; never invokes passport/registry generation.

Run with Python 3.12 on Windows with Excel. Inputs are pinned, read-locked and
hashed. All generated workbooks, print PDFs and diagnostics go to a NEW output.
"""
import argparse
import ast
from collections import Counter
from copy import copy
from dataclasses import asdict
from decimal import Decimal
import hashlib
import gc
import json
from pathlib import Path
import shutil
import subprocess
import sys
import traceback

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from akz_stage6 import core as c, part2
from stage9_app import akz_auto as auto
from stage9_app.ownership import owned_office
from stage8_closure import storage as storage

COUNTS = (1, 7, 17, 18, 104, 200)
FIELDS = ('mark', 'name', 'quantity', 'mass_unit', 'mass_total', 'area_unit', 'area_total')

def save(path, data):
    with path.open('x', encoding='utf-8') as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2, default=str)

def fingerprint(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def check(ok, message):
    if not ok:
        raise AssertionError(message)

def syntax():
    files = subprocess.check_output(['git', 'diff', '--name-only', 'origin/main...HEAD'], cwd=ROOT, text=True).splitlines()
    files += ['tests/akz_dynamic/native_check.py']
    for name in dict.fromkeys(files):
        if name.endswith('.py'):
            ast.parse((ROOT/name).read_text(encoding='utf-8-sig'), filename=name)
    return [f for f in dict.fromkeys(files) if f.endswith('.py')]

def payload(variant, count, template):
    rows = []
    assignments = {}
    for i in range(count):
        mark = f'A{variant}M{i+1:04}'
        qty = 1+i % 5
        mass, area = Decimal('3.25')+Decimal(i)/100, Decimal('1.125')+Decimal(i)/1000
        rows.append(dict(row=i+4, mark=mark, name='Балка TEST', quantity=qty,
                         mass_unit=mass, mass_total=mass*qty, area_unit=area, area_total=area*qty))
        assignments[mark] = dict(row=i+2, assignment=variant)
    for n in [None, *[v for v in range(4) if v != variant]]:
        mark = 'UNASSIGNED' if n is None else f'OTHER{n}'
        rows.append(dict(rows[0], mark=mark, row=1000+len(rows)))
        assignments[mark] = dict(row=1000+len(rows), assignment=n)
    selected, excluded, totals = (c.matching({'rows': rows}, assignments) if variant == 0
                                else part2.matching({'rows': rows}, assignments, variant))
    check(len(selected) == count and len(excluded) == 4, 'GROUP_ISOLATION')
    check([r['mark'] for r in selected] == [r['mark'] for r in rows[:count]], 'MATCHING_NO_LOSS')
    check(next(r for r in excluded if r['mark']=='UNASSIGNED')['reason']=='NO_ASSIGNMENT', 'BLANK_IS_NOT_ZERO')
    data = dict(mode='TEST_ONLY', synthetic=True, variant=variant,
                passport=dict(number='9/2026', date='2026-10-05', contract='TEST contract',
                              customer='TEST customer', object='TEST object', cipher='TEST cipher'),
                templates=template, act_number=f'SYNTHETIC-TEST 9/2026-A{variant}', date='2026-10-05',
                sources={}, selected=selected, excluded=excluded, totals=totals,
                climate=asdict(c.ClimateDay('2026-10-05')), official='STOP')
    return c.Preview(data, c.digest(data))

def inspect_application(app, path, pdf, p, template):
    n = p.payload['variant']; count = len(p.payload['selected'])
    first = 4 if n == 0 else 5; total = first+count
    original_total = 21 if n == 0 else part2.layout(n).total
    old = load_workbook(template['application']); new = load_workbook(path); cached = load_workbook(path, data_only=True)
    try:
        a, b, v = old.active, new.active, cached.active
        check(new.sheetnames == old.sheetnames, 'SHEET_NAMES')
        for r, expected in enumerate(p.payload['selected'], first):
            for col, field in enumerate(FIELDS, 1):
                actual = b.cell(r, col).value
                check(abs(Decimal(str(actual))-expected[field]) <= Decimal('0.00000001')
                      if isinstance(expected[field], Decimal) else actual==expected[field], f'CELL_VALUE {r}/{col}')
                before, after = a.cell(min(r, original_total-1), col), b.cell(r, col)
                for key in ('font', 'fill', 'border', 'alignment', 'protection'):
                    check(copy(getattr(before,key))==copy(getattr(after,key)), f'BODY_STYLE {r}/{col}/{key}')
        check(b.cell(total,1).value=='Итого', 'TOTAL_IMMEDIATELY_AFTER_MARKS')
        check(not any(b.cell(r,1).value for r in range(total+1,b.max_row+1)), 'STALE_MARKS_BELOW_TOTAL')
        for col in 'ABCDEFG':
            for key in ('font','fill','border','alignment','protection'):
                check(copy(getattr(a[f'{col}{original_total}'],key))==copy(getattr(b[f'{col}{total}'],key)), f'TOTAL_STYLE {col}/{key}')
        for r in range(1,first):
            for col in range(1,8):
                for key in ('font','fill','border','alignment'):
                    check(copy(getattr(a.cell(r,col),key))==copy(getattr(b.cell(r,col),key)), f'HEADER_STYLE {r}/{col}/{key}')
        for col, field in [('C','quantity'),('E','mass_total'),('G','area_total')]:
            check(b[f'{col}{total}'].value==f'=SUM({col}{first}:{col}{total-1})', 'SUM_RANGE')
            check(abs(Decimal(str(v[f'{col}{total}'].value))-p.payload['totals'][field])<Decimal('0.00000001'), 'CACHED_TOTAL')
        check(b.print_title_rows==f'$1:${first-1}', 'REPEATED_HEADER_SETTING')
        check(str(b.print_area).replace('$','').endswith(f'G{total}'), 'FULL_PRINT_AREA')
        check(b.page_setup.fitToWidth in (None,1) and not b.page_setup.fitToHeight, 'MULTIPAGE_SETTING')
        check(all(not b.row_dimensions[r].hidden for r in range(1,total+1)), 'VISIBLE_ROWS')
    finally:
        old.close(); new.close(); cached.close()
    book = c.native_open(app,path,read_only=True)
    try:
        pages = int(book.Worksheets(1).PageSetup.Pages.Count)
        book.Worksheets(1).ExportAsFixedFormat(0, str(pdf))
    finally:
        book.Close(SaveChanges=False)
    import fitz
    with fitz.open(pdf) as document:
        check(len(document)==pages, 'NATIVE_PDF_PAGE_COUNT')
        texts=[page.get_text() for page in document]
        printed=Counter(__import__('re').findall(r'A[0-3]M\d{4}', '\n'.join(texts)))
        check(printed==Counter(r['mark'] for r in p.payload['selected']), 'PRINT_NO_LOSS_OR_DUPLICATES')
        check(all('марк' in text.lower() and 'масса' in text.lower() for text in texts), 'PRINT_REPEATED_HEADER')
        check(sum(text.count('Итого') for text in texts)==1 and 'Итого' in texts[-1], 'PRINT_TOTAL_LAST_PAGE')
        if count >= 104:
            check(pages>1, 'LARGE_APPENDIX_MUST_SPAN_PAGES')
        if count in (1, 18, 200):
            for i in sorted({0,len(document)-1}):
                document[i].get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(pdf.with_name(pdf.stem+f'-p{i+1}.png'))
    return dict(pages=pages, total_row=total, marks=count, no_loss=True, formulas=True,
                body_and_total_styles=True, print_headers=True, pdf=str(pdf))

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--variant', type=int, choices=range(4))
    parser.add_argument('--count', type=int, choices=COUNTS)
    args=parser.parse_args(); out=args.output.resolve();out.mkdir(parents=True,exist_ok=False)
    source=args.source.resolve();project=source.parent/'Гараж казахстан'
    c.PROJECT=project;c.ACT=project/'Паспорт/Акт № 4-2026.xlsx';c.APPLICATION=project/'Паспорт/1_2026 - Приложение к АКЗ.xlsx'
    c.PINS={project/'Паспорт'/p.name:h for p,h in c.PINS.items()}
    protected=set(c.PINS)
    for n in (1,2,3):
        protected.update((part2.layout(n).act, part2.layout(n).application))
    # Existing registry/catalog files are read-only inputs; no generator that
    # appends or publishes registry records is called by this harness.
    for folder in (source/'DEV', project):
        protected.update(p for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in ('.xls','.xlsx','.xlsm') and not p.is_symlink())
    protected.add(source/'stage9_dist_0912/VZMK-Passport-0.9.12/VZMK-Passport-0.9.12.exe')
    production_registry=source.parent.parent/'ВЗМК.xlsx'
    if production_registry.is_file():
        protected.add(production_registry)
    before={str(p):fingerprint(p) for p in sorted(protected)}
    save(out/'BEFORE.json', before)
    results=[]; result={'syntax':syntax(),'cases':results,'official_issue':False,'full_stage_audit':'NOT_RUN','work':'NOT_RUN'}
    try:
        with storage.frozen(protected), owned_office(out/'office.json'):
          with c.native_excel() as app:
            try:
              for n in ([args.variant] if args.variant is not None else range(4)):
                template=c.templates() if n==0 else part2.template(n)
                for count in ([args.count] if args.count else COUNTS):
                    folder=out/f'akz{n}_{count:03}';folder.mkdir()
                    p=payload(n,count,template)
                    act=folder/'act.xlsx'; application=folder/'application.xlsx'; standard=folder/'stage6-act.xlsx'
                    for src,dst in [(template['act'],act),(template['act'],standard),(template['application'],application)]:
                        shutil.copyfile(src,dst)
                    shapes=(auto.shape(act),auto.shape(application))
                    (c.write_act if n==0 else part2.write_act)(app,standard,p)
                    cfg={'climate':{'air_c':[20,20],'surface_c':[25,25],'relative_humidity':[50,50],'formula':auto.FORMULA},
                         'thickness':{str(k):[list(pair) for pair in pairs] for k,pairs in auto.DEFAULT_THICKNESS.items()}}
                    climate=folder/'simulation';climate.mkdir()
                    proof=auto.write_act(app,act,p,cfg,climate)
                    save(folder/'preview.json',p.payload)
                    save(folder/'proof.json',proof)
                    auto.write_application(app,application,p,shapes[1])
                    (c.validate_saved_pair if n==0 else part2.validate_saved_pair)(app,standard,application,p)
                    auto.validate_pair(app,act,application,p,proof,*shapes)
                    detail=inspect_application(app,application,folder/'print.pdf',p,template)
                    results.append(dict(variant=n,count=count,status='PASS',**detail))
                    print(f'PASS AKZ={n} rows={count} pages={detail["pages"]}',flush=True)
            except Exception:
                result['error']=traceback.format_exc()
          del app
          gc.collect()
        if 'error' in result:
            raise AssertionError(result['error'])
        result['status']='PASS'
    except Exception:
        result['status']='FAIL';result['error']=traceback.format_exc()
        print(result['error'],flush=True)
    finally:
        after={p:fingerprint(Path(p)) for p in before}
        result['protected_files']=len(before)
        result['protected_changes']=[p for p in before if before[p]!=after[p]]
        if result['protected_changes']:result['status']='FAIL'
        save(out/'AFTER.json',after);save(out/'RESULT.json',result)
    return 0 if result['status']=='PASS' else 1

if __name__=='__main__':
    sys.exit(main())
