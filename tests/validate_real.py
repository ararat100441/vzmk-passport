"""Acceptance on disposable copies of audited workbooks; never writes production files."""
import hashlib
import json
import shutil
import sys
import time
from datetime import date, datetime
from pathlib import Path
from collections import Counter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from passport_app.config import Settings, PassportError
from passport_app.excel import excel_session, open_excel, sha256, locked_file
from passport_app.registry import read_registry, select_project
from passport_app.service import validate, generate
from passport_app.vom import read_vom

BASE = Path(__file__).resolve().parents[1]
PRODUCTION = Settings.load(BASE / 'settings.json')
AUDIT = json.loads((BASE / 'audit' / 'audit_inventory.json').read_text(encoding='utf-8'))
SOURCE_VOM = Path(AUDIT['vom']['path'])
SOURCE_TEMPLATE = PRODUCTION.documents_root / PRODUCTION.projects['Гараж Казахстан']['template']
OUT = BASE / 'validation' / datetime.now().strftime('%Y%m%d_%H%M%S')
CHECKS = []


def check(label, condition, detail=''):
    CHECKS.append({'check': label, 'pass': bool(condition), 'detail': detail})
    if not condition:
        raise AssertionError(label + ': ' + str(detail))


def copy_case(name):
    root = OUT / name
    docs = root / 'Документы'
    passports = docs / 'Гараж казахстан' / 'Паспорт'
    passports.mkdir(parents=True)
    registry = root / 'ВЗМК.xlsx'
    shutil.copy2(PRODUCTION.registry_path, registry)
    template = passports / SOURCE_TEMPLATE.name
    shutil.copy2(SOURCE_TEMPLATE, template)
    vom = root / 'ВОМ отгрузочный.xlsx'
    shutil.copy2(SOURCE_VOM, vom)
    settings = Settings(docs, registry, PRODUCTION.projects)
    (root / 'settings.json').write_text(json.dumps({'documents_root': 'Документы', 'registry_path': 'ВЗМК.xlsx',
                                                   'projects': PRODUCTION.projects}, ensure_ascii=False, indent=2), encoding='utf-8')
    return settings, vom, template


def picture(excel, path, region, output, sheet_index=1):
    # PNGs are rendered from saved XLSX, not from values rebuilt into another table.
    # Native Excel save/recalculation is independently checked by this test.
    import subprocess
    if __import__('os').environ.get('PASSPORT_NO_PREVIEWS'):
        return
    node = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
    completed = subprocess.run([str(node), str(BASE / 'tests/render_excel.mjs'), str(path), region, str(output)],
                               capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=90)
    if completed.returncode:
        raise AssertionError('Не удалось отрисовать сохранённый XLSX: ' + completed.stderr[-1500:])


def formula_normal(value):
    if isinstance(value, str) and value.startswith('='):
        import re
        return re.sub(r'\[\[#This Row\],\[([^]]+)\]\]', r'[@[\1]]', value)
    return value


def validate_output(settings, vom_path, template, result, count):
    incoming = read_vom(vom_path)
    wb = openpyxl.load_workbook(settings.registry_path, data_only=False)
    values = openpyxl.load_workbook(settings.registry_path, data_only=True)
    original = openpyxl.load_workbook(PRODUCTION.registry_path, data_only=False)
    try:
        s, v, old = wb.worksheets[0], values.worksheets[0], original.worksheets[0]
        row = result['row']
        check(f'{count}: номер +1', s.cell(row, 2).value == '9/2026')
        check(f'{count}: текущая дата Windows', v.cell(row, 9).value.date() == date.today())
        check(f'{count}: масса из ВОМ', abs(v.cell(row, 7).value - float(incoming.mass)) < 1e-8)
        check(f'{count}: соединения из ВОМ', s.cell(row, 16).value == incoming.connections)
        changed = {2, 7, 9, 16}
        check(f'{count}: постоянные данные первого паспорта', all(
            formula_normal(s.cell(row, c).value) == formula_normal(old.cell(353, c).value)
            for c in range(1, 18) if c not in changed))
        check(f'{count}: договор взят из первого, а не последнего', s.cell(row, 3).value == old['C353'].value and s.cell(row, 3).value != old['C396'].value)
        check(f'{count}: формулы M:N:O сохранены', all(s.cell(row, c).data_type == 'f' and
                                                       formula_normal(s.cell(row, c).value) == formula_normal(old.cell(353, c).value)
                                                       for c in (13, 14, 15)))
        check(f'{count}: формулы рассчитаны Excel', all(v.cell(row, c).value.date() == date.today() - __import__('datetime').timedelta(days=offset)
                                                       for c, offset in [(13, 7), (14, 5), (15, 1)]))
        check(f'{count}: формат новой строки', all(s.cell(row, c)._style == s.cell(353, c)._style for c in range(1, 18)))
        check(f'{count}: высота строки', s.row_dimensions[row].height == s.row_dimensions[353].height)
        check(f'{count}: таблица Excel расширена', s.tables['Таблица1'].ref == 'A1:P400')
        unchanged = []
        for old_sheet in original:
            new_sheet = wb[old_sheet.title]
            for cells in old_sheet:
                for cell in cells:
                    if formula_normal(cell.value) != formula_normal(new_sheet[cell.coordinate].value):
                        unchanged.append(f'{old_sheet.title}!{cell.coordinate}')
        check(f'{count}: существующие данные всех листов реестра сохранены', not unchanged, unchanged[:10])
    finally:
        wb.close(); values.close(); original.close()
    folder = Path(result['folder'])
    expected_folder = f'9-2026 от {date.today():%d.%m.%Y}'
    check(f'{count}: имя папки', folder.name == expected_folder)
    check(f'{count}: только Приложение №1', [p.name for p in folder.iterdir()] == ['9_2026 - Приложение № 1.xlsx'])
    book = openpyxl.load_workbook(result['appendix'], data_only=False)
    book_values = openpyxl.load_workbook(result['appendix'], data_only=True)
    src = openpyxl.load_workbook(template, data_only=False)
    try:
        a, av, reference = book.worksheets[0], book_values.worksheets[0], src.worksheets[0]
        total = count + 5
        check(f'{count}: ровно строк ВОМ', a.max_row == total and a.cell(total, 1).value == 'Итого')
        rows = list(av.iter_rows(min_row=5, max_row=total - 1, max_col=7, values_only=True))
        expected = [tuple(float(x) if i >= 2 else x for i, x in enumerate(r)) for r in incoming.rows]
        check(f'{count}: порядок и все исходные значения', rows == expected)
        check(f'{count}: повторяющиеся марки не объединены', Counter(r[0] for r in rows) == Counter(r[0] for r in expected))
        check(f'{count}: нет пустых строк перед Итого', bool(a.cell(total - 1, 1).value))
        check(f'{count}: вся строка Итого жирная', all(a.cell(total, c).font.bold for c in range(1, 8)))
        check(f'{count}: суммы количества, массы, площади', all(abs(av.cell(total, c).value - float(e)) < 1e-8
                                                                for c, e in [(3, incoming.quantity), (5, incoming.mass), (7, incoming.area)]))
        check(f'{count}: итоговые формулы', all(a.cell(total, c).data_type == 'f' for c in (3, 5, 7)))
        check(f'{count}: номер и дата шапки из новой строки', f'№ 9/2026 от {date.today():%d.%m.%Y} г.' in a['A1'].value)
        check(f'{count}: объект шапки из новой строки', a['A2'].value == 'Объект: ' + read_registry(settings.registry_path).projects['Гараж Казахстан'][-1]['object'])
        check(f'{count}: область печати', str(a.print_area) == f"'Приложение №1'!$A$1:$G${total}")
        check(f'{count}: объединения шапки', set(map(str, a.merged_cells.ranges)) == set(map(str, reference.merged_cells.ranges)))
        check(f'{count}: ширины столбцов сохранены', {k: x.width for k, x in a.column_dimensions.items()} == {k: x.width for k, x in reference.column_dimensions.items()})
        check(f'{count}: высоты строк данных', all(a.row_dimensions[r].height == reference.row_dimensions[5].height for r in range(5, total)))
        check(f'{count}: параметры печати', a.page_setup == reference.page_setup and a.page_margins == reference.page_margins and a.print_options == reference.print_options)
    finally:
        book.close(); book_values.close(); src.close()
    check(f'{count}: первый паспорт-файл не изменён', sha256(template) == sha256(SOURCE_TEMPLATE))
    check(f'{count}: резервная копия исходного реестра', sha256(result['backup']) == sha256(PRODUCTION.registry_path))


def derive_vom(excel, path, count):
    book = open_excel(excel, path)
    try:
        sheet = book.Worksheets(1)
        if count < 32:
            sheet.Rows(f'{4 + count}:35').Delete()
        else:
            blocks, remainder = divmod(count, 32)
            for block in range(1, blocks):
                start = 4 + block * 32
                sheet.Range('A4:H35').Copy(Destination=sheet.Range(f'A{start}:H{start + 31}'))
            if remainder:
                start = 4 + blocks * 32
                sheet.Range(f'A4:H{3 + remainder}').Copy(Destination=sheet.Range(f'A{start}:H{start + remainder - 1}'))
            sheet.Rows(f'36:{3 + count}').RowHeight = sheet.Rows(4).RowHeight
        excel.CalculateFull()
        book.Save()
    finally:
        book.Close(SaveChanges=False)


def expect_error(label, action, contains=None):
    try:
        action()
    except PassportError as exc:
        check(label, contains is None or contains in str(exc), str(exc))
    else:
        check(label, False, 'ожидалась ошибка')


def main():
    import faulthandler
    faulthandler.dump_traceback_later(45, repeat=True)
    sources = [PRODUCTION.registry_path, SOURCE_VOM, SOURCE_TEMPLATE]
    initial = {str(p): sha256(p) for p in sources}
    check('Checkpoint аудита совпадает', all(initial[str(p)] == AUDIT['hashes_after'][str(p)] for p in sources))
    OUT.mkdir(parents=True)
    settings, vom, template = copy_case('main_32')
    images = OUT / 'images'
    images.mkdir()
    reuse = __import__('os').environ.get('PASSPORT_REUSE_SOURCE_PREVIEWS')
    if reuse:
        for name in ('01_template.png', '02_registry_first_left.png', '03_registry_first_right.png'):
            shutil.copy2(Path(reuse) / name, images / name)
    else:
        with excel_session() as excel:
            picture(excel, template, 'A1:G45', images / '01_template.png')
            picture(excel, settings.registry_path, 'A353:G353', images / '02_registry_first_left.png')
            picture(excel, settings.registry_path, 'H353:P353', images / '03_registry_first_right.png')
    file_snapshot = {str(p): sha256(p) for p in (settings.registry_path, vom, template)}
    paths_before = set(settings.documents_root.rglob('*'))
    plan = validate(settings, 'Гараж Казахстан', vom)
    check('Проверить не создаёт паспорт/файлы', paths_before == set(settings.documents_root.rglob('*')) and file_snapshot == {str(p): sha256(p) for p in (settings.registry_path, vom, template)})
    check('Контрольные значения ВОМ', len(plan.vom.rows) == 32 and str(plan.vom.quantity) == '50' and str(plan.vom.mass) == '15071.44' and str(plan.vom.area) == '477.749' and plan.vom.connections == 'Балка')
    result = generate(settings, plan)
    print('Generated primary 32', flush=True)
    validate_output(settings, vom, template, result, 32)
    with excel_session() as excel:
        picture(excel, result['appendix'], 'A1:G37', images / '04_result_32.png')
        picture(excel, settings.registry_path, 'A400:G400', images / '05_registry_new_left.png')
        picture(excel, settings.registry_path, 'H400:P400', images / '06_registry_new_right.png')
    # Existing row makes the next run 10/2026; collisions are tested on an isolated copy.
    check('Повторное формирование предлагает 10/2026', validate(settings, 'Гараж Казахстан', vom).project.next_number == '10/2026')
    for count in (7, 100):
        print(f'Start dynamic {count}', flush=True)
        setting, v, t = copy_case(f'dynamic_{count}')
        fixtures = __import__('os').environ.get('PASSPORT_DERIVED_FIXTURES')
        if fixtures:
            shutil.copy2(Path(fixtures) / f'dynamic_{count}' / 'ВОМ отгрузочный.xlsx', v)
        else:
            with excel_session() as excel:
                derive_vom(excel, v, count)
        print(f'Derived VOM {count}', flush=True)
        p = validate(setting, 'Гараж Казахстан', v)
        r = generate(setting, p)
        print(f'Generated dynamic {count}', flush=True)
        validate_output(setting, v, t, r, count)
        with excel_session() as excel:
            picture(excel, r['appendix'], 'A1:G12' if count == 7 else 'A96:G105', images / f'07_result_{count}.png')
    safe, v, t = copy_case('safety')
    p = validate(safe, 'Гараж Казахстан', v)
    p.folder.mkdir()
    expect_error('Конфликт существующей папки', lambda: validate(safe, 'Гараж Казахстан', v), 'уже существует')
    # Preserve even the disposable user-like conflicting directory. Use another case.
    safe, v, t = copy_case('failure_before_commit')
    p = validate(safe, 'Гараж Казахстан', v)
    before = sha256(safe.registry_path)
    expect_error('Отказ до публикации', lambda: generate(safe, p, fail_before_commit=True), 'Тестовая остановка')
    check('При отказе исходный реестр и целевая папка не изменены', sha256(safe.registry_path) == before and not p.folder.exists())
    safe, v, t = copy_case('changed_input')
    p = validate(safe, 'Гараж Казахстан', v)
    with v.open('ab') as stream:
        stream.write(b'changed-after-check')
    expect_error('Изменённый ВОМ после проверки', lambda: generate(safe, p), 'ВОМ изменился')
    check('Изменённый ВОМ не привёл к записи', sha256(safe.registry_path) == initial[str(PRODUCTION.registry_path)] and not p.folder.exists())
    safe, v, t = copy_case('locked_registry')
    with locked_file(safe.registry_path, writable=True):
        expect_error('Заблокированный реестр', lambda: validate(safe, 'Гараж Казахстан', v))
    registry = read_registry(safe.registry_path)
    expect_error('Нестандартный номер не угадывается', lambda: select_project(registry, 'ТЭМПО 2'), 'не соответствует')
    expect_error('Отсутствие первого паспорта не угадывается', lambda: select_project(registry, 'Порт альянс 2,7'), 'первый паспорт')
    check('Производственные файлы неизменны', all(sha256(Path(p)) == value for p, value in initial.items()))
    (OUT / 'results.json').write_text(json.dumps({'checks': CHECKS, 'main_result': result,
                                                   'sources_sha256': initial, 'visual_review': 'PENDING'}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(OUT), 'checks': len(CHECKS), 'pass': all(c['pass'] for c in CHECKS), 'main_result': result}, ensure_ascii=False))
    faulthandler.cancel_dump_traceback_later()


if __name__ == '__main__':
    try:
        main()
    finally:
        if OUT.exists():
            (OUT / 'checks_partial.json').write_text(json.dumps(CHECKS, ensure_ascii=False, indent=2), encoding='utf-8')
