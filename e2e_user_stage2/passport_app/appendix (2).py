import re
from pathlib import Path
from openpyxl.utils import get_column_letter
from .config import PassportError
from .excel import load_book, open_excel, sha256
from .tables import locate_layout, norm
from .vom import FIELDS


def same_object(value, expected):
    value = re.sub(r'^\s*Объект\s*:\s*', '', str(value or ''), flags=re.I)
    # Only comparison normalizes spaces. The generated object uses the new registry value.
    return ' '.join(value.split()) == ' '.join(expected.split())


def inspect_template(path, project):
    original_digest = sha256(path)
    book = load_book(path)
    try:
        layout = locate_layout(book, appendix=True)
        sheet = book[layout.sheet]
        if sheet.protection.sheet or len(book.worksheets) != 1:
            raise PassportError('Шаблон должен содержать один незапрещённый для записи лист Приложения №1.')
        headers, objects = [], []
        for row in sheet.iter_rows(max_row=layout.header_row - 1):
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                if re.search(r'№\s*\S+\s+от\s+\d{2}\.\d{2}\.\d{4}\s*г\.', cell.value):
                    headers.append(cell)
                if re.match(r'^\s*Объект\s*:', cell.value, re.I):
                    objects.append(cell)
        if len(headers) != 1 or len(objects) != 1:
            raise PassportError('В шаблоне не найдена однозначная шапка с номером, датой и объектом.')
        match = re.search(r'№\s*(\S+)\s+от\s+(\d{2}\.\d{2}\.\d{4})\s*г\.', headers[0].value)
        if match[1] != project.first_number or match[2] != project.first_date.strftime('%d.%m.%Y'):
            raise PassportError('Шаблон не соответствует номеру и дате ПЕРВОГО паспорта проекта.')
        if not same_object(objects[0].value, project.object_name):
            raise PassportError('Объект в шаблоне не соответствует первой строке проекта.')
        if layout.total_row <= layout.start:
            raise PassportError('В шаблоне нет строки данных для копирования оформления.')
        if len(set(layout.columns.values())) != 7:
            raise PassportError('Неоднозначные столбцы приложения.')
        if sheet.max_row != layout.total_row:
            raise PassportError('Ниже «Итого» есть дополнительные строки. Требуется отдельная проверка шаблона.')
        for m in sheet.merged_cells.ranges:
            if m.min_row >= layout.start and m.max_row > m.min_row:
                raise PassportError('В таблице приложения есть вертикальные объединения. Безопасное изменение длины не подтверждено.')
            if layout.start <= m.min_row < layout.total_row and sum(m.min_col <= col <= m.max_col for col in layout.columns.values()) > 1:
                raise PassportError('В строке данных шаблона объединены разные поля таблицы.')
        if sha256(path) != original_digest:
            raise PassportError('Первое приложение изменилось во время чтения.')
        return layout, headers[0].coordinate, objects[0].coordinate
    finally:
        book.close()


def find_template(root, configured, project):
    if configured:
        inspect_template(configured, project)
        return configured
    matches = []
    for path in root.rglob('*.xlsx'):
        if path.name.startswith('~$') or not re.search(r'Приложение\s*№\s*1', path.name, re.I):
            continue
        try:
            inspect_template(path, project)
            matches.append(path)
        except PassportError:
            continue
    if len(matches) != 1:
        raise PassportError(f'Найдено шаблонов первого паспорта: {len(matches)}. Укажите точный template в settings.json.')
    return matches[0]


def generate_appendix(excel, path, project, vom, requisites):
    layout, header, obj = inspect_template(path, project)
    book = open_excel(excel, path)
    try:
        sheet = book.Worksheets(layout.sheet)
        original_area = sheet.PageSetup.PrintArea
        area_match = None
        if original_area:
            area_match = re.fullmatch(r'\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)', original_area)
            if not area_match or int(area_match[4]) != layout.total_row:
                raise PassportError('Сложная область печати шаблона: безопасное изменение не подтверждено.')
        original_count = layout.total_row - layout.start
        needed = len(vom.rows)
        if needed < original_count:
            sheet.Rows(f'{layout.start + needed}:{layout.total_row - 1}').Delete()
        elif needed > original_count:
            additional = needed - original_count
            sheet.Rows(f'{layout.total_row}:{layout.total_row + additional - 1}').Insert()
        total = layout.start + needed
        # Copy only a data row, never the heading or total. Excel adjusts native styles.
        mincol, maxcol = min(layout.columns.values()), max(layout.columns.values())
        row_height = sheet.Rows(layout.start).RowHeight
        row_hidden = sheet.Rows(layout.start).Hidden
        if needed > 1:
            # One native fill copies the first row's contents AND formatting to the body.
            # Data are then replaced below. Avoid an OS clipboard operation per mark.
            sheet.Range(sheet.Cells(layout.start, mincol), sheet.Cells(total - 1, maxcol)).FillDown()
            sheet.Rows(f'{layout.start + 1}:{total - 1}').RowHeight = row_height
            sheet.Rows(f'{layout.start + 1}:{total - 1}').Hidden = row_hidden
        for field_index, field in enumerate(FIELDS):
            col = layout.columns[field]
            sheet.Range(sheet.Cells(layout.start, col), sheet.Cells(total - 1, col)).Value2 = tuple(
                (float(row[field_index]) if field_index >= 2 else row[field_index],) for row in vom.rows)
        for field in ('quantity', 'mass_total', 'area_total'):
            col = layout.columns[field]
            letter = get_column_letter(col)
            sheet.Cells(total, col).Formula = f'=SUM({letter}{layout.start}:{letter}{total - 1})'
        sheet.Range(sheet.Cells(total, mincol), sheet.Cells(total, maxcol)).Font.Bold = True
        current = str(sheet.Range(header).Value2)
        replaced, count = re.subn(r'(№\s*)\S+(\s+от\s+)\d{2}\.\d{2}\.\d{4}(\s*г\.)',
                                  lambda m: m[1] + str(requisites['number']) + m[2] + requisites['date'].strftime('%d.%m.%Y') + m[3], current)
        if count != 1:
            raise PassportError('Не удалось безопасно заменить реквизиты шапки приложения.')
        sheet.Range(header).Value2 = replaced
        prefix = re.match(r'^(\s*Объект\s*:\s*)', str(sheet.Range(obj).Value2), re.I)[1]
        sheet.Range(obj).Value2 = prefix + str(requisites['object'])
        # Keep page settings and column widths; update the bottom of the print area.
        if area_match:
            sheet.PageSetup.PrintArea = f'${area_match[1]}${area_match[2]}:${area_match[3]}${total}'
        else:
            letter = get_column_letter(maxcol)
            sheet.PageSetup.PrintArea = f'$A$1:${letter}${total}'
        excel.CalculateFull()
        for field, expected, tolerance in [('quantity', vom.quantity, 0), ('mass_total', vom.mass, .01), ('area_total', vom.area, .001)]:
            actual = sheet.Cells(total, layout.columns[field]).Value2
            if not isinstance(actual, (float, int)) or abs(actual - float(expected)) > tolerance + 1e-8:
                raise PassportError(f'Итог приложения не совпадает с ВОМ: {field}.')
        book.Save()
        return total
    finally:
        book.Close(SaveChanges=False)
