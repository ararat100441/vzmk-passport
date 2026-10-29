"""Explicit accepted inputs; no latest-file discovery, production or auto-sync."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from .provenance import Stage7Error, sha, read, verify, digest

ROOT = Path(__file__).resolve().parents[1]
PROJECT = 'Гараж Казахстан'
NUMBER = '9/2026'
REGISTRY = ROOT / 'DEV/ВЗМК_DEV.xlsx'
PASSPORTS = ROOT.parent / 'Гараж казахстан/Паспорт'
SHIPMENT = PASSPORTS / '9 от 06.10.2026'
VOM = SHIPMENT / 'ВОМ отгрузочный счет 183  (5).xls'
APPENDIX = SHIPMENT / '9_2026 - Приложение № 1.xlsx'
TEMPLATE = PASSPORTS / '1_2026 - Документ о качестве.docx'
VIK_TEMPLATE = PASSPORTS / '1_2026 - ВИК.docx'
AKZ_ROOT = ROOT / 'validation/stage6_fix_r1_r2_20261009_094707'
PAIR = AKZ_ROOT / 'NEW_TEST_9_2026/PAIR.json'
AKZ = PAIR.parent / 'ТЕСТ АКЗ №0.xlsx'
AKZ_APP = PAIR.parent / 'ТЕСТ Приложение к АКЗ №0.xlsx'
PINS = {REGISTRY: '8e56f68518c70cade9efb09084dd67cd5d72e07e3fdc813b8515030e9cb98e80',
    VOM: 'aae5e5b8c0246c8fe97b3dec9cbcc039bc3e3cedb9d245b474b2f871b90596ae',
    APPENDIX: 'b1c6a950b932373e37b203ccff1cedfa5a11e528552a8e6e06692f666af63ff6',
    TEMPLATE: '4df7234d37030188030142ed6a519dc2a923aada4f90e58a7b5140c0fd54f9cc',
    AKZ: 'b640796570a8d04b6bca404205d99eae4c7005c700587812ed25bc99e715eb1a',
    AKZ_APP: '7766c92df82df2484eb82cd6917769f5d9ede912b3f0e8f26e3f8aaabdd357f2',
    ROOT / 'passport_app/vom.py': 'c63bbd362763ba99f85de1610c01622ac230f3303b7a1e297345d02259ea954f'}


def current_fields():
    from passport_app.registry import read_registry, select_project
    registry = read_registry(REGISTRY)
    select_project(registry, PROJECT)
    entries = [e for e in registry.projects[PROJECT] if e['number'] == NUMBER]
    if len(entries) != 1:
        raise Stage7Error('CURRENT_PASSPORT_AMBIGUOUS')
    row = entries[0]['row']
    book = load_workbook(REGISTRY, read_only=True, data_only=True, keep_links=False)
    try:
        s = book[registry.sheet]
        values = [s.cell(row, c).value for c in range(2, 17)]
    finally:
        book.close()
    keys = ('number', 'contract', 'customer', 'object', 'cipher', 'mass', 'start_date', 'end_date',
            'km', 'kmd_company', 'steel', 'vik_date', 'uzk_date', 'akz_date', 'connections')
    fields = dict(zip(keys, values)) | {'row': row, 'sheet': registry.sheet, 'project': PROJECT}
    for k in ('start_date', 'end_date', 'vik_date', 'uzk_date', 'akz_date'):
        v = fields[k]
        if not isinstance(v, (date, datetime)):
            raise Stage7Error('REGISTRY_DATE_INVALID: ' + k)
        fields[k] = v.strftime('%d.%m.%Y')
    if (fields['number'] != NUMBER or Decimal(str(fields['mass'])) != Decimal('15822.14')
            or [fields[k] for k in ('end_date', 'vik_date', 'uzk_date', 'akz_date')]
            != ['06.10.2026', '29.09.2026', '01.10.2026', '05.10.2026']):
        raise Stage7Error('CURRENT_CONTRACT_CHANGED')
    for k in ('contract', 'customer', 'object', 'cipher', 'km', 'kmd_company', 'steel', 'connections'):
        if not isinstance(fields[k], str) or not fields[k].strip():
            raise Stage7Error('REGISTRY_TEXT_INVALID: ' + k)
    return fields


def aggregate(rows):
    result = defaultdict(lambda: [Decimal(0), Decimal(0)])
    for mark, quantity, mass in rows:
        key = str(mark).strip().casefold()
        result[key][0] += Decimal(str(quantity))
        result[key][1] += Decimal(str(mass))
    return dict(result)


def validate_appendix(vom, fields):
    book = load_workbook(APPENDIX, read_only=True, data_only=True, keep_links=False)
    try:
        s = book.active
        if NUMBER not in str(s['A1'].value) or '06.10.2026' not in str(s['A1'].value) or fields['object'] not in str(s['A2'].value):
            raise Stage7Error('APPENDIX_MEMBERSHIP_MISMATCH')
        data = list(s.iter_rows(values_only=True))
        totals = [i for i, r in enumerate(data) if str(r[0] or '').strip().casefold().startswith('итого')]
        if len(totals) != 1:
            raise Stage7Error('APPENDIX_TOTAL_AMBIGUOUS')
        end = totals[0]
        body = [(r[0], r[2], r[4]) for r in data[4:end] if r[0] is not None]
        expected = aggregate((r['mark'], r['quantity'], r['mass_kg']) for r in vom['rows'])
        actual = aggregate(body)
        if set(actual) != set(expected) or any(actual[k][0] != expected[k][0] or abs(actual[k][1] - expected[k][1]) > Decimal('.005') for k in expected):
            raise Stage7Error('APPENDIX_VOM_PER_MARK_MISMATCH')
        if Decimal(str(data[end][2])) != 288 or abs(Decimal(str(data[end][4])) - Decimal('15822.14')) > Decimal('.005'):
            raise Stage7Error('APPENDIX_TOTAL_MISMATCH')
        return {'status': 'PASS', 'body_rows': len(body), 'total_row': end + 1,
                'visible_vom_marks': len(expected), 'quantity': 288, 'mass_kg': '15822.14', 'per_mark_equal': True}
    finally:
        book.close()


def validate_akz(vom, fields):
    pair = read(PAIR)
    preview = read(AKZ_ROOT / 'current_preview.json')
    p = preview['payload']
    if (preview['seal'] != digest(p) or pair['act_id'] != preview['seal'][:16]
            or pair['state'] != 'COMPLETE_TEST_PAIR' or pair['number'] != NUMBER
            or pair['date'] != '2026-10-05' or pair['synthetic'] is not False):
        raise Stage7Error('AKZ_PAIR_PROVENANCE_MISMATCH')
    verify(p['sources'])
    verify({str(PAIR.parent / n): v for n, v in pair['sha256'].items()})
    if p['sources'].get(str(REGISTRY)) != PINS[REGISTRY] or p['sources'].get(str(VOM)) != PINS[VOM]:
        raise Stage7Error('AKZ_SOURCE_CYCLE_MISMATCH')
    book = load_workbook(AKZ, read_only=True, data_only=True, keep_links=False)
    try:
        s = book.active
        if (str(s['E4'].value).strip() != '9/2026 от' or s['H4'].value.strftime('%d.%m.%Y') != fields['akz_date']
                or s['E17'].value != 1 or 'Pirocor Protect 057' not in str(s['E18'].value)
                or s['E21'].value != 'RAL 7004' or s['E25'].value != 80
                or 'ТЕСТ' not in str(s['B3'].value) or 'Выдача запрещена' not in str(s['B43'].value)):
            raise Stage7Error('AKZ_POINT12_OR_CONTENT_CONFLICT')
        if any('НЕТ ДАННЫХ / ТЕСТ' not in str(s[cell].value) for cell in ('E26', 'E27', 'E28')):
            raise Stage7Error('AKZ_UNCONFIRMED_STATUS_CHANGED')
    finally:
        book.close()
    book = load_workbook(AKZ_APP, read_only=True, data_only=True, keep_links=False)
    try:
        s = book.active
        if 'Приложение № 9/2026 к акту АКЗ № 9/2026 от 05.10.2026' not in str(s['A1'].value) or 'ТЕСТ' not in str(s['A1'].value):
            raise Stage7Error('AKZ_APPLICATION_PARENT_MISMATCH')
        body = [(s.cell(i, 1).value, s.cell(i, 3).value, s.cell(i, 5).value) for i in (4, 5, 6)]
        available = aggregate((r['mark'], r['quantity'], r['mass_kg']) for r in vom['rows'])
        a = aggregate(body)
        if set(a) != {'бп2-5', 'бп2-18', 'бп2-19'} or any(a[k][0]!=available[k][0] or abs(a[k][1]-available[k][1])>Decimal('.005') for k in a) or s['C21'].value != 3:
            raise Stage7Error('AKZ_SUBSET_MEMBERSHIP_MISMATCH')
        return {'status': 'PASS', 'pair_id': pair['act_id'], 'preview_seal': preview['seal'],
                'variant': 0, 'assigned_marks': [r[0] for r in body], 'quantity': 3,
                'shipment_quantity': 288, 'actual_measurements_confirmed': False,
                'scope': 'THREE_ASSIGNED_MARKS_ONLY', 'point12_compatible': True}
    finally:
        book.close()


def validate_uzk_plan(plan):
    expected = {'Бп2-1': (8, [('1280', '5'), ('1024.0', '8.4')]),
                'Пс1-9': (9, [('3600', '4')]), 'Пс1-10': (15, [('6000', '4')]), 'Пс2-16': (1, [('480', '3')])}
    marks = plan.data['marks']
    if len(marks) != 4 or plan.data['controlled_units'] != 33 or plan.data['passport']['date'] != '01.10.2026':
        raise Stage7Error('UZK_PLAN_CONTRACT_MISMATCH')
    for m in marks:
        if m['mark'] not in expected or m['Q'] != expected[m['mark']][0]:
            raise Stage7Error('UZK_PLAN_MARK_MISMATCH')
        segments = m['segments']
        e = expected[m['mark']][1]
        if len(e) != len(segments) or any(Decimal(str(s['L'])) != Decimal(v[0]) or Decimal(str(s['S'])) != Decimal(v[1])
                or Decimal(str(s['frequency_mhz'])) != 5 or Decimal(str(s['angle_degrees'])) != 70 for s, v in zip(segments, e)):
            raise Stage7Error('UZK_PLAN_GEOMETRY_MISMATCH')
    return {'status': 'PASS', 'marks': 4, 'quantity': 33, 'date': '01.10.2026', 'mode': 'SAMPLE_ONLY'}
