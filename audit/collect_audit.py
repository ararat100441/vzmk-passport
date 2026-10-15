"""Read-only Stage 1 inventory. Never saves an Excel workbook."""
from pathlib import Path
from collections import defaultdict, Counter
from decimal import Decimal
import hashlib
import json
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
REGISTRY = ROOT.parent / 'ВЗМК.xlsx'
OUT = Path(__file__).resolve().parent

def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def normal(v):
    return ' '.join(str(v or '').split()).casefold().rstrip(' .')

files = [p for p in ROOT.rglob('*.xlsx') if not p.name.startswith('~$')]
templates = [p for p in files if p.name == '1_2026 - Приложение № 1.xlsx']
garage = ROOT / 'Гараж казахстан' / 'Паспорт'
vom = garage / '8 от 02.10.2026' / 'ВОМ отгрузочный.xlsx'
app = garage / '1_2026 - Приложение № 1.xlsx'
tracked = [REGISTRY, vom, *templates]
before = {str(p): digest(p) for p in tracked}
w = openpyxl.load_workbook(REGISTRY, data_only=False)
s = w.worksheets[0]
headers = {c.column_letter: c.value for c in s[1] if c.value is not None}
projects = defaultdict(list)
for row in s.iter_rows(min_row=2, max_col=16):
    if row[0].value:
        projects[str(row[0].value)].append({
            'row': row[0].row, 'number': row[1].value,
            'object': row[4].value, 'mass': row[6].value,
            'date': str(row[8].value), 'connections': row[15].value,
        })
template_info = []
for p in templates:
    book = openpyxl.load_workbook(p, read_only=True, data_only=False)
    first = book.worksheets[0]
    template_info.append({'path': str(p), 'header': str(first['A1'].value),
                          'object_header': str(first['A2'].value)})
    book.close()
lines = [
    'ВЗМК — аудит Этапа 1, 05.10.2026',
    'СТАТУС: ОСТАНОВКА ДО РЕАЛИЗАЦИИ. Источник поля «Соединения» требует уточнения.',
    '', 'A. Реестр и структура', str(REGISTRY),
    f'Лист: {s.title}; таблица: {[(t.name, t.ref) for t in s.tables.values()]}.',
    f'Непустых строк проектов: {sum(map(len, projects.values()))}; уникальных проектов: {len(projects)}.',
    'Столбцы: ' + '; '.join(f'{k}: {v}' for k, v in headers.items()),
    'M, N, O — существующие формулы дат, а не поля для собственной логики программы.',
    f'M353: {s["M353"].value}', f'N353: {s["N353"].value}', f'O353: {s["O353"].value}',
    f'I353: числовой формат {s["I353"].number_format}; G353: {s["G353"].number_format}; B353: {s["B353"].number_format}.',
    '', 'Гараж Казахстан: первый 1/2026, строка 353, дата 09.09.2026.',
    'Последний 8/2026, строка 396, дата 02.10.2026. Следующий 9/2026.',
    f'Первое приложение: {app}',
    'Первое приложение лежит непосредственно в папке «Паспорт». Папки «1-2026 от 09.09.2026» нет.',
    'Отличие постоянных данных: C353 содержит «№ 06/26 от 03.06.2026», C396 — «06/26 от 03.06.2026».',
    'При последующей реализации постоянные данные должны копироваться из строки 353.',
]
a = openpyxl.load_workbook(app, data_only=False).worksheets[0]
lines.extend([
    f'Шаблон: лист {a.title}; шапка A1:G1 и A2:G2; заголовки строки 3–4; данные 5–44 (40 строк); Итого 45.',
    'Данные приложения: A марка, B наименование, C количество, D масса 1 шт., E общая масса, F площадь 1 шт., G общая площадь.',
    f'Объединения: {a.merged_cells}. Область печати: {a.print_area}.',
    'Итого: C45=SUM(C5:C44), E45=SUM(E5:E44), G45=SUM(G5:G44).',
    '', 'Отгрузочный ВОМ', str(vom),
    'Заголовки определены по тексту и объединениям строк 2–3.',
    'A: Лист №; B: Отправочная марка; C: Наименование марки; D: Количество;',
    'E/F: Масса (изделие), кг — 1 дет./Всех; G/H: Площадь окраски, м2 — 1 дет./Всех.',
])
vf = openpyxl.load_workbook(vom, data_only=False)
vv = openpyxl.load_workbook(vom, data_only=True)
vs = vv.worksheets[0]
rows = list(vs.iter_rows(min_row=4, max_col=8, values_only=True))
totals = {name: str(sum(Decimal(str(r[col])) for r in rows))
          for name, col in [('quantity', 3), ('mass', 5), ('area', 7)]}
lines.extend([
    f'Лист ВОМ: {vs.title}; данные строки 4–35; строк {len(rows)}.',
    f'Изделий: {totals["quantity"]}; масса: {totals["mass"]} кг; площадь: {totals["area"]} м².',
    'Отдельной строки «Итого» в этом ВОМ нет. Итоги сверены суммированием исходных строк.',
    'F4:F35 — формулы E*D; сохранённые значения имеются. H — исходные числовые значения.',
    'Площадь «Всех» нельзя заменять произведением округлённой площади 1 шт.: G10=7.349, D10=2, H10=14.697.',
    f'Повторяющиеся марки: {[(k, v) for k, v in Counter(r[1] for r in rows).items() if v > 1]}.',
    '', 'Неоднозначность «Соединения»',
    'В ВОМ нет отдельного поля/заголовка «Соединения».',
    f'Кандидат: C4:C35 «Наименование марки»; уникальные значения: {list(dict.fromkeys(r[2] for r in rows))}.',
    'Нужно подтвердить правило: брать уникальные непустые наименования в порядке первого появления и соединять через «, ».',
    'Это только кандидат правила. Автоматически оно не принято.',
    'Второй возможный источник — отдельный реквизит другого отгрузочного ВОМ; в проверенном файле такого реквизита нет.',
    'В P396 реестра: «Колонна, Прогон, Балка, Воротный ригель, Прогон стеновой». Масса G396: 16541.88 кг.',
    'Проверенный ВОМ имеет только «Балка» и массу 15071.44 кг. Совпадение расположения с папкой №8 не доказывает, что это ВОМ паспорта №8.',
    'Уточнение: использовать найденный ВОМ как тестовый вход независимо от существующего паспорта №8 либо указать другой тестовый ВОМ.',
    '', 'Проекты и первые строки реестра',
    'Для обычных номеров первым кандидатом является единственная строка 1/год. Для прочих форматов правило не выбиралось.',
    'Связь с приложением ниже: точное совпадение объекта, номера и даты в шапке. Это кандидаты, не утверждённая карта папок.',
])
inventory = []
for name, entries in projects.items():
    first = [e for e in entries if re.fullmatch(r'1/\d{4}', str(e['number']))]
    invalid = [e['number'] for e in entries if not re.fullmatch(r'[1-9]\d*/\d{4}', str(e['number']))]
    candidates = []
    if len(first) == 1:
        e = first[0]
        date = e['date'][:10].split('-')
        date = '.'.join(reversed(date))
        for t in template_info:
            obj = re.sub(r'^Объект\s*:\s*', '', t['object_header'], flags=re.I)
            if normal(obj) == normal(e['object']) and f'№ {e["number"]} ' in t['header'] and date in t['header']:
                candidates.append(t['path'])
    item = {'project': name, 'entries': entries, 'first_candidates': first,
            'nonstandard_numbers': invalid, 'template_candidates': candidates}
    inventory.append(item)
    first_text = ', '.join(f'{e["number"]} (строка {e["row"]})' for e in first) or 'однозначно не определён'
    lines.append(f'{name!r}: строк {len(entries)}; первый {first_text}; номера: {", ".join(str(e["number"]) for e in entries)}')
    lines.append('  Приложения-кандидаты: ' + ('; '.join(candidates) or 'точное совпадение не найдено в 37 файлах «1_2026 - Приложение № 1.xlsx»'))
    if invalid:
        lines.append('  Нестандартные номера: ' + ', '.join(map(str, invalid)))
lines.extend([
    '', 'Ограничения поиска',
    'Для ряда проектов наименование в A не совпадает с папкой: например «Титул 6 КМ» и «НАВИГАТОР Титул 6 склад хим\\Паспорт\\км\\1».',
    'В A встречается символ /, недопустимый в имени одной папки Windows: «Русоко Б2/Метал инвест».',
    'Нужна явная карта путей для несовпадающих проектов. Автоматический нечёткий выбор папок не принят.',
    'Некоторые проекты имеют номера с дополнительными частями, буквами, нулём или без 1/год. Их нельзя считать поддержанными обычным правилом.',
    'Аудит первых паспортов всех проектов не завершён: точные соответствия отсутствующих/нестандартных шаблонов требуют уточнения.',
    '', 'B. Созданные файлы',
    'Только отчёт аудита, JSON с найденными строками/кандидатами/хешами и диагностический скрипт чтения.',
    'Прикладные модули и GUI не создавались: действует условие остановки из постановки задачи.',
    '', 'C. Запуск Этапа 1',
    'Недоступен до уточнения источника «Соединения» и реализации программы.',
    '', 'D. Реальные файлы', str(REGISTRY), str(app), str(vom),
    'Также прочитаны шапки 37 найденных файлов «1_2026 - Приложение № 1.xlsx».',
    '', 'E. Проверки',
    'Выполнено только чтение и сверка исходных данных. Генерация на копиях и 14 приёмочных проверок: НЕ ВЫПОЛНЕНЫ.',
    'Визуальная проверка итогового Excel: НЕ ВЫПОЛНЕНА, итоговый файл не создавался.',
    'Рабочие книги не сохранялись, исходный ВОМ и первые паспорта не изменялись. PDF не создавались.',
    '', 'F. Продолжение',
    'Сначала уточнить «Соединения» и карту папок, затем продолжить только Этап 1. Этап 2 не начат.',
])
w.close()
vf.close()
vv.close()
after = {str(p): digest(p) for p in tracked}
unchanged = before == after
lines.append(f'Контроль SHA-256 до и после сбора отчёта: {"PASS" if unchanged else "FAIL"}, файлов {len(before)}.')
OUT.mkdir(exist_ok=True)
(OUT / 'Аудит Этапа 1.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8-sig')
(OUT / 'audit_inventory.json').write_text(json.dumps({
    'registry': str(REGISTRY), 'headers': headers, 'projects': inventory,
    'templates': template_info, 'vom': {'path': str(vom), 'totals': totals, 'rows': len(rows)},
    'hashes_before': before, 'hashes_after': after, 'unchanged': unchanged,
}, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(json.dumps({'projects': len(projects), 'templates': len(templates),
                  'matched_projects': sum(bool(i['template_candidates']) for i in inventory),
                  'unchanged': unchanged, 'report': str(OUT / 'Аудит Этапа 1.txt')}, ensure_ascii=False))
