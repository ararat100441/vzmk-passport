"""Automatic DESIGN characteristics for UZK from current VOM and Stage4 profiles.

Separates WHAT must be inspected (mark, length, wall thickness, frequency and
angle according to approved profile geometry) from WHETHER a weld passed.

The current UZK.xlsm supplies only planned inspected item quantities. It does
not prove an actual inspection result. This module NEVER assigns defects,
conclusion 'ГОДЕН', welders, inspection date or an evidence/protocol number.

The single shared catalog is a read-only verified 'Профили.xlsx' built from
accepted Stage4 rules. Do not silently fall back to an unapproved geometry,
another standard, the first passport's historical measured outcome or DEV.
"""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from profiles_stage4.model import (
    normalize_profile,calculate_uzk,ProfileError,
    SUPPORTED_RULES,
)
from .storage import Stop,sha,safe

PROFILE_HEADINGS=(
    'Профиль','Ключ','Тип','НД','h','b','b2','s','t','D','r',
    'Источник','Дата получения','Статус','УЗК правило','Радиус R',
)
DIMENSIONS={'h':4,'b':5,'b2':6,'s':7,'t':8,'D':9,'r':10,'R':15}


def _decimal(value,what):
    try: n=Decimal(str(value))
    except Exception as exc:raise Stop('УЗК: неверное число '+what) from exc
    if not n.is_finite() or n<=0:
        raise Stop('УЗК: неверное положительное число '+what)
    return n


def _string_decimal(n):
    return format(Decimal(str(n)).normalize(),'f')


def _geometry_identity(profile):
    """All dimensions and the rule must agree, not just today's S/L result."""
    geometry=tuple(sorted((field,Decimal(str(value)))
                          for field,value in profile['geometry'].items()))
    if any(not value.is_finite() for _,value in geometry):
        raise ValueError('non-finite geometry')
    return profile['family'],profile['uzk_rule'],geometry


def _approved_rows(workbook):
    """Read only profile geometry and unambiguous user-approved rule rows."""
    if not workbook.is_file() or workbook.suffix.lower()!='.xlsx':
        raise Stop('УЗК: не найден проверенный справочник геометрии Профили.xlsx.')
    w=load_workbook(workbook,read_only=True,data_only=True,keep_links=False)
    try:
        if 'Профили' not in w.sheetnames:
            raise Stop('УЗК: в Профили.xlsx отсутствует лист «Профили».')
        sheet=w['Профили']
        data=sheet.iter_rows(values_only=True)
        headers=next(data,None)
        if headers is None or tuple(headers[:len(PROFILE_HEADINGS)])!=PROFILE_HEADINGS:
            raise Stop('УЗК: изменилась структура профилей Stage4.')
        index=defaultdict(list)
        for n,row in enumerate(data,2):
            if not row[1]:continue
            if len(row)<=15 or not isinstance(row[1],str):
                raise Stop('УЗК: некорректная строка справочника '+str(n))
            if row[14] not in SUPPORTED_RULES:
                continue  # No accepted UZK geometry rule for this profile.
            if row[13]=='КОНФЛИКТ':
                continue
            profile={
                'key':normalize_profile(row[1]),
                'family':row[2], 'standard':row[3],
                'uzk_rule':row[14],
                'geometry':{
                    field:row[col] for field,col in DIMENSIONS.items()
                    if row[col] is not None
                },
            }
            index[profile['key']].append((n,profile))
        return dict(index)
    finally:w.close()


def find_approved_geometry(cfg,*,require=False):
    """Only a verified current working or preinstalled copy is acceptable.

    No new user-facing path: accepted 'Профили.xlsx' may be next to the
    company working registry under Справочники, or inside installed
    clean-form storage. An explicit internal pin may specify a different
    approved copy. Ambiguous different matches require review.
    """
    candidates=[]
    pinned=cfg.get('profiles_workbook')
    if pinned:
        candidates=[safe(pinned)]
    else:
        if cfg.get('registry'):
            candidates.append(safe(Path(cfg['registry']).parent/
                                   'Справочники'/'Профили.xlsx'))
        if cfg.get('documents'):
            candidates.append(safe(Path(cfg['documents'])/'Профили.xlsx'))
    found=[p for p in candidates if p.is_file()]
    if not found:
        if require:
            raise Stop('УЗК: отсутствует утверждённый справочник геометрии '
                       '«Профили.xlsx» в рабочих Справочниках или во '
                       'внутреннем комплекте программы. '
                       'Справочник марок не содержит размеров профиля.')
        return None
    hashes={sha(p) for p in found}
    if len(hashes)!=1:
        raise Stop('УЗК: найдено несколько различных справочников '
                   '«Профили.xlsx». Выберите один проверенный источник.')
    path=found[0]
    expected=cfg.get('profiles_workbook_sha256')
    if expected and sha(path)!=expected:
        raise Stop('УЗК: контрольная сумма Профили.xlsx изменилась.')
    return path


def calculate_design_characteristics(vom,facts,profile_workbook,*,expected_sha=None):
    """For every actually selected UZK mark: rows of independent S/L.

    No summation between marks. Repeated VOM rows of SAME mark are already
    aggregated into a single inspected quantity by the shared UZK system.
    The Stage4 rule calculates per-item joined portions by equal thickness,
    and this module multiplies each per-item length by the inspected count.
    """
    path=safe(profile_workbook)
    before=sha(path)
    if expected_sha and before!=expected_sha:
        raise Stop('УЗК: справочник геометрии изменился после выбора.')
    known=_approved_rows(path)
    mark_facts={}
    for item in facts.get('marks',[]):
        mark=str(item.get('mark') or '').strip()
        if not mark or mark in mark_facts:
            raise Stop('УЗК: отсутствует/повторяется отправочная марка.')
        mark_facts[mark]=item
    expected={str(r[0]).strip() for r in vom.rows}
    if set(mark_facts)!=expected:
        raise Stop('УЗК: назначенные марки отличаются от текущей ВОМ.')
    output=[]
    for mark in dict.fromkeys(str(r[0]).strip() for r in vom.rows):
        item=mark_facts[mark]
        quantity=Decimal(str(item.get('uzk_quantity',0)))
        if quantity!=quantity.to_integral_value() or quantity<0:
            raise Stop('УЗК: неверное количество изделий по марке '+mark)
        if quantity==0:
            continue
        assigned=item.get('profile')
        if not isinstance(assigned,str) or not assigned.strip():
            raise Stop('УЗК: у выбранной марки '+mark+
                       ' не указан профиль в справочнике проекта.')
        profile_key=normalize_profile(assigned)
        candidates=known.get(profile_key,[])
        # Stage4 accepted channel alias: [16П => 16П, but ONLY when its
        # profile is the unique unambiguous channel identity.
        if not candidates and profile_key.startswith('['):
            alt=known.get(profile_key[1:],[])
            if alt and all(p['family']=='Швеллер' for _,p in alt):
                candidates=alt
        standard=str(item.get('standard') or '').strip()
        if standard:
            candidates=[(line,p) for line,p in candidates if p['standard']==standard]
            # Standard order remains a UI default. A repeated row is safe to
            # collapse only if every dimension, family and rule is equivalent.
            if len(candidates)>1:
                lines=', '.join(str(line) for line,_ in candidates)
                try:identities={_geometry_identity(p) for _,p in candidates}
                except (ValueError,ArithmeticError,TypeError) as exc:
                    raise Stop('УЗК: неверная геометрия в Профили.xlsx, строки '+lines+'.') from exc
                if len(identities)!=1:
                    raise Stop('УЗК: неоднозначная геометрия марки '+mark+
                               ', профиль '+assigned+', ГОСТ «'+standard+'». '
                               'Профили.xlsx, строки '+lines+
                               ': отличаются размеры, тип или правило расчёта. '
                               'Требуется однозначный утверждённый источник.')
                candidates=candidates[:1]
        if len(candidates)!=1:
            if len(candidates)>1 and not standard:
                raise Stop('Для марки '+mark+' выберите ГОСТ в столбце ГОСТ / НД.')
            raise Stop('УЗК: для марки '+mark+' профиль '+assigned+
                       ' не имеет единственной подтверждённой геометрии. '
                       'Найдено вариантов: '+str(len(candidates))+'.')
        line,profile=candidates[0]
        try:
            calculated=calculate_uzk({'status':'FOUND','profile':profile})
        except ProfileError as exc:
            raise Stop('УЗК: неверная геометрия '+assigned+
                       ' в Профили.xlsx, строка '+str(line)+': '+str(exc)) from exc
        if calculated['status']!='CALCULATED':
            raise Stop('УЗК: у '+assigned+' отсутствует принятое правило.')
        for part in calculated['rows']:
            s=_decimal(part['S'],'S профиля '+assigned)
            l=_decimal(part['L'],'L профиля '+assigned)*quantity
            output.append({
                'mark':mark,
                'profile':str(assigned),
                'quantity':int(quantity),
                'thickness_mm':_string_decimal(s),
                'length_mm':_string_decimal(l),
                'segments':tuple(part['parts']),
                'frequency_mhz':str(part['frequency_mhz']),
                'angle_deg':str(part['angle_degrees']),
                'geometry_reference':{
                    'path':str(path),'sha256':before,
                    'sheet':'Профили','row':line,
                    'standard':profile['standard'],
                    'uzk_rule':profile['uzk_rule'],
                },
            })
    if sha(path)!=before:
        raise Stop('УЗК: Профили.xlsx изменился во время расчёта.')
    return output


def profile_standards(path,assigned):
    known=_approved_rows(Path(path));key=normalize_profile(assigned)
    found=known.get(key,[])
    if not found and key.startswith('['):
        found=[x for x in known.get(key[1:],[]) if x[1]['family']=='Швеллер']
    return list(dict.fromkeys(p['standard'] for _,p in found))


def joint_text(record):
    """Deterministic descriptive field, NOT an actual defect conclusion."""
    return ('Профиль '+record['profile']+'; '+
            '/'.join(record['segments'])+
            '; толщина '+record['thickness_mm']+' мм; длина контроля '+
            record['length_mm']+' мм')


def ordered_vik_names(vom):
    """Only current visible VOM, exact mark spelling and order."""
    return tuple(dict.fromkeys(str(row[0]).strip() for row in vom.rows))
