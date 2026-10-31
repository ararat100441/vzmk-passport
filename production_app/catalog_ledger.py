"""Project shipment ledger. Migration/candidates never overwrite their input.

Only registered, identity-checked Appendix #1 quantities establish history.
Missing history and missing total are UNKNOWN, never an invented zero stock.
"""
from pathlib import Path
from datetime import date, datetime
from decimal import Decimal
from collections import Counter
from uuid import uuid4
import json, shutil, hashlib
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from .storage import Stop,sha,mutex
from .catalog_assignments import read_assignments

SCHEMA='VZMK_PROJECT_LEDGER_1'
META='_VZMK_LEDGER'
SOURCE='_VZMK_SOURCE'
HEADERS=('Марка','Количество','Профиль','№ АКЗ','Остаток')


def quantity(raw,label,allow_zero=False):
    if raw is None or isinstance(raw,bool) or isinstance(raw,str) and (not raw.strip() or raw.startswith('=')):
        raise Stop(label+': не заполнено подтверждённое целое количество.')
    try:value=Decimal(str(raw).replace(',','.'))
    except Exception as exc:raise Stop(label+': нечисловое количество.') from exc
    if not value.is_finite() or value!=value.to_integral_value() or value<(0 if allow_zero else 1):
        raise Stop(label+': требуется '+('неотрицательное' if allow_zero else 'положительное')+' целое количество.')
    return int(value)


def shipment_quantities(vom):
    result=Counter()
    for row in vom.rows:result[str(row[0]).strip()]+=quantity(row[2],'ВОМ '+str(row[0]))
    return dict(result)


def is_ledger(book):
    return META in book.sheetnames and book[META]['A1'].value==SCHEMA


def metadata(book):
    if not is_ledger(book):raise Stop('Справочник: подготовьте структуру учёта в Настройках; исходный файл пока не изменён.')
    try:data=json.loads(''.join(str(row[0] or '') for row in book[META].iter_rows(min_row=2,max_col=1,values_only=True)))
    except Exception as exc:raise Stop('Повреждена история справочника.') from exc
    s=book.worksheets[0]
    if tuple(s.cell(1,c).value for c in range(1,6))!=HEADERS:raise Stop('Изменена структура учёта справочника.')
    numbers=[h['number'] for h in data['history']]
    if len(numbers)!=len(set(numbers)):raise Stop('Дубли паспортов в истории справочника.')
    return data


def write_metadata(sheet,data):
    # Excel cells are limited to 32767 characters; long project history
    # must never be silently truncated by the XLSX writer.
    text=json.dumps(data,ensure_ascii=False)
    for row in range(2,sheet.max_row+1):sheet.cell(row,1).value=None
    for no,start in enumerate(range(0,len(text),30000),2):sheet.cell(no,1,text[start:start+30000])


def registered_history(cfg,project_name):
    """Read only; exact passport number/date/object and mass must agree."""
    from . import model
    from .first_passport import discover_first_passport_for_config
    from passport_app.appendix import inspect_template
    r=model.read_registry(cfg['registry']);p=model.select(r,project_name)
    result=[]
    for entry in r.projects[project_name]:
        n=str(entry['number']);day=entry['date'].date() if isinstance(entry['date'],datetime) else entry['date']
        history={'number':n,'date':day.isoformat(),'complete':False,'quantities':{},'reason':''}
        try:
            selected=cfg|{'source_passports':{**cfg.get('source_passports',{}),project_name:n}}
            inventory=discover_first_passport_for_config(selected,p)
            source=inventory['sources'].get('appendix')
            if not source:raise Stop('Нет единственного Приложения №1 паспорта '+n)
            sample=model.reference_project(p,{'first_passport_sources':inventory})
            layout,_,_=inspect_template(source['path'],sample)
            b=load_workbook(source['path'],read_only=True,data_only=True)
            try:
                sheet=b[layout.sheet];counts=Counter();mass=Decimal(0)
                for no in range(layout.start,layout.total_row):
                    mark=sheet.cell(no,layout.columns['mark']).value
                    if not isinstance(mark,str) or not mark.strip():raise Stop('Нет марки в историческом приложении '+n)
                    counts[mark.strip()]+=quantity(sheet.cell(no,layout.columns['quantity']).value,'История '+n+' / '+mark)
                    value=sheet.cell(no,layout.columns['mass_total']).value
                    if value is None or isinstance(value,bool):raise Stop('Нет массы исторического приложения '+n)
                    mass+=Decimal(str(value))
                if abs(mass-Decimal(str(entry['mass'])))>Decimal('.005'):raise Stop('Масса исторического приложения не совпала с реестром '+n)
            finally:b.close()
            history.update(complete=True,quantities=dict(counts),source=source['path'],sha256=source['sha256'])
        except (Stop,ValueError,KeyError,ArithmeticError) as exc:history['reason']=str(exc)
        result.append(history)
    return result


def create_migrated(source,target,project,history):
    source=Path(source);target=Path(target)
    if target.exists() or source.resolve()==target.resolve():raise Stop('Миграция требует отдельного нового файла.')
    if source.suffix.lower()!='.xlsx' or target.suffix.lower()!='.xlsx':raise Stop('Миграция справочника поддерживает только XLSX; макросы не преобразуем.')
    assignments=read_assignments(source);before=sha(source)
    book=load_workbook(source)
    try:
        if is_ledger(book):raise Stop('Справочник уже переведён на новую структуру.')
        if len(book.worksheets)!=1:raise Stop('Неизвестная исходная структура справочника.')
        original=book.worksheets[0];original.title=SOURCE;original.sheet_state='hidden'
        s=book.create_sheet('Марки и отгрузки',0)
        s.append(HEADERS)
        for mark,values in assignments.items():s.append((mark,None,values.get('profile'),values.get('akz'),None))
        for h in history:
            for mark in h['quantities']:
                if mark not in assignments:
                    assignments[mark]={};s.append((mark,None,None,None,None))
        numbers=set()
        for column,h in enumerate(history,6):
            if h['number'] in numbers:raise Stop('История содержит дублирующийся паспорт '+h['number'])
            numbers.add(h['number'])
            s.cell(1,column,'Паспорт '+h['number']+' от '+date.fromisoformat(h['date']).strftime('%d.%m.%Y'))
            for row,mark in enumerate(assignments,2):
                if h['complete']:s.cell(row,column,quantity(h['quantities'].get(mark,0),'История '+mark,True))
        m=book.create_sheet(META);m.sheet_state='veryHidden';m['A1']=SCHEMA
        write_metadata(m,{'project':project,'history':history,'source_sha256':before})
        s.freeze_panes='F2';s.auto_filter.ref=s.dimensions
        for cell in s[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='DCE6F1');cell.alignment=__import__('openpyxl').styles.Alignment(wrap_text=True)
        s.row_dimensions[1].height=42
        for col,width in [('A',24),('B',16),('C',22),('D',12),('E',16)]:s.column_dimensions[col].width=width
        for column in range(6,s.max_column+1):s.column_dimensions[get_column_letter(column)].width=25
        for row in range(2,s.max_row+1):update_balance(s,row,len(history),all(h['complete'] for h in history))
        book.active=0;book.save(target)
    finally:book.close()
    if read_assignments(target)!=assignments:raise Stop('Миграция изменила назначение марки/профиля/краски/АКЗ.')
    if sha(source)!=before:raise Stop('Исходный справочник изменился при миграции.')
    return {'marks':len(assignments),'complete_history':all(h['complete'] for h in history),
            'missing_history':[h['number'] for h in history if not h['complete']]}


def update_balance(sheet,row,count,complete):
    # A blank total/history is explicitly not a confirmed stock.
    if not complete:sheet.cell(row,5).value=None;return
    if sheet.cell(row,2).value is None:
        expression='B'+str(row)+('-SUM(F'+str(row)+':'+get_column_letter(5+count)+str(row)+')' if count else '')
        sheet.cell(row,5,'=IF(ISNUMBER(B'+str(row)+'),'+expression+',"")');return
    total=quantity(sheet.cell(row,2).value,'Общее количество '+str(sheet.cell(row,1).value),True)
    values=[quantity(sheet.cell(row,c).value,'История '+str(sheet.cell(row,1).value),True) for c in range(6,6+count)]
    if total<sum(values):raise Stop('Отрицательный остаток '+str(sheet.cell(row,1).value))
    sheet.cell(row,5,'=B'+str(row)+('-SUM(F'+str(row)+':'+get_column_letter(5+count)+str(row)+')' if count else ''))


def validate_stock(path,project,number,quantities,registered=None):
    source=Path(path);before=sha(source)
    b=load_workbook(path,data_only=False)
    try:
        data=metadata(b);s=b.worksheets[0]
        if data['project']!=project:raise Stop('Справочник учёта другого проекта.')
        if number in [h['number'] for h in data['history']]:raise Stop('Паспорт '+number+' уже записан в историю справочника.')
        if registered is not None:
            expected={str(e['number']):(e['date'].date() if isinstance(e['date'],datetime) else e['date']).isoformat() for e in registered}
            actual={h['number']:h['date'] for h in data['history']}
            if actual!=expected:raise Stop('История справочника не соответствует номерам и датам реестра. Обновите подтверждённую историю.')
        from .history_binding import reconcile_history
        reconcile_history(data,registered)
        incomplete=[h['number'] for h in data['history'] if not h['complete']]
        if incomplete:raise Stop('Остаток не подтверждён: отсутствует история паспортов '+', '.join(incomplete))
        rows={}
        for no in range(2,s.max_row+1):
            raw=s.cell(no,1).value
            if not raw:continue
            mark=str(raw).strip()
            if mark in rows:raise Stop('Дублирующаяся марка в справочнике '+mark)
            rows[mark]=no
        if s.max_column!=5+len(data['history']):raise Stop('В справочнике есть неучтённые столбцы отгрузок.')
        for col,h in enumerate(data['history'],6):
            if s.cell(1,col).value!='Паспорт '+h['number']+' от '+date.fromisoformat(h['date']).strftime('%d.%m.%Y'):raise Stop('Изменён заголовок истории '+h['number'])
        available={}
        for mark,no in rows.items():
            shipped=0
            for col,h in enumerate(data['history'],6):
                value=quantity(s.cell(no,col).value,'История '+mark,True)
                if value!=h['quantities'].get(mark,0):raise Stop('Количество истории '+h['number']+' / '+mark+' не совпало с подтверждённым приложением.')
                shipped+=value
            # B is the current editable project quantity. _VZMK_SOURCE is
            # retained only for technical assignments, never a frozen total.
            raw=s.cell(no,2).value
            if raw is None and mark not in quantities:continue
            total=quantity(raw,'Марка '+mark+': общее проектное количество',True)
            if total<shipped:raise Stop(f'Марка {mark}: отрицательный остаток {total-shipped}; проектное количество {total}, уже отгружено {shipped}.')
            available[mark]=total-shipped
        for mark,requested in quantities.items():
            if mark not in rows:raise Stop('В справочнике отсутствует марка '+mark)
            requested=quantity(requested,'Заявка '+mark,True)
            if requested>available[mark]:raise Stop(f'Марка {mark}: доступный остаток {available[mark]}, запрошено {requested}.')
        if sha(source)!=before:raise Stop('Справочник изменился во время проверки остатков.')
        return data
    finally:b.close()


def shipment_candidate(source,target,project,number,day,quantities,registered=None):
    before=sha(source);data=validate_stock(source,project,number,quantities,registered)
    if Path(target).exists():raise Stop('Кандидат справочника уже существует.')
    b=load_workbook(source)
    try:
        s=b.worksheets[0];col=6+len(data['history'])
        if s.max_column>=col:raise Stop('В справочнике есть неучтённые столбцы отгрузок.')
        s.cell(1,col,'Паспорт '+number+' от '+date.fromisoformat(day).strftime('%d.%m.%Y'))
        for no in range(2,s.max_row+1):
            mark=str(s.cell(no,1).value or '').strip()
            if mark:s.cell(no,col,quantities.get(mark,0));update_balance(s,no,len(data['history'])+1,True)
        data['history'].append({'number':number,'date':day,'complete':True,'quantities':quantities,'origin':'CONFIRMED_CURRENT_ISSUE'})
        write_metadata(b[META],data);s.auto_filter.ref=s.dimensions;b.save(target)
    finally:b.close()
    if sha(source)!=before or read_assignments(source)!=read_assignments(target):raise Stop('Изменились назначения при подготовке отгрузки.')
    return {'number':number,'column':col,'sha256':sha(target),
            'history_binding_updates':[{'number':h['number'],**u} for h in data['history'] for u in h.get('binding_updates',[])]}


def refresh_history(source,target,project,history):
    if Path(target).exists():raise Stop('Кандидат истории уже существует.')
    b=load_workbook(source)
    try:
        data=metadata(b);sheet=b.worksheets[0]
        if data['project']!=project:raise Stop('История другого объекта.')
        if sheet.max_column!=5+len(data['history']):raise Stop('Есть неучтённые столбцы истории.')
        old={h['number']:h for h in data['history']}
        if not set(old).issubset({h['number'] for h in history}):raise Stop('Зарегистрированный паспорт исчез из реестра.')
        for h in history:
            previous=old.get(h['number'])
            if previous and (previous['date']!=h['date'] or previous['complete'] and (not h['complete'] or previous['quantities']!=h['quantities'])):
                raise Stop('Изменилась подтверждённая история '+h['number']+'; автоматическая замена запрещена.')
        known={str(sheet.cell(row,1).value or '').strip() for row in range(2,sheet.max_row+1)}
        for h in history:
            for mark in h['quantities']:
                if mark not in known:
                    sheet.append((mark,None,None,None,None));known.add(mark)
        for col,h in enumerate(history,6):
            sheet.cell(1,col,'Паспорт '+h['number']+' от '+date.fromisoformat(h['date']).strftime('%d.%m.%Y'))
            for row in range(2,sheet.max_row+1):
                mark=str(sheet.cell(row,1).value or '').strip()
                sheet.cell(row,col).value=h['quantities'].get(mark,0) if h['complete'] else None
        for row in range(2,sheet.max_row+1):update_balance(sheet,row,len(history),all(h['complete'] for h in history))
        data['history']=history;write_metadata(b[META],data);b.save(target)
    finally:b.close()
    before=read_assignments(source);after=read_assignments(target)
    if any(after.get(mark)!=values for mark,values in before.items()):raise Stop('Обновление истории изменило назначения.')
    return {'marks':len(read_assignments(target)),'complete_history':all(h['complete'] for h in history),'missing_history':[h['number'] for h in history if not h['complete']]}


def migrate_in_place(path,cfg,project):
    """Operator-requested migration only; backups survive success or failure."""
    from passport_app.excel import locked_file,atomic_replace
    p=Path(path);token=uuid4().hex;candidate=p.with_name('.ledger-'+token+'.xlsx');backup=p.with_name(p.stem+'.before-ledger-'+token+'.xlsx')
    history=registered_history(cfg,project)
    with mutex(p),locked_file(p,writable=True) as original:
        backup.write_bytes(original)
        if sha(backup)!=hashlib.sha256(original).hexdigest():raise Stop('Резервная копия справочника не совпала.')
        # Use verified backup because the working path is exclusively locked.
        check=load_workbook(backup,read_only=True)
        try:converted=is_ledger(check)
        finally:check.close()
        result=(refresh_history if converted else create_migrated)(backup,candidate,project,history)
        atomic_replace(candidate,p,backup=p.with_name('.replaced-ledger-'+token+'.xlsx'))
    return result|{'backup':str(backup),'path':str(p),'notice':'Заполните общее количество; неполная история блокирует официальный выпуск.'}
