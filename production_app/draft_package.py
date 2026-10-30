"""Read-only, explicitly INCOMPLETE document preparation; never an issue plan.

Uses the accepted first-form adapters and native Office/PDF workers. Unknown
values missing from current inputs remain REFERENCE CONTENT, never certified. This module
does not call plan/confirm/release/commit and cannot authorize registry writes.
"""
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4
import re
import shutil

from docx import Document
from docx.shared import Pt
from io import BytesIO
from zipfile import ZipFile

from passport_app.excel import load_book, excel_session, open_excel
from passport_app.appendix import generate_appendix, inspect_template
from stage8_closure.storage import frozen
from . import model, office
from .storage import Stop, safe, local_root, pins, verify_pins, save, sha
from .shared_uzk import find_shared_uzk, planned_control
from .catalog_assignments import current_assignments
from .automatic_acts import find_approved_geometry, calculate_design_characteristics
from .highlighted_reference import (
    supports_vik_first, supports_uzk_first, _set_sdt, _replace_vik_names,
    _replace_uzk_heading, _copy_uzk_rows, _save_preserving_parts, _plain,
)
from .quality_from_first import (
    is_exact_first_quality, _one_highlighted_cell, _fill_highlighted,
    DATE_RUN,
)
from .akz_from_first import is_first_akz, is_first_akz_appendix, fill_first_akz_appendix

BANNER='ЧЕРНОВИК — НЕ ДЛЯ ВЫПУСКА'
REFERENCE='ОБРАЗЕЦ: содержание выбранного паспорта требует проверки; новый контроль не подтверждён'


def shipment_summary(vom,planned):
    return {'marks':len(vom.rows),'quantity':str(vom.quantity),
            'mass_kg':format(vom.mass,'.2f'),'area_m2':format(vom.area,'.3f'),
            'uzk_quantity':sum(p['uzk_quantity'] for p in planned.values())}


def collect_marks(vom, planned, assigned, saved):
    """Only engineering inputs; never inherit hidden verdicts or approvals."""
    old={str(r.get('mark')):r for r in (saved or {}).get('marks',[])}
    expected=set(planned)
    if old and set(old)!=expected:
        raise Stop('Сохранённые марки относятся к другой ВОМ. Откройте фактические результаты заново.')
    rows=[]
    for mark,control in planned.items():
        previous=old.get(mark,{})
        source=assigned.get(mark,{})
        raw=previous.get('akz',source.get('akz'))
        if raw is None or str(raw).strip()=='':variant=None
        elif isinstance(raw,bool) or str(raw).strip() not in ('0','1','2','3'):
            raise Stop('Некорректный № АКЗ: '+mark)
        else:variant=int(raw)
        rows.append({'mark':mark,'profile':str(previous.get('profile',source.get('profile')) or '').strip(),
                     'standard':str(previous.get('standard') or ''),
                     'akz':variant,'paint':str(source.get('paint') or ''),
                     'uzk_quantity':control['uzk_quantity']})
    return rows


def draft_climates(saved, rows):
    from .simple_facts import prepare_for_check
    from .measurement_ranges import FIELDS, midrange, magnus_dew_point,climate_midrange,LABELS
    climates={};notes=[]
    for variant in sorted({r['akz'] for r in rows if r['akz'] is not None}):
        inputs=((saved or {}).get('akz_drafts') or {}).get(str(variant)) or {}
        current={}
        for key in ('air','surface','humidity'):
            raw=inputs.get(key) or {}
            if str(raw.get('min','')).strip() and str(raw.get('max','')).strip():
                current[key]=climate_midrange(variant,key,raw)
        layers=inputs.get('layer_intervals') or []
        wet=inputs.get('wet') or {}
        if inputs.get('coats')==1 and all(str(wet.get(k,'')).strip() for k in ('min','max')):
            current['wet_thickness']=climate_midrange(variant,'wet',wet)
        if inputs.get('coats')==1 and len(layers)==1 and layers[0].get('layer')==1:
            raw=layers[0]
            if str(raw.get('min','')).strip() and str(raw.get('max','')).strip():
                current['thickness']=climate_midrange(variant,'thickness',raw)
        if {'air','surface','humidity','thickness'}.issubset(current):
            data={'simple_ui':True,'marks':[r for r in rows if r['akz']==variant],
                  'akz_drafts':{str(variant):deepcopy(inputs)}}
            current=prepare_for_check(data)['climates'][0]
        elif 'air' in current and 'humidity' in current:
            current['dew']=format(magnus_dew_point(current['air'],current['humidity']),'f')
        if current:climates[variant]=current
        missing=[k for k in ('wet_thickness','thickness','surface','air','humidity','dew') if k not in current]
        if missing:
            readable={'wet_thickness':'толщина мокрого слоя','thickness':'толщина сухого слоя','surface':'температура поверхности','air':'температура воздуха','humidity':'влажность','dew':'точка росы'}
            notes.append('АКЗ №'+str(variant)+': нет полных текущих диапазонов '+', '.join(readable[k] for k in missing)+
                         '; сохранены значения первого акта как образец, требуют проверки.')
    return climates,notes


def _dates(registry,project,number,new):
    """Read dates only from the requested registry row; missing dates retain the form sample."""
    keys={'date':'END_DATE','vik_date':'VIK_DATE','uzk_date':'UZK_DATE','akz_date':'AKZ_DATE'}
    result={k:None for k in keys}
    cols=model.columns(registry);w=load_book(registry.path,values=True)
    try:
        result['start_date']=w[registry.sheet].cell(project.first_row,cols['START_DATE']).value
        if not new:
            rows=[e['row'] for e in registry.projects[project.name] if e['number']==number]
            if len(rows)!=1:raise Stop('Паспорт не найден однозначно в реестре.')
            for key,col in keys.items():
                value=w[registry.sheet].cell(rows[0],cols[col]).value
                if isinstance(value,(date,datetime)):result[key]=value
    finally:w.close()
    return result


def _day(value):
    return value.strftime('%d.%m.%Y') if isinstance(value,(date,datetime)) else None


def _quality(source,target,number,dates,vom):
    doc=Document(source);t=doc.tables[0]
    run=_one_highlighted_cell(t,1,9,r'[1-9]\d*/20\d\d');old=run.text
    _fill_highlighted(run,number)
    mass=f'{vom.mass:,.2f}'.replace(',',' ').replace('.',',')
    _fill_highlighted(_one_highlighted_cell(t,16,6,r'[\d\s]+,\d\d'),mass)
    if dates['date'] is not None:
        _fill_highlighted(_one_highlighted_cell(t,18,8,r'\d\d\.\d\d\.20\d\d'),_day(dates['date']))
    from .quality_from_first import _appendix_acts, _manufacturing_start
    for row in (50,53):
        paragraph=t.rows[row].cells[0].paragraphs[0]
        # When an old API caller supplies no dates, retain sample dates with
        # the explicit draft warning. The normal GUI resolves current dates.
        values=[_day(dates[k]) for k in ('vik_date','uzk_date','akz_date')]
        if all(values):
            _appendix_acts(paragraph,old,number,*values,variants=dates.get('akz_variants',[0]))
        else:
            from docx.enum.text import WD_COLOR_INDEX
            for run in paragraph.runs:
                if run.font.highlight_color==WD_COLOR_INDEX.YELLOW:
                    if run.text==old:run.text=number
                    run.font.highlight_color=None
    _manufacturing_start(t,dates.get('start_date'))
    if dates['date'] is not None:
        _fill_highlighted(_one_highlighted_cell(doc.tables[1],4,0,r'\d\d\.\d\d\.20\d\d'),_day(dates['date']))
    _save_preserving_parts(doc,source,target)


def _word_copy(source,target,role,number,dates,vom,design):
    if role=='quality':return _quality(source,target,number,dates,vom)
    doc=Document(source)
    if role=='vik':
        _set_sdt(doc,'PASSPORT_NO',number,2)
        if dates['vik_date'] is not None:_set_sdt(doc,'VIK_DATE',_day(dates['vik_date']),2)
        _replace_vik_names(doc,[row[1] for row in vom.rows])
    else:
        heading=next(p.text for p in doc.paragraphs if p.text.startswith('Акт № ') and ' от ' in p.text)
        _replace_uzk_heading(doc,number,_day(dates['uzk_date']) or heading.rsplit(' от ',1)[1])
        _copy_uzk_rows(doc,design,_day(dates['uzk_date']) or doc.tables[0].rows[5].cells[4].text)
    _save_preserving_parts(doc,source,target,relationships=role=='uzk')


def _akz_draft(sheet,number,day,climate):
    """Change ONLY agreed variable cells; keep original formulas and sample data."""
    sample_wet=sheet.Range('E27').Value2
    sheet.Range('E4').Value2=number+' от '
    if day is not None:
        sheet.Range('H4').Value2=((day.date() if isinstance(day,datetime) else day)-date(1899,12,30)).days
    for position,key in [('E27','wet_thickness'),('E28','thickness'),('E31','surface'),('E32','air'),('E33','humidity'),('E34','dew')]:
        value=(climate or {}).get(key)
        if value is not None:sheet.Range(position).Value2=float(value)
    if (climate or {}).get('wet_thickness') is None:
        sample=sample_wet
        if not isinstance(sample,(int,float)) or sample<=0:raise Stop('АКЗ: нет текущей мокрой толщины и числового образца E27.')
        sheet.Range('E27').Value2=sample


def _stamp_word_draft(path):
    """Append a draft warning in headers; preserve body and opaque source parts."""
    doc=Document(path);seen=set()
    for section in doc.sections:
        headers=[section.header]
        if section.different_first_page_header_footer:headers.append(section.first_page_header)
        if doc.settings.odd_and_even_pages_header_footer:headers.append(section.even_page_header)
        for header in headers:
            part=str(header.part.partname)
            if part in seen:continue
            seen.add(part)
            p=header.add_paragraph();run=p.add_run(BANNER)
            run.font.size=Pt(8)
    stream=BytesIO();doc.save(stream);stream.seek(0)
    original=Path(path).read_bytes();output=BytesIO()
    with ZipFile(BytesIO(original)) as old,ZipFile(stream) as new,ZipFile(output,'w') as dst:
        changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
        changed.update(name for name in new.namelist() if name.startswith('word/header') and name.endswith('.xml'))
        for item in old.infolist():dst.writestr(item,new.read(item.filename) if item.filename in changed else old.read(item.filename))
        for name in new.namelist():
            if name not in old.namelist() and name in changed:dst.writestr(name,new.read(name))
    Path(path).write_bytes(output.getvalue())


def prepare_draft(cfg,project,number,vom,facts=None,new=True):
    """Create a separate package under PRIVATE drafts; source files read only."""
    registry,_,_=model.validate_paths(cfg)
    from .core import training_guard, pending
    training_guard(cfg)
    if pending(registry):raise Stop('Есть незавершённая запись. Сначала выполните восстановление.')
    reg=model.read_registry(registry);selected=model.select(reg,project)
    if new and number!=selected.next_number:raise Stop('Следующий номер изменился: '+selected.next_number)
    model.parse_number(number)
    v=model.visible_vom(safe(vom));profile=model.profile(cfg,project)
    if v.area is None:raise Stop('Площадь ВОМ не указана: заполните площадь для Приложения №1 и приложений АКЗ. Неизвестная площадь не заменяется нулём.')
    sample=model.reference_project(selected,profile)
    if not profile.get('catalog'):raise Stop('Не найден Справочник марок.xlsx выбранного объекта '+project+'. Проверьте путь в Настройках; профили и АКЗ не подставлены.')
    shared=find_shared_uzk(cfg);planned=planned_control(shared,project,v)
    assigned=current_assignments(profile['catalog'],profile['catalog_sha256'],planned)
    rows=collect_marks(v,planned,assigned,facts)
    climates,notes=draft_climates(facts,rows)
    for item in rows:
        choice=cfg.get('geometry_choices',{}).get(project,{}).get(item['mark'])
        if choice:item['standard']=choice
    dates=_dates(reg,selected,number,new)
    if new and (facts or {}).get('shipment_date'):
        from .registry_dates import resolve_registry_dates
        chosen=date.fromisoformat(facts['shipment_date'])
        resolved=resolve_registry_dates(cfg,project,number,v,True,day=chosen)
        dates={k:date.fromisoformat(resolved[k]) for k in ('start_date','date','vik_date','uzk_date','akz_date')}
    dates['akz_variants']=sorted({r['akz'] for r in rows if r['akz'] is not None})
    for key,label in [('date','окончания изготовления / отгрузки'),('vik_date','ВИК'),('uzk_date','УЗК'),('akz_date','АКЗ')]:
        if dates[key] is None:notes.append('В реестре '+str(registry)+' ['+reg.sheet+'] нет определённой даты '+label+' для '+number+
                                         '; сохранена дата соответствующего первого бланка как образец, требует проверки.')
    notes.append(REFERENCE+'. ФИО, реквизиты и формулировки сохранены из выбранного паспорта. Официальный выпуск запрещён.')
    missing=[r['mark'] for r in rows if r['akz'] is None]
    if missing:notes.append('Без подтверждённого назначения АКЗ: '+', '.join(missing)+'. В приложения АКЗ эти марки не включены; №0 не присвоен.')
    geometry=find_approved_geometry(cfg,require=True)
    design=calculate_design_characteristics(v,{'marks':rows},geometry,
              expected_sha=cfg.get('profiles_workbook_sha256') or sha(geometry))
    sources=profile.get('first_passport_sources',{}).get('sources',{})
    if profile.get('first_passport_sources') is not None:
        from .working_facts import selected_form
        for role in ['appendix','vik','uzk','quality',*[r+str(n) for n in dates['akz_variants'] for r in ('akz','akz_appendix')]]:
            selected_form(profile,role)
    inputs=pins([registry,v.path,shared,geometry,profile['catalog'],
                 *[t['path'] for t in profile['templates'].values()],*[s['path'] for s in sources.values()]])
    if inputs[str(registry)]!=reg.digest or inputs[str(v.path)]!=v.digest:
        raise Stop('Источник изменился во время подготовки черновика.')
    for ref in [*profile['templates'].values(),*sources.values(),
                {'path':profile['catalog'],'sha256':profile['catalog_sha256']}]:
        if inputs[str(safe(ref['path']))]!=ref['sha256']:
            raise Stop('Источник изменился после выбора: '+str(ref['path']))
    job=local_root()/'drafts'/uuid4().hex
    title='ЧЕРНОВИК_'+model.filename(number)
    if new and (facts or {}).get('shipment_date'):title+=' от '+_day(dates['date'])
    folder=job/title;folder.mkdir(parents=True,exist_ok=False)
    outputs=[];blocked={}
    save(job/'state.json',{'schema':'VZMK_DRAFT_STATE_1','phase':'PREPARING','official_issue':False,'inputs':inputs})
    with frozen(inputs):
        verify_pins(inputs)
        with excel_session() as excel:
            ref=profile['templates']['appendix']['path'];target=folder/(model.filename(number)+' - appendix.xlsx')
            layout,_,object_cell=inspect_template(ref,sample);shutil.copyfile(ref,target)
            original=load_book(ref,values=True)
            try:source_object=original[layout.sheet][object_cell].value
            finally:original.close()
            generate_appendix(excel,target,sample,v,{'number':number,'date':dates['date'] or selected.first_date,'object':selected.object_name},undated_draft=True)
            book=open_excel(excel,target)
            try:
                s=book.Worksheets(layout.sheet)
                s.Range(object_cell).Value2=source_object
                office.fit_appendix_body(s,layout.start,layout.start+len(v.rows)-1)
                office.set_appendix_print_titles(s,layout.start)
                s.Cells(layout.start+len(v.rows),layout.columns['area_total']).NumberFormatLocal='0'+str(excel.DecimalSeparator)+'000'
                book.Save()
            finally:book.Close(SaveChanges=False)
            outputs.append({'role':'appendix','path':str(target)})
            variants=sorted({r['akz'] for r in rows if r['akz'] is not None})
            for variant in variants:
                act,appendix='akz'+str(variant),'akz_appendix'+str(variant)
                a=sources.get(act);z=sources.get(appendix)
                if not a or not z or not is_first_akz(a['path']) or not is_first_akz_appendix(z['path']):
                    blocked[act]='Нет распознанного собственного бланка с проверенными переменными полями для АКЗ №'+str(variant)
                    blocked[appendix]=blocked[act];continue
                selected_rows=[row for row in v.rows if next(r for r in rows if r['mark']==str(row[0]))['akz']==variant]
                from .akz_from_first import act_number
                actual_number=act_number(number,variant)
                for role,source in ((act,a['path']),(appendix,z['path'])):
                    target=folder/(actual_number.replace('/','_')+' - '+role+'.xlsx');shutil.copyfile(source,target)
                    book=open_excel(excel,target)
                    try:
                        s=book.Worksheets(1)
                        if role==act:_akz_draft(s,actual_number,dates['akz_date'],climates.get(variant))
                        else:
                            from .akz_from_first import TITLE
                            sample_day=datetime.strptime(TITLE.fullmatch(str(s.Range('A1').Value2)).group(3),'%d.%m.%Y').date()
                            fill_first_akz_appendix(s,number=actual_number,akz_date=dates['akz_date'] or sample_day,rows=selected_rows,undated_draft=True)
                            office.fit_appendix_body(s,4,3+len(selected_rows))
                        s.PageSetup.CenterHeader=BANNER
                        s.PageSetup.CenterFooter='АКЗ №'+str(variant)+': '+str(len(selected_rows))+' марок. Без назначения: '+str(len(missing))+'.'
                        excel.CalculateFullRebuild();book.Save()
                    finally:book.Close(SaveChanges=False)
                    outputs.append({'role':role,'path':str(target)})
        for role,recognize in [('quality',is_exact_first_quality),('vik',supports_vik_first),('uzk',supports_uzk_first)]:
            ref=sources.get(role)
            if not ref or not recognize(ref['path']):
                blocked[role]='Нет распознанного первого бланка '+role+' с проверенными переменными полями.';continue
            if role=='uzk' and not design:
                blocked[role]='По текущей ВОМ нет марок с положительным планом УЗК.';continue
            target=folder/(model.filename(number)+' - '+role+'.docx')
            _word_copy(ref['path'],target,role,number,dates,v,design)
            outputs.append({'role':role,'path':str(target)})
        for item in outputs:
            if Path(item['path']).suffix=='.docx':_stamp_word_draft(item['path'])
        notes += [role+': '+reason for role,reason in blocked.items()]
        from .print_pdf import create_print_pdf
        result=create_print_pdf(outputs,folder,folder.name,partial_draft=True)
        verify_pins(inputs)
    manifest={'schema':'VZMK_DRAFT_PACKAGE_1','status':'DRAFT_READY','official_issue':False,
              'reference_content_unverified':True,'dates_from_registry':dates,
              'project':project,'number':number,'folder':str(folder),'preview_pdf':result['output']['path'],
              'outputs':outputs,'print_pdf':result['manifest'],'blocked_documents':blocked,'notes':notes,
              'missing_akz_assignments':missing,'marks':rows,'fingerprints':inputs,
              'output_fingerprints':pins([x['path'] for x in outputs]+[result['output']['path']]),
              'summary':shipment_summary(v,planned)}
    save(folder/'DRAFT.json',manifest)
    (folder/'ПРОЧИТАТЬ — ЧЕРНОВИК.txt').write_text(BANNER+'\n\n'+'\n'.join(notes)+'\n\nЧерновик нельзя подтвердить или записать в реестр. Результаты контроля заполняются пользователем самостоятельно.\n',encoding='utf-8-sig')
    save(job/'state.json',{'schema':'VZMK_DRAFT_STATE_1','phase':'DRAFT_READY','official_issue':False,'inputs':inputs})
    return {'status':'PASS',**{k:manifest[k] for k in ('schema','official_issue','folder','preview_pdf','summary','notes','blocked_documents')},
            'files':len(outputs),'pages':result['manifest']['pages'],'manifest':str(folder/'DRAFT.json')}
