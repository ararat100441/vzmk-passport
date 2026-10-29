"""Daily Tk interface over the existing isolated production worker pipeline."""
from copy import deepcopy
from datetime import date
from pathlib import Path
from uuid import uuid4
import os, sys, re, tkinter as tk, traceback
from tkinter import ttk, filedialog, messagebox
from .gui import App as EngineUI
from . import model, core, akz_presets, VERSION, control_book
from .storage import read,save,sha,Stop,local_root
from .real_forms import DOC_FIELD_LABELS

def explain_error(message):
    first,separator,rest=str(message).partition('\n')
    prefix='Не подтверждена дата: '
    if first.startswith(prefix):first='Укажите «'+control_book.META.get(first[len(prefix):],first[len(prefix):])+'» в данных текущего комплекта.'
    prefix='Не подтверждено '
    if first.startswith(prefix):
        field,_,suffix=first[len(prefix):].partition(' слоя ')
        if field in control_book.LAYERS:first='Заполните «'+control_book.LAYERS[field]+'» слоя '+suffix+'.'
    prefix='Не заполнено поле '
    if first.startswith(prefix):
        field,_,suffix=first[len(prefix):].partition(' в УЗК')
        labels=control_book.SHEETS['uzk_details'][1]
        if field in labels:first='Заполните «'+labels[field]+'» в УЗК'+suffix+'.'
    return first+separator+rest

class DailyApp(EngineUI):
    def __init__(self,root):
        self.root=root;root.title('ВЗМК — Паспортизация '+VERSION);self.fit(root,1040,730)
        root.minsize(850,620)
        settings=local_root()/'settings.json'
        initial=Path(sys.executable).parent/'initial_settings.json'
        try:self.cfg=read(settings) if settings.exists() else read(initial) if getattr(sys,'frozen',False) and initial.exists() else {'profiles':{}}
        except Exception:self.cfg={'profiles':{}}
        self.plan=None;self.facts={};self.draft={};self.active=None;self.result_path=None
        self.phase='input';self.preview_path=None;self.final_pdf=None;self.final_folder=None;self.result_project=None;self.opened=False
        self.cancel_requested=False;self.close_when_safe=False
        self.project=tk.StringVar();self.number=tk.StringVar();self.vom=tk.StringVar();self.new=tk.BooleanVar(value=True)
        self.lockable=[]
        self.day=tk.StringVar(value=date.today().strftime('%d.%m.%Y'));self.operator=tk.StringVar(value=self.cfg.get('operator',''))
        self.approved=tk.BooleanVar(value=False);self.reviewed=tk.BooleanVar(value=False)
        self.control_path=tk.StringVar();self.summary=tk.StringVar(value='Выберите проект и ВОМ.')
        self.data_summary=tk.StringVar(value='Добавьте результаты текущего контроля из Excel или введите недостающие данные.')
        self.project.trace_add('write',lambda *_:self.inputs_changed())
        self.vom.trace_add('write',lambda *_:self.inputs_changed())
        self.day.trace_add('write',lambda *_:self.metadata_changed())
        self.operator.trace_add('write',lambda *_:self.metadata_changed())
        top=ttk.Frame(root,padding=14);top.pack(fill='x')
        ttk.Label(top,text='Паспортизация',font=('Segoe UI',19,'bold')).pack(side='left')
        ttk.Button(top,text='Настройки',command=self.settings_menu).pack(side='right')
        self.paths=ttk.Label(root,wraplength=980,justify='left');self.paths.pack(fill='x',padx=16)
        form=ttk.Frame(root,padding=16);form.pack(fill='x');form.columnconfigure(1,weight=1)
        for n,label in enumerate(('Проект','ВОМ','Дата паспорта','Номер паспорта','Ответственный')):
            ttk.Label(form,text=label).grid(row=n,column=0,sticky='w',pady=5)
        self.projects=ttk.Combobox(form,textvariable=self.project,state='readonly');self.projects.grid(row=0,column=1,sticky='ew',padx=10)
        self.projects.bind('<<ComboboxSelected>>',lambda e:self.select_project())
        self.vom_widget=ttk.Entry(form,textvariable=self.vom);self.vom_widget.grid(row=1,column=1,sticky='ew',padx=10)
        ttk.Button(form,text='Выбрать файл…',command=self.choose_vom).grid(row=1,column=2)
        date_entry=ttk.Entry(form,textvariable=self.day,width=20);date_entry.grid(row=2,column=1,sticky='w',padx=10);self.lockable.append(date_entry)
        self.numbers=ttk.Entry(form,textvariable=self.number,state='readonly');self.numbers.grid(row=3,column=1,sticky='w',padx=10)
        operator_entry=ttk.Entry(form,textvariable=self.operator);operator_entry.grid(row=4,column=1,sticky='ew',padx=10);self.lockable.append(operator_entry)
        self.new_widget=ttk.Checkbutton(root,variable=self.new)  # engine compatibility, never shown
        ttk.Label(root,textvariable=self.summary,font=('Segoe UI',11,'bold')).pack(fill='x',padx=16,pady=5)
        data=ttk.LabelFrame(root,text='Данные текущего комплекта',padding=12);data.pack(fill='x',padx=16,pady=10);self.data_frame=data
        ttk.Label(data,textvariable=self.data_summary,wraplength=920).pack(fill='x')
        bar=ttk.Frame(data);bar.pack(fill='x',pady=(10,0))
        ttk.Button(bar,text='Загрузить ведомость Excel…',command=self.import_controls).pack(side='left')
        ttk.Button(bar,text='Ввести / изменить данные',command=self.edit_data).pack(side='left',padx=8)
        ttk.Button(bar,text='Новая ведомость Excel…',command=self.new_control_book).pack(side='left')
        approve_widget=ttk.Checkbutton(root,text='Подтверждаю данные текущих протоколов, свои полномочия и применяемые формы',variable=self.approved);approve_widget.pack(anchor='w',padx=16,pady=4);self.lockable.append(approve_widget);self.approve_widget=approve_widget
        self.confirm_frame=ttk.LabelFrame(root,text='Просмотр перед записью',padding=12)
        self.confirm_text=ttk.Label(self.confirm_frame,wraplength=920,justify='left');self.confirm_text.pack(fill='x')
        ttk.Checkbutton(self.confirm_frame,text='Я просмотрел(а) весь PDF и подтверждаю выпуск указанного комплекта',variable=self.reviewed).pack(anchor='w',pady=6)
        actions=ttk.Frame(root,padding=16);actions.pack(fill='x')
        self.primary=ttk.Button(actions,text='Сформировать пакет',command=self.primary_action);self.primary.pack(side='left',fill='x',expand=True)
        self.buttons=[self.primary]
        ttk.Button(actions,text='Открыть PDF',command=self.open_pdf).pack(side='left',padx=8)
        ttk.Button(actions,text='Открыть папку',command=self.open_folder).pack(side='left')
        self.progress=ttk.Label(root,text='Реестр изменяется только после просмотра и подтверждения.');self.progress.pack(fill='x',padx=16)
        self.log=tk.Text(root,height=5,wrap='word',state='disabled',font=('Segoe UI',9))
        ttk.Label(root,text='Готовые Word и Excel остаются редактируемыми. PDF сохраняется рядом с ними.',wraplength=940).pack(fill='x',padx=16,pady=12)
        root.protocol('WM_DELETE_WINDOW',self.close)
        root.bind('<Escape>',lambda e:self.cancel())
        self.refresh()
        configured=list(self.cfg.get('profiles',{}))
        name=self.cfg.get('last_project') or (configured[0] if len(configured)==1 else '')
        if name in self.projects['values']:self.project.set(name);self.select_project()
        self.vom_widget.bind('<FocusOut>',lambda e:self.update_summary())

    def invalidate(self):
        if self.active:return
        self.plan=None;self.phase='input';self.opened=False;self.preview_path=None
        if hasattr(self,'confirm_frame'):self.confirm_frame.pack_forget()
        if hasattr(self,'primary'):self.primary.configure(text='Сформировать пакет',state='disabled' if self.active else 'normal')
        if hasattr(self,'primary') and hasattr(self,'data_frame'):
            self.data_frame.pack(fill='x',padx=16,pady=10,before=self.primary.master)
            self.approve_widget.pack(anchor='w',padx=16,pady=4,before=self.primary.master)
        if hasattr(self,'reviewed'):self.reviewed.set(False)

    def inputs_changed(self):
        if self.active:return
        self.draft={};self.facts={}
        if hasattr(self,'approved'):self.approved.set(False)
        if hasattr(self,'data_summary'):self.data_summary.set('Добавьте результаты текущего контроля. Старые результаты не перенесены.')
        self.invalidate()

    def metadata_changed(self):
        if self.active:return
        self.approved.set(False);self.invalidate()

    def refresh(self):
        if self.active:return
        self.paths.configure(text=('УЧЕБНЫЙ РЕЖИМ — только отдельная копия\n' if self.cfg.get('training') else 'РАБОЧИЙ РЕЖИМ — запись после подтверждения\n')+'Реестр: '+self.cfg.get('registry','не выбран'))
        if not self.cfg.get('registry'):self.progress.configure(text='Выберите рабочие пути в настройках.');return
        try:
            registry=model.read_registry(self.cfg['registry']);self.projects['values']=sorted(registry.projects)
            if self.project.get():self.select_project()
        except Exception as exc:self.error(exc)

    def select_project(self):
        if self.active:return
        self.inputs_changed()
        self.number.set('');self.control_path.set('');self.vom.set('')
        self.summary.set('Выберите ВОМ для объекта '+self.project.get()+'.')
        if self.result_project!=self.project.get():self.final_pdf=None;self.final_folder=None
        try:
            p=model.select(model.read_registry(self.cfg['registry']),self.project.get());self.number.set(p.next_number)
            profile=self.cfg.get('profiles',{}).get(self.project.get(),{})
            previous=profile.get('last_vom','')
            if previous and Path(previous).is_file():self.vom.set(previous);self.update_summary()
            if not profile:self.progress.configure(text='Для этого проекта один раз подключите его шаблоны в настройках.')
            else:self.progress.configure(text='Следующий номер определён по актуальному реестру. Результаты новой отгрузки не подтверждены.')
        except Exception as exc:self.error(exc)

    def choose_vom(self):
        self.choose(self.vom,False)
        if self.vom.get():self.update_summary()

    def update_summary(self):
        if not self.vom.get():self.summary.set('Выберите ВОМ.');return
        try:
            v=model.visible_vom(self.vom.get());self.summary.set(f'{len(v.rows)} строк / {v.quantity} изделий   •   Масса {v.mass:.2f} кг   •   Площадь {v.area:.2f} м²')
        except Exception as exc:self.error(exc)

    def current_preferences(self):return deepcopy(self.cfg.get('profiles',{}).get(self.project.get(),{}).get('preferences',{}))

    def fresh_draft(self,vom):
        prefs=self.current_preferences()
        d={'project':self.project.get(),'number':self.number.get(),'date':control_book.text(self.day.get(),'date'),'responsible':self.operator.get(),
           'marks':[{'mark':mark,**prefs.get('assignments',{}).get(mark,{})} for mark in dict.fromkeys(str(r[0]) for r in vom.rows)],
           'doc_fields':prefs.get('doc_fields',{}),'climates':[]}
        if not self.cfg.get('training'):
            presets=akz_presets.last_for_project(self.project.get())
            for n in sorted({r['akz'] for r in d['marks'] if type(r.get('akz')) is int}):
                cached=presets['variants'].get(str(n),presets.get('latest',{}))
                if cached:d['climates'].append({'variant':n,**cached.get('climate',{}),'layers':cached.get('layers',[])})
        return d

    def import_controls(self):
        if self.active:return
        if not self.project.get() or not self.vom.get():self.error('Сначала выберите проект и ВОМ.');return
        path=filedialog.askopenfilename(title='Ведомость текущего контроля',filetypes=[('Ведомость Excel','*.xlsx *.xlsm')])
        if not path:return
        try:
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,keep_links=False);full='Комплект' in book.sheetnames;book.close()
            if full:d=control_book.read_book(path,self.project.get(),self.number.get(),self.cfg.get('training',False),self.current_preferences())
            else:
                details,proof=control_book.read_uzk_book(path)
                d=deepcopy(self.draft) if self.draft else self.fresh_draft(model.visible_vom(self.vom.get()))
                d['uzk_details']=details;d['evidence']=d.get('evidence',[])+[proof]
                dates={r['date'] for r in details}
                if len(dates)==1:d['uzk_date']=next(iter(dates))
                d.setdefault('climates',[])
            if d.get('date'):self.day.set(date.fromisoformat(d['date']).strftime('%d.%m.%Y'))
            if d.get('responsible'):self.operator.set(d['responsible'])
            self.draft=d;self.control_path.set(path);self.approved.set(False);self.invalidate()
            self.data_summary.set(f'Загружено: {len(d["marks"])} марок, {len(d["uzk_details"])} соединений УЗК, {len(d["climates"])} вариантов АКЗ. Значения доступны для изменения.\n'+str(path))
            if not full:self.data_summary.set('Соединения УЗК загружены без повторного ввода. Дополните недостающие сведения ВИК/АКЗ и количество проконтролированных изделий в «Ввести / изменить данные».\n'+str(path))
        except Exception as exc:self.error(exc)

    def new_control_book(self):
        if self.active:return
        try:
            v=model.visible_vom(self.vom.get())
            path=filedialog.asksaveasfilename(title='Новая ведомость текущего контроля',initialfile='Ведомость '+self.number.get().replace('/','_')+'.xlsx',defaultextension='.xlsx',filetypes=[('Excel','*.xlsx')])
            if not path:return
            control_book.write_book(path,self.project.get(),self.number.get(),v,self.current_preferences(),self.draft or self.fresh_draft(v))
            os.startfile(path)
            self.data_summary.set('Новая ведомость открыта в Excel. Внесите только фактические результаты, сохраните и загрузите её этой кнопкой. '+path)
        except Exception as exc:self.error(exc)

    def edit_data(self):
        """Small tabbed editor; only used AKZ variants, no developer fields."""
        if self.active:return
        try:
            v=model.visible_vom(self.vom.get());profile=model.profile(self.cfg,self.project.get())
            folder=local_root()/'input-drafts'/uuid4().hex;folder.mkdir(parents=True)
            path=folder/'Ведомость.xlsx'
            base=self.draft or self.fresh_draft(v)
            control_book.write_book(path,self.project.get(),self.number.get(),v,self.current_preferences(),base)
            from openpyxl import load_workbook,Workbook
            source=load_workbook(path);window=tk.Toplevel(self.root);window.title('Данные текущей отгрузки — '+self.project.get());self.fit(window,1040,680)
            window.transient(self.root);window.grab_set();window.bind('<Escape>',lambda e:window.destroy())
            ttk.Label(window,text='Изменение ячейки — двойной щелчок. '+('Данные взяты из ведомости текущей отгрузки.' if self.draft else 'Реквизиты, назначения, климат и диапазоны — изменяемые подсказки выбранного проекта. Новые замеры и протоколы не подставляются.'),wraplength=980).pack(padx=10,pady=8)
            notebook=ttk.Notebook(window);notebook.pack(fill='both',expand=True,padx=10);trees={}
            required=set()
            from .word_controls import inspect_word_fields
            for role in ('vik','uzk','quality'):a,b=inspect_word_fields(profile['templates'][role]['path']);required|=(a|b)&set(DOC_FIELD_LABELS)
            used={r['akz'] for r in base['marks'] if type(r.get('akz')) is int}
            for role,ref in profile['templates'].items():
                if not role.startswith('akz') or 'appendix' in role:continue
                if int(role[-1]) not in used:continue
                from zipfile import ZipFile
                import re
                with ZipFile(ref['path']) as z:
                    required|={n for name in z.namelist() if name.endswith('.xml') for n in re.findall(r'\{\{([A-Z_0-9]+)\}\}',z.read(name).decode('utf8'))}&set(DOC_FIELD_LABELS)
            for sheet in source:
                frame=ttk.Frame(notebook);notebook.add(frame,text=sheet.title)
                values=list(sheet.values);columns=[str(n) for n in range(len(values[0]))]
                tree=ttk.Treeview(frame,columns=columns,show='headings');trees[sheet.title]=tree
                for i,label in enumerate(values[0]):tree.heading(str(i),text=label);tree.column(str(i),width=200 if len(columns)<4 else 145)
                for row in values[1:]:
                    if sheet.title=='Реквизиты' and row[0] not in {DOC_FIELD_LABELS[k] for k in required}:continue
                    tree.insert('','end',values=[control_book.text(x) for x in row])
                sy=ttk.Scrollbar(frame,orient='vertical',command=tree.yview);sx=ttk.Scrollbar(frame,orient='horizontal',command=tree.xview)
                tree.configure(yscrollcommand=sy.set,xscrollcommand=sx.set);sy.pack(side='right',fill='y');sx.pack(side='bottom',fill='x');tree.pack(fill='both',expand=True)
                self.editable(tree,immutable={0} if sheet.title in ('Комплект','Реквизиты','Контроль') else set())
                if sheet.title=='Соединения УЗК':ttk.Button(frame,text='Добавить соединение',command=lambda t=tree:t.insert('','end',values=['']*len(t['columns']))).pack()
            source.close();proofs=deepcopy(self.draft.get('evidence',[]))
            def add_proof():
                for f in filedialog.askopenfilenames(title='Протоколы текущего контроля'):proofs.append({'path':str(Path(f).resolve()),'sha256':sha(f)})
                proof_label.configure(text='Приложено протоколов: '+str(len(proofs)))
            proof_label=ttk.Label(window,text='Приложено протоколов: '+str(len(proofs)));proof_label.pack()
            ttk.Button(window,text='Приложить протоколы…',command=add_proof).pack()
            def variants():
                marks=trees['Контроль'];idx=list(control_book.MARKS).index('akz');used=sorted({int(marks.item(i,'values')[idx]) for i in marks.get_children() if marks.item(i,'values')[idx]!=''})
                for name,keys in [('Климат',control_book.CLIMATE),('Слои АКЗ',control_book.LAYERS)]:
                    t=trees[name];existing={tuple(t.item(i,'values')[:2 if name=='Слои АКЗ' else 1]):t.item(i,'values') for i in t.get_children()}
                    for item in t.get_children():t.delete(item)
                    for n in used:
                        for layer in ((1,2,3) if name=='Слои АКЗ' else (None,)):
                            identity=(str(n),str(layer)) if layer else (str(n),)
                            t.insert('','end',values=existing.get(identity,list(identity)+['']*(len(keys)-len(identity))))
            ttk.Button(window,text='Показать только назначенные варианты АКЗ',command=variants).pack(pady=3)
            def accept():
                try:
                    if not proofs:raise Stop('Приложите первичные протоколы текущего контроля. Старые акты не подходят.')
                    output=Workbook();output.remove(output.active)
                    for name,t in trees.items():
                        s=output.create_sheet(name);s.append([t.heading(c)['text'] for c in t['columns']])
                        for item in t.get_children():
                            row=list(t.item(item,'values'))
                            for i,value in enumerate(row):
                                cell=s.cell(s.max_row+1 if i==0 else s.max_row,i+1,value)
                                if isinstance(value,str):cell.data_type='s'
                    edited=folder/'Введённые данные.xlsx';output.save(edited);output.close()
                    d=control_book.read_book(edited,self.project.get(),self.number.get(),self.cfg.get('training',False),self.current_preferences());d['evidence']=proofs
                    self.draft=d;self.approved.set(False);self.invalidate();self.data_summary.set(f'Введено {len(d["marks"])} марок. Подтвердите актуальность данных перед формированием.');window.destroy()
                except Exception as exc:self.error(exc)
            ttk.Button(window,text='Сохранить данные текущего комплекта',command=accept).pack(pady=8)
        except Exception as exc:self.error(exc)

    def primary_action(self):
        if self.active:return
        if self.phase=='review':return self.authorize()
        self.cancel_requested=False
        try:
            if not self.draft:raise Stop('Добавьте ведомость текущего контроля Excel или введите недостающие данные.')
            if not self.approved.get():raise Stop('Подтвердите данные текущих протоколов и свои полномочия флажком над кнопкой.')
            p=model.select(model.read_registry(self.cfg['registry']),self.project.get())
            if p.next_number!=self.number.get():self.number.set(p.next_number);raise Stop('Реестр обновился. Предложен актуальный номер '+p.next_number+'. Проверьте ведомость и повторите формирование.')
            if self.draft.get('number')!=p.next_number:raise Stop('Загрузите ведомость для актуального номера '+p.next_number+'. Ранее загруженные результаты не переназначаются автоматически.')
            self.facts=deepcopy(self.draft)|{'project':self.project.get(),'number':p.next_number,'date':control_book.text(self.day.get(),'date'),'responsible':self.operator.get(),
                     'kind':'TRAINING' if self.cfg.get('training') else 'ACTUAL','confirmed':True,'authorized':True,'forms_approved':True}
            from .facts import validate
            validate(self.facts,model.visible_vom(self.vom.get()),self.project.get(),self.number.get(),self.cfg.get('training',False))
            self.run('check',{'cfg':self.cfg,'project':self.project.get(),'number':self.number.get(),'vom':self.vom.get(),'facts':self.facts,'new':True},self.checked)
        except Exception as exc:self.error(exc)

    def checked(self,result):
        self.plan=result['plan'];self.run('prepare',{'path':self.plan},lambda r:self.run('verify',{'path':self.plan},lambda v:self.run('preview',{'path':self.plan},self.show_review)))

    def show_review(self,result):
        self.phase='review';self.preview_path=result['pdf'];self.opened=False;self.reviewed.set(False)
        self.data_frame.pack_forget();self.approve_widget.pack_forget()
        d=read(self.plan);self.confirm_text.configure(text=f"{d['project']} • № {d['number']} от {d['date']}\n{d['quantity']} изделий / {float(d['mass_kg']):.2f} кг / {float(d['area_m2']):.2f} м²\nПапка: {d['target']}\nРеестр: {d['registry']}\nPDF: {result['pages']} страниц. Резервная копия создаётся перед записью. Откройте PDF и просмотрите все страницы.")
        self.confirm_frame.pack(fill='x',padx=16,pady=8,before=self.primary.master)
        self.primary.configure(text='Подтвердить выпуск и запись',state='normal');self.progress.configure(text='Пакет подготовлен. Реестр ещё не изменён. Требуется ваш просмотр и подтверждение.')

    def authorize(self):
        if not self.approved.get() or not self.opened or not self.reviewed.get():self.error('Подтвердите данные протоколов, откройте PDF, просмотрите весь комплект и отметьте подтверждение просмотра.');return
        self.run('confirm',{'path':self.plan,'responsible':self.facts['responsible'],'reviewed':True},lambda r:self.run('release',{'path':self.plan},self.released))

    def completed(self,result):
        self.result_project=self.project.get();self.final_folder=result['folder'];self.final_pdf=result['print_pdf'];self.write('ГОТОВО. '+result['folder'])
        if not self.cfg.get('training'):
            try:
                name=self.project.get();p=deepcopy(self.cfg['profiles'][name]);p['preferences']=control_book.preferences(self.facts,p.get('preferences'));p['last_vom']=self.vom.get()
                self.cfg['profiles'][name]=p;self.cfg['operator']=self.operator.get();self.cfg['last_project']=name;save(local_root()/'settings.json',self.cfg)
                akz_presets.remember(name,self.facts.get('climates',[]),self.facts.get('akz_date'))
            except Exception as exc:self.write('Комплект сохранён; параметры проекта не сохранены: '+str(exc))
        self.invalidate();self.draft={};self.facts={};self.approved.set(False);self.refresh()
        self.progress.configure(text='ГОТОВО. Word, Excel и общий PDF сохранены. Реестр повторно проверен.')
        self.data_summary.set('Комплект завершён. Для новой отгрузки добавьте её фактические результаты.')
        if self.cancel_requested:self.progress.configure(text='Отмена получена после начала записи. Запись и контроль безопасно завершены; результат сохранён.')
        if self.close_when_safe:self.root.destroy()

    def cancel(self):
        self.cancel_requested=True
        self.approved.set(False)
        if self.active:
            self.progress.configure(text='Отмена принята. Дожидаемся безопасного завершения текущей операции Office.');return
        self.invalidate();self.progress.configure(text='Формирование отменено. Новый выпуск не запускался.')

    def close(self):
        if self.active:
            self.close_when_safe=True;self.cancel();return
        self.root.destroy()

    def after_stage(self,task,callback,result):
        if self.cancel_requested and task not in ('commit','readback','complete'):
            self.invalidate();self.progress.configure(text='Операция отменена до записи. Реестр не изменён.')
            if self.close_when_safe:self.root.destroy()
            return
        callback(result)

    def open_pdf(self):
        if self.active:self.progress.configure(text='Дождитесь подготовки текущего комплекта.');return
        path=self.preview_path if self.phase=='review' else self.final_pdf
        if not path:self.error('PDF появится после формирования и проверки комплекта.');return
        try:os.startfile(path);self.opened=True
        except Exception as exc:self.error('Не удалось открыть PDF. Выберите программу просмотра PDF в Windows. '+str(exc))

    def open_folder(self):
        path=(str(Path(self.preview_path).parent) if self.phase=='review' and self.preview_path else
              self.final_folder or (str(Path(self.plan).parent/'documents') if self.plan else self.cfg.get('output')))
        if path and Path(path).is_dir():os.startfile(path)

    def settings_menu(self):
        if self.active:return
        menu=tk.Menu(self.root,tearoff=False)
        menu.add_command(label='Рабочий реестр и папки',command=self.settings)
        menu.add_command(label='Шаблоны выбранного проекта',command=self.profile)
        menu.add_separator();menu.add_command(label='Незавершённые операции / восстановление',command=self.recovery)
        menu.add_command(label='Открыть журнал',command=lambda:os.startfile(local_root()))
        menu.tk_popup(self.root.winfo_pointerx(),self.root.winfo_pointery())

    def error(self,exc):
        detail=str(exc);folder=local_root()/'diagnostics';folder.mkdir(parents=True,exist_ok=True)
        with (folder/'application.log').open('a',encoding='utf8') as log:log.write(detail+'\n'+traceback.format_exc()+'\n')
        if isinstance(exc,FileNotFoundError):detail='Файл недоступен: '+str(exc.filename)+'. Проверьте его расположение в настройках или выберите файл заново.'
        elif re.search(r'(?<!\d)-214\d{7}(?!\d)|\bcom_error\b',detail,re.I):detail='Microsoft Word или Excel не выполнил операцию. Закройте окно редактирования нужной книги или диалог Office и повторите. Чужие процессы программа не закрывает.'
        else:detail=explain_error(detail)
        self.progress.configure(text='Операция остановлена. '+detail[:200]);self.write(detail)
        messagebox.showerror('Не удалось сформировать комплект',detail+'\n\nПодробности сохранены в журнале программы.')

    def run(self,task,request,callback):
        super().run(task,request,lambda result:self.after_stage(task,callback,result))
        for widget in self.lockable:widget.configure(state='disabled')
        labels={'check':'Проверяем исходные данные','prepare':'Формируем Word и Excel','verify':'Проверяем комплект','preview':'Готовим PDF для просмотра','confirm':'Сохраняем подтверждение','release':'Создаём выпускной PDF','commit':'Сохраняем комплект и резервную копию','readback':'Повторно проверяем реестр','complete':'Завершаем выпуск'}
        self.progress.configure(text=labels.get(task,'Выполняется операция')+'. Дождитесь завершения Office.')

    def poll(self,root,task,callback,states):
        super().poll(root,task,callback,states)
        try:
            if self.active is None and self.root.winfo_exists():
                for widget in self.lockable:widget.configure(state='normal')
        except tk.TclError:pass  # normal close after a safely completed worker

def launch():
    root=tk.Tk();ttk.Style(root).theme_use('clam');DailyApp(root);root.mainloop()
