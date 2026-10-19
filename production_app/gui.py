import os,subprocess,sys,traceback
from pathlib import Path
from uuid import uuid4
import tkinter as tk
from tkinter import ttk,filedialog,messagebox
from .storage import local_root,read,save,sha,Stop
from . import model,core,akz_presets,VERSION

class App:
    def __init__(self,root):
        self.root=root;root.title('ВЗМК — Паспортизация '+VERSION);self.fit(root,1050,760);root.minsize(min(950,root.winfo_screenwidth()-40),min(650,root.winfo_screenheight()-80))
        config_error=None
        try:self.cfg=read(local_root()/'settings.json') if (local_root()/'settings.json').exists() else {'profiles':{}}
        except Exception as exc:self.cfg={'profiles':{}};config_error=str(exc)
        self.plan=None;self.facts={};self.active=None;self.result_path=None;self.buttons=[]
        self.project=tk.StringVar();self.number=tk.StringVar();self.vom=tk.StringVar();self.new=tk.BooleanVar(value=False)
        self.project.trace_add('write',lambda *_:self.inputs_changed())
        self.number.trace_add('write',lambda *_:self.inputs_changed())
        self.vom.trace_add('write',lambda *_:self.inputs_changed())
        top=ttk.Frame(root,padding=12);top.pack(fill='x')
        ttk.Label(top,text='Паспорта металлоконструкций',font=('Segoe UI',17,'bold')).pack(side='left')
        ttk.Button(top,text='Настройки путей',command=self.settings).pack(side='right')
        self.paths=ttk.Label(root,wraplength=1050,justify='left');self.paths.pack(fill='x',padx=14)
        form=ttk.Frame(root,padding=14);form.pack(fill='x');form.columnconfigure(1,weight=1)
        ttk.Label(form,text='Объект из реестра').grid(row=0,column=0,sticky='w')
        self.projects=ttk.Combobox(form,textvariable=self.project,state='readonly');self.projects.grid(row=0,column=1,sticky='ew',padx=8,pady=5)
        self.projects.bind('<<ComboboxSelected>>',lambda e:self.select_project())
        ttk.Button(form,text='Обновить реестр',command=self.refresh).grid(row=0,column=2)
        ttk.Label(form,text='№ паспорта, исходная серия').grid(row=1,column=0,sticky='w')
        self.numbers=ttk.Combobox(form,textvariable=self.number,state='readonly');self.numbers.grid(row=1,column=1,sticky='ew',padx=8,pady=5)
        self.new_widget=ttk.Checkbutton(form,text='Новый паспорт',variable=self.new,command=self.select_project);self.new_widget.grid(row=1,column=2)
        ttk.Label(form,text='Отгрузочная ВОМ').grid(row=2,column=0,sticky='w')
        self.vom_widget=ttk.Entry(form,textvariable=self.vom);self.vom_widget.grid(row=2,column=1,sticky='ew',padx=8,pady=5)
        ttk.Button(form,text='Выбрать ВОМ…',command=lambda:self.choose(self.vom,False)).grid(row=2,column=2)
        tools=ttk.Frame(root,padding=(14,0));tools.pack(fill='x')
        for text,command in [('Шаблоны объекта',self.profile),('Фактические результаты',self.fact_dialog),('Черновик справочника',self.catalog),('Открыть подготовленные документы',self.open_documents),('Восстановление',self.recovery)]:ttk.Button(tools,text=text,command=command).pack(side='left',padx=(0,6),pady=5)
        steps=ttk.Frame(root,padding=14);steps.pack(fill='x')
        for label,command in [('1. Проверить данные',self.check),('2. Подготовить документы',self.prepare),('3. Проверить комплект',self.verify),('4. Подтвердить выпуск и запись',self.issue)]:
            b=ttk.Button(steps,text=label,command=command);b.pack(side='left',expand=True,fill='x',padx=3);self.buttons.append(b)
        self.progress=ttk.Label(root,text='Выберите рабочие пути. Реестр пока не изменяется.');self.progress.pack(fill='x',padx=15)
        ttk.Label(root,text='TEST/SAMPLE_ONLY не являются результатами контроля. Выпуск требует протоколов, проверки Office и отдельного подтверждения.',wraplength=940).pack(side='bottom',padx=14,pady=(0,12))
        self.log=tk.Text(root,wrap='word',font=('Segoe UI',10),state='disabled',height=10);self.log.pack(fill='both',expand=True,padx=14,pady=12)
        root.protocol('WM_DELETE_WINDOW',self.close)
        self.invalidate();self.refresh()
        if not self.cfg.get('registry'):root.after(300,self.settings)
        if config_error:root.after(500,lambda:self.error('Не удалось загрузить настройки. Выберите пути заново. Исходники не изменены. '+config_error))

    def fit(self,window,width,height):window.geometry(f'{min(width,window.winfo_screenwidth()-40)}x{min(height,window.winfo_screenheight()-80)}')

    def scroll_content(self,dialog):
        outer=ttk.Frame(dialog);outer.pack(fill='both',expand=True)
        canvas=tk.Canvas(outer,highlightthickness=0)
        sy=ttk.Scrollbar(outer,orient='vertical',command=canvas.yview)
        sx=ttk.Scrollbar(outer,orient='horizontal',command=canvas.xview)
        canvas.configure(yscrollcommand=sy.set,xscrollcommand=sx.set)
        sy.pack(side='right',fill='y');sx.pack(side='bottom',fill='x');canvas.pack(fill='both',expand=True)
        body=ttk.Frame(canvas);canvas.create_window((0,0),window=body,anchor='nw')
        body.bind('<Configure>',lambda e:canvas.configure(scrollregion=canvas.bbox('all')))
        return body

    def close(self):
        if self.active and self.active.poll() is None:
            messagebox.showinfo('Операция Office','Дождитесь штатного завершения Office. Программа не будет принудительно закрывать процессы.');return
        self.root.destroy()

    def write(self,text):
        self.log.configure(state='normal');self.log.insert('end',str(text)+'\n');self.log.see('end');self.log.configure(state='disabled')

    def invalidate(self):
        self.plan=None
        for n,b in enumerate(self.buttons):b.configure(state='normal' if n==0 and not self.active else 'disabled')

    def inputs_changed(self):
        self.facts={};self.invalidate()

    def error(self,exc):self.write(str(exc));messagebox.showerror('Остановлено — исходники не исправляются автоматически',str(exc))

    def choose(self,var,directory):
        if self.active:return
        path=filedialog.askdirectory() if directory else filedialog.askopenfilename(filetypes=[('Excel/Word','*.xlsx *.xlsm *.docx'),('Все файлы','*.*')])
        if path:var.set(path)

    def settings(self):
        if self.active:return
        dialog=tk.Toplevel(self.root);dialog.title('Рабочие пути — явный выбор');self.fit(dialog,1000,350);dialog.transient(self.root);dialog.grab_set()
        vars={k:tk.StringVar(value=self.cfg.get(k,'')) for k in ('registry','documents','output')}
        for row,(key,label) in enumerate([('registry','Рабочий реестр ВЗМК.xlsx'),('documents','Корень документов'),('output','Папка сохранения паспортов')]):
            ttk.Label(dialog,text=label).grid(row=row,column=0,padx=12,pady=12,sticky='w');ttk.Entry(dialog,textvariable=vars[key],width=85).grid(row=row,column=1)
            ttk.Button(dialog,text='Выбрать…',command=lambda k=key:self.choose(vars[k],k!='registry')).grid(row=row,column=2,padx=8)
        ttk.Label(dialog,text='Никакие пути DEV автоматически не используются. Для испытаний создайте отдельную учебную копию.',wraplength=900).grid(row=3,column=0,columnspan=3,pady=12)
        def store(training=False):
            try:
                cfg=self.cfg|{k:v.get() for k,v in vars.items()}|{'training':False}
                model.validate_paths(cfg)
                if training:
                    parent=filedialog.askdirectory(title='Куда создать НОВУЮ учебную копию?')
                    if not parent:return
                    cfg=core.training_copy(cfg,parent)
                self.cfg=cfg;save(local_root()/'settings.json',cfg);self.invalidate();self.facts={};dialog.destroy();self.refresh()
            except Exception as exc:self.error(exc)
        ttk.Button(dialog,text='Сохранить выбранные рабочие пути',command=store).grid(row=4,column=1,sticky='w')
        ttk.Button(dialog,text='Создать и выбрать учебную копию',command=lambda:store(True)).grid(row=4,column=1,sticky='e')

    def refresh(self):
        if self.active:return
        self.paths.configure(text=('УЧЕБНАЯ КОПИЯ\n' if self.cfg.get('training') else 'РАБОЧИЙ РЕЖИМ — запись только после подтверждения\n')+'\n'.join(f'{label}: {self.cfg.get(k,"не выбран")}' for k,label in [('registry','Реестр'),('documents','Документы'),('output','Результат')]))
        if not self.cfg.get('registry'):return
        try:
            r=model.read_registry(self.cfg['registry']);self.projects['values']=sorted(r.projects);self.invalidate()
            self.progress.configure(text='Реестр прочитан. Выберите объект; запись не производилась.')
            if self.project.get():self.select_project()
        except Exception as exc:self.error(exc)

    def select_project(self):
        if self.active:return
        self.invalidate();self.facts={}
        try:
            r=model.read_registry(self.cfg['registry']);name=self.project.get();entries=r.projects[name]
            self.numbers['values']=[str(e['number']) for e in entries]
            self.number.set(str(entries[0]['number']))
            p=model.select(r,name)
            if self.new.get():self.numbers['values']=[p.next_number];self.number.set(p.next_number)
            self.write(f'{name}: первый {p.first_number}; последний {p.last_number}; следующий {p.next_number}')
            aliases=model.approved_legacy_aliases(name,entries)
            if aliases:
                self.write('ВНИМАНИЕ: в реестре есть подтверждённые опечатки: '
                           +'; '.join(f"строка {e['row']}: {e['number']} (для расчёта серии {alias})"
                                     for e,alias in aliases)
                           +'. Исходные записи НЕ изменяются.')
        except Exception as exc:self.error(exc)

    def profile(self):
        if not self.project.get() or self.active:return
        name=self.project.get();old=self.cfg.get('profiles',{}).get(name,{})
        d=tk.Toplevel(self.root);d.title('Собственные источники: '+name);self.fit(d,1070,720);d.transient(self.root)
        d.grab_set();d.bind('<Escape>',lambda e:d.destroy())
        root=tk.StringVar(value=old.get('root',''));fields={}
        ttk.Label(d,text='Папка только выбранного объекта').grid(row=0,column=0,padx=8,pady=8)
        ttk.Entry(d,textvariable=root,width=85).grid(row=0,column=1);ttk.Button(d,text='Выбрать',command=lambda:self.choose(root,True)).grid(row=0,column=2)
        roles=['appendix','vik','uzk','quality']+[r+str(n) for n in range(4) for r in ('akz','akz_appendix')]
        for row,role in enumerate(roles,1):
            var=tk.StringVar(value=old.get('templates',{}).get(role,{}).get('path',''));fields[role]=var
            label={'appendix':'Приложение №1','vik':'ВИК','uzk':'УЗК','quality':'Документ о качестве'}.get(role,('Приложение АКЗ №' if role.startswith('akz_appendix') else 'Акт АКЗ №')+role[-1])
            ttk.Label(d,text=label).grid(row=row,column=0,padx=8,pady=5,sticky='w');ttk.Entry(d,textvariable=var,width=85).grid(row=row,column=1)
            ttk.Button(d,text='Файл…',command=lambda v=var:self.choose(v,False)).grid(row=row,column=2)
        catalog=tk.StringVar(value=old.get('catalog',''))
        ttk.Label(d,text='Справочник марок (если есть)').grid(row=13,column=0,padx=8,sticky='w');ttk.Entry(d,textvariable=catalog,width=85).grid(row=13,column=1);ttk.Button(d,text='Файл…',command=lambda:self.choose(catalog,False)).grid(row=13,column=2)
        ttk.Label(d,text='appendix — первое приложение объекта. Word: утверждённые поля {{NUMBER}}, {{OBJECT}} и др.; УЗК: {{UZK_TABLE}}.\nСтарые акты с готовыми результатами не превращаются в новые подтверждённые результаты автоматически.',wraplength=990).grid(row=14,column=0,columnspan=3,pady=8)
        def save_profile():
            try:
                p={'root':str(model.contained(root.get(),self.cfg['documents'])),'templates':{}}
                for role,var in fields.items():
                    if var.get():
                        f=model.contained(var.get(),p['root']);p['templates'][role]={'path':str(f),'sha256':sha(f)}
                if catalog.get():p['catalog']=str(model.contained(catalog.get(),p['root']));p['catalog_sha256']=sha(p['catalog'])
                cfg=self.cfg|{'profiles':self.cfg.get('profiles',{})|{name:p}}
                model.profile(cfg,name)
                from .templates import validate
                project=model.select(model.read_registry(cfg['registry']),name)
                variants=sorted({int(role[-1]) for role in p['templates'] if role.startswith('akz')})
                validate(p,project,variants,cfg.get('training',False))
                self.cfg=cfg;save(local_root()/'settings.json',cfg);self.facts={};self.invalidate();d.destroy()
            except Exception as exc:self.error(exc)
        def forms():
            if not root.get():self.error('Сначала выберите папку объекта.');return
            folder=Path(root.get())/('Формы_на_утверждение_'+uuid4().hex[:8])
            if not messagebox.askyesno('Новые формы','Создать отдельные незаполненные формы?\n'+str(folder)+'\nИх необходимо утвердить ответственному лицу перед рабочим применением.'):return
            d.destroy();self.run('forms',{'folder':str(folder)},lambda result:self.write('Созданы формы без результатов контроля: '+result['folder']+'. Выберите их в настройках объекта; первое приложение выберите из собственных источников.'))
        ttk.Button(d,text='Создать отдельные формы на утверждение',command=forms).grid(row=15,column=1,sticky='w',pady=10)
        ttk.Button(d,text='Сохранить привязки и SHA-256',command=save_profile).grid(row=15,column=1,sticky='e')

    def fact_dialog(self):
        if self.active:return
        try:v=model.visible_vom(self.vom.get())
        except Exception as exc:self.error(exc);return
        presets={'saved_at':'','variants':{}}
        if not self.cfg.get('training'):
            try:presets=akz_presets.last_for_project(self.project.get())
            except Stop as exc:
                # A damaged optional cache must never block new ACTUAL data.
                self.write('Подсказки АКЗ не загружены (ввод доступен вручную): '+str(exc))
        d=tk.Toplevel(self.root);d.title('Фактические результаты — '+self.project.get()+' '+self.number.get());self.fit(d,1160,760)
        d.transient(self.root);d.grab_set()
        window=d;window.bind('<Escape>',lambda e:window.destroy());d=self.scroll_content(window)
        variables={k:tk.StringVar(value=self.facts.get(k,'')) for k in ('responsible','start_date','date','vik_date','uzk_date','akz_date')}
        top=ttk.Frame(d);top.pack(fill='x',padx=10,pady=8)
        for n,(key,label) in enumerate([('responsible','Ответственный'),('start_date','Начало YYYY-MM-DD'),('date','Окончание'),('vik_date','ВИК'),('uzk_date','УЗК'),('akz_date','АКЗ')]):
            ttk.Label(top,text=label).grid(row=0,column=n);ttk.Entry(top,textvariable=variables[key],width=25 if n==0 else 16).grid(row=1,column=n,padx=5)
        ttk.Label(d,text='Двойной щелчок по ячейке для ввода. Пустой АКЗ остаётся пустым. УЗК: ГОДЕН / НЕ ТРЕБУЕТСЯ; для исключений обязательно основание.').pack(pady=4)
        cols=('mark','profile','paint','akz','uzk_quantity','uzk_result','basis','vik_result','vik_basis')
        table_frame=ttk.Frame(d);table_frame.pack(fill='both',expand=True,padx=10)
        grid=ttk.Treeview(table_frame,columns=cols,show='headings',height=12)
        for c in cols:grid.heading(c,text={'mark':'Марка','profile':'Профиль','paint':'Краска','akz':'№ АКЗ','uzk_quantity':'УЗК, шт','uzk_result':'Результат УЗК','basis':'Основание / протокол УЗК','vik_result':'Результат ВИК','vik_basis':'Протокол ВИК'}.get(c,c));grid.column(c,width=90 if c in ('akz','uzk_quantity') else 160)
        y=ttk.Scrollbar(table_frame,orient='vertical',command=grid.yview);grid.configure(yscrollcommand=y.set);y.pack(side='right',fill='y')
        grid.pack(side='left',fill='both',expand=True)
        old={r['mark']:r for r in self.facts.get('marks',[])}
        for mark in dict.fromkeys(str(r[0]) for r in v.rows):grid.insert('', 'end',values=[mark]+[old.get(mark,{}).get(c,'') if old.get(mark,{}).get(c) is not None else '' for c in cols[1:]])
        self.editable(grid,immutable={0})
        def import_assignments():
            try:
                profile=self.cfg.get('profiles',{}).get(self.project.get(),{});path=profile.get('catalog')
                if not path:raise Stop('У объекта нет привязанного справочника. Создайте черновик и подтвердите назначения.')
                if sha(path)!=profile['catalog_sha256']:raise Stop('Справочник изменился. Повторно подтвердите его в настройках объекта.')
                if not messagebox.askyesno('Подтверждение источника назначений',str(path)+'\nПодтверждаете профиль, краску и АКЗ из этого справочника? Результаты УЗК и замеры АКЗ не заполняются.'):return
                from openpyxl import load_workbook
                w=load_workbook(path,data_only=True)
                try:
                    s=w.active;header=None;aliases={'марка':'mark','отправочная марка':'mark','профиль':'profile','краска':'paint','№ акз':'akz'}
                    for row in s.iter_rows(max_row=10):
                        found={aliases[str(c.value).strip().lower()]:c.column for c in row if str(c.value).strip().lower() in aliases}
                        if 'mark' in found:header=(row[0].row,found);break
                    if header is None:raise Stop('Не найдена таблица справочника.')
                    hr,cols=header;mapping={}
                    for row in range(hr+1,s.max_row+1):
                        mark=s.cell(row,cols['mark']).value
                        if mark is None:continue
                        key=str(mark)
                        if key in mapping:raise Stop('Неоднозначная марка в справочнике: '+key)
                        mapping[key]={c:s.cell(row,n).value for c,n in cols.items()}
                    for item in grid.get_children():
                        values=list(grid.item(item,'values'));source=mapping.get(values[0])
                        if source:
                            for c in ('profile','paint','akz'):
                                if source.get(c) is not None:values[cols_tuple.index(c)]=str(source[c])
                            grid.item(item,values=values)
                finally:w.close()
            except Exception as exc:self.error(exc)
        cols_tuple=cols
        ttk.Button(d,text='Импортировать подтверждённые назначения из справочника',command=import_assignments).pack(pady=3)
        ttk.Label(d,text='АКЗ: один набор климата на три слоя в день. Для трёхслойного акта средние значения сухой толщины вводятся ниже, по этапам 1/2/3.').pack(pady=6)
        if presets['variants']:
            ttk.Label(d,text='Предзаполнены ПОСЛЕДНИЕ показатели выбранного объекта. Дата прежнего контроля: '+
                      ', '.join(sorted({v.get('last_akz_date','') for v in presets['variants'].values() if v.get('last_akz_date')}))+
                      '. Это подсказки: перед новым выпуском сверьте их с фактическими измерениями.',wraplength=1080).pack(pady=4)
        ccols=('variant','air','surface','humidity','dew','thickness','material','basis')
        climates=ttk.Treeview(d,columns=ccols,show='headings',height=4)
        for c in ccols:
            climates.heading(c,text={'variant':'№ АКЗ','air':'Воздух °C','surface':'Поверхн. °C','humidity':'Влажность %','dew':'Роса °C','thickness':'Толщина (старый формат), мкм','material':'Материал','basis':'Послойный протокол'}[c]);climates.column(c,width=130)
        climates.pack(fill='x',padx=10);self.editable(climates,immutable={0})
        prev={int(c['variant']):c for c in self.facts.get('climates',[])}
        defaults=presets['variants']
        for n in range(4):
            fresh=prev.get(n)
            cached=(defaults.get(str(n)) or presets.get('latest',{})).get('climate',{})
            values=(fresh if fresh is not None else cached)
            climates.insert('','end',values=[n]+[values.get(c,'') if c in akz_presets.CLIMATE_FIELDS or fresh is not None else '' for c in ccols[1:]])
        ttk.Label(d,text='АКЗ: допустимый диапазон каждого слоя и фактически измеренная СРЕДНЯЯ СУХАЯ ТОЛЩИНА после слоя (нарастающим итогом). Минимальные и максимальные фактические замеры НЕ требуются.',wraplength=1070).pack(pady=6)
        lcols=('variant','layer','min','max','average')
        layer_frame=ttk.Frame(d);layer_frame.pack(fill='x',padx=10)
        layers_grid=ttk.Treeview(layer_frame,columns=lcols,show='headings',height=7)
        for col,label,width in [('variant','№ АКЗ',80),('layer','Слой',80),('min','От, мкм',160),('max','До, мкм',160),('average','Средняя сухая после слоя, мкм',250)]:
            layers_grid.heading(col,text=label);layers_grid.column(col,width=width)
        layer_scroll=ttk.Scrollbar(layer_frame,orient='vertical',command=layers_grid.yview)
        layers_grid.configure(yscrollcommand=layer_scroll.set)
        layer_scroll.pack(side='right',fill='y');layers_grid.pack(fill='x')
        self.editable(layers_grid,immutable={0,1})
        for n in range(4):
            current={int(x['layer']):x for x in prev.get(n,{}).get('layers',[])}
            observed={int(x['layer']):x for x in prev.get(n,{}).get('dry_averages',[])}
            cached={int(x['layer']):x for x in (defaults.get(str(n)) or presets.get('latest',{})).get('layers',[])}
            for layer in (1,2,3):
                value=current.get(layer,cached.get(layer,{}))
                # Measured averages are NEVER loaded from old project presets.
                layers_grid.insert('','end',values=(n,layer,value.get('min',''),value.get('max',''),
                                                   observed.get(layer,{}).get('average','')))
        proofs=list(self.facts.get('evidence',[]));proof_label=ttk.Label(d,text='Протоколов: '+str(len(proofs)));proof_label.pack()
        def add():
            for f in filedialog.askopenfilenames(title='Первичные протоколы фактического контроля'):
                proofs.append({'path':str(Path(f).resolve()),'sha256':sha(f)})
            proof_label.configure(text='Протоколов: '+str(len(proofs)))
        ttk.Button(d,text='Приложить первичные протоколы…',command=add).pack()
        approved=tk.BooleanVar(value=self.facts.get('confirmed',False));rights=tk.BooleanVar(value=self.facts.get('authorized',False));forms=tk.BooleanVar(value=self.facts.get('forms_approved',False))
        for variable,text in [(approved,'Подтверждаю фактические результаты; это не TEST/SAMPLE_ONLY'),(rights,'Подтверждаю полномочия ответственного исполнителя'),(forms,'Формы документов утверждены ответственным лицом')]:ttk.Checkbutton(d,text=text,variable=variable).pack(anchor='w',padx=12)
        def store():
            try:
                data={k:var.get() for k,var in variables.items()}|{'project':self.project.get(),'number':self.number.get(),'kind':'TRAINING' if self.cfg.get('training') else 'ACTUAL','confirmed':approved.get(),'authorized':rights.get(),'forms_approved':forms.get(),'evidence':proofs,'marks':[],'climates':[]}
                for item in grid.get_children():
                    row=dict(zip(cols,grid.item(item,'values')))
                    row['akz']=int(row['akz']) if row['akz']!='' else None;data['marks'].append(row)
                ranges_by_variant={n:[] for n in range(4)}
                averages_by_variant={n:[] for n in range(4)}
                for item in layers_grid.get_children():
                    row=dict(zip(lcols,layers_grid.item(item,'values')))
                    if any(str(row[k]).strip() for k in ('min','max','average')):
                        n=int(row['variant'])
                        ranges_by_variant[n].append({'layer':int(row['layer']),'min':row['min'],'max':row['max']})
                        averages_by_variant[n].append({'layer':int(row['layer']),'average':row['average']})
                used={r['akz'] for r in data['marks'] if r['akz'] is not None}
                for item in climates.get_children():
                    row=dict(zip(ccols,climates.item(item,'values')))
                    n=int(row['variant'])
                    if n not in used:continue
                    row['variant']=n
                    if ranges_by_variant[n]:
                        row['layers']=ranges_by_variant[n]
                        row['dry_averages']=averages_by_variant[n]
                    data['climates'].append(row)
                from .facts import validate
                validate(data,v,self.project.get(),self.number.get(),self.cfg.get('training',False))
                self.facts=data;self.invalidate();window.destroy();self.write('Фактические данные сохранены локально в текущей заявке. Исходники не изменены.')
            except Exception as exc:self.error(exc)
        ttk.Button(d,text='Проверить и принять введённые данные',command=store).pack(pady=10)

    def editable(self,tree,immutable=set()):
        def edit(event):
            item=tree.identify_row(event.y);column=tree.identify_column(event.x)
            if not item or not column:return
            col=int(column[1:])-1
            if col in immutable:return
            x,y,w,h=tree.bbox(item,column);entry=ttk.Entry(tree);entry.place(x=x,y=y,width=w,height=h)
            entry.insert(0,tree.item(item,'values')[col]);entry.focus_set()
            def done(e=None):
                if not entry.winfo_exists():return
                values=list(tree.item(item,'values'));values[col]=entry.get();tree.item(item,values=values);entry.destroy()
            entry.bind('<Return>',done);entry.bind('<FocusOut>',done)
        tree.bind('<Double-1>',edit)

    def catalog(self):
        if self.active:return
        target=filedialog.asksaveasfilename(title='НОВЫЙ черновик справочника',defaultextension='.xlsx',initialfile='Черновик справочника марок.xlsx')
        if target:self.run('catalog',{'vom':self.vom.get(),'target':target},lambda r:self.write('Черновик: '+r['path']+'; профиль/краска/АКЗ пустые.'))

    def check(self):
        if not all(self.cfg.get(k) for k in ('registry','documents','output')):self.error('Сначала выберите рабочие пути в «Настройки путей».');return
        if not self.project.get() or not self.number.get() or not self.vom.get():self.error('Выберите объект, паспорт и отгрузочную ВОМ.');return
        if not self.facts:self.error('Введите и подтвердите фактические результаты контроля. TEST-данные для рабочего выпуска запрещены.');return
        self.run('check',{'cfg':self.cfg,'project':self.project.get(),'number':self.number.get(),'vom':self.vom.get(),'facts':self.facts,'new':self.new.get()},self.checked)
    def checked(self,result):
        self.plan=result['plan'];self.buttons[1].configure(state='normal');d=result['data']
        self.write(f"PASS. {d['project']} | {d['number']} | {d['marks']} марок / {d['quantity']} изделий / {d['mass_kg']} кг\nРеестр: {d['registry']}\nНовая папка: {d['target']}\nАКЗ: {d['variants']}. Источники закреплены SHA-256.")
    def prepare(self):self.run('prepare',{'path':self.plan},lambda r:(self.write('Подготовлено файлов: '+str(r['files'])+'; '+r['folder']),self.buttons[1].configure(state='disabled'),self.buttons[2].configure(state='normal')))
    def verify(self):self.run('verify',{'path':self.plan},lambda r:(self.write('Комплект PASS. Просмотрите каждый файл перед подтверждением.\n'+str(r['documents'])),self.buttons[3].configure(state='normal')))
    def open_documents(self):
        if self.plan:
            folder=Path(self.plan).parent/'documents'
            if folder.is_dir():os.startfile(folder)
    def issue(self):
        if not self.plan:return
        d=read(self.plan)
        message=f"{'УЧЕБНАЯ КОПИЯ' if self.cfg.get('training') else 'РАБОЧИЙ ВЫПУСК'}\nОбъект: {d['project']}\nПаспорт: {d['number']}\nРеестр: {d['registry']}\nРезервная копия: {Path(d['registry']).parent/'.vzmk-operations'/d['id']/'backup.xlsx'}\nРезультат: {d['target']}\n"+('Будет добавлена ОДНА новая строка.' if d['new'] else 'Существующий паспорт: строки реестра не добавляются.')+'\n\nПодтверждаете, что просмотрели ВСЕ документы, фактические результаты и полномочия ответственного?'
        if not messagebox.askyesno('Отдельное подтверждение выпуска',message):return
        self.run('confirm',{'path':self.plan,'responsible':self.facts['responsible'],'reviewed':True},lambda r:self.run('release',{'path':self.plan},self.released))
    def released(self,r):self.run('commit',{'path':self.plan,'release_result':self.result_path},self.committed)
    def committed(self,r):self.run('readback',{'journal':r['journal']},lambda v:self.run('complete',{'journal':r['journal'],'readback_result':self.result_path},self.completed))
    def completed(self,r):
        self.write('ГОТОВО. '+str(r))
        # Store preferences only for an actually completed production issue.
        # Old protocols, measurement thicknesses, dates and approvals never
        # become facts of the next passport.
        if r.get('state')=='ГОТОВО' and not self.cfg.get('training') and self.facts.get('kind')=='ACTUAL':
            try:
                if akz_presets.remember(self.project.get(),self.facts.get('climates',[]),self.facts.get('akz_date')):
                    self.write('Последние климатические показатели и диапазоны слоёв сохранены только для объекта '+self.project.get()+'.')
            except Exception as exc:
                self.write('Выпуск завершён, но сохранить подсказки АКЗ не удалось: '+str(exc))
        messagebox.showinfo('Проверенный результат',str(r));self.invalidate();self.refresh()
    def recovery(self):
        if self.active:return
        try:
            paths=core.pending(self.cfg['registry'])
            if not paths:messagebox.showinfo('Восстановление','Незавершённых операций нет.');return
            journal=str(paths[0])
            if messagebox.askyesno('Откат незавершённой операции',journal+'\nВосстановить только при совпадении SHA? Изменённые пользователем документы/реестр не будут перезаписаны.'):
                self.run('recover',{'journal':journal},lambda r:self.write(r))
        except Exception as exc:self.error(exc)

    def run(self,task,request,callback):
        if self.active:return
        root=local_root()/'gui-jobs'/uuid4().hex;root.mkdir(parents=True)
        save(root/'request.json',request)
        prefix=[sys.executable] if getattr(sys,'frozen',False) else [sys.executable,'-B',str(Path(__file__).resolve().parents[1]/'production.py')]
        self.active=subprocess.Popen(prefix+['--task',task,'--request',str(root/'request.json'),'--result',str(root/'result.json')],creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        for widget in (self.projects,self.numbers,self.new_widget,self.vom_widget):widget.configure(state='disabled')
        states=[b['state'] for b in self.buttons]
        for b in self.buttons:b.configure(state='disabled')
        self.progress.configure(text='Выполняется: '+task+'. Дождитесь штатного закрытия Office.');self.root.after(500,lambda:self.poll(root,task,callback,states))
    def poll(self,root,task,callback,states):
        if self.active.poll() is None:self.root.after(500,lambda:self.poll(root,task,callback,states));return
        code=self.active.returncode;self.active=None
        self.projects.configure(state='readonly');self.numbers.configure(state='readonly');self.new_widget.configure(state='normal');self.vom_widget.configure(state='normal')
        try:
            r=read(root/'result.json');self.result_path=str(root/'result.json')
            if code or r['status']!='PASS':raise Stop(str(r.get('error',r))+'\nЖурнал: '+str(root))
            for b,state in zip(self.buttons,states):b.configure(state=state)
            self.progress.configure(text='PASS: '+task);callback(r)
        except Exception as exc:self.invalidate();self.progress.configure(text='FAIL: '+task);self.error(exc)

def launch():
    root=tk.Tk();ttk.Style(root).theme_use('clam');App(root);root.mainloop()
