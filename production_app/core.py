from pathlib import Path
from dataclasses import replace
from datetime import date,datetime,timezone
from uuid import uuid4
from decimal import Decimal
import hashlib,os,shutil
from copy import copy
from passport_app.registry import read_registry,append_row
from passport_app.excel import excel_session,open_excel,locked_file,atomic_replace,load_book
from . import model,facts as factual,office
from .storage import Stop,safe,contained,sha,digest,save,read,local_root,pins,verify_pins,mutex
from stage8_closure.storage import frozen

def journal_write(path,state):save(path,state|{'seal':digest(state)})

def journal_read(path):
    d=read(path);seal=d.pop('seal',None)
    if digest(d)!=seal:raise Stop('Журнал повреждён, автоматическая запись/восстановление запрещены: '+str(path))
    return d

def training_copy(cfg,parent):
    """A dedicated, explicitly chosen copy. Never infer production destinations."""
    source,documents,output=model.validate_paths(cfg)
    root=safe(parent)/('Учебный_'+uuid4().hex[:8]);root.mkdir(exist_ok=False)
    target=root/'COPY_ВЗМК.xlsx'
    with locked_file(source) as payload:target.write_bytes(payload)
    mark={'schema':'VZMK_TRAINING_COPY_1','registry':str(target),'source':str(source),'original_sha256':sha(source),'id':root.name}
    save(root/'TRAINING_COPY.json',mark)
    # Preserve the separately selected historical source root READ ONLY.
    # Before training, `output` is a user-selected folder containing the
    # original passports; after creating the sandbox, output points to TEMP.
    archive=cfg.get('first_passport_documents') or cfg['output']
    additions={}
    if cfg.get('use_shared_uzk'):
        from .shared_uzk import find_shared_uzk
        source_uzk=find_shared_uzk(cfg);copy_uzk=root/'УЗК.xlsm'
        with locked_file(source_uzk) as payload:copy_uzk.write_bytes(payload)
        additions['uzk_registry']=str(copy_uzk)
        name=cfg.get('last_project')
        if name:
            original_profile=model.profile(cfg,name);catalog=original_profile.get('catalog')
            if catalog:
                private=root/'Accounting'/name;private.mkdir(parents=True)
                destination=private/'Справочник марок.xlsx'
                with locked_file(catalog) as payload:destination.write_bytes(payload)
                additions['accounting_catalogs']={name:str(destination)}
    return cfg|{'registry':str(target),'output':str(root),
                'first_passport_documents':str(archive),'training':True,**additions}

def training_guard(cfg):
    if not cfg.get('training'):return
    p=safe(cfg['registry']);mark=read(p.parent/'TRAINING_COPY.json')
    if p.name!='COPY_ВЗМК.xlsx' or mark.get('registry')!=str(p) or safe(mark['source'])==p:raise Stop('Учебный режим разрешён только на отдельной созданной программой копии реестра.')
    if safe(cfg['output'])!=p.parent:raise Stop('Учебные документы должны оставаться в папке учебной копии.')

def pending(registry):
    root=safe(registry).parent/'.vzmk-operations'
    return [p for p in root.glob('*/journal.json') if journal_read(p).get('phase') not in ('COMPLETE','ROLLED_BACK','ABORTED')]

def plan(cfg,project,number,vom,facts,new):
    registry,documents,output=model.validate_paths(cfg);training_guard(cfg)
    if pending(registry):raise Stop('Есть незавершённая операция. Откройте «Восстановление»: '+str(pending(registry)[0]))
    r=read_registry(registry);p=model.select(r,project);profile=model.profile(cfg,project)
    if new and number!=p.next_number:raise Stop('Следующий номер изменился: '+p.next_number)
    if not new and number==p.first_number:raise Stop('Первый паспорт оформляется владельцем самостоятельно и не выпускается программой повторно.')
    existing=[e for e in r.projects[project] if e['number']==number]
    if new and existing or not new and len(existing)!=1:raise Stop('Дубль/не найден паспорт: '+number)
    facts={**facts,'marks':[dict(row) for row in facts.get('marks',[])]}
    for row in facts['marks']:
        choice=cfg.get('geometry_choices',{}).get(project,{}).get(str(row['mark']))
        if choice:row['standard']=choice
    v=model.visible_vom(safe(vom))
    if v.area is None:raise Stop('Для Приложения №1 и приложений АКЗ заполните площадь текущей ВОМ. Неизвестная площадь не равна нулю.')
    shared_uzk=None;shared_uzk_sha=None
    if cfg.get('use_shared_uzk'):
        from .shared_uzk import (
            find_shared_uzk,planned_control,verify_planned_quantities,
        )
        shared_uzk=find_shared_uzk(cfg)
        shared_uzk_sha=sha(shared_uzk)
        expected=planned_control(shared_uzk,project,v)
        verify_planned_quantities(facts,expected)
    if not new and abs(v.mass-Decimal(str(existing[0]['mass'])))>Decimal('.005'):
        raise Stop('Масса ВОМ не совпадает с существующим паспортом.')
    # Derive UZK DESIGN characteristics from current marks and the accepted
    # Stage4 profile geometry when its approved workbook is available.
    # Calculated S/L/frequency/angle NEVER imply inspection GОДЕН.
    profile_geometry_source=None;profile_geometry_sha=None
    from .word_controls import inspect_word_fields
    from .working_facts import selected_form
    effective_uzk=selected_form(profile,'uzk')
    tokens,controls=inspect_word_fields(effective_uzk['path'])
    first_uzk=(profile.get('first_passport_sources',{})
               .get('sources',{}).get('uzk'))
    historical_uzk=False
    if first_uzk:
        from .highlighted_reference import supports_uzk_first
        historical_uzk=supports_uzk_first(first_uzk['path'])
    wants_design=('UZK_DESIGN' in (tokens|controls) or
                  cfg.get('auto_uzk_characteristics') is True or
                  historical_uzk)
    if wants_design:
        from .automatic_acts import (
            find_approved_geometry,calculate_design_characteristics,
        )
        # Only explicitly adapted/approved Word forms opt in to this newer
        # source. Do not change the behavior of existing 383-pass Windows
        # E2E that uses historical five/eight-column forms without the marker.
        profile_geometry_source=find_approved_geometry(cfg,require=True)
        profile_geometry_sha=sha(profile_geometry_source)
        computed=calculate_design_characteristics(
            v,facts,profile_geometry_source,
            expected_sha=cfg.get('profiles_workbook_sha256') or profile_geometry_sha)
        facts={**facts,'uzk_design':computed}
    if facts.get('auto_registry_dates'):
        from .registry_dates import resolve_registry_dates
        # For a new row, today's ending date is a PROPOSAL.
        # Manufacturing START_DATE intentionally stays equal to the
        # FIRST passport of this project, including when H353 is a fixed
        # date. Current VIK/UZK/AKZ dates require valid formulas in the
        # private registry copy; don't copy historical control results.
        # The ending date is displayed for final confirmation.
        chosen=date.fromisoformat(facts['shipment_date']) if new and facts.get('shipment_date') else None
        resolved=resolve_registry_dates(cfg,project,number,v,new,day=chosen)
        facts={**facts,**resolved}
    variants=factual.validate(
        facts,v,project,number,cfg.get('training',False),
        defer_approvals=bool(facts.get('simple_ui')),
        preparing=bool(facts.get('simple_ui') and facts.get('owner_confirmation_required')),
    )
    if shared_uzk is not None and sha(shared_uzk)!=shared_uzk_sha:
        raise Stop('Общий УЗК изменился после расчёта количества. Повторите проверку.')
    accounting=None
    if new:
        from .coordinated import accounting_plan
        accounting=accounting_plan(cfg,project,number,v,facts,profile,r.projects[project])
    for n in variants:
        for role in (f'akz{n}',f'akz_appendix{n}'):
            if role not in profile['templates']:raise Stop('Нет собственного шаблона '+role)
    from .templates import validate as validate_templates
    validate_templates(profile,p,variants,cfg.get('training',False))
    # The historically detailed UZK form requires each real welded-joint row.
    from .word_controls import inspect_word_fields
    from .real_forms import validate_uzk_details
    from .working_facts import selected_form
    effective_uzk=selected_form(profile,'uzk')
    uzk_tokens=set() if effective_uzk.get('exact_reference') else inspect_word_fields(effective_uzk['path'])[0]
    validate_uzk_details(facts,v,required=bool({'UZK_DETAILS_TABLE','UZK_NATIVE_ROW'} & uzk_tokens))
    if 'UZK_NATIVE_ROW' in uzk_tokens:
        from .native_acts import native_rows
        native_rows(facts)
    from .real_forms import require_form_inputs
    require_form_inputs(profile,facts)
    from .real_forms import require_akz_form_inputs
    require_akz_form_inputs(profile,facts,variants)
    # Pin the exact FIRST passport source files alongside inputs. Even the
    # historical acts used for visual provenance (not generation) must not
    # silently change between CHECK and PREPARE. Dedupe pinned paths.
    first_refs=profile.get('first_passport_sources',{}).get('sources',{})
    inputs=list(dict.fromkeys([
        registry,v.path,
        *[t['path'] for t in profile['templates'].values()],
        *[t['path'] for t in first_refs.values()],
        *[e['path'] for e in facts['evidence']],
    ]))
    catalog=profile.get('catalog')
    if catalog and Path(catalog).is_file():inputs.append(catalog)
    if shared_uzk is not None:inputs.append(shared_uzk)
    if accounting:inputs.extend(accounting.get('history_sources',[]))
    if profile_geometry_source is not None:inputs.append(profile_geometry_source)
    # Frozen human entries supply explicitly confirmed assignments when no catalog exists.
    passport_folder=model.output_passport_directory(cfg,project,profile)
    target=passport_folder/(model.filename(number)+' от '+date.fromisoformat(facts['date']).strftime('%d.%m.%Y'))
    safe(target)
    if target.exists():raise Stop('Папка готового комплекта уже существует: '+str(target))
    hashes=pins(inputs)
    if profile_geometry_source is not None and hashes[str(profile_geometry_source)]!=profile_geometry_sha:
        raise Stop('Справочник геометрии изменился после расчёта УЗК. Повторите проверку.')
    if shared_uzk is not None and hashes[str(shared_uzk)]!=shared_uzk_sha:
        raise Stop('Общий УЗК изменился после расчёта количества. Повторите проверку текущих данных.')
    if hashes[str(registry)]!=r.digest or hashes[str(v.path)]!=v.digest:raise Stop('Источник изменился во время проверки.')
    model.profile(cfg,project)
    for path in inputs:
        with locked_file(path):pass
    data={'schema':'VZMK_PRODUCTION_PLAN_1','id':uuid4().hex,'cfg':cfg,'project':project,'number':number,
          'new':bool(new),'facts':facts,'vom':str(v.path),'fingerprints':hashes,'target':str(target),
          'marks':len(v.rows),'quantity':str(v.quantity),'mass_kg':str(v.mass),'area_m2':str(v.area),
          'variants':variants,'registry':str(registry),'date':facts['date'],'accounting':accounting}
    run=local_root()/'jobs'/data['id'];run.mkdir(parents=True,exist_ok=False)
    data['run']=str(run);data['seal']=digest(data)
    save(run/'plan.json',data)
    return {'plan':str(run/'plan.json'),'data':data,'status':'PASS'}

def load_plan(path):
    d=read(path)
    if d.get('schema')!='VZMK_PRODUCTION_PLAN_1':
        raise Stop('Черновой пакет не является планом официального выпуска.')
    seal=d.pop('seal')
    if digest(d)!=seal:raise Stop('План повреждён.')
    d['seal']=seal
    if safe(path)!=local_root()/'jobs'/d['id']/'plan.json':raise Stop('План вне частной рабочей папки.')
    verify_pins(d['fingerprints']);training_guard(d['cfg'])
    return d

def prepare(path):
    d=load_plan(path)
    with frozen(d['fingerprints']):
        verify_pins(d['fingerprints'])
        result=_prepare(d)
        verify_pins(d['fingerprints'])
        return result

def _prepare(d):
    run=safe(d['run']);outputs=run/'documents'
    outputs.mkdir(exist_ok=False)
    registry=Path(d['registry']);snapshot=run/'registry.xlsx'
    payload=registry.read_bytes()
    if hashlib.sha256(payload).hexdigest()!=d['fingerprints'][str(registry)]:raise Stop('SHA реестра изменился.')
    snapshot.write_bytes(payload)
    v=model.visible_vom(d['vom']);r=read_registry(snapshot);p=model.select(r,d['project'])
    if d['new']:
        with excel_session() as app:
            row,_=append_row(app,snapshot,r,p,v,date.fromisoformat(d['date']))
    else:row=next(e['row'] for e in r.projects[d['project']] if e['number']==d['number'])
    # Recalculate formulas only in the private registry.
    with excel_session() as app:
        b=open_excel(app,snapshot)
        try:
            if d['new']:
                col=model.columns(r)['START_DATE'];cell=b.Worksheets(r.sheet).Cells(row,col)
                if not cell.HasFormula:cell.Value2=(date.fromisoformat(d['facts']['start_date'])-date(1899,12,30)).days
            app.CalculateFullRebuild();b.Save()
        finally:b.Close(SaveChanges=False)
    staged=read_registry(snapshot)
    if d['new']:verify_registry(registry,snapshot,d['project'],d['number'],v,date.fromisoformat(d['date']))
    out=office.render(model.profile(d['cfg'],d['project']),p,staged,row,v,d['facts'],outputs,d['cfg'].get('training',False))
    from .coordinated import prepare_candidates
    accounting=prepare_candidates(d)
    save(run/'prepared.json',{'plan_seal':d['seal'],'outputs':out,'output_fingerprints':pins([o['path'] for o in out]),'registry':str(snapshot),'row':row,'snapshot_sha256':sha(snapshot),'accounting':accounting})
    return {'status':'PASS','folder':str(outputs),'files':len(out)}

def verify_registry(original,after,project,number,vom,expected_date=None):
    r0=read_registry(original);r1=read_registry(after)
    entries=[e for e in r1.projects[project] if e['number']==number]
    if len(entries)!=1 or sum(map(len,r1.projects.values()))!=sum(map(len,r0.projects.values()))+1:raise Stop('Реестр должен содержать ровно одну новую строку.')
    e=entries[0]
    if abs(Decimal(str(e['mass']))-vom.mass)>Decimal('.000001') or e['connections']!=vom.connections:raise Stop('Неверные масса/соединения новой строки.')
    if expected_date is not None:
        actual=e['date'].date() if isinstance(e['date'],datetime) else e['date']
        if actual!=expected_date:raise Stop('Дата новой строки не совпала с подтверждённой датой паспорта.')
    a=load_book(original);b=load_book(after)
    try:
        if a.sheetnames!=b.sheetnames:raise Stop('Изменилась структура листов реестра.')
        for sa in a.worksheets:
            sb=b[sa.title]
            def print_ranges(sheet,extend=False):
                from openpyxl.worksheet.print_settings import PrintArea
                ranges=[]
                for area in PrintArea.from_string(str(sheet.print_area)).ranges if sheet.print_area else []:
                    bounds=list(area.bounds)
                    # Excel extends an existing print area ending at the last
                    # table row when ListRows.Add inserts the single new row.
                    if extend and sa.title==r0.sheet and bounds[3]==e['row']-1:bounds[3]+=1
                    ranges.append(tuple(bounds))
                return sorted(ranges)
            if print_ranges(sa,True)!=print_ranges(sb) or sa.print_title_rows!=sb.print_title_rows or sa.print_title_cols!=sb.print_title_cols:
                raise Stop('Изменилась область печати/повторяемые строки реестра: '+sa.title)
            for row in sa:
                for ca in row:
                    # Insertion must not change any pre-existing cell's value/formula/style.
                    cb=sb.cell(ca.row,ca.column)
                    if ca.value!=cb.value or ca.number_format!=cb.number_format or copy(ca.font)!=copy(cb.font) or copy(ca.border)!=copy(cb.border) or copy(ca.alignment)!=copy(cb.alignment) or copy(ca.fill)!=copy(cb.fill) or copy(ca.protection)!=copy(cb.protection):
                        raise Stop('Изменилась исходная ячейка реестра: '+sa.title+'!'+ca.coordinate)
        source=model.select(r0,project)
        s0=a[r0.sheet];s1=b[r1.sheet]
        for col in range(1,r0.width+1):
            c=s0.cell(source.first_row,col);n=s1.cell(e['row'],col)
            if c.data_type=='f' and n.data_type!='f':raise Stop('Утрачена формула новой строки.')
        if len(b[r1.sheet].tables)!=len(a[r0.sheet].tables):raise Stop('Утрачена Excel-таблица.')
        from openpyxl.utils.cell import range_boundaries
        for sa in a.worksheets:
            sb=b[sa.title]
            if set(sa.tables)!=set(sb.tables):raise Stop('Изменён набор Excel-таблиц.')
            for name,table in sa.tables.items():
                old=sa.tables[name];new=sb.tables[name]
                bounds=list(range_boundaries(old.ref))
                if sa.title==r0.sheet and name==r0.table:bounds[3]+=1
                if tuple(bounds)!=range_boundaries(new.ref) or old.tableStyleInfo!=new.tableStyleInfo:
                    raise Stop('Неверное расширение/оформление Excel-таблицы: '+name)
    finally:a.close();b.close()
    return e['row']

def verify(path,preview=False):
    d=load_plan(path);run=Path(d['run']);prepared=read(run/'prepared.json')
    if prepared['plan_seal']!=d['seal'] or sha(prepared['registry'])!=prepared['snapshot_sha256']:
        raise Stop('Изменена подготовленная копия реестра.')
    if not prepared.get('output_fingerprints'):
        raise Stop('Комплект подготовлен предыдущей версией. Повторите подготовку новой заявкой.')
    verify_pins(prepared['output_fingerprints'])
    report=office.verify_documents(prepared['outputs'],d['number'],d['cfg'].get('training',False))
    hashes=pins([o['path'] for o in prepared['outputs']])
    preview_path=None
    if preview:
        from .print_pdf import create_print_pdf
        draft=run/'preview'
        draft.mkdir(exist_ok=False)
        copied=[]
        for item in prepared['outputs']:
            source=Path(item['path'])
            target=draft/source.name
            if target.exists():
                raise Stop('Дублирующийся документ чернового PDF: '+target.name)
            shutil.copyfile(source,target)
            copied.append({'role':item['role'],'path':str(target)})
        # PDF marking is applied after pagination, on every preview page.
        printout=create_print_pdf(copied,draft,Path(d['target']).name,preview=True)
        preview_path=printout['output']['path']
        hashes.update(pins([preview_path]))
    save(run/'verified.json',{
        'plan_seal':d['seal'],'documents':report,'fingerprints':hashes,
        'preview_pdf':preview_path,
    })
    return {'status':'PASS','files':len(report),'documents':report,
            'preview_pdf':preview_path,'ready_to_confirm':True}

def preview(path):
    """Read-only PDF of verified drafts before the operator authorizes release."""
    d=load_plan(path);run=Path(d['run']);prepared=read(run/'prepared.json');verified=read(run/'verified.json')
    if prepared['plan_seal']!=d['seal'] or verified['plan_seal']!=d['seal']:raise Stop('Другой подготовленный комплект.')
    verify_pins(prepared['output_fingerprints']);verify_pins(verified['fingerprints'])
    folder=run/('preview-'+uuid4().hex[:8]);folder.mkdir(exist_ok=False)
    from .print_pdf import create_print_pdf
    copies=[]
    for item in prepared['outputs']:
        target=folder/Path(item['path']).name;shutil.copyfile(item['path'],target)
        if sha(target)!=verified['fingerprints'][item['path']]:raise Stop('Изменилась копия для просмотра.')
        copies.append(item|{'path':str(target)})
    result=create_print_pdf(copies,folder,Path(d['target']).name,preview=True)
    data={'plan_seal':d['seal'],'verified_sha256':sha(run/'verified.json'),
          'pdf':result['output']['path'],'sha256':result['output']['sha256'] if 'sha256' in result['output'] else sha(result['output']['path']),
          'pages':result['manifest']['pages']}
    save(run/'preview.json',data)
    return {'status':'PASS',**data}

def confirm(path,responsible,reviewed,approved=False,current_control=False):
    d=load_plan(path);run=Path(d['run']);verified=read(run/'verified.json')
    if not reviewed or responsible!=d['facts']['responsible']:
        raise Stop('Подтвердите просмотр всех документов и ответственное лицо.')
    if d['facts'].get('simple_ui') and approved is not True:
        raise Stop('Не подтверждены достоверность текущих результатов, '
                   'полномочия и утверждённые формы.')
    verify_pins(verified['fingerprints'])
    if (run/'preview.json').exists():
        shown=read(run/'preview.json')
        if shown['plan_seal']!=d['seal'] or shown['verified_sha256']!=sha(run/'verified.json'):raise Stop('Просмотрен другой комплект. Сформируйте PDF снова.')
        verify_pins({shown['pdf']:shown['sha256']})
    confirmation={'plan_seal':d['seal'],'verified_sha256':sha(run/'verified.json'),'responsible':responsible,'reviewed':True,'approved':approved is True,'created_at':datetime.now(timezone.utc).isoformat()}
    if d['facts'].get('owner_confirmation_required'):
        if current_control is not True:raise Stop('Явно подтвердите выполненный текущий контроль. Старый образец не является подтверждением.')
        from .owner_control import declaration
        confirmation['current_control']=declaration(d,confirmation['verified_sha256'])
    save(run/'confirmation.json',confirmation)
    return {'status':'PASS','confirmation':str(run/'confirmation.json')}

def supervisor_proof(result_path,task,path):
    result=read(result_path)
    if result.get('status')!='PASS' or result.get('task')!=task or result.get('plan')!=str(path):raise Stop('Нет PASS завершённого Office-запуска '+task)
    sup_path=safe(result['office_supervisor']);summary=read(sup_path)
    request=read(sup_path.parent/'JOB_REQUEST.json')
    arguments=request['arguments']
    if '--task' not in arguments or arguments[arguments.index('--task')+1]!=task:raise Stop('Office-доказательство другой операции.')
    if summary.get('owned_orphans')!=0 or summary.get('ownership_errors') or summary.get('worker_returncode')!=0:raise Stop('Office не закрыт корректно.')
    if any(r.get('exit_code')!=0 or r.get('still_running') or r.get('forced_termination_attempted') for r in summary['instances']):raise Stop('Office: аварийное завершение/остаточный процесс.')
    return result

def check_release_graphics(outputs):
    """A DOCX drawing is not evidence that it is a signature or an approved logo.

    Without an independently approved classification keep the release barrier.
    Include headers/footers: an old signature must not bypass it by location.
    This is a read-only check and never strips objects from a source/output.
    """
    from zipfile import ZipFile
    from xml.etree import ElementTree
    namespace='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    blocked=[]
    for item in outputs:
        if Path(item['path']).suffix.lower()!='.docx':continue
        count=0
        with ZipFile(item['path']) as package:
            for name in package.namelist():
                if not (name in ('word/document.xml','word/footnotes.xml','word/endnotes.xml')
                        or name.startswith(('word/header','word/footer')) and name.endswith('.xml')):continue
                root=ElementTree.fromstring(package.read(name))
                count+=len(list(root.iter(namespace+'drawing')))+len(list(root.iter(namespace+'pict')))
        if count:blocked.append('«'+Path(item['path']).name+'» ('+item['role']+'): '+str(count))
    if blocked:
        raise Stop('Назначение графических объектов не подтверждено: '+
                   '; '.join(blocked)+'. Объект может быть логотипом, схемой, '
                   'подписью или печатью; автоматическое распознавание ненадёжно. '
                   'Выпуск остановлен. Требуется утверждённый образец без '
                   'старых графических подписей и непроверенных объектов. '
                   'Изображения автоматически не удаляются.')


def release(path):
    d=load_plan(path);run=Path(d['run']);prepared=read(run/'prepared.json');verified=read(run/'verified.json');confirmation=read(run/'confirmation.json')
    if confirmation['plan_seal']!=d['seal'] or confirmation['verified_sha256']!=sha(run/'verified.json'):raise Stop('Нет подтверждения.')
    if d['facts'].get('owner_confirmation_required'):
        from .owner_control import require_declaration
        require_declaration(d,confirmation)
        check_release_graphics(prepared['outputs'])
    verify_pins(verified['fingerprints'])
    output=run/'release';output.mkdir(exist_ok=False)
    files=office.release_copies(prepared['outputs'],output,d['cfg'].get('training',False))
    office.verify_documents(files,d['number'],d['cfg'].get('training',False),draft=False)
    # Only the FINAL, unwatermarked release copies enter the printable PDF.
    # A failed native Office export stops before any registry modification.
    from .print_pdf import create_print_pdf, verify_package_print
    printed=create_print_pdf(files,output,Path(d['target']).name)
    files.append(printed['output'])
    result={'plan_seal':d['seal'],'outputs':files,
            'fingerprints':pins([o['path'] for o in files]),
            'print_pdf':printed['manifest']}
    verify_package_print(result,Path(d['target']).name)
    save(run/'release.json',result)
    return {'status':'PASS','plan':str(path),'task':'release',
            'print_pdf':printed['manifest']['filename'],
            'print_pdf_pages':printed['manifest']['pages']}

def commit(path,release_result):
    d=load_plan(path);run=Path(d['run']);registry=safe(d['registry']);target=safe(d['target'])
    prepared=read(run/'prepared.json');verified=read(run/'verified.json');confirmation=read(run/'confirmation.json')
    supervisor_proof(release_result,'release',path)
    released=read(run/'release.json')
    if released['plan_seal']!=d['seal']:raise Stop('Другой выпуск.')
    verify_pins(released['fingerprints'])
    from .print_pdf import verify_package_print
    verify_package_print(released,target.name)
    if confirmation['plan_seal']!=d['seal'] or confirmation['verified_sha256']!=sha(run/'verified.json') or not confirmation['reviewed']:raise Stop('Нет действительного подтверждения выпуска.')
    verify_pins(verified['fingerprints']);v=model.visible_vom(d['vom'])
    approved_facts=d['facts']|{'documents_reviewed':True}
    if d['facts'].get('simple_ui'):
        if confirmation.get('approved') is not True:
            raise Stop('Нет итогового подтверждения полномочий и текущих результатов.')
        approved_facts.update(confirmed=True,authorized=True,forms_approved=True)
    owner=None
    if d['facts'].get('owner_confirmation_required'):
        from .owner_control import require_declaration
        owner=require_declaration(d,confirmation)
        approved_facts['owner_control_confirmation']=owner
    factual.validate(approved_facts,v,d['project'],d['number'],
                     d['cfg'].get('training',False),True,control_confirmation=owner)
    if d.get('accounting'):
        from .coordinated import commit as commit_accounting
        return commit_accounting(d,prepared,released,confirmation,approved_facts)
    with mutex(registry),frozen(p for p in d['fingerprints'] if p!=str(registry)):
        if pending(registry):raise Stop('Незавершённая операция: используйте восстановление.')
        if target.exists():raise Stop('Папка уже существует, перезапись запрещена.')
        token=d['id'];stage=registry.parent/'.vzmk-operations'/token
        stage.mkdir(parents=True,exist_ok=False)
        backup=stage/'backup.xlsx';candidate=stage/'candidate.xlsx';journal=stage/'journal.json'
        state={'id':token,'phase':'STARTED','registry':str(registry),'target':str(target),
            'backup':str(backup),'original_sha256':d['fingerprints'][str(registry)],'new':d['new'],'plan':str(path)}
        journal_write(journal,state)
        try:
            with locked_file(registry,writable=d['new']) as original:
                if hashlib.sha256(original).hexdigest()!=state['original_sha256']:raise Stop('SHA реестра изменился перед выпуском.')
                verify_pins({p:h for p,h in d['fingerprints'].items() if p!=str(registry)})
                with backup.open('xb') as f:f.write(original);f.flush();os.fsync(f.fileno())
                if sha(backup)!=state['original_sha256']:raise Stop('SHA резервной копии не совпал.')
                state['phase']='BACKUP_VERIFIED';journal_write(journal,state)
                target.parent.mkdir(parents=True,exist_ok=True)
                private=target.parent/('.pending-'+token);private.mkdir(exist_ok=False)
                out=[]
                for o in released['outputs']:
                    destination=private/Path(o['path']).name;shutil.copyfile(o['path'],destination)
                    if sha(destination)!=released['fingerprints'][o['path']]:raise Stop('Изменилась копия документа.')
                    out.append(o|{'path':str(destination)})
                state['output_hashes']={Path(o['path']).name:sha(o['path']) for o in out}
                state['pending']=str(private)
                if d['new']:
                    shutil.copyfile(prepared['registry'],candidate)
                    if sha(candidate)!=prepared['snapshot_sha256']:raise Stop('Копия реестра изменилась.')
                    verify_registry(backup,candidate,d['project'],d['number'],v,date.fromisoformat(d['date']))
                    state['candidate_sha256']=sha(candidate)
                save(private/'ISSUE.json',{'project':d['project'],'number':d['number'],'training':bool(d['cfg'].get('training')),
                    'facts':approved_facts if d['facts'].get('simple_ui') else d['facts'],
                    'approved_at':confirmation['created_at'],
                    'files':state['output_hashes'],'source_registry_sha256':state['original_sha256']})
                state['issue_sha256']=sha(private/'ISSUE.json')
                state['phase']='PUBLISH_INTENT';journal_write(journal,state)
                private.rename(target)
                state['phase']='DOCUMENTS_PUBLISHED';journal_write(journal,state)
                if d['new']:
                    state['phase']='REPLACE_INTENT';journal_write(journal,state)
                    atomic_replace(candidate,registry,backup=stage/'replaced_original.xlsx')
                    state['phase']='REGISTRY_REPLACED';journal_write(journal,state)
            if d['new']:verify_registry(backup,registry,d['project'],d['number'],v,date.fromisoformat(d['date']))
            elif sha(registry)!=state['original_sha256']:raise Stop('Изменился реестр существующего паспорта.')
            state['phase']='NEEDS_READBACK';journal_write(journal,state)
            return {'status':'PASS','state':'ТРЕБУЕТСЯ ПРОВЕРКА OFFICE','folder':str(target),'backup':str(backup),'journal':str(journal),'training':bool(d['cfg'].get('training'))}
        except BaseException as exc:
            state['error']=str(exc)
            if state['phase']=='STARTED':state['phase']='ABORTED'
            # Keep intent: a crash at ANY boundary is recoverable by hash inspection.
            journal_write(journal,state)
            raise Stop('Выпуск не завершён. Используйте восстановление: '+str(journal)+'; '+str(exc)) from exc

def readback(journal):
    s=journal_read(journal);d=read(s['plan']);registry=safe(s['registry']);target=safe(s['target'])
    if s['phase']!='NEEDS_READBACK':raise Stop('Неверная фаза проверки.')
    if s.get('participants'):
        from .coordinated import verify_current,native_readback
        verify_current(s);native_readback(s)
    if sha(registry)!=(s['candidate_sha256'] if s['new'] else s['original_sha256']):raise Stop('Реестр изменился после записи.')
    for name,h in s['output_hashes'].items():
        if sha(target/name)!=h:raise Stop('Изменён документ: '+name)
    with excel_session() as app:
        b=app.Workbooks.Open(str(registry),UpdateLinks=0,ReadOnly=True,AddToMru=False)
        try:
            r=read_registry(registry)
            found=[e for e in r.projects[d['project']] if e['number']==d['number']]
            if len(found)!=1 or b.Worksheets(r.sheet).Cells(found[0]['row'],r.columns['number']).Value2!=d['number']:raise Stop('Excel: новая строка не подтверждена.')
        finally:b.Close(SaveChanges=False)
    released=read(Path(s['plan']).parent/'release.json')
    from .print_pdf import verify_package_print,verify_print_pdf
    print_manifest=verify_package_print(released,target.name)
    name=print_manifest['filename']
    if name not in s['output_hashes'] or s['output_hashes'][name]!=print_manifest['sha256']:
        raise Stop('Сводный PDF не совпадает с манифестом выпуска.')
    verify_print_pdf(target/name,expected_pages=print_manifest['pages'])
    docs=[{'path':str(target/name),'role':name.rsplit(' - ',1)[-1].rsplit('.',1)[0]}
          for name in s['output_hashes'] if not name.lower().endswith('.pdf')]
    office.verify_documents(docs,d['number'],d['cfg'].get('training',False),False)
    return {'status':'PASS','task':'readback','plan':s['plan'],'journal':str(journal),
            'print_pdf_pages':print_manifest['pages']}

def complete(journal,readback_result):
    journal=safe(journal);s=journal_read(journal);registry=safe(s['registry'])
    supervisor_proof(readback_result,'readback',s['plan'])
    if s.get('participants'):
        if s['phase']!='NEEDS_READBACK':raise Stop('Неверная фаза завершения.')
        from .coordinated import complete as complete_accounting
        return complete_accounting(journal,s)
    with mutex(registry),locked_file(registry) as current:
        if s['phase']!='NEEDS_READBACK':raise Stop('Неверная фаза завершения.')
        if hashlib.sha256(current).hexdigest()!=(s['candidate_sha256'] if s['new'] else s['original_sha256']):raise Stop('Реестр изменился после проверки.')
        for name,h in s['output_hashes'].items():
            if sha(safe(s['target'])/name)!=h:raise Stop('Документ изменился после проверки.')
        print_name=Path(s['target']).name+'.pdf'
        if print_name not in s['output_hashes']:
            raise Stop('Не найден проверенный сводный PDF. Завершение запрещено.')
        s['phase']='COMPLETE';journal_write(journal,s)
    return {'status':'PASS','state':'ГОТОВО','folder':s['target'],'journal':str(journal),
            'backup':s['backup'],'print_pdf':str(Path(s['target'])/print_name)}

def recover(journal):
    journal=safe(journal);s=journal_read(journal);registry=safe(s['registry'])
    if journal.parent.parent!=registry.parent/'.vzmk-operations' or journal.parent.name!=s['id']:raise Stop('Журнал вне папки операций выбранного реестра.')
    if s.get('participants'):
        from .coordinated import recover as recover_accounting
        return recover_accounting(journal,s)
    if s['phase'] in ('COMPLETE','ROLLED_BACK','ABORTED'):return {'status':'PASS','phase':s['phase']}
    backup=journal.parent/'backup.xlsx'
    if not backup.exists() and s['phase']=='STARTED':
        with mutex(registry),locked_file(registry,writable=True) as current:
            if hashlib.sha256(current).hexdigest()!=s['original_sha256'] or Path(s['target']).exists():raise Stop('Неизвестное состояние до резервного копирования.')
            s['phase']='ABORTED';journal_write(journal,s)
        return {'status':'PASS','phase':'ABORTED'}
    if safe(s['backup'])!=backup or not backup.is_file() or sha(backup)!=s['original_sha256']:raise Stop('Резервная копия отсутствует/повреждена. Ручная проверка; запись запрещена.')
    with mutex(registry),locked_file(registry,writable=True) as current:
        h=hashlib.sha256(current).hexdigest()
        if h not in (s['original_sha256'],s.get('candidate_sha256')):raise Stop('Реестр изменён другим пользователем. Откат запрещён.')
        target=safe(s['target'])
        if target.exists():
            if not s.get('output_hashes') or sha(target/'ISSUE.json')!=s.get('issue_sha256'):raise Stop('Комплект изменён после сбоя. Ручная проверка.')
            expected=set(s['output_hashes'])|{'ISSUE.json'}
            if {p.name for p in target.iterdir()}!=expected:raise Stop('В папке появились чужие файлы. Откат запрещён.')
            for name,expected_hash in s['output_hashes'].items():
                if sha(target/name)!=expected_hash:raise Stop('Готовый документ изменён. Откат запрещён.')
        if h!=s['original_sha256']:
            restore=journal.parent/('restore_'+uuid4().hex+'.xlsx');shutil.copyfile(backup,restore)
            atomic_replace(restore,registry,backup=journal.parent/('pre_restore_'+uuid4().hex+'.xlsx'))
        if target.exists():target.rename(journal.parent/'unpublished_documents')
        if h!=s['original_sha256'] and sha(registry)!=s['original_sha256']:raise Stop('Не совпал SHA восстановленного реестра.')
        s['phase']='ROLLED_BACK';journal_write(journal,s)
        return {'status':'PASS','phase':'ROLLED_BACK','journal':str(journal),'registry_sha256':s['original_sha256']}
