"""Coordinated, journalled registry/catalog/UZK replacement.

Windows cannot atomically replace three files. All participants are locked
and backed up before the first replacement; an interrupted intent is blocked
until hash-based recovery restores the complete original set. No retry can
spend quantities twice. Unknown/external changes prohibit every rollback.
"""
from contextlib import ExitStack
from pathlib import Path
from uuid import uuid4
import hashlib,os,shutil
from .storage import Stop,safe,sha,save,read,verify_pins,mutex
from passport_app.excel import locked_file,atomic_replace
from stage8_closure.storage import frozen


def accounting_plan(cfg,project,number,vom,facts,profile,registry_entries):
    if not cfg.get('use_shared_uzk'):
        if cfg.get('auto_first_passport'):raise Stop('Для официального выпуска подключите общий УЗК.xlsm; учёт остатков нельзя пропустить.')
        return None # Compatibility for the older, explicitly configured engine.
    from .shared_uzk import find_shared_uzk,read_shared_uzk,_key
    from .catalog_ledger import validate_stock,shipment_quantities
    for climate in facts.get('climates',[]):
        if not climate.get('layers'):
            from .facts import numeric
            numeric(climate.get('wet_thickness'),'Средняя мокрая толщина',__import__('decimal').Decimal('.001'))
    catalog=profile.get('catalog')
    if not catalog:raise Stop('Нет справочника остатков выбранного проекта.')
    uzk=find_shared_uzk(cfg);paths=[safe(catalog),uzk,safe(cfg['registry'])]
    if len(set(paths))!=3:raise Stop('Реестр, справочник и УЗК должны быть разными файлами.')
    if cfg.get('training'):
        sandbox=safe(cfg['registry']).parent
        if any(not path.is_relative_to(sandbox) for path in paths):
            raise Stop('Учебный официальный цикл разрешён только с копиями всех трёх файлов внутри учебной папки.')
    quantities=shipment_quantities(vom)
    history=validate_stock(catalog,project,number,quantities,registry_entries)
    controls={str(row['mark']):int(row['uzk_quantity']) for row in facts['marks']}
    index=read_shared_uzk(uzk,project=project,marks=controls)
    for mark,count in controls.items():
        if count and len(index[(_key(project),_key(mark))]['rows'])!=1:raise Stop('Общий УЗК: дубли строк марки '+mark+'; запись запрещена.')
    return {'catalog':str(paths[0]),'uzk':str(paths[1]),'quantities':quantities,'controls':controls,'history_sources':[h['source'] for h in history['history'] if h.get('source')],
            'history_fingerprints':{h['source']:h['sha256'] for h in history['history'] if h.get('source')},
            'history_binding_updates':[{'number':h['number'],**u} for h in history['history'] for u in h.get('binding_updates',[])]}


def prepare_candidates(d):
    if not d.get('accounting'):return []
    from .catalog_ledger import shipment_candidate
    from .shared_uzk import decrement_candidate
    from .model import read_registry
    from passport_app.excel import excel_session
    a=d['accounting'];root=Path(d['run'])/'accounting';root.mkdir(exist_ok=False)
    ledger=root/'catalog.xlsx';uzk=root/'УЗК.xlsm'
    ledger_result=shipment_candidate(a['catalog'],ledger,d['project'],d['number'],d['date'],a['quantities'],read_registry(d['registry']).projects[d['project']])
    decrement_candidate(a['uzk'],uzk,d['project'],a['controls'])
    # Native Save/Open acceptance applies only to private candidates. XLSM
    # is opened read-only: keep its opaque VBA ZIP parts exactly unchanged.
    with excel_session() as app:
        for candidate in (ledger,uzk):
            b=app.Workbooks.Open(str(candidate),UpdateLinks=0,ReadOnly=True,AddToMru=False)
            try:
                if b.Worksheets.Count<1:raise Stop('Пустой кандидат '+candidate.name)
                if candidate==ledger:
                    app.CalculateFull()
                    s=b.Worksheets(1)
                    for row in range(2,int(s.UsedRange.Rows.Count)+1):
                        if str(s.Cells(row,1).Value2 or '') in a['quantities']:
                            value=s.Cells(row,5).Value2
                            if not isinstance(value,(float,int)) or value<0:raise Stop('Excel: не подтверждён остаток справочника.')
            finally:b.Close(SaveChanges=False)
    return [{'role':role,'path':a[role],'candidate':str(candidate),'original_sha256':d['fingerprints'][a[role]],'candidate_sha256':sha(candidate),
             **({'history_binding_updates':ledger_result['history_binding_updates']} if role=='catalog' else {})} for role,candidate in [('catalog',ledger),('uzk',uzk)]]


def _lock_all(stack,participants,writable):
    payloads={}
    for item in sorted(participants,key=lambda x:x['path'].casefold()):
        stack.enter_context(mutex(item['path']))
        payloads[item['path']]=stack.enter_context(locked_file(item['path'],writable=writable))
    return payloads


def _write_bytes(path,payload):
    with Path(path).open('xb') as file:file.write(payload);file.flush();os.fsync(file.fileno())


def commit(d,prepared,released,confirmation,approved_facts):
    from .core import pending,journal_write,verify_registry
    from .model import visible_vom
    from datetime import date
    registry=safe(d['registry']);target=safe(d['target']);stage=registry.parent/'.vzmk-operations'/d['id'];journal=stage/'journal.json'
    participants=[{'role':'registry','path':str(registry),'candidate':prepared['registry'],'original_sha256':d['fingerprints'][str(registry)],'candidate_sha256':prepared['snapshot_sha256']},*prepared['accounting']]
    if len(participants)!=3:raise Stop('Нет трёх проверенных кандидатов учёта.')
    for item in participants:
        expected=d['fingerprints'].get(item['path'])
        if expected!=item['original_sha256'] or sha(item['candidate'])!=item['candidate_sha256']:raise Stop('Кандидат учёта изменён.')
        item['backup']=str(stage/('backup-'+item['role']+Path(item['path']).suffix))
    paths={i['path'] for i in participants}
    with ExitStack() as stack:
        payloads=_lock_all(stack,participants,True)
        stack.enter_context(frozen(p for p in d['fingerprints'] if p not in paths))
        verify_pins({p:h for p,h in d['fingerprints'].items() if p not in paths})
        for item in participants:
            if hashlib.sha256(payloads[item['path']]).hexdigest()!=item['original_sha256']:raise Stop('Другой пользователь изменил '+item['path'])
        if pending(registry):raise Stop('Есть незавершённая операция; требуется восстановление.')
        if target.exists():raise Stop('Комплект уже существует; повторная запись запрещена.')
        stage.mkdir(parents=True,exist_ok=False)
        state={'id':d['id'],'phase':'STARTED','registry':str(registry),'target':str(target),'plan':str(Path(d['run'])/'plan.json'),
               'backup':participants[0]['backup'],'original_sha256':participants[0]['original_sha256'],'candidate_sha256':participants[0]['candidate_sha256'],
               'new':True,'participants':participants,
               'current_control':confirmation.get('current_control')}
        journal_write(journal,state)
        try:
            for item in participants:
                _write_bytes(item['backup'],payloads[item['path']])
                if sha(item['backup'])!=item['original_sha256']:raise Stop('Повреждена резервная копия '+item['role'])
                local=Path(item['path']).with_name('.vzmk-candidate-'+d['id']+Path(item['path']).suffix)
                _write_bytes(local,Path(item['candidate']).read_bytes())
                if sha(local)!=item['candidate_sha256']:raise Stop('Повреждён кандидат '+item['role'])
                item['local_candidate']=str(local)
            state['phase']='BACKUP_VERIFIED';journal_write(journal,state)
            private=target.parent/('.pending-'+d['id']);private.mkdir(parents=True,exist_ok=False)
            state['pending']=str(private);state['output_hashes']={}
            for output in released['outputs']:
                dest=private/Path(output['path']).name;shutil.copyfile(output['path'],dest)
                if sha(dest)!=released['fingerprints'][output['path']]:raise Stop('Изменён выпускаемый документ.')
                state['output_hashes'][dest.name]=sha(dest)
            save(private/'ISSUE.json',{'project':d['project'],'number':d['number'],'training':bool(d['cfg'].get('training')),'facts':approved_facts,
                                     'approved_at':confirmation['created_at'],'files':state['output_hashes'],'source_registry_sha256':state['original_sha256']})
            state['issue_sha256']=sha(private/'ISSUE.json');state['phase']='PUBLISH_INTENT';journal_write(journal,state)
            private.rename(target);state['phase']='DOCUMENTS_PUBLISHED';journal_write(journal,state)
            for item in participants:
                state['phase']='REPLACE_INTENT_'+item['role'].upper();journal_write(journal,state)
                atomic_replace(item['local_candidate'],item['path'],backup=Path(item['path']).with_name('.vzmk-replaced-'+d['id']+Path(item['path']).suffix))
                state['phase']='REPLACED_'+item['role'].upper();journal_write(journal,state)
            state['phase']='NEEDS_READBACK';journal_write(journal,state)
        except BaseException as exc:
            state['error']=str(exc);journal_write(journal,state)
            raise Stop('Запись трёх файлов не завершена. До восстановления работа заблокирована: '+str(journal)+'; '+str(exc)) from exc
    verify_current(state)
    verify_registry(participants[0]['backup'],registry,d['project'],d['number'],visible_vom(d['vom']),date.fromisoformat(d['date']))
    return {'status':'PASS','state':'ТРЕБУЕТСЯ ПРОВЕРКА OFFICE','folder':str(target),'backup':state['backup'],'journal':str(journal),'training':bool(d['cfg'].get('training'))}


def verify_current(state):
    for item in state['participants']:
        if sha(item['path'])!=item['candidate_sha256']:raise Stop('После записи изменён файл '+item['path'])
    return True


def native_readback(state):
    from passport_app.excel import excel_session
    from .catalog_ledger import metadata,quantity
    from .shared_uzk import read_shared_uzk,_key
    from openpyxl import load_workbook
    data=read(state['plan']);a=data['accounting']
    catalog=next(i['path'] for i in state['participants'] if i['role']=='catalog')
    book=load_workbook(catalog,data_only=False)
    try:
        history=metadata(book)['history'];matches=[(column,h) for column,h in enumerate(history,6) if h['number']==data['number']]
        if len(matches)!=1 or matches[0][1]['date']!=data['date']:raise Stop('Readback: нет единственного паспорта в справочнике.')
        column,_=matches[0];sheet=book.worksheets[0];expected={}
        for no in range(2,sheet.max_row+1):
            mark=str(sheet.cell(no,1).value or '')
            if mark in a['quantities']:
                if sheet.cell(no,column).value!=a['quantities'][mark]:raise Stop('Readback: неверная отгрузка '+mark)
                expected[no]=quantity(sheet.cell(no,2).value,mark,True)-sum(h['quantities'].get(mark,0) for h in history)
    finally:book.close()
    with excel_session() as app:
        native=app.Workbooks.Open(str(catalog),UpdateLinks=0,ReadOnly=True,AddToMru=False)
        try:
            if not native.ReadOnly:raise Stop('Readback: справочник открыт не только для чтения.')
            app.CalculateFull()
            for no,value in expected.items():
                if native.Worksheets(1).Cells(no,5).Value2!=value:raise Stop('Readback Excel: не совпал остаток.')
        finally:native.Close(SaveChanges=False)
        uzk=next(i for i in state['participants'] if i['role']=='uzk')
        native=app.Workbooks.Open(uzk['path'],UpdateLinks=0,ReadOnly=True,AddToMru=False)
        try:
            if not native.ReadOnly:raise Stop('Readback: общий УЗК открыт не только для чтения.')
            before=read_shared_uzk(uzk['backup'],project=data['project'],marks=a['controls'])
            after=read_shared_uzk(uzk['path'],project=data['project'],marks=a['controls'])
            for mark,consumed in a['controls'].items():
                if not consumed:continue
                key=(_key(data['project']),_key(mark));remaining=before[key]['remaining']-consumed
                if remaining and after.get(key,{}).get('remaining')!=remaining or not remaining and key in after:
                    raise Stop('Readback: неверный остаток общего УЗК '+mark)
        finally:native.Close(SaveChanges=False)
    return True


def complete(journal,state):
    from .core import journal_write
    with ExitStack() as stack:
        payloads=_lock_all(stack,state['participants'],False)
        for item in state['participants']:
            if hashlib.sha256(payloads[item['path']]).hexdigest()!=item['candidate_sha256']:raise Stop('После readback изменён '+item['path'])
        target=Path(state['target'])
        for name,h in state['output_hashes'].items():
            if sha(target/name)!=h:raise Stop('После readback изменён документ '+name)
        if target.name+'.pdf' not in state['output_hashes']:raise Stop('Отсутствует итоговый PDF.')
        if sha(target/'ISSUE.json')!=state.get('issue_sha256'):raise Stop('После readback изменены сведения о выпуске.')
        state['phase']='COMPLETE';journal_write(journal,state)
    return {'status':'PASS','state':'ГОТОВО','folder':str(target),'journal':str(journal),'backup':state['backup'],'print_pdf':str(target/(target.name+'.pdf'))}


def recover(journal,state):
    from .core import journal_write
    if state['phase'] in ('COMPLETE','ROLLED_BACK','ABORTED'):return {'status':'PASS','phase':state['phase']}
    stage=Path(journal).parent;registry=safe(state['registry'])
    if stage.parent!=registry.parent/'.vzmk-operations' or stage.name!=state['id']:raise Stop('Журнал вне папки операций.')
    participants=state['participants']
    if len(participants)!=3 or len({i['path'] for i in participants})!=3 or participants[0]['path']!=str(registry):raise Stop('Повреждены участники транзакции.')
    with ExitStack() as stack:
        payloads=_lock_all(stack,participants,True)
        changed=[]
        # Validate EVERY participant and EVERY document BEFORE restoring any.
        for item in participants:
            current=hashlib.sha256(payloads[item['path']]).hexdigest()
            if current not in (item['original_sha256'],item['candidate_sha256']):raise Stop('Внешнее изменение '+item['path']+'; автоматический откат всех файлов запрещён.')
            if current!=item['original_sha256']:
                backup=safe(item['backup'])
                if backup.parent!=stage or not backup.is_file() or sha(backup)!=item['original_sha256']:raise Stop('Резервная копия повреждена '+item['role'])
                changed.append(item)
        target=safe(state['target'])
        if target.exists():
            if sha(target/'ISSUE.json')!=state.get('issue_sha256') or {p.name for p in target.iterdir()}!=set(state.get('output_hashes',{}))|{'ISSUE.json'}:raise Stop('Комплект изменён после сбоя; откат запрещён.')
            for name,h in state['output_hashes'].items():
                if sha(target/name)!=h:raise Stop('Изменён документ '+name+'; откат всех файлов запрещён.')
        for item in reversed(changed):
            path=Path(item['path']);restore=path.with_name('.vzmk-restore-'+uuid4().hex+path.suffix)
            _write_bytes(restore,Path(item['backup']).read_bytes())
            atomic_replace(restore,path,backup=stage/('pre-restore-'+uuid4().hex+path.suffix))
        if target.exists():target.rename(stage/'unpublished_documents')
    for item in participants:
        if sha(item['path'])!=item['original_sha256']:raise Stop('Не подтверждено восстановление '+item['path'])
    state['phase']='ROLLED_BACK';journal_write(journal,state)
    return {'status':'PASS','phase':'ROLLED_BACK','journal':str(journal),'participants':len(participants)}
