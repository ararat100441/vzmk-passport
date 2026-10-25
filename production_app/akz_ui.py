"""Small testable rules controlling which АКЗ controls the GUI shows."""
import re
from pathlib import Path
from openpyxl import load_workbook
from .catalog_assignments import used_akz_variants
from .storage import Stop,sha


def available_variants(profile):
    templates=profile.get('templates',{})
    return [n for n in range(4) if
            'akz'+str(n) in templates and
            'akz_appendix'+str(n) in templates]


def display_variants(mark_data,profile):
    """Actual current assignments dominate; never silently assign zero."""
    used=used_akz_variants(mark_data)
    available=available_variants(profile)
    # First time in garage, no assignments yet: show sole *available form*,
    # but do NOT change any mark to variant0.
    if not used and len(available)==1:
        return available
    return used


def source_coats(ref):
    """Determine coats from approved empty template, not variant number."""
    path=Path(ref['path'])
    if not path.is_file() or sha(path)!=ref['sha256']:
        raise Stop('Изменена утверждённая форма АКЗ.')
    w=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        if len(w.worksheets)!=1:
            raise Stop('В АКЗ должна быть одна форма Excel.')
        tokens=set()
        for row in w.active.iter_rows():
            for cell in row:
                if isinstance(cell.value,str):
                    tokens.update(re.findall(r'\{\{([A-Z_0-9]+)\}\}',cell.value))
        if any(t.startswith('LAYER_3_') or t.startswith('AKZ_LAYER_3_')
               for t in tokens):
            return 3
        if any(t.startswith('LAYER_2_') or t.startswith('AKZ_LAYER_2_')
               for t in tokens):
            return 2
        return 1
    finally:
        w.close()
