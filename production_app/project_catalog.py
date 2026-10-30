"""Project-local ledger naming and explicit conflict resolution."""
from pathlib import Path
from .storage import Stop,sha

def catalog_name(project):
    if not str(project).strip() or any(c in str(project) for c in '<>:"/\\|?*') or str(project).endswith((' ','.')):
        raise Stop('Небезопасное название проекта для справочника.')
    return 'Справочник Марок '+str(project)+'.xlsx'

def resolve_catalog(root,project,chosen=None):
    root=Path(root).resolve(strict=True)
    if not root.is_dir():raise Stop('Выберите существующую папку проекта.')
    allowed={catalog_name(project).casefold(),'справочник марок.xlsx'}
    found=[p for p in root.iterdir() if p.is_file() and not p.is_symlink() and p.name.casefold() in allowed]
    if chosen:
        p=Path(chosen).resolve(strict=True)
        if p.parent!=root or p not in found:raise Stop('Выбранный справочник не принадлежит существующей папке этого проекта.')
        return p
    if len(found)>1:raise Stop('Конфликт справочников проекта '+str(project)+': найдены старое и новое имена. В настройках явно выберите действующий справочник. Ничего не переименовано.')
    return found[0] if found else None

def create_catalog(root,project,cfg=None):
    """Only create a file in an ALREADY existing directory, never a project."""
    from openpyxl import Workbook
    from .catalog_ledger import HEADERS,META,SOURCE,SCHEMA,write_metadata,update_balance,registered_history
    root=Path(root).resolve(strict=True)
    if not root.is_dir():raise Stop('Папка проекта должна существовать заранее.')
    if resolve_catalog(root,project):raise Stop('Справочник уже существует. Для старого файла используйте подготовку структуры с резервной копией.')
    history=[]
    if cfg:
        from .model import read_registry
        registry=read_registry(cfg['registry'])
        if project in registry.projects:history=registered_history(cfg,project)
    b=Workbook();s=b.active;s.title='Марки и отгрузки';s.append(HEADERS)
    marks=list(dict.fromkeys(m for h in history for m in h['quantities']))
    for mark in marks:s.append((mark,None,None,None,None))
    for col,h in enumerate(history,6):
        from datetime import date
        s.cell(1,col,'Паспорт '+h['number']+' от '+date.fromisoformat(h['date']).strftime('%d.%m.%Y'))
        for row,mark in enumerate(marks,2):s.cell(row,col,h['quantities'].get(mark,0) if h['complete'] else None)
    technical=b.create_sheet(SOURCE);technical.append(('Марка','Профиль','Краска','№ АКЗ'));technical.sheet_state='hidden'
    meta=b.create_sheet(META);meta.sheet_state='veryHidden';meta['A1']=SCHEMA
    write_metadata(meta,{'project':project,'history':history,'source_sha256':None})
    for row in range(2,s.max_row+1):update_balance(s,row,len(history),all(h['complete'] for h in history))
    s.freeze_panes='F2'
    for key,width in [('A',24),('B',16),('C',24),('D',12),('E',16)]:s.column_dimensions[key].width=width
    target=root/catalog_name(project)
    try:
        with target.open('xb') as f:b.save(f)
    finally:b.close()
    return {'path':str(target),'sha256':sha(target),'marks':len(marks),'registered_history':len(history),'notice':'Общее количество и назначения заполняются владельцем. Неизвестный остаток не равен нулю.'}
