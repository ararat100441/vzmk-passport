"""Read-only project catalogue for approved profile/paint/AKZ assignments.

A catalogue is an engineering assignment source, NOT evidence that VIK/UZK/AKZ
inspections took place. Blank means unassigned, while numeric 0 is AKZ №0.
"""
from pathlib import Path
from openpyxl import load_workbook
from .storage import Stop, sha

ALIASES={
    'mark':('марка','отправочная марка'),
    'profile':('профиль','профиль металлопроката'),
    'paint':('краска','лакокрасочный материал','лкм'),
    'akz':('№ акз','номер акз','акз'),
}


def _caption(value):
    return ' '.join(str(value or '').replace('ё','е').casefold().split())


def read_assignments(path, expected_sha=None):
    """Load all assignments without opening workbook in Excel or saving it."""
    path=Path(path).resolve(strict=True)
    before=sha(path)
    if expected_sha is not None and before!=expected_sha:
        raise Stop('Изменился справочник назначений: '+str(path))
    if path.suffix.lower() not in ('.xlsx','.xlsm'):
        raise Stop('Справочник должен быть XLSX/XLSM.')
    book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    result={}
    try:
        if len(book.worksheets)!=1:
            raise Stop('В справочнике должен быть один лист с назначениями.')
        sheet=book.active
        columns=None
        for row_index,cells in enumerate(sheet.iter_rows(min_row=1,max_row=min(sheet.max_row,20),values_only=True),start=1):
            captions=[_caption(value) for value in cells]
            matches={}
            for key,aliases in ALIASES.items():
                pos=[i for i,value in enumerate(captions) if value in aliases]
                if len(pos)>1:
                    raise Stop('В справочнике повторяется столбец '+key)
                if pos:
                    matches[key]=pos[0]
            if 'mark' in matches and any(key in matches for key in ('profile','paint','akz')):
                columns=matches
                break
        if columns is None:
            raise Stop('Не найден заголовок «Марка» и назначения профиля/краски/АКЗ.')
        for row in sheet.iter_rows(min_row=row_index+1,values_only=True):
            if columns['mark']>=len(row):continue
            mark=row[columns['mark']]
            if mark is None or not str(mark).strip():
                continue
            key=str(mark).strip()
            if key in result:
                raise Stop('Дублирующаяся марка в справочнике: '+key)
            values={}
            for field in ('profile','paint','akz'):
                if field not in columns:
                    continue
                raw=row[columns[field]] if columns[field]<len(row) else None
                if raw is None or isinstance(raw,str) and not raw.strip():
                    continue
                if field=='akz':
                    if isinstance(raw,bool):
                        raise Stop('Неверный номер АКЗ у марки '+key)
                    try:
                        if str(raw).strip() not in ('0','1','2','3'):
                            raise ValueError()
                        raw=int(raw)
                    except (ValueError,TypeError):
                        raise Stop('№ АКЗ должен быть 0, 1, 2 или 3: '+key)
                else:
                    if not isinstance(raw,str):
                        raise Stop('Неверное текстовое назначение '+field+': '+key)
                    raw=raw.strip()
                values[field]=raw
            result[key]=values
    finally:
        book.close()
    if sha(path)!=before:
        raise Stop('Справочник изменился при чтении.')
    return result


def current_assignments(path, expected_sha, marks):
    """Apply only exact VOM marks; missing mappings must stay blank."""
    mapping=read_assignments(path,expected_sha)
    result={}
    for mark in marks:
        name=str(mark).strip()
        result[name]=mapping.get(name,{})
    return result


def used_akz_variants(rows):
    """Return only explicitly assigned variants; never assume AKZ0."""
    found=set()
    for row in rows:
        value=row.get('akz')
        if value is None or value=='':
            continue
        if type(value) is not int or value not in range(4):
            raise Stop('Некорректный номер АКЗ в данных: '+str(value))
        found.add(value)
    return sorted(found)
