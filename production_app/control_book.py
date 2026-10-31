"""Current-shipment Excel input. Reads facts; never infers control results."""
from copy import deepcopy
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .storage import Stop, sha
from .real_forms import DOC_FIELD_LABELS, UZK_COLUMNS, UZK_HEADERS, UZK_ALIASES

META = {'project':'Проект','number':'Номер паспорта','responsible':'Ответственный',
        'start_date':'Начало изготовления','date':'Дата паспорта','vik_date':'Дата ВИК',
        'uzk_date':'Дата УЗК','akz_date':'Дата АКЗ'}
MARKS = {'mark':'Марка','profile':'Профиль','paint':'Краска','akz':'№ АКЗ',
         'uzk_quantity':'Проконтролировано УЗК, шт','uzk_result':'Результат УЗК',
         'basis':'Протокол УЗК / основание','vik_result':'Результат ВИК','vik_basis':'Протокол ВИК'}
CLIMATE = {'variant':'№ АКЗ','air':'Воздух °C','surface':'Поверхность °C','humidity':'Влажность %',
           'dew':'Точка росы °C','thickness':'Средняя толщина (однослойный акт)',
           'material':'Материал','basis':'Протокол АКЗ'}
LAYERS = {'variant':'№ АКЗ','layer':'Слой','min':'Допуск от, мкм','max':'Допуск до, мкм',
          'average':'Средняя сухая после слоя, мкм','material':'ЛКМ','paint_batch':'Партия ЛКМ',
          'hardener_batch':'Партия отвердителя','color':'Цвет','application_date':'Дата нанесения',
          'time_start':'Начало HH:MM','time_end':'Окончание HH:MM','wet_average':'Средняя мокрая, мкм'}
SHEETS = {'marks':('Контроль',MARKS),'uzk_details':('Соединения УЗК',dict(zip(UZK_COLUMNS,UZK_HEADERS))),
          'climates':('Климат',CLIMATE),'layers':('Слои АКЗ',LAYERS)}
CURRENT_FIELDS = {k for k in DOC_FIELD_LABELS if any(s in k for s in
    ('CONCLUSION','DEFECTS','ROUGHNESS','CLEANING_GRADE','PREPARATION_GRADE','DEGREASING','DUST_REMOVAL','APPLICATION_METHOD'))}
CURRENT_FIELDS.add('QUALITY_PREPARATION')

def text(value, key=''):
    if value is None:return ''
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    if isinstance(value,time):return value.strftime('%H:%M')
    if isinstance(value,float) and value.is_integer():return str(int(value))
    result=str(value).strip()
    if result and (key=='date' or key.endswith('_date')):
        try:
            try:return date.fromisoformat(result).isoformat()
            except ValueError:return datetime.strptime(result,'%d.%m.%Y').date().isoformat()
        except ValueError as exc:raise Stop('Неверная дата: '+result+'. Пример: 24.10.2026.') from exc
    if result and key in ('time_start','time_end'):
        try:return datetime.strptime(result,'%H:%M').strftime('%H:%M')
        except ValueError as exc:raise Stop('Неверное время: '+result+'. Пример: 08:00.') from exc
    return result

def key(value):return ' '.join(text(value).casefold().split())

def rows(book, name, columns):
    if name not in book.sheetnames:return []
    sheet=book[name]; aliases={key(v):k for k,v in columns.items()}|{key(k):k for k in columns}
    header=None
    for row in sheet.iter_rows(max_row=min(10,sheet.max_row)):
        found={i:aliases[key(c.value)] for i,c in enumerate(row) if key(c.value) in aliases}
        if len(found)>=min(2,len(columns)):
            if len(set(found.values()))!=len(found):raise Stop('Повтор заголовка на листе '+name)
            header=(row[0].row,found);break
    if header is None:raise Stop('На листе «'+name+'» нет понятных заголовков. Используйте ведомость из программы.')
    number,mapping=header;result=[]
    for row in sheet.iter_rows(min_row=number+1):
        data={field:text(row[i].value,field) for i,field in mapping.items()}
        if any(data.values()):result.append(data)
    return result

def read_book(path, project, number, training=False, defaults=None):
    path=Path(path).resolve()
    if path.suffix.lower() not in ('.xlsx','.xlsm'):raise Stop('Выберите ведомость контроля Excel (.xlsx или .xlsm).')
    before=sha(path)
    formulas=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        # No silent None for formulas without Excel's calculated cache.
        for sheet in formulas:
            cached=book[sheet.title]
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f' and cached[cell.coordinate].value is None:
                        raise Stop('Пересчитайте и сохраните ведомость в Excel: '+sheet.title+'!'+cell.coordinate)
        if 'Комплект' not in book.sheetnames:raise Stop('В ведомости нужен лист «Комплект» с проектом и данными отгрузки. Кнопка «Новая ведомость» создаёт его без результатов контроля.')
        aliases={key(v):k for k,v in META.items()}|{key(k):k for k in META}
        meta={aliases[key(r[0])]:text(r[1],aliases[key(r[0])]) for r in book['Комплект'].iter_rows(values_only=True)
              if len(r)>1 and key(r[0]) in aliases}
        if meta.get('project')!=project:raise Stop('Ведомость относится к другому проекту: '+meta.get('project','не указан'))
        if meta.get('number') and meta['number']!=number:raise Stop('Номер в ведомости не совпадает с актуальным реестром: '+number)
        data=meta|{'project':project,'number':number,'kind':'TRAINING' if training else 'ACTUAL',
                   'confirmed':False,'authorized':False,'forms_approved':False,
                   'evidence':[{'path':str(path),'sha256':before}]}
        for field,(name,columns) in SHEETS.items():data[field]=rows(book,name,columns)
        for mark in data['marks']:
            raw=mark.get('akz','')
            try:mark['akz']=int(raw) if raw else None
            except ValueError:raise Stop('№ АКЗ должен быть 0, 1, 2, 3 либо пустым: '+mark.get('mark',''))
        layers=data.pop('layers');by_variant={}
        for row in layers:
            try:n=int(row['variant']);layer=int(row['layer'])
            except (ValueError,KeyError):raise Stop('На листе «Слои АКЗ» укажите номер АКЗ и слой.')
            if not any(row.get(k) for k in LAYERS if k not in ('variant','layer')):continue
            by_variant.setdefault(n,[]).append(row|{'layer':layer})
        for climate in data['climates']:
            try:n=int(climate['variant']);climate['variant']=n
            except (KeyError,ValueError):raise Stop('На листе «Климат» укажите № АКЗ.')
            if n in by_variant:
                entries=by_variant[n]
                climate['layers']=[{k:r.get(k,'') for k in ('layer','min','max')} for r in entries]
                climate['dry_averages']=[{k:r.get(k,'') for k in ('layer','average')} for r in entries]
                climate['layer_details']=[{k:r.get(k,'') for k in LAYERS if k not in ('variant','min','max','average')} for r in entries]
        used={r['akz'] for r in data['marks'] if r['akz'] is not None}
        data['climates']=[c for c in data['climates'] if c['variant'] in used]
        field_aliases={key(v):k for k,v in DOC_FIELD_LABELS.items()}|{key(k):k for k in DOC_FIELD_LABELS}
        data['doc_fields']={k:v for k,v in (defaults or {}).get('doc_fields',{}).items() if k not in CURRENT_FIELDS}
        if 'Реквизиты' in book.sheetnames:
            for row in book['Реквизиты'].iter_rows(values_only=True):
                if len(row)>1 and key(row[0]) in field_aliases:
                    name=field_aliases[key(row[0])];value=text(row[1])
                    if value:data['doc_fields'][name]=value
                    else:data['doc_fields'].pop(name,None)
    finally:formulas.close();book.close()
    if sha(path)!=before:raise Stop('Ведомость изменилась во время чтения. Сохраните её и выберите снова.')
    return data

def preferences(facts, previous=None):
    """Reusable assignments/requisites only; never persist shipment verdicts."""
    return {'doc_fields':{k:v for k,v in facts.get('doc_fields',{}).items() if k not in CURRENT_FIELDS},
            'assignments':deepcopy((previous or {}).get('assignments',{}))|
                {r['mark']:{k:r.get(k,'') for k in ('profile','paint','akz')} for r in facts.get('marks',[])}}

def read_uzk_book(path):
    """Read a normal Excel joint ledger, preserving order and repeated marks.

    A joint count is deliberately NOT converted into a controlled-unit count.
    """
    path=Path(path).resolve();before=sha(path)
    book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    aliases={key(k):v for k,v in UZK_ALIASES.items()}|{key(v):k for k,v in zip(UZK_COLUMNS,UZK_HEADERS)}|{key(k):k for k in UZK_COLUMNS}
    found=[]
    try:
        for sheet in book:
            for row in sheet.iter_rows(max_row=min(sheet.max_row,20)):
                mapping={i:aliases[key(c.value)] for i,c in enumerate(row) if key(c.value) in aliases}
                if not (set(UZK_COLUMNS)-{'basis'}).issubset(set(mapping.values())):continue
                if len(mapping)!=len(set(mapping.values())):raise Stop('Неоднозначные заголовки УЗК: '+sheet.title)
                for values in sheet.iter_rows(min_row=row[0].row+1,values_only=True):
                    item={name:text(values[i],name) for i,name in mapping.items()}
                    if not any(item.values()):continue
                    if not item.get('basis'):item['basis']=path.name
                    found.append(item)
                break
        if not found:raise Stop('В Excel не найдена подробная ведомость УЗК: нужны марка, соединение, сварщик, дата, частота, угол, дефекты и заключение. Исторические заполненные акты не используются как новые результаты.')
    finally:book.close()
    if sha(path)!=before:raise Stop('Ведомость УЗК изменилась при чтении. Выберите сохранённый файл повторно.')
    return found,{'path':str(path),'sha256':before}

def write_book(path, project, number, vom, defaults=None, facts=None):
    """New editable input, including only supplied facts; no invented verdicts."""
    path=Path(path)
    if path.exists():raise Stop('Ведомость уже существует. Выберите новое имя: '+str(path))
    defaults=defaults or {};facts=deepcopy(facts or {})
    wb=Workbook();wb.remove(wb.active)
    meta=wb.create_sheet('Комплект');meta.append(['Параметр','Значение'])
    for k,label in META.items():meta.append([label,project if k=='project' else number if k=='number' else facts.get(k,'')])
    marks=facts.get('marks')
    if marks is None:marks=[{'mark':mark,**defaults.get('assignments',{}).get(mark,{})} for mark in dict.fromkeys(str(r[0]) for r in vom.rows)]
    flat=[]
    for c in facts.get('climates',[]):
        for n in sorted({r['layer'] for key in ('layer_details','layers') for r in c.get(key,[])}):
            l=next((r for r in c.get('layer_details',[]) if r['layer']==n),{'layer':n})
            flat.append({'variant':c['variant'],**l,
                **next((r for r in c.get('layers',[]) if r['layer']==n),{}),
                'average':next((r['average'] for r in c.get('dry_averages',[]) if r['layer']==n),'')})
    datasets={'marks':marks,'uzk_details':facts.get('uzk_details',[]),'climates':facts.get('climates',[]),'layers':flat}
    for field,(name,columns) in SHEETS.items():
        sheet=wb.create_sheet(name);sheet.append(list(columns.values()))
        for row in datasets[field]:sheet.append([row.get(k,'') for k in columns])
    sheet=wb.create_sheet('Реквизиты');sheet.append(['Реквизит','Значение'])
    values=defaults.get('doc_fields',{})|facts.get('doc_fields',{})
    for k,label in DOC_FIELD_LABELS.items():sheet.append([label,values.get(k,'')])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='27465A')
        for column in sheet.columns:
            letter=column[0].column_letter;sheet.column_dimensions[letter].width=min(62,max(20,len(str(column[0].value or ''))+3))
            for cell in column:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):cell.data_type='s'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
    wb.save(path);wb.close();return str(path)
