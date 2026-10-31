"""Read-only reconciliation of issued Appendix #1 accounting and location.

A byte change is never accepted from a hash alone. The registry and the
previously confirmed quantities remain authoritative. Rebindings are staged
in the ledger candidate; coordinated commit retains the original ledger backup.
"""
from pathlib import Path
from datetime import date, datetime, timezone
from decimal import Decimal
from collections import Counter
from types import SimpleNamespace
import re
from openpyxl import load_workbook
from .storage import Stop, sha, safe


def appendix_accounting(path, entry):
    from passport_app.appendix import inspect_template, same_object
    from passport_app.tables import locate_layout
    from .catalog_ledger import quantity
    path=safe(path);before=sha(path);number=str(entry['number'])
    day=entry['date'].date() if isinstance(entry['date'],datetime) else entry['date']
    if isinstance(day,str):day=date.fromisoformat(day)
    label='История '+number+': '
    book=load_workbook(path,data_only=False)
    try:
        layout=locate_layout(book,appendix=True);sheet=book[layout.sheet]
        headers=[];objects=[]
        for cells in sheet.iter_rows(max_row=layout.header_row-1):
            for cell in cells:
                if not isinstance(cell.value,str):continue
                match=re.search(r'№\s*(\S+)\s+от\s+(\d{2}\.\d{2}\.\d{4})\s*г\.',cell.value)
                if match:headers.append(match.groups())
                if re.match(r'^\s*Объект\s*:',cell.value,re.I):objects.append(cell.value)
        if len(headers)!=1 or len(objects)!=1:raise Stop(label+'неоднозначные номер, дата или объект приложения.')
        differences=[]
        if headers[0][0]!=number:differences.append(f'номер {headers[0][0]!r}, в реестре {number!r}')
        if headers[0][1]!=day.strftime('%d.%m.%Y'):differences.append(f'дата {headers[0][1]}, в реестре {day:%d.%m.%Y}')
        if not same_object(objects[0],str(entry['object'])):differences.append(f'объект {objects[0]!r}, в реестре {entry["object"]!r}')
        if differences:raise Stop(label+'; '.join(differences))
        # Reuse the established structural guard after reporting exact identity differences.
        inspect_template(path,SimpleNamespace(first_number=number,first_date=day,object_name=str(entry['object'])))
        counts=Counter();mass=Decimal(0);mass_by_mark=Counter()
        for cells in sheet.iter_rows(min_row=layout.start,max_row=layout.total_row-1):
            mark=cells[layout.columns['mark']-1].value
            if not isinstance(mark,str) or not mark.strip():raise Stop(label+'пустая марка в строке '+str(cells[0].row))
            mark=mark.strip();counts[mark]+=quantity(cells[layout.columns['quantity']-1].value,label+mark)
            raw=cells[layout.columns['mass_total']-1].value
            try:
                value=Decimal(str(raw))
                if isinstance(raw,bool) or not value.is_finite() or value<=0:raise ValueError()
            except (ValueError,ArithmeticError):raise Stop(label+'не подтверждена числовая масса '+mark+' (формула без проверенного значения недопустима).')
            mass+=value;mass_by_mark[mark]+=value
        expected=Decimal(str(entry['mass']))
        if not expected.is_finite() or abs(mass-expected)>Decimal('.005'):
            raise Stop(label+f'масса приложения {mass} кг, в реестре {expected} кг.')
        from openpyxl.utils import get_column_letter
        for field,total in [('quantity',sum(counts.values())),('mass_total',mass)]:
            column=layout.columns[field];raw=sheet.cell(layout.total_row,column).value
            letter=get_column_letter(column)
            formula=f'=SUM({letter}{layout.start}:{letter}{layout.total_row-1})'
            if isinstance(raw,str) and raw.startswith('='):
                if raw.replace('$','').replace(' ','').upper()!=formula:
                    raise Stop(label+'изменена формула итога '+field+': '+raw)
            else:
                try:matches=not isinstance(raw,bool) and abs(Decimal(str(raw))-Decimal(total))<=Decimal('.005')
                except ArithmeticError:matches=False
                if not matches:raise Stop(label+'итог '+field+' не совпадает со строками приложения.')
    finally:book.close()
    if sha(path)!=before:raise Stop(label+'приложение изменилось во время проверки.')
    return {'quantities':dict(counts),'mass':str(mass),'mass_by_mark':{m:str(v.normalize()) for m,v in mass_by_mark.items()},'sha256':before,'source':str(path)}


def relocated_source(historical):
    """Only one exact basename under the same project's Passport directory."""
    old=Path(historical['source'])
    if old.is_file():return safe(old)
    roots=[p for p in old.parents if p.name.casefold()=='паспорт']
    if len(roots)!=1 or not roots[0].is_dir():
        raise Stop('История '+historical['number']+': отсутствует источник '+str(old))
    root=safe(roots[0]);matches=[]
    for folder in [root,*[p for p in root.iterdir() if p.is_dir()]]:
        folder=safe(folder)
        if not folder.is_relative_to(root):raise Stop('История: источник вне папки паспортов.')
        for path in folder.iterdir():
            if path.name.casefold()==old.name.casefold() and path.is_file():matches.append(safe(path))
    if len(matches)!=1:
        raise Stop('История '+historical['number']+': исходный путь отсутствует; найдено приложений с точным именем '+str(len(matches))+'. Автоматический выбор запрещён.')
    return matches[0]


def reconcile_history(data, registered):
    entries={str(e['number']):e for e in registered} if registered is not None else {}
    for historical in data['history']:
        if not historical.get('complete') or not historical.get('source'):continue
        path=relocated_source(historical);current=sha(path)
        moved=str(path)!=historical['source'];changed=current!=historical.get('sha256')
        if not moved and not changed:continue
        number=historical['number'];entry=entries.get(number)
        if entry is None:raise Stop('История '+number+': для обновления привязки необходима строка реестра с объектом, датой и массой.')
        checked=appendix_accounting(path,entry)
        old=historical['quantities'];new=checked['quantities']
        differences=[f'{m}: было {old.get(m,0)}, стало {new.get(m,0)}' for m in sorted(set(old)|set(new)) if old.get(m,0)!=new.get(m,0)]
        if differences:raise Stop('Изменена историческая отгрузка '+number+': '+'; '.join(differences))
        if historical.get('mass_by_mark') and historical['mass_by_mark']!=checked['mass_by_mark']:
            differences=[f'{m}: было {historical["mass_by_mark"].get(m)}, стало {checked["mass_by_mark"].get(m)} кг' for m in sorted(set(historical['mass_by_mark'])|set(checked['mass_by_mark'])) if historical['mass_by_mark'].get(m)!=checked['mass_by_mark'].get(m)]
            raise Stop('Изменена историческая масса '+number+': '+'; '.join(differences))
        reason='SOURCE_RELOCATED_SAME_BYTES' if moved and not changed else 'ACCOUNTING_IDENTICAL_BYTES_CHANGED'
        historical.setdefault('binding_updates',[]).append({'at':datetime.now(timezone.utc).isoformat(),'reason':reason,'previous_source':historical['source'],'previous_sha256':historical.get('sha256'),'source':str(path),'sha256':checked['sha256'],'registry_mass':str(entry['mass'])})
        historical.update(checked)
    return data
